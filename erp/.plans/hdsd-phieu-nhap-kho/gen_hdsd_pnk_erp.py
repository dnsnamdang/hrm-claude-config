# -*- coding: utf-8 -*-
"""Dung HDSD man 'Phieu nhap kho' (WarehouseImport / PNK) - ERP TanPhatDev.

Nguon: khao sat controller Warehouse\\WarehouseImportsController (2505 dong) + model
WarehouseImport (+ ImportModel TYPES + WarehouseImportLot/LotDetail + Tab/TabProduct
+ TmpWarehouseImportLot + ArrangeDelivery/Executor + ExcShortageWarehouseImport) +
routes warehouse_imports + views warehouse/warehouse_imports (index/all/forAccounting
/discrepancy/create/form/edit/show/arrange). KHONG chen anh -> placeholder.

Chay:  /usr/local/bin/python3 gen_hdsd_pnk_erp.py
Output: ERP/HDSD_luongchinh/HDSD_PhieuNhapKho.docx
"""
import os
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

ERP = "/Users/nguyentrancu/DEV/code/ERP-HRM/ERP"
sys.path.insert(0, os.path.join("/Users/nguyentrancu/DEV/code/ERP-HRM/HRM",
                                ".claude", "skills", "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUT_DIR = os.path.join(ERP, "HDSD_luongchinh")
os.makedirs(OUT_DIR, exist_ok=True)
OUTPUT = os.path.join(OUT_DIR, "HDSD_PhieuNhapKho.docx")

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)


def img_placeholder(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


def clear_toc_cache(path):
    """Xoa text cache trong moi field TOC (muc luc + danh muc hinh) -> Word tu dung lai khi mo."""
    doc = Document(path)
    body = doc.element.body
    fields = []
    for el in body.iter():
        tag = el.tag
        if tag == qn('w:fldChar'):
            t = el.get(qn('w:fldCharType'))
            if t == 'begin':
                fields.append({'instr': '', 'in_result': False})
            elif t == 'separate':
                if fields:
                    fields[-1]['in_result'] = True
            elif t == 'end':
                if fields:
                    fields.pop()
        elif tag == qn('w:instrText'):
            if fields:
                fields[-1]['instr'] += (el.text or '')
        elif tag == qn('w:t'):
            if any(f['in_result'] and 'TOC' in f['instr'] for f in fields):
                el.text = ''
    doc.save(path)


def finish_macos(b):
    b.doc.save(b.output)
    clear_toc_cache(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    my_headings = {p.text.strip() for p in chk.paragraphs
                   if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in my_headings]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables),
          "| Anh:", len(chk.inline_shapes), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


# ============================================================ NOI DUNG
b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Phiếu nhập kho)",
    doc_title="HDSD - Phiếu nhập kho",
)

# ------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Phiếu nhập kho (PNK)", "Chứng từ bước 3 (cuối) của quy trình nhập kho: do Thủ kho lập từ một Đề nghị nhập kho để nhận hàng thực tế, dự kiến vị trí xếp hàng rồi nhập vào kho — đây là bước LÀM TĂNG TỒN KHO thực tế. Đây chính là màn hình tài liệu này hướng dẫn. Mã phiếu có dạng PNK-xxxxx."],
    ["Đề nghị nhập kho (ĐNNK)", "Chứng từ bước 2 (WarehouseImportRequest). Mỗi phiếu nhập kho luôn được lập từ một đề nghị nhập kho; toàn bộ danh sách hàng hóa được kế thừa từ đề nghị này."],
    ["Yêu cầu nhập hàng (YCNH)", "Chứng từ bước 1 (ProductImportRequest). Là gốc của cả chuỗi nhập kho."],
    ["Phiếu nhập hàng", "Chứng từ hạch toán/giá vốn (ProductImport) — lập SAU phiếu nhập kho qua nút \"Tạo phiếu nhập hàng\". Phần Nợ/Có, mã phí, vụ việc, giá nhập, ngày hạch toán nằm ở đây, KHÔNG ở phiếu nhập kho."],
    ["Thủ kho", "Người phụ trách một hoặc nhiều kho; là người lập phiếu nhập kho, đi xếp hàng và gửi phiếu."],
    ["Kế toán kho", "Vai trò duyệt phiếu nhập kho (qua tạo phiếu nhập hàng) hoặc từ chối; có màn danh sách riêng."],
    ["Đi xếp hàng / Dự kiến vị trí", "Bước sau khi lưu nháp: gán từng hàng hóa vào vị trí kho cụ thể (Nhà/Dãy/Khoang) trước khi hoàn tất nhập."],
    ["Vị trí kho", "Ô/kệ chứa hàng trong kho (position). Mỗi hàng hóa có thể được chia vào nhiều vị trí, mỗi vị trí một số lượng."],
    ["Lô (số lô)", "Mã lô hàng (lot_number) hệ thống tự sinh khi hoàn tất nhập, dùng để truy vết tồn theo lô. Người dùng không tự nhập ở màn này."],
    ["Tồn kho (Stock)", "Số lượng hàng thực có trong kho; phiếu nhập kho khi hoàn tất sẽ cộng vào tồn kho (total_qty)."],
    ["Loại nhập (type)", "Phân loại mục đích nhập (mua trong/ngoài nước, mượn/bán trả lại, điều chuyển kho, tách/ghép, sản xuất theo hợp đồng…); kế thừa từ đề nghị nhập kho, người dùng không tự chọn."],
    ["Tiến trình (step)", "Cột riêng cho biết phiếu đang ở giai đoạn nào: Đang đi xếp hàng → Đang nhập kho → Hoàn thành nhập kho."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "11/09/2026", "Khởi tạo HDSD màn Phiếu nhập kho (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Màn hình Phiếu nhập kho là bước cuối của quy trình nhập kho: sau khi có Đề nghị nhập kho, Thủ kho lập phiếu nhập kho để nhận hàng thực tế. Thủ kho chọn đề nghị nhập kho nguồn, khai người giao hàng, dự kiến vị trí xếp hàng cho từng mặt hàng, (nếu có) khai vận chuyển và bốc xếp, rồi Gửi để hoàn tất — đây chính là thời điểm hệ thống TĂNG TỒN KHO thực tế. Sau đó Kế toán kho tạo Phiếu nhập hàng để hạch toán. Toàn bộ danh sách hàng hóa được kế thừa từ đề nghị nhập kho — người lập KHÔNG tự thêm/xóa mặt hàng; chỉ phân bổ số lượng vào các vị trí kho.")
b.bullet("menu Kho › Nhập kho › \"Phiếu nhập kho\". Có nhiều màn danh sách theo ngữ cảnh: phiếu theo quyền của tôi, tất cả (theo quyền), màn dành cho Kế toán kho, và màn Chờ báo hàng thừa/thiếu.", bold_prefix="Đường dẫn danh sách: ")
b.bullet("phiếu nhập kho được lập từ một Đề nghị nhập kho đã gửi (chờ duyệt): bấm \"Tạo mới\" rồi chọn đề nghị nhập kho nguồn. Có thể vào thẳng từ nút \"Tạo phiếu nhập kho\" trên màn chi tiết đề nghị nhập kho.", bold_prefix="Điểm tạo phiếu: ")
b.para("Vị trí trong quy trình nhập kho (chuỗi chứng từ nối tiếp):")
b.table([
    ["Bước", "Chứng từ", "Người thực hiện", "Ghi chú"],
    ["1", "Yêu cầu nhập hàng", "Người yêu cầu / kế toán", "Đề nghị nhập hàng, đã được duyệt."],
    ["2", "Đề nghị nhập kho", "Kế toán kho", "Chọn kho nhập, gửi Thủ kho."],
    ["3", "Phiếu nhập kho", "Thủ kho", "Màn hình này — nhận hàng, dự kiến vị trí, nhập kho thực tế (TĂNG TỒN)."],
    ["4", "Phiếu nhập hàng", "Kế toán kho", "Hạch toán / giá vốn (chứng từ riêng, tạo từ phiếu nhập kho)."],
])

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Phạm vi dữ liệu nhìn thấy và các thao tác đều phụ thuộc quyền. Màn này hầu như KHÔNG chặn truy cập qua đường dẫn (route) — chỉ màn Kế toán kho gắn quyền \"Kế toán kho\"; phần còn lại kiểm soát trong xử lý nghiệp vụ và ẩn/hiện nút theo trạng thái. Các nhóm quyền chính:")
b.bullet("3 quyền xem theo cấp (tổng công ty / công ty / phòng ban) và quyền \"Xem tất cả phiếu xuất nhập kho\" (lọc theo kho mình phụ trách) quyết định thấy phiếu của phạm vi nào ở màn \"Tất cả\". Không có quyền nào thì chỉ thấy phiếu do chính mình lập. Lưu ý: phiếu nhập kho KHÔNG có mức xem \"theo bộ phận\".", bold_prefix="Xem danh sách: ")
b.bullet("người phụ trách kho — lập phiếu nhập kho, đi xếp hàng, gửi phiếu và báo hàng thừa/thiếu.", bold_prefix="Thủ kho: ")
b.bullet("vai trò duyệt phiếu (tạo phiếu nhập hàng) / từ chối; có màn danh sách \"Kế toán kho\" riêng.", bold_prefix="Kế toán kho: ")
b.bullet("cho phép tạo / duyệt phiếu Báo hàng thừa, thiếu nhập kho.", bold_prefix="Thừa/thiếu: ")
b.para("Các thao tác Sửa/Hủy phụ thuộc quyền sở hữu + trạng thái: chỉ người lập mới sửa/hủy phiếu ở trạng thái Đang tạo. Duyệt (Tạo phiếu nhập hàng) / Không duyệt chỉ dành cho Kế toán kho cùng công ty khi phiếu ở trạng thái Chờ duyệt. Chi tiết ở PHẦN 2.")
b.para("Cảnh báo quan trọng: khi Thủ kho bấm \"Lưu & Nhập\" (gửi phiếu), hệ thống CỘNG số lượng vào tồn kho thực tế và sinh số lô — hãy kiểm tra kỹ vị trí và số lượng trước khi gửi. Với các loại điều chuyển kho, thao tác này còn cập nhật phiếu điều chuyển liên quan.")

# ------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")

b.h2("1.1. Các màn danh sách")
b.para("Cùng một tiêu đề \"Danh sách phiếu nhập kho\", hệ thống có nhiều màn danh sách phục vụ ngữ cảnh khác nhau (khác nhau ở phạm vi dữ liệu và một vài nút/cột):")
b.table([
    ["Màn", "Dùng cho", "Nội dung / khác biệt"],
    ["Danh sách (index)", "Người lập / theo quyền", "Hiển thị phiếu theo phạm vi phân quyền của user. Có nút Tạo mới và Xuất Excel; có cột Tiến trình."],
    ["Tất cả (all)", "Quản lý / người có quyền xem", "Hiển thị phiếu theo phạm vi phân quyền (tổng công ty / công ty / phòng ban / theo kho). Có nút Tạo mới, Xuất Excel; có cột Tiến trình. Ẩn phiếu nháp của người khác."],
    ["Kế toán kho (forAccounting)", "Kế toán kho", "Chỉ hiển thị phiếu đã gửi (khác nháp) trong công ty mình (route gắn quyền \"Kế toán kho\"). BỎ cột Tiến trình và KHÔNG có nút Xuất Excel. Có nút Tạo phiếu nhập hàng khi đủ điều kiện."],
    ["Chờ báo hàng thừa/thiếu (discrepancy)", "Thủ kho / người xử lý thừa thiếu", "Danh sách phiếu đã hoàn tất cần đối chiếu thừa/thiếu. Mỗi dòng chỉ có nút \"Báo hàng thừa thiếu\"."],
])
img_placeholder(b, "Ảnh menu Phiếu nhập kho.")

b.h2("1.2. Bố cục màn hình danh sách")
b.bullet("tiêu đề \"Danh sách phiếu nhập kho\".", bold_prefix="Thanh tiêu đề: ")
b.bullet("các ô lọc theo cột (Kho hàng, Tên/mã hàng hóa, Mã phiếu, Loại, Nhà cung cấp, Phiếu ĐNNK, Người đề nghị, Người yêu cầu, Người lập, Trạng thái, Người duyệt) và lọc theo khoảng thời gian.", bold_prefix="Khu vực lọc: ")
b.bullet("bảng liệt kê phiếu theo phạm vi quyền; hai cột riêng biệt Trạng thái và Tiến trình; cột Hành động ở cuối mỗi dòng.", bold_prefix="Bảng danh sách: ")
b.bullet("nút Tạo mới (+), nút Xuất Excel (màn Danh sách / Tất cả), chọn số dòng/trang và chuyển trang.", bold_prefix="Công cụ & phân trang: ")
img_placeholder(b, "Ảnh tổng quan màn danh sách phiếu nhập kho.")

# ------------------------------------------------- PHAN 2
b.h1("PHẦN 2: DANH SÁCH PHIẾU NHẬP KHO")
img_placeholder(b, "Ảnh màn danh sách với dữ liệu mẫu và cột Hành động.")

b.h2("2.1. Phân quyền & hướng dẫn theo quyền")
b.para("Danh sách hiển thị theo phạm vi quyền của người đăng nhập. Bảng dưới liệt kê đầy đủ các quyền liên quan tới màn. Người dùng thường chỉ có một vài quyền — hãy đối chiếu đúng phần áp dụng cho mình.")
b.table([
    ["Tên quyền", "Cho phép làm gì", "Nút/màn tương ứng"],
    ["Xem phiếu nhập kho theo tổng công ty", "Thấy toàn bộ phiếu (trạng thái khác nháp) của mọi công ty ở màn Tất cả.", "Danh sách / Tất cả"],
    ["Xem phiếu nhập kho theo công ty", "Thấy phiếu trong công ty của mình.", "Danh sách / Tất cả"],
    ["Xem phiếu nhập kho theo phòng ban", "Thấy phiếu của các phòng ban mình quản lý (hoặc do mình lập).", "Danh sách / Tất cả"],
    ["Xem tất cả phiếu xuất nhập kho", "Thấy phiếu thuộc các kho mình được gán làm thủ kho (hoặc do mình lập).", "Danh sách / Tất cả"],
    ["Thủ kho", "Lập phiếu nhập kho, đi xếp hàng, gửi phiếu; báo hàng thừa/thiếu.", "Nút Tạo mới / Phiếu xếp hàng / Lưu & Nhập / Báo hàng thừa thiếu"],
    ["Kế toán kho", "Duyệt (Tạo phiếu nhập hàng) / Không duyệt phiếu; có màn danh sách riêng.", "Màn \"Kế toán kho\"; nút Tạo phiếu nhập hàng / Không duyệt"],
    ["Tạo phiếu báo hàng thừa, thiếu nhập kho", "Tạo phiếu báo chênh lệch thừa/thiếu khi nhập.", "Nút Báo hàng thừa thiếu"],
    ["Duyệt phiếu báo hàng thừa, thiếu nhập kho", "Duyệt phiếu báo thừa/thiếu.", "Màn duyệt thừa/thiếu"],
])
b.para("Lưu ý phạm vi: ở màn Tất cả, hệ thống xét quyền xem theo thứ tự tổng công ty → công ty → phòng ban → theo kho phụ trách; nếu không có quyền nào, người dùng chỉ thấy phiếu do chính mình lập. Phiếu nháp (Đang tạo) của người khác luôn bị ẩn — chỉ người lập mới thấy phiếu nháp của mình. Màn \"Kế toán kho\" lọc theo công ty người dùng và chỉ hiện phiếu đã gửi (khác nháp).")

b.h3("Người dùng có quyền \"Xem phiếu nhập kho theo tổng công ty\"")
b.para("Ở màn Tất cả, thấy toàn bộ phiếu (trạng thái khác nháp) của mọi công ty, cộng phiếu do chính mình lập. Thực hiện đầy đủ thao tác xem/lọc/in/xuất Excel; sửa/hủy vẫn theo quy tắc sở hữu + trạng thái.")

b.h3("Người dùng có quyền \"Xem phiếu nhập kho theo công ty / phòng ban\"")
b.para("Thấy phiếu trong phạm vi công ty / phòng ban tương ứng mà mình phụ trách, cộng thêm phiếu do chính mình lập. Cấp phòng ban dựa trên các phòng người dùng quản lý.")

b.h3("Người dùng có quyền \"Xem tất cả phiếu xuất nhập kho\"")
b.para("Thấy phiếu thuộc các kho mà người dùng được gán làm thủ kho, cộng phiếu do chính mình lập. Đây là góc nhìn theo kho phụ trách chứ không theo cấp tổ chức.")

b.h3("Người dùng là Kế toán kho")
b.para("Vào màn \"Kế toán kho\" để xem danh sách phiếu đã gửi trong công ty mình (chờ duyệt/đã xử lý). Với phiếu ở trạng thái Chờ duyệt, Kế toán kho được: Không duyệt (kèm lý do) hoặc Tạo phiếu nhập hàng để hạch toán.")

b.h3("Người dùng là Thủ kho")
b.para("Vai trò lập phiếu: bấm Tạo mới, chọn đề nghị nhập kho nguồn, khai người giao hàng, dự kiến vị trí, rồi Lưu (nháp) / Lưu & Nhập (gửi). Sau khi hoàn tất còn có thể vào màn Chờ báo hàng thừa/thiếu để đối chiếu.")

b.h2("2.2. Tìm kiếm & lọc")
b.table([
    ["Tiêu chí", "Cách dùng"],
    ["Kho hàng", "Chọn kho để lọc phiếu theo kho nhập (chỉ các kho người dùng được phép)."],
    ["Tên, mã hàng hóa", "Nhập để lọc phiếu có chứa hàng theo tên hoặc mã."],
    ["Mã phiếu", "Nhập mã phiếu nhập kho (PNK-…)."],
    ["Loại", "Chọn loại nhập."],
    ["Nhà cung cấp", "Chọn nhà cung cấp."],
    ["Phiếu ĐNNK", "Nhập mã đề nghị nhập kho nguồn."],
    ["Người đề nghị", "Chọn (tìm theo tên) người lập đề nghị/yêu cầu."],
    ["Người yêu cầu", "Chọn (tìm theo tên) người yêu cầu nhập hàng."],
    ["Người lập", "Chọn (tìm theo tên) người lập phiếu nhập kho."],
    ["Trạng thái", "Chọn trạng thái phiếu."],
    ["Người duyệt", "Chọn (tìm theo tên) người duyệt."],
    ["Khoảng thời gian", "Lọc theo khoảng ngày lập phiếu."],
])

b.h2("2.3. Các cột trong danh sách")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Mã phiếu", "Mã phiếu nhập kho; bấm để mở chi tiết."],
    ["Loại", "Loại nhập (kế thừa từ đề nghị nhập kho)."],
    ["Nhà cung cấp", "Nhà cung cấp của phiếu (nếu có)."],
    ["Ngày lập", "Ngày tạo phiếu."],
    ["Phiếu ĐNNK", "Mã đề nghị nhập kho nguồn."],
    ["Người đề nghị", "Người lập đề nghị nhập kho."],
    ["Người yêu cầu", "Người yêu cầu nhập hàng."],
    ["Người lập", "Người lập phiếu nhập kho (chỉ ở màn Tất cả)."],
    ["Trạng thái", "Trạng thái phiếu, hiển thị nhãn màu (xem 2.4)."],
    ["Tiến trình", "Giai đoạn xếp hàng/nhập kho (xem 2.5). BỎ ở màn Kế toán kho."],
    ["Người duyệt", "Người duyệt phiếu."],
    ["Ngày nhận", "Thời điểm nhận hàng (khi gửi)."],
    ["Ngày duyệt", "Thời điểm duyệt."],
    ["Hành động", "Các nút thao tác (xem mục 2.6)."],
])

b.h2("2.4. Bảng trạng thái phiếu")
b.para("Trong bộ lọc và nhãn hiển thị, cột Trạng thái gồm các giá trị sau:")
b.table([
    ["Trạng thái", "Màu nhãn", "Ý nghĩa"],
    ["Đang tạo", "Đỏ", "Phiếu nháp — mới Lưu, chưa gửi. Chỉ người lập thấy và sửa/hủy được."],
    ["Chờ duyệt", "Đỏ", "Đã Gửi (Lưu & Nhập), đã tăng tồn kho, chờ Kế toán kho tạo phiếu nhập hàng / duyệt."],
    ["Không duyệt", "Đỏ", "Bị Kế toán kho từ chối."],
    ["Đã xuất lại", "Đỏ", "Đã tạo phiếu xuất lại cho hàng đã nhập nhầm/trả lại."],
    ["Đã nhập lại", "Xanh nhạt", "Đã tạo phiếu nhập lại từ phiếu gốc."],
    ["Đang hạch toán", "Xanh", "Đã được duyệt (đã tạo phiếu nhập hàng), đang hạch toán."],
    ["Đã hạch toán", "Xanh", "Đã hạch toán xong."],
    ["Đã hủy", "Đỏ", "Phiếu bị hủy."],
])
b.para("Vòng đời cơ bản người dùng thấy khi thao tác: Đang tạo (Lưu nháp) → đi xếp hàng (dự kiến vị trí) → Chờ duyệt (Lưu & Nhập — tăng tồn kho) → Đang hạch toán (Kế toán tạo phiếu nhập hàng) → Đã hạch toán. Các nhánh phụ: Không duyệt (bị từ chối), Đã xuất lại / Đã nhập lại (xử lý nhập nhầm).")

b.h2("2.5. Cột Tiến trình (giai đoạn xếp hàng/nhập kho)")
b.para("Ngoài cột Trạng thái, danh sách còn cột Tiến trình cho biết phiếu đang ở giai đoạn nào của việc xếp/nhập hàng:")
b.table([
    ["Tiến trình", "Màu nhãn", "Ý nghĩa"],
    ["Đang đi xếp hàng", "Xanh nhạt", "Phiếu vừa lập, đang ở bước dự kiến vị trí (chưa nhập kho thực tế)."],
    ["Đang nhập kho", "Vàng", "Đã chuyển sang bước nhập kho nhưng chưa gửi hoàn tất."],
    ["Hoàn thành nhập kho", "Xanh", "Đã gửi và hoàn tất nhập kho (đã tăng tồn)."],
])
b.para("Cột Tiến trình được ẩn ở màn \"Kế toán kho\".")

b.h2("2.6. Thao tác trên từng dòng (cột Hành động)")
b.para("Các nút nằm trong menu thao tác và hiển thị tùy trạng thái, loại phiếu và quyền:")
b.table([
    ["Nút", "Điều kiện hiển thị", "Kết quả"],
    ["Phiếu xếp hàng", "Phiếu Đang tạo và đang ở bước đi xếp hàng.", "Mở màn dự kiến vị trí / in phiếu xếp hàng."],
    ["Sửa / Nhập kho", "Phiếu Đang tạo do bạn lập (hoặc Super Admin), ở bước xếp/nhập.", "Mở màn sửa để hoàn thiện và nhập kho."],
    ["Tạo phiếu nhập hàng", "Bạn là Kế toán kho cùng công ty và phiếu ở Chờ duyệt (loại không phải điều chuyển cùng thủ kho).", "Chuyển sang lập Phiếu nhập hàng để hạch toán."],
    ["Tạo phiếu xuất lại", "Bạn là người lập và phiếu ở trạng thái Không duyệt.", "Lập phiếu xuất lại cho hàng đã nhập nhầm."],
    ["Tạo phiếu nhập lại", "Bạn là người lập và phiếu ở trạng thái Đã xuất lại.", "Lập phiếu nhập lại."],
    ["Hủy phiếu", "Phiếu Đang tạo do bạn lập, loại KHÔNG thuộc nhóm điều chuyển cùng thủ kho / nội bộ / chi nhánh / nhập tách.", "Hủy phiếu (chuyển sang Đã hủy) và hủy luôn đề nghị nhập kho nguồn."],
    ["Báo hàng thừa thiếu", "Phiếu Đã hạch toán, bạn có quyền tạo báo thừa thiếu hoặc là Thủ kho, và chưa có phiếu báo.", "Lập phiếu Báo hàng thừa, thiếu (xem PHẦN 7)."],
    ["In phiếu", "Luôn có (theo quyền xem).", "In mẫu Phiếu nhập kho."],
    ["In barcode", "Theo quyền xem.", "In mã vạch hàng hóa của phiếu."],
])
b.para("Ngoài ra danh sách (Danh sách / Tất cả) có nút Tạo mới (+) mở màn lập phiếu, và nút Xuất Excel.")

# ------------------------------------------------- PHAN 3
b.h1("PHẦN 3: LẬP PHIẾU NHẬP KHO & DỰ KIẾN VỊ TRÍ")
b.para("Phiếu nhập kho do Thủ kho lập từ một Đề nghị nhập kho đã gửi (chờ duyệt). Bấm Tạo mới ở màn danh sách (hoặc vào thẳng từ nút \"Tạo phiếu nhập kho\" trên chi tiết đề nghị nhập kho). Màn tạo mới có tiêu đề \"Dự kiến vị trí\". Loại nhập được suy ra từ đề nghị nhập kho, người dùng KHÔNG tự chọn.")
img_placeholder(b, "Ảnh màn lập phiếu nhập kho (Dự kiến vị trí).")

b.h2("3.1. Bố cục form")
b.para("Form gồm 2 thẻ (card):")
b.bullet("gồm các tab: \"Thông tin chung\" (bắt buộc), \"Vận chuyển\" (chỉ hiện với loại phát sinh vận chuyển), \"Bốc xếp\". Xem PHẦN 4 cho hai tab Vận chuyển & Bốc xếp.", bold_prefix="Thẻ Thông tin chung: ")
b.bullet("tab \"Hàng hóa\" (bảng dự kiến vị trí) và các tab theo hợp đồng/nhóm (nếu có).", bold_prefix="Thẻ Chi tiết: ")
b.para("Header phiếu hiển thị tên loại nhập và \"Người lập - Ngày lập\".")

b.h2("3.2. Nguồn tạo phiếu & cách chọn Đề nghị nhập kho")
b.para("Mỗi phiếu nhập kho luôn gắn với một Đề nghị nhập kho nguồn. Tại trường \"Chọn phiếu đề nghị nhập kho\", bấm nút kính lúp để mở cửa sổ \"Phiếu đề nghị nhập kho\":")
b.table([
    ["Thành phần", "Nội dung"],
    ["Cột danh sách", "STT, Mã phiếu, Người tạo, Ngày tạo."],
    ["Bộ lọc", "Mã phiếu (nhập), Người tạo (chọn)."],
    ["Điều kiện hiển thị", "Chỉ liệt kê đề nghị nhập kho đã gửi (chờ duyệt) đủ điều kiện lập phiếu."],
    ["Chọn phiếu", "Bấm chọn 1 dòng → hệ thống nạp thông tin + danh sách hàng hóa của đề nghị, rồi đóng cửa sổ."],
])
b.para("Khi phiếu đã lưu (đã có mã), ô chọn phiếu nguồn bị khóa và nhãn đổi thành \"Phiếu đề nghị nhập kho\". Nếu chưa chọn phiếu nguồn, bảng hàng hóa hiển thị dòng nhắc \"Chưa chọn phiếu đề nghị\".")

b.h2("3.3. Các trường trong tab \"Thông tin chung\"")
b.table([
    ["Trường", "Bắt buộc", "Giá trị điền sẵn / ghi chú"],
    ["Chọn phiếu đề nghị nhập kho", "Có (khi tạo mới)", "Hiển thị mã đề nghị; khóa khi phiếu đã lưu."],
    ["Người giao hàng", "Có", "Điền tên người giao hàng; ở màn sửa/xem bị khóa."],
    ["Kho nhập", "—", "Chỉ đọc — tự điền từ đề nghị nhập kho."],
    ["Người đề nghị", "—", "Chỉ đọc — người lập đề nghị nhập kho nguồn."],
    ["Ghi chú", "Không", "Tối đa 255 ký tự; mặc định lấy ghi chú của đề nghị nguồn."],
    ["File đính kèm", "Có (ở màn sửa/nhập)", "Tải lên file (pdf, jpg, jpeg, png, doc, docx, xls, xlsx, csv). Bắt buộc trước khi Lưu & Nhập."],
])
b.para("Màn này KHÔNG có trường Nhà cung cấp riêng, KHÔNG có chọn Lô, KHÔNG có Ngày nhập kế toán / Giá nhập / phân bổ kho kế toán / hạch toán — các phần đó thuộc Phiếu nhập hàng. Tại đây chỉ dự kiến VỊ TRÍ và số lượng nhập.")

b.h2("3.4. Bảng hàng hóa & dự kiến vị trí (tab \"Hàng hóa\")")
b.para("Bảng hàng hóa kế thừa nguyên trạng từ đề nghị nhập kho: người lập KHÔNG thêm/xóa mặt hàng. Việc cần làm là gán mỗi mặt hàng vào một hoặc nhiều vị trí kho và nhập số lượng cho từng vị trí.")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Tên hàng hóa", "Tên sản phẩm (chỉ đọc)."],
    ["Model", "Model sản phẩm (chỉ đọc)."],
    ["Mã hàng hóa", "Mã sản phẩm (chỉ đọc)."],
    ["Thương hiệu", "Thương hiệu (chỉ đọc)."],
    ["Đơn vị tính", "Đơn vị tính (chỉ đọc)."],
    ["SL đề nghị", "Số lượng đề nghị nhập (lấy từ đề nghị nhập kho, chỉ hiển thị)."],
    ["SL thực nhập", "Tổng tự tính từ số lượng các dòng vị trí bên nhóm Chi tiết (không nhập trực tiếp)."],
    ["Vị trí (nhóm Chi tiết)", "Bấm để mở cửa sổ \"Chọn vị trí\"; mỗi mặt hàng có thể chia nhiều dòng vị trí."],
    ["Tồn (nhóm Chi tiết)", "Tồn hiện tại của vị trí đã chọn (chỉ đọc)."],
    ["SL nhập (nhóm Chi tiết)", "Số lượng nhập vào vị trí đó (nhập được)."],
])
b.bullet("khi tạo mới có tùy chọn \"Áp dụng tất cả\" để gán cùng một vị trí cho mọi mặt hàng, tiết kiệm thao tác.", bold_prefix="Áp dụng tất cả: ")
b.bullet("hệ thống tự đảm bảo tổng SL nhập các vị trí khớp với SL đề nghị; nếu còn thiếu sẽ tự thêm dòng cho phần còn lại, nếu vượt sẽ cắt bớt.", bold_prefix="Kiểm tra số lượng: ")
b.bullet("ở màn sửa, mỗi dòng vị trí có nút xóa và cuối nhóm có link \"+ Thêm\" để thêm dòng vị trí.", bold_prefix="Thêm/xóa dòng vị trí: ")
b.para("Với hàng theo hợp đồng, ngoài tab \"Hàng hóa\" mặc định còn có các tab theo hợp đồng/nhóm (tên tab lấy từ phiếu nguồn); ở màn sửa các tab này cho nhập SL thực nhập.")
img_placeholder(b, "Ảnh bảng hàng hóa và dự kiến vị trí trong form phiếu nhập kho.")

b.h2("3.5. Cửa sổ \"Chọn vị trí\"")
b.para("Bấm nút chọn vị trí ở cột Vị trí để mở cửa sổ \"Chọn vị trí\":")
b.table([
    ["Thành phần", "Nội dung"],
    ["Cột danh sách", "STT, Vị trí, Nhà, Dãy, Khoang, Tồn."],
    ["Bộ lọc", "Mã vị trí, Nhà, Dãy, Khoang."],
    ["Chọn vị trí", "Bấm chọn 1 vị trí → gán vào dòng; nếu vị trí đã chọn, hệ thống báo \"Vị trí đã được chọn\"."],
])

b.h2("3.6. Các nút lưu")
b.table([
    ["Nút", "Màn", "Hành động"],
    ["Lưu", "Tạo mới", "Lưu phiếu ở trạng thái Đang tạo (nháp) rồi chuyển sang màn dự kiến vị trí (/arrange). Hệ thống nhắc \"Phiếu đã được lưu. Hãy tạo phiếu đi xếp hàng\"."],
    ["Lưu", "Sửa", "Giữ phiếu ở trạng thái Đang tạo (chưa gửi). Hệ thống nhắc cần gửi để phiếu được xử lý."],
    ["Lưu & Nhập", "Sửa", "Gửi phiếu (trạng thái Chờ duyệt) — đây là bước TĂNG TỒN KHO thực tế và sinh số lô. Hệ thống báo \"Phiếu đã được gửi\"."],
    ["In dự kiến", "Sửa", "In bản dự kiến vị trí (xem trước)."],
    ["Hủy", "Cả hai", "Quay lại danh sách, không lưu."],
])
b.para("Lưu ý: nếu loại phiếu có vận chuyển bằng công ty và còn chuyến xe chưa duyệt, khi bấm Lưu & Nhập hệ thống sẽ giữ phiếu ở Đang tạo và nhắc \"Phiếu đã được lưu. Bạn cần duyệt chuyến xe trước khi gửi\".")

b.h2("3.7. Quy tắc kiểm tra dữ liệu (validate)")
b.table([
    ["Trường / tình huống", "Yêu cầu / thông báo"],
    ["Phiếu đề nghị nhập kho", "Bắt buộc chọn (\"Bắt buộc phải chọn\")."],
    ["Người giao hàng", "Bắt buộc; tối đa 255 ký tự (\"Bắt buộc phải nhập\" / \"Không được vượt quá 255 ký tự\")."],
    ["Ghi chú", "Tối đa 255 ký tự (\"Không được vượt quá 255 ký tự\")."],
    ["Vị trí / SL nhập", "Mỗi mặt hàng phải có ít nhất một dòng vị trí; vị trí bắt buộc; SL nhập phải nhập (0–999999)."],
    ["File đính kèm", "Bắt buộc trước khi Lưu & Nhập (đúng định dạng cho phép)."],
    ["Điều chuyển kho chi nhánh", "Số lượng thực nhập phải bằng số lượng đề nghị (\"Số lượng thực nhập phải bằng số lượng đề nghị\")."],
    ["Vị trí trùng", "Không được chọn trùng vị trí cho cùng mặt hàng (\"Hàng hóa có vị trị trùng nhau\")."],
    ["Đang kiểm kho", "Nếu hàng/vị trí đang kiểm kho: \"Hàng … đang kiểm kho!\" / \"Vị trí … đang kiểm kho\"."],
    ["Quyền lập phiếu", "Đề nghị nguồn phải đủ điều kiện, nếu không: \"Không đủ quyền!\"."],
])
b.para("Thiếu dữ liệu bắt buộc sẽ báo lỗi ngay tại ô tương ứng (viền đỏ) kèm tooltip \"Vui lòng kiểm tra lại mục này\" ở tab lỗi, và không lưu cho tới khi sửa đủ.")

# ------------------------------------------------- PHAN 4
b.h1("PHẦN 4: VẬN CHUYỂN & BỐC XẾP")
b.para("Hai tab này chỉ áp dụng cho các loại nhập có phát sinh vận chuyển/bốc xếp. Tab \"Vận chuyển\" chỉ hiện khi loại phiếu có vận chuyển.")

b.h2("4.1. Tab Vận chuyển")
b.para("Chọn đối tượng vận chuyển ở trường \"Vận chuyển\":")
b.table([
    ["Đối tượng vận chuyển", "Ý nghĩa / trường phát sinh"],
    ["Nhân viên vận chuyển", "Nhân viên tự vận chuyển."],
    ["Công ty vận chuyển", "Hiện thêm: Tuyến đường (bắt buộc), Số Km dự kiến (khóa), Số Km chốt (bắt buộc), Ghi chú vận chuyển (bắt buộc khi Km chốt khác Km dự kiến) và khối \"Chọn chuyến xe\" (thêm/xóa dòng, có nút tạo chuyến xe)."],
    ["Khách hàng vận chuyển", "Khách hàng tự vận chuyển."],
])
b.para("Khi chọn chuyến xe: không được chọn trùng chuyến (\"Không được chọn trùng\"); chuyến xe không hợp lệ báo \"Chuyến xe không hợp lệ\". Nếu còn chuyến xe chưa duyệt thì chưa gửi được phiếu (xem 3.6).")
img_placeholder(b, "Ảnh tab Vận chuyển (Công ty vận chuyển) với khối chọn chuyến xe.")

b.h2("4.2. Tab Bốc xếp")
b.para("Khai chi phí bốc xếp (nếu có). Chọn cách tính rồi loại hình thuê:")
b.table([
    ["Trường", "Ghi chú"],
    ["Cách tính", "Chọn \"Tính theo khối lượng\" hoặc \"Tính theo thời gian\" (loại trừ nhau)."],
    ["Loại hình thuê", "Công ty (nội bộ) hoặc Thuê ngoài."],
    ["Nhà cung cấp", "Bắt buộc khi Thuê ngoài; khóa khi thuê Công ty."],
    ["Loại bốc xếp", "Chọn loại bốc xếp."],
    ["Khối lượng (Tấn) / Số giờ", "Nhãn động theo cách tính; bắt buộc, khác 0."],
    ["Đơn giá / Số tiền / VAT / Số tiền sau VAT", "Tự tính, chỉ đọc."],
    ["Bảng người nhận việc", "Khi thuê Công ty: thêm nhân viên và chia số tiền; tổng tiền chia phải bằng tổng tiền bốc xếp (\"Tổng số tiền chia cho các nhân viên phải bằng tổng số tiền bốc xếp\")."],
])
b.para("Thêm nhân viên nhận việc: mỗi nhân viên một số tiền (bắt buộc, khác 0); trùng nhân viên báo \"Nhân viên đã tồn tại!\", thêm thành công báo \"Thêm nhân viên thành công!\".")

# ------------------------------------------------- PHAN 5
b.h1("PHẦN 5: XEM CHI TIẾT & XỬ LÝ CỦA KẾ TOÁN KHO")
b.para("Bấm Mã phiếu để xem chi tiết. Trang chi tiết chỉ đọc, gồm các thẻ: Thông tin chung (tabs Thông tin chung / Vận chuyển / Bốc xếp), Chi tiết (bảng hàng hóa + vị trí + tab hợp đồng, có dòng Tổng cộng) và Ghi chú duyệt (chỉ hiện khi phiếu có ghi chú/bị từ chối).")
b.bullet("Phiếu đề nghị nhập kho (link mở phiếu gốc), Người đề nghị, Kho nhập, Người giao hàng, Ghi chú; nếu ứng với phiếu điều chuyển kho sẽ có link phiếu điều chuyển — đều chỉ đọc.", bold_prefix="Thông tin chung: ")
b.bullet("bảng hàng hóa (SL đề nghị / SL thực nhập, vị trí, SL nhập) và các tab hợp đồng, chỉ để xem; KHÔNG có tab hạch toán/giá vốn/lô.", bold_prefix="Chi tiết: ")
b.bullet("hiển thị lý do khi Kế toán kho không duyệt (nếu có).", bold_prefix="Ghi chú duyệt: ")
b.para("Các nút cuối màn chi tiết phụ thuộc quyền và trạng thái:")
b.table([
    ["Nút", "Điều kiện", "Kết quả"],
    ["Không duyệt", "Bạn là Kế toán kho cùng công ty, phiếu Chờ duyệt (loại không phải điều chuyển cùng thủ kho / nội bộ / nhập tách / đã hủy).", "Mở cửa sổ \"Ghi chú duyệt\" để nhập lý do rồi từ chối."],
    ["Tạo phiếu nhập hàng", "Bạn là Kế toán kho cùng công ty và phiếu Chờ duyệt.", "Chuyển sang lập Phiếu nhập hàng để hạch toán (với HĐ HRM loại 20/21 sẽ mở sang giao diện HRM)."],
    ["Sửa", "Phiếu Đang tạo do bạn lập (hoặc Super Admin).", "Mở màn sửa phiếu."],
    ["Hủy", "Phiếu Đang tạo do bạn lập, loại cho phép hủy.", "Hủy phiếu (và hủy đề nghị nguồn)."],
    ["Quay lại", "Luôn có.", "Về màn danh sách."],
])
img_placeholder(b, "Ảnh màn xem chi tiết phiếu nhập kho với các nút xử lý.")

b.h2("5.1. Không duyệt (từ chối) — dành cho Kế toán kho")
b.para("Khi Kế toán kho bấm \"Không duyệt\", hệ thống mở cửa sổ \"Ghi chú duyệt\":")
b.table([
    ["Trường", "Bắt buộc", "Ghi chú"],
    ["Ghi chú (lý do)", "Có", "Nhập lý do từ chối; tối đa 255 ký tự (\"Bắt buộc phải nhập\" / \"Không được vượt quá 255 ký tự\")."],
])
b.para("Bấm Xác nhận: phiếu chuyển sang trạng thái Không duyệt; người lập nhận thông báo. Bấm Hủy để đóng cửa sổ, không làm gì. Yêu cầu quyền Kế toán kho cùng công ty; nếu không đủ quyền: \"Không đủ quyền!\".")

b.h2("5.2. Tạo phiếu nhập hàng (hạch toán) — dành cho Kế toán kho")
b.para("Với phiếu ở trạng thái Chờ duyệt, Kế toán kho bấm \"Tạo phiếu nhập hàng\" để chuyển sang lập Phiếu nhập hàng — nơi thực hiện hạch toán, giá vốn, mã phí, vụ việc. Sau khi tạo, phiếu nhập kho chuyển sang Đang hạch toán rồi Đã hạch toán. Hạch toán KHÔNG thực hiện trên màn phiếu nhập kho.")

# ------------------------------------------------- PHAN 6
b.h1("PHẦN 6: SỬA, HỦY, XUẤT LẠI & NHẬP LẠI")
b.h2("6.1. Sửa phiếu")
b.para("Chỉ sửa được phiếu ở trạng thái Đang tạo do chính mình lập (hoặc Super Admin). Mở phiếu và bấm \"Sửa\". Khi sửa, không đổi được phiếu nguồn (đã khóa). Có thể cập nhật vị trí/số lượng nhập, ghi chú, file đính kèm; sau đó bấm \"Lưu\" (giữ nháp) hoặc \"Lưu & Nhập\" (gửi, tăng tồn kho). Nếu không đủ quyền sửa, hệ thống báo \"Bạn không có quyền sửa phiếu này\".")
b.h2("6.2. Hủy phiếu")
b.para("Chỉ hủy được phiếu ở trạng thái Đang tạo do chính mình lập, và loại nhập KHÔNG thuộc nhóm điều chuyển cùng thủ kho / điều chuyển nội bộ / điều chuyển chi nhánh / nhập tách / đã hủy. Bấm \"Hủy\": nếu hợp lệ, phiếu chuyển sang Đã hủy (\"Hủy thành công!\") và hủy luôn đề nghị nhập kho nguồn; ngược lại báo \"Không thể hủy!\".")
b.h2("6.3. Xuất lại (nhập nhầm)")
b.para("Với phiếu bị Không duyệt, người lập có thể \"Tạo phiếu xuất lại\" để xuất trả phần hàng đã nhập; phiếu gốc chuyển sang Đã xuất lại.")
b.h2("6.4. Nhập lại")
b.para("Với phiếu ở trạng thái Đã xuất lại, người lập có thể \"Tạo phiếu nhập lại\" để nhập lại; phiếu nhập lại nếu gửi (Lưu & Nhập) cũng làm tăng tồn kho. Phiếu gốc chuyển sang Đã nhập lại.")

# ------------------------------------------------- PHAN 7
b.h1("PHẦN 7: BÁO HÀNG THỪA, THIẾU NHẬP KHO")
b.para("Sau khi phiếu Đã hạch toán, nếu số lượng thực nhận lệch so với chứng từ, người dùng lập phiếu Báo hàng thừa, thiếu để ghi nhận và duyệt chênh lệch. Vào màn \"Chờ báo hàng thừa/thiếu\" từ danh sách, mỗi dòng có nút \"Báo hàng thừa thiếu\".")
b.table([
    ["Nội dung", "Ghi chú"],
    ["Điều kiện tạo", "Phiếu Đã hạch toán; bạn có quyền \"Tạo phiếu báo hàng thừa, thiếu nhập kho\" hoặc là Thủ kho; phiếu chưa có phiếu báo thừa/thiếu."],
    ["Trạng thái phiếu báo", "Đang tạo → Chờ duyệt → Đã duyệt (hoặc Đã hủy)."],
    ["Duyệt", "Người có quyền \"Duyệt phiếu báo hàng thừa, thiếu nhập kho\" duyệt phiếu báo."],
    ["In / Xuất Excel", "Có in danh sách và in chi tiết, xuất Excel danh sách chờ báo thừa/thiếu."],
])
b.para("Cần quyền tương ứng để tạo/duyệt; nếu thiếu quyền, nút sẽ không hiển thị.")

# ------------------------------------------------- PHAN 8
b.h1("PHẦN 8: IN ẤN & XUẤT EXCEL")
b.table([
    ["Chức năng", "Nội dung"],
    ["In phiếu nhập kho", "In mẫu Phiếu nhập kho (thông tin phiếu + bảng hàng hóa: Tên hàng hóa, Model, Mã hàng hóa, Code đặt hàng, Thương hiệu, Đơn vị tính, SL đề nghị, SL thực nhập, Vị trí, SL nhập)."],
    ["In dự kiến / phiếu xếp hàng", "In bản dự kiến vị trí xếp hàng (bước đi xếp hàng)."],
    ["In barcode", "In mã vạch hàng hóa của phiếu."],
    ["Xuất Excel danh sách", "Tải file Excel danh sách phiếu theo bộ lọc & phạm vi hiện tại (màn Danh sách / Tất cả). Danh sách quá lớn (≥ 2000 dòng) sẽ được gửi qua email: \"Danh sách phiếu nhập kho sẽ được gửi về mail sau ít phút!\"."],
])
b.para("Lưu ý: màn \"Kế toán kho\" không có nút Xuất Excel. Màn phiếu nhập kho KHÔNG có mục Lịch sử phiếu riêng (khác với màn Đề nghị nhập kho).")

# ------------------------------------------------- PHAN 9
b.h1("PHẦN 9: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Lập phiếu nhập kho ở đâu?", "Bấm Tạo mới ở màn danh sách rồi chọn một Đề nghị nhập kho đã gửi; hoặc bấm \"Tạo phiếu nhập kho\" trên chi tiết đề nghị nhập kho. Vai trò Thủ kho."],
    ["Vì sao không chọn được loại nhập / mặt hàng?", "Loại nhập và toàn bộ hàng hóa được kế thừa từ đề nghị nhập kho; tại màn này chỉ dự kiến vị trí và số lượng nhập."],
    ["Bấm Lưu rồi mà tồn kho chưa tăng?", "Bấm \"Lưu\" chỉ tạo nháp (đi xếp hàng). Tồn kho chỉ tăng khi bấm \"Lưu & Nhập\" (gửi phiếu) và không còn chuyến xe chưa duyệt."],
    ["Không thấy Ngày nhập / Giá nhập / phần hạch toán?", "Các phần này thuộc Phiếu nhập hàng (chứng từ hạch toán), lập sau bằng nút \"Tạo phiếu nhập hàng\"."],
    ["Không sửa được \"SL thực nhập\" trực tiếp", "SL thực nhập là tổng tự tính từ số lượng các dòng vị trí. Hãy nhập/điều chỉnh ở cột SL nhập của từng vị trí."],
    ["Báo \"Số lượng thực nhập phải bằng số lượng đề nghị\"", "Với loại điều chuyển kho chi nhánh, tổng thực nhập phải khớp đúng số đề nghị."],
    ["Không gửi được phiếu (còn chuyến xe chưa duyệt)", "Với vận chuyển bằng công ty, cần duyệt chuyến xe trước khi bấm Lưu & Nhập."],
    ["Không thấy nút Tạo phiếu nhập hàng / Không duyệt", "Hai nút này chỉ hiện với Kế toán kho cùng công ty khi phiếu ở trạng thái Chờ duyệt."],
    ["Không hủy được phiếu", "Chỉ hủy được phiếu Đang tạo do chính mình lập và không thuộc nhóm điều chuyển kho / nhập tách."],
    ["Nhập nhầm hàng thì xử lý sao?", "Nếu phiếu bị Không duyệt: Tạo phiếu xuất lại; sau đó có thể Tạo phiếu nhập lại. Nếu lệch số lượng khi nhập: dùng Báo hàng thừa/thiếu."],
    ["Danh sách trống dù có nhiều phiếu", "Có thể bạn chưa có quyền xem cấp phù hợp (chỉ thấy phiếu của mình). Liên hệ quản trị để cấp quyền xem theo công ty/phòng ban/kho."],
])

# ============================================================ FINISH
finish_macos(b)
print("XONG.")
