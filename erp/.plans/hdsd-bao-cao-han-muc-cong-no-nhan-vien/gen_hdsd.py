# -*- coding: utf-8 -*-
"""Dung HDSD man 'Bao cao han muc cong no nhan vien' (ERP, nhanh gop_db).

Bam khuon file mau HDSD_YeuCauXuatHang.docx (user chi dinh 15/10/2026) = khuon chi tiet
HdsdBuilder cua skill hrm/.claude/skills/hdsd-documenter. KHONG chen anh -> o "[Vi tri chen anh]".

Nguon khao sat (doc truc tiep code):
  ERP routes/web.php :5107-5113, Accounting/ReportLimitDebtController.php (searchDataEmployee,
      getDebtEmployeeByCustomer), app/Employee.php :134, Model/Accounting/AccountDetail.php :2317, :2401
  ERP views accounting/report_limit_debts/employee-index.blade.php,
      partials/classes/base/Datatable.blade.php, layouts/topmenubar.blade.php :1208

Hanh vi nghi loi KHONG ghi vao HDSD - xem testcase ../bao-cao-han-muc-cong-no-nhan-vien/.

Chay:  python gen_hdsd.py   (Windows, can Word de cap nhat muc luc)
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "..", "hrm", ".claude", "skills",
                                "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUTPUT = os.path.join(HERE, "HDSD_BaoCaoHanMucCongNoNhanVien.docx")
SHOTS = os.path.join(HERE, "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)

MENU = "menu Kế toán › Kế toán Công nợ › Nhân viên › \"Báo cáo hạn mức công nợ nhân viên\""


def img(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Báo cáo hạn mức công nợ nhân viên)",
    doc_title="HDSD - Báo cáo hạn mức công nợ nhân viên",
)

# ------------------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Báo cáo hạn mức công nợ nhân viên", "Màn hình tài liệu này hướng dẫn: so công nợ khách hàng do từng nhân viên phụ trách với hạn mức của nhân viên đó."],
    ["Hạn mức công nợ xuất hàng (nhân viên)", "Ngưỡng công nợ (VNĐ) khai báo ở ô \"Hạn mức công nợ xuất hàng\" trong hồ sơ nhân viên."],
    ["Tổng công nợ khách hàng hiện tại", "Tổng Nợ trừ tổng Có tài khoản phải thu khách hàng (131) trên các hợp đồng do chính nhân viên đó lập, cộng cho mọi khách."],
    ["Trong hạn mức", "Tổng công nợ nhỏ hơn hạn mức."],
    ["Vượt hạn mức", "Tổng công nợ lớn hơn hạn mức."],
    ["Hợp đồng do nhân viên lập", "Hợp đồng bán hàng, hợp đồng nguyên tắc / mở, biên bản thanh lý, hợp đồng dịch vụ, hợp đồng dự án mà nhân viên là người tạo."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "15/10/2026", "Khởi tạo HDSD màn Báo cáo hạn mức công nợ nhân viên (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Báo cáo hạn mức công nợ nhân viên giúp kế toán và quản lý theo dõi nhân viên kinh doanh nào "
       "đang để khách hàng nợ trong hoặc vượt hạn mức được giao, và xem chi tiết công nợ theo từng "
       "khách của nhân viên đó. Màn hình chỉ để xem.")
b.bullet(MENU + ".", bold_prefix="Đường dẫn: ")
b.bullet("ô \"Hạn mức công nợ xuất hàng\" trong hồ sơ nhân viên (danh mục nhân viên).",
         bold_prefix="Nơi đặt hạn mức: ")
b.bullet("mục \"Báo cáo hạn mức công nợ khách hàng\" ở nhóm Khách hàng cùng menu là một báo cáo khác "
         "(theo khách hàng).", bold_prefix="Phân biệt: ")

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Báo cáo không yêu cầu quyền riêng. Các quyền xem theo cấp dùng chung quyết định ô lọc theo "
       "đơn vị nào xuất hiện trong khối Bộ lọc:")
b.table([
    ["Quyền", "Ô lọc theo đơn vị hiển thị"],
    ["Xem tất cả phiếu", "Công ty, Phòng ban (mọi công ty)"],
    ["Xem tất cả phiếu của công ty", "Phòng ban (trong công ty mình)"],
    ["Xem tất cả phiếu của phòng ban", "Phòng ban (các phòng mình quản lý), Bộ phận"],
    ["Không có các quyền trên", "Không có ô lọc theo đơn vị"],
])

# ------------------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")
b.h2("1.1. Cách truy cập")
b.para("Vào " + MENU + ". Báo cáo tự tải danh sách nhân viên có hạn mức.")
img(b, "Ảnh menu Kế toán › Kế toán Công nợ › Nhân viên › Báo cáo hạn mức công nợ nhân viên.")

b.h2("1.2. Bố cục màn hình")
b.bullet("tiêu đề \"Báo cáo hạn mức công nợ nhân viên\".", bold_prefix="Thanh tiêu đề: ")
b.bullet("nút \"Bộ lọc\" để mở / thu gọn khối lọc; trong khối lọc có nút tìm kiếm (kính lúp) và nút "
         "làm mới.", bold_prefix="Bộ lọc: ")
b.bullet("mỗi dòng là một nhân viên có hạn mức.", bold_prefix="Bảng báo cáo: ")
b.bullet("dòng \"Hiển thị … của … mục\", ô chọn số dòng mỗi trang và thanh chuyển trang ở dưới bảng.",
         bold_prefix="Phân trang: ")
img(b, "Ảnh tổng quan Báo cáo hạn mức công nợ nhân viên.")

# ------------------------------------------------------------- PHAN 2
b.h1("PHẦN 2: XEM BÁO CÁO")
img(b, "Ảnh bảng báo cáo có dữ liệu.")

b.h2("2.1. Nhân viên được đưa vào báo cáo")
b.bullet("nhân viên có Hạn mức công nợ xuất hàng lớn hơn 0 trong hồ sơ nhân viên, thuộc mọi công ty.",
         bold_prefix="Có trong báo cáo: ")
b.bullet("nhân viên có hạn mức để trống hoặc bằng 0.", bold_prefix="Không có trong báo cáo: ")

b.h2("2.2. Các cột trong báo cáo")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Mã nhân viên", "Mã nhân viên; biểu tượng chữ i bên cạnh để mở chi tiết theo khách hàng (PHẦN 4)."],
    ["Tên nhân viên", "Họ tên nhân viên."],
    ["Bộ phận / Phòng ban", "Theo hồ sơ nhân sự hiện tại của nhân viên."],
    ["Hạn mức công nợ xuất hàng", "Hạn mức đặt trong hồ sơ nhân viên (VNĐ). Bấm tiêu đề cột để sắp xếp."],
    ["Tổng công nợ khách hàng hiện tại", "Tổng Nợ - Có tài khoản 131 trên các hợp đồng do nhân viên lập, trong khoảng ngày đang lọc. Có thể âm khi khách trả thừa."],
])

# ------------------------------------------------------------- PHAN 3
b.h1("PHẦN 3: TÌM KIẾM & LỌC")
b.para("Bấm nút \"Bộ lọc\" để mở khối lọc, chọn điều kiện rồi bấm kính lúp để lọc.")
img(b, "Ảnh khối Bộ lọc đang mở.")

b.h2("3.1. Các tiêu chí lọc")
b.table([
    ["Tiêu chí", "Cách dùng"],
    ["Công ty / Phòng ban / Bộ phận", "Chọn đơn vị theo hồ sơ nhân sự hiện tại của nhân viên (hiện theo quyền, xem Tổng quan mục 4). Chọn Công ty thì ô Phòng ban chỉ còn phòng của công ty đó."],
    ["Nhân viên", "Tìm theo tên nhân viên, gồm cả nhân viên đã nghỉ."],
    ["Chỉ tiêu vượt hạn mức", "Trong hạn mức: công nợ nhỏ hơn hạn mức. Vượt hạn mức: công nợ lớn hơn hạn mức."],
    ["Từ ngày / Đến ngày", "Lọc theo ngày ghi sổ của bút toán công nợ (dạng ngày/tháng/năm). Bỏ trống là lấy toàn bộ phát sinh từ trước tới nay."],
])

b.h2("3.2. Lưu ý khi lọc theo ngày")
b.bullet("hệ thống tính tới hết ngày liền trước \"Đến ngày\". Muốn lấy trọn ngày 30/09 thì nhập "
         "Đến ngày 01/10.", bold_prefix="Đến ngày: ")
b.bullet("lọc ngày chỉ thay đổi cột Tổng công nợ và kết quả chỉ tiêu Trong / Vượt hạn mức, không làm "
         "ẩn nhân viên.", bold_prefix="Phạm vi tác động: ")

b.h2("3.3. Làm mới & ghi nhớ bộ lọc")
b.bullet("bấm nút làm mới (mũi tên tròn, xanh lá) để xóa mọi điều kiện lọc.", bold_prefix="Làm mới: ")
b.bullet("hệ thống ghi nhớ bộ lọc lần trước khi mở lại báo cáo trên cùng trình duyệt.",
         bold_prefix="Ghi nhớ: ")

# ------------------------------------------------------------- PHAN 4
b.h1("PHẦN 4: XEM CHI TIẾT CÔNG NỢ THEO KHÁCH HÀNG")
b.para("Bấm biểu tượng chữ i cạnh Mã nhân viên để mở cửa sổ \"Thông tin chi tiết hạn mức dư nợ nhân "
       "viên <Mã> - <Tên> theo khách hàng\".")
img(b, "Ảnh cửa sổ chi tiết công nợ nhân viên theo khách hàng.")

b.h2("4.1. Các cột trong cửa sổ chi tiết")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Mã khách / Tên khách", "Khách hàng có phát sinh công nợ trên hợp đồng do nhân viên lập."],
    ["Công nợ", "Số tiền công nợ (VNĐ)."],
])

b.h2("4.2. Lọc trong cửa sổ chi tiết")
b.table([
    ["Tiêu chí", "Cách dùng"],
    ["Khách hàng", "Chọn một khách để xem riêng."],
    ["Số tiền từ / Số tiền đến", "Giữ các dòng có Công nợ trong khoảng đã nhập (tính cả số đúng bằng mốc)."],
])
b.para("Đóng cửa sổ để quay về báo cáo chính.")

# ------------------------------------------------------------- PHAN 5
b.h1("PHẦN 5: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Nhân viên không có trong báo cáo", "Hồ sơ nhân viên chưa đặt \"Hạn mức công nợ xuất hàng\" hoặc đang để 0. Cập nhật hồ sơ nhân viên rồi mở lại báo cáo."],
    ["Muốn đổi hạn mức của nhân viên", "Sửa ô \"Hạn mức công nợ xuất hàng\" trong hồ sơ nhân viên. Báo cáo không sửa được hạn mức."],
    ["Công nợ thấp hơn số trên sổ chi tiết công nợ khách", "Báo cáo chỉ tính bút toán gắn với hợp đồng do chính nhân viên lập; bút toán không chọn hợp đồng hoặc thuộc hợp đồng người khác lập không được tính."],
    ["Công nợ bằng đúng hạn mức không có trong chỉ tiêu nào", "Trong hạn mức là nhỏ hơn, Vượt hạn mức là lớn hơn hạn mức. Bỏ chọn chỉ tiêu để thấy nhân viên đó."],
    ["Không thấy ô lọc Công ty / Phòng ban", "Tài khoản chưa có quyền xem theo cấp tương ứng (Tổng quan mục 4)."],
    ["Cần xuất Excel", "Báo cáo chưa có chức năng xuất Excel / in."],
])


def finish_win(b):
    """Nhu b.finish() nhung bo qua tieu de trung voi tieu de cua chinh tai lieu
    (vd "TONG QUAN PHAN MEM" co ca o khung lan o noi dung that)."""
    import zipfile
    from docx import Document
    b.doc.save(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    b._update_fields_by_word()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    mine = {p.text.strip() for p in chk.paragraphs
            if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in mine]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


finish_win(b)
