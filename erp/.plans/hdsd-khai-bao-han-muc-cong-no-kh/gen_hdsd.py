# -*- coding: utf-8 -*-
"""Dung HDSD man 'Khai bao han muc cong no khach hang' (ERP, nhanh gop_db).

Bam khuon file mau HDSD_YeuCauXuatHang.docx (user chi dinh 15/10/2026) = khuon chi tiet
HdsdBuilder cua skill hrm/.claude/skills/hdsd-documenter (bia, muc luc, danh muc hinh,
TONG QUAN 4 muc, cac PHAN co bang). KHONG chen anh -> o "[Vi tri chen anh]".

Nguon khao sat (doc truc tiep code):
  ERP routes/web.php :5050-5054, Accounting/DeclareLimitDebtController.php
  ERP app/Model/Sale/Customer.php :1709 getDeclareLimitDebt, Model/Accounting/DeclareLimitDebt.php
  ERP views accounting/declare-limit-debt/index.blade.php,
      partials/classes/sale/{DeclareLimitDebt,CustomerDeclareLimitDebt}.blade.php
  ERP noi dung han muc: Warehouse/ProductExportRequestsController :2100,
      Sale/CustomersController :1907, Accounting/ReportLimitDebtController :32
  ERP layouts/topmenubar.blade.php :1192, PermissionsTableSeeder :402-403

Hanh vi nghi loi KHONG ghi vao HDSD - xem testcase ../khai-bao-han-muc-cong-no-kh/.

Chay:  python gen_hdsd.py   (Windows, can Word de cap nhat muc luc)
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "..", "hrm", ".claude", "skills",
                                "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUTPUT = os.path.join(HERE, "HDSD_KhaiBaoHanMucCongNoKhachHang.docx")
SHOTS = os.path.join(HERE, "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)

MENU = "menu Kế toán › Kế toán Công nợ › Khách hàng › \"Khai báo hạn mức công nợ KH\""


def img(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Khai báo hạn mức công nợ khách hàng)",
    doc_title="HDSD - Khai báo hạn mức công nợ khách hàng",
)

# ------------------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Khai báo hạn mức công nợ KH", "Màn hình tài liệu này hướng dẫn: đặt hạn mức công nợ cho từng khách hàng, riêng theo từng công ty."],
    ["Khai báo theo công ty", "Mỗi khách hàng có một bộ khai báo riêng cho mỗi công ty trong tập đoàn. Công ty nào chỉ dùng khai báo của công ty đó."],
    ["Cách tính lãi", "Cách tính lãi chậm trả áp cho khách: \"Theo vượt hạn mức công nợ\" hoặc \"Theo thời gian thanh toán\". Được dùng khi lập Phiếu tính lãi."],
    ["Hạn mức công nợ tính lãi", "Ngưỡng công nợ (VNĐ) dùng cho việc tính lãi trên Phiếu tính lãi."],
    ["Hạn mức công nợ xuất hàng", "Ngưỡng công nợ (VNĐ) tối đa được phép khi gửi duyệt Phiếu yêu cầu xuất hàng. Vượt ngưỡng thì phiếu phải qua duyệt vượt hạn mức."],
    ["Công nợ khách hàng", "Số dư Nợ trừ Có tài khoản phải thu khách hàng (131) của khách."],
    ["YCXH", "Phiếu yêu cầu xuất hàng."],
    ["TP / BGĐ", "Trưởng phòng / Ban giám đốc."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "15/10/2026", "Khởi tạo HDSD màn Khai báo hạn mức công nợ khách hàng (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Màn hình Khai báo hạn mức công nợ khách hàng dùng để đặt, cho từng cặp khách hàng - công ty, "
       "cách tính lãi và hai hạn mức công nợ. Khai báo có hiệu lực ngay sau khi Lưu và được dùng tiếp "
       "ở các màn nghiệp vụ khác.")
b.bullet(MENU + ".", bold_prefix="Đường dẫn: ")
b.para("Các màn sử dụng số liệu khai báo:")
b.table([
    ["Màn sử dụng", "Thông tin được dùng", "Tác dụng"],
    ["Phiếu yêu cầu xuất hàng", "Hạn mức công nợ xuất hàng", "Khi gửi duyệt, nếu công nợ của khách cộng giá trị phiếu vượt hạn mức thì hệ thống hỏi xác nhận gửi duyệt vượt hạn mức."],
    ["Phiếu tính lãi", "Cách tính lãi, Hạn mức công nợ tính lãi", "Tự điền khi chọn khách hàng; khách chưa có Cách tính lãi thì không chọn được."],
    ["Báo cáo hạn mức công nợ khách hàng", "Cả hai hạn mức", "Hiển thị ở cột Hạn mức tính lãi / Hạn mức được nợ để so với công nợ thực tế."],
])

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Màn hình có 2 quyền thuộc nhóm \"Quản lý khai báo hạn mức công nợ khách hàng\":")
b.bullet("được sửa và Lưu khai báo của CÔNG TY MÌNH.",
         bold_prefix="Quản lý khai báo hạn mức công nợ khách hàng: ")
b.bullet("được xem khai báo của MỌI công ty.",
         bold_prefix="Xem khai báo hạn mức công nợ khách hàng: ")
b.para("Dòng của công ty khác luôn ở chế độ chỉ xem (ô xám), kể cả khi người dùng có đủ 2 quyền. "
       "Công ty của người dùng lấy theo hồ sơ nhân sự hiện tại.")

# ------------------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")
b.h2("1.1. Cách truy cập")
b.para("Vào " + MENU + ". Màn hình mở ra và tự tải trang 1 của danh sách khách hàng.")
img(b, "Ảnh menu Kế toán › Kế toán Công nợ › Khách hàng › Khai báo hạn mức công nợ KH.")

b.h2("1.2. Bố cục màn hình")
b.bullet("tiêu đề \"Khai báo hạn mức công nợ khách hàng\".", bold_prefix="Thanh tiêu đề: ")
b.bullet("ô \"Chọn khách hàng\", nút tìm kiếm (kính lúp, xanh dương) và nút làm mới (mũi tên tròn, "
         "xanh lá).", bold_prefix="Khu vực lọc: ")
b.bullet("mỗi khách hàng là một khối dòng: dòng đầu là thông tin khách, các dòng dưới là từng công ty "
         "kèm 3 ô khai báo; nút Lưu nằm ở cột Hành động.", bold_prefix="Bảng khai báo: ")
b.bullet("thanh chuyển trang ở góc phải dưới bảng, 10 khách hàng mỗi trang.", bold_prefix="Phân trang: ")
img(b, "Ảnh tổng quan màn Khai báo hạn mức công nợ khách hàng.")

# ------------------------------------------------------------- PHAN 2
b.h1("PHẦN 2: DANH SÁCH KHÁCH HÀNG")
img(b, "Ảnh bảng khai báo với một khách hàng có nhiều dòng công ty.")

b.h2("2.1. Tìm kiếm khách hàng")
b.table([
    ["Bước", "Thao tác"],
    ["1", "Bấm ô \"Chọn khách hàng\", gõ mã khách, tên khách hoặc số điện thoại."],
    ["2", "Chọn khách trong danh sách gợi ý (dạng \"Mã - Tên khách\")."],
    ["3", "Bấm biểu tượng kính lúp để lọc. Bảng chỉ còn khách đã chọn."],
    ["4", "Bấm nút làm mới để xóa điều kiện lọc và hiện lại toàn bộ khách hàng."],
])
b.para("Gõ dưới 3 ký tự, danh sách gợi ý chỉ hiện tối đa 20 khách mới nhất; gõ thêm ký tự để tìm "
       "chính xác hơn.")

b.h2("2.2. Các cột trong bảng")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự khách hàng trong trang đang xem."],
    ["Mã khách / Tên khách / MST/CMND", "Thông tin của khách hàng."],
    ["Khai báo theo công ty", "Tên công ty của dòng khai báo."],
    ["Cách tính lãi", "Ô chọn: Theo vượt hạn mức công nợ / Theo thời gian thanh toán."],
    ["Hạn mức công nợ tính lãi", "Ô nhập số tiền (VNĐ)."],
    ["Hạn mức công nợ xuất hàng", "Ô nhập số tiền (VNĐ)."],
    ["Hành động", "Nút Lưu (biểu tượng đĩa mềm, chú thích \"Cập nhật\") lưu khai báo của khách hàng đó."],
])

b.h2("2.3. Các dòng công ty hiển thị")
b.table([
    ["Trường hợp", "Dòng công ty hiển thị"],
    ["Khách chưa từng được khai báo", "Đủ các công ty mà người dùng được xem; hai ô hạn mức hiện 0."],
    ["Khách đã có khai báo", "Các công ty đã được khai báo cho khách đó."],
    ["Người dùng chỉ có quyền Quản lý", "Chỉ dòng công ty của mình."],
])

# ------------------------------------------------------------- PHAN 3
b.h1("PHẦN 3: KHAI BÁO / SỬA HẠN MỨC")
b.para("Khai báo mới và sửa khai báo làm cùng một cách: sửa trực tiếp trên bảng rồi bấm Lưu ở khách "
       "hàng đó.")
img(b, "Ảnh dòng công ty đang được nhập hạn mức.")

b.h2("3.1. Các bước khai báo")
b.table([
    ["Bước", "Thao tác"],
    ["1", "Tìm khách hàng cần khai báo (mục 2.1)."],
    ["2", "Ở dòng công ty của mình, bấm ô \"Cách tính lãi\" và chọn một trong hai cách tính."],
    ["3", "Nhập \"Hạn mức công nợ tính lãi\"."],
    ["4", "Nhập \"Hạn mức công nợ xuất hàng\"."],
    ["5", "Bấm nút Lưu (đĩa mềm) ở cột Hành động của khách hàng đó."],
    ["6", "Hệ thống báo \"Cập nhật thành công!\"."],
])

b.h2("3.2. Quy tắc nhập")
b.table([
    ["Trường", "Bắt buộc", "Quy tắc"],
    ["Cách tính lãi", "Không", "Phải bấm chọn để giá trị được lưu. Khách chưa có Cách tính lãi thì không lập được Phiếu tính lãi."],
    ["Hạn mức công nợ tính lãi", "Không", "Số nguyên VNĐ, nhập số liền; hệ thống tự thêm dấu phẩy ngăn cách hàng nghìn khi gõ."],
    ["Hạn mức công nợ xuất hàng", "Không", "Như trên. Để trống hoặc 0 nghĩa là không giới hạn khi gửi duyệt Phiếu yêu cầu xuất hàng."],
])
b.para("Không gõ dấu chấm để ngăn cách hàng nghìn (vd 500.000.000): dấu chấm được hiểu là dấu thập "
       "phân. Hãy nhập 500000000 hoặc 500,000,000.")

b.h2("3.3. Lưu khai báo")
b.bullet("mỗi khách hàng có nút Lưu riêng; nút chỉ lưu khối dòng của khách hàng đó.",
         bold_prefix="Lưu theo khách: ")
b.bullet("chuyển trang hoặc tải lại trang trước khi Lưu sẽ làm mất số vừa nhập.",
         bold_prefix="Chưa Lưu: ")
b.bullet("thông báo vàng \"Cập nhật thất bại!\" hoặc đỏ \"Đã có lỗi xảy ra\": kiểm tra kết nối rồi "
         "bấm Lưu lại.", bold_prefix="Lưu không thành công: ")
b.bullet("tải lại trang trước khi khai báo để thấy số liệu mới nhất do người khác vừa lưu.",
         bold_prefix="Nhiều người cùng khai báo: ")

b.h2("3.4. Bỏ hạn mức")
b.para("Màn hình không có nút Xóa. Muốn bỏ hạn mức, xóa trắng (hoặc nhập 0) ô hạn mức tương ứng rồi "
       "bấm Lưu.")

# ------------------------------------------------------------- PHAN 4
b.h1("PHẦN 4: TÁC ĐỘNG TỚI CÁC MÀN NGHIỆP VỤ")
b.h2("4.1. Phiếu yêu cầu xuất hàng")
b.para("Khi nhân viên gửi duyệt Phiếu yêu cầu xuất hàng, hệ thống lấy Hạn mức công nợ xuất hàng của "
       "khách theo công ty của người lập phiếu và so sánh:")
b.table([
    ["Thành phần", "Cách tính"],
    ["Công nợ hiện tại", "Công nợ của khách trên các hợp đồng do chính người lập phiếu tạo."],
    ["Giá trị phiếu", "Tổng tiền sau thuế của phiếu đang gửi."],
    ["Kết quả", "Công nợ hiện tại + Giá trị phiếu lớn hơn hạn mức: hiện hộp \"Xác nhận!\" với nội dung \"Yêu cầu vượt hạn mức công nợ xuất hàng cho khách hàng (… VNĐ)\". Bấm Đồng ý để gửi phiếu sang luồng duyệt vượt hạn mức."],
])
img(b, "Ảnh hộp xác nhận vượt hạn mức công nợ xuất hàng trên Phiếu yêu cầu xuất hàng.")

b.h2("4.2. Phiếu tính lãi")
b.para("Khi chọn khách hàng trên Phiếu tính lãi, hệ thống tự điền Cách tính lãi và Hạn mức công nợ tính "
       "lãi theo khai báo của công ty người lập. Khách chưa có Cách tính lãi ở công ty đó sẽ báo "
       "\"Khách hàng chưa cấu hình cách tính lãi!\".")

b.h2("4.3. Báo cáo hạn mức công nợ khách hàng")
b.para("Báo cáo liệt kê các khách của công ty người xem có ít nhất một hạn mức lớn hơn 0, kèm công nợ "
       "hiện tại để theo dõi khách trong hạn mức / vượt hạn mức.")

# ------------------------------------------------------------- PHAN 5
b.h1("PHẦN 5: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Các ô bị xám, không nhập được", "Dòng đó thuộc công ty khác, hoặc tài khoản chưa có quyền \"Quản lý khai báo hạn mức công nợ khách hàng\". Liên hệ quản trị hệ thống nếu cần quyền."],
    ["Chỉ thấy dòng công ty mình", "Tài khoản chỉ có quyền Quản lý. Cần thêm quyền \"Xem khai báo hạn mức công nợ khách hàng\" để xem khai báo của các công ty khác."],
    ["Phiếu tính lãi báo \"Khách hàng chưa cấu hình cách tính lãi!\"", "Vào màn này, bấm chọn Cách tính lãi ở dòng công ty mình cho khách đó rồi bấm Lưu."],
    ["Đã đặt hạn mức nhưng YCXH không hỏi vượt hạn mức", "Kiểm tra: hạn mức đặt đúng công ty của người lập phiếu chưa; công nợ trên hợp đồng của người lập cộng giá trị phiếu có thật sự vượt hạn mức không; hạn mức có đang là 0 không."],
    ["Số hiển thị khác số vừa nhập", "Kiểm tra cách gõ: không dùng dấu chấm ngăn cách hàng nghìn; số lẻ thập phân được làm tròn."],
    ["Nhập xong chuyển trang thì mất số", "Phải bấm Lưu ở từng khách hàng trước khi chuyển trang."],
])

def finish_win(b):
    """Nhu b.finish() nhung bo qua tieu de trung voi tieu de cua chinh tai lieu
    (vd "TONG QUAN PHAN MEM" co ca o khung lan o noi dung that)."""
    import zipfile
    from docx import Document
    b.doc.save(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    b._update_fields_by_word()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    mine = {p.text.strip() for p in chk.paragraphs
            if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in mine]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


finish_win(b)
