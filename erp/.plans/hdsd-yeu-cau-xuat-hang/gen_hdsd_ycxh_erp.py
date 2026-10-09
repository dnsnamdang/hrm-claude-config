# -*- coding: utf-8 -*-
"""Dung HDSD man 'Yeu cau xuat hang' (Phieu Yeu cau xuat hang) - ERP TanPhatDev.

Nguon: khao sat controller Warehouse\\ProductExportRequestsController + model
ProductExportRequest + views warehouse/product_export_requests + PermissionsTableSeeder.
KHONG chen anh (user chot khong can) -> dung o placeholder.

Chay:  /usr/local/bin/python3 gen_hdsd_ycxh_erp.py
Output: ERP/HDSD_luongchinh/HDSD_YeuCauXuatHang.docx
"""
import os
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

ERP = "/Users/nguyentrancu/DEV/code/ERP-HRM/ERP"
sys.path.insert(0, os.path.join("/Users/nguyentrancu/DEV/code/ERP-HRM/HRM",
                                ".claude", "skills", "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUT_DIR = os.path.join(ERP, "HDSD_luongchinh")
os.makedirs(OUT_DIR, exist_ok=True)
OUTPUT = os.path.join(OUT_DIR, "HDSD_YeuCauXuatHang.docx")

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)


def img_placeholder(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


def clear_toc_cache(path):
    """Xoa text cache trong moi field TOC (muc luc + danh muc hinh) -> Word tu dung lai khi mo."""
    doc = Document(path)
    body = doc.element.body
    fields = []
    for el in body.iter():
        tag = el.tag
        if tag == qn('w:fldChar'):
            t = el.get(qn('w:fldCharType'))
            if t == 'begin':
                fields.append({'instr': '', 'in_result': False})
            elif t == 'separate':
                if fields:
                    fields[-1]['in_result'] = True
            elif t == 'end':
                if fields:
                    fields.pop()
        elif tag == qn('w:instrText'):
            if fields:
                fields[-1]['instr'] += (el.text or '')
        elif tag == qn('w:t'):
            if any(f['in_result'] and 'TOC' in f['instr'] for f in fields):
                el.text = ''
    doc.save(path)


def finish_macos(b):
    b.doc.save(b.output)
    clear_toc_cache(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    my_headings = {p.text.strip() for p in chk.paragraphs
                   if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in my_headings]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables),
          "| Anh:", len(chk.inline_shapes), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


# ============================================================ NOI DUNG
b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Yêu cầu xuất hàng)",
    doc_title="HDSD - Yêu cầu xuất hàng",
)

# ------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Phiếu Yêu cầu xuất hàng (YCXH)", "Chứng từ đầu tiên của quy trình xuất hàng, đề nghị xuất hàng ra khỏi kho. Đây chính là màn hình tài liệu này hướng dẫn."],
    ["Đề nghị xuất kho", "Chứng từ bước 2, do Kế toán kho lập từ YCXH đã duyệt."],
    ["Phiếu xuất kho", "Chứng từ bước 3, kho thực hiện xuất."],
    ["Phiếu xuất hàng", "Chứng từ bước 4 (cuối), ghi nhận & hạch toán."],
    ["Loại yêu cầu (loại xuất)", "Phân loại mục đích xuất hàng: xuất bán, xuất mượn, xuất trả NCC, điều chuyển kho, xuất bảo hành, xuất sản xuất… Loại quyết định các trường phải nhập."],
    ["Xuất thẳng", "Xuất trực tiếp không qua nghiệp vụ giữ hàng/đề nghị xuất kho theo cách thông thường; khi tick, không cần chọn Kho xuất và cho phép tạo Phiếu xuất hàng ngay."],
    ["Hàng giữ (prepick)", "Hàng đã được giữ trước cho phiếu; hệ thống kiểm soát hạn giữ."],
    ["Hạn mức công nợ", "Ngưỡng công nợ của khách hàng/nhân viên. Xuất vượt hạn mức phải được Trưởng phòng (TP) hoặc Ban giám đốc (BGĐ) duyệt."],
    ["TP duyệt / BGĐ duyệt", "Cấp duyệt phiếu khi vượt hạn mức: Trưởng phòng duyệt trước, vượt tiếp thì chuyển Ban giám đốc duyệt."],
    ["Kế toán kho", "Vai trò tiếp nhận phiếu đã duyệt để lập đề nghị xuất kho / phiếu xuất hàng."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "09/09/2026", "Khởi tạo HDSD màn Yêu cầu xuất hàng (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Màn hình Yêu cầu xuất hàng dùng để lập và quản lý các phiếu đề nghị xuất hàng ra khỏi kho cho nhiều mục đích khác nhau (bán hàng, cho mượn, trả nhà cung cấp, điều chuyển kho, bảo hành, sản xuất, thực hiện hợp đồng…). Đây là bước khởi đầu của quy trình xuất hàng gồm 4 chứng từ nối tiếp.")
b.bullet("menu Khởi tạo › Hàng hóa › Nhập - xuất hàng › \"Phiếu Yêu cầu xuất hàng\".", bold_prefix="Đường dẫn: ")
b.bullet("một mục menu riêng \"Phiếu yêu cầu xuất hàng chờ duyệt\" dẫn tới danh sách các phiếu đang chờ duyệt (dành cho TP/BGĐ/Kế toán kho).", bold_prefix="Chờ duyệt: ")
b.para("Vị trí trong quy trình xuất hàng (4 chứng từ nối tiếp):")
b.table([
    ["Bước", "Chứng từ", "Ghi chú"],
    ["1", "Yêu cầu xuất hàng", "Màn hình này — lập, gửi duyệt."],
    ["2", "Đề nghị xuất kho", "Kế toán kho lập từ YCXH đã duyệt."],
    ["3", "Phiếu xuất kho", "Kho thực hiện xuất."],
    ["4", "Phiếu xuất hàng", "Ghi nhận & hạch toán, hoàn tất."],
])

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Phạm vi dữ liệu nhìn thấy và các thao tác đều phụ thuộc quyền. Có các nhóm quyền chính:")
b.bullet("4 quyền xem theo cấp (tổng công ty / công ty / phòng ban / bộ phận) quyết định thấy phiếu của phạm vi nào; không có quyền nào thì chỉ thấy phiếu do chính mình tạo.", bold_prefix="Xem danh sách: ")
b.bullet("cần để lập Đề nghị xuất kho / Phiếu xuất hàng và xem phiếu trong công ty.", bold_prefix="Kế toán kho: ")
b.bullet("Trưởng phòng / Ban giám đốc duyệt xuất hàng vượt hạn mức công nợ; Trưởng phòng / Ban giám đốc duyệt hàng mượn.", bold_prefix="Quyền duyệt: ")
b.para("Các thao tác Sửa / Hủy / Xóa không phụ thuộc quyền mà phụ thuộc quyền sở hữu + trạng thái: chỉ người tạo mới sửa/hủy/xóa được phiếu ở trạng thái Đang tạo. Chi tiết ở PHẦN 2 và PHẦN 5.")

# ------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")

b.h2("1.1. Cách truy cập & các màn danh sách")
b.para("Vào menu Khởi tạo › Hàng hóa › Nhập - xuất hàng › \"Phiếu Yêu cầu xuất hàng\". Hệ thống có nhiều màn danh sách phục vụ các vai trò khác nhau:")
b.table([
    ["Màn", "Dùng cho", "Nội dung"],
    ["Danh sách (tất cả)", "Người lập / người xem", "Toàn bộ phiếu trong phạm vi quyền của người dùng."],
    ["Chờ duyệt", "TP / BGĐ / Kế toán kho", "Các phiếu đang chờ duyệt hoặc chờ xử lý bước tiếp."],
    ["Màn duyệt của Trưởng phòng/BGĐ", "TP / BGĐ", "Phiếu vượt hạn mức cần cấp trên duyệt; có thêm cột Phòng ban."],
    ["Màn kế toán", "Kế toán kho", "Phiếu đã duyệt để lập đề nghị xuất kho / phiếu xuất hàng."],
])
img_placeholder(b, "Ảnh menu Khởi tạo › Hàng hóa › Nhập - xuất hàng › Phiếu Yêu cầu xuất hàng.")

b.h2("1.2. Bố cục màn hình danh sách")
b.bullet("tiêu đề \"Danh sách yêu cầu xuất hàng\".", bold_prefix="Thanh tiêu đề: ")
b.bullet("các ô tìm theo cột (Kho hàng, Tên hàng, Model, Mã phiếu, Loại, Số hợp đồng, Trạng thái, Khách hàng, Người tạo, Người duyệt…) và lọc theo khoảng thời gian.", bold_prefix="Khu vực lọc: ")
b.bullet("bảng liệt kê phiếu theo phạm vi quyền, cột Hành động ở cuối mỗi dòng.", bold_prefix="Bảng danh sách: ")
b.bullet("chọn số dòng/trang và chuyển trang.", bold_prefix="Phân trang: ")
img_placeholder(b, "Ảnh tổng quan màn danh sách yêu cầu xuất hàng.")

# ------------------------------------------------- PHAN 2
b.h1("PHẦN 2: DANH SÁCH YÊU CẦU XUẤT HÀNG")
img_placeholder(b, "Ảnh màn danh sách với dữ liệu mẫu và cột Hành động.")

b.h2("2.1. Phân quyền & hướng dẫn theo quyền")
b.para("Danh sách hiển thị theo phạm vi quyền của người đăng nhập. Bảng dưới liệt kê đầy đủ các quyền liên quan tới màn. Người dùng thường chỉ có một vài quyền — hãy đối chiếu đúng phần áp dụng cho mình.")
b.table([
    ["Tên quyền", "Cho phép làm gì", "Nút/màn tương ứng"],
    ["Xem yêu cầu xuất hàng theo tổng công ty", "Thấy toàn bộ phiếu của mọi công ty.", "Danh sách"],
    ["Xem yêu cầu xuất hàng theo công ty", "Thấy phiếu trong công ty của mình.", "Danh sách"],
    ["Xem yêu cầu xuất hàng theo phòng ban", "Thấy phiếu trong phòng ban mình quản lý.", "Danh sách"],
    ["Xem yêu cầu xuất hàng theo bộ phận", "Thấy phiếu trong bộ phận mình quản lý.", "Danh sách"],
    ["Kế toán kho", "Xem phiếu trong công ty; lập Đề nghị xuất kho / Phiếu xuất hàng từ phiếu đã duyệt.", "Nút \"Tạo đề nghị xuất kho\", \"Tạo phiếu xuất hàng\"; màn kế toán"],
    ["Trưởng phòng duyệt xuất hàng vượt hạn mức công nợ", "TP duyệt phiếu vượt hạn mức công nợ.", "Nút \"TP duyệt\" / \"Chuyển BGĐ duyệt\""],
    ["Ban giám đốc duyệt xuất hàng vượt hạn mức công nợ", "BGĐ duyệt phiếu vượt hạn mức công nợ.", "Nút \"BGĐ duyệt\""],
    ["Trưởng phòng duyệt hàng mượn", "TP duyệt phiếu xuất mượn.", "Nút \"TP duyệt\" (phiếu mượn)"],
    ["Ban giám đốc duyệt hàng mượn", "BGĐ duyệt phiếu xuất mượn.", "Nút \"BGĐ duyệt\" (phiếu mượn)"],
])
b.para("Lưu ý phạm vi: nếu không có quyền xem cấp nào, người dùng chỉ thấy phiếu do chính mình tạo. Việc lập phiếu (Tạo mới) thì ai cũng làm được. Nếu thiếu quyền, các nút tương ứng sẽ không hiển thị; truy cập trực tiếp bằng đường dẫn sẽ báo không có quyền.")

b.h3("Người dùng có quyền \"Xem yêu cầu xuất hàng theo tổng công ty\"")
b.para("Thấy toàn bộ phiếu của mọi công ty. Thực hiện đầy đủ thao tác xem/lọc; sửa/hủy/xóa vẫn theo quy tắc sở hữu + trạng thái (PHẦN 5).")

b.h3("Người dùng có quyền \"Xem yêu cầu xuất hàng theo công ty / phòng ban / bộ phận\"")
b.para("Thấy phiếu trong phạm vi công ty / phòng ban / bộ phận tương ứng mà mình phụ trách, cộng thêm phiếu do chính mình tạo. Phòng ban/bộ phận được xác định theo phân công quản lý của nhân viên.")

b.h3("Người dùng là Kế toán kho")
b.para("Thấy các phiếu trong công ty. Với phiếu đã ở trạng thái Chờ duyệt: nếu là xuất thường (không xuất thẳng) sẽ có nút \"Tạo đề nghị xuất kho\"; nếu là xuất thẳng sẽ có nút \"Tạo phiếu xuất hàng\". Đây là vai trò xử lý bước tiếp của quy trình.")

b.h3("Người dùng là Trưởng phòng / Ban giám đốc (có quyền duyệt)")
b.para("Với phiếu đang chờ cấp mình duyệt (Đợi TP duyệt / Đợi BGĐ duyệt), có nút duyệt tương ứng và nút \"Không duyệt\" (từ chối). Xem chi tiết luồng duyệt ở PHẦN 4.")

b.h2("2.2. Tìm kiếm & lọc")
b.para("Các ô lọc theo cột (nhập/chọn ngay trên đầu bảng):")
b.table([
    ["Tiêu chí", "Cách dùng"],
    ["Kho hàng", "Chọn kho (theo kho người dùng được phép)."],
    ["Tên hàng hóa / Model", "Nhập để lọc phiếu có chứa hàng theo tên hoặc model."],
    ["Mã phiếu", "Nhập mã phiếu."],
    ["Loại", "Chọn loại yêu cầu xuất (xem mục 3.1)."],
    ["Số hợp đồng", "Nhập số hợp đồng liên quan."],
    ["Xuất thẳng", "Lọc Có / Không."],
    ["Cần lắp đặt", "Lọc Có / Không."],
    ["Trạng thái", "Chọn trạng thái phiếu (xem mục 2.5)."],
    ["Khách hàng / Người tạo / Người duyệt", "Chọn (tìm theo tên) để lọc."],
    ["Khoảng thời gian", "Lọc theo khoảng ngày lập phiếu."],
])

b.h2("2.3. Các cột trong danh sách")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Ngày lập", "Ngày tạo phiếu."],
    ["Mã phiếu", "Mã yêu cầu xuất hàng (bấm để mở chi tiết)."],
    ["Loại", "Loại yêu cầu xuất."],
    ["Số hợp đồng", "Hợp đồng liên quan (có liên kết mở chi tiết theo loại hợp đồng)."],
    ["Khách hàng", "Khách hàng của phiếu."],
    ["Trạng thái", "Trạng thái phiếu, hiển thị nhãn màu."],
    ["Người lập", "Người tạo phiếu (mã phòng - tên)."],
    ["Người duyệt", "Người đã duyệt phiếu."],
    ["Ngày nhận", "Thời điểm phiếu được gửi/nhận để xử lý."],
    ["Ngày duyệt", "Thời điểm duyệt."],
    ["Hành động", "Các nút thao tác (xem mục 2.4)."],
])

b.h2("2.4. Thao tác trên từng dòng (cột Hành động)")
b.para("Các nút hiển thị tùy trạng thái, loại phiếu, quyền và quyền sở hữu:")
b.table([
    ["Nút", "Điều kiện hiển thị", "Kết quả"],
    ["Sửa yêu cầu", "Phiếu Đang tạo do chính bạn lập.", "Mở màn sửa phiếu."],
    ["Xóa", "Phiếu Đang tạo do chính bạn lập.", "Xóa hẳn phiếu (cần xác nhận)."],
    ["Tạo đề nghị xuất kho", "Bạn là Kế toán kho; phiếu Chờ duyệt; không xuất thẳng; cùng công ty/kho.", "Mở màn lập Đề nghị xuất kho (bước 2)."],
    ["Tạo phiếu xuất hàng", "Bạn là Kế toán kho; phiếu Chờ duyệt; xuất thẳng.", "Mở màn lập Phiếu xuất hàng."],
    ["Tạo yêu cầu lắp đặt", "Phiếu cần lắp đặt, chưa có yêu cầu lắp đặt, trạng thái Đang tạo, do bạn lập.", "Mở màn tạo yêu cầu lắp đặt."],
    ["TP duyệt", "Bạn có quyền TP duyệt; phiếu Đợi TP duyệt thuộc phòng ban bạn quản lý.", "Mở chi tiết để duyệt."],
    ["BGĐ duyệt", "Bạn có quyền BGĐ duyệt; phiếu Đợi BGĐ duyệt.", "Mở chi tiết để duyệt."],
    ["In yêu cầu", "Luôn có.", "In phiếu yêu cầu xuất hàng."],
    ["In bộ giấy đi đường", "Luôn có.", "In lệnh điều động + phiếu xuất kho đi đường."],
    ["In biên bản giao nhận", "Luôn có.", "In biên bản giao nhận (chọn nhân viên)."],
    ["Xuất excel danh sách hàng", "Luôn có.", "Tải Excel danh sách hàng của phiếu."],
])
b.para("Nếu không thấy nút Sửa/Xóa: phiếu đã qua trạng thái Đang tạo, hoặc không phải phiếu của bạn. Không thấy nút duyệt: phiếu chưa ở đúng trạng thái chờ cấp bạn, hoặc bạn chưa có quyền duyệt tương ứng.")

b.h2("2.5. Bảng trạng thái phiếu")
b.table([
    ["Trạng thái", "Ý nghĩa"],
    ["Đang tạo", "Phiếu nháp, chưa gửi duyệt. Chỉ người tạo sửa/hủy/xóa được."],
    ["Chờ duyệt", "Đã gửi; chờ Kế toán kho xử lý bước tiếp."],
    ["Đợi TP duyệt", "Vượt hạn mức, chờ Trưởng phòng duyệt."],
    ["Đợi BGĐ duyệt", "Chờ Ban giám đốc duyệt."],
    ["Đã đề nghị", "Đã tạo đề nghị xuất kho từ phiếu."],
    ["Đang lập đề nghị", "Đang lập đề nghị xuất kho."],
    ["Đang xuất kho", "Kho đang xuất."],
    ["Đã xuất kho", "Kho đã xuất xong."],
    ["Đang hạch toán / Đã hạch toán", "Đang / đã ghi nhận hạch toán."],
    ["Đã hủy phiếu", "Phiếu đã bị hủy."],
])

# ------------------------------------------------- PHAN 3
b.h1("PHẦN 3: TẠO YÊU CẦU XUẤT HÀNG")
b.para("Trên màn danh sách bấm nút tạo mới để mở màn lập yêu cầu xuất hàng. Việc lập phiếu không đòi hỏi quyền đặc biệt; phiếu nháp thuộc về người lập.")
img_placeholder(b, "Ảnh màn Tạo yêu cầu xuất hàng (form).")

b.h2("3.1. Chọn loại yêu cầu")
b.para("Trường \"Loại yêu cầu\" (bắt buộc) quyết định các trường phía dưới và bảng hàng hóa. Các loại có thể tạo:")
b.table([
    ["Loại yêu cầu", "Mục đích"],
    ["Xuất bán hàng", "Xuất bán cho khách theo hợp đồng hãng."],
    ["Xuất khuyến mại", "Xuất hàng khuyến mại."],
    ["Xuất mượn", "Cho mượn hàng (có ngày mượn/ngày trả)."],
    ["Xuất trả nhà cung cấp", "Trả hàng lại nhà cung cấp (theo phiếu yêu cầu nhập)."],
    ["Xuất điều chuyển kho nội bộ", "Chuyển hàng giữa các kho nội bộ (có kho nhập)."],
    ["Xuất điều chuyển kho chi nhánh", "Chuyển hàng sang công ty/chi nhánh khác."],
    ["Xuất hàng gửi", "Xuất hàng gửi cho khách hàng."],
    ["Xuất bán bảo hành", "Xuất theo hợp đồng dịch vụ bảo hành."],
    ["Xuất bán SC - BD", "Xuất theo hợp đồng sửa chữa - bảo dưỡng."],
    ["Xuất sản xuất", "Xuất phục vụ sản xuất."],
    ["Xuất thực hiện hợp đồng", "Xuất theo hợp đồng thực hiện."],
    ["Xuất khác", "Các trường hợp còn lại (kèm ghi nhận chi phí HĐ)."],
])
b.para("Lưu ý: một số loại như Xuất ghép/tách, Xuất bán (hợp đồng dự án - cũ), Xuất sản xuất/bán theo hợp đồng chỉ dùng để hiển thị/lọc, không tạo tay tại màn này. Khi sửa phiếu, trường Loại yêu cầu bị khóa (không đổi loại sau khi đã tạo).")

b.h2("3.2. Các trường thông tin chung")
b.para("Tùy loại yêu cầu, các trường sau hiển thị (cột \"Hiện khi\" cho biết loại nào áp dụng):")
b.table([
    ["Trường", "Bắt buộc", "Hiện khi / ghi chú"],
    ["Loại yêu cầu", "Có", "Luôn hiện. Khóa khi đang sửa."],
    ["Xuất thẳng", "Không", "Mọi loại trừ Xuất mượn và Xuất khuyến mại. Tick để xuất thẳng (không cần Kho xuất)."],
    ["Cần lắp đặt", "Không", "Chỉ khi loại Xuất bán hàng. Chọn Có nếu đơn có lắp đặt."],
    ["Chọn hợp đồng / đơn hàng", "Có (khi hiện)", "Các loại xuất bán theo hợp đồng/đơn hàng. Khóa khi đang sửa."],
    ["Chọn hợp đồng thực hiện", "Có (khi hiện)", "Loại Xuất thực hiện hợp đồng."],
    ["Chọn HĐ SC - BH", "Có (khi hiện)", "Loại Xuất bán bảo hành / Xuất bán SC - BD."],
    ["Chọn hợp đồng hãng + Nhóm VAT + Nhóm hàng/khuyến mại", "Có (khi hiện)", "Loại Xuất bán hàng / Xuất khuyến mại. Nhóm VAT chỉ với Xuất bán hàng."],
    ["Chọn hợp đồng dịch vụ", "Có (khi hiện)", "Loại xuất bán theo hợp đồng dịch vụ."],
    ["Khách hàng", "Có (khi hiện)", "Loại Xuất hàng gửi."],
    ["Chọn phiếu yêu cầu nhập hàng", "Có (khi hiện)", "Loại Xuất trả nhà cung cấp."],
    ["Gửi đến công ty + Kiểu nhập kho + Phiếu yêu cầu chuyển hàng", "Có (khi hiện)", "Loại Xuất điều chuyển kho chi nhánh. Kiểu nhập: Nhập thẳng / Nhập về kho."],
    ["Ngày cần mượn / Ngày trả", "Có (khi hiện)", "Loại Xuất mượn."],
    ["Kho xuất", "Có (khi hiện)", "Mọi loại trừ điều chuyển chi nhánh, và khi không tick Xuất thẳng."],
    ["Kho nhập", "Có (khi hiện)", "Loại điều chuyển kho nội bộ; khác với Kho xuất."],
    ["Vận chuyển", "Không", "Các loại có giao hàng: Tự vận chuyển / Công ty vận chuyển / Khách hàng vận chuyển."],
    ["Số km dự kiến", "Không", "Chỉ khi chọn \"Công ty vận chuyển\"."],
    ["Lịch lắp đặt / Ghi chú / File đính kèm", "Không", "Ghi chú tối đa 255 ký tự; đính kèm pdf, ảnh, doc, xls, ppt."],
])

b.h2("3.3. Bảng chi tiết hàng hóa")
b.para("Phần dưới form là bảng hàng hóa (có các tab: Hàng hóa, theo Nhóm hợp đồng hãng, Khuyến mại, Tổng hợp khuyến mại — tùy loại). Mỗi dòng hàng gồm: Tên hàng hóa, Model, Mã, ĐVT, các cột số lượng (Tồn, Đang giữ, Hợp đồng, Đã xuất, Đang xuất, Đề nghị), và với các loại xuất bán còn có Giá niêm yết, Giá bán, Thành tiền, Chiết khấu, Đơn giá sau giảm, Thành tiền sau giảm, VAT, Tiền VAT, Thành tiền sau VAT, Hình ảnh. Cột giá được ẩn với loại Xuất KM (nguyên tắc) và Xuất khuyến mại. Nhập Số lượng đề nghị cho từng dòng cần xuất.")

b.h2("3.4. Các nút lưu")
b.table([
    ["Nút", "Hành động"],
    ["Lưu", "Lưu phiếu ở trạng thái Đang tạo (nháp)."],
    ["Lưu & Gửi", "Gửi duyệt: với đa số loại chuyển sang Chờ duyệt; với loại Xuất mượn chuyển sang Đợi TP duyệt."],
    ["Hủy", "Quay lại danh sách, không lưu."],
])
b.para("Nếu phiếu cần lắp đặt mà chưa tạo yêu cầu lắp đặt, hệ thống sẽ giữ phiếu ở trạng thái Đang tạo và nhắc phải tạo yêu cầu lắp đặt trước khi gửi.")

b.h2("3.5. Quy tắc kiểm tra dữ liệu (validate)")
b.table([
    ["Trường / tình huống", "Yêu cầu"],
    ["Loại yêu cầu", "Bắt buộc, thuộc danh sách loại hợp lệ."],
    ["Danh sách hàng", "Bắt buộc ít nhất 1 dòng; số lượng mỗi dòng từ 0 đến 999.999."],
    ["Ghi chú / đính kèm", "Ghi chú tối đa 255 ký tự; tệp đúng định dạng cho phép."],
    ["Kho xuất", "Bắt buộc (trừ điều chuyển chi nhánh hoặc khi tick Xuất thẳng)."],
    ["Điều chuyển chi nhánh", "Bắt buộc: Gửi đến công ty, Kiểu nhập kho, Phiếu yêu cầu chuyển hàng."],
    ["Xuất bảo hành / SC - BD", "Bắt buộc chọn hợp đồng SC - BH."],
    ["Xuất bán hàng (HĐ hãng)", "Bắt buộc chọn hợp đồng hãng, nhóm VAT và nhóm hàng."],
    ["Xuất mượn", "Bắt buộc Ngày cần mượn và Ngày trả."],
    ["Xuất trả nhà cung cấp", "Bắt buộc chọn phiếu yêu cầu nhập hàng."],
    ["Xuất thực hiện hợp đồng", "Bắt buộc chọn hợp đồng thực hiện."],
    ["Điều chuyển kho nội bộ", "Kho nhập phải khác Kho xuất."],
    ["Xuất hàng gửi", "Bắt buộc chọn khách hàng."],
    ["Vận chuyển bằng công ty", "Bắt buộc nhập Số km dự kiến."],
])
b.para("Thiếu dữ liệu bắt buộc sẽ báo lỗi ngay tại trường tương ứng và không lưu cho tới khi bổ sung đủ.")

b.h2("3.6. Cảnh báo hàng quá hạn khi lưu")
b.para("Khi lưu/gửi, nếu người dùng còn hàng giữ, hàng mượn hoặc hàng nhập thẳng quá hạn, hệ thống chặn với thông báo \"Có hàng giữ/mượn/nhập thẳng quá hạn!\". Trường hợp toàn bộ hàng của phiếu là hàng đã giữ quá hạn hoặc là xuất thẳng thì được cho qua. Quản trị hệ thống (Super Admin) không bị chặn.")

# ------------------------------------------------- PHAN 4
b.h1("PHẦN 4: XEM CHI TIẾT & DUYỆT PHIẾU")

b.h2("4.1. Xem chi tiết")
b.para("Bấm Mã phiếu (hoặc mở chi tiết) để xem phiếu. Trang chi tiết hiển thị đầy đủ thông tin chung, danh sách hàng, số lượng, giá trị, trạng thái, người lập, người duyệt và ghi chú duyệt (nếu có). Cuối trang là các nút thao tác theo trạng thái và quyền.")

b.h2("4.2. Luồng duyệt Trưởng phòng → Ban giám đốc")
b.para("Phiếu xuất vượt hạn mức công nợ phải qua duyệt. Luồng:")
b.bullet("phiếu ở trạng thái Đợi TP duyệt. Người có quyền TP duyệt mở chi tiết, bấm \"TP duyệt\". Nếu phiếu vượt hạn mức tiếp, hệ thống hỏi và chuyển sang \"Đợi BGĐ duyệt\"; nếu không, phiếu chuyển Chờ duyệt (sang Kế toán kho).", bold_prefix="Trưởng phòng: ")
b.bullet("với phiếu không phải hàng mượn, thay vì duyệt thẳng, TP có thể bấm \"Chuyển BGĐ duyệt\" (bắt buộc nhập ý kiến) để đẩy lên Ban giám đốc.", bold_prefix="Chuyển BGĐ: ")
b.bullet("phiếu ở trạng thái Đợi BGĐ duyệt. Người có quyền BGĐ duyệt bấm \"BGĐ duyệt\" để duyệt.", bold_prefix="Ban giám đốc: ")
b.para("Khi duyệt, hệ thống có thể cảnh báo vượt hạn mức theo khách hàng hoặc theo nhân viên; người duyệt xác nhận để tiếp tục. Ngoài ra, nếu phòng ban của người duyệt có nhân viên đang có hàng quá hạn, việc duyệt bị chặn (cấu hình \"Duyệt yêu cầu xuất hàng\"). Quản trị hệ thống không bị chặn.")

b.h2("4.3. Không duyệt (từ chối)")
b.para("Người có quyền duyệt (hoặc được từ chối) bấm \"Không duyệt\", nhập lý do (bắt buộc, tối đa 255 ký tự). Phiếu trả về trạng thái Đang tạo để người lập chỉnh sửa; hệ thống báo cho người tạo.")

b.h2("4.4. Các nút chuyển bước tiếp")
b.table([
    ["Nút", "Điều kiện", "Kết quả"],
    ["Tạo phiếu đề nghị xuất kho", "Kế toán kho; phiếu Chờ duyệt; không xuất thẳng.", "Mở màn lập Đề nghị xuất kho, mang theo dữ liệu phiếu."],
    ["Tạo phiếu xuất hàng", "Kế toán kho; phiếu Chờ duyệt; xuất thẳng.", "Mở màn lập Phiếu xuất hàng."],
    ["Tạo yêu cầu lắp đặt", "Phiếu cần lắp đặt, chưa có yêu cầu lắp đặt, Đang tạo, do bạn lập.", "Mở màn tạo yêu cầu lắp đặt."],
])

# ------------------------------------------------- PHAN 5
b.h1("PHẦN 5: HỦY & XÓA PHIẾU")
b.h2("5.1. Hủy phiếu")
b.para("Chỉ hủy được phiếu ở trạng thái Đang tạo do chính mình lập. Sau khi hủy, phiếu chuyển sang trạng thái Đã hủy phiếu và được đánh dấu hoàn tất (không đi tiếp quy trình).")
b.h2("5.2. Xóa phiếu")
b.para("Chỉ xóa được phiếu ở trạng thái Đang tạo do chính mình lập. Xóa loại bỏ hoàn toàn phiếu (và các dòng liên quan) khỏi hệ thống — khác với Hủy (vẫn giữ phiếu ở trạng thái Đã hủy). Cân nhắc kỹ trước khi xóa.")

# ------------------------------------------------- PHAN 6
b.h1("PHẦN 6: IN ẤN & XUẤT EXCEL")
b.table([
    ["Chức năng", "Nội dung"],
    ["In yêu cầu", "In mẫu Phiếu yêu cầu xuất hàng."],
    ["In bộ giấy đi đường", "In Lệnh điều động và Phiếu xuất kho đi đường."],
    ["In biên bản giao nhận", "In biên bản giao nhận; chọn nhân viên giao nhận trước khi in."],
    ["Xuất excel danh sách hàng", "Tải file Excel danh sách hàng của phiếu."],
])

# ------------------------------------------------- PHAN 7
b.h1("PHẦN 7: MÀN DÀNH CHO NGƯỜI DUYỆT & KẾ TOÁN")
b.para("Ngoài danh sách chung, hệ thống có các màn chuyên biệt:")
b.bullet("liệt kê các phiếu đang chờ duyệt để TP/BGĐ/Kế toán kho xử lý nhanh.", bold_prefix="Chờ duyệt: ")
b.bullet("dành cho Trưởng phòng/Ban giám đốc duyệt phiếu vượt hạn mức; có thêm cột Phòng ban để lọc.", bold_prefix="Màn duyệt của TP/BGĐ: ")
b.bullet("dành cho Kế toán kho tiếp nhận phiếu đã duyệt để lập đề nghị xuất kho / phiếu xuất hàng.", bold_prefix="Màn kế toán: ")

# ------------------------------------------------- PHAN 8
b.h1("PHẦN 8: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Danh sách trống dù có nhiều phiếu", "Có thể bạn chưa có quyền xem cấp phù hợp (chỉ thấy phiếu của mình). Liên hệ quản trị để cấp quyền xem theo công ty/phòng ban/bộ phận."],
    ["Không thấy nút Sửa / Xóa", "Hai thao tác này chỉ dành cho phiếu Đang tạo do chính bạn lập."],
    ["Không thấy nút Tạo đề nghị xuất kho / Tạo phiếu xuất hàng", "Cần quyền Kế toán kho và phiếu phải ở trạng thái Chờ duyệt; xuất thẳng thì dùng \"Tạo phiếu xuất hàng\", ngược lại dùng \"Tạo đề nghị xuất kho\"."],
    ["Không lưu/gửi được, báo hàng quá hạn", "Bạn còn hàng giữ/mượn/nhập thẳng quá hạn. Xử lý các phiếu quá hạn trước, hoặc dùng đúng trường hợp được phép (hàng đã giữ quá hạn / xuất thẳng)."],
    ["Phiếu bị chuyển lên TP/BGĐ duyệt", "Do xuất vượt hạn mức công nợ. Chờ cấp có quyền duyệt xử lý; nếu bị Không duyệt, phiếu quay về Đang tạo để chỉnh sửa."],
    ["Chọn nhầm loại xuất", "Không đổi được loại sau khi tạo. Nếu chọn nhầm, hãy Xóa phiếu nháp và lập lại với loại đúng."],
    ["Tick Xuất thẳng khác gì không tick?", "Xuất thẳng bỏ qua bước đề nghị xuất kho: không cần chọn Kho xuất và cho phép Kế toán kho tạo Phiếu xuất hàng ngay khi phiếu Chờ duyệt."],
])

# ============================================================ FINISH
finish_macos(b)
print("XONG.")
