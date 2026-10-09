# -*- coding: utf-8 -*-
"""Dung HDSD man 'De nghi nhap kho' (WarehouseImportRequest / PDNNK) - ERP TanPhatDev.

Nguon: khao sat controller Warehouse\\WarehouseImportRequestsController + model
WarehouseImportRequest (+ ImportModel + WarehouseImportRequestDetail + Tab/TabProduct
+ DetailAccounting) + views warehouse/warehouse_import_requests (index/all/forWarehouse
/create/form/edit/show/history). KHONG chen anh -> placeholder.

Chay:  /usr/local/bin/python3 gen_hdsd_dnnk_erp.py
Output: ERP/HDSD_luongchinh/HDSD_DeNghiNhapKho.docx
"""
import os
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

ERP = "/Users/nguyentrancu/DEV/code/ERP-HRM/ERP"
sys.path.insert(0, os.path.join("/Users/nguyentrancu/DEV/code/ERP-HRM/HRM",
                                ".claude", "skills", "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUT_DIR = os.path.join(ERP, "HDSD_luongchinh")
os.makedirs(OUT_DIR, exist_ok=True)
OUTPUT = os.path.join(OUT_DIR, "HDSD_DeNghiNhapKho.docx")

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)


def img_placeholder(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


def clear_toc_cache(path):
    """Xoa text cache trong moi field TOC (muc luc + danh muc hinh) -> Word tu dung lai khi mo."""
    doc = Document(path)
    body = doc.element.body
    fields = []
    for el in body.iter():
        tag = el.tag
        if tag == qn('w:fldChar'):
            t = el.get(qn('w:fldCharType'))
            if t == 'begin':
                fields.append({'instr': '', 'in_result': False})
            elif t == 'separate':
                if fields:
                    fields[-1]['in_result'] = True
            elif t == 'end':
                if fields:
                    fields.pop()
        elif tag == qn('w:instrText'):
            if fields:
                fields[-1]['instr'] += (el.text or '')
        elif tag == qn('w:t'):
            if any(f['in_result'] and 'TOC' in f['instr'] for f in fields):
                el.text = ''
    doc.save(path)


def finish_macos(b):
    b.doc.save(b.output)
    clear_toc_cache(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    my_headings = {p.text.strip() for p in chk.paragraphs
                   if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in my_headings]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables),
          "| Anh:", len(chk.inline_shapes), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


# ============================================================ NOI DUNG
b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Đề nghị nhập kho)",
    doc_title="HDSD - Đề nghị nhập kho",
)

# ------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Đề nghị nhập kho (PDNNK)", "Chứng từ bước 2 của quy trình nhập kho: do Kế toán kho lập từ một Yêu cầu nhập hàng để đề nghị Thủ kho nhận và nhập hàng vào một kho cụ thể. Đây chính là màn hình tài liệu này hướng dẫn. Mã phiếu có dạng PDNNK-xxxxx."],
    ["Yêu cầu nhập hàng (YCNH)", "Chứng từ bước 1 (Product Import Request). Mỗi đề nghị nhập kho luôn được lập từ một yêu cầu nhập hàng đã duyệt; toàn bộ danh sách hàng hóa được kế thừa từ yêu cầu này."],
    ["Phiếu nhập kho (PNK)", "Chứng từ bước 3 (cuối), do Thủ kho lập từ đề nghị nhập kho để nhập hàng thực tế vào kho (tăng tồn kho)."],
    ["Thủ kho", "Người phụ trách một hoặc nhiều kho; là người nhận, duyệt (hoặc từ chối) đề nghị nhập kho và tạo phiếu nhập kho."],
    ["Kế toán kho", "Vai trò được phép lập đề nghị nhập kho từ yêu cầu nhập hàng."],
    ["Kho nhập", "Kho mà hàng sẽ được nhập vào; quyết định Thủ kho nào nhận và xử lý phiếu. Chọn kho nhập sẽ tự xác định người duyệt (thủ kho của kho đó)."],
    ["Loại nhập (type)", "Phân loại mục đích nhập (mua trong/ngoài nước, mượn/bán trả lại, điều chuyển kho, tách/ghép…); kế thừa từ yêu cầu nhập hàng, người dùng không tự chọn."],
    ["Nhà cung cấp", "Đơn vị cung cấp hàng; ở màn này chỉ là cột/bộ lọc trong danh sách, không nhập trong form."],
    ["Trạng thái hiển thị", "Chuỗi trạng thái người dùng thấy theo mốc thời gian: Đã nhận đề nghị → Đang xếp hàng → Đã hoàn thành."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "11/09/2026", "Khởi tạo HDSD màn Đề nghị nhập kho (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Màn hình Đề nghị nhập kho là bước trung gian trong quy trình nhập kho: sau khi có Yêu cầu nhập hàng đã duyệt, Kế toán kho lập đề nghị nhập kho để chỉ định kho nhập và gửi cho Thủ kho. Thủ kho nhận đề nghị, kiểm tra rồi tạo Phiếu nhập kho để nhập hàng thực tế (hoặc từ chối trả lại). Toàn bộ danh sách hàng hóa được kế thừa nguyên trạng từ yêu cầu nhập hàng — người lập KHÔNG tự thêm/xóa dòng hay sửa số lượng.")
b.bullet("menu Kho › Nhập kho › \"Đề nghị nhập kho\". Có 3 màn danh sách theo ngữ cảnh: đề nghị của tôi, tất cả (theo quyền), và màn dành cho Thủ kho.", bold_prefix="Đường dẫn danh sách: ")
b.bullet("đề nghị nhập kho được lập từ một Yêu cầu nhập hàng đã duyệt: bấm \"Tạo mới\" rồi chọn phiếu yêu cầu nhập hàng nguồn. Có thể vào thẳng từ chính phiếu yêu cầu nhập hàng.", bold_prefix="Điểm tạo phiếu: ")
b.para("Vị trí trong quy trình nhập kho (3 chứng từ nối tiếp):")
b.table([
    ["Bước", "Chứng từ", "Người thực hiện", "Ghi chú"],
    ["1", "Yêu cầu nhập hàng", "Người yêu cầu / kế toán", "Đề nghị nhập hàng, đã được duyệt."],
    ["2", "Đề nghị nhập kho", "Kế toán kho", "Màn hình này — chọn kho nhập, gửi Thủ kho."],
    ["3", "Phiếu nhập kho", "Thủ kho", "Nhập hàng thực tế vào kho (tăng tồn kho)."],
])

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Phạm vi dữ liệu nhìn thấy và các thao tác đều phụ thuộc quyền. Màn này KHÔNG chặn truy cập qua đường dẫn (route) mà kiểm soát trong xử lý nghiệp vụ. Các nhóm quyền chính:")
b.bullet("4 quyền xem theo cấp (tổng công ty / công ty / phòng ban / bộ phận) quyết định thấy đề nghị của phạm vi nào ở màn \"Tất cả\". Không có quyền nào thì chỉ thấy đề nghị do chính mình lập.", bold_prefix="Xem danh sách: ")
b.bullet("người phụ trách kho — nhận, duyệt/từ chối đề nghị và tạo phiếu nhập kho cho các kho mình quản lý. Thủ kho có màn danh sách riêng.", bold_prefix="Thủ kho: ")
b.bullet("vai trò được phép lập đề nghị nhập kho (bắt buộc để bấm Lưu/Gửi phiếu).", bold_prefix="Kế toán kho: ")
b.para("Các thao tác Sửa/Hủy phụ thuộc quyền sở hữu + trạng thái: chỉ người lập mới sửa/hủy phiếu ở trạng thái Đang tạo. Duyệt/Từ chối/Tạo phiếu nhập kho chỉ dành cho Thủ kho của kho khi phiếu ở trạng thái Chờ duyệt. Chi tiết ở PHẦN 2.")
b.para("Cảnh báo quan trọng: khi Thủ kho \"Tạo phiếu nhập kho\" từ đề nghị, hệ thống chuyển sang lập chứng từ nhập kho thực tế (tăng tồn kho). Với các loại bán/mượn trả lại, thao tác Gửi/Hủy còn ảnh hưởng đến số lượng đã xuất theo hợp đồng — cần kiểm tra kỹ trước khi thao tác.")

# ------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")

b.h2("1.1. Ba màn danh sách")
b.para("Cùng một tiêu đề \"Danh sách đề nghị nhập kho\", hệ thống có 3 màn danh sách phục vụ ngữ cảnh khác nhau (khác nhau ở phạm vi dữ liệu và một vài nút/cột):")
b.table([
    ["Màn", "Dùng cho", "Nội dung / khác biệt"],
    ["Đề nghị của tôi (index)", "Người lập (kế toán kho)", "Chỉ hiển thị đề nghị do chính mình lập. Có nút Tạo mới."],
    ["Tất cả (all)", "Quản lý / người có quyền xem", "Hiển thị đề nghị theo phạm vi phân quyền (tổng công ty / công ty / phòng ban / bộ phận). Có thêm cột \"Người lập\" và nút Xuất Excel. Ẩn phiếu nháp của người khác."],
    ["Cho Thủ kho (forWarehouse)", "Thủ kho", "Chỉ hiển thị đề nghị thuộc các kho mình phụ trách và đã gửi (ẩn phiếu Đang tạo). Không có nút Tạo mới / Xuất Excel — đây là hàng chờ xử lý."],
])
img_placeholder(b, "Ảnh menu Đề nghị nhập kho.")

b.h2("1.2. Bố cục màn hình danh sách")
b.bullet("tiêu đề \"Danh sách đề nghị nhập kho\".", bold_prefix="Thanh tiêu đề: ")
b.bullet("các ô lọc theo cột (Kho hàng, Tên/mã hàng hóa, Mã phiếu, Loại, Nhà cung cấp, Phiếu YCNH, Người yêu cầu, Trạng thái, Người lập, Người duyệt) và lọc theo khoảng thời gian.", bold_prefix="Khu vực lọc: ")
b.bullet("bảng liệt kê đề nghị theo phạm vi quyền, cột Hành động ở cuối mỗi dòng.", bold_prefix="Bảng danh sách: ")
b.bullet("nút Tạo mới (+) (màn của tôi / tất cả), nút Xuất Excel (chỉ màn Tất cả), chọn số dòng/trang và chuyển trang.", bold_prefix="Công cụ & phân trang: ")
img_placeholder(b, "Ảnh tổng quan màn danh sách đề nghị nhập kho.")

# ------------------------------------------------- PHAN 2
b.h1("PHẦN 2: DANH SÁCH ĐỀ NGHỊ NHẬP KHO")
img_placeholder(b, "Ảnh màn danh sách với dữ liệu mẫu và cột Hành động.")

b.h2("2.1. Phân quyền & hướng dẫn theo quyền")
b.para("Danh sách hiển thị theo phạm vi quyền của người đăng nhập. Bảng dưới liệt kê đầy đủ các quyền liên quan tới màn. Người dùng thường chỉ có một vài quyền — hãy đối chiếu đúng phần áp dụng cho mình.")
b.table([
    ["Tên quyền", "Cho phép làm gì", "Nút/màn tương ứng"],
    ["Xem đề nghị nhập kho theo tổng công ty", "Thấy toàn bộ đề nghị (trạng thái khác nháp) của mọi công ty ở màn Tất cả.", "Danh sách (Tất cả)"],
    ["Xem đề nghị nhập kho theo công ty", "Thấy đề nghị trong công ty của mình.", "Danh sách (Tất cả)"],
    ["Xem đề nghị nhập kho theo phòng ban", "Thấy đề nghị của các phòng ban mình quản lý (hoặc do mình lập).", "Danh sách (Tất cả)"],
    ["Xem đề nghị nhập kho theo bộ phận", "Thấy đề nghị của các bộ phận mình quản lý (hoặc do mình lập).", "Danh sách (Tất cả)"],
    ["Thủ kho", "Nhận, duyệt/từ chối đề nghị và tạo phiếu nhập kho cho kho mình phụ trách.", "Màn \"Cho Thủ kho\"; nút Không duyệt / Tạo phiếu nhập kho"],
    ["Kế toán kho", "Lập (Lưu/Gửi) đề nghị nhập kho từ yêu cầu nhập hàng.", "Nút Tạo mới / Lưu / Lưu & Gửi"],
])
b.para("Lưu ý phạm vi: ở màn Tất cả, hệ thống xét quyền xem theo thứ tự tổng công ty → công ty → phòng ban → bộ phận; nếu không có quyền nào, người dùng chỉ thấy đề nghị do chính mình lập. Phiếu nháp (Đang tạo) của người khác luôn bị ẩn — chỉ người lập mới thấy phiếu nháp của mình. Màn \"Cho Thủ kho\" thì lọc theo các kho người dùng phụ trách và chỉ hiện phiếu đã gửi.")

b.h3("Người dùng có quyền \"Xem đề nghị nhập kho theo tổng công ty\"")
b.para("Ở màn Tất cả, thấy toàn bộ đề nghị (trạng thái khác nháp) của mọi công ty, cộng phiếu do chính mình lập. Thực hiện đầy đủ thao tác xem/lọc/in/xuất Excel; sửa/hủy vẫn theo quy tắc sở hữu + trạng thái.")

b.h3("Người dùng có quyền \"Xem đề nghị nhập kho theo công ty / phòng ban / bộ phận\"")
b.para("Thấy đề nghị trong phạm vi công ty / phòng ban / bộ phận tương ứng mà mình phụ trách, cộng thêm đề nghị do chính mình lập. Cấp phòng ban dựa trên các phòng người dùng quản lý; cấp bộ phận dựa trên các bộ phận người dùng quản lý.")

b.h3("Người dùng là Thủ kho")
b.para("Vào màn \"Cho Thủ kho\" để xem danh sách đề nghị của các kho mình phụ trách (đã gửi, chờ xử lý). Với phiếu ở trạng thái Chờ duyệt, Thủ kho được: Từ chối (Không duyệt, kèm lý do) hoặc Tạo phiếu nhập kho để nhập hàng thực tế.")

b.h3("Người dùng là Kế toán kho")
b.para("Vai trò lập phiếu: bấm Tạo mới, chọn phiếu yêu cầu nhập hàng nguồn, chọn kho nhập rồi Lưu (nháp) hoặc Lưu & Gửi. Nếu tài khoản không có quyền Kế toán kho, khi bấm Lưu/Gửi hệ thống báo \"Không đủ quyền!\".")

b.h2("2.2. Tìm kiếm & lọc")
b.table([
    ["Tiêu chí", "Cách dùng"],
    ["Kho hàng", "Chọn kho để lọc đề nghị theo kho nhập (chỉ các kho người dùng được phép)."],
    ["Tên, mã hàng hóa", "Nhập để lọc phiếu có chứa hàng theo tên hoặc mã."],
    ["Mã phiếu", "Nhập mã đề nghị nhập kho (PDNNK-…)."],
    ["Loại", "Chọn loại nhập."],
    ["Nhà cung cấp", "Chọn nhà cung cấp."],
    ["Phiếu YCNH", "Nhập mã yêu cầu nhập hàng nguồn."],
    ["Người yêu cầu", "Chọn (tìm theo tên) người lập yêu cầu nhập hàng."],
    ["Trạng thái", "Chọn trạng thái phiếu."],
    ["Người lập", "Chọn (tìm theo tên) người lập đề nghị nhập kho."],
    ["Người duyệt", "Chọn (tìm theo tên) thủ kho duyệt."],
    ["Khoảng thời gian", "Lọc theo khoảng ngày lập phiếu."],
])

b.h2("2.3. Các cột trong danh sách")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Mã phiếu", "Mã đề nghị nhập kho; bấm để mở chi tiết."],
    ["Loại", "Loại nhập (kế thừa từ yêu cầu nhập hàng)."],
    ["Nhà cung cấp", "Nhà cung cấp của phiếu (nếu có)."],
    ["Phiếu YCNH", "Mã yêu cầu nhập hàng nguồn."],
    ["Người yêu cầu", "Người lập yêu cầu nhập hàng."],
    ["Ngày lập", "Ngày tạo đề nghị."],
    ["Trạng thái", "Trạng thái phiếu, hiển thị nhãn màu."],
    ["Người lập", "Người lập đề nghị (chỉ ở màn Tất cả)."],
    ["Người duyệt", "Thủ kho duyệt phiếu."],
    ["Ngày nhận", "Thời điểm Thủ kho nhận đề nghị (khi gửi)."],
    ["Ngày duyệt", "Thời điểm duyệt."],
    ["Hành động", "Các nút thao tác (xem mục 2.5)."],
])

b.h2("2.4. Bảng trạng thái phiếu")
b.para("Trong bộ lọc và nhãn hiển thị, trạng thái gồm các giá trị sau:")
b.table([
    ["Trạng thái", "Màu nhãn", "Ý nghĩa"],
    ["Đang tạo", "Đỏ", "Phiếu nháp — mới Lưu, chưa gửi (hoặc bị Thủ kho trả lại). Chỉ người lập thấy và sửa/hủy được."],
    ["Chờ duyệt", "Đỏ", "Đã Gửi, chờ Thủ kho duyệt / xử lý."],
    ["Đang nhập kho", "Xanh", "Đã được Thủ kho duyệt, đang chờ tạo/hoàn tất phiếu nhập kho."],
    ["Đã nhập kho", "Xanh", "Đã hoàn tất nhập kho (đã tăng tồn)."],
    ["Đang hạch toán", "Xanh", "Đang trong bước hạch toán."],
    ["Đã hạch toán", "Xanh", "Đã hạch toán xong."],
    ["Đã hủy", "Đỏ", "Phiếu bị hủy."],
])
b.para("Vòng đời cơ bản người dùng thấy khi thao tác: Đang tạo (Lưu nháp) → Chờ duyệt (Lưu & Gửi) → Đang nhập kho (Thủ kho duyệt) → Đã nhập kho (tạo xong phiếu nhập kho). Ngoài ra, ở màn dành cho Thủ kho hệ thống còn hiển thị nhãn động theo mốc thời gian: \"Đã nhận đề nghị\" → \"Đang xếp hàng\" → \"Đã hoàn thành\".")

b.h2("2.5. Thao tác trên từng dòng (cột Hành động)")
b.para("Các nút nằm trong menu thao tác và hiển thị tùy trạng thái, loại phiếu và quyền:")
b.table([
    ["Nút", "Điều kiện hiển thị", "Kết quả"],
    ["Sửa đề nghị", "Phiếu Đang tạo do chính bạn lập.", "Mở màn sửa đề nghị nhập kho."],
    ["Hủy đề nghị", "Phiếu Đang tạo do chính bạn lập, và loại KHÔNG thuộc nhóm điều chuyển kho / nhập tách / nhập nước ngoài (mới).", "Hủy phiếu (chuyển sang Đã hủy)."],
    ["Tạo phiếu nhập kho", "Bạn là Thủ kho của kho đó và phiếu ở trạng thái Chờ duyệt.", "Chuyển sang màn lập Phiếu nhập kho, nạp sẵn dữ liệu từ đề nghị."],
    ["In đề nghị", "Luôn có.", "In mẫu Đề nghị nhập kho."],
    ["Lịch sử phiếu", "Luôn có.", "Xem lịch sử thay đổi của phiếu."],
])
b.para("Ngoài ra danh sách (màn của tôi / Tất cả) có nút Tạo mới (+) mở màn lập phiếu, và nút Xuất Excel (chỉ ở màn Tất cả).")

# ------------------------------------------------- PHAN 3
b.h1("PHẦN 3: LẬP ĐỀ NGHỊ NHẬP KHO")
b.para("Đề nghị nhập kho do Kế toán kho lập từ một Yêu cầu nhập hàng đã duyệt. Bấm Tạo mới ở màn danh sách (hoặc vào thẳng từ phiếu yêu cầu nhập hàng). Loại nhập được suy ra từ yêu cầu nhập hàng, người dùng KHÔNG tự chọn.")
img_placeholder(b, "Ảnh màn lập đề nghị nhập kho.")

b.h2("3.1. Nguồn tạo phiếu")
b.para("Mỗi đề nghị nhập kho luôn gắn với một phiếu nguồn:")
b.bullet("phiếu nguồn là Yêu cầu nhập hàng — luồng chuẩn. Toàn bộ hàng hóa được kế thừa từ yêu cầu này.", bold_prefix="Chọn phiếu yêu cầu nhập hàng: ")
b.bullet("với loại Nhập tách, phiếu nguồn là Yêu cầu xuất tách (chỉ hiển thị, không chọn được).", bold_prefix="Phiếu yêu cầu xuất tách: ")
b.para("Khi phiếu đã lưu (đã có mã), ô chọn phiếu nguồn bị khóa, không đổi được. Nếu chưa chọn phiếu nguồn, bảng hàng hóa hiển thị dòng nhắc \"Chưa chọn phiếu yêu cầu\".")

b.h2("3.2. Cách chọn phiếu Yêu cầu nhập hàng")
b.para("Tại trường \"Chọn phiếu yêu cầu nhập hàng\", bấm nút kính lúp để mở cửa sổ \"Phiếu yêu cầu nhập hàng\". Cửa sổ liệt kê các yêu cầu nhập hàng đã duyệt (đủ điều kiện lập đề nghị):")
b.table([
    ["Thành phần", "Nội dung"],
    ["Cột danh sách", "STT, Mã phiếu, Người tạo, Ngày tạo."],
    ["Bộ lọc", "Mã phiếu (nhập), Người tạo (chọn)."],
    ["Chọn phiếu", "Bấm chọn 1 dòng → hệ thống nạp thông tin + danh sách hàng hóa của yêu cầu, rồi đóng cửa sổ."],
])
b.para("Nút X bên cạnh dùng để bỏ lựa chọn phiếu nguồn (chỉ khi đang tạo mới, chưa lưu).")

b.h2("3.3. Các trường trong form (Thông tin chung)")
b.table([
    ["Trường", "Bắt buộc", "Giá trị điền sẵn / ghi chú"],
    ["Chọn phiếu yêu cầu nhập hàng", "Có", "Hiển thị mã yêu cầu nhập hàng; khóa khi phiếu đã lưu. (Ẩn với loại Nhập tách — thay bằng \"Phiếu yêu cầu xuất tách\".)"],
    ["Người yêu cầu", "—", "Chỉ đọc — tự điền người lập yêu cầu nhập hàng nguồn (hiện khi đã chọn phiếu nguồn)."],
    ["Chọn kho nhập", "Có", "Chọn kho sẽ nhập hàng. Chọn kho sẽ tự xác định người duyệt = Thủ kho của kho đó. Có thể được điền sẵn theo phiếu nguồn."],
    ["Ghi chú", "Không", "Tối đa 255 ký tự; mặc định lấy ghi chú của yêu cầu nhập hàng nguồn."],
])
b.para("Header phiếu hiển thị tên loại nhập và \"Người lập - Ngày lập\" (khi tạo mới: người đang đăng nhập và ngày hôm nay). Màn này KHÔNG có trường Nhà cung cấp, KHÔNG có đính kèm file, KHÔNG có ô nhập số lượng — mọi thông tin hàng hóa lấy sẵn từ yêu cầu nhập hàng.")

b.h2("3.4. Bảng hàng hóa (chỉ đọc)")
b.para("Bảng hàng hóa được kế thừa nguyên trạng từ yêu cầu nhập hàng: người lập KHÔNG thêm/xóa dòng và KHÔNG sửa số lượng. Bảng chỉ hiện khi đã chọn phiếu nguồn.")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Tên hàng hóa", "Tên sản phẩm."],
    ["Model", "Model sản phẩm."],
    ["Mã hàng hóa", "Mã sản phẩm."],
    ["Thương hiệu", "Thương hiệu."],
    ["Đơn vị tính", "Đơn vị tính."],
    ["SL đề nghị", "Số lượng đề nghị nhập (lấy từ yêu cầu nhập hàng, chỉ hiển thị)."],
])
b.para("Với hàng theo hợp đồng, ngoài tab \"Hàng hóa\" mặc định còn có các tab theo hợp đồng/nhóm (tên tab lấy từ phiếu nguồn); các tab này cũng chỉ đọc.")
img_placeholder(b, "Ảnh bảng hàng hóa (chỉ đọc) trong form đề nghị nhập kho.")

b.h2("3.5. Các nút lưu")
b.table([
    ["Nút", "Hành động"],
    ["Lưu", "Lưu phiếu ở trạng thái Đang tạo (nháp) — chưa gửi cho Thủ kho. Hệ thống nhắc \"Đề nghị của bạn đã được lưu. Bạn cần gửi để đề nghị được xử lý\"."],
    ["Lưu & Gửi", "Gửi phiếu cho Thủ kho (trạng thái Chờ duyệt). Thủ kho nhận được thông báo. Hệ thống báo \"Đề nghị của bạn đã được gửi\"."],
    ["Hủy", "Quay lại danh sách, không lưu."],
])
b.para("Cần có quyền Kế toán kho để lưu/gửi. Nếu chọn kho không hợp lệ (kho không được phép thao tác), hệ thống báo \"Không thể thao tác trên kho này!\".")

b.h2("3.6. Quy tắc kiểm tra dữ liệu (validate)")
b.table([
    ["Trường / tình huống", "Yêu cầu / thông báo"],
    ["Phiếu yêu cầu nhập hàng", "Bắt buộc chọn (\"Bắt buộc phải chọn\")."],
    ["Kho nhập", "Bắt buộc chọn (\"Bắt buộc phải nhập\")."],
    ["Ghi chú", "Tối đa 255 ký tự (\"Không được vượt quá 255 ký tự\")."],
    ["Quyền lập phiếu", "Phải có quyền Kế toán kho, nếu không: \"Không đủ quyền!\"."],
    ["Kho thao tác", "Kho phải hợp lệ, nếu không: \"Không thể thao tác trên kho này!\"."],
])
b.para("Thiếu dữ liệu bắt buộc sẽ báo lỗi ngay tại ô tương ứng (viền đỏ) và không lưu cho tới khi sửa đủ.")

# ------------------------------------------------- PHAN 4
b.h1("PHẦN 4: XEM CHI TIẾT & XỬ LÝ CỦA THỦ KHO")
b.para("Bấm Mã phiếu để xem chi tiết. Trang chi tiết chỉ đọc, gồm 3 khối: Thông tin chung, Chi tiết (bảng hàng hóa + tab hợp đồng) và Ghi chú duyệt (chỉ hiện khi phiếu bị từ chối/ghi chú).")
b.bullet("Phiếu yêu cầu nhập hàng (link mở phiếu gốc), Người yêu cầu, Ngày lập, Người lập, Kho nhập, Ghi chú — đều chỉ đọc.", bold_prefix="Thông tin chung: ")
b.bullet("bảng hàng hóa và các tab hợp đồng giống form lập, chỉ để xem.", bold_prefix="Chi tiết: ")
b.bullet("hiển thị lý do khi Thủ kho không duyệt (nếu có).", bold_prefix="Ghi chú duyệt: ")
b.para("Các nút cuối màn chi tiết phụ thuộc quyền và trạng thái:")
b.table([
    ["Nút", "Điều kiện", "Kết quả"],
    ["Không duyệt", "Bạn là Thủ kho của kho, phiếu Chờ duyệt, loại không phải nhập nước ngoài (mới).", "Mở cửa sổ \"Ghi chú duyệt\" để nhập lý do rồi từ chối."],
    ["Tạo phiếu nhập kho", "Bạn là Thủ kho của kho và phiếu Chờ duyệt.", "Chuyển sang lập Phiếu nhập kho thực tế."],
    ["Quay lại", "Luôn có.", "Về màn danh sách."],
])
img_placeholder(b, "Ảnh màn xem chi tiết đề nghị nhập kho với các nút xử lý của Thủ kho.")

b.h2("4.1. Không duyệt (từ chối) — dành cho Thủ kho")
b.para("Khi Thủ kho bấm \"Không duyệt\", hệ thống mở cửa sổ \"Ghi chú duyệt\":")
b.table([
    ["Trường", "Bắt buộc", "Ghi chú"],
    ["Ghi chú (lý do)", "Có", "Nhập lý do từ chối; tối đa 255 ký tự (\"Bắt buộc phải nhập\" / \"Không được vượt quá 255 ký tự\")."],
])
b.para("Bấm Xác nhận: phiếu được trả về trạng thái Đang tạo để người lập sửa lại; người lập nhận thông báo \"… vừa từ chối đề nghị nhập kho: <mã phiếu>\". Bấm Hủy để đóng cửa sổ, không làm gì. Yêu cầu quyền Thủ kho của kho; nếu không đủ quyền: \"Không đủ quyền!\".")

b.h2("4.2. Tạo phiếu nhập kho — dành cho Thủ kho")
b.para("Với phiếu ở trạng thái Chờ duyệt, Thủ kho bấm \"Tạo phiếu nhập kho\" để chuyển sang màn lập Phiếu nhập kho. Hệ thống nạp sẵn danh sách hàng hóa và thông tin từ đề nghị. Đây là bước cuối làm tăng tồn kho thực tế (xem tài liệu HDSD Phiếu nhập kho). Nếu bạn không phải Thủ kho của kho hoặc phiếu không ở trạng thái Chờ duyệt, hệ thống báo \"Không thể chọn phiếu này\".")

# ------------------------------------------------- PHAN 5
b.h1("PHẦN 5: SỬA & HỦY PHIẾU")
b.h2("5.1. Sửa phiếu")
b.para("Chỉ sửa được phiếu ở trạng thái Đang tạo (nháp) do chính mình lập. Mở phiếu và bấm \"Sửa đề nghị\". Khi sửa, không đổi được phiếu nguồn (đã khóa). Có thể cập nhật: kho nhập và ghi chú. Bấm \"Lưu\" để giữ nháp hoặc \"Lưu & Gửi\" để gửi Thủ kho. Nếu không đủ quyền sửa, hệ thống báo \"Bạn không có quyền sửa đề nghị này\".")
b.h2("5.2. Hủy phiếu")
b.para("Chỉ hủy được phiếu ở trạng thái Đang tạo do chính mình lập, và loại nhập KHÔNG thuộc nhóm điều chuyển kho (nội bộ / cùng thủ kho / chi nhánh), nhập tách, hoặc nhập nước ngoài (mới). Bấm \"Hủy đề nghị\" ở danh sách: nếu hợp lệ, phiếu chuyển sang Đã hủy (\"Hủy thành công!\"); ngược lại báo \"Không thể hủy!\". Với loại bán/mượn trả lại, hủy sẽ hoàn lại số lượng đã xuất theo hợp đồng.")

# ------------------------------------------------- PHAN 6
b.h1("PHẦN 6: LỊCH SỬ, IN ẤN & XUẤT EXCEL")
b.h2("6.1. Lịch sử phiếu")
b.para("Bấm \"Lịch sử phiếu\" để xem timeline thay đổi theo từng phiên bản, mỗi mốc có người thay đổi + thời gian. Mỗi mốc là một bảng: STT | Thông tin | Giá trị. Các mốc được theo dõi:")
b.table([
    ["Thông tin", "Ý nghĩa"],
    ["Thời gian duyệt", "Thời điểm Thủ kho duyệt."],
    ["Ghi chú", "Ghi chú của phiếu."],
    ["Ghi chú duyệt", "Lý do khi từ chối/duyệt."],
    ["Thời gian gửi", "Thời điểm gửi (Thủ kho nhận)."],
    ["Thời gian hoàn thành", "Thời điểm hoàn tất nhập/xếp hàng."],
    ["Thời gian trả lại", "Thời điểm bị trả lại (từ chối)."],
])
b.para("Nếu chưa có thay đổi nào, màn hiển thị \"Chưa có dữ liệu\".")

b.h2("6.2. In ấn & Xuất Excel")
b.table([
    ["Chức năng", "Nội dung"],
    ["In đề nghị", "In mẫu Đề nghị nhập kho (thông tin phiếu + bảng hàng hóa: Tên hàng hóa, Model, Mã hàng hóa, Thương hiệu, Số lượng, Đơn vị tính)."],
    ["Xuất Excel danh sách", "Tải file Excel danh sách đề nghị theo bộ lọc & phạm vi hiện tại (chỉ ở màn Tất cả). Danh sách quá lớn (≥ 2000 dòng) sẽ được gửi qua email."],
])

# ------------------------------------------------- PHAN 7
b.h1("PHẦN 7: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Lập đề nghị nhập kho ở đâu?", "Bấm Tạo mới ở màn danh sách rồi chọn một Yêu cầu nhập hàng đã duyệt; hoặc vào thẳng từ phiếu yêu cầu nhập hàng. Cần quyền Kế toán kho."],
    ["Vì sao không chọn được loại nhập / số lượng?", "Loại nhập và toàn bộ hàng hóa (kể cả số lượng) được kế thừa từ yêu cầu nhập hàng, không sửa tại màn này."],
    ["Không thấy nút Tạo mới / Xuất Excel", "Nút Tạo mới có ở màn \"của tôi\" và \"Tất cả\"; nút Xuất Excel chỉ có ở màn \"Tất cả\". Màn \"Cho Thủ kho\" không có hai nút này."],
    ["Báo \"Không đủ quyền!\" khi Lưu/Gửi", "Tài khoản chưa có quyền Kế toán kho. Liên hệ quản trị để được cấp quyền."],
    ["Báo \"Không thể thao tác trên kho này!\"", "Kho nhập đã chọn không hợp lệ để thao tác. Chọn lại kho phù hợp."],
    ["Phiếu bị Thủ kho từ chối thì sao?", "Phiếu quay về trạng thái Đang tạo kèm ghi chú duyệt; người lập sửa lại (kho/ghi chú) rồi Gửi lại."],
    ["Không hủy được phiếu", "Chỉ hủy được phiếu Đang tạo do chính mình lập và không thuộc nhóm điều chuyển kho / nhập tách / nhập nước ngoài (mới)."],
    ["Không thấy nút Tạo phiếu nhập kho / Không duyệt", "Hai nút này chỉ hiện với Thủ kho của kho khi phiếu ở trạng thái Chờ duyệt (Không duyệt còn yêu cầu loại không phải nhập nước ngoài mới)."],
    ["Danh sách trống dù có nhiều phiếu", "Có thể bạn chưa có quyền xem cấp phù hợp (chỉ thấy phiếu của mình). Liên hệ quản trị để cấp quyền xem theo công ty/phòng ban/bộ phận."],
])

# ============================================================ FINISH
finish_macos(b)
print("XONG.")
