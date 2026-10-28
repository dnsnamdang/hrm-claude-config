# -*- coding: utf-8 -*-
"""Dung HDSD man 'Cau hinh chan luong hang qua han' (ERP, nhanh gop_db).

Bam khuon file mau HDSD_YeuCauXuatHang.docx (user chi dinh 14/10/2026) = khuon chi tiet
HdsdBuilder cua skill hrm/.claude/skills/hdsd-documenter (bia, muc luc, danh muc hinh,
TONG QUAN 4 muc, cac PHAN co bang). KHONG chen anh -> o "[Vi tri chen anh]".

Nguon khao sat (doc truc tiep code):
  ERP routes/web.php :6945-6949, Common/DueConfigsController.php
  ERP Middleware/CheckDueConfigs.php (tab 1), CheckDueConfigsManager.php +
      Services/DueConfig/DueConfigBlockService.php (tab 2), Api/DueConfigController.php
  ERP views common/due_configs/{edit,overdue-modal}.blade.php, layouts/topmenubar.blade.php :2001
  HRM hrm-api Middleware/CheckDueConfigsManager.php, Modules/Finance CheckDueConfig,
      cac route chan (Timesheet, Assign, PaymentBusinessRequest, PaymentProfile, SettlementContract)

Chay:  python gen_hdsd.py   (Windows, can Word de cap nhat muc luc)
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "..", "..", "hrm", ".claude", "skills",
                                "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUTPUT = os.path.join(HERE, "HDSD_CauHinhChanLuongHangQuaHan.docx")
SHOTS = os.path.join(HERE, "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)

MENU = "menu QTTT › Cấu hình ERP › \"Cấu hình chặn luồng hàng quá hạn\""


def img(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Cấu hình chặn luồng hàng quá hạn)",
    doc_title="HDSD - Cấu hình chặn luồng hàng quá hạn",
)

# ------------------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Cấu hình chặn luồng hàng quá hạn", "Màn hình tài liệu này hướng dẫn: chọn, cho từng công ty, những thao tác nào bị chặn khi có hàng quá hạn."],
    ["Hàng giữ quá hạn", "Hàng nhân viên đang giữ (số lượng giữ còn lớn hơn 0) mà ngày hết hạn giữ đã qua (trước ngày hôm nay)."],
    ["Hàng mượn quá hạn", "Hàng xuất mượn trên phiếu yêu cầu xuất hàng do nhân viên lập, đã mượn, quá ngày hẹn trả mà chưa trả hết."],
    ["Hàng nhập thẳng quá hạn (NXT)", "Hàng nhập thẳng của nhân viên còn chưa xuất hết, đã quá số ngày ở ô \"Hạn xuất hàng nhập thẳng (Số ngày)\" trong quy chế công ty."],
    ["Chặn nhân viên (Tab 1)", "Chặn chính người đang thao tác nếu BẢN THÂN người đó có hàng quá hạn."],
    ["Chặn trưởng phòng duyệt (Tab 2)", "Chặn trưởng phòng duyệt / tạo phiếu nếu PHÒNG BAN mình quản lý có nhân viên quá hạn."],
    ["Trưởng phòng", "Người được gán quản lý phòng ban trong công ty của mình."],
    ["TP / BGĐ", "Trưởng phòng / Ban giám đốc."],
    ["HRM", "Hệ thống quản trị nhân sự dùng chung cấu hình này cho các phiếu chấm công, giao việc, thanh toán."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "14/10/2026", "Khởi tạo HDSD màn Cấu hình chặn luồng hàng quá hạn (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Màn hình Cấu hình chặn luồng hàng quá hạn dùng để bật / tắt việc chặn các thao tác nghiệp vụ "
       "(lập báo giá, lập hợp đồng, lập yêu cầu xuất / nhập hàng, duyệt phiếu…) khi nhân viên hoặc "
       "phòng ban còn hàng giữ, hàng mượn, hàng nhập thẳng quá hạn. Cấu hình làm riêng cho từng công ty "
       "và có hiệu lực ngay sau khi Lưu, trên cả ERP và HRM.")
b.bullet(MENU + ".", bold_prefix="Đường dẫn: ")
b.bullet("mỗi lần Lưu có thay đổi được ghi vào bảng Lịch sử chỉnh sửa ở cuối trang.",
         bold_prefix="Lịch sử: ")
b.para("Hai tab cấu hình và đối tượng bị chặn:")
b.table([
    ["Tab", "Ai bị chặn", "Khi nào bị chặn"],
    ["1. Chặn nhân viên", "Người đang thao tác", "Chính người đó có hàng quá hạn thuộc nhóm đang được tích"],
    ["2. Chặn trưởng phòng duyệt", "Trưởng phòng", "Phòng ban mình quản lý có nhân viên quá hạn bất kỳ loại nào (giữ / mượn / nhập thẳng)"],
])

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Quyền liên quan tới màn hình:")
b.bullet("cần quyền \"Cấu hình hệ thống\" để mở màn và Lưu cấu hình. Tài khoản Super Admin luôn vào được.",
         bold_prefix="Vào màn & Lưu: ")
b.bullet("cấu hình của công ty nào chỉ áp cho nhân viên / trưởng phòng THUỘC công ty đó (theo hồ sơ nhân viên).",
         bold_prefix="Phạm vi áp dụng: ")
b.bullet("tài khoản Super Admin không bao giờ bị chặn dù cấu hình thế nào.", bold_prefix="Ngoại lệ: ")
b.para("Người không có quyền vẫn thấy mục menu, nhưng khi bấm vào hệ thống hiện trang \"Không có quyền "
       "truy cập\".")

# ------------------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")
b.h2("1.1. Cách truy cập")
b.para("Vào " + MENU + ". Màn hình mở ở tab \"1. Chặn nhân viên\", tab con công ty đầu tiên (xếp theo "
       "tên A-Z).")
img(b, "Ảnh menu QTTT › Cấu hình ERP › Cấu hình chặn luồng hàng quá hạn.")

b.h2("1.2. Bố cục màn hình")
b.bullet("tiêu đề \"Chi tiết chặn (*)\", góc phải có nút \"Ẩn tất cả\" / \"Hiện tất cả\".",
         bold_prefix="Khung trên: ")
b.bullet("\"1. Chặn nhân viên\" và \"2. Chặn trưởng phòng duyệt\".", bold_prefix="Hai tab chính: ")
b.bullet("dưới mỗi tab chính là thanh tab con, mỗi tab con là một công ty đang hoạt động.",
         bold_prefix="Tab con công ty: ")
b.bullet("các ô tích thao tác; tab 1 chia 3 nhóm theo loại hàng quá hạn, tab 2 là danh sách phẳng.",
         bold_prefix="Vùng ô tích: ")
b.bullet("nút Lưu (xanh lá) và Hủy (đỏ) ở chân khung.", bold_prefix="Nút thao tác: ")
b.bullet("khung \"Lịch sử chỉnh sửa\" ở cuối trang.", bold_prefix="Lịch sử: ")
img(b, "Ảnh tổng quan màn Cấu hình chặn luồng hàng quá hạn (tab 1).")

# ------------------------------------------------------------- PHAN 2
b.h1("PHẦN 2: CẤU HÌNH CHẶN NHÂN VIÊN (TAB 1)")
b.para("Tab 1 chặn chính người đang thao tác nếu bản thân họ có hàng quá hạn. Mỗi công ty có 21 ô tích "
       "chia 3 nhóm; tên nhóm chuyển màu xanh và hiện số ô đang tích khi có ít nhất một ô được tích.")
img(b, "Ảnh tab 1 với 3 nhóm Hàng mượn quá hạn, Hàng giữ quá hạn, Hàng nhập thẳng quá hạn.")

b.h2("2.1. Ba nhóm loại hàng quá hạn")
b.table([
    ["Nhóm", "Số ô", "Cách xác định quá hạn"],
    ["Hàng mượn quá hạn", "6", "Phiếu xuất mượn do chính nhân viên lập, đã mượn, ngày hẹn trả trước hôm nay, chưa trả hết."],
    ["Hàng giữ quá hạn", "6", "Hàng giữ của nhân viên còn số lượng, ngày hết hạn giữ trước hôm nay."],
    ["Hàng nhập thẳng quá hạn", "9", "Hàng nhập thẳng của nhân viên còn tồn, ngày nhập + \"Hạn xuất hàng nhập thẳng (Số ngày)\" của công ty đã qua. Công ty để trống số ngày thì mọi hàng nhập thẳng còn tồn đều tính là quá hạn."],
])

b.h2("2.2. Các thao tác có thể chặn")
b.table([
    ["Ô tích", "Có ở nhóm", "Thao tác bị chặn"],
    ["Lập báo giá", "Cả 3 nhóm", "Tạo / lưu Báo giá và Báo giá dịch vụ bảo hành - sửa chữa"],
    ["Lập hợp đồng", "Cả 3 nhóm", "Tạo / lưu Hợp đồng bán và Hợp đồng dịch vụ bảo hành - sửa chữa"],
    ["Lập yêu cầu xuất hàng", "Cả 3 nhóm", "Lưu phiếu Yêu cầu xuất hàng (tạo mới và sửa)"],
    ["Lập yêu cầu nhập hàng", "Cả 3 nhóm", "Lưu phiếu Yêu cầu nhập hàng; nhóm Hàng giữ và Hàng nhập thẳng chặn ngay khi bấm Tạo mới"],
    ["Cập nhật hàng hóa", "Cả 3 nhóm", "Thêm, sửa, sao chép, xóa, khôi phục hàng hóa và cập nhật danh sách hàng"],
    ["Duyệt hàng tạm", "Cả 3 nhóm", "Sửa, duyệt, từ chối hàng tạm"],
    ["Lập yêu cầu xuất giữ", "Hàng nhập thẳng", "Tạo / lưu phiếu Yêu cầu xuất giữ"],
    ["Gia hạn hàng giữ", "Hàng nhập thẳng", "Tạo / lưu phiếu Yêu cầu gia hạn hàng giữ"],
    ["Điều chuyển hàng giữ", "Hàng nhập thẳng", "Tạo / lưu phiếu Điều chuyển hàng giữ"],
])

b.h2("2.3. Các trường hợp được cho qua")
b.para("Để nhân viên vẫn xử lý được hàng quá hạn của mình, hệ thống KHÔNG chặn các trường hợp sau:")
b.bullet("phiếu Yêu cầu xuất hàng có tích \"Xuất thẳng\" (nhóm Hàng giữ và Hàng nhập thẳng).",
         bold_prefix="Xuất thẳng: ")
b.bullet("phiếu Yêu cầu xuất hàng chỉ gồm đúng những hàng đang giữ quá hạn của chính người lập.",
         bold_prefix="Giải phóng hàng giữ: ")
b.bullet("phiếu Yêu cầu nhập hàng loại \"Mượn trả lại\" không bị chặn bởi nhóm Hàng mượn quá hạn.",
         bold_prefix="Trả hàng mượn: ")

b.h2("2.4. Người dùng thấy gì khi bị chặn")
b.para("Thao tác không thực hiện được, góc màn hình hiện thông báo màu vàng:")
b.table([
    ["Loại quá hạn", "Thông báo"],
    ["Hàng giữ quá hạn", "Có hàng giữ quá hạn!"],
    ["Hàng mượn quá hạn", "Có hàng mượn quá hạn!"],
    ["Hàng nhập thẳng quá hạn", "Có hàng nhập thẳng quá hạn!"],
])
b.para("Khi người dùng dính nhiều loại cùng lúc, hệ thống xét theo thứ tự Hàng giữ → Hàng mượn → Hàng "
       "nhập thẳng và chỉ hiện một thông báo của loại đầu tiên bị chặn. Tab 1 chỉ xét hàng của chính "
       "người thao tác, không xét đồng nghiệp cùng phòng.")

# ------------------------------------------------------------- PHAN 3
b.h1("PHẦN 3: CẤU HÌNH CHẶN TRƯỞNG PHÒNG DUYỆT (TAB 2)")
b.para("Tab 2 chặn trưởng phòng thực hiện các thao tác duyệt / tạo phiếu nếu phòng ban do mình quản lý "
       "có nhân viên quá hạn hàng mượn / hàng giữ / hàng NXT. Mỗi công ty có 31 ô tích, kèm liên kết "
       "\"Tích tất cả | Bỏ tích tất cả\".")
img(b, "Ảnh tab 2 Chặn trưởng phòng duyệt với dòng chú thích và liên kết Tích tất cả | Bỏ tích tất cả.")

b.h2("3.1. Phạm vi xét của trưởng phòng")
b.bullet("tất cả nhân viên thuộc các phòng ban trưởng phòng được gán quản lý trong công ty của mình, "
         "kể cả chính trưởng phòng.", bold_prefix="Nhân viên được xét: ")
b.bullet("chỉ cần một nhân viên quá hạn BẤT KỲ loại nào (giữ, mượn, nhập thẳng) là chặn; tab 2 không "
         "chia nhóm.", bold_prefix="Loại quá hạn: ")
b.bullet("người không được gán quản lý phòng ban nào thì tab 2 không chặn.", bold_prefix="Không quản lý phòng: ")

b.h2("3.2. Các thao tác có thể chặn")
b.table([
    ["Ô tích", "Hệ thống", "Thao tác bị chặn"],
    ["Duyệt yêu cầu xuất giữ", "ERP, HRM", "Duyệt phiếu Yêu cầu xuất giữ"],
    ["Duyệt gia hạn hàng giữ", "ERP", "Duyệt phiếu Yêu cầu gia hạn hàng giữ"],
    ["Duyệt điều chuyển hàng giữ", "ERP", "Duyệt phiếu Điều chuyển hàng giữ"],
    ["Duyệt gia hạn hàng mượn", "ERP, HRM", "Duyệt phiếu Yêu cầu gia hạn hàng mượn"],
    ["Duyệt yêu cầu xuất hàng", "ERP", "TP duyệt và Chuyển BGĐ duyệt phiếu Yêu cầu xuất hàng"],
    ["Duyệt yêu cầu xuất bán hàng mượn", "ERP", "TP duyệt phiếu Yêu cầu xuất bán hàng mượn"],
    ["Duyệt yêu cầu nhập hàng", "ERP", "Các bước duyệt của TP / ban kiểm soát / BGĐ trên phiếu Yêu cầu nhập hàng"],
    ["Duyệt hợp đồng bán", "ERP", "Duyệt Hợp đồng bán, phụ lục hợp đồng, hợp đồng nguyên tắc"],
    ["Duyệt hợp đồng mua", "ERP", "Duyệt hợp đồng mua trong nước, nước ngoài, nguyên tắc, mua dịch vụ"],
    ["Duyệt yêu cầu đặt hàng", "ERP", "Duyệt yêu cầu đặt hàng (phòng ban)"],
    ["Duyệt đề nghị thanh toán", "ERP", "Lập / sửa phiếu Đề nghị thanh toán"],
    ["Duyệt yêu cầu lắp đặt bàn giao", "ERP", "Duyệt yêu cầu lắp đặt bàn giao"],
    ["Duyệt giá vận chuyển quốc tế", "ERP", "Duyệt giá vận chuyển quốc tế"],
    ["Duyệt yêu cầu mua dịch vụ", "ERP, HRM", "Duyệt yêu cầu mua dịch vụ (phòng ban)"],
    ["Duyệt quyết toán hợp đồng", "ERP", "Duyệt quyết toán hợp đồng"],
    ["Duyệt kết quả công việc", "ERP", "Lưu kết quả duyệt công việc bảo hành - sửa chữa"],
    ["Tạo phiếu giao việc (phân hệ chấm công)", "HRM", "Tạo phiếu giao việc"],
    ["Tạo phiếu đi công tác (phân hệ chấm công)", "HRM", "Tạo phiếu đi công tác"],
    ["Duyệt đơn xin nghỉ", "HRM", "Duyệt đơn xin nghỉ"],
    ["Duyệt đề nghị tra soát công", "HRM", "Duyệt đề nghị tra soát công"],
    ["Duyệt đăng ký làm thêm", "HRM", "Duyệt đăng ký làm thêm"],
    ["Duyệt đơn đi muộn về sớm", "HRM", "Duyệt đơn đi muộn về sớm"],
    ["Duyệt phiếu yêu cầu làm thêm", "HRM", "Duyệt phiếu yêu cầu làm thêm"],
    ["Duyệt phiếu giao việc (phân hệ chấm công)", "HRM", "Duyệt phiếu giao việc"],
    ["Duyệt phiếu đi công tác (phân hệ chấm công)", "HRM", "Duyệt phiếu đi công tác"],
    ["Tạo phiếu công tác (phân hệ giao việc)", "HRM", "Tạo phiếu công tác"],
    ["Chặn gia hạn và kết thúc phiếu công tác (phân hệ giao việc)", "HRM", "Duyệt yêu cầu gia hạn / kết thúc sớm phiếu công tác"],
    ["Duyệt đề nghị thanh toán HRM", "HRM", "Duyệt đề nghị thanh toán công tác"],
    ["Duyệt hồ sơ thanh toán", "HRM", "Duyệt hồ sơ thanh toán"],
    ["Duyệt quyết toán công theo HĐ", "HRM", "Duyệt quyết toán công theo hợp đồng"],
    ["Duyệt kết quả công tác trên HRM", "HRM", "Duyệt kết quả công tác"],
])

b.h2("3.3. Người dùng thấy gì khi bị chặn")
b.para("Thao tác không thực hiện được. Thay vì thông báo chữ, hệ thống (cả ERP và HRM) mở cửa sổ "
       "\"Phòng ban bạn quản lý có NV hàng quá hạn như sau:\" gồm:")
b.table([
    ["Thành phần", "Nội dung"],
    ["Cột STT", "Số thứ tự"],
    ["Cột Nhân viên", "Mã - tên nhân viên quá hạn"],
    ["Cột Loại quá hạn", "Hàng mượn, Hàng NXT, Hàng giữ (một nhân viên có thể nhiều loại)"],
    ["Nút Xuất Excel", "Tải danh sách về file Excel để nhắc nhân viên xử lý"],
])
b.para("Cửa sổ liệt kê mọi nhân viên quá hạn của phòng, mọi loại quá hạn. Trưởng phòng nhắc nhân viên "
       "xử lý xong hàng quá hạn rồi thực hiện lại thao tác.")
img(b, "Ảnh cửa sổ Phòng ban bạn quản lý có NV hàng quá hạn như sau.")

# ------------------------------------------------------------- PHAN 4
b.h1("PHẦN 4: THAO TÁC CẤU HÌNH & LƯU")
b.h2("4.1. Tích / bỏ tích theo công ty")
b.bullet("chọn tab chính (1 hoặc 2), rồi bấm tab con công ty cần cấu hình.", bold_prefix="Bước 1: ")
b.bullet("tích ô thao tác muốn chặn, bỏ tích ô muốn thôi chặn (bấm vào ô hoặc vào chữ của ô).",
         bold_prefix="Bước 2: ")
b.bullet("lặp lại cho công ty khác nếu cần; có thể đổi nhiều công ty, cả 2 tab rồi Lưu một lần.",
         bold_prefix="Bước 3: ")
b.bullet("bấm Lưu.", bold_prefix="Bước 4: ")
b.para("Mỗi công ty cấu hình độc lập; cùng tên thao tác ở các nhóm khác nhau (ví dụ \"Lập báo giá\" của "
       "Hàng mượn và của Hàng giữ) là các ô riêng.")

b.h2("4.2. Tích tất cả / Bỏ tích tất cả (tab 2)")
b.para("Ở tab 2, mỗi tab con công ty có liên kết \"Tích tất cả | Bỏ tích tất cả\": tích hoặc bỏ tích "
       "cùng lúc 31 ô của công ty đang xem, không ảnh hưởng công ty khác và tab 1. Vẫn phải bấm Lưu để "
       "ghi nhận.")

b.h2("4.3. Thu gọn / mở rộng nhóm (tab 1)")
b.bullet("bấm vào tên nhóm để thu gọn hoặc mở rộng riêng nhóm đó.", bold_prefix="Một nhóm: ")
b.bullet("bấm \"Ẩn tất cả\" để thu gọn mọi nhóm; nút đổi thành \"Hiện tất cả\" để mở lại.",
         bold_prefix="Tất cả nhóm: ")

b.h2("4.4. Các nút lưu")
b.table([
    ["Nút", "Hành động"],
    ["Lưu", "Ghi nhận cấu hình của mọi công ty, cả 2 tab. Thành công hiện \"Thao tác thành công!\" và trang tự tải lại sau khoảng 1 giây. Cấu hình có hiệu lực ngay."],
    ["Hủy", "Bỏ các thay đổi chưa Lưu và quay về trang chủ."],
])
b.para("Tải lại trang (F5) trước khi Lưu cũng làm mất các thay đổi chưa Lưu. Khi hai người cùng sửa, "
       "người Lưu sau ghi đè cấu hình — nên tải lại trang trước khi bắt đầu chỉnh.")

# ------------------------------------------------------------- PHAN 5
b.h1("PHẦN 5: LỊCH SỬ CHỈNH SỬA")
b.para("Khung \"Lịch sử chỉnh sửa\" ở cuối trang hiển thị 50 lần Lưu có thay đổi gần nhất, mới nhất ở "
       "trên cùng. Lưu mà không thay đổi gì thì không sinh dòng lịch sử.")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự, 1 là lần sửa mới nhất"],
    ["Thời gian", "Thời điểm Lưu, dạng ngày/tháng/năm giờ:phút"],
    ["Nội dung chỉnh sửa", "Mỗi thay đổi một dòng, dạng: [Tên công ty] Tích: <Nhóm> - <Thao tác> hoặc [Tên công ty] Bỏ tích: <Nhóm> - <Thao tác>. Các ô tab 2 có nhóm là \"Chặn trưởng phòng duyệt\"."],
    ["Người cập nhật", "Mã phòng ban - họ tên người Lưu"],
])
b.para("Một lần Lưu gộp mọi thay đổi của mọi công ty vào một dòng lịch sử. Khi chưa có thay đổi nào, "
       "bảng hiện \"Chưa có lịch sử\".")
img(b, "Ảnh khung Lịch sử chỉnh sửa có dữ liệu.")

# ------------------------------------------------------------- PHAN 6
b.h1("PHẦN 6: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Bấm mục menu nhưng hiện \"Không có quyền truy cập\"", "Tài khoản chưa có quyền \"Cấu hình hệ thống\". Liên hệ quản trị hệ thống để được cấp."],
    ["Đã tích ô nhưng nhân viên vẫn thao tác được", "Kiểm tra: nhân viên có thuộc đúng công ty vừa cấu hình không; hàng có thật sự quá hạn (hết hạn trước hôm nay) không; tài khoản có phải Super Admin không; phiếu có thuộc trường hợp được cho qua (Xuất thẳng, giải phóng hàng giữ, Mượn trả lại) không."],
    ["Nhân viên báo \"Có hàng nhập thẳng quá hạn!\" dù hàng mới nhập", "Công ty chưa khai báo \"Hạn xuất hàng nhập thẳng (Số ngày)\" trong quy chế công ty; khi để trống, mọi hàng nhập thẳng còn tồn đều tính quá hạn ở tab 1. Khai báo số ngày để áp đúng."],
    ["Trưởng phòng bị chặn dù nhân viên không quá hạn", "Chính trưởng phòng cũng được tính là thành viên phòng; kiểm tra hàng quá hạn của trưởng phòng trong cửa sổ danh sách."],
    ["Muốn biết ai trong phòng đang quá hạn", "Thực hiện thao tác bị chặn để mở cửa sổ danh sách, bấm Xuất Excel."],
    ["Muốn tắt chặn tạm thời", "Bỏ tích ô tương ứng của công ty rồi Lưu; tích lại khi cần chặn tiếp. Mọi lần thay đổi đều có trong Lịch sử chỉnh sửa."],
    ["Lỡ tích nhầm chưa Lưu", "Bấm Hủy hoặc tải lại trang để bỏ thay đổi."],
])

def finish_win(b):
    """Nhu b.finish() nhung bo qua tieu de trung voi tieu de cua chinh tai lieu
    (vd "TONG QUAN PHAN MEM" co ca o khung lan o noi dung that)."""
    import zipfile
    from docx import Document
    b.doc.save(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    b._update_fields_by_word()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    mine = {p.text.strip() for p in chk.paragraphs
            if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in mine]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


finish_win(b)
