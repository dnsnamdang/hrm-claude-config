# -*- coding: utf-8 -*-
"""Dung HDSD theo mau "van xuoi + bang" (mau HDSD_YeuCauXuatHang.docx user dua 12/10/2026).

Khung: HDSD_MAU_VAN_XUOI.docx (ban sao nguyen van file mau). Giu bia, MUC LUC, DANH MUC HINH ANH; xoa toan bo
noi dung tu heading "TONG QUAN PHAN MEM" tro di roi viet lai bang cac ham duoi day, dung dung style cua mau:
Heading 1/2/3, Normal (can deu), List Bullet, bang 'Light Grid Accent 1', chu thich 'Caption' + truong SEQ Hinh
(de Danh muc hinh anh tu gom). Cuoi cung cap_nhat_truong() mo Word cap nhat Muc luc + Danh muc hinh.

Dung:
    d = HdsdVanXuoi('Bảng đối chiếu công nợ 1 khách hàng')
    d.h1('TỔNG QUAN PHẦN MỀM'); d.h2('1. Thuật ngữ'); d.p('...'); d.bullet('Nhãn: ', 'nội dung')
    d.table([['Cột', 'Ý nghĩa'], [...]], widths=[2, 4.5]); d.image(path, 'chú thích')
    d.save(out); cap_nhat_truong(out)
"""
import copy
import os
import subprocess

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches

HERE = os.path.dirname(os.path.abspath(__file__))
TPL = os.path.join(HERE, 'HDSD_MAU_VAN_XUOI.docx')
TEN_MAN_MAU = 'Yêu cầu xuất hàng'


class HdsdVanXuoi:
    def __init__(self, ten_man):
        self.d = docx.Document(TPL)
        body = self.d.element.body
        # 1. doi ten man tren bia
        for p in self.d.paragraphs[:25]:
            if '(Màn hình:' in p.text:
                for t in p._p.iter(qn('w:t')):
                    if TEN_MAN_MAU in (t.text or ''):
                        t.text = t.text.replace(TEN_MAN_MAU, ten_man)
                assert ten_man in p.text, 'Không thay được tên màn trên bìa'
                break
        else:
            raise AssertionError('Không thấy dòng (Màn hình: ...) trên bìa')
        # 2. xoa noi dung tu TONG QUAN PHAN MEM tro di (giu sectPr)
        start = None
        for el in list(body.iterchildren()):
            if start is None and el.tag == qn('w:p') and ''.join(t.text or '' for t in el.iter(qn('w:t'))).strip() == 'TỔNG QUAN PHẦN MỀM':
                start = el
            if start is not None and el.tag != qn('w:sectPr'):
                body.remove(el)
        assert start is not None, 'Không thấy heading TỔNG QUAN PHẦN MỀM trong file khung'
        self.n_hinh = 0

    # ---------------------------------------------------------------- văn bản
    def _para(self, style=None, segs=(), align=WD_ALIGN_PARAGRAPH.JUSTIFY):
        p = self.d.add_paragraph(style=style)
        if align is not None:
            p.alignment = align
        if isinstance(segs, str):
            segs = [segs]
        for s in segs:
            if isinstance(s, tuple):
                r = p.add_run(s[0])
                r.bold = 'b' in s[1]
                r.italic = 'i' in s[1]
            else:
                p.add_run(s)
        return p

    def h1(self, text):
        self._para('Heading 1', text, align=None)

    def h2(self, text):
        self._para('Heading 2', text, align=None)

    def h3(self, text):
        self._para('Heading 3', text, align=None)

    def p(self, segs):
        """Đoạn thường, căn đều. segs: chuỗi, hoặc list chuỗi / (chuỗi, 'b'|'i'|'bi')."""
        return self._para(None, segs)

    def note(self, nhan, text):
        """Đoạn có nhãn đậm đầu dòng, vd note('Lưu ý: ', '...')."""
        return self._para(None, [(nhan, 'b'), text])

    def bullet(self, nhan, text=None):
        """Gạch đầu dòng; có text thì nhan in đậm (kiểu mẫu: 'Đường dẫn: ' + nội dung)."""
        segs = [(nhan, 'b'), text] if text is not None else [nhan]
        return self._para('List Bullet', segs, align=None)

    def blank(self):
        self.d.add_paragraph()

    # ---------------------------------------------------------------- bảng
    def table(self, rows, widths=None):
        ncol = len(rows[0])
        t = self.d.add_table(rows=len(rows), cols=ncol)
        t.style = self.d.styles['Light Grid Accent 1']
        t.alignment = 1
        for i, row in enumerate(rows):
            assert len(row) == ncol, 'Dòng %d lệch số cột: %r' % (i, row)
            for j, val in enumerate(row):
                c = t.cell(i, j)
                c.text = ''
                para = c.paragraphs[0]
                parts = val.split('\n') if isinstance(val, str) else [val]
                for k, part in enumerate(parts):
                    if k:
                        para = c.add_paragraph()
                    para.add_run(part)
        if widths:
            # Word đọc độ rộng ở tblGrid (python-docx chỉ ghi tcW) -> ghi cả hai, tắt tự co giãn
            t.autofit = False
            grid = t._tbl.tblGrid
            for j, gc in enumerate(grid.findall(qn('w:gridCol'))):
                gc.set(qn('w:w'), str(int(widths[j] * 1440)))
            for row in t.rows:
                for j, w in enumerate(widths):
                    row.cells[j].width = Inches(w)
        self.blank()
        return t

    # ---------------------------------------------------------------- ảnh
    def image(self, path, caption, width=6.0):
        assert os.path.exists(path), 'Thiếu ảnh ' + path
        p = self.d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(path, width=Inches(width))
        self.n_hinh += 1
        cap = self.d.add_paragraph(style='Caption')
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.add_run('Hình ')
        self._field(cap, 'SEQ Hình \\* ARABIC', str(self.n_hinh))
        cap.add_run(': ' + caption)

    @staticmethod
    def _field(p, instr, shown):
        def r(tag, **attr):
            run = OxmlElement('w:r')
            el = OxmlElement(tag)
            for k, v in attr.items():
                el.set(qn(k), v)
            run.append(el)
            p._p.append(run)
            return el
        r('w:fldChar', **{'w:fldCharType': 'begin'})
        it = r('w:instrText')
        it.set(qn('xml:space'), 'preserve')
        it.text = ' ' + instr + ' '
        r('w:fldChar', **{'w:fldCharType': 'separate'})
        t = r('w:t')
        t.text = shown
        r('w:fldChar', **{'w:fldCharType': 'end'})

    def save(self, out):
        self.d.save(out)
        return out


def cap_nhat_truong(path):
    """Mở Word (COM) cập nhật Mục lục + Danh mục hình ảnh + số hình, lưu đè. Trả về số trang."""
    ps = r'''
$ErrorActionPreference = 'Stop'
$w = New-Object -ComObject Word.Application
$w.Visible = $false
$w.DisplayAlerts = 0
try {
  $doc = $w.Documents.Open('%s')
  $doc.Fields.Update() | Out-Null
  foreach ($t in $doc.TablesOfContents) { $t.Update() }
  foreach ($t in $doc.TablesOfFigures) { $t.Update() }
  $n = $doc.ComputeStatistics(2)
  $doc.Save()
  $doc.Close()
  Write-Output "PAGES=$n"
} finally { try { $w.Quit() } catch {} }
''' % path.replace("'", "''")
    r = subprocess.run(['powershell', '-NoProfile', '-Command', ps], capture_output=True, text=True, encoding='utf-8', errors='replace')
    out = (r.stdout or '') + (r.stderr or '')
    for line in out.splitlines():
        if line.startswith('PAGES='):
            return int(line[6:])
    print('Cảnh báo: không cập nhật được trường bằng Word —', out.strip()[:300])
    return None
