# -*- coding: utf-8 -*-
"""Dung HDSD man 'Phieu xuat hang' (ProductExport / PXH) - ERP TanPhatDev.

Nguon: khao sat controller Warehouse\\ProductExportsController + model ProductExport
(+ ExportModel + ProductExportDetail + ProductExportDetailAccounting + Tab) + views
warehouse/product_exports + PermissionsTableSeeder (918/919/920). KHONG chen anh -> placeholder.

Chay:  /usr/local/bin/python3 gen_hdsd_pxh_erp.py
Output: ERP/HDSD_luongchinh/HDSD_PhieuXuatHang.docx
"""
import os
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

ERP = "/Users/nguyentrancu/DEV/code/ERP-HRM/ERP"
sys.path.insert(0, os.path.join("/Users/nguyentrancu/DEV/code/ERP-HRM/HRM",
                                ".claude", "skills", "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUT_DIR = os.path.join(ERP, "HDSD_luongchinh")
os.makedirs(OUT_DIR, exist_ok=True)
OUTPUT = os.path.join(OUT_DIR, "HDSD_PhieuXuatHang.docx")

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)


def img_placeholder(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


def clear_toc_cache(path):
    """Xoa text cache trong moi field TOC (muc luc + danh muc hinh) -> Word tu dung lai khi mo."""
    doc = Document(path)
    body = doc.element.body
    fields = []
    for el in body.iter():
        tag = el.tag
        if tag == qn('w:fldChar'):
            t = el.get(qn('w:fldCharType'))
            if t == 'begin':
                fields.append({'instr': '', 'in_result': False})
            elif t == 'separate':
                if fields:
                    fields[-1]['in_result'] = True
            elif t == 'end':
                if fields:
                    fields.pop()
        elif tag == qn('w:instrText'):
            if fields:
                fields[-1]['instr'] += (el.text or '')
        elif tag == qn('w:t'):
            if any(f['in_result'] and 'TOC' in f['instr'] for f in fields):
                el.text = ''
    doc.save(path)


def finish_macos(b):
    b.doc.save(b.output)
    clear_toc_cache(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    my_headings = {p.text.strip() for p in chk.paragraphs
                   if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in my_headings]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables),
          "| Anh:", len(chk.inline_shapes), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


# ============================================================ NOI DUNG
b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Phiếu xuất hàng)",
    doc_title="HDSD - Phiếu xuất hàng",
)

# ------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Phiếu xuất hàng (PXH)", "Chứng từ bước 4 (cuối) của quy trình xuất hàng: do Kế toán lập từ Phiếu xuất kho để ghi nhận số liệu và HẠCH TOÁN (doanh thu, giá vốn, thuế, chiết khấu…). Đây chính là màn hình tài liệu này hướng dẫn. Mã phiếu có dạng PXH-xxxxx."],
    ["Phiếu xuất kho (PXK)", "Chứng từ bước 3, do Thủ kho lập và xuất kho thực tế. Phiếu xuất hàng thường được lập từ một Phiếu xuất kho đã xuất."],
    ["Đề nghị xuất kho (ĐNXK)", "Chứng từ bước 2. Khi hoàn tất phiếu xuất hàng, đề nghị xuất kho gốc được đánh dấu đã hoàn thành."],
    ["Yêu cầu xuất hàng (YCXH)", "Chứng từ bước 1. Trường hợp xuất thẳng (không qua kho vật lý), phiếu xuất hàng lập trực tiếp từ yêu cầu xuất hàng."],
    ["Hạch toán", "Ghi nhận bút toán kế toán (Nợ / Có) cho nghiệp vụ xuất: doanh thu, giá vốn, thuế GTGT, chiết khấu, hoa hồng… Đây là điểm khác biệt cốt lõi của phiếu xuất hàng so với phiếu xuất kho."],
    ["Kho kế toán", "Kho dùng để ghi nhận giá vốn xuất. Mỗi dòng hàng trên phiếu phải phân bổ số lượng xuất vào một hoặc nhiều kho kế toán."],
    ["Giá vốn (đơn giá vốn)", "Giá xuất của hàng, tính theo giá tồn kho bình quân; dùng để hạch toán giá vốn (Nợ 632 / Có 156x). Chỉ hiển thị khi phiếu đã hoàn thành."],
    ["Tài khoản Nợ / Có", "Số hiệu tài khoản kế toán dùng trong bút toán. Với một số loại xuất, kế toán nhập trực tiếp trên phiếu."],
    ["Mã phí (mã chi phí)", "Mã phân loại chi phí, bắt buộc khi hạch toán vào một số tài khoản chi phí."],
    ["Vụ việc (mã vụ việc)", "Mã theo dõi chi phí/doanh thu theo từng công trình/vụ việc; bắt buộc khi hạch toán vào tài khoản 154."],
    ["Loại xuất (type)", "Phân loại mục đích xuất; kế thừa từ phiếu nguồn, không tự chọn."],
    ["Xuất thẳng", "Xuất không qua kho vật lý — phiếu xuất hàng lập trực tiếp từ yêu cầu xuất hàng / yêu cầu xuất ghép."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "11/09/2026", "Khởi tạo HDSD màn Phiếu xuất hàng (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Màn hình Phiếu xuất hàng là công cụ của Kế toán để chốt số liệu và hạch toán nghiệp vụ xuất: bổ sung đơn giá bán, thành tiền, thuế, giá vốn và sinh toàn bộ bút toán (doanh thu, giá vốn, chiết khấu, hoa hồng…). Đây là bước cuối biến việc xuất hàng thực tế thành số liệu kế toán hoàn chỉnh.")
b.bullet("menu Khởi tạo › Hàng hóa › Nhập - xuất hàng › \"Phiếu xuất hàng\" (danh sách tất cả phiếu theo phạm vi quyền — nơi có nút Xuất Excel).", bold_prefix="Đường dẫn danh sách: ")
b.bullet("phiếu xuất hàng được lập từ một Phiếu xuất kho đã xuất — Kế toán mở phiếu xuất kho ở trạng thái Chờ duyệt và bấm \"Tạo phiếu xuất hàng\". Trường hợp xuất thẳng thì lập từ Yêu cầu xuất hàng / Yêu cầu xuất ghép.", bold_prefix="Điểm tạo phiếu: ")
b.para("Vị trí trong quy trình xuất hàng (4 chứng từ nối tiếp):")
b.table([
    ["Bước", "Chứng từ", "Người thực hiện", "Ghi chú"],
    ["1", "Yêu cầu xuất hàng", "Nhân viên kinh doanh", "Đề nghị xuất, gửi duyệt."],
    ["2", "Đề nghị xuất kho", "Kế toán kho", "Lập từ yêu cầu, gửi Thủ kho."],
    ["3", "Phiếu xuất kho", "Thủ kho", "Chọn lô, đi lấy hàng, xuất kho (trừ tồn thực tế)."],
    ["4", "Phiếu xuất hàng", "Kế toán", "Màn hình này — ghi nhận đơn giá/thành tiền/giá vốn và HẠCH TOÁN, hoàn tất."],
])

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Phạm vi dữ liệu nhìn thấy và các thao tác đều phụ thuộc quyền. Màn này KHÔNG chặn truy cập qua đường dẫn (route) mà kiểm soát trong xử lý nghiệp vụ. Các nhóm quyền chính:")
b.bullet("3 quyền xem theo cấp (tổng công ty / công ty / phòng ban) quyết định thấy phiếu của phạm vi nào ở màn \"Phiếu xuất hàng\". Không có quyền nào thì chỉ thấy phiếu do chính mình lập.", bold_prefix="Xem danh sách: ")
b.bullet("người có quyền này được xem chi tiết/in mọi phiếu trong công ty mình; đồng thời là vai trò lập phiếu xuất hàng (từ phiếu xuất kho).", bold_prefix="Kế toán kho: ")
b.para("Các thao tác Sửa phụ thuộc quyền sở hữu + trạng thái: chỉ người lập mới sửa phiếu ở trạng thái Đang tạo. Lưu ý: quyền xem phiếu xuất hàng CHỈ có 3 cấp (tổng công ty / công ty / phòng ban), KHÔNG có cấp bộ phận. Chi tiết ở PHẦN 2.")
b.para("Cảnh báo quan trọng: khi \"Lưu & Duyệt\" (đưa phiếu về trạng thái Đã hoàn thành), hệ thống sinh bút toán, trừ tồn kho kế toán, cập nhật phiếu xuất kho gốc và tạo nhắc nợ. Đây là hành động khó thu hồi — cần kiểm tra kỹ số liệu trước khi duyệt.")

# ------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")

b.h2("1.1. Cách truy cập & các màn danh sách")
b.table([
    ["Màn", "Dùng cho", "Nội dung"],
    ["Phiếu xuất hàng (danh sách chính)", "Kế toán / người xem", "Tất cả phiếu trong phạm vi quyền; có nút Xuất Excel và phân quyền phạm vi (tổng công ty / công ty / phòng ban)."],
    ["Danh sách cho Kế toán kho", "Kế toán kho", "Lọc theo các kho kế toán mà người dùng phụ trách."],
])
img_placeholder(b, "Ảnh menu Phiếu xuất hàng.")

b.h2("1.2. Bố cục màn hình danh sách")
b.bullet("tiêu đề \"Danh sách phiếu xuất hàng\".", bold_prefix="Thanh tiêu đề: ")
b.bullet("các ô tìm theo cột (Mã phiếu, Loại, Phiếu xuất kho, Người đề nghị, Người yêu cầu, Phòng ban yêu cầu, Số HĐ, Khách hàng, Trạng thái, Người lập, Chọn kho, Tên/mã hàng hóa) và lọc theo khoảng thời gian.", bold_prefix="Khu vực lọc: ")
b.bullet("bảng liệt kê phiếu theo phạm vi quyền, cột Hành động ở cuối mỗi dòng.", bold_prefix="Bảng danh sách: ")
b.bullet("nút Tạo mới (+), nút Xuất Excel (màn danh sách chính), chọn số dòng/trang và chuyển trang.", bold_prefix="Công cụ & phân trang: ")
img_placeholder(b, "Ảnh tổng quan màn danh sách phiếu xuất hàng.")

# ------------------------------------------------- PHAN 2
b.h1("PHẦN 2: DANH SÁCH PHIẾU XUẤT HÀNG")
img_placeholder(b, "Ảnh màn danh sách với dữ liệu mẫu và cột Hành động.")

b.h2("2.1. Phân quyền & hướng dẫn theo quyền")
b.para("Danh sách hiển thị theo phạm vi quyền của người đăng nhập. Bảng dưới liệt kê đầy đủ các quyền liên quan tới màn. Người dùng thường chỉ có một vài quyền — hãy đối chiếu đúng phần áp dụng cho mình.")
b.table([
    ["Tên quyền", "Cho phép làm gì", "Nút/màn tương ứng"],
    ["Xem phiếu xuất hàng theo tổng công ty", "Thấy toàn bộ phiếu (trạng thái khác nháp) của mọi công ty.", "Danh sách"],
    ["Xem phiếu xuất hàng theo công ty", "Thấy phiếu trong công ty của mình.", "Danh sách"],
    ["Xem phiếu xuất hàng theo phòng ban", "Thấy phiếu do nhân viên thuộc các phòng ban mình quản lý lập.", "Danh sách"],
    ["Kế toán kho", "Xem chi tiết/in mọi phiếu trong công ty mình; lập phiếu xuất hàng từ phiếu xuất kho.", "Nút \"Tạo phiếu xuất hàng\" (trên màn Phiếu xuất kho); Xem; In"],
])
b.para("Lưu ý phạm vi: ở màn danh sách chính, hệ thống xét quyền xem theo thứ tự tổng công ty → công ty → phòng ban; nếu không có quyền nào, người dùng chỉ thấy phiếu do chính mình lập. Phiếu nháp (Đang tạo) của người khác luôn bị ẩn — chỉ người lập mới thấy phiếu nháp của mình. Màn \"cho Kế toán kho\" thì lọc theo các kho kế toán mà người dùng phụ trách.")

b.h3("Người dùng có quyền \"Xem phiếu xuất hàng theo tổng công ty\"")
b.para("Thấy toàn bộ phiếu (trạng thái khác nháp) của mọi công ty, cộng phiếu do chính mình lập. Thực hiện đầy đủ thao tác xem/lọc/in; sửa vẫn theo quy tắc sở hữu + trạng thái.")

b.h3("Người dùng có quyền \"Xem phiếu xuất hàng theo công ty / phòng ban\"")
b.para("Thấy phiếu trong phạm vi công ty / phòng ban tương ứng mà mình phụ trách, cộng thêm phiếu do chính mình lập.")

b.h3("Người dùng là Kế toán kho")
b.para("Đây là vai trò chính của màn: lập phiếu xuất hàng từ phiếu xuất kho, nhập hạch toán và duyệt. Kế toán kho được xem chi tiết/in mọi phiếu trong công ty mình (dù không phải người lập).")

b.h2("2.2. Tìm kiếm & lọc")
b.table([
    ["Tiêu chí", "Cách dùng"],
    ["Mã phiếu", "Nhập mã phiếu xuất hàng (PXH-…)."],
    ["Loại", "Chọn loại xuất."],
    ["Phiếu xuất kho", "Nhập mã phiếu xuất kho gốc (PXK-…)."],
    ["Người đề nghị", "Chọn (tìm theo tên) người tạo phiếu nguồn."],
    ["Người yêu cầu", "Chọn (tìm theo tên) người lập yêu cầu xuất hàng."],
    ["Phòng ban yêu cầu", "Chọn phòng ban của người yêu cầu (màn danh sách chính)."],
    ["Số HĐ", "Nhập mã hợp đồng liên quan."],
    ["Khách hàng", "Chọn (tìm theo tên/mã) khách hàng."],
    ["Trạng thái", "Chọn trạng thái phiếu (Đã hoàn thành / Đang tạo)."],
    ["Người lập", "Chọn (tìm theo tên) người lập phiếu xuất hàng."],
    ["Chọn kho", "Chọn kho để lọc phiếu theo kho xuất."],
    ["Tên, mã hàng hóa", "Nhập để lọc phiếu có chứa hàng theo tên hoặc mã."],
    ["Khoảng thời gian", "Lọc theo khoảng ngày lập phiếu."],
])

b.h2("2.3. Các cột trong danh sách")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Mã phiếu", "Mã phiếu xuất hàng; bấm để mở chi tiết."],
    ["Loại", "Loại xuất (kế thừa từ phiếu nguồn)."],
    ["Phiếu Xuất kho", "Mã phiếu xuất kho gốc (bấm để mở); trống nếu là xuất thẳng."],
    ["Người đề nghị", "Người tạo phiếu nguồn."],
    ["Người yêu cầu", "Phòng ban + người lập yêu cầu xuất hàng."],
    ["Số HĐ", "Hợp đồng liên quan (bấm để mở)."],
    ["Khách hàng", "Mã - tên khách hàng."],
    ["Trạng thái", "Trạng thái phiếu, hiển thị nhãn màu."],
    ["Ngày lập", "Ngày tạo phiếu."],
    ["Người lập", "Người lập phiếu xuất hàng (Kế toán)."],
    ["Hành động", "Các nút thao tác (xem mục 2.5)."],
])

b.h2("2.4. Bảng trạng thái phiếu")
b.table([
    ["Trạng thái", "Ý nghĩa"],
    ["Đang tạo", "Phiếu nháp — chưa hạch toán, chưa trừ kho. Chỉ người lập thấy và sửa được."],
    ["Đã hoàn thành", "Đã duyệt — đã sinh bút toán, trừ tồn kho kế toán, cập nhật phiếu xuất kho và tạo nhắc nợ."],
    ["Đang quyết toán", "Đang trong quy trình quyết toán."],
    ["Đã quyết toán", "Đã quyết toán xong."],
])
b.para("Khi lập/sửa, phiếu chỉ nhận 2 trạng thái: Đang tạo (Lưu nháp) hoặc Đã hoàn thành (Lưu & Duyệt). Trạng thái Đang quyết toán / Đã quyết toán phát sinh từ quy trình quyết toán khác, không đặt tại màn này.")

b.h2("2.5. Thao tác trên từng dòng (cột Hành động)")
b.para("Các nút nằm trong menu thao tác và hiển thị tùy trạng thái, loại phiếu và quyền sở hữu:")
b.table([
    ["Nút", "Điều kiện hiển thị", "Kết quả"],
    ["Sửa phiếu", "Phiếu Đang tạo do chính bạn lập.", "Mở màn sửa phiếu xuất hàng."],
    ["In phiếu", "Luôn có.", "In mẫu Phiếu xuất hàng."],
    ["In hạch toán", "Chỉ với loại Xuất ghép.", "In mẫu Hạch toán phiếu xuất hàng."],
])
b.para("Ngoài ra danh sách có nút Tạo mới (+) mở màn lập phiếu, và nút Xuất Excel (ở màn danh sách chính).")

# ------------------------------------------------- PHAN 3
b.h1("PHẦN 3: LẬP & HẠCH TOÁN PHIẾU XUẤT HÀNG")
b.para("Phiếu xuất hàng do Kế toán lập. Điểm vào lập phiếu thường là từ màn Phiếu xuất kho: mở phiếu xuất kho đã xuất (Chờ duyệt) rồi bấm \"Tạo phiếu xuất hàng\". Với hàng xuất thẳng, lập trực tiếp từ Yêu cầu xuất hàng / Yêu cầu xuất ghép. Loại xuất được suy ra từ phiếu nguồn, người dùng KHÔNG tự chọn.")
img_placeholder(b, "Ảnh màn lập phiếu xuất hàng.")

b.h2("3.1. Nguồn tạo phiếu")
b.para("Mỗi phiếu xuất hàng luôn gắn với một phiếu nguồn:")
b.bullet("phiếu nguồn là Phiếu xuất kho — luồng chuẩn, xuất qua kho vật lý.", bold_prefix="Chọn phiếu xuất kho: ")
b.bullet("phiếu nguồn là Yêu cầu xuất hàng (khi tích Xuất thẳng, không qua kho vật lý).", bold_prefix="Chọn phiếu yêu cầu xuất hàng: ")
b.bullet("với loại Xuất ghép, phiếu nguồn là Yêu cầu xuất ghép.", bold_prefix="Phiếu yêu cầu xuất ghép: ")
b.para("Khi phiếu đã lưu (đã có mã), ô chọn phiếu nguồn bị khóa, không đổi được. Nếu chưa chọn phiếu nguồn, bảng hàng hóa hiển thị dòng nhắc \"Chưa chọn phiếu…\".")

b.h2("3.2. Tab Thông tin chung")
b.table([
    ["Trường", "Bắt buộc", "Hiện khi / ghi chú"],
    ["Loại xuất", "—", "Chỉ hiển thị với xuất sản xuất / xuất thực hiện hợp đồng; luôn khóa."],
    ["Chọn phiếu xuất kho", "Có", "Hiển thị mã phiếu xuất kho; khóa khi phiếu đã lưu."],
    ["Xuất thẳng", "—", "Chỉ hiển thị (khóa), cho biết phiếu có xuất qua kho vật lý hay không."],
    ["Ngày hạch toán", "Có", "Với loại xuất sản xuất / thực hiện HĐ / ghép / tách; mặc định ngày hiện tại; khóa."],
    ["Tài khoản Nợ", "Có", "Với loại xuất sản xuất / thực hiện HĐ / ghép / tách — kế toán nhập. Xuất thực hiện HĐ không được dùng TK 1561."],
    ["Tài khoản Có", "Có", "Với loại xuất sản xuất — kế toán nhập."],
    ["Mã phí", "Tùy TK", "Với các loại trên; bắt buộc khi hạch toán vào tài khoản chi phí (642, tài khoản bắt đầu số 6). Có nút \"+\" tạo nhanh mã phí."],
    ["Vụ việc", "Tùy TK", "Với các loại trên; bắt buộc khi hạch toán vào tài khoản 154. Có nút \"+\" tạo nhanh vụ việc."],
    ["Hệ số điều chuyển", "Có", "Chỉ với loại điều chuyển kho chi nhánh."],
    ["Hợp đồng / Khách hàng / NVKD", "—", "Chỉ đọc — chép từ phiếu nguồn (hiện tùy loại xuất)."],
    ["Thủ kho / Kho / Người nhận hàng", "—", "Chỉ đọc — từ phiếu xuất kho."],
    ["Ghi chú", "Không", "Tối đa 255 ký tự; mặc định lấy ghi chú của phiếu xuất kho."],
    ["File đính kèm", "Không", "Với loại xuất sản xuất / thực hiện HĐ; chấp nhận pdf, png, jpg, doc, docx, xls, xlsx."],
])
b.para("Dưới các trường, hệ thống còn hiển thị khối \"gợi ý bút toán\" (chỉ để tham khảo, không nhập) tùy loại xuất — ví dụ với xuất bán thường: TK nợ 1311, doanh thu 5111, TK kho 1561, thuế 33311, giá vốn 632, chiết khấu 5211, giảm giá 5213.")

b.h2("3.3. Chọn kho kế toán & số lượng (bảng Hàng hóa)")
b.para("Đây là thao tác trung tâm của việc hạch toán: với mỗi dòng hàng, kế toán phân bổ số lượng xuất vào một hoặc nhiều kho kế toán để ghi nhận giá vốn.")
b.table([
    ["Cột", "Nhập được?", "Ghi chú"],
    ["Kho kế toán", "Chọn", "Bắt buộc. Chọn kho kế toán để lấy giá vốn. Chọn kho ở dòng đầu của sản phẩm đầu tiên sẽ tự điền cho dòng đầu của mọi sản phẩm. Không được chọn trùng kho trong cùng một hàng."],
    ["SL được xuất", "Không", "Số lượng có thể xuất từ kho kế toán đã chọn (tồn + giữ), lấy tự động."],
    ["Số lượng", "Có (nhập số)", "Số lượng xuất từ kho kế toán đó. Tổng số lượng các dòng kho phải bằng số lượng thực xuất của hàng và không được vượt \"SL được xuất\"."],
    ["Tài khoản Có", "Tùy loại", "Với loại xuất ghép/xuất tách thì nhập được; xuất thực hiện HĐ thì hiển thị (khóa)."],
])
b.para("Nút \"+ Thêm kho\" (hiện khi số lượng > 0) thêm một dòng kho kế toán cho hàng; nút xóa (X) bỏ một dòng. Ngoài phần kho kế toán, tùy loại xuất, bảng còn hiển thị các cột giá trị: Giá niêm yết, Đơn giá bán, Thành tiền bán, Chiết khấu, Đơn giá sau giảm, Thành tiền sau giảm, % VAT, Tiền VAT, Thành tiền sau VAT, và SL đề nghị / SL thực xuất. Cuối bảng có các dòng tổng (Tổng tiền bán, Tổng giảm giá, Tổng tiền trước thuế, Tiền VAT, Tổng tiền sau thuế…). Với hợp đồng hãng, hàng được chia theo các tab hợp đồng (chỉ đọc).")
img_placeholder(b, "Ảnh bảng Hàng hóa với phần chọn Kho kế toán / Số lượng và các cột giá/thành tiền/VAT.")

b.h2("3.4. Tab Vận chuyển & tab Bốc xếp")
b.bullet("chỉ hiện khi phiếu xuất kho do công ty vận chuyển. Toàn bộ chỉ đọc (chép từ phiếu xuất kho): số Km dự kiến/chốt, chi phí vận chuyển, thuế, nhân viên/công ty chịu phí, danh sách chuyến xe.", bold_prefix="Tab Vận chuyển: ")
b.bullet("khai chi phí bốc xếp để hạch toán: tính theo khối lượng hay thời gian, loại hình thuê (Công ty / Thuê ngoài), nhà cung cấp, loại bốc xếp, khối lượng/số giờ, đơn giá, VAT, tài khoản Nợ, mã phí, vụ việc. Nếu thuê Công ty thì lập bảng chia tiền cho từng nhân viên (tổng tiền chia phải bằng tổng tiền bốc xếp).", bold_prefix="Tab Bốc xếp: ")

b.h2("3.5. Các nút lưu")
b.table([
    ["Nút", "Hành động"],
    ["Lưu", "Lưu phiếu ở trạng thái Đang tạo (nháp) — chưa hạch toán, chưa trừ kho."],
    ["Lưu & Duyệt", "Chốt phiếu (Đã hoàn thành): sinh bút toán hạch toán, trừ tồn kho kế toán, cập nhật phiếu xuất kho gốc và tạo nhắc nợ."],
    ["Hủy", "Quay lại danh sách, không lưu."],
])
b.para("Ở màn tạo mới, nút Lưu bị khóa khi còn hàng chưa hợp lệ (tổng số lượng các kho kế toán không bằng số lượng thực xuất, hoặc xuất quá số được phép). Nếu chỉ Lưu nháp, hệ thống nhắc \"Phiếu đã được lưu. Bạn cần duyệt để số liệu được cập nhật\".")

b.h2("3.6. Quy tắc kiểm tra dữ liệu (validate)")
b.table([
    ["Trường / tình huống", "Yêu cầu"],
    ["Phiếu nguồn", "Bắt buộc — phiếu xuất kho (luồng chuẩn) hoặc yêu cầu xuất hàng / yêu cầu xuất ghép (xuất thẳng)."],
    ["Danh sách hàng", "Bắt buộc ít nhất 1 dòng; mỗi dòng có phân bổ kho kế toán."],
    ["Kho kế toán / Số lượng", "Bắt buộc chọn kho và nhập số lượng; tổng số lượng các kho phải khớp số thực xuất; không vượt số được xuất; không chọn trùng kho trong cùng hàng."],
    ["Tài khoản Nợ / Có", "Bắt buộc với loại xuất sản xuất / thực hiện HĐ / ghép / tách (theo loại)."],
    ["Mã phí", "Bắt buộc khi hạch toán vào tài khoản chi phí (642 / tài khoản bắt đầu số 6)."],
    ["Vụ việc", "Bắt buộc khi hạch toán vào tài khoản 154."],
    ["Hệ số điều chuyển", "Bắt buộc (> 0) với loại điều chuyển kho chi nhánh."],
    ["Ghi chú", "Tối đa 255 ký tự."],
    ["Bốc xếp thuê Công ty", "Tổng tiền chia cho các nhân viên phải bằng tổng tiền bốc xếp."],
])
b.para("Thiếu dữ liệu hoặc số lượng không khớp sẽ báo lỗi ngay tại tab/dòng tương ứng (tab có lỗi hiện dấu chấm than đỏ) và không lưu cho tới khi sửa đủ.")

# ------------------------------------------------- PHAN 4
b.h1("PHẦN 4: XEM CHI TIẾT & TAB HẠCH TOÁN")
b.para("Bấm Mã phiếu để xem chi tiết. Trang chi tiết chỉ đọc, gồm các tab: Thông tin chung, Vận chuyển (nếu có), Bốc xếp và đặc biệt là tab Hạch toán.")
b.bullet("hiển thị bảng bút toán của phiếu, gồm các cột: Tên tài khoản, Số tài khoản, Nợ, Có, Mã đối tượng, Tên đối tượng, Mã phí, Mã vụ việc, Tên vụ việc. Đây là toàn bộ số liệu kế toán mà phiếu đã sinh (doanh thu, giá vốn, thuế, chiết khấu, hoa hồng…).", bold_prefix="Tab Hạch toán: ")
b.bullet("khi phiếu đã hoàn thành, bảng hàng hóa hiển thị thêm cột Đơn giá vốn / Thành tiền vốn và dòng Tổng tiền vốn.", bold_prefix="Cột giá vốn: ")
b.para("Các nút cuối màn chi tiết: Sửa (chỉ khi phiếu Đang tạo và do chính mình lập), In, In hạch toán (chỉ loại Xuất ghép), Quay lại.")
img_placeholder(b, "Ảnh màn xem chi tiết với tab Hạch toán (bảng bút toán Nợ/Có).")

# ------------------------------------------------- PHAN 5
b.h1("PHẦN 5: SỬA PHIẾU")
b.para("Chỉ sửa được phiếu ở trạng thái Đang tạo (nháp) do chính mình lập. Mở phiếu và bấm \"Sửa phiếu\".")
b.para("Khi sửa, không đổi được loại xuất và phiếu nguồn (đã cố định). Có thể cập nhật: ghi chú, tài khoản Nợ/Có, mã phí, vụ việc, phân bổ kho kế toán, bốc xếp, đính kèm. Bấm \"Lưu\" để giữ nháp hoặc \"Lưu & Duyệt\" để chốt và hạch toán. Phiếu đã hoàn thành không sửa tại màn này.")

# ------------------------------------------------- PHAN 6
b.h1("PHẦN 6: IN ẤN & XUẤT EXCEL")
b.table([
    ["Chức năng", "Nội dung"],
    ["In phiếu", "In mẫu Phiếu xuất hàng (thông tin phiếu + bảng hàng hóa)."],
    ["In hạch toán", "In mẫu Hạch toán phiếu xuất hàng — chỉ áp dụng với loại Xuất ghép."],
    ["Xuất Excel danh sách", "Tải file Excel danh sách phiếu theo bộ lọc & phạm vi hiện tại (ở màn danh sách chính). Danh sách quá lớn (≥ 2000 dòng) sẽ được gửi qua email."],
])

# ------------------------------------------------- PHAN 7
b.h1("PHẦN 7: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Lập phiếu xuất hàng ở đâu?", "Từ màn Phiếu xuất kho: mở phiếu xuất kho đã xuất (Chờ duyệt) rồi bấm \"Tạo phiếu xuất hàng\". Cần là Kế toán kho cùng công ty. Hàng xuất thẳng thì lập từ Yêu cầu xuất hàng / xuất ghép."],
    ["Vì sao không chọn được loại xuất?", "Loại xuất được suy ra từ phiếu nguồn, không tự chọn. Chỉ loại xuất sản xuất / thực hiện HĐ mới hiện ô \"Loại xuất\" (và cũng khóa)."],
    ["Nút Lưu bị mờ (không bấm được)", "Còn hàng chưa hợp lệ: tổng số lượng các kho kế toán chưa bằng số thực xuất, hoặc nhập vượt \"SL được xuất\". Sửa lại số lượng/kho cho khớp."],
    ["Báo \"Số lượng xuất kho kế toán không khớp\"", "Tổng số lượng phân bổ vào các kho kế toán của một dòng hàng phải đúng bằng số lượng thực xuất của dòng đó."],
    ["Báo bắt buộc nhập Mã phí / Vụ việc", "Do hạch toán vào tài khoản chi phí (642, TK bắt đầu số 6) thì bắt buộc mã phí; vào tài khoản 154 thì bắt buộc vụ việc. Chọn/thêm mã phí hoặc vụ việc tương ứng."],
    ["Lưu & Duyệt rồi có sửa lại được không?", "Không sửa tại màn này. Duyệt đã sinh bút toán, trừ kho và cập nhật phiếu xuất kho — hãy kiểm tra kỹ số liệu trước khi duyệt."],
    ["Không thấy nút \"In hạch toán\"", "Nút này chỉ hiện với loại Xuất ghép."],
    ["Danh sách trống dù có nhiều phiếu", "Có thể bạn chưa có quyền xem cấp phù hợp (chỉ thấy phiếu của mình). Liên hệ quản trị để cấp quyền xem theo công ty/phòng ban."],
])

# ============================================================ FINISH
finish_macos(b)
print("XONG.")
