# -*- coding: utf-8 -*-
"""Dung HDSD man 'Phieu nhap hang' (ProductImport / PNH) - ERP TanPhatDev.

Nguon: khao sat controller Warehouse\\ProductImportsController (~1820 dong) + model
ProductImport (2284 dong, ke thua ImportModel TYPES) + ProductImportDetail /
ProductImportDetailAccounting (kho ke toan) / ProductImportDetailCustomer (phan bo KH)
+ ArrangeDelivery/Executor (boc xep) + routes product_imports + views
warehouse/product_imports (index/all/forAccounting/create/edit/form/show/formJs).
KHONG chen anh -> placeholder.

Chay:  /usr/local/bin/python3 gen_hdsd_pnh_erp.py
Output: ERP/HDSD_luongchinh/HDSD_PhieuNhapHang.docx
"""
import os
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn

ERP = "/Users/nguyentrancu/DEV/code/ERP-HRM/ERP"
sys.path.insert(0, os.path.join("/Users/nguyentrancu/DEV/code/ERP-HRM/HRM",
                                ".claude", "skills", "hdsd-documenter", "assets"))
from hdsd_engine import HdsdBuilder  # noqa: E402

OUT_DIR = os.path.join(ERP, "HDSD_luongchinh")
os.makedirs(OUT_DIR, exist_ok=True)
OUTPUT = os.path.join(OUT_DIR, "HDSD_PhieuNhapHang.docx")

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "_shots_empty")
os.makedirs(SHOTS, exist_ok=True)


def img_placeholder(b, mo_ta):
    b.para(mo_ta, bold_prefix="[Vị trí chèn ảnh] ")


def clear_toc_cache(path):
    """Xoa text cache trong moi field TOC (muc luc + danh muc hinh) -> Word tu dung lai khi mo."""
    doc = Document(path)
    body = doc.element.body
    fields = []
    for el in body.iter():
        tag = el.tag
        if tag == qn('w:fldChar'):
            t = el.get(qn('w:fldCharType'))
            if t == 'begin':
                fields.append({'instr': '', 'in_result': False})
            elif t == 'separate':
                if fields:
                    fields[-1]['in_result'] = True
            elif t == 'end':
                if fields:
                    fields.pop()
        elif tag == qn('w:instrText'):
            if fields:
                fields[-1]['instr'] += (el.text or '')
        elif tag == qn('w:t'):
            if any(f['in_result'] and 'TOC' in f['instr'] for f in fields):
                el.text = ''
    doc.save(path)


def finish_macos(b):
    b.doc.save(b.output)
    clear_toc_cache(b.output)
    b._set_update_fields()
    removed = b._purge_orphan_media()
    chk = Document(b.output)
    raw = zipfile.ZipFile(b.output).read("word/document.xml").decode("utf-8")
    my_headings = {p.text.strip() for p in chk.paragraphs
                   if p.style.name in ("Heading 1", "Heading 2", "Heading 3")}
    sot = [h for h in b.template_headings if h and h in raw and h not in my_headings]
    h1 = sum(1 for p in chk.paragraphs if p.style.name == "Heading 1")
    print("File:", b.output)
    print("Heading 1:", h1, "| Bang:", len(chk.tables),
          "| Anh:", len(chk.inline_shapes), "| Media mo coi da xoa:", len(removed))
    print("Tieu de khung con sot:", sot or "khong (OK)")
    assert not sot, "Van con tieu de khung: %s" % sot


# ============================================================ NOI DUNG
b = HdsdBuilder(
    output=OUTPUT,
    shots_dir=SHOTS,
    cover_title="(Màn hình: Phiếu nhập hàng)",
    doc_title="HDSD - Phiếu nhập hàng",
)

# ------------------------------------------------- TONG QUAN
b.h1("TỔNG QUAN PHẦN MỀM")

b.h2("1. Thuật ngữ & viết tắt")
b.para("Bảng thuật ngữ giúp người dùng hiểu đúng các bước hướng dẫn.")
b.table([
    ["Thuật ngữ", "Ý nghĩa"],
    ["Phiếu nhập hàng (PNH)", "Chứng từ HẠCH TOÁN / GIÁ VỐN — bước cuối (bước 4) của quy trình nhập kho. Do Kế toán kho lập từ một Phiếu nhập kho đã gửi (hoặc lập thẳng từ Yêu cầu nhập hàng). Đây là nơi khai GIÁ nhà cung cấp, thuế, chi phí, phân bổ kho kế toán và sinh BÚT TOÁN Nợ/Có. Đây chính là màn hình tài liệu này hướng dẫn. Mã phiếu có dạng PNH-xxxxx."],
    ["Phiếu nhập kho (PNK)", "Chứng từ bước 3 (WarehouseImport) — làm TĂNG TỒN KHO thực tế. Phiếu nhập hàng thường được tạo từ phiếu nhập kho (nút \"Tạo phiếu nhập hàng\"). Phiếu nhập kho KHÔNG lưu giá; giá và hạch toán nằm ở phiếu nhập hàng."],
    ["Yêu cầu nhập hàng (YCNH)", "Chứng từ bước 1 (ProductImportRequest). Khi \"Nhập thẳng\", phiếu nhập hàng được lập trực tiếp từ yêu cầu nhập hàng, không qua phiếu nhập kho."],
    ["Nhập thẳng", "Cờ đánh dấu phiếu nhập hàng lập trực tiếp từ Yêu cầu nhập hàng (không qua phiếu nhập kho). Trường hiển thị dạng ô đánh dấu chỉ đọc."],
    ["Kế toán kho", "Vai trò chính lập/duyệt phiếu nhập hàng; có màn danh sách riêng và được lọc phiếu theo các kho kế toán mình phụ trách."],
    ["Giá NCC (giá nhà cung cấp)", "Đơn giá mua từ nhà cung cấp (supplier_price). Là căn cứ tính giá vốn. Chỉ nhập được khi phiếu CHƯA chốt giá."],
    ["Kho kế toán", "Kho dùng để hạch toán/theo dõi giá vốn (AccountingWarehouse). Mỗi mặt hàng được phân bổ số lượng vào một hoặc nhiều kho kế toán; khi duyệt sẽ cập nhật tồn kho kế toán."],
    ["Bút toán (hạch toán)", "Các dòng Nợ/Có sinh ra khi duyệt phiếu (Đã hoàn thành). Ví dụ mua trong nước: Nợ 1541 (giá trước VAT) / Nợ 1331 (VAT) / Có 3311 (phải trả NCC); kết chuyển giá vốn Nợ 1561 / Có 1541."],
    ["Phân bổ chi phí", "Chia các chi phí (nội địa / nhập khẩu / bốc xếp) cho từng mặt hàng theo Số lượng hoặc theo Giá trị, để tính đúng giá nhập kho từng mặt hàng."],
    ["Phân bổ hàng giữ", "Chia số lượng đã về cho từng khách hàng/đơn hàng đang giữ (có hạn giữ)."],
    ["Mã phí / Vụ việc", "Mã chi phí (cost_debt) và mã vụ việc/công trình (work) gắn vào bút toán bốc xếp và một số loại nhập (nhập ghép/tách)."],
    ["Loại nhập (type)", "Phân loại mục đích nhập (mua trong/ngoài nước, mượn/bán trả lại, điều chuyển, tách/ghép, nhập gửi…); kế thừa từ phiếu nguồn, người dùng không tự chọn."],
])

b.h2("2. Lịch sử cập nhật tài liệu")
b.table([
    ["Phiên bản", "Ngày", "Nội dung", "Người thực hiện"],
    ["1.0", "11/09/2026", "Khởi tạo HDSD màn Phiếu nhập hàng (ERP)", "Phòng Phát triển"],
])

b.h2("3. Giới thiệu chung & đường dẫn truy cập")
b.para("Màn hình Phiếu nhập hàng là bước cuối của quy trình nhập kho: sau khi hàng đã được nhập kho thực tế (Phiếu nhập kho), Kế toán kho lập Phiếu nhập hàng để HẠCH TOÁN — khai giá nhà cung cấp, thuế nhập khẩu/VAT, các chi phí, phân bổ hàng vào kho kế toán và (khi duyệt) sinh bút toán Nợ/Có, cập nhật giá vốn/tồn kho kế toán. Danh sách hàng hóa được kế thừa từ phiếu nguồn — người lập KHÔNG tự thêm/xóa mặt hàng; công việc chính là khai giá/chi phí và phân bổ.")
b.bullet("menu Kho › Nhập kho › \"Phiếu nhập hàng\". Có nhiều màn danh sách theo ngữ cảnh: phiếu theo quyền của tôi (index), tất cả (all — theo cấp phân quyền), và màn dành cho Kế toán kho (forAccounting — lọc theo kho kế toán).", bold_prefix="Đường dẫn danh sách: ")
b.bullet("phiếu nhập hàng thường được tạo từ một Phiếu nhập kho đã gửi: bấm \"Tạo phiếu nhập hàng\" trên màn chi tiết phiếu nhập kho. Khi nhập thẳng, phiếu được lập trực tiếp từ Yêu cầu nhập hàng. Ngoài ra vẫn có nút \"Tạo mới\" ở màn danh sách.", bold_prefix="Điểm tạo phiếu: ")
b.para("Vị trí trong quy trình nhập kho (chuỗi chứng từ nối tiếp):")
b.table([
    ["Bước", "Chứng từ", "Người thực hiện", "Ghi chú"],
    ["1", "Yêu cầu nhập hàng", "Người yêu cầu / kế toán", "Đề nghị nhập hàng, đã được duyệt. Khi \"Nhập thẳng\" là nguồn của phiếu nhập hàng."],
    ["2", "Đề nghị nhập kho", "Kế toán kho", "Chọn kho nhập, gửi Thủ kho."],
    ["3", "Phiếu nhập kho", "Thủ kho", "Nhận hàng, nhập kho thực tế (TĂNG TỒN)."],
    ["4", "Phiếu nhập hàng", "Kế toán kho", "Màn hình này — HẠCH TOÁN / giá vốn: khai giá, thuế, chi phí, phân bổ kho kế toán, sinh bút toán Nợ/Có."],
])

b.h2("4. Quyền & phạm vi (tổng quan)")
b.para("Phạm vi dữ liệu nhìn thấy và các thao tác đều phụ thuộc quyền. Màn này KHÔNG gắn quyền ở đường dẫn (route) — toàn bộ kiểm soát nằm trong xử lý nghiệp vụ (hàm can*() của phiếu) và ẩn/hiện nút theo trạng thái. Các nhóm quyền chính:")
b.bullet("3 quyền xem theo cấp (tổng công ty / công ty / phòng ban) quyết định thấy phiếu của phạm vi nào ở màn \"Tất cả\". Không có quyền nào thì chỉ thấy phiếu do chính mình lập. Lưu ý: phiếu nhập hàng KHÔNG có mức xem \"theo bộ phận\".", bold_prefix="Xem danh sách: ")
b.bullet("vai trò cốt lõi — lập, sửa và duyệt (Lưu & Duyệt) phiếu nhập hàng; có màn danh sách \"Kế toán kho\" riêng, lọc phiếu theo các kho kế toán được gán.", bold_prefix="Kế toán kho: ")
b.bullet("quyền duyệt giá đối với hàng bán/mượn trả lại (Ban kiểm soát / Ban giám đốc) và quyền báo cáo chi phí nhập hàng theo hợp đồng — liên quan gián tiếp, không hiển thị nút trực tiếp trên màn này.", bold_prefix="Quyền liên quan: ")
b.para("Các thao tác Sửa phụ thuộc quyền sở hữu + trạng thái: chỉ người lập mới sửa phiếu ở trạng thái Đang tạo. Màn này KHÔNG có nút Duyệt/Từ chối/Hủy/Xóa riêng — việc \"duyệt\" chính là bấm \"Lưu & Duyệt\" (đưa phiếu về trạng thái Đã hoàn thành). Chi tiết ở PHẦN 2.")
b.para("Cảnh báo quan trọng: khi bấm \"Lưu & Duyệt\" (phiếu Đã hoàn thành), hệ thống SINH BÚT TOÁN, cập nhật tồn kho kế toán và giá vốn, đóng phiếu nhập kho nguồn và yêu cầu nhập hàng — hãy kiểm tra kỹ giá, thuế, chi phí và phân bổ trước khi duyệt. Sau khi duyệt, phiếu chỉ còn xem, không sửa được.")

# ------------------------------------------------- PHAN 1
b.h1("PHẦN 1: TRUY CẬP & BỐ CỤC MÀN HÌNH")

b.h2("1.1. Các màn danh sách")
b.para("Cùng một tiêu đề \"Danh sách phiếu nhập hàng\", hệ thống có 3 màn danh sách phục vụ ngữ cảnh khác nhau (khác nhau ở phạm vi dữ liệu và một vài nút/cột):")
b.table([
    ["Màn", "Dùng cho", "Nội dung / khác biệt"],
    ["Danh sách (index)", "Người lập / theo quyền", "Hiển thị phiếu theo phạm vi phân quyền của user. Có nút Tạo mới và Xuất Excel; có thêm cột \"Phiếu yc nhập hàng\"."],
    ["Tất cả (all)", "Quản lý / người có quyền xem", "Hiển thị phiếu theo cấp phân quyền (tổng công ty / công ty / phòng ban). Có nút Tạo mới, Xuất Excel; nhiều bộ lọc hơn. Ẩn phiếu nháp (Đang tạo) của người khác. KHÔNG có cột \"Phiếu yc nhập hàng\"."],
    ["Kế toán kho (forAccounting)", "Kế toán kho", "Chỉ hiển thị phiếu thuộc các kho kế toán người dùng phụ trách. Cột gọn hơn (đối tác đổi nhãn \"Nhà cung cấp\", bỏ Số hợp đồng / Phiếu yc / Người yêu cầu). KHÔNG có nút Xuất Excel."],
])
img_placeholder(b, "Ảnh menu Phiếu nhập hàng.")

b.h2("1.2. Bố cục màn hình danh sách")
b.bullet("tiêu đề \"Danh sách phiếu nhập hàng\".", bold_prefix="Thanh tiêu đề: ")
b.bullet("các ô lọc theo cột (Kho hàng, Loại, Tên/mã hàng hóa, Mã phiếu, Đối tác/Nhà cung cấp, Người đề nghị, Người yêu cầu, Phòng ban yêu cầu, Người lập, Trạng thái, Số hợp đồng, Mã phiếu nhập kho) và lọc theo khoảng thời gian.", bold_prefix="Khu vực lọc: ")
b.bullet("bảng liệt kê phiếu theo phạm vi quyền; cột Trạng thái (nhãn màu); cột Hành động ở cuối mỗi dòng.", bold_prefix="Bảng danh sách: ")
b.bullet("nút Tạo mới (+), nút Xuất Excel (màn Danh sách / Tất cả), chọn số dòng/trang và chuyển trang.", bold_prefix="Công cụ & phân trang: ")
img_placeholder(b, "Ảnh tổng quan màn danh sách phiếu nhập hàng.")

# ------------------------------------------------- PHAN 2
b.h1("PHẦN 2: DANH SÁCH PHIẾU NHẬP HÀNG")
img_placeholder(b, "Ảnh màn danh sách với dữ liệu mẫu và cột Hành động.")

b.h2("2.1. Phân quyền & hướng dẫn theo quyền")
b.para("Danh sách hiển thị theo phạm vi quyền của người đăng nhập. Bảng dưới liệt kê đầy đủ các quyền liên quan tới màn. Người dùng thường chỉ có một vài quyền — hãy đối chiếu đúng phần áp dụng cho mình.")
b.table([
    ["Tên quyền", "Cho phép làm gì", "Nút/màn tương ứng"],
    ["Xem phiếu nhập hàng theo tổng công ty", "Thấy toàn bộ phiếu (khác nháp) của mọi công ty ở màn Tất cả.", "Danh sách / Tất cả"],
    ["Xem phiếu nhập hàng theo công ty", "Thấy phiếu trong công ty của mình.", "Danh sách / Tất cả"],
    ["Xem phiếu nhập hàng theo phòng ban", "Thấy phiếu của các phòng ban mình quản lý (hoặc do mình lập).", "Danh sách / Tất cả"],
    ["Kế toán kho", "Lập/sửa/duyệt (Lưu & Duyệt) phiếu; có màn danh sách riêng, lọc theo kho kế toán được gán.", "Màn \"Kế toán kho\"; nút Tạo mới / Sửa / Lưu & Duyệt"],
    ["Ban kiểm soát / Ban giám đốc duyệt giá nhập hàng trả lại", "Tham gia duyệt giá với hàng bán/mượn trả lại (nghiệp vụ liên quan).", "(không có nút trực tiếp trên màn này)"],
    ["Xem báo cáo chi phí nhập hàng theo hợp đồng", "Xem báo cáo chi phí liên quan (nghiệp vụ báo cáo).", "(màn báo cáo riêng)"],
])
b.para("Lưu ý phạm vi: ở màn Tất cả, hệ thống xét quyền xem theo thứ tự tổng công ty → công ty → phòng ban; nếu không có quyền nào, người dùng chỉ thấy phiếu do chính mình lập. Phiếu nháp (Đang tạo) của người khác luôn bị ẩn — chỉ người lập mới thấy phiếu nháp của mình. Màn \"Kế toán kho\" lọc theo các kho kế toán mà người dùng được gán.")

b.h3("Người dùng có quyền \"Xem phiếu nhập hàng theo tổng công ty\"")
b.para("Ở màn Tất cả, thấy toàn bộ phiếu (trạng thái khác nháp) của mọi công ty, cộng phiếu do chính mình lập. Thực hiện đầy đủ thao tác xem/lọc/in/xuất Excel; sửa vẫn theo quy tắc sở hữu + trạng thái (chỉ sửa phiếu Đang tạo do mình lập).")

b.h3("Người dùng có quyền \"Xem phiếu nhập hàng theo công ty / phòng ban\"")
b.para("Thấy phiếu trong phạm vi công ty / phòng ban tương ứng mà mình phụ trách, cộng thêm phiếu do chính mình lập. Cấp phòng ban dựa trên các phòng người dùng quản lý.")

b.h3("Người dùng là Kế toán kho")
b.para("Vào màn \"Kế toán kho\" để xem danh sách phiếu thuộc các kho kế toán mình phụ trách. Đây là vai trò lập và duyệt phiếu nhập hàng: bấm Tạo mới (hoặc \"Tạo phiếu nhập hàng\" từ phiếu nhập kho), khai giá/chi phí, phân bổ rồi \"Lưu\" (nháp) hoặc \"Lưu & Duyệt\" (hoàn thành, sinh bút toán). Ngoài ra Kế toán kho cùng công ty còn xem được phiếu đã hoàn thành ngay cả khi không có quyền xem theo cấp.")
b.para("Cảnh báo chuẩn: nếu không có quyền phù hợp, phiếu sẽ không hiển thị trong danh sách; trường hợp truy cập trực tiếp bằng đường dẫn tới phiếu không thuộc phạm vi, hệ thống báo lỗi không có quyền (\"Không đủ quyền!\").")

b.h2("2.2. Tìm kiếm & lọc")
b.para("Bộ lọc thay đổi theo màn (Danh sách ít tiêu chí, Tất cả nhiều nhất, Kế toán kho rút gọn). Tổng hợp các tiêu chí:")
b.table([
    ["Tiêu chí", "Cách dùng", "Màn áp dụng"],
    ["Kho hàng", "Chọn kho (chỉ các kho người dùng được phép).", "Danh sách, Tất cả, Kế toán kho"],
    ["Loại (loại yêu cầu)", "Chọn loại nhập.", "Cả 3 màn"],
    ["Tên, mã hàng hóa", "Nhập để lọc phiếu có chứa hàng theo tên hoặc mã.", "Cả 3 màn"],
    ["Mã phiếu", "Nhập mã phiếu nhập hàng (PNH-…).", "Cả 3 màn"],
    ["Đối tác / Nhà cung cấp", "Chọn đối tác (tìm theo từ khóa) / nhà cung cấp.", "Cả 3 màn"],
    ["Người đề nghị", "Chọn (tìm theo tên) người đề nghị.", "Tất cả"],
    ["Người yêu cầu", "Chọn (tìm theo tên) người yêu cầu nhập hàng.", "Tất cả"],
    ["Phòng ban yêu cầu", "Chọn phòng ban của người yêu cầu.", "Tất cả"],
    ["Người lập", "Chọn (tìm theo tên) người lập phiếu.", "Tất cả, Kế toán kho"],
    ["Trạng thái", "Chọn: Đã hoàn thành / Đang tạo.", "Tất cả, Kế toán kho"],
    ["Số hợp đồng", "Nhập mã hợp đồng liên quan.", "Tất cả"],
    ["Mã phiếu nhập kho", "Nhập mã phiếu nhập kho nguồn.", "Tất cả"],
    ["Khoảng thời gian", "Lọc theo khoảng ngày lập phiếu.", "Cả 3 màn"],
])

b.h2("2.3. Các cột trong danh sách")
b.table([
    ["Cột", "Ý nghĩa"],
    ["STT", "Số thứ tự dòng."],
    ["Mã phiếu", "Mã phiếu nhập hàng; bấm để mở chi tiết."],
    ["Loại", "Loại nhập (kế thừa từ phiếu nguồn)."],
    ["Số hợp đồng", "Mã hợp đồng liên quan (màn Danh sách / Tất cả)."],
    ["Đối tác / Nhà cung cấp", "Nhà cung cấp hoặc khách hàng của phiếu (màn Kế toán kho hiển thị nhãn \"Nhà cung cấp\")."],
    ["Phiếu nhập kho", "Mã phiếu nhập kho nguồn."],
    ["Phiếu yc nhập hàng", "Mã yêu cầu nhập hàng nguồn (chỉ ở màn Danh sách)."],
    ["Trạng thái", "Trạng thái phiếu, hiển thị nhãn màu (xem 2.4)."],
    ["Người yêu cầu", "Người yêu cầu nhập hàng (màn Danh sách / Tất cả)."],
    ["Ngày lập", "Ngày tạo phiếu."],
    ["Người lập", "Người lập phiếu nhập hàng."],
    ["Hành động", "Các nút thao tác (xem mục 2.5)."],
])

b.h2("2.4. Bảng trạng thái phiếu")
b.para("Phiếu nhập hàng CHỈ có 2 trạng thái:")
b.table([
    ["Trạng thái", "Màu nhãn", "Ý nghĩa"],
    ["Đang tạo", "Đỏ", "Phiếu nháp — mới Lưu, chưa duyệt. Chỉ người lập thấy và sửa được. Chưa sinh bút toán, chưa cập nhật giá vốn."],
    ["Đã hoàn thành", "Xanh", "Đã duyệt (Lưu & Duyệt) — đã sinh bút toán, đã cập nhật tồn kho kế toán/giá vốn và đóng phiếu nguồn. Chỉ còn xem, không sửa."],
])
b.para("Vòng đời: Đang tạo (Lưu nháp) → Đã hoàn thành (Lưu & Duyệt). Không có bước từ chối/hủy riêng; muốn sửa số liệu phải sửa khi còn ở Đang tạo. Sau khi Đã hoàn thành, một số loại (mua nước ngoài / điều chuyển chi nhánh) sẽ chuyển tiếp sang màn Phân bổ hàng (xem PHẦN 7).")

b.h2("2.5. Thao tác trên từng dòng (cột Hành động)")
b.para("Các nút hiển thị tùy trạng thái, loại phiếu và quyền:")
b.table([
    ["Nút", "Điều kiện hiển thị", "Kết quả"],
    ["Xem", "Luôn có (theo quyền xem).", "Mở màn chi tiết phiếu (chỉ đọc), gồm cả tab Hạch toán."],
    ["Sửa", "Phiếu Đang tạo do chính bạn lập.", "Mở màn sửa để hoàn thiện giá/chi phí/phân bổ và duyệt."],
    ["Tạo phân bổ hàng", "Phiếu Đã hoàn thành, do bạn lập, loại mua nước ngoài / điều chuyển chi nhánh, chưa phân bổ.", "Chuyển sang màn Phân bổ hàng (xem PHẦN 7)."],
    ["In phiếu", "Theo quyền xem.", "In mẫu Phiếu nhập hàng (có cột Giá NCC)."],
])
b.para("Ngoài ra danh sách (Danh sách / Tất cả) có nút Tạo mới (+) và nút Xuất Excel. Màn này không có nút Duyệt/Từ chối/Hủy/Xóa trên dòng — việc duyệt thực hiện bằng \"Lưu & Duyệt\" trong form.")

# ------------------------------------------------- PHAN 3
b.h1("PHẦN 3: LẬP & SỬA PHIẾU NHẬP HÀNG — THÔNG TIN CHUNG")
b.para("Phiếu nhập hàng do Kế toán kho lập. Nguồn phổ biến nhất là bấm \"Tạo phiếu nhập hàng\" trên chi tiết Phiếu nhập kho (đã gửi); khi \"Nhập thẳng\" thì lập trực tiếp từ Yêu cầu nhập hàng. Tiêu đề khi tạo là \"Tạo phiếu nhập hàng\", khi sửa là \"Phiếu nhập hàng: {mã}\". Loại nhập và danh sách hàng hóa được suy ra từ phiếu nguồn, người dùng KHÔNG tự chọn/không tự thêm mặt hàng.")
img_placeholder(b, "Ảnh màn lập phiếu nhập hàng.")

b.h2("3.1. Bố cục form")
b.para("Form gồm 2 thẻ (card):")
b.bullet("header là tên loại nhập, góc phải hiển thị \"Người lập - Ngày lập\". Gồm 3 tab: \"Thông tin chung\" (bắt buộc), \"Vận chuyển\" (chỉ hiện khi có phát sinh vận chuyển — chỉ đọc), \"Bốc xếp\".", bold_prefix="Thẻ 1: ")
b.bullet("gồm tab \"Hàng hóa\" (bảng khai giá & phân bổ) và các tab chi phí/phân bổ (Chi phí HĐ, Chi phí quốc tế, Chi phí nội địa, Chi phí bốc hạ, Phân bổ chi phí, Phân bổ hàng giữ…) tùy loại nhập. Xem PHẦN 4 và PHẦN 5.", bold_prefix="Thẻ 2 (\"Chi tiết\"): ")
b.para("Tiêu đề tab sẽ chuyển đỏ kèm biểu tượng cảnh báo (tooltip \"Vui lòng kiểm tra lại mục này\") nếu tab đó có lỗi nhập liệu.")

b.h2("3.2. Các trường trong tab \"Thông tin chung\"")
b.para("Hầu hết trường hạch toán là chỉ đọc / điền sẵn từ phiếu nguồn; người dùng chủ yếu nhập Ghi chú và (với vài loại) chọn Tài khoản có / Vụ việc.")
b.table([
    ["Trường", "Bắt buộc", "Giá trị điền sẵn / ghi chú"],
    ["Chọn phiếu nhập kho", "Có (khi không nhập thẳng)", "Mở cửa sổ chọn phiếu nhập kho nguồn; khóa khi phiếu đã lưu."],
    ["Chọn phiếu yêu cầu nhập hàng", "Có (khi nhập thẳng)", "Nguồn khi \"Nhập thẳng\"; thường được nạp sẵn từ đường dẫn."],
    ["Nhập thẳng", "—", "Ô đánh dấu chỉ đọc — cho biết phiếu lập thẳng từ yêu cầu nhập hàng hay không."],
    ["Kho nhập", "—", "Chỉ đọc — tự điền từ phiếu nguồn."],
    ["Thủ kho", "—", "Chỉ đọc — người lập phiếu nhập kho nguồn."],
    ["Nhà cung cấp / Khách hàng", "—", "Chỉ đọc — lấy từ phiếu nhập kho / yêu cầu nhập hàng."],
    ["Người giao hàng", "—", "Chỉ đọc — lấy từ phiếu nhập kho."],
    ["Ngày hạch toán", "—", "Chỉ đọc — mặc định là ngày lập phiếu (hiển thị với loại nhập ghép/tách)."],
    ["Tài khoản có", "Có (loại nhập ghép/tách)", "Chọn tài khoản Có cho nhập ghép (type 8) / nhập tách (type 10)."],
    ["Mã Vụ việc", "—", "Chọn vụ việc/công trình (loại nhập ghép/tách)."],
    ["Không trả hóa đơn", "—", "Ô đánh dấu (loại bán/mượn trả lại); khóa khi phiếu đã lưu."],
    ["Ghi chú", "Không", "Tối đa 255 ký tự."],
])
b.para("Với các loại đặc thù, tab Thông tin chung còn hiển thị KHỐI TÀI KHOẢN HẠCH TOÁN cố định (chỉ để tham khảo, không nhập):")
b.table([
    ["Loại nhập", "Tài khoản hiển thị (cố định)"],
    ["Mua trong nước (type 2)", "TK có: 3311 · TK kho: 1561 · TK thuế: 1331 · TK chi phí thu mua: 1562."],
    ["Mua nước ngoài mới (type 11)", "TK có: 3311 · TK kho: 1561 · TK thuế: 1331."],
    ["Bán trả lại (type 4)", "TK doanh thu: 5212 · TK kho: 1561 · TK giá vốn: 632 · TK chuyển khoản: 5211 · TK giảm giá: 5213 · TK thuế: 33311."],
    ["Nhập bán (khi mượn) trả lại (type 9)", "TK doanh thu: 5212 · TK kho: 157 · TK giá vốn: 632 · TK chuyển khoản: 5211 · TK giảm giá: 5213 · TK thuế: 33311."],
])
b.para("Màn này KHÔNG có phần dự kiến vị trí / số lô (đó là của Phiếu nhập kho). Tại đây tập trung vào giá, thuế, chi phí, phân bổ kho kế toán và bút toán.")

b.h2("3.3. Cửa sổ \"Chọn phiếu nhập kho\"")
b.para("Khi không nhập thẳng, bấm nút kính lúp ở trường \"Chọn phiếu nhập kho\" để mở cửa sổ \"Phiếu nhập kho\":")
b.table([
    ["Thành phần", "Nội dung"],
    ["Cột danh sách", "STT, Mã phiếu, Người tạo, Ngày tạo."],
    ["Bộ lọc", "Mã phiếu (nhập), Người tạo (chọn)."],
    ["Điều kiện hiển thị", "Chỉ liệt kê phiếu nhập kho đã gửi (chờ duyệt) đủ điều kiện lập phiếu nhập hàng."],
    ["Chọn phiếu", "Bấm chọn 1 dòng → hệ thống nạp thông tin + danh sách hàng hóa từ phiếu nhập kho, rồi đóng cửa sổ."],
])
b.para("Khi phiếu đã lưu, ô chọn phiếu nguồn bị khóa. Nếu chưa chọn phiếu nguồn, thẻ Chi tiết hiển thị dòng nhắc \"Chưa chọn phiếu nhập kho\".")

b.h2("3.4. Tab \"Bốc xếp\"")
b.para("Khai chi phí bốc xếp (nếu có). Chọn cách tính rồi loại hình thuê:")
b.table([
    ["Trường", "Bắt buộc", "Ghi chú"],
    ["Cách tính", "—", "Chọn \"Tính theo trọng lượng\" hoặc \"Tính theo thời gian\" (loại trừ nhau)."],
    ["Loại hình thuê", "—", "Công ty (nội bộ) hoặc Thuê ngoài."],
    ["Nhà cung cấp", "Có (khi Thuê ngoài)", "Khóa khi thuê Công ty."],
    ["Loại bốc xếp", "—", "Chọn loại bốc xếp."],
    ["Khối lượng (Tấn) / Số giờ", "Có, khác 0", "Nhãn động theo cách tính."],
    ["Đơn giá / Số tiền / VAT / Số tiền sau VAT", "—", "Tự tính, chỉ đọc."],
    ["Tài khoản nợ", "Có", "Chọn tài khoản Nợ cho chi phí bốc xếp."],
    ["Mã phí", "—", "Chọn mã phí; có nút \"+\" thêm nhanh."],
    ["Vụ việc", "—", "Chọn vụ việc; có nút \"+\" thêm nhanh."],
])
b.para("Khi thuê Công ty: hiện \"Bảng người nhận việc\" — thêm nhân viên và chia số tiền; tổng tiền chia phải bằng tổng tiền bốc xếp (\"Tổng số tiền chia cho các nhân viên phải bằng tổng số tiền bốc xếp\"). Bút toán bốc xếp: Nợ (tài khoản đã chọn, thường 642/154) / Có 33481, kèm mã phí và vụ việc.")

# ------------------------------------------------- PHAN 4
b.h1("PHẦN 4: BẢNG HÀNG HÓA, GIÁ & PHÂN BỔ")
b.para("Đây là phần trọng tâm của phiếu nhập hàng: khai giá, thuế và phân bổ hàng vào kho kế toán. Bảng hàng hóa kế thừa nguyên trạng từ phiếu nguồn — người lập KHÔNG thêm/xóa mặt hàng.")

b.h2("4.1. Bảng hàng hóa (tab \"Hàng hóa\")")
b.para("Các cột thay đổi theo loại nhập; nhìn chung gồm cột thông tin (chỉ đọc), cột giá (nhập khi chưa chốt giá) và cột phân bổ (nhập được):")
b.table([
    ["Cột", "Nhập / đọc", "Ý nghĩa"],
    ["STT, Tên hàng hóa, Model, Mã hàng hóa, Đơn vị tính", "Chỉ đọc", "Thông tin sản phẩm (kế thừa từ phiếu nguồn)."],
    ["SL nhập kho", "Chỉ đọc", "Số lượng đã nhập từ phiếu nhập kho."],
    ["Giá NCC / Giá nhập kho / Đơn giá", "Nhập (chỉ khi chưa chốt giá)", "Đơn giá mua từ nhà cung cấp. Nếu phiếu đã chốt giá thì chỉ hiển thị."],
    ["Thành tiền", "Chỉ đọc (tự tính)", "Đơn giá × số lượng (× tỉ giá nếu ngoại tệ)."],
    ["VAT % / Tiền VAT", "Nhập % (loại mua trong nước) / tự tính tiền", "Thuế GTGT."],
    ["Kho kế toán", "Nhập (bắt buộc)", "Chọn kho kế toán để phân bổ hàng. Có nút \"Thêm kho\" để chia vào nhiều kho kế toán."],
    ["Số lượng", "Nhập (bắt buộc)", "Số lượng phân bổ vào kho kế toán đó (tổng phải khớp)."],
    ["Tài khoản nợ", "Nhập (loại nhập ghép/tách)", "Tài khoản Nợ theo từng kho kế toán (type 8/10)."],
    ["Hạn gửi", "Nhập (loại nhập gửi)", "Hạn gửi hàng (type 14)."],
])
b.para("Kết luận nhập liệu: chỉ Giá (khi chưa chốt giá), Kho kế toán, Số lượng phân bổ và Tài khoản nợ (nhập ghép/tách) nhập được; Thành tiền, VAT tiền, các giá trị đều tự tính (chỉ đọc). Dòng tổng hiển thị Thành tiền / VAT / Tổng cộng.")
img_placeholder(b, "Ảnh bảng hàng hóa (khai giá & phân bổ kho kế toán).")

b.h2("4.2. Phân bổ chi phí")
b.para("Với loại nhập nội địa/nhập khẩu có chi phí, tab \"Phân bổ chi phí\" chia các chi phí cho từng mặt hàng để tính đúng giá nhập kho:")
b.bullet("chọn cách phân bổ theo \"Số lượng\" hoặc \"Giá trị\", rồi bấm nút \"Phân bổ\".", bold_prefix="Cách phân bổ: ")
b.bullet("STT, Hàng hóa, ĐVT, SL, Đơn giá, Thành tiền, Chi phí NCC, Chi phí nội địa, (Thuế BVMT với loại nhập khẩu), Giá trị nhập kho, Đơn giá nhập kho.", bold_prefix="Bảng phân bổ: ")
b.bullet("hệ thống hiển thị \"Đã phân bổ …\" màu xanh (khớp) hoặc đỏ (chưa khớp). Hai nút Lưu bị KHÓA cho tới khi phân bổ đủ.", bold_prefix="Trạng thái: ")
b.para("Nếu phân bổ chưa khớp, khi lưu hệ thống báo \"Phân bổ chi phí không khớp!\".")

b.h2("4.3. Phân bổ số lượng vào kho kế toán")
b.para("Mỗi mặt hàng phải được phân bổ đủ số lượng vào các kho kế toán (cột Kho kế toán + Số lượng, nút \"Thêm kho\"). Tổng số lượng phân bổ phải bằng số lượng nhập; nếu lệch, khi lưu hệ thống báo \"Phân bổ số lượng kho kế toán không khớp!\". Đây là dữ liệu để cập nhật tồn kho kế toán và giá vốn khi duyệt.")

b.h2("4.4. Phân bổ hàng giữ")
b.para("Với các loại có giữ hàng cho khách, tab \"Phân bổ hàng giữ\" chia số lượng đã về cho từng khách hàng/đơn hàng: nhập \"Hạn giữ\" (bắt buộc), \"SL giữ\" và tích \"Cần giữ\". Nếu số liệu chưa hợp lệ, hệ thống nhắc \"Vui lòng phân bổ lại hàng giữ\".")

# ------------------------------------------------- PHAN 5
b.h1("PHẦN 5: CÁC LOẠI CHI PHÍ")
b.para("Tùy loại nhập, thẻ Chi tiết có thêm các tab chi phí. Các chi phí này được cộng vào giá vốn (qua phân bổ) và sinh bút toán tương ứng khi duyệt.")
b.table([
    ["Tab chi phí", "Áp dụng", "Nội dung chính"],
    ["Chi phí HĐ", "Mua trong nước (2/15/16) & nhập khẩu", "Nhập: Chi phí, Giá trị trước VAT (bắt buộc), VAT %, Giá trị sau VAT. Với nhập khẩu là bảng chỉ đọc kèm phụ lục."],
    ["Chi phí quốc tế (có/không tính thuế)", "Nhập khẩu (type 11)", "Tên chi phí, Giá trị, Đơn vị tiền, Tỉ giá, Giá trị (VNĐ), VAT, Nhà cung cấp, File đính kèm (bắt buộc)."],
    ["Chi phí nội địa", "Loại có chi phí nội địa", "Chi phí, Giá trị trước VAT (bắt buộc), VAT %, Giá trị sau VAT, Nhà cung cấp (bắt buộc), Ghi chú, File đính kèm (bắt buộc), nút \"Thêm chi phí\"."],
    ["Chi phí bốc hạ hàng tại kho", "Nhập khẩu", "Tương tự chi phí nội địa (pick_up_costs)."],
])
b.para("File đính kèm chi phí nhận các định dạng pdf, png, jpg, jpeg, doc, docx, xls, xlsx. Dòng chi phí bốc xếp tự sinh sẽ bị khóa (khai ở tab Bốc xếp). Thiếu dữ liệu bắt buộc sẽ báo lỗi ngay tại ô (viền đỏ) và tô đỏ tiêu đề tab.")

# ------------------------------------------------- PHAN 6
b.h1("PHẦN 6: LƯU, DUYỆT & XEM CHI TIẾT (HẠCH TOÁN)")

b.h2("6.1. Các nút lưu")
b.para("Cuối form có các nút:")
b.table([
    ["Nút", "Hành động", "Kết quả"],
    ["Lưu", "Lưu phiếu ở trạng thái Đang tạo (nháp).", "Phiếu được lưu, chưa sinh bút toán. Hệ thống báo \"Phiếu đã được lưu. Bạn cần duyệt để số liệu được cập nhật\"."],
    ["Lưu & Duyệt", "Duyệt phiếu (trạng thái Đã hoàn thành).", "Sinh bút toán, cập nhật tồn kho kế toán/giá vốn, đóng phiếu nguồn. Báo \"Số liệu đã được cập nhật\" (hoặc \"… Hãy phân bổ hàng\" nếu cần phân bổ tiếp — xem PHẦN 7)."],
    ["Hủy", "Rời form, không lưu.", "Quay lại màn danh sách Tất cả."],
])
b.para("Cả hai nút Lưu đều bị KHÓA cho tới khi phân bổ chi phí đủ (\"Đã phân bổ\" màu xanh). Thông báo khi lưu: thành công → toast xanh; lỗi kiểm tra dữ liệu → toast vàng kèm thông báo (\"Tạo thất bại!\" / \"Sửa thất bại!\" / \"Phân bổ chi phí không khớp!\" / \"Phân bổ số lượng kho kế toán không khớp!\"); lỗi hệ thống → \"Đã có lỗi xảy ra\". Nếu hàng đang kiểm kho: \"Hàng {tên} đang kiểm kho!\".")

b.h2("6.2. Xem chi tiết & tab Hạch toán")
b.para("Bấm Mã phiếu (hoặc nút Xem) để mở chi tiết. Trang chi tiết chỉ đọc, bố cục 2 thẻ giống form nhưng thẻ 1 có THÊM tab \"Hạch toán\" (không có ở form nhập liệu).")
b.bullet("bảng bút toán Nợ/Có với các cột: Tên tài khoản, Số tài khoản, Nợ, Có, Mã đối tượng, Tên đối tượng, Mã vụ việc, Tên vụ việc. Đây là toàn bộ bút toán giá vốn/hạch toán do phiếu sinh ra.", bold_prefix="Tab Hạch toán: ")
b.bullet("bảng hàng hóa, giá, chi phí và các phân bổ — chỉ để xem; với vài loại có ô \"Xem chi tiết\" để mở rộng.", bold_prefix="Thẻ Chi tiết: ")
b.para("Các nút cuối màn chi tiết: \"Sửa\" (chỉ hiện khi phiếu Đang tạo và do chính bạn lập) và \"Quay lại\". KHÔNG có nút Duyệt/Từ chối/Hủy/Xóa ở màn chi tiết.")
img_placeholder(b, "Ảnh màn xem chi tiết — tab Hạch toán (bảng bút toán Nợ/Có).")

b.h2("6.3. Ý nghĩa bút toán theo loại nhập (tham khảo)")
b.para("Khi duyệt, hệ thống sinh bút toán theo loại nhập. Một số mẫu chính:")
b.table([
    ["Loại nhập", "Bút toán chính"],
    ["Mua trong nước / Tự do+Hãng / Zitec (2/15/16)", "Giá trị hàng: Nợ 1541 (trước VAT) / Nợ 1331 (VAT) / Có 3311 (NCC). Chi phí NCC & nội địa: tương tự. Kết chuyển giá vốn: Nợ 1561 / Có 1541."],
    ["Mua nước ngoài mới (11)", "Chi phí (bốc xếp/nhập khẩu): Nợ 1541 / Nợ 1331 / Có 3311. Kết chuyển giá vốn: Nợ 1561 / Có 1541."],
    ["Nhập ghép (8) / Nhập tách (10)", "Nợ (tài khoản nợ theo kho kế toán) / Có (tài khoản có đã chọn), giá trị = giá NCC × SL, kèm mã vụ việc."],
    ["Bán trả lại (4) / Nhập bán (khi mượn) trả lại (9)", "Doanh thu (1311/5112/33311), chiết khấu (5211), giảm trừ (5213), giá vốn (157/632/1561), hoa hồng/thưởng/thuế TNCN qua các TK 35241/6411/3335/3341."],
    ["Bốc xếp", "Nợ 642/154 / Có 33481 (kèm mã phí & vụ việc)."],
])

# ------------------------------------------------- PHAN 7
b.h1("PHẦN 7: SỬA PHIẾU & PHÂN BỔ HÀNG SAU DUYỆT")
b.h2("7.1. Sửa phiếu")
b.para("Chỉ sửa được phiếu ở trạng thái Đang tạo do chính mình lập. Mở phiếu → bấm \"Sửa\". Khi sửa không đổi được phiếu nguồn (đã khóa); có thể cập nhật giá, thuế, chi phí, phân bổ kho kế toán/hàng giữ, ghi chú, file đính kèm; rồi bấm \"Lưu\" (giữ nháp) hoặc \"Lưu & Duyệt\" (duyệt). Nếu không đủ quyền sửa, hệ thống báo \"Bạn không có quyền sửa phiếu này\". Sau khi Đã hoàn thành thì không sửa được nữa.")
b.h2("7.2. Phân bổ hàng sau duyệt")
b.para("Với loại Mua nước ngoài hoặc Điều chuyển kho chi nhánh, sau khi \"Lưu & Duyệt\" hệ thống báo \"Số liệu đã được cập nhật. Hãy phân bổ hàng\" và chuyển sang màn Phân bổ hàng (phân bổ nội bộ / phân bổ điều chuyển). Người lập cũng có thể vào lại bằng nút \"Tạo phân bổ hàng\" ở cột Hành động (khi phiếu Đã hoàn thành và chưa phân bổ). Đây là bước phân chia hàng đã nhập cho các đơn hàng/chi nhánh liên quan.")

# ------------------------------------------------- PHAN 8
b.h1("PHẦN 8: IN ẤN & XUẤT EXCEL")
b.table([
    ["Chức năng", "Nội dung"],
    ["In phiếu nhập hàng", "In mẫu Phiếu nhập hàng: thông tin phiếu (công ty, số phiếu, loại nhập, mã phiếu nhập kho, thủ kho, ngày lập, người lập, nhà cung cấp, người giao hàng, ghi chú) + bảng hàng hóa (STT, Tên hàng hóa, Model, Mã hàng hóa, Thương hiệu, Số lượng, Đơn vị tính, Giá NCC). Đây là mẫu in CÓ cột giá (khác phiếu nhập kho)."],
    ["Xuất Excel danh sách", "Tải file Excel danh sách phiếu theo bộ lọc & phạm vi hiện tại (màn Danh sách / Tất cả). Cột: STT, Mã phiếu, Loại, Số hợp đồng, Đối tác, Phiếu nhập kho, Trạng thái, Người yêu cầu, Ngày lập, Người lập. Danh sách quá lớn (≥ 2000 dòng) sẽ được gửi qua email: \"Danh sách phiếu nhập hàng sẽ được gửi về mail sau ít phút!\"."],
])
b.para("Lưu ý: màn \"Kế toán kho\" không có nút Xuất Excel. Màn phiếu nhập hàng KHÔNG có mục Lịch sử phiếu riêng; tab \"Hạch toán\" ở màn chi tiết là bảng bút toán, không phải lịch sử thao tác.")

# ------------------------------------------------- PHAN 9
b.h1("PHẦN 9: CÂU HỎI THƯỜNG GẶP & LƯU Ý")
b.table([
    ["Tình huống", "Giải thích / cách xử lý"],
    ["Lập phiếu nhập hàng ở đâu?", "Thường bấm \"Tạo phiếu nhập hàng\" trên chi tiết Phiếu nhập kho (đã gửi); khi nhập thẳng thì lập từ Yêu cầu nhập hàng. Vai trò Kế toán kho."],
    ["Vì sao không chọn được loại nhập / mặt hàng?", "Loại nhập và toàn bộ hàng hóa được kế thừa từ phiếu nguồn; tại màn này chỉ khai giá/chi phí và phân bổ."],
    ["Bấm Lưu rồi mà chưa lên sổ / chưa sinh bút toán?", "Bấm \"Lưu\" chỉ tạo nháp (Đang tạo). Bút toán và giá vốn chỉ được cập nhật khi bấm \"Lưu & Duyệt\" (Đã hoàn thành)."],
    ["Nút Lưu bị mờ, không bấm được", "Phải phân bổ chi phí đủ (\"Đã phân bổ\" màu xanh) mới cho lưu. Kiểm tra tab Phân bổ chi phí và phân bổ kho kế toán."],
    ["Báo \"Phân bổ chi phí không khớp!\" / \"Phân bổ số lượng kho kế toán không khớp!\"", "Tổng chi phí phân bổ / tổng số lượng vào kho kế toán chưa khớp với số liệu hàng. Bấm lại nút \"Phân bổ\" và kiểm tra cột Số lượng / Kho kế toán."],
    ["Không nhập được Giá NCC", "Giá chỉ nhập khi phiếu chưa chốt giá. Nếu giá đã được chốt (từ nghiệp vụ trước) thì cột giá chỉ hiển thị."],
    ["Sau khi duyệt bị chuyển sang màn Phân bổ hàng", "Đúng quy trình với loại Mua nước ngoài / Điều chuyển chi nhánh — cần phân bổ hàng cho các đơn/chi nhánh. Có thể vào lại qua nút \"Tạo phân bổ hàng\"."],
    ["Không sửa được phiếu", "Chỉ sửa được phiếu Đang tạo do chính mình lập. Phiếu Đã hoàn thành không sửa được."],
    ["Muốn hủy/từ chối phiếu", "Màn này không có nút hủy/từ chối. Khi còn Đang tạo, sửa lại số liệu rồi duyệt; nghiệp vụ điều chỉnh sau khi hoàn thành xử lý ở chứng từ khác."],
    ["Danh sách trống dù có nhiều phiếu", "Có thể bạn chưa có quyền xem cấp phù hợp (chỉ thấy phiếu của mình) hoặc phiếu không thuộc kho kế toán bạn phụ trách. Liên hệ quản trị để cấp quyền."],
    ["Xem bút toán Nợ/Có ở đâu?", "Mở chi tiết phiếu → tab \"Hạch toán\": bảng liệt kê đầy đủ các dòng Nợ/Có, mã đối tượng và vụ việc."],
])

# ============================================================ FINISH
finish_macos(b)
print("XONG.")
