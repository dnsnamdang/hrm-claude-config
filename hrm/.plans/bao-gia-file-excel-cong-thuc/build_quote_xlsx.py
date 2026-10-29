"""Sinh file Excel 'Bảng tính báo giá' — công thức bám đúng hrm-client/pages/assign/quotations/_id/edit.vue."""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule

OUT = sys.argv[1]

F_INPUT = PatternFill('solid', fgColor='FFF9DB')     # ô nhập
F_FORM = PatternFill('solid', fgColor='EEF2F7')      # ô công thức
F_HEAD = PatternFill('solid', fgColor='B8CCE4')
F_SEC = PatternFill('solid', fgColor='DDEBF7')
F_TOTAL = PatternFill('solid', fgColor='FCE4D6')
F_BAD = PatternFill('solid', fgColor='F8CBAD')
THIN = Side(style='thin', color='9CA3AF')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
MONEY = '#,##0'
PCT = '0.00%'
B = Font(bold=True)
DESC_FONT = Font(italic=True, size=9, color='4B5563')

CFG = "'Cấu hình'"
NA, NB = 25, 12          # số dòng sẵn cho phần A / phần B

DEMO_A = [
    ('1', 'Hàng cha', 'TD-01', 'Tủ điện điều khiển', 2, 'Bộ', None, 30000000, 10),
    ('1.1', 'Hàng con', 'BT-75', 'Biến tần 7.5kW', 2, 'Cái', 6500000, 8000000, None),
    ('1.2', 'Hàng con', 'VT-01', 'Vỏ tủ + phụ kiện', 2, 'Bộ', 3200000, 4000000, None),
    ('2', 'Hàng độc lập', 'CB-AS', 'Cảm biến áp suất', 5, 'Cái', 1200000, 1650000, 8),
    ('3', 'Hàng độc lập', 'CAP-DK', 'Cáp điều khiển', 100, 'Mét', 25000, 35000, 10),
]
DEMO_B = [
    ('1', 'Dịch vụ', 'DV-LD', 'Lắp đặt, chạy thử', 1, 'Gói', 5000000, 8000000, 8),
    ('2', 'Chi phí', 'CP-QL', 'Chi phí quản lý dự án', 1, 'Gói', 2000000, 2500000, 8),
]
DEMO_GG1 = {0: (5, None), 3: (None, 50000)}          # A: dòng idx -> (GG%, GG₫)
DEMO_GG1_B = {0: (10, None)}
DEMO_SHIP = (1500000, 2000000)
DEMO_SHIP_GG1 = (10, None)
DEMO_GG2 = [('Chiết khấu theo dự án', '%', 3), ('Chiết khấu gói thầu', 'đ', 500000)]

MODES = {
    0: 'Không giảm giá',
    1: 'Giảm giá mặt hàng',
    2: 'Giảm giá tổng',
}


def style(c, fill=None, fmt=None, bold=False, align=None, font=None):
    c.border = BORDER
    if fill:
        c.fill = fill
    if fmt:
        c.number_format = fmt
    if bold:
        c.font = B
    if font:
        c.font = font
    if align:
        c.alignment = align


def build_columns(mode):
    cols = [
        ('stt', 'STT', 'Nhập. 1, 1.1, 1.2…', 7, 'in'),
        ('loai', 'Loại dòng', 'Chọn: Hàng cha / Hàng con / Hàng độc lập (A) — Dịch vụ / Chi phí (B)', 14, 'in'),
        ('ma', 'Mã', 'Nhập', 11, 'in'),
        ('ten', 'Tên hàng', 'Nhập', 30, 'in'),
        ('sl', 'SL', 'Nhập (> 0)', 8, 'in'),
        ('dvt', 'ĐVT', 'Nhập', 8, 'in'),
        ('gn', 'Giá nhập', 'Hàng cha: TỰ TÍNH = ROUND(Σ Thành tiền nhập con ÷ SL cha, làm tròn). Dòng khác: nhập', 14, 'in'),
        ('ttn', 'Thành tiền nhập', 'Tự tính: SL × Giá nhập', 15, 'f'),
        ('gb', 'Giá bán', 'Nhập. Hàng con: chỉ để kiểm Giá bán cha ≥ Σ con', 14, 'in'),
    ]
    if mode == 1:
        cols += [
            ('ggp', 'GG(%)', 'Nhập 0–100. Nhập 1 trong 2 ô GG', 8, 'in'),
            ('ggd', 'GG(₫)', 'Nhập (₫/đơn vị). Có số thì ưu tiên hơn GG(%)', 12, 'in'),
            ('dgsg', 'Đơn giá sau GG', 'Tự tính: Giá bán − GG₫/đv (GG₫/đv = ROUND(Giá bán × GG% ÷ 100, 2) nếu không nhập ₫)', 14, 'f'),
        ]
    if mode == 2:
        cols += [
            ('auto', 'GG phân bổ tự động', 'Tự tính: ROUND(Tổng GG × (SL × Giá bán) ÷ Σ(SL × Giá bán) mọi dòng + CP vận chuyển, 2)', 15, 'f'),
            ('pb', 'Phân bổ GG', 'Để trống = lấy GG phân bổ tự động. Nhập số = phân bổ tay', 14, 'in'),
        ]
    cols += [
        ('ttb', 'Thành tiền bán', 'Tự tính: SL × Giá bán − giảm giá của dòng' if mode else 'Tự tính: SL × Giá bán', 15, 'f'),
        ('tsln', 'Tỷ suất LN', 'Tự tính: (Thành tiền bán − Thành tiền nhập) ÷ Thành tiền nhập', 10, 'f'),
        ('vat', 'VAT(%)', 'Nhập 0–100 (chỉ dòng cha / độc lập / dịch vụ)', 8, 'in'),
        ('tvat', 'Tiền VAT', 'Tự tính: Thành tiền bán × VAT% ÷ 100', 14, 'f'),
        ('tsv', 'Thành tiền sau VAT', 'Tự tính: Thành tiền bán + Tiền VAT', 16, 'f'),
        # cột phụ (ẩn)
        ('grp', '[phụ] Nhóm cha', 'ROW() của dòng cha', 8, 'h'),
        ('gross', '[phụ] SL×Giá bán', '', 12, 'h'),
        ('chs', '[phụ] Bán con', '', 12, 'h'),
    ]
    if mode == 1:
        cols.append(('ggu', '[phụ] GG₫/đv', '', 12, 'h'))
    if mode == 2:
        cols.append(('alloc', '[phụ] Phân bổ áp dụng', '', 12, 'h'))
    return cols


def build_sheet(wb, mode):
    ws = wb.create_sheet(MODES[mode])
    cols = build_columns(mode)
    L = {k: get_column_letter(i + 1) for i, (k, *_r) in enumerate(cols)}
    last_vis = max(i for i, c in enumerate(cols) if c[4] != 'h') + 1

    ws['A1'] = f'BẢNG TÍNH BÁO GIÁ — {MODES[mode].upper()}'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A2'] = 'Ô nền VÀNG = nhập · ô nền XÁM = công thức tự tính (không gõ đè). Thêm dòng: chèn dòng TRONG khối rồi copy nguyên 1 dòng trống có sẵn.'
    ws['A2'].font = DESC_FONT

    # Cấu hình chung
    ws['A3'], ws['C3'] = 'Loại tiền', 'VNĐ'
    ws['E3'], ws['G3'] = 'Tỷ giá (VND)', 1
    ws['I3'], ws['K3'] = 'Làm tròn', '0'
    ws['L3'] = '=IF(K3="Mặc định",2,VALUE(K3))'
    ws['M3'] = '← số chữ số làm tròn (Mặc định = 2 số lẻ; VNĐ = 0)'
    for a in ('A3', 'E3', 'I3'):
        ws[a].font = B
    for a in ('C3', 'G3', 'K3'):
        style(ws[a], F_INPUT)
    style(ws['G3'], F_INPUT, '#,##0.####')
    style(ws['L3'], F_FORM)
    ws['M3'].font = DESC_FONT
    dv_round = DataValidation(type='list', formula1='"Mặc định,-3,-2,-1,0,1,2"', allow_blank=False)
    ws.add_data_validation(dv_round)
    dv_round.add('K3')
    PREC = '$L$3'

    HR, DR = 5, 6
    for i, (k, title, desc, w, kind) in enumerate(cols):
        col = i + 1
        h = ws.cell(HR, col, title)
        style(h, F_HEAD, bold=True, align=Alignment(horizontal='center', vertical='center', wrap_text=True))
        d = ws.cell(DR, col, desc)
        style(d, None, font=DESC_FONT, align=Alignment(vertical='top', wrap_text=True))
        ws.column_dimensions[get_column_letter(col)].width = w
        if kind == 'h':
            ws.column_dimensions[get_column_letter(col)].hidden = True
    ws.row_dimensions[DR].height = 75

    # Bố cục khối
    a_head = DR + 1
    a0, a1 = a_head + 1, a_head + NA
    b_head = a1 + 1
    b0, b1 = b_head + 1, b_head + NB
    tot = b1 + 1

    def rng(k, r0, r1):
        return f'${L[k]}${r0}:${L[k]}${r1}'

    # Ô dùng chung — đặt chỗ, gán sau
    gg_start = tot + 3
    ship_start = gg_start + (10 if mode == 2 else 0)
    TOTGG = f'$G${gg_start + 7}' if mode == 2 else '0'
    SHIP_COST = f'$E${ship_start + 2}'
    GRAND = f'(SUM({rng("gross", a0, a1)})+SUM({rng("gross", b0, b1)})+IF(N({SHIP_COST})>0,{SHIP_COST},0))'

    dv_a = DataValidation(type='list', formula1='"Hàng cha,Hàng con,Hàng độc lập"', allow_blank=True)
    dv_b = DataValidation(type='list', formula1='"Dịch vụ,Chi phí"', allow_blank=True)
    ws.add_data_validation(dv_a)
    ws.add_data_validation(dv_b)

    def item_row(r, sec_end, demo, gg1):
        def c(k):
            return f'{L[k]}{r}'
        isP = f'AND({c("loai")}<>"",{c("loai")}<>"Hàng con")'
        sub = lambda k: f'{L[k]}{r + 1}:{L[k]}${sec_end}'   # vùng các dòng BÊN DƯỚI (tránh tham chiếu vòng)
        f = {
            'gn': f'=IF({c("loai")}="Hàng cha",IF(N({c("sl")})>0,ROUND(SUMIFS({sub("ttn")},{sub("grp")},ROW(),{sub("loai")},"Hàng con")/{c("sl")},{PREC}),0),"")',
            'ttn': f'=IF({c("loai")}="","",N({c("sl")})*N({c("gn")}))',
            'grp': f'=IF({c("loai")}="Hàng con",{L["grp"]}{r - 1},ROW())',
            'gross': f'=IF({isP},N({c("sl")})*N({c("gb")}),0)',
            'chs': f'=IF({c("loai")}="Hàng con",N({c("sl")})*N({c("gb")}),0)',
        }
        if mode == 0:
            after = c('gross')
        elif mode == 1:
            f['ggu'] = f'=IF({isP},IF({c("ggd")}<>"",N({c("ggd")}),ROUND(N({c("gb")})*N({c("ggp")})/100,2)),0)'
            f['dgsg'] = f'=IF({isP},N({c("gb")})-{c("ggu")},"")'
            after = f'({c("gross")}-{c("ggu")}*N({c("sl")}))'
        else:
            f['auto'] = f'=IF({isP},IF(AND(N({TOTGG})>0,{GRAND}>0),ROUND({TOTGG}*{c("gross")}/{GRAND},2),0),"")'
            f['alloc'] = f'=IF({isP},IF({c("pb")}<>"",N({c("pb")}),N({c("auto")})),0)'
            after = f'({c("gross")}-{c("alloc")})'
        f['ttb'] = f'=IF({isP},{after},"")'
        f['tsln'] = f'=IF(AND({isP},N({c("ttn")})>0),ROUND(({c("ttb")}-{c("ttn")})/{c("ttn")},4),"—")'
        f['tvat'] = f'=IF({isP},{c("ttb")}*N({c("vat")})/100,"")'
        f['tsv'] = f'=IF({isP},{c("ttb")}+{c("tvat")},"")'

        vals = {}
        if demo:
            stt, loai, ma, ten, sl, dvt, gn, gb, vat = demo
            vals = dict(stt=stt, loai=loai, ma=ma, ten=ten, sl=sl, dvt=dvt, gb=gb, vat=vat)
            if gn is not None:
                vals['gn'] = gn
            if gg1:
                vals['ggp'], vals['ggd'] = gg1
        for k, _t, _d, _w, kind in cols:
            cell = ws[c(k)]
            if k in vals and vals[k] is not None:
                cell.value = vals[k]
            elif k in f:
                cell.value = f[k]
            fmt = PCT if k == 'tsln' else (None if k in ('stt', 'loai', 'ma', 'ten', 'dvt', 'grp') else MONEY)
            if k in ('sl', 'vat', 'ggp'):
                fmt = '#,##0.##'
            style(cell, F_INPUT if kind == 'in' else F_FORM, fmt)
        ws[c('stt')].number_format = '@'

    def sec_head(r, label, r0, r1, crit):
        ws.cell(r, 1, label)
        for i in range(1, last_vis + 1):
            style(ws.cell(r, i), F_SEC, bold=True)
        ws[f'{L["ttn"]}{r}'] = f'=SUMIFS({rng("ttn", r0, r1)},{rng("loai", r0, r1)},"<>Hàng con")'
        for k in ('ttb', 'tvat', 'tsv'):
            ws[f'{L[k]}{r}'] = f'=SUM({rng(k, r0, r1)})'
        for k in ('ttn', 'ttb', 'tvat', 'tsv'):
            ws[f'{L[k]}{r}'].number_format = MONEY

    sec_head(a_head, 'A — Hàng hoá', a0, a1, None)
    for i, r in enumerate(range(a0, a1 + 1)):
        item_row(r, a1, DEMO_A[i] if i < len(DEMO_A) else None, DEMO_GG1.get(i) if mode == 1 else None)
        dv_a.add(f'{L["loai"]}{r}')
    sec_head(b_head, 'B — Dịch vụ & Chi phí khác', b0, b1, None)
    for i, r in enumerate(range(b0, b1 + 1)):
        item_row(r, b1, DEMO_B[i] if i < len(DEMO_B) else None, DEMO_GG1_B.get(i) if mode == 1 else None)
        dv_b.add(f'{L["loai"]}{r}')

    # Dòng TỔNG (bảng HH/DV — KHÔNG gồm vận chuyển)
    ws.cell(tot, 1, 'TỔNG (Hàng hoá + Dịch vụ & Chi phí, chưa gồm vận chuyển)')
    for i in range(1, last_vis + 1):
        style(ws.cell(tot, i), F_TOTAL, bold=True)
    ws[f'{L["ttn"]}{tot}'] = f'={L["ttn"]}{a_head}+{L["ttn"]}{b_head}'
    if mode == 1:
        # Σ GG₫/đv × SL các dòng = Σ(SL × Giá bán) − Σ Thành tiền bán (sau GG)
        ws[f'{L["ggd"]}{tot}'] = f'=SUM({rng("gross", a0, a1)})+SUM({rng("gross", b0, b1)})-SUM({rng("ttb", a0, a1)})-SUM({rng("ttb", b0, b1)})'
    if mode == 2:
        ws[f'{L["auto"]}{tot}'] = f'=SUM({rng("auto", a0, a1)})+SUM({rng("auto", b0, b1)})'
        ws[f'{L["pb"]}{tot}'] = f'=SUM({rng("alloc", a0, a1)})+SUM({rng("alloc", b0, b1)})'
    ws[f'{L["ttb"]}{tot}'] = f'={L["ttb"]}{a_head}+{L["ttb"]}{b_head}'
    ws[f'{L["tsln"]}{tot}'] = f'=IF({L["ttn"]}{tot}>0,ROUND(({L["ttb"]}{tot}-{L["ttn"]}{tot})/{L["ttn"]}{tot},4),0)'
    ws[f'{L["tvat"]}{tot}'] = f'={L["tvat"]}{a_head}+{L["tvat"]}{b_head}'
    ws[f'{L["tsv"]}{tot}'] = f'={L["tsv"]}{a_head}+{L["tsv"]}{b_head}'
    for k in L:
        if ws[f'{L[k]}{tot}'].value is not None and k != 'stt':
            ws[f'{L[k]}{tot}'].number_format = PCT if k == 'tsln' else MONEY

    # ===== Khối Giảm giá tổng (mode 2) =====
    if mode == 2:
        r = gg_start
        ws.cell(r, 1, 'GIẢM GIÁ TỔNG').font = Font(bold=True, size=12)
        hdr = ['#', 'Loại GG', '', 'Kiểu', 'Giá trị', '', 'Thành tiền GG']
        for i, t in enumerate(hdr):
            if t:
                style(ws.cell(r + 1, i + 1, t), F_HEAD, bold=True)
        dv_k = DataValidation(type='list', formula1='"%,đ"', allow_blank=True)
        ws.add_data_validation(dv_k)
        gross_all = f'(SUM({rng("gross", a0, a1)})+SUM({rng("gross", b0, b1)})+N({SHIP_COST}))'
        for j in range(5):
            rr = r + 2 + j
            demo = DEMO_GG2[j] if j < len(DEMO_GG2) else (None, None, None)
            style(ws.cell(rr, 1, j + 1))
            style(ws.cell(rr, 2, demo[0]), F_INPUT)
            style(ws.cell(rr, 4, demo[1]), F_INPUT)
            style(ws.cell(rr, 5, demo[2]), F_INPUT, '#,##0.##')
            dv_k.add(f'D{rr}')
            style(ws.cell(rr, 7, f'=IF(D{rr}="%",ROUND({gross_all}*N(E{rr})/100,2),IF(D{rr}="đ",N(E{rr}),0))'), F_FORM, MONEY)
        # ⚠ Chỉ 5 dòng khoản GG → hàng tổng ở gg_start+7
        rr = r + 7
        ws.cell(rr, 2, 'Tổng giảm giá').font = B
        style(ws.cell(rr, 7, f'=SUM(G{r + 2}:G{r + 6})'), F_TOTAL, MONEY, bold=True)
        ws.cell(rr + 1, 2, 'Còn phải phân bổ (phải = 0)').font = B
        ws.cell(rr + 1, 3, '% tính trên Σ(SL × Giá bán) HH + DV + CP + Chi phí vận chuyển').font = DESC_FONT
        CON_LAI = f'G{rr + 1}'

    # ===== Chi phí vận chuyển =====
    r = ship_start
    ws.cell(r, 1, 'CHI PHÍ VẬN CHUYỂN').font = Font(bold=True, size=12)
    rows = [('Giá nhập vận chuyển', DEMO_SHIP[0], 'in'), ('Chi phí vận chuyển (trước VAT)', DEMO_SHIP[1], 'in')]
    for j, (lab, v, kind) in enumerate(rows):
        ws.cell(r + 1 + j, 2, lab)
        style(ws.cell(r + 1 + j, 5, v), F_INPUT, MONEY)
    SHIP_IMP = f'$E${r + 1}'
    nr = r + 3
    if mode == 1:
        ws.cell(nr, 2, 'GG vận chuyển (%)')
        style(ws.cell(nr, 5, DEMO_SHIP_GG1[0]), F_INPUT, '#,##0.##')
        ws.cell(nr + 1, 2, 'GG vận chuyển (₫) — có số thì ưu tiên hơn %')
        style(ws.cell(nr + 1, 5, DEMO_SHIP_GG1[1]), F_INPUT, MONEY)
        ws.cell(nr + 2, 2, 'GG vận chuyển áp dụng')
        style(ws.cell(nr + 2, 5, f'=IF(E{nr + 1}<>"",N(E{nr + 1}),ROUND(N({SHIP_COST})*N(E{nr})/100,2))'), F_FORM, MONEY)
        CKVC = f'$E${nr + 2}'
        nr += 3
    elif mode == 2:
        ws.cell(nr, 2, 'GG phân bổ tự động (vận chuyển)')
        style(ws.cell(nr, 5, f'=IF(AND(N({TOTGG})>0,{GRAND}>0,N({SHIP_COST})>0),ROUND({TOTGG}*{SHIP_COST}/{GRAND},2),0)'), F_FORM, MONEY)
        ws.cell(nr + 1, 2, 'Phân bổ GG vận chuyển (trống = lấy tự động)')
        style(ws.cell(nr + 1, 5), F_INPUT, MONEY)
        ws.cell(nr + 2, 2, 'GG vận chuyển áp dụng')
        style(ws.cell(nr + 2, 5, f'=IF(E{nr + 1}<>"",N(E{nr + 1}),E{nr})'), F_FORM, MONEY)
        CKVC = f'$E${nr + 2}'
        nr += 3
    else:
        CKVC = '0'
    ws.cell(nr, 2, 'VAT vận chuyển (%) — cố định 8%')
    style(ws.cell(nr, 5, 8), F_FORM, '0.##')
    SHIP_VAT = f'$E${nr}'
    if mode == 2:
        ws[CON_LAI] = f'={TOTGG}-({L["pb"]}{tot}+{CKVC})'
        style(ws[CON_LAI], F_FORM, MONEY, bold=True)

    # ===== Tổng hợp giá trị báo giá =====
    s = nr + 3
    ws.cell(s, 1, 'TỔNG HỢP GIÁ TRỊ BÁO GIÁ').font = Font(bold=True, size=12)
    heads = ['STT', 'Nhóm', 'Thành tiền nhập', 'Thành tiền trước VAT'] + (['Giảm giá', 'Thành tiền sau GG'] if mode else []) + ['Thuế VAT', 'Thành tiền sau VAT']
    for i, t in enumerate(heads):
        style(ws.cell(s + 1, i + 1, t), F_HEAD, bold=True, align=Alignment(wrap_text=True, horizontal='center'))
    A = lambda k: rng(k, a0, a1)
    Bm = lambda k: rng(k, b0, b1)
    groups = [
        ('I', 'Hàng hoá', f'SUMIFS({A("ttn")},{A("loai")},"<>Hàng con")', f'SUM({A("gross")})', f'SUM({A("ttb")})', f'SUM({A("tvat")})'),
        ('II', 'Dịch vụ', f'SUMIFS({Bm("ttn")},{Bm("loai")},"Dịch vụ")', f'SUMIFS({Bm("gross")},{Bm("loai")},"Dịch vụ")', f'SUMIFS({Bm("ttb")},{Bm("loai")},"Dịch vụ")', f'SUMIFS({Bm("tvat")},{Bm("loai")},"Dịch vụ")'),
        ('III', 'Chi phí', f'SUMIFS({Bm("ttn")},{Bm("loai")},"Chi phí")', f'SUMIFS({Bm("gross")},{Bm("loai")},"Chi phí")', f'SUMIFS({Bm("ttb")},{Bm("loai")},"Chi phí")', f'SUMIFS({Bm("tvat")},{Bm("loai")},"Chi phí")'),
        ('IV', 'Chi phí vận chuyển', f'IF(N({SHIP_COST})>0,N({SHIP_IMP}),0)', f'IF(N({SHIP_COST})>0,{SHIP_COST},0)',
         f'IF(N({SHIP_COST})>0,{SHIP_COST}-{CKVC},0)', f'IF(N({SHIP_COST})>0,({SHIP_COST}-{CKVC})*{SHIP_VAT}/100,0)'),
    ]
    col_of = {}
    ci = 3
    for k in ['nhap', 'tv'] + (['ck', 'sck'] if mode else []) + ['vat', 'sv']:
        col_of[k] = get_column_letter(ci)
        ci += 1
    for j, (stt, name, nhap, tv, sck, vat) in enumerate(groups):
        rr = s + 2 + j
        style(ws.cell(rr, 1, stt), align=Alignment(horizontal='center'))
        style(ws.cell(rr, 2, name))
        vals = {'nhap': '=' + nhap, 'tv': '=' + tv, 'vat': '=' + vat}
        if mode:
            vals['sck'] = '=' + sck
            vals['ck'] = f'={col_of["tv"]}{rr}-{col_of["sck"]}{rr}'
            vals['sv'] = f'={col_of["sck"]}{rr}+{col_of["vat"]}{rr}'
        else:
            vals['sv'] = f'={col_of["tv"]}{rr}+{col_of["vat"]}{rr}'
        for k, v in vals.items():
            style(ws[f'{col_of[k]}{rr}'], F_FORM, MONEY)
            ws[f'{col_of[k]}{rr}'] = v
    tr = s + 6
    style(ws.cell(tr, 1, 'V'), F_TOTAL, bold=True)
    style(ws.cell(tr, 2, 'Tổng giá trị báo giá'), F_TOTAL, bold=True)
    for k, cl in col_of.items():
        ws[f'{cl}{tr}'] = f'=SUM({cl}{s + 2}:{cl}{s + 5})'
        style(ws[f'{cl}{tr}'], F_TOTAL, MONEY, bold=True)
    T = {k: f'${cl}${tr}' for k, cl in col_of.items()}
    T_SCK = T['sck'] if mode else T['tv']

    # ===== Chỉ số =====
    k0 = tr + 2
    kpis = [
        ('Tỷ suất LN trước giảm giá', f'=IF({T["nhap"]}>0,ROUND(({T["tv"]}-{T["nhap"]})/{T["nhap"]},4),0)', PCT),
        ('Tỷ suất LN (sau giảm giá)', f'=IF({T["nhap"]}>0,ROUND(({T_SCK}-{T["nhap"]})/{T["nhap"]},4),0)', PCT),
        ('Doanh thu sau GG quy VND', f'={T_SCK}*N($G$3)', MONEY),
    ]
    for j, (lab, fm, nf) in enumerate(kpis):
        ws.cell(k0 + j, 2, lab).font = B
        style(ws.cell(k0 + j, 4, fm), F_FORM, nf, bold=True)
    LN = f'$D${k0 + 1}'
    VND = f'$D${k0 + 2}'

    def level(v, base):   # base: dòng đầu bảng cấu hình (cấp 1) trong sheet Cấu hình, cột C=min, D=max
        parts = []
        for lv in (1, 2):
            mn, mx = f'{CFG}!$C${base + lv - 1}', f'{CFG}!$D${base + lv - 1}'
            parts.append((lv, f'AND(OR({mn}="",{v}>={mn}),OR({mx}="",{v}<{mx}))'))
        return f'IF({parts[0][1]},1,IF({parts[1][1]},2,3))'
    ws.cell(k0 + 3, 2, 'Cấp duyệt (dự kiến)').font = B
    style(ws.cell(k0 + 3, 4, f'=IF({T_SCK}<=0,"—",MAX({level(VND, 5)},{level(LN + "*100", 10)}))'), F_FORM, bold=True)
    ws.cell(k0 + 3, 5, '1 = Tự duyệt · 2 = Trưởng phòng · 3 = TP & BGĐ (ngưỡng ở sheet Cấu hình)').font = DESC_FONT

    # ===== Tô màu cảnh báo (giống viền đỏ trên màn) =====
    def cf(cells, formula, fill=F_BAD, font=None):
        ws.conditional_formatting.add(cells, FormulaRule(formula=[formula], fill=fill, font=font, stopIfTrue=False))
    for r0, r1 in ((a0, a1), (b0, b1)):
        for r in range(r0, r1 + 1):
            lo, sl, gn, gb = (f'${L[k]}{r}' for k in ('loai', 'sl', 'gn', 'gb'))
            isP = f'AND({lo}<>"",{lo}<>"Hàng con")'
            cf(f'{L["gb"]}{r}', f'AND({lo}="Hàng cha",{L["gross"]}{r}<SUMIFS({L["chs"]}{r + 1}:{L["chs"]}${r1},{L["grp"]}{r + 1}:{L["grp"]}${r1},ROW())-0.01)')
            cf(f'{L["gb"]}{r}', f'AND({isP},N({L["ttb"]}{r})<=0)')
            cf(f'{L["gn"]}{r}', f'AND(OR({lo}="Hàng độc lập",{lo}="Hàng con"),N({gn})<=0)')
            cf(f'{L["sl"]}{r}', f'AND({lo}<>"",N({sl})<=0)')
            cf(f'{L["vat"]}{r}', f'AND({isP},OR(N({L["vat"]}{r})<0,N({L["vat"]}{r})>100))')
            if mode == 1:
                cf(f'{L["ggp"]}{r}', f'AND({isP},OR(N({L["ggp"]}{r})<0,N({L["ggp"]}{r})>100))')
                cf(f'{L["ggd"]}{r}', f'AND({isP},{L["ggu"]}{r}>N({gb}))')
        rr = f'{L["tsln"]}{r0}:{L["tsln"]}{r1}'
        cf(rr, f'AND(ISNUMBER({L["tsln"]}{r0}),{L["tsln"]}{r0}*100<{CFG}!$C$15)', fill=PatternFill(), font=Font(color='DC2626', bold=True))
        cf(rr, f'AND(ISNUMBER({L["tsln"]}{r0}),{L["tsln"]}{r0}*100>={CFG}!$C$15)', fill=PatternFill(), font=Font(color='16A34A', bold=True))
    if mode == 2:
        cf(CON_LAI, f'ABS(N({CON_LAI}))>=1')

    ws.freeze_panes = None
    ws.sheet_view.zoomScale = 90
    return ws


def build_config(wb):
    ws = wb.create_sheet('Cấu hình')
    ws['A1'] = 'CẤU HÌNH DÙNG CHO 3 SHEET (lấy theo cấu hình duyệt giá trên hệ thống — sửa nếu cấu hình thay đổi)'
    ws['A1'].font = Font(bold=True, size=12)
    ws['A3'] = 'Theo giá trị đơn hàng (VND, sau giảm giá)'
    ws['A3'].font = B
    for i, t in enumerate(['Cấp', '', 'Từ (≥)', 'Đến (<)']):
        if t:
            style(ws.cell(4, i + 1, t), F_HEAD, bold=True)
    for j, (lv, mn, mx) in enumerate([(1, 0, 1000000000), (2, 1000000000, 20000000000), (3, 20000000000, None)]):
        style(ws.cell(5 + j, 1, lv))
        style(ws.cell(5 + j, 3, mn), F_INPUT, MONEY)
        style(ws.cell(5 + j, 4, mx), F_INPUT, MONEY)
    ws['A8'] = 'Theo tỷ suất lợi nhuận (%)'
    ws['A8'].font = B
    for i, t in enumerate(['Cấp', '', 'Từ (≥)', 'Đến (<)']):
        if t:
            style(ws.cell(9, i + 1, t), F_HEAD, bold=True)
    for j, (lv, mn, mx) in enumerate([(1, 35, None), (2, 20, 35), (3, 0, 20)]):
        style(ws.cell(10 + j, 1, lv))
        style(ws.cell(10 + j, 3, mn), F_INPUT, '0.##')
        style(ws.cell(10 + j, 4, mx), F_INPUT, '0.##')
    ws['A14'] = 'Cấp duyệt = MAX(cấp theo giá trị, cấp theo tỷ suất). Không khớp khoảng nào (vd tỷ suất âm) → cấp 3.'
    ws['A14'].font = DESC_FONT
    ws['A15'] = 'Tỷ suất LN mức sàn (%)'
    ws['A15'].font = B
    style(ws['C15'], F_INPUT, '0.##')
    ws['C15'] = 0
    ws['D15'] = 'Tỷ suất < mức sàn → chữ đỏ, ≥ → chữ xanh'
    ws['D15'].font = DESC_FONT
    for c, w in zip('ABCD', (10, 4, 18, 18)):
        ws.column_dimensions[c].width = w
    ws.column_dimensions['A'].width = 12


FORMULAS = [
    ('Dòng hàng', 'Thành tiền nhập', 'SL × Giá nhập', 'Mọi dòng (cả hàng con)'),
    ('Dòng hàng', 'Giá nhập hàng cha', 'ROUND(Σ (SL con × Giá nhập con) ÷ SL cha, số chữ số làm tròn)', 'Hàng cha tự tạo. Cha hàng ERP giữ giá ERP → gõ đè giá ERP vào ô'),
    ('Dòng hàng', 'Thành tiền bán (trước GG)', 'SL × Giá bán', 'Dòng cha / độc lập / dịch vụ / chi phí. Hàng con không tính'),
    ('Dòng hàng', 'Thành tiền bán (cột trên bảng)', 'SL × Giá bán − giảm giá của dòng', 'Không GG: = SL × Giá bán'),
    ('Dòng hàng', 'Tỷ suất LN dòng', '(Thành tiền bán − Thành tiền nhập) ÷ Thành tiền nhập, làm tròn 2 số lẻ (%)', 'Thành tiền nhập ≤ 0 → "—"'),
    ('Dòng hàng', 'Tiền VAT', 'Thành tiền bán × VAT% ÷ 100', 'VAT tính trên số SAU giảm giá'),
    ('Dòng hàng', 'Thành tiền sau VAT', 'Thành tiền bán + Tiền VAT', ''),
    ('GG mặt hàng', 'GG ₫/đơn vị', 'Nhập ₫, hoặc = ROUND(Giá bán × GG% ÷ 100, 2)', 'Trên màn 2 ô % ↔ ₫ tự quy đổi; Excel: nhập 1 trong 2, ₫ được ưu tiên'),
    ('GG mặt hàng', 'Đơn giá sau GG', 'Giá bán − GG ₫/đơn vị', ''),
    ('GG mặt hàng', 'Giảm giá của dòng', 'GG ₫/đơn vị × SL', ''),
    ('GG mặt hàng', 'GG vận chuyển', 'Nhập ₫, hoặc = ROUND(CP vận chuyển × % ÷ 100, 2)', ''),
    ('GG tổng', 'Thành tiền 1 khoản GG', 'Kiểu đ: = Giá trị · Kiểu %: = ROUND(Σ(SL × Giá bán) HH + DV + CP + CP vận chuyển × % ÷ 100, 2)', 'Nền % GỒM cả vận chuyển'),
    ('GG tổng', 'Tổng giảm giá', 'Σ Thành tiền các khoản GG', ''),
    ('GG tổng', 'GG phân bổ tự động (dòng)', 'Tổng GG × (SL × Giá bán của dòng) ÷ [Σ(SL × Giá bán) mọi dòng + CP vận chuyển]', 'Màn hình dùng phương pháp phần dư lớn nhất tới 0,01 → Excel có thể lệch 0,01 ở vài dòng'),
    ('GG tổng', 'Phân bổ GG (dòng)', 'Mặc định = GG phân bổ tự động (như đã bấm "Phân bổ tự động"); được sửa tay', ''),
    ('GG tổng', 'Còn phải phân bổ', 'Tổng GG − Σ Phân bổ GG (gồm dòng vận chuyển)', 'Phải = 0 (lệch ≥ 1 → tô đỏ)'),
    ('Vận chuyển', 'VC sau GG', 'CP vận chuyển − GG vận chuyển', ''),
    ('Vận chuyển', 'VAT vận chuyển', '(CP vận chuyển − GG) × 8%', 'VAT vận chuyển cố định 8%'),
    ('Tổng hợp', 'Thành tiền nhập nhóm', 'Σ Thành tiền nhập dòng cha/độc lập (KHÔNG cộng hàng con)', 'Vận chuyển: Giá nhập vận chuyển'),
    ('Tổng hợp', 'Thành tiền trước VAT', 'Σ SL × Giá bán (trước GG)', 'Dịch vụ / Chi phí tách theo "Có tính doanh thu"'),
    ('Tổng hợp', 'Giảm giá nhóm', 'Thành tiền trước VAT − Thành tiền sau GG', ''),
    ('Tổng hợp', 'Thành tiền sau VAT', 'Thành tiền sau GG + Thuế VAT', ''),
    ('Chỉ số', 'Tỷ suất LN trước GG', '(Tổng trước VAT − Tổng nhập) ÷ Tổng nhập', 'Gồm vận chuyển'),
    ('Chỉ số', 'Tỷ suất LN (sau GG)', '(Tổng sau GG − Tổng nhập) ÷ Tổng nhập', 'Dùng để xét cấp duyệt'),
    ('Chỉ số', 'Tỷ suất dòng TỔNG bảng', '(Σ sau GG − Σ nhập) của HH + DV + CP, KHÔNG gồm vận chuyển', ''),
    ('Chỉ số', 'Doanh thu quy VND', 'Tổng sau GG × Tỷ giá', ''),
    ('Chỉ số', 'Cấp duyệt', 'MAX(cấp theo doanh thu quy VND, cấp theo tỷ suất LN sau GG)', 'Ngưỡng ở sheet Cấu hình'),
    ('Làm tròn', 'Số chữ số', 'Mặc định = 2 số lẻ; VNĐ = 0; chọn −3 (nghìn) … 2', 'Áp vào Giá nhập hàng cha tự tính. Giá nhập/Giá bán còn lại nhập sẵn số đã làm tròn (trên màn = bấm "Áp dụng")'),
]
RULES = [
    'Giá bán hàng cha ≥ Σ (SL × Giá bán) hàng con → sai thì ô Giá bán cha tô đỏ',
    'Giá / Thành tiền bán sau GG phải > 0 → sai thì ô Giá bán tô đỏ',
    'Giá nhập hàng độc lập / hàng con phải > 0; SL phải > 0',
    'VAT và GG(%) phải trong 0–100; GG(₫) không lớn hơn Giá bán',
    'Giá nhập cha tự tạo ≥ Σ giá nhập con (cha hàng ERP được phép nhỏ hơn) — Excel không phân biệt hàng ERP nên không tô',
]


def build_guide(wb):
    ws = wb.active
    ws.title = 'Hướng dẫn'
    ws['A1'] = 'BẢNG TÍNH BÁO GIÁ — HƯỚNG DẪN & DANH SÁCH CÔNG THỨC'
    ws['A1'].font = Font(bold=True, size=14)
    lines = [
        'Công thức bám đúng màn Báo giá đang chạy (Tạo / Sửa báo giá). Chọn sheet theo kiểu giảm giá: Không giảm giá · Giảm giá mặt hàng · Giảm giá tổng.',
        'Ô nền VÀNG là ô nhập; ô nền XÁM là công thức. Dữ liệu sẵn trong file là dữ liệu MẪU — xoá ô vàng rồi nhập số của mình.',
        'Phần A (Hàng hoá): mỗi cụm là 1 dòng "Hàng cha" + các dòng "Hàng con" NGAY BÊN DƯỚI. Hàng không có con chọn "Hàng độc lập".',
        'Phần B: "Dịch vụ" = dịch vụ có tính doanh thu, "Chi phí" = dịch vụ & chi phí khác không tính doanh thu.',
        'Thêm dòng: chèn dòng ở GIỮA khối (không chèn dưới dòng cuối khối) rồi copy 1 dòng trống có sẵn dán vào. Ô Giá nhập của dòng trống có sẵn công thức hàng cha — dòng không phải hàng cha thì gõ đè số.',
        'Không tự lấy được giá ERP, tỷ giá, cấu hình duyệt: tự nhập. Khi Import lên hệ thống, giá hàng ERP vẫn lấy lại theo ERP.',
    ]
    for i, t in enumerate(lines):
        ws.cell(3 + i, 1, f'{i + 1}. {t}')
    r = 3 + len(lines) + 1
    ws.cell(r, 1, 'DANH SÁCH CÔNG THỨC').font = Font(bold=True, size=12)
    for i, t in enumerate(['Khối', 'Trường', 'Công thức', 'Ghi chú']):
        style(ws.cell(r + 1, i + 1, t), F_HEAD, bold=True)
    for j, row in enumerate(FORMULAS):
        for i, t in enumerate(row):
            style(ws.cell(r + 2 + j, i + 1, t), align=Alignment(wrap_text=True, vertical='top'))
    r = r + 2 + len(FORMULAS) + 1
    ws.cell(r, 1, 'QUY TẮC KIỂM TRA (tô đỏ như viền đỏ trên màn)').font = Font(bold=True, size=12)
    for j, t in enumerate(RULES):
        ws.cell(r + 1 + j, 1, f'• {t}')
    for c, w in zip('ABCD', (16, 30, 70, 55)):
        ws.column_dimensions[c].width = w


wb = Workbook()
build_guide(wb)
for m in (0, 1, 2):
    build_sheet(wb, m)
build_config(wb)
wb.save(OUT)
print('saved', OUT)
