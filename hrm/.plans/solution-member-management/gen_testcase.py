# -*- coding: utf-8 -*-
"""
Sinh test case cho Redmine #11354 — Quản lý nhân sự giải pháp (Cập nhật / Khóa / Xóa).

KHÁC với khuôn của skill `testcase-documenter` (file .xlsx độc lập có 9 mục mô tả + 2 khối
TEST SUMMARY): lần này test case phải CHÈN TIẾP vào tab "Testcase _ Quản lý chi tiết GP" của
file Google Sheets "Testcase _Quản lý dự án" mà team đang dùng, nên chỉ sinh ĐÚNG phần dòng dữ
liệu theo 19 cột sẵn có của tab đó, không kèm khối mô tả/summary để khỏi phá bố cục cũ.

Quy ước bám theo tab hiện tại:
  - Cột A "Module" = "Quản lý giải pháp"
  - Mã TC nối tiếp TC-ROLE-91 (mã cuối cùng đang có) -> bắt đầu từ TC-ROLE-92
  - Dòng tiêu đề nhóm: chỉ điền cột "Chức năng", các cột khác để trống
  - 3 cột trạng thái DNS mặc định "Not Executed"; 3 cột TPE để trống
"""

import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

MODULE = "Quản lý giải pháp"
START_NO = 92

HEADERS = [
    "Module", "Nhóm chức năng", "TC ID", "Chức năng", "Priority",
    "Tiền điều kiện", "Bước thực hiện", "Test Data", "Expected Result (chi tiết)",
    "KQ thực tế", "trạng tháicheck lần 1.", "trạng tháicheck lần2", "trạng tháicheck lần3",
    "Ghi chú", "TPE TEST Lần 1.", "TPE TEST Lần2", "TPE TEST Lần3",
    "Kết quả hiện tại", "Ghi chú TPE",
]

# (Chức năng, Priority, Tiền điều kiện, Bước thực hiện, Test Data, Expected Result)
GROUPS = [
    ("Tab Nhân sự — Quyền thao tác trên thành viên", [
        (
            "Trưởng phòng giải pháp thao tác được mọi thành viên", "P0",
            "Giải pháp G1 đang triển khai, có 6 thành viên. Đăng nhập bằng tài khoản đã tạo giải pháp G1 (Trưởng phòng giải pháp).",
            "1. Vào Dự án TKT → Danh sách làm giải pháp\n2. Bấm Quản lý ở giải pháp G1\n3. Mở tab Nhân sự\n4. Quan sát cột Thao tác ở từng dòng",
            "—",
            "- Bảng có thêm cột Thao tác ở cuối\n- Cả 6 dòng thành viên đều hiện đủ 3 nút: Sửa phân công (bút chì), Khóa thành viên (ổ khóa), Xóa khỏi giải pháp (thùng rác đỏ)",
        ),
        (
            "PM giải pháp thao tác được mọi thành viên", "P0",
            "Giải pháp G1 có 6 thành viên. Đăng nhập bằng tài khoản đang là PM của G1.",
            "1. Mở tab Nhân sự của giải pháp G1\n2. Quan sát cột Thao tác",
            "—",
            "- Cả 6 dòng thành viên đều hiện đủ 3 nút thao tác",
        ),
        (
            "Leader hạng mục chỉ thao tác được thành viên hạng mục mình phụ trách", "P0",
            "Giải pháp G2 có 2 hạng mục: HM1 do ông A làm Leader, HM2 do ông B làm Leader. Mỗi hạng mục có 1 thành viên. Đăng nhập bằng tài khoản ông A (KHÔNG phải PM, KHÔNG phải người tạo giải pháp).",
            "1. Mở tab Nhân sự của giải pháp G2\n2. Quan sát cột Thao tác ở dòng thành viên HM1\n3. Quan sát cột Thao tác ở dòng thành viên HM2",
            "—",
            "- Dòng thành viên thuộc HM1: hiện đủ 3 nút\n- ⚠️ Dòng thành viên thuộc HM2: KHÔNG có nút nào, cột Thao tác hiển thị dấu gạch ngang",
        ),
        (
            "Thành viên thường không thao tác được ai", "P0",
            "Giải pháp G1 có 6 thành viên. Đăng nhập bằng tài khoản là thành viên thường của G1 (không phải PM, không phải Leader, không phải người tạo).",
            "1. Mở tab Nhân sự của giải pháp G1\n2. Quan sát cột Thao tác ở tất cả các dòng",
            "—",
            "- Toàn bộ các dòng đều hiển thị dấu gạch ngang ở cột Thao tác, không có nút nào",
        ),
        (
            "Dòng PM và dòng Leader hạng mục luôn ẩn hết nút", "P0",
            "Giải pháp G2 có chia hạng mục, có PM và có Leader từng hạng mục. Đăng nhập bằng tài khoản PM.",
            "1. Mở tab Nhân sự của giải pháp G2\n2. Tìm dòng có vai trò PM làm giải pháp\n3. Tìm dòng có vai trò Leader hạng mục (người này KHÔNG được thêm làm thành viên)\n4. Quan sát cột Thao tác của 2 dòng đó",
            "—",
            "- ⚠️ Cả dòng PM lẫn dòng Leader hạng mục đều hiển thị dấu gạch ngang, không có nút Sửa/Khóa/Xóa\n- Lý do: 2 dòng này lấy từ cấu hình giải pháp và hạng mục, không phải bản ghi thành viên; muốn đổi thì dùng chức năng Phân công",
        ),
        (
            "Người vừa là Leader hạng mục vừa là thành viên chỉ hiện 1 dòng", "P0",
            "Giải pháp G2, hạng mục HM1 có Leader là ông A. Thêm chính ông A làm thành viên HM1 với vai trò Thiết kế. Đăng nhập bằng tài khoản PM.",
            "1. Mở tab Nhân sự của giải pháp G2\n2. Đếm số dòng mang tên ông A",
            "—",
            "- ⚠️ Chỉ có 1 dòng duy nhất của ông A, không bị lặp 2 dòng\n- Cột Vai trò ghép cả hai: Leader hạng mục, Thiết kế\n- Dòng này CÓ nút thao tác (vì có bản ghi thành viên)",
        ),
        (
            "Giải pháp đã đóng thì ẩn hết nút thao tác", "P0",
            "Giải pháp G3 ở trạng thái đã đóng, có 3 thành viên. Đăng nhập bằng tài khoản PM của G3.",
            "1. Mở tab Nhân sự của giải pháp G3\n2. Quan sát cột Thao tác và nút Phân công phía trên",
            "—",
            "- Không dòng nào hiện nút Sửa/Khóa/Xóa\n- Nút Phân công cũng không hiện — cách ứng xử giống nhau ở cả hai chỗ",
        ),
        (
            "Gọi thẳng chức năng Sửa, bỏ qua giao diện", "P0",
            "Đăng nhập bằng tài khoản thành viên thường của giải pháp G1.",
            "1. Dùng công cụ kiểm thử API gọi thẳng chức năng Sửa phân công cho một thành viên của G1, bỏ qua giao diện",
            "—",
            "- Hệ thống từ chối, báo không có quyền thao tác trên thành viên này\n- ⚠️ Dữ liệu thành viên giữ nguyên, không bị sửa",
        ),
        (
            "Gọi thẳng chức năng Xóa với thành viên của giải pháp KHÁC", "P0",
            "Đăng nhập bằng tài khoản PM của giải pháp G1. Giải pháp G2 là giải pháp khác, có thành viên riêng.",
            "1. Dùng công cụ kiểm thử API gọi chức năng Xóa thành viên, truyền mã giải pháp G1 nhưng mã dòng thành viên lại thuộc G2",
            "—",
            "- Hệ thống báo không tìm thấy thành viên trong giải pháp này, đề nghị tải lại danh sách\n- ⚠️ Thành viên bên G2 KHÔNG bị xóa",
        ),
    ]),

    ("Tab Nhân sự — Cập nhật phân công thành viên", [
        (
            "Mở cửa sổ Sửa phân công, dữ liệu điền sẵn đúng", "P0",
            "Giải pháp G1 có thành viên ông C: vai trò Thành viên, ngày bắt đầu 25/05/2026, chưa có ngày kết thúc, mô tả công việc 'Theo dõi hồ sơ'. Đăng nhập bằng tài khoản PM.",
            "1. Mở tab Nhân sự\n2. Bấm nút Sửa phân công ở dòng ông C\n3. Quan sát các ô trong cửa sổ",
            "—",
            "- Cửa sổ Sửa phân công thành viên mở ra, dòng đầu ghi đúng tên ông C\n- Vai trò dự án: Thành viên\n- Mô tả công việc: Theo dõi hồ sơ\n- Ngày bắt đầu: 25/05/2026\n- Ngày kết thúc: để trống",
        ),
        (
            "Sửa ngày kết thúc và lưu thành công", "P0",
            "Như trên, đang mở cửa sổ Sửa phân công của ông C.",
            "1. Chọn Ngày kết thúc là 31/12/2026\n2. Bấm Lưu\n3. Quan sát bảng",
            "Ngày kết thúc: 31/12/2026",
            "- Hiện thông báo Đã cập nhật phân công\n- Cửa sổ đóng lại\n- Cột Ngày kết thúc của dòng ông C hiển thị 31/12/2026 ngay, không cần tải lại trang",
        ),
        (
            "Sửa vai trò dự án", "P1",
            "Giải pháp G1 có thành viên ông C vai trò Thành viên. Danh mục vai trò có sẵn giá trị Thiết kế.",
            "1. Bấm Sửa phân công ở dòng ông C\n2. Đổi Vai trò dự án sang Thiết kế\n3. Bấm Lưu",
            "Vai trò dự án: Thiết kế",
            "- Hiện thông báo Đã cập nhật phân công\n- Cột Vai trò của dòng ông C đổi thành Thiết kế",
        ),
        (
            "Sửa mô tả công việc", "P1",
            "Giải pháp G1 có thành viên ông C.",
            "1. Bấm Sửa phân công ở dòng ông C\n2. Nhập lại Mô tả công việc\n3. Bấm Lưu\n4. Mở lại cửa sổ Sửa phân công của ông C",
            "Mô tả công việc: Phụ trách bản vẽ mặt bằng",
            "- Lưu thành công\n- Mở lại thấy đúng nội dung vừa nhập",
        ),
        (
            "Ngày kết thúc trước ngày bắt đầu bị chặn", "P0",
            "Giải pháp G1 có thành viên ông C, ngày bắt đầu 25/05/2026.",
            "1. Bấm Sửa phân công ở dòng ông C\n2. Chọn Ngày kết thúc là 01/01/2026 (trước ngày bắt đầu)\n3. Bấm Lưu",
            "Ngày bắt đầu: 25/05/2026\nNgày kết thúc: 01/01/2026",
            "- ⚠️ Hệ thống báo lỗi đỏ ngay dưới ô Ngày kết thúc: Ngày kết thúc phải từ ngày bắt đầu trở đi\n- Cửa sổ KHÔNG đóng, dữ liệu đã nhập vẫn còn\n- Bảng không thay đổi",
        ),
        (
            "Ngày kết thúc trùng ngày bắt đầu vẫn lưu được", "P1",
            "Giải pháp G1 có thành viên ông C, ngày bắt đầu 25/05/2026.",
            "1. Bấm Sửa phân công ở dòng ông C\n2. Chọn Ngày kết thúc đúng bằng 25/05/2026\n3. Bấm Lưu",
            "Ngày kết thúc: 25/05/2026",
            "- Lưu thành công, không báo lỗi (bằng nhau là hợp lệ)",
        ),
        (
            "Đổi hạng mục phụ trách ở giải pháp có chia hạng mục", "P0",
            "Giải pháp G2 có 2 hạng mục HM1 và HM2. Ông D đang là thành viên HM1. Đăng nhập bằng tài khoản PM.",
            "1. Bấm Sửa phân công ở dòng ông D\n2. Đổi ô Hạng mục sang HM2\n3. Bấm Lưu\n4. Quan sát cột Hạng mục của dòng ông D",
            "Hạng mục: HM2",
            "- Lưu thành công\n- Cột Hạng mục của ông D đổi thành HM2\n- ⚠️ Vẫn là cùng một dòng, lịch sử ngày bắt đầu/kết thúc giữ nguyên, không bị tạo dòng mới",
        ),
        (
            "Chuyển vào hạng mục đã có sẵn chính người đó thì bị chặn", "P0",
            "Giải pháp G2: ông D là thành viên của CẢ HM1 và HM2 (2 dòng).",
            "1. Bấm Sửa phân công ở dòng ông D thuộc HM1\n2. Đổi Hạng mục sang HM2\n3. Bấm Lưu",
            "Hạng mục: HM2",
            "- ⚠️ Hệ thống báo lỗi đỏ dưới ô Hạng mục: Nhân sự đã có trong hạng mục này\n- Cửa sổ không đóng, không tạo dòng trùng",
        ),
        (
            "Giải pháp KHÔNG chia hạng mục thì không có ô Hạng mục", "P1",
            "Giải pháp G1 không chia hạng mục.",
            "1. Bấm Sửa phân công ở một dòng thành viên bất kỳ\n2. Quan sát các ô trong cửa sổ",
            "—",
            "- Cửa sổ chỉ có: Vai trò dự án, Mô tả công việc, Ngày bắt đầu, Ngày kết thúc\n- Không hiển thị ô Hạng mục",
        ),
        (
            "Bấm Hủy không lưu thay đổi", "P1",
            "Giải pháp G1 có thành viên ông C, ngày kết thúc đang trống.",
            "1. Bấm Sửa phân công ở dòng ông C\n2. Chọn Ngày kết thúc 31/12/2026\n3. Bấm Hủy\n4. Quan sát bảng",
            "Ngày kết thúc: 31/12/2026",
            "- Cửa sổ đóng\n- Cột Ngày kết thúc của ông C vẫn trống, thay đổi không được lưu",
        ),
    ]),

    ("Tab Nhân sự — Khóa và Mở khóa thành viên", [
        (
            "Khóa thành viên không còn công việc tồn đọng", "P0",
            "Giải pháp G1 có thành viên ông E, cột Task đang phụ trách hiển thị 0, không có vấn đề nào đang xử lý. Đăng nhập bằng tài khoản PM.",
            "1. Mở tab Nhân sự\n2. Bấm nút Khóa thành viên (ổ khóa) ở dòng ông E\n3. Đọc nội dung cửa sổ xác nhận\n4. Bấm Khóa",
            "—",
            "- Cửa sổ xác nhận ghi rõ tên ông E và cảnh báo sẽ không tạo thêm được hồ sơ hàng hóa, nhiệm vụ hay vấn đề trong giải pháp này, dữ liệu cũ vẫn giữ nguyên\n- Sau khi xác nhận: hiện thông báo Đã khóa thành viên\n- Cột Trạng thái của ông E đổi thành Đã khóa, nền xám\n- Nút giữa đổi sang biểu tượng Mở khóa",
        ),
        (
            "Khóa bị chặn khi còn công việc tồn đọng", "P0",
            "Giải pháp G1 có thành viên ông F đang được giao 2 nhiệm vụ ở trạng thái Đang thực hiện.",
            "1. Bấm nút Khóa thành viên ở dòng ông F\n2. Bấm Khóa ở cửa sổ xác nhận\n3. Quan sát",
            "—",
            "- ⚠️ Hệ thống KHÔNG khóa, mở cửa sổ Bàn giao công việc tồn đọng trước khi khóa thành viên\n- Câu mở đầu ghi đúng tên ông F và số 2 công việc chưa hoàn tất\n- Cột Trạng thái của ông F vẫn là Đang hoạt động",
        ),
        (
            "Danh sách công việc tồn đọng hiển thị đủ thông tin", "P0",
            "Như trên, đang mở cửa sổ Bàn giao công việc tồn đọng của ông F với 2 nhiệm vụ.",
            "1. Quan sát bảng Task trong cửa sổ",
            "—",
            "- Có nhóm Task kèm số lượng trong ngoặc\n- Mỗi dòng có: Mã, Tên task, Hạn, Bàn giao\n- Cột Hạn hiển thị dạng ngày/tháng/năm đủ 2 chữ số, ví dụ 31/05/2026",
        ),
        (
            "Cột Bàn giao cho biết việc đã nằm trong phiếu nào", "P0",
            "Ông F có 2 nhiệm vụ tồn đọng. Nhiệm vụ thứ nhất đã được ông F đưa vào một phiếu bàn giao đang Chờ duyệt; nhiệm vụ thứ hai chưa bàn giao.",
            "1. Bấm Khóa thành viên ở dòng ông F\n2. Xác nhận Khóa để mở cửa sổ tồn đọng\n3. Đọc cột Bàn giao của từng dòng",
            "—",
            "- Dòng nhiệm vụ 1: Đã nằm trong phiếu <mã phiếu> — Chờ duyệt\n- Dòng nhiệm vụ 2: Chưa bàn giao\n- ⚠️ Cả 2 vẫn tính là tồn đọng, vì công việc chưa thực sự sang tay người khác",
        ),
        (
            "Vấn đề đang xử lý cũng chặn khóa", "P0",
            "Giải pháp G1 có thành viên ông G không có nhiệm vụ nào, nhưng đang được giao 1 vấn đề ở trạng thái Đang xử lý.",
            "1. Bấm Khóa thành viên ở dòng ông G\n2. Xác nhận Khóa",
            "—",
            "- Cửa sổ tồn đọng mở ra, có nhóm Issue kèm số lượng 1\n- Dòng hiển thị đúng mã và tiêu đề vấn đề",
        ),
        (
            "Việc đã hoàn thành không tính là tồn đọng", "P1",
            "Giải pháp G1 có thành viên ông H chỉ được giao 3 nhiệm vụ, cả 3 đều ở trạng thái Hoàn thành.",
            "1. Bấm Khóa thành viên ở dòng ông H\n2. Xác nhận Khóa",
            "—",
            "- Khóa thành công ngay, không mở cửa sổ tồn đọng\n- Cột Trạng thái đổi thành Đã khóa",
        ),
        (
            "Bấm Nhắc bàn giao gửi thông báo cho thành viên", "P0",
            "Đang mở cửa sổ Bàn giao công việc tồn đọng của ông F (2 công việc). Chuẩn bị sẵn phiên đăng nhập của ông F ở trình duyệt khác.",
            "1. Bấm nút Nhắc bàn giao\n2. Chuyển sang phiên của ông F\n3. Mở chuông thông báo",
            "—",
            "- Người quản lý thấy thông báo Đã gửi nhắc bàn giao cho thành viên\n- Ông F nhận được 1 thông báo mới, nội dung nêu tên giải pháp và số công việc cần bàn giao, kèm tên người nhắc\n- Bấm vào thông báo mở đúng màn Bàn giao công việc",
        ),
        (
            "Bấm Mở màn Bàn giao chuyển đúng màn", "P1",
            "Đang mở cửa sổ Bàn giao công việc tồn đọng của ông F.",
            "1. Bấm nút Mở màn Bàn giao",
            "—",
            "- Hệ thống chuyển sang màn Bàn giao công việc",
        ),
        (
            "Khóa được ngay sau khi công việc đã sang tay người khác", "P0",
            "Ông F có 2 nhiệm vụ tồn đọng. Ông F lập phiếu bàn giao cho ông C, phiếu được duyệt và ông C đã tiếp nhận cả 2 nhiệm vụ.",
            "1. Đăng nhập lại bằng tài khoản PM\n2. Mở tab Nhân sự, quan sát cột Task đang phụ trách của ông F\n3. Bấm Khóa thành viên ở dòng ông F\n4. Xác nhận Khóa",
            "—",
            "- ⚠️ Lần này KHÔNG mở cửa sổ tồn đọng nữa\n- Khóa thành công, cột Trạng thái đổi thành Đã khóa",
        ),
        (
            "Mở khóa thành viên", "P0",
            "Giải pháp G1 có thành viên ông E đang ở trạng thái Đã khóa.",
            "1. Bấm nút Mở khóa thành viên ở dòng ông E\n2. Đọc cửa sổ xác nhận\n3. Bấm Mở khóa",
            "—",
            "- Cửa sổ xác nhận ghi đúng tên ông E\n- Hiện thông báo Đã mở khóa thành viên\n- Cột Trạng thái trở lại Đang hoạt động, nền xanh",
        ),
        (
            "Hủy ở cửa sổ xác nhận thì không khóa", "P1",
            "Giải pháp G1 có thành viên ông E đang hoạt động, không có việc tồn đọng.",
            "1. Bấm Khóa thành viên ở dòng ông E\n2. Bấm Hủy",
            "—",
            "- Cửa sổ đóng, không có thông báo nào\n- Cột Trạng thái vẫn là Đang hoạt động",
        ),
    ]),

    ("Tab Nhân sự — Xóa thành viên khỏi giải pháp", [
        (
            "Xóa được thành viên chưa phát sinh dữ liệu", "P0",
            "Giải pháp G1 có thành viên ông K vừa được thêm vào, chưa tạo và chưa được giao bất kỳ hồ sơ hàng hóa, nhiệm vụ hay vấn đề nào. Danh sách đang có 7 dòng.",
            "1. Bấm nút Xóa khỏi giải pháp ở dòng ông K\n2. Đọc cửa sổ xác nhận\n3. Bấm Xóa",
            "—",
            "- Cửa sổ xác nhận ghi đúng tên ông K, nút xác nhận màu đỏ\n- Hiện thông báo Đã xóa thành viên khỏi giải pháp\n- Danh sách còn 6 dòng, không còn ông K",
        ),
        (
            "Chặn xóa thành viên đã được giao nhiệm vụ", "P0",
            "Giải pháp G1 có thành viên ông F đang được giao 2 nhiệm vụ.",
            "1. Bấm Xóa khỏi giải pháp ở dòng ông F\n2. Bấm Xóa ở cửa sổ xác nhận",
            "—",
            "- ⚠️ Hệ thống báo đúng nguyên văn: Thành viên đã phát sinh dữ liệu (BOM/Task/Issue). Vui lòng sử dụng chức năng Khóa để bảo toàn dữ liệu hệ thống.\n- Ông F vẫn còn trong danh sách",
        ),
        (
            "Chặn xóa thành viên đã tạo hồ sơ hàng hóa", "P0",
            "Giải pháp G1 có thành viên ông L, không được giao nhiệm vụ nào nhưng đã từng tạo 1 hồ sơ hàng hóa trong giải pháp.",
            "1. Bấm Xóa khỏi giải pháp ở dòng ông L\n2. Bấm Xóa",
            "—",
            "- Hệ thống chặn, báo đúng câu hướng dẫn dùng chức năng Khóa\n- Ông L vẫn còn trong danh sách",
        ),
        (
            "Việc đã hoàn thành vẫn chặn xóa", "P0",
            "Giải pháp G1 có thành viên ông H từng được giao 3 nhiệm vụ, cả 3 đã Hoàn thành.",
            "1. Bấm Xóa khỏi giải pháp ở dòng ông H\n2. Bấm Xóa",
            "—",
            "- ⚠️ Vẫn bị chặn, dù việc đã xong: quy tắc là chỉ xóa được người CHƯA TỪNG phát sinh dữ liệu\n- Hệ thống gợi ý dùng chức năng Khóa",
        ),
        (
            "Hủy ở cửa sổ xác nhận thì không xóa", "P1",
            "Giải pháp G1 có thành viên ông K chưa phát sinh dữ liệu.",
            "1. Bấm Xóa khỏi giải pháp ở dòng ông K\n2. Bấm Hủy",
            "—",
            "- Cửa sổ đóng, ông K vẫn còn trong danh sách",
        ),
        (
            "Xóa thành viên đang bị khóa", "P1",
            "Giải pháp G1 có thành viên ông E ở trạng thái Đã khóa, chưa phát sinh dữ liệu nào.",
            "1. Bấm Xóa khỏi giải pháp ở dòng ông E\n2. Bấm Xóa",
            "—",
            "- Xóa thành công, ông E biến mất khỏi danh sách",
        ),
    ]),

    ("Tab Nhân sự — Hiệu lực của trạng thái Đã khóa", [
        (
            "Thành viên đã khóa không mở được giải pháp", "P0",
            "Ông E là thành viên giải pháp G1 và đang ở trạng thái Đã khóa. Đăng nhập bằng tài khoản ông E.",
            "1. Vào Dự án TKT → Danh sách làm giải pháp\n2. Bấm Quản lý ở giải pháp G1",
            "—",
            "- ⚠️ Hệ thống từ chối, báo đã bị khóa khỏi giải pháp này và hướng dẫn liên hệ PM hoặc Trưởng phòng giải pháp\n- Không mở được màn quản lý giải pháp G1",
        ),
        (
            "Thành viên đã khóa vẫn vào được giải pháp KHÁC", "P0",
            "Ông E bị khóa ở giải pháp G1 nhưng là thành viên đang hoạt động của giải pháp G4.",
            "1. Đăng nhập bằng tài khoản ông E\n2. Mở màn quản lý giải pháp G4",
            "—",
            "- ⚠️ Vào bình thường, hiển thị đầy đủ dữ liệu G4\n- Khóa chỉ có hiệu lực trong đúng giải pháp bị khóa",
        ),
        (
            "Thành viên đã khóa không tạo được nhiệm vụ trong giải pháp đó", "P0",
            "Ông E bị khóa ở giải pháp G1. Đăng nhập bằng tài khoản ông E.",
            "1. Dùng công cụ kiểm thử API gọi thẳng chức năng Tạo nhiệm vụ, gắn vào giải pháp G1, bỏ qua giao diện",
            "—",
            "- Hệ thống từ chối, báo đã bị khóa khỏi giải pháp này\n- ⚠️ Bị chặn ngay cả khi dữ liệu gửi lên còn thiếu trường — tức là chặn quyền diễn ra trước bước kiểm tra dữ liệu",
        ),
        (
            "Thành viên đã khóa không tạo được vấn đề và hồ sơ hàng hóa trong giải pháp đó", "P0",
            "Ông E bị khóa ở giải pháp G1. Đăng nhập bằng tài khoản ông E.",
            "1. Dùng công cụ kiểm thử API gọi thẳng chức năng Tạo vấn đề gắn vào giải pháp G1\n2. Làm tương tự với chức năng Tạo hồ sơ hàng hóa",
            "—",
            "- Cả hai đều bị từ chối với cùng thông báo đã bị khóa khỏi giải pháp",
        ),
        (
            "Người không phải thành viên không bị ảnh hưởng", "P0",
            "Giải pháp G1 có ít nhất 1 thành viên đang bị khóa. Đăng nhập bằng tài khoản PM của G1 (PM không nằm trong danh sách thành viên).",
            "1. Mở màn quản lý giải pháp G1\n2. Mở lần lượt các tab",
            "—",
            "- ⚠️ PM vào bình thường, không bị chặn\n- Việc khóa thành viên không làm ảnh hưởng người quản lý",
        ),
        (
            "Mở khóa xong thì truy cập lại được ngay", "P0",
            "Ông E đang bị khóa ở giải pháp G1.",
            "1. PM mở tab Nhân sự và bấm Mở khóa cho ông E\n2. Đăng nhập bằng tài khoản ông E\n3. Mở màn quản lý giải pháp G1",
            "—",
            "- Ông E vào được bình thường\n- Tạo nhiệm vụ trong giải pháp G1 trở lại bình thường",
        ),
        (
            "Dữ liệu cũ của người bị khóa vẫn còn nguyên", "P0",
            "Ông F đã bàn giao hết việc và bị khóa ở giải pháp G1. Trước đó ông F từng tạo 1 hồ sơ hàng hóa và 2 nhiệm vụ trong G1.",
            "1. Đăng nhập bằng tài khoản PM\n2. Mở các tab Nhiệm vụ và hồ sơ hàng hóa của giải pháp G1\n3. Tìm các bản ghi do ông F tạo",
            "—",
            "- ⚠️ Toàn bộ dữ liệu cũ còn nguyên, vẫn ghi nhận ông F là người tạo\n- Khóa chỉ chặn thao tác mới, không xóa lịch sử",
        ),
    ]),

    ("Tab Nhân sự — Hiển thị và dữ liệu bảng", [
        (
            "Cột Trạng thái hiển thị đúng chữ và màu", "P0",
            "Giải pháp G1 có 1 thành viên Đã khóa và 5 thành viên đang hoạt động.",
            "1. Mở tab Nhân sự\n2. Quan sát cột Trạng thái",
            "—",
            "- Thành viên bình thường: Đang hoạt động, nền xanh dương\n- Thành viên bị khóa: Đã khóa, nền xám\n- ⚠️ Trước đây mọi dòng đều hiện Active kể cả khi đã khóa — nay phải phân biệt được",
        ),
        (
            "Cột Ngày kết thúc lấy đúng của từng thành viên", "P0",
            "Giải pháp G1 có ngày kết thúc dự kiến là 30/06/2026. Thành viên ông C đã đặt ngày kết thúc riêng 31/12/2026; ông M chưa đặt ngày kết thúc.",
            "1. Mở tab Nhân sự\n2. So cột Ngày kết thúc của ông C và ông M",
            "—",
            "- Ông C: 31/12/2026\n- ⚠️ Ông M: dấu gạch ngang, KHÔNG được lấy ngày kết thúc của giải pháp điền vào",
        ),
        (
            "Thêm nhân sự mới rồi sửa ngay", "P1",
            "Giải pháp G1, đăng nhập bằng tài khoản PM.",
            "1. Bấm Thêm nhân sự, chọn một nhân sự chưa có trong giải pháp, chọn vai trò và ngày bắt đầu, bấm lưu\n2. Tìm dòng vừa thêm\n3. Bấm Sửa phân công, đặt ngày kết thúc, bấm Lưu",
            "Ngày bắt đầu: ngày hiện tại\nNgày kết thúc: 31/12/2026",
            "- Dòng mới có đủ 3 nút thao tác ngay, không cần tải lại trang\n- Sửa và lưu thành công",
        ),
        (
            "Sơ đồ cấu trúc nhân sự cập nhật theo", "P2",
            "Giải pháp G1 có 6 thành viên.",
            "1. Mở tab Nhân sự, ghi lại số ở nhãn Nhân sự phía trên sơ đồ\n2. Xóa 1 thành viên chưa phát sinh dữ liệu\n3. Quan sát lại sơ đồ và nhãn",
            "—",
            "- Số nhân sự giảm 1\n- Sơ đồ không còn ô của người vừa xóa\n- Không có lỗi hiển thị",
        ),
        (
            "Hai người cùng thao tác trên một thành viên", "P1",
            "Hai tài khoản PM và Trưởng phòng giải pháp cùng mở tab Nhân sự của G1 trên 2 trình duyệt, cùng nhìn thấy ông K.",
            "1. Tài khoản 1 xóa ông K khỏi giải pháp\n2. Tài khoản 2 (chưa tải lại trang) bấm Sửa phân công ở dòng ông K và bấm Lưu",
            "—",
            "- ⚠️ Tài khoản 2 nhận thông báo không tìm thấy thành viên trong giải pháp này, đề nghị tải lại danh sách\n- Không treo trang, không tạo lại bản ghi đã xóa",
        ),
    ]),
]


THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FONT = "Times New Roman"


def build(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "TC 11354"

    # --- Header ---
    for col, name in enumerate(HEADERS, start=1):
        c = ws.cell(row=1, column=col, value=name)
        c.font = Font(name=FONT, size=12, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BORDER
    ws.row_dimensions[1].height = 40

    row = 2
    no = START_NO
    total = 0
    p0 = 0

    for group_name, cases in GROUPS:
        # Dòng tiêu đề nhóm: chỉ điền cột Chức năng, giống cách tab hiện tại đang làm
        c = ws.cell(row=row, column=4, value=group_name)
        c.font = Font(name=FONT, size=12, bold=True, color="1F4E79")
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=row, column=col)
            cell.fill = PatternFill("solid", fgColor="D6E4F0")
            cell.border = BORDER
        ws.row_dimensions[row].height = 26
        row += 1

        for func, prio, pre, steps, data, expected in cases:
            values = [
                MODULE, "", f"TC-ROLE-{no}", func, prio,
                pre, steps, data, expected,
                "", "Not Executed", "Not Executed", "Not Executed",
                "", "", "", "", "", "",
            ]
            for col, v in enumerate(values, start=1):
                cell = ws.cell(row=row, column=col, value=v)
                cell.font = Font(name=FONT, size=12)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                cell.border = BORDER
            ws.row_dimensions[row].height = 34
            row += 1
            no += 1
            total += 1
            if prio == "P0":
                p0 += 1

    widths = {
        'A': 22, 'B': 18, 'C': 14, 'D': 34, 'E': 9, 'F': 34, 'G': 40, 'H': 20,
        'I': 52, 'J': 20, 'K': 14, 'L': 14, 'M': 14, 'N': 16, 'O': 11, 'P': 11,
        'Q': 11, 'R': 14, 'S': 16,
    }
    for k, v in widths.items():
        ws.column_dimensions[k].width = v

    last = row + 40
    dv1 = DataValidation(type="list", formula1='"Passed,Failed,Pending,Not Executed"', allow_blank=True)
    dv1.showDropDown = False
    ws.add_data_validation(dv1)
    dv1.add(f"K2:M{last}")

    dv2 = DataValidation(type="list", formula1='"P,F,PE"', allow_blank=True)
    dv2.showDropDown = False
    ws.add_data_validation(dv2)
    dv2.add(f"O2:Q{last}")

    wb.save(path)
    return total, p0, no - 1


if __name__ == "__main__":
    out = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                       "testcase - 11354 Quan ly nhan su giai phap.xlsx")
    total, p0, last_id = build(out)

    # Bộ kiểm tra thuật ngữ kỹ thuật (bắt buộc theo skill testcase-documenter)
    import re
    BANNED = [
        r"`[a-z_]{3,}`", r"\bBE\b", r"\bFE\b", r"\bHTTP\b",
        r"trả (400|403|404|422)", r"\b(400|403|404|422)\b",
        r"permission id", r"/api/v1", r"\bAPI /", r"localStorage",
        r"number_format", r"meta\.", r"sort_by", r"per_page",
        r"role_has_permissions", r"current_company_role",
    ]
    text_all = "\n".join(
        "\n".join([g] + ["\n".join(map(str, tc)) for tc in cases])
        for g, cases in GROUPS
    )
    found = {p: len(re.findall(p, text_all)) for p in BANNED if re.findall(p, text_all)}
    print("!!! CON THUAT NGU KY THUAT:", found) if found else print("OK - sach")

    print(f"File: {out}")
    print(f"Tong TC: {total} | P0: {p0} ({p0 * 100 // total}%) | Ma TC: TC-ROLE-{START_NO} -> TC-ROLE-{last_id}")
