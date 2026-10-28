# -*- coding: utf-8 -*-
"""HDSD bản GỌN — Phiếu đề nghị thu tiền (form chuẩn Danh mục quốc gia, 28/10/2026).

Chạy: python gen_hdsd.py  → out/HDSD_Phiếu đề nghị thu tiền.docx (+ .pdf, render/ để soát)
Khuôn: màn 27 (finance-product-import-direct-transfer/form_qg/gen_hdsd.py); phần chung ở
../../phieu-theo-form-quoc-gia/hdsd_phieu_lib.py. Thuật ngữ + dòng phiên bản cũ đọc từ bản Drive.
"""
import os
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'phieu-theo-form-quoc-gia'))
from hdsd_phieu_lib import PhieuHdsd, finalize_word, toc_lines, render_check  # noqa: E402

TEN = 'Phiếu đề nghị thu tiền'
SHOTS = os.path.join(HERE, 'hdsd_shots')
OUT = os.path.join(HERE, 'out', 'HDSD_%s.docx' % TEN)
DRIVE = os.path.join(HERE, 'drive_0710', 'HDSD_Drive.docx')


def drive_tables():
    d = docx.Document(DRIVE)
    t0, t1 = d.tables[0], d.tables[1]
    thuat_ngu = [[r.cells[0].text.strip(), r.cells[1].text.strip()] for r in t0.rows[1:]]
    cu = [[r.cells[0].text.strip(), r.cells[1].text.strip(), r.cells[2].text.strip()] for r in t1.rows[1:]]
    return thuat_ngu, cu


QUYEN = [
    ['Xem danh sách, xem chi tiết, in phiếu, xem lịch sử',
     'Một trong các quyền: Xem tất cả phiếu đề nghị thu của tổng công ty / của công ty / của phòng ban / của bộ '
     'phận. Phiếu do chính mình lập luôn xem được.'],
    ['Lập phiếu', 'Không cần quyền riêng.'],
    ['Sửa, xóa phiếu', 'Người lập phiếu, phiếu đang ở trạng thái Đang tạo hoặc Không duyệt.'],
    ['Màn chờ duyệt, Tạo phiếu thu, Không duyệt',
     'Quyền Kế toán thanh toán; phiếu đang ở trạng thái Chờ KT duyệt.'],
]


def build():
    thuat_ngu, cu = drive_tables()
    phien_ban = cu + [['1.1', '28/10/2026', 'Chuyển sang form chuẩn']]
    d = PhieuHdsd(TEN, SHOTS)
    d.tong_quan(thuat_ngu, phien_ban, QUYEN, os.path.join(HERE, 'uml', 'overview_srs.png'))

    # PHẦN 1
    d.truy_cap(
        menu=[('Tài chính', 'menu_phanhe'), ('Khởi tạo', 'menu_nhom'),
              ('Khởi tạo phiếu yêu cầu - Công nợ - Thu - Chi', 'menu_khoitao'), ('Đề nghị thu tiền', 'menu_muc')],
        o_nhanh='mã phiếu',
        bo_loc=[('Công ty', 'o_congty'), ('Phòng ban', 'o_phongban'), ('Bộ phận', 'o_bophan'),
                ('Loại thu', 'o_loaithu'), ('Trạng thái', 'o_trangthai'), ('Số đơn hàng/hợp đồng', 'o_hopdong'),
                ('Khách hàng', 'o_khachhang'), ('Nhà cung cấp', 'o_ncc'), ('Người tạo', 'o_nguoitao'),
                ('Số tiền đề nghị', 'o_sotien'), ('Ngày tạo', 'o_ngaytao')])

    # PHẦN 2 — Lập phiếu
    d.phan('LẬP PHIẾU ĐỀ NGHỊ THU TIỀN')
    d.step(d.vao)
    d.step(['Bước 2: Bấm ', '@btn_taomoi'])
    d.img('10_create', 'Màn hình Thêm phiếu đề nghị thu tiền')
    d.step('Bước 3: Nhập thông tin chung:')
    d.sub('Loại thu (bắt buộc): chọn Thu bán hàng hoặc Thu nhà cung cấp; mặc định Thu bán hàng.')
    d.sub('Loại tiền (bắt buộc): mặc định VNĐ; chọn ngoại tệ thì Tỷ giá tự điền và sửa được.')
    d.sub('Tỷ giá (VND) (bắt buộc): số lớn hơn 0; loại tiền VNĐ thì cố định bằng 1.')
    d.sub('Lý do thu (bắt buộc khi gửi duyệt, tùy chọn khi lưu nháp): nhập chữ.')
    d.step(['Bước 4: Ở khối Chi tiết, bấm ', '@btn_themdong', ' để thêm một dòng.'])
    d.step(['Bước 5: Bấm vào ô ', '@o_chonkhachhang', '. Hệ thống mở cửa sổ “Chọn khách hàng” (loại Thu nhà cung cấp '
            'thì mở cửa sổ “Chọn nhà cung cấp”).'])
    d.img('11_customer', 'Cửa sổ Chọn khách hàng')
    d.step('Bước 6: Nhập Tên / Mã khách hàng, Mã số thuế hoặc Số điện thoại, bấm Tìm kiếm rồi bấm vào dòng khách hàng '
           'cần chọn.')
    d.sub(['Hoặc bấm ', '@btn_dong_kh', ' để không chọn.'])
    d.step(['Bước 7: Bấm vào ô ', '@o_chonhopdong', ', bấm vào dòng hợp đồng cần thu. Số tiền còn nợ tự điền.'])
    d.sub('Hợp đồng đã có ở dòng khác của phiếu không chọn lại được.')
    d.step('Bước 8: Nhập Số tiền đề nghị thu (bắt buộc, số không âm; dấu phẩy ngăn hàng nghìn, ví dụ 1,200,000) và '
           'Ghi chú của dòng (tùy chọn).')
    d.sub(['Hoặc bấm ', '@btn_xoadong', ' để bỏ dòng khỏi phiếu.'])
    d.step('Bước 9: Nhập Ghi chú của phiếu (tùy chọn) ở khối Ghi chú.')
    d.step(['Bước 10: Bấm ', '@btn_luuguiduyet', ', chọn ', '@btn_confirm_xacnhan', ' để gửi Kế toán thanh toán duyệt '
            '(bắt buộc có Lý do thu và ít nhất một dòng chi tiết),'])
    d.sub(['Hoặc ', '@btn_luunhap', ' để lưu nháp.'])
    d.sub(['Hoặc ', '@btn_luutieptuc', ' để lưu nháp rồi lập tiếp phiếu khác.'])
    d.sub(['Hoặc ', '@btn_quaylai', ' để không lưu thông tin.'])

    # PHẦN 3 — Chỉnh sửa
    d.phan('CHỈNH SỬA')
    d.step(d.vao)
    d.step(['Bước 2: Bấm nút ', '@btn_sua_row', ' ở cột Hành động của phiếu Đang tạo hoặc Không duyệt do mình lập.'])
    d.sub(['Hoặc bấm vào Mã phiếu để mở màn chi tiết, chọn ', '@btn_sua'])
    d.step('Bước 3: Hệ thống hiển thị màn hình “Sửa phiếu đề nghị thu tiền”.')
    d.img('19_edit', 'Màn hình Sửa phiếu đề nghị thu tiền')
    d.step('Bước 4: Nhập thông tin cần sửa (quy tắc nhập như Lập phiếu ở PHẦN 2). Phiếu Không duyệt: đọc lý do ở ô '
           'Ghi chú rồi sửa lại.')
    d.step(['Bước 5: Bấm ', '@btn_luuguiduyet', ' để lưu và gửi duyệt,'])
    d.sub(['Hoặc ', '@btn_luunhap', ' để lưu nháp.'])
    d.sub(['Hoặc ', '@btn_quaylai', ' nếu muốn hủy.'])

    # PHẦN 4 — Xóa (Xóa nằm trong nút ba chấm khi dòng có trên 3 nút)
    d.phan('XÓA')
    d.step(['Bước 1: Ở cột Hành động của phiếu Đang tạo hoặc Không duyệt do mình lập, bấm ', '@btn_bacham_row',
            ' rồi chọn ', '@item_xoa'])
    d.sub(['Hoặc bấm vào Mã phiếu để mở màn chi tiết, chọn ', '@btn_xoa'])
    d.step('Bước 2: Hệ thống hiện hộp xác nhận “Xác nhận xóa”')
    d.img('13_delete', 'Hộp xác nhận xóa phiếu')
    d.step(['Bước 3: Bấm ', '@btn_confirm_xoa', ' để xác nhận xóa'])
    d.sub(['Hoặc bấm ', '@btn_huy_confirm', ' nếu bấm nhầm'])

    # PHẦN 5 — Xem chi tiết
    d.phan('XEM CHI TIẾT PHIẾU')
    d.step(d.vao)
    d.step('Bước 2: Bấm vào Mã phiếu. Hệ thống mở màn “Chi tiết phiếu đề nghị thu tiền”.')
    d.img('17_detail', 'Màn hình Chi tiết phiếu đề nghị thu tiền')
    d.step(['Bước 3: Bấm ', '@btn_quaylai', ' để về màn danh sách.'])

    # PHẦN 6 — Màn chờ duyệt + Tạo phiếu thu
    d.phan('DUYỆT PHIẾU (TẠO PHIẾU THU)')
    d.step(['Bước 1: Từ phân hệ Tài chính ', '@menu_phanhe', ', tại menu Chờ duyệt ', '@menu_choduyet', ', nhóm '
            'Phê duyệt - Công nợ - Thu - Chi ', '@menu_pheduyet', ', chọn Phiếu đề nghị thu tiền chờ duyệt ',
            '@menu_muc_choduyet'])
    d.img('21_pending', 'Màn hình Phiếu đề nghị thu tiền chờ duyệt')
    d.step('Bước 2: Bấm vào Mã phiếu để mở màn chi tiết, kiểm tra các dòng chi tiết.')
    d.step(['Bước 3: Bấm ', '@btn_taophieuthu', '. Hệ thống chuyển sang màn lập Phiếu thu, gắn sẵn phiếu đề nghị này.'])
    d.sub(['Hoặc bấm ', '@btn_taophieuthu_row', ' ở cột Hành động của phiếu.'])
    d.step('Bước 4: Lập và gửi duyệt Phiếu thu; phiếu đề nghị chuyển sang trạng thái Đã tạo phiếu thu.')

    # PHẦN 7 — Không duyệt
    d.phan('KHÔNG DUYỆT PHIẾU')
    d.step('Bước 1: Mở màn chi tiết của phiếu Chờ KT duyệt (như Bước 1, 2 PHẦN 6).')
    d.step(['Bước 2: Bấm ', '@btn_khongduyet', '. Hệ thống mở cửa sổ “Không duyệt phiếu”.'])
    d.img('15_reject', 'Cửa sổ Không duyệt phiếu')
    d.step('Bước 3: Nhập Lý do không duyệt (bắt buộc).')
    d.step(['Bước 4: Bấm ', '@btn_confirm_khongduyet', ' để xác nhận'])
    d.sub(['Hoặc bấm ', '@btn_dong_reject', ' nếu muốn hủy.'])

    # PHẦN 8 — In
    d.phan('IN PHIẾU')
    d.step(['Bước 1: Ở cột Hành động, bấm ', '@btn_in_row'])
    d.sub(['Hoặc tại màn chi tiết phiếu, bấm ', '@btn_in'])
    d.step('Bước 2: Hệ thống mở cửa sổ xem trước bản in.')
    d.img('18_print', 'Cửa sổ Xem trước phiếu đề nghị thu tiền')
    d.step(['Bước 3: Bấm ', '@btn_in_preview', ' để in.'])

    # PHẦN 9..11
    d.lich_su('Mã phiếu', img='16_history', cap='Cửa sổ Lịch sử thay đổi của phiếu')
    d.cai_dat_bo_loc()
    d.tuy_chinh_cot()
    return d.save(OUT)


if __name__ == '__main__':
    out = build()
    pdf = os.path.splitext(out)[0] + '.pdf'
    size = finalize_word(out, pdf)
    print('docx', out, size)
    print('TOC:')
    for line in toc_lines(out):
        print('  ', line)
    n, warn, sheets = render_check(pdf, os.path.join(HERE, 'out', 'render'))
    print('pages', n)
    for w_ in warn:
        print('WARN', w_)
    for s in sheets:
        print(s)
