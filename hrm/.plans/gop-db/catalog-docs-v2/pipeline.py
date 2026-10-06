"""python3 pipeline.py <slug> cap|imp|copy|build|up
cap   → sinh .playwright-mcp/cap_<slug>.js từ CFG['capture'] (chạy bằng browser_run_code_unsafe)
imp   → dựng .playwright-mcp/cat/<slug>/import_test.xlsx từ file mẫu vừa tải + CFG['import_test'], sinh imp_<slug>.js
copy  → chép ảnh + icon chụp được vào catalog-docs-v2/<slug>/
build → dựng HDSD + SRS (catalog_v2 all)
up    → đẩy Drive giữ ID (upload.py)
Chạy từ thư mục HRM.
"""
import json, os, shutil, subprocess, sys
HRM = '/Users/manhcuong/Desktop/dns/HRM'
H = os.path.join(HRM, '.plans/gop-db/catalog-docs-v2')
LIB = os.path.join(HRM, '.plans/gop-db/_catalog_docs_lib')
sys.path.insert(0, LIB)
from catalog_v2 import load_cfg

slug, step = sys.argv[1], sys.argv[2]
cfg = load_cfg(slug)
cat = os.path.join(HRM, '.playwright-mcp', 'cat', slug)
os.chdir(HRM)

if step == 'cappage':
    c = dict(cfg['capture']); c.pop('manual', None)
    c.setdefault('lockSearch', c.get('lockName'))
    c['base'] = 'http://127.0.0.1:3002'; c['dir'] = '.playwright-mcp/cat/%s/shots/' % slug
    os.makedirs(os.path.join(cat, 'shots'), exist_ok=True); os.makedirs(os.path.join(cat, 'icons'), exist_ok=True)
    tpl = open(os.path.join(LIB, 'capture_page_tpl.js')).read().replace('__PARAMS__', json.dumps(c, ensure_ascii=False))
    js = os.path.join(HRM, '.playwright-mcp', 'capp_%s.js' % slug); open(js, 'w').write(tpl); print(js)
elif step == 'cap':
    subprocess.run(['python3', os.path.join(LIB, 'make_capture.py'), slug, json.dumps(cfg['capture'], ensure_ascii=False)], check=True)
elif step == 'imp':
    import openpyxl
    src = os.path.join(cat, 'shots', 'mau_import.xlsx')
    wb = openpyxl.load_workbook(src)
    ws = wb.worksheets[0]
    has_stt = str(ws.cell(1, 1).value or '').strip().upper() == 'STT'
    # hàng 2 là dòng mô tả (chữ nghiêng) theo khuôn chung; file mẫu cũ không có → dữ liệu bắt đầu từ hàng 2
    row2_desc = any(str(rg).split(':')[0].endswith('2') and str(rg).split(':')[0][:-1].isalpha() for rg in ws.merged_cells.ranges) \
        or (ws.cell(2, 2).font and ws.cell(2, 2).font.i) or (ws.cell(2, 1).font and ws.cell(2, 1).font.i)
    first = 3 if row2_desc else 2
    for r in range(first, ws.max_row + 1):
        for c in range(1, ws.max_column + 1):
            ws.cell(r, c).value = None
    rows_in = cfg['import_test']['rows']
    # config đã tự ghi cột STT (giá trị đầu = 1, 2, …) thì không chèn thêm
    stt_in_cfg = all(str(r[0]).strip() == str(i + 1) for i, r in enumerate(rows_in) if r)
    for i, row in enumerate(rows_in):
        c0 = 1
        if has_stt and not stt_in_cfg:
            ws.cell(first + i, 1).value = i + 1
            c0 = 2
        for j, v in enumerate(row):
            ws.cell(first + i, c0 + j).value = v
    out = os.path.join(cat, 'import_test.xlsx')
    wb.save(out)
    P = {'route': cfg['capture']['route'], 'base': 'http://127.0.0.1:3002', 'dir': '.playwright-mcp/cat/%s/shots/' % slug,
         'file': out, 'doImport': True}
    tpl = open(os.path.join(LIB, 'import_tpl.js')).read().replace('__PARAMS__', json.dumps(P, ensure_ascii=False))
    js = os.path.join(HRM, '.playwright-mcp', 'imp_%s.js' % slug)
    open(js, 'w').write(tpl)
    print(js, 'header:', [ws.cell(1, c).value for c in range(1, ws.max_column + 1)])
elif step == 'copy':
    for sub in ('shots', 'icons'):
        d = os.path.join(H, slug, sub)
        os.makedirs(d, exist_ok=True)
        for f in os.listdir(os.path.join(cat, sub)):
            if f.endswith('.png'):
                shutil.copy(os.path.join(cat, sub, f), d)
    print(sorted(os.listdir(os.path.join(H, slug, 'shots'))))
elif step == 'build':
    subprocess.run(['python3', os.path.join(LIB, 'catalog_v2.py'), slug, 'all'], check=True)
elif step == 'up':
    subprocess.run(['python3', os.path.join(H, 'upload.py'), slug], check=True)
