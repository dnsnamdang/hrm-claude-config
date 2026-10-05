"""HDSD Danh mục khách hàng — bản 1.3 (28/09/2026): viết lại chi tiết theo khuôn HDSD Quốc gia.

Sửa THẲNG trên bản tải từ Drive (đã có người sửa tay) chứ không build lại từ config, để giữ phần
người khác đã sửa. Clone đúng paragraph/table có sẵn trong file làm khuôn định dạng.

    python3 patch_hdsd_v13.py <vào.docx> <ra.docx>

Phạm vi:
  - PHẦN 1 mục 4 Cột Hành động: viết lại, mỗi nút kèm ảnh nút
  - PHẦN 2 / 9 / 10: chèn ảnh nút vào các bước đang chỉ có chữ
  - PHẦN 4 → PHẦN 8: viết lại toàn bộ (Chỉnh sửa, Khoá/Mở khoá, Lịch sử, Chi tiết, Quản lý)
"""
import copy, os, re, sys
from docx import Document
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

SRC, OUT = sys.argv[1], sys.argv[2]
HERE = os.path.dirname(os.path.abspath(__file__))
SHOTS = os.path.join(HERE, 'shots')
ICON_DIRS = [os.path.join(HERE, 'icons'), os.path.join(HERE, '..', 'icons')]
d = Document(SRC)
body = d.element.body
XML_SPACE = '{http://www.w3.org/XML/1998/namespace}space'


def P(el):
    return Paragraph(el, d)


def paras():
    return list(body.iterchildren(qn('w:p')))


def find_p(prefix, style=None, after=None):
    started = after is None
    for el in body.iterchildren():
        if el is after:
            started = True
            continue
        if not started or el.tag != qn('w:p'):
            continue
        p = P(el)
        if p.text.startswith(prefix) and (style is None or p.style.name == style):
            return el
    raise SystemExit('Không thấy đoạn: ' + prefix)


def find_table(header_first):
    for t in d.tables:
        if t.rows[0].cells[0].text.strip() == header_first:
            return t._tbl
    raise SystemExit('Không thấy bảng: ' + header_first)


# ── khuôn định dạng (lấy trước khi xoá) ──
H1_4 = find_p('PHẦN 4: CHỈNH SỬA', 'Heading 1')
T_PB = copy.deepcopy(H1_4.getprevious())                  # đoạn ngắt trang trước Heading 1
assert 'w:type="page"' in T_PB.xml
T_H1 = copy.deepcopy(H1_4)
T_H2 = copy.deepcopy(H1_4.getnext())
T_STEP = copy.deepcopy(find_p('Bước 1: Truy cập vào màn hình Danh mục khách hàng.', after=H1_4))
T_BUL = copy.deepcopy(find_p('Mã khách hàng không sửa được'))
_cap = find_p('Hình 8:')
T_CAP = copy.deepcopy(_cap)
T_IMG_PPR = copy.deepcopy(_cap.getprevious().find(qn('w:pPr')))
T_TBL2 = copy.deepcopy(find_table('Nút'))
T_TBL5 = copy.deepcopy(find_table('Trường'))
T_RPR = copy.deepcopy(T_BUL.find('.//' + qn('w:rPr'), ) if False else None)
_r0 = next(r for r in T_BUL.iter(qn('w:r')) if r.find(qn('w:t')) is not None)
T_RPR = copy.deepcopy(_r0.find(qn('w:rPr')))
T_PLAIN = copy.deepcopy(T_STEP)
_ppr = T_PLAIN.find(qn('w:pPr'))
_ppr.remove(_ppr.find(qn('w:numPr')))
_ppr.remove(_ppr.find(qn('w:ind')))


def icon_path(name):
    for dd in ICON_DIRS:
        p = os.path.join(dd, name + '.png')
        if os.path.exists(p):
            return p
    raise SystemExit('Thiếu icon: ' + name)


TOKEN = re.compile(r'(\{icon:[a-z0-9_]+\}|\*\*[^*]+\*\*)')


def clear(el):
    for c in list(el):
        if c.tag != qn('w:pPr'):
            el.remove(c)
    return el


def add_text(el, text, bold=False):
    r = el.makeelement(qn('w:r'), {})
    if T_RPR is not None:
        r.append(copy.deepcopy(T_RPR))
    if bold:
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = r.makeelement(qn('w:rPr'), {})
            r.insert(0, rpr)
        rpr.append(rpr.makeelement(qn('w:b'), {}))
    t = r.makeelement(qn('w:t'), {})
    t.text = text
    t.set(XML_SPACE, 'preserve')
    r.append(t)
    el.append(r)


def fill(el, text):
    """Ghi chuỗi có {icon:x} và **đậm** vào paragraph el (giữ pPr + định dạng chữ của run đầu)."""
    global T_RPR
    keep = T_RPR
    r0 = next((r for r in el.iter(qn('w:r')) if r.find(qn('w:t')) is not None), None)
    if r0 is not None and r0.find(qn('w:rPr')) is not None:
        T_RPR = copy.deepcopy(r0.find(qn('w:rPr')))
    clear(el)
    for part in TOKEN.split(text):
        if not part:
            continue
        if part.startswith('{icon:'):
            run = P(el).add_run()
            run.add_picture(icon_path(part[6:-1]), height=Inches(0.22))
        elif part.startswith('**'):
            add_text(el, part[2:-2], bold=True)
        else:
            add_text(el, part)
    T_RPR = keep
    return el


class Seq:
    def __init__(self, anchor):
        self.a = anchor

    def put(self, el):
        self.a.addnext(el)
        self.a = el
        return el

    def h1(self, text):
        self.put(copy.deepcopy(T_PB))
        el = copy.deepcopy(T_H1)
        for b in el.findall(qn('w:bookmarkStart')) + el.findall(qn('w:bookmarkEnd')):
            el.remove(b)
        return self.put(fill(el, text))

    def h2(self, text):
        return self.put(fill(copy.deepcopy(T_H2), text))

    def h3(self, text):
        el = fill(copy.deepcopy(T_H2), text)
        el.find(qn('w:pPr')).find(qn('w:pStyle')).set(qn('w:val'), 'Heading3')
        return self.put(el)

    def p(self, text):
        return self.put(fill(copy.deepcopy(T_PLAIN), text))

    def step(self, text):
        return self.put(fill(copy.deepcopy(T_STEP), text))

    def b(self, text):
        return self.put(fill(copy.deepcopy(T_BUL), text))

    def img(self, fname, caption):
        path = os.path.join(SHOTS, fname)
        assert os.path.exists(path), path
        tmp = d.add_paragraph()
        tmp.add_run().add_picture(path, width=Inches(6.0))
        tmp._p.insert(0, copy.deepcopy(T_IMG_PPR))
        self.put(tmp._p)
        self.put(fill(copy.deepcopy(T_CAP), 'Hình 0: ' + caption))

    def table(self, rows, five=False):
        tpl = T_TBL5 if five else T_TBL2
        tbl = copy.deepcopy(tpl)
        trs = tbl.findall(qn('w:tr'))
        row_tpl = trs[1]
        for tr in trs[1:]:
            tbl.remove(tr)
        for _ in rows[1:]:
            tbl.append(copy.deepcopy(row_tpl))
        t = Table(tbl, d)
        for ri, r in enumerate(rows):
            for ci, txt in enumerate(r):
                cell_text(t.rows[ri].cells[ci], txt)
        return self.put(tbl)


def cell_text(cell, text):
    ps = cell._tc.findall(qn('w:p'))
    for p in ps[1:]:
        cell._tc.remove(p)
    p = ps[0]
    rpr = None
    r0 = p.find(qn('w:r'))
    if r0 is not None and r0.find(qn('w:rPr')) is not None:
        rpr = copy.deepcopy(r0.find(qn('w:rPr')))
    clear(p)
    r = p.makeelement(qn('w:r'), {})
    if rpr is not None:
        r.append(rpr)
    t = r.makeelement(qn('w:t'), {})
    t.text = text
    t.set(XML_SPACE, 'preserve')
    r.append(t)
    p.append(r)


def cut(start, end):
    """Xoá các phần tử body từ start (gồm) tới end (không gồm); trả phần tử đứng trước start."""
    before = start.getprevious()
    el = start
    while el is not end:
        nxt = el.getnext()
        body.remove(el)
        el = nxt
    return before


# ═════════════ PHẦN 1 — mục 4 Cột Hành động ═════════════
a = find_p('4. Cột Hành động', 'Heading 2')
z = find_p('5. Phân trang và sắp xếp', 'Heading 2')
s = Seq(a)
el = a.getnext()
while el is not z:
    nxt = el.getnext()
    body.remove(el)
    el = nxt
s.p('Trên mỗi dòng, cột Hành động có tối đa bốn nút. Nút nào hiện phụ thuộc quyền của Người dùng và '
    'trạng thái của khách hàng:')
s.b('{icon:btn_sua_kh} **Sửa** — mở trang Chỉnh sửa khách hàng (xem PHẦN 4). Chỉ hiện khi có quyền '
    '“Sửa khách hàng” VÀ khách hàng đang Hoạt động.')
s.b('{icon:btn_khoa_row} **Khóa** (khách hàng đang Hoạt động) hoặc {icon:btn_mokhoa_row} **Mở khóa** '
    '(khách hàng đang Khóa) — xem PHẦN 5. Chỉ hiện khi có quyền “Xóa khách hàng”.')
s.b('{icon:btn_quanly_row} **Quản lý** — mở màn Quản lý khách hàng (xem PHẦN 8). Chỉ hiện khi có quyền '
    '“Xem khách hàng”.')
s.b('{icon:btn_lichsu_row} **Lịch sử** — mở cửa sổ Lịch sử khách hàng (xem PHẦN 6). Chỉ hiện khi có quyền '
    '“Xem lịch sử khách hàng”.')
s.p('Cách các nút được sắp xếp trên dòng:')
s.b('Dòng có đủ 4 nút (khách hàng đang Hoạt động, Người dùng đủ quyền): Sửa và Khóa hiện thẳng trên dòng; '
    'Quản lý và Lịch sử được gom vào nút ba chấm {icon:btn_bacham_kh} “Hành động khác”. Bấm nút ba chấm, '
    'danh sách mở ra gồm {icon:item_quanly} và {icon:item_lichsu}.')
s.b('Dòng chỉ có từ 3 nút trở xuống (ví dụ khách hàng đang Khóa nên không có nút Sửa, hoặc Người dùng thiếu '
    'quyền): mọi nút hiện thẳng trên dòng dạng biểu tượng, KHÔNG có nút ba chấm.')
s.img('03_rowmenu.png', 'Dòng khách hàng đang Hoạt động — nút ba chấm mở ra Quản lý và Lịch sử')
s.img('03b_row_locked.png', 'Dòng khách hàng đang Khóa — ba nút Mở khóa, Quản lý, Lịch sử hiện thẳng trên dòng')
s.p('Lưu ý quan trọng: nút nào Người dùng không có quyền, hoặc chưa đủ điều kiện (ví dụ khách hàng ĐÃ KHÓA '
    'thì không còn nút Sửa), sẽ BIẾN MẤT hẳn chứ không bị làm mờ. Muốn sửa khách hàng đã khóa, phải Mở khóa trước.')
s.p('Màn hình không có chức năng Xóa khách hàng — khách hàng không dùng nữa thì Khóa lại.')

# ═════════════ PHẦN 2 — ảnh nút ở bước 2 và Cài đặt bộ lọc ═════════════
p2 = find_p('Bước 2: Nhập từ khóa vào ô tìm kiếm nhanh')
fill(p2, 'Bước 2: Nhập từ khóa vào ô tìm kiếm nhanh rồi nhấn Enter hoặc bấm {icon:btn_timkiem}; hoặc bấm '
         '{icon:btn_timkiemnangcao} để mở khối bộ lọc chi tiết và chọn tiêu chí.')
pc = find_p('Cài đặt bộ lọc:')
fill(pc, 'Cài đặt bộ lọc: trong khối Tìm kiếm nâng cao bấm {icon:btn_caidatboloc} để bỏ tích (ẩn) bớt ô lọc '
         'hoặc kéo đổi thứ tự; cấu hình lưu riêng cho từng tài khoản. Ô lọc bị ẩn mà đang có giá trị thì giá '
         'trị đó vẫn còn tác dụng cho tới khi bấm Làm mới.')

# ═════════════ PHẦN 4 → PHẦN 8: viết lại ═════════════
start = find_p('PHẦN 4: CHỈNH SỬA', 'Heading 1').getprevious()    # tính cả đoạn ngắt trang
end = find_p('PHẦN 9:', 'Heading 1').getprevious()
assert 'w:type="page"' in start.xml and 'w:type="page"' in end.xml
s = Seq(cut(start, end))

# ── PHẦN 4 ──
s.h1('PHẦN 4: CHỈNH SỬA')
s.h2('1. Các bước chỉnh sửa')
s.step('Bước 1: Truy cập vào màn hình Danh mục khách hàng (xem PHẦN 1).')
s.step('Bước 2: Bấm {icon:btn_sua_kh} ở cột Hành động của khách hàng cần sửa, hoặc bấm {icon:btn_sua_footer_kh} '
       'ở chân trang Chi tiết khách hàng. Nút chỉ hiện khi có quyền “Sửa khách hàng” và khách hàng đang Hoạt động.')
s.step('Bước 3: Hệ thống mở trang “Chỉnh sửa khách hàng”, mọi ô đã điền sẵn thông tin hiện tại của khách hàng.')
s.img('12_edit.png', 'Trang Chỉnh sửa khách hàng')
s.step('Bước 4: Sửa các ô cần thay đổi. Quy tắc nhập từng ô giống Thêm mới (PHẦN 3, mục 2). Các phần CHỈ có ở '
       'trang Sửa được hướng dẫn riêng ở mục 3, 4 và 5 dưới đây.')
s.step('Bước 5: Bấm {icon:btn_luu_kh} ở chân trang. Lưu thành công, hệ thống báo “Cập nhật khách hàng thành '
       'công” và quay về màn danh sách. Còn ô thiếu/sai, hệ thống báo “Bạn chưa nhập đầy đủ thông tin”, ô sai '
       'có viền đỏ kèm lỗi ngay dưới ô, trang không đóng và dữ liệu đã nhập vẫn còn.')
s.step('Hoặc bấm {icon:btn_quaylai_kh} để giữ nguyên thông tin cũ. Nếu đã sửa dở, hệ thống hỏi “Thông tin chưa '
       'lưu” (Thoát / Ở lại) — chọn Ở lại thì dữ liệu đang sửa vẫn còn nguyên.')

s.h2('2. Những điểm khác với trang Thêm mới')
s.table([
    ['Nội dung', 'Trên trang Chỉnh sửa'],
    ['Dữ liệu khi mở trang', 'Mọi ô đã điền sẵn thông tin hiện tại; các khối theo Loại hình tổ chức (cá nhân / '
     'tổ chức) hiện sẵn, không phải chọn lại.'],
    ['Mã khách hàng', 'Không có ô Mã trên trang. Mã giữ nguyên sau khi lưu, kể cả khi đổi Tỉnh/Thành phố hoặc '
     'Phường/Xã (mã chỉ tự sinh một lần khi Thêm mới).'],
    ['Nhân viên phụ trách đại lý, Cấp đại lý', 'CHỈ có ở trang Chỉnh sửa, nằm ở hàng cuối khối Thông tin khách '
     'hàng — xem mục 3.'],
    ['Khối Địa chỉ giao hàng', 'CHỈ có ở trang Chỉnh sửa, nằm sau khối địa chỉ và tài khoản ngân hàng — xem mục 4.'],
    ['Người liên hệ đã có (khách hàng tổ chức)', 'Người liên hệ còn giữ trên trang được cập nhật tại chỗ. Người '
     'liên hệ bị bỏ khỏi trang KHÔNG phải lúc nào cũng bị xóa — xem mục 5.'],
    ['Người đại diện, Tài khoản ngân hàng, Nhóm khách hàng, Hãng xe', 'Khi lưu, hệ thống ghi lại đúng danh sách '
     'đang có trên trang: dòng đã bấm xóa thì mất hẳn, dòng mới thêm được ghi vào.'],
    ['Trùng mã số thuế / email / số CCCD', 'Giữ nguyên giá trị của chính khách hàng đang sửa thì không bị báo '
     'trùng; chỉ đổi sang giá trị của khách hàng khác mới bị chặn.'],
])

s.h2('3. Nhân viên phụ trách đại lý và Cấp đại lý')
s.p('Hai ô này dùng cho khách hàng là đại lý: ghi nhận nhân viên phụ trách chăm sóc đại lý và cấp của đại lý. '
    'Cấp đại lý còn là một cột bật được trên màn danh sách, một tiêu chí lọc và một trường khi xuất file.')
s.img('13b_edit_agent.png', 'Hai ô Nhân viên phụ trách đại lý và Cấp đại lý trên trang Chỉnh sửa')
s.table([
    ['Trường', 'Kiểu nhập', 'Bắt buộc', 'Giá trị ban đầu', 'Ghi chú'],
    ['Nhân viên phụ trách đại lý', 'Ô chọn giá trị, gõ chữ để tìm', 'Không', 'Nhân viên đang lưu (trống nếu chưa có)',
     'Danh sách là nhân viên đang làm việc, hiển thị dạng “Tên nhân viên - Mã phòng - Mã nhân viên”, xếp theo tên '
     '(hiển thị tối đa 100 người đầu tiên). Nhân viên đang được chọn luôn hiện đúng tên. Bấm × để bỏ chọn.'],
    ['Cấp đại lý', 'Ô chọn giá trị', 'Không', 'Cấp đang lưu (trống nếu chưa có)',
     'Chọn một trong: Cấp 1 / Cấp 2 / Cấp 3. Bấm × để bỏ chọn.'],
], five=True)
s.b('Đổi hai ô này được ghi vào Lịch sử (“Nhân viên phụ trách đại lý”, “Cấp đại lý”, giá trị cũ → giá trị mới).')

s.h2('4. Khối Địa chỉ giao hàng')
s.p('Khối này khai các nơi nhận hàng của khách hàng (kho, chi nhánh, công trình…). Một khách hàng có thể có '
    'nhiều địa chỉ giao hàng. Địa chỉ giao hàng KHÁC với địa chỉ của khách hàng ở khối phía trên (địa chỉ dùng '
    'để sinh mã khách hàng); sửa địa chỉ giao hàng không ảnh hưởng tới mã.')
s.img('13a_edit_delivery.png', 'Khối Địa chỉ giao hàng trên trang Chỉnh sửa (khách hàng có 2 địa chỉ)')
s.h3('4.1. Cách hiển thị')
s.b('Mỗi địa chỉ là một khung đánh số “Địa chỉ 1”, “Địa chỉ 2”… theo thứ tự từ trên xuống; góc phải mỗi khung '
    'có nút {icon:btn_xoa_diachigh} để bỏ địa chỉ đó.')
s.b('Khách hàng chưa có địa chỉ giao hàng nào thì khối chỉ có nút {icon:btn_themdiachigh}.')
s.b('Mỗi khung có các ô: Quốc gia · Tỉnh/Thành phố · Quận/Huyện (chỉ với nước ngoài) · Phường/Xã/Thị trấn · '
    'Đường/Thôn · Số nhà.')
s.h3('4.2. Thêm một địa chỉ giao hàng')
s.step('Bước 1: Cuộn xuống khối Địa chỉ giao hàng, bấm {icon:btn_themdiachigh}. Một khung mới hiện ở cuối '
       'khối, mọi ô để trống.')
s.step('Bước 2: Chọn Quốc gia. Nếu bỏ trống, hệ thống hiểu là Việt Nam.')
s.step('Bước 3: Chọn Tỉnh/Thành phố.')
s.step('Bước 4: Việt Nam — chọn luôn Phường/Xã/Thị trấn (ô này chỉ mở sau khi đã chọn Tỉnh/Thành phố). Nước '
       'ngoài — chọn Quận/Huyện trước, sau đó mới chọn được Phường/Xã.')
s.step('Bước 5: Chọn Đường/Thôn (chỉ mở sau khi đã chọn Phường/Xã) và nhập Số nhà nếu cần.')
s.step('Bước 6: Bấm {icon:btn_luu_kh} ở chân trang. Địa chỉ chỉ được ghi vào hệ thống khi bấm Lưu; bấm '
       '{icon:btn_quaylai_kh} và chọn Thoát thì địa chỉ vừa thêm bị bỏ.')
s.h3('4.3. Các trường của một địa chỉ giao hàng')
s.table([
    ['Trường', 'Kiểu nhập', 'Bắt buộc', 'Giá trị ban đầu', 'Ghi chú'],
    ['Quốc gia', 'Ô chọn giá trị', 'Không', 'Trống (khung mới) / giá trị đang lưu',
     'Để trống = Việt Nam. Đổi Quốc gia thì Tỉnh/Thành phố, Quận/Huyện, Phường/Xã, Đường/Thôn của khung đó bị '
     'xóa trắng.'],
    ['Tỉnh/Thành phố', 'Ô chọn giá trị', 'Không có dấu *, nhưng THIẾU thì khung không được lưu',
     'Trống / giá trị đang lưu',
     'Danh sách tỉnh/thành phố lấy theo Quốc gia ở địa chỉ chính của khách hàng (khối phía trên). Đổi tỉnh thì '
     'các ô bên dưới của khung đó bị xóa trắng và nạp lại danh sách.'],
    ['Quận/Huyện', 'Ô chọn giá trị', 'Không', 'Trống / giá trị đang lưu',
     'Chỉ hiện khi Quốc gia của khung KHÁC Việt Nam (Việt Nam đã bỏ cấp huyện). Mở sau khi chọn Tỉnh/Thành phố.'],
    ['Phường/Xã/Thị trấn', 'Ô chọn giá trị', 'Không có dấu *, nhưng THIẾU thì khung không được lưu',
     'Trống / giá trị đang lưu',
     'Việt Nam: liệt kê theo Tỉnh/Thành phố; nước ngoài: theo Quận/Huyện. Đổi phường/xã thì ô Đường/Thôn bị '
     'xóa trắng.'],
    ['Đường/Thôn', 'Ô chọn giá trị', 'Không', 'Trống / giá trị đang lưu', 'Liệt kê theo Phường/Xã đã chọn.'],
    ['Số nhà', 'Ô nhập giá trị', 'Không', 'Trống / giá trị đang lưu', 'Tối đa 255 ký tự.'],
], five=True)
s.h3('4.4. Sửa hoặc bỏ một địa chỉ giao hàng')
s.b('Sửa: chọn lại giá trị trong khung cần sửa rồi bấm {icon:btn_luu_kh}.')
s.b('Bỏ: bấm {icon:btn_xoa_diachigh} ở góc phải khung. Khung biến mất NGAY, không có hộp xác nhận; các khung '
    'phía sau được đánh số lại. Nếu bấm nhầm: bấm {icon:btn_quaylai_kh}, chọn Thoát ở hộp “Thông tin chưa lưu” '
    'để bỏ toàn bộ thay đổi rồi mở lại trang.')
s.b('Bỏ hết mọi khung rồi bấm Lưu thì khách hàng không còn địa chỉ giao hàng nào.')
s.h3('4.5. Quy tắc khi lưu — cần đặc biệt lưu ý')
s.b('Khối Địa chỉ giao hàng không có ô bắt buộc và KHÔNG báo lỗi đỏ. Tuy nhiên, khung nào thiếu Tỉnh/Thành phố '
    'HOẶC thiếu Phường/Xã/Thị trấn sẽ bị hệ thống BỎ QUA khi lưu mà không báo gì — mở lại trang sẽ không thấy '
    'khung đó. Vì vậy luôn chọn đủ hai ô này cho mọi khung.')
s.b('Mỗi lần Lưu, hệ thống ghi lại toàn bộ danh sách địa chỉ đang có trên trang (thay cho danh sách cũ), giữ '
    'đúng thứ tự từ trên xuống.')
s.b('Thêm, sửa, bỏ địa chỉ giao hàng đều được ghi Lịch sử ở mục “Địa điểm giao hàng”, mỗi địa chỉ hiển thị dạng '
    '“Số nhà, Đường/Thôn, Phường/Xã, Tỉnh/Thành phố”.')
s.h3('4.6. Địa chỉ giao hàng được dùng ở đâu')
s.b('Khi báo giá của khách hàng được duyệt, hệ thống đồng bộ báo giá sang ERP và lấy “Địa chỉ 1” (khung trên '
    'cùng) làm địa chỉ giao hàng. Vì vậy nên đặt địa chỉ giao hàng chính ở khung đầu tiên.')
s.b('Khách hàng CHƯA có địa chỉ giao hàng nào thì báo giá đã duyệt KHÔNG đồng bộ được sang ERP (lỗi đồng bộ '
    'ghi “chưa có địa chỉ giao hàng”). Cần vào Chỉnh sửa khách hàng khai ít nhất một địa chỉ rồi đồng bộ lại.')
s.b('Trang Chi tiết khách hàng hiển thị khối này ở chế độ chỉ đọc; màn Quản lý khách hàng sửa được khối này ở '
    'thẻ Thông tin chung.')

s.h2('5. Người liên hệ đã có khi sửa (khách hàng tổ chức)')
s.b('Thêm người liên hệ bằng nút {icon:btn_themnguoilh}; bỏ bằng nút {icon:btn_xoa_diachigh} ở góc khung người '
    'liên hệ (chỉ hiện khi có từ 2 người liên hệ trở lên, bấm là bỏ ngay không hỏi xác nhận).')
s.b('Người liên hệ còn giữ trên trang: được cập nhật thông tin tại chỗ (không tạo người mới).')
s.b('Người liên hệ bị bỏ khỏi trang chỉ bị xóa hẳn khi: Người dùng là Giám đốc; hoặc người liên hệ đó đã có '
    'báo giá do chính Người dùng lập (Trưởng phòng: do nhân viên các phòng mình quản lý lập).')
s.b('Các trường hợp còn lại, người liên hệ vẫn được GIỮ LẠI — mở lại trang Chỉnh sửa sẽ thấy người đó xuất hiện '
    'trở lại. Đây là quy tắc bảo vệ người liên hệ đang gắn với báo giá của người khác, không phải lỗi.')

s.h2('6. Các lưu ý khác')
s.b('Khách hàng đang Khóa không có nút Sửa. Nếu khách hàng vừa bị người khác khóa trong lúc đang sửa, bấm Lưu '
    'sẽ nhận thông báo “Khách hàng đang bị khoá, vui lòng mở khoá trước khi cập nhật.”')
s.b('Không có quyền “Sửa khách hàng” mà mở thẳng trang Sửa thì hệ thống báo “Bạn không có quyền sửa khách hàng” '
    'và chuyển về trang Chi tiết.')
s.b('Mọi thay đổi đều được ghi vào Lịch sử (xem PHẦN 6).')

# ── PHẦN 5 ──
s.h1('PHẦN 5: KHÓA VÀ MỞ KHÓA')
s.h2('1. Ý nghĩa')
s.p('Khi một khách hàng không còn giao dịch nhưng không muốn xóa, hãy dùng thao tác Khóa. Sau khi khóa:')
s.b('Khách hàng VẪN nằm trong danh sách, cột Trạng thái hiện chữ Khóa.')
s.b('Vẫn xem chi tiết, lịch sử, màn Quản lý và vẫn xuất ra file bình thường; báo giá, hợp đồng, trang thiết bị '
    'của khách hàng vẫn còn nguyên.')
s.b('Khách hàng đã khóa không còn xuất hiện ở ô chọn khách hàng của các màn nghiệp vụ khác (lập báo giá, hợp '
    'đồng…).')
s.b('Nút Sửa của dòng đó biến mất; hệ thống cũng chặn mọi thao tác cập nhật (kể cả trên màn Quản lý) cho tới '
    'khi Mở khóa.')
s.h2('2. Các bước khóa')
s.p('Tại màn hình danh sách khách hàng:')
s.step('Bước 1: Ở cột Hành động của khách hàng đang Hoạt động, bấm {icon:btn_khoa_row} (Khóa). Hoặc tại chân '
       'trang Chi tiết khách hàng, bấm {icon:btn_khoa_footer}.')
s.step('Bước 2: Hệ thống hiện hộp “Xác nhận khóa” với câu “Bạn có chắc muốn khóa khách hàng \'<Mã KH> - <Tên '
       'khách hàng>\'?”. Đọc kỹ mã và tên để chắc không bấm nhầm dòng.')
s.img('14_lock.png', 'Hộp xác nhận khóa khách hàng')
s.step('Bước 3: Bấm {icon:btn_confirm_khoa} để xác nhận. Hệ thống báo “Khóa khách hàng thành công!”, cột Trạng '
       'thái đổi thành Khóa; trên dòng, nút Sửa biến mất và nút Khóa đổi thành Mở khóa.')
s.step('Hoặc bấm {icon:btn_huy_kh} nếu bấm nhầm — không có gì thay đổi.')
s.h2('3. Mở khóa')
s.p('Tại màn hình danh sách khách hàng:')
s.step('Bước 1: Ở dòng khách hàng đang Khóa, bấm {icon:btn_mokhoa_row} (Mở khóa). Hoặc tại chân trang Chi tiết '
       'khách hàng, bấm {icon:btn_mokhoa_footer_kh}.')
s.step('Bước 2: Hệ thống hiện hộp “Xác nhận mở khóa” với câu “Bạn có chắc muốn mở khóa khách hàng \'<Mã KH> - '
       '<Tên khách hàng>\'?”.')
s.img('15_unlock.png', 'Hộp xác nhận mở khóa khách hàng')
s.step('Bước 3: Bấm {icon:btn_confirm_mokhoa_kh} để xác nhận. Hệ thống báo “Mở khóa khách hàng thành công!”, cột '
       'Trạng thái đổi thành Hoạt động, nút Sửa hiện lại và khách hàng chọn lại được ở các màn nghiệp vụ.')
s.step('Hoặc bấm {icon:btn_huy_kh} nếu bấm nhầm.')
s.p('Hai thao tác Khóa / Mở khóa cần quyền “Xóa khách hàng”; không có quyền thì nút không hiển thị. Mỗi lần đổi '
    'trạng thái được ghi lịch sử nhóm “Thay đổi trạng thái”.')

# ── PHẦN 6 ──
s.h1('PHẦN 6: XEM LỊCH SỬ THAY ĐỔI')
s.h2('1. Xem lịch sử từ màn danh sách')
s.step('Bước 1: Truy cập vào màn hình Danh mục khách hàng.')
s.step('Bước 2: Ở cột Hành động của khách hàng cần xem, bấm nút ba chấm {icon:btn_bacham_kh} rồi chọn '
       '{icon:item_lichsu}. Nếu dòng không có nút ba chấm thì bấm thẳng {icon:btn_lichsu_row} trên dòng. Cần quyền '
       '“Xem lịch sử khách hàng”.')
s.step('Bước 3: Hệ thống mở cửa sổ “Lịch sử khách hàng”, dòng dưới tiêu đề ghi tên khách hàng đang xem. Các lần '
       'thay đổi xếp theo thời gian, mới nhất trên cùng.')
s.img('16_history.png', 'Cửa sổ Lịch sử khách hàng')
s.step('Bước 4 (khi cần thu hẹp): Bấm {icon:btn_boloc_lichsu} để mở bộ lọc, chọn Loại hành động, Người thực hiện, '
       'Từ ngày, Đến ngày. Danh sách lọc NGAY khi chọn, không có nút Tìm kiếm. Bấm {icon:btn_lammoi_lichsu} để bỏ '
       'hết điều kiện lọc.')
s.img('16b_history_filter.png', 'Bộ lọc trong cửa sổ Lịch sử khách hàng')
s.step('Bước 5: Bấm {icon:btn_dong_lichsu} (hoặc dấu × góc trên) để đóng cửa sổ.')
s.h2('2. Xem lịch sử từ màn Chi tiết')
s.step('Bước 1: Mở trang Chi tiết khách hàng (xem PHẦN 7).')
s.step('Bước 2: Cuộn xuống cuối trang — khối “Lịch sử” hiển thị ngay trong trang, không cần bấm nút.')
s.h2('3. Cách đọc một mốc lịch sử')
s.b('Thời điểm thay đổi, dạng dd/mm/yyyy hh:mm.')
s.b('Loại thay đổi (tô màu) — Tạo khách hàng, Chỉnh sửa thông tin, Khóa khách hàng, Mở khóa khách hàng…')
s.b('“Người thực hiện: <tên> — <phòng ban>”.')
s.b('Từng trường đã đổi: giá trị cũ → giá trị mới.')
s.b('Trường dạng danh sách (Hãng xe, Nhóm khách hàng, Người đại diện, Người liên hệ, Tài khoản, Địa điểm giao '
    'hàng…): mỗi phần tử một dòng — dòng bắt đầu bằng “+” là phần được THÊM, “-” là phần bị BỎ, “~” là phần '
    'được SỬA (kèm trường nào đổi từ gì sang gì).')
s.b('Khách hàng chưa có thay đổi nào: cửa sổ hiện “Chưa có lịch sử thay đổi.” (và không có nút Bộ lọc). Lọc '
    'không ra kết quả: hiện “Không có lịch sử phù hợp bộ lọc.”')

# ── PHẦN 7 ──
s.h1('PHẦN 7: XEM CHI TIẾT KHÁCH HÀNG')
s.step('Bước 1: Truy cập vào màn hình Danh mục khách hàng.')
s.step('Bước 2: Bấm vào Mã KH ở cột Mã KH (cần quyền “Xem khách hàng”).')
s.step('Bước 3: Hệ thống mở trang “Chi tiết khách hàng: <Mã KH>”.')
s.img('17_detail.png', 'Trang Chi tiết khách hàng ở chế độ chỉ đọc')
s.p('Mọi ô ở chế độ chỉ đọc (kể cả Nhân viên phụ trách đại lý, Cấp đại lý và khối Địa chỉ giao hàng), không có '
    'nút Lưu; khối Lịch sử nằm ở cuối trang. Chân trang có các nút giống cột Hành động của đúng dòng đó:')
s.b('{icon:btn_sua_footer_kh} — mở trang Chỉnh sửa (PHẦN 4). Chỉ hiện khi có quyền “Sửa khách hàng” và khách '
    'hàng đang Hoạt động.')
s.b('{icon:btn_quanly_footer} — mở màn Quản lý khách hàng (PHẦN 8).')
s.b('{icon:btn_khoa_footer} hoặc {icon:btn_mokhoa_footer_kh} — Khóa / Mở khóa (PHẦN 5). Chỉ hiện khi có quyền '
    '“Xóa khách hàng”.')
s.b('{icon:btn_quaylai_kh} — quay về nơi đã mở trang (thường là danh sách khách hàng).')
s.p('Thiếu quyền “Xem khách hàng” mà mở thẳng trang thì hệ thống báo “Bạn không có quyền xem khách hàng” và quay '
    'về danh sách. Khách hàng nằm ngoài phạm vi dữ liệu của Người dùng cũng không xem được (“Bạn không có quyền '
    'xem khách hàng này”).')

# ── PHẦN 8 ──
s.h1('PHẦN 8: MÀN QUẢN LÝ KHÁCH HÀNG')
s.p('Màn Quản lý gom toàn bộ thông tin của MỘT khách hàng vào các thẻ: thông tin chung, người liên hệ, báo giá, '
    'hợp đồng, trang thiết bị và hình ảnh / tài liệu / video. Khác với trang Chi tiết, màn này SỬA ĐƯỢC.')
s.h2('1. Mở màn Quản lý khách hàng')
s.step('Bước 1: Tại màn hình danh sách, ở cột Hành động bấm nút ba chấm {icon:btn_bacham_kh} rồi chọn '
       '{icon:item_quanly} (dòng không có nút ba chấm thì bấm thẳng {icon:btn_quanly_row}). Hoặc tại chân trang '
       'Chi tiết khách hàng, bấm {icon:btn_quanly_footer}.')
s.step('Bước 2: Hệ thống mở trang “Quản lý khách hàng”, thẻ Thông tin chung được chọn sẵn.')
s.img('30_mgr_0.png', 'Màn Quản lý khách hàng — thẻ Thông tin chung')
s.p('Cần quyền “Xem khách hàng”; thiếu quyền thì hệ thống báo “Bạn không có quyền xem khách hàng” và quay về danh '
    'sách. Các thẻ trên màn hình:')
s.table([
    ['Thẻ', 'Nội dung và khi nào hiện'],
    ['Thông tin chung', 'Luôn hiện. Sửa được thông tin khách hàng, gồm cả Nhân viên phụ trách đại lý, Cấp đại lý '
     'và Địa chỉ giao hàng — mục 3.'],
    ['Thông tin liên hệ', 'Chỉ hiện với khách hàng tổ chức (Doanh nghiệp tư nhân, Doanh nghiệp nước ngoài, Tổ chức '
     'phi chính phủ, Cơ quan nhà nước). Sửa được danh sách người liên hệ — mục 4.'],
    ['Báo giá', 'Chỉ xem: báo giá hàng hóa và báo giá dịch vụ của khách hàng; in, xuất Excel — mục 5.'],
    ['Hợp đồng', 'Chỉ xem: hợp đồng hàng hóa và hợp đồng dịch vụ; in, xuất Excel — mục 6.'],
    ['Danh sách trang thiết bị', 'Thiết bị khách hàng đang có; thêm / sửa / xóa thiết bị khai thêm, quản lý '
     'serial — mục 7.'],
    ['Thông tin khác', 'Hình ảnh, tài liệu PDF, video của khách hàng; có nút Lưu riêng — mục 8.'],
])
s.p('Các thẻ Báo giá, Hợp đồng, Danh sách trang thiết bị, Thông tin khác chỉ tải dữ liệu khi bấm vào thẻ lần đầu.')

s.h2('2. Lưu dữ liệu trên màn Quản lý — đọc kỹ trước khi thao tác')
s.b('{icon:btn_luu_kh} ở chân trang: lưu thẻ Thông tin chung và Thông tin liên hệ. Thành công hệ thống báo '
    '“Cập nhật khách hàng thành công” và CHUYỂN VỀ màn danh sách khách hàng. Cần quyền “Sửa khách hàng”.')
s.b('{icon:btn_luu_tab} ở cuối thẻ Thông tin khác: CHỈ lưu hình ảnh, tài liệu, video.')
s.b('Hai nút Lưu độc lập nhau: sửa ở thẻ Thông tin khác rồi bấm Lưu chân trang sẽ KHÔNG lưu ảnh / tài liệu / '
    'video, và ngược lại. Hãy lưu xong từng phần trước khi chuyển.')
s.b('Các thao tác trên thẻ Danh sách trang thiết bị được lưu ngay trong từng cửa sổ (nút Lưu lại), không cần bấm '
    'Lưu chân trang.')
s.b('{icon:btn_quaylai_kh}: về màn danh sách khách hàng. Màn Quản lý KHÔNG hỏi “Thông tin chưa lưu” khi rời '
    'trang — thay đổi chưa bấm Lưu sẽ mất.')
s.b('Khách hàng đang Khóa: màn Quản lý vẫn hiện đủ ô và nút, nhưng mọi thao tác lưu / thêm / sửa / xóa đều bị '
    'hệ thống từ chối (báo “… thất bại” hoặc “Khách hàng đang bị khoá, vui lòng mở khoá trước khi cập nhật.”). '
    'Muốn cập nhật phải Mở khóa trước (PHẦN 5). Xem, In, Xuất Excel vẫn làm được.')
s.b('Thiếu quyền “Sửa khách hàng”: vẫn mở được màn và thấy các ô, nhưng bấm Lưu (và các thao tác sửa / xóa thiết '
    'bị, serial, tài liệu) hệ thống báo không có quyền.')

s.h2('3. Thẻ Thông tin chung')
s.step('Bước 1: Thẻ Thông tin chung được chọn sẵn khi mở màn (hoặc bấm vào tên thẻ).')
s.step('Bước 2: Sửa các ô cần thay đổi. Các khối và quy tắc nhập giống trang Chỉnh sửa: Thông tin khách hàng '
       '(gồm Nhân viên phụ trách đại lý, Cấp đại lý — PHẦN 4 mục 3), Thông tin cá nhân hoặc Thông tin tổ chức, '
       'địa chỉ, tài khoản, và khối Địa chỉ giao hàng với nút {icon:btn_themdiachigh} (PHẦN 4 mục 4).')
s.step('Bước 3: Bấm {icon:btn_luu_kh} ở chân trang để lưu. Hệ thống báo “Cập nhật khách hàng thành công” và '
       'chuyển về màn danh sách. Còn ô thiếu/sai: báo “Bạn chưa nhập đầy đủ thông tin”, lỗi đỏ ngay dưới ô.')

s.h2('4. Thẻ Thông tin liên hệ')
s.step('Bước 1: Bấm thẻ “Thông tin liên hệ” (chỉ có với khách hàng tổ chức).')
s.img('31_mgr_1.png', 'Thẻ Thông tin liên hệ')
s.step('Bước 2: Mỗi người liên hệ là một khung “Người liên hệ N”. Dòng “NVKD: … / Phòng ban: …” (chỉ đọc) cho '
       'biết nhân viên kinh doanh đang phụ trách người liên hệ đó. Sửa các ô Họ tên (*), Chức vụ (*), Sinh nhật, '
       'Email, CCCD/CMT, Số điện thoại (*) và khối Tài khoản cá nhân.')
s.step('Bước 3: Thêm người bằng {icon:btn_themnguoilh}; bỏ người bằng {icon:btn_xoa_diachigh} ở góc khung (chỉ '
       'hiện khi có từ 2 người trở lên, bỏ ngay không hỏi xác nhận).')
s.step('Bước 4: Bấm {icon:btn_luu_kh} ở chân trang. Quy tắc giữ / xóa người liên hệ bị bỏ giống PHẦN 4 mục 5.')

def tab_docs(num, kind, tbl1, tbl2, so_cot, extra_col, quyen):
    s.p('Thẻ gồm 2 bảng: “%s” và “%s”. Mỗi bảng có bộ lọc và nút riêng, thao tác giống nhau.' % (tbl1, tbl2))
    s.h3('%d.1. Các bước tra cứu' % num)
    s.step('Bước 1: Bấm thẻ “%s”.' % kind)
    s.step('Bước 2: Ở bảng cần tra, nhập điều kiện: Từ ngày, Đến ngày, Nhân viên, Phòng ban, Công ty (chỉ có ở '
           'bảng hàng hóa).')
    s.step('Bước 3: Bấm {icon:btn_loc_tab}. Bảng CHỈ lọc khi bấm Lọc — đổi giá trị ô lọc không tự tải lại. Bấm '
           '{icon:btn_lammoi_tab} để xóa hết điều kiện và tải lại.')
    s.table([
        ['Ô lọc', 'Cách dùng'],
        ['Từ ngày / Đến ngày', 'Lọc theo Ngày lập, chọn ngày dạng dd/mm/yyyy.'],
        ['Nhân viên', 'Người lập chứng từ, hiển thị “Tên nhân viên - Mã phòng - Mã nhân viên” (danh sách gồm tối '
         'đa 100 nhân viên đầu tiên theo tên).'],
        ['Phòng ban', 'Phòng ban của người lập.'],
        ['Công ty', 'Công ty của người lập. Chỉ có ở bảng hàng hóa.'],
    ])
    s.h3('%d.2. Các cột của bảng' % num)
    cols = [['Cột', 'Nội dung'], ['STT', 'Số thứ tự.'], ['Ngày lập', 'Ngày tạo chứng từ, dd/mm/yyyy.'],
            [so_cot, 'Mã chứng từ.'],
            ['Tổng thanh toán', 'Tổng tiền sau thuế, dấu phẩy ngăn cách hàng nghìn; bằng 0 thì để trống.']]
    if extra_col:
        cols.append(extra_col)
    cols += [['Người tạo', 'Người lập chứng từ.'], ['Phòng ban', 'Phòng ban của người lập.']]
    s.table(cols)
    s.b('Dòng “Tổng cộng” cuối bảng là tổng tiền của TOÀN BỘ kết quả lọc, không chỉ trang đang xem.')
    s.b('Mỗi trang 10 dòng; thanh chuyển trang chỉ hiện khi có từ 2 trang. Bảng chỉ để xem, không bấm mở chứng từ '
        'được. Không có dữ liệu thì hiện “Không có dữ liệu”.')
    s.h3('%d.3. Chứng từ nào được liệt kê' % num)
    s.b('Chỉ chứng từ của đúng khách hàng đang mở, mới nhất ở trên.')
    s.b('Bảng hàng hóa: mọi trạng thái. Bảng dịch vụ: không hiện chứng từ đang ở trạng thái “Đang tạo” của người '
        'khác (chứng từ “Đang tạo” của chính mình vẫn hiện).')
    for q in quyen:
        s.b(q)
    s.b('Không có quyền nào ở trên: chỉ thấy chứng từ do chính mình lập.')
    s.h3('%d.4. In và Xuất Excel' % num)
    s.step('In: bấm {icon:btn_in_tab} ở bảng cần in. Hệ thống mở tab mới với bản in (đầu trang là tiêu đề của công '
           'ty người đang in, tên bảng, các cột như trên và dòng Tổng cộng) rồi tự bật hộp thoại in. Bản in lấy '
           'TOÀN BỘ kết quả lọc. Nếu trình duyệt chặn: “Trình duyệt chặn cửa sổ in, vui lòng cho phép popup”.')
    s.step('Xuất Excel: bấm {icon:btn_xuatexcel_tab}. Hệ thống tải về file “<Tên bảng>.xls” gồm các cột như bảng và '
           'dòng Tổng cộng. Lỗi thì báo “Xuất Excel thất bại”.')

s.h2('5. Thẻ Báo giá')
s.img('32_mgr_2.png', 'Thẻ Báo giá — bảng báo giá hàng hóa')
tab_docs(5, 'Báo giá', 'Bảng báo giá hàng hóa', 'Bảng báo giá dịch vụ', 'Số báo giá', None, [
    'Báo giá hàng hóa — phạm vi theo quyền (xét lần lượt): “Xem tất cả báo giá” thấy mọi báo giá; “Xem báo giá '
    'công ty” thấy báo giá do người cùng công ty lập; “Xem báo giá phòng” thấy báo giá của các phòng mình quản '
    'lý; “Xem báo giá bộ phận” thấy báo giá của các bộ phận mình quản lý.',
    'Báo giá dịch vụ — “Xem báo giá dịch vụ SC - BH theo tổng công ty” thấy tất cả; “… theo công ty” thấy báo giá '
    'của công ty mình; “… theo phòng ban” thấy báo giá của phòng mình.',
])

s.h2('6. Thẻ Hợp đồng')
s.img('33_mgr_3.png', 'Thẻ Hợp đồng — bảng hợp đồng hàng hóa')
tab_docs(6, 'Hợp đồng', 'Bảng hợp đồng hàng hóa', 'Bảng hợp đồng dịch vụ', 'Số hợp đồng',
         ['Ngày hiệu lực', 'Ngày hợp đồng được duyệt, dd/mm/yyyy.'], [
    'Hợp đồng hàng hóa — phạm vi theo quyền (xét lần lượt): “Xem tất cả hợp đồng”, “Xem hợp đồng công ty”, “Xem '
    'hợp đồng phòng”, “Xem hợp đồng bộ phận” — cùng cách tính như báo giá.',
    'Hợp đồng dịch vụ — “Xem hợp đồng dịch vụ SC - BH theo tổng công ty” / “… theo công ty” / “… theo phòng ban”.',
])

# ── Thẻ trang thiết bị ──
s.h2('7. Thẻ Danh sách trang thiết bị')
s.p('Thẻ liệt kê mọi thiết bị khách hàng đang có, gồm hàng Tân Phát đã bán qua hệ thống và thiết bị khai thêm '
    '(thiết bị cũ, thiết bị của nhà cung cấp khác), kèm số serial của từng thiết bị.')
s.img('34_mgr_4.png', 'Thẻ Danh sách trang thiết bị')
s.h3('7.1. Tìm kiếm, lọc')
s.step('Bước 1: Bấm thẻ “Danh sách trang thiết bị”.')
s.step('Bước 2: Chọn / nhập điều kiện rồi bấm {icon:btn_loc_tab} (hoặc nhấn Enter trong ô Mã/Tên hàng hóa). Đổi '
       'hai ô chọn KHÔNG tự lọc. Bấm {icon:btn_lammoi_tab} để xóa hết điều kiện và tải lại.')
s.table([
    ['Ô lọc', 'Cách dùng'],
    ['Nhà cung cấp thiết bị', '“Tân Phát” (ra cả hàng Tân Phát đã bán lẫn thiết bị cũ) hoặc “NCC khác”. Để trống = '
     'tất cả.'],
    ['Nhóm hàng hóa', 'Chọn một nhóm hàng. Để trống = tất cả nhóm.'],
    ['Mã/Tên hàng hóa', 'Gõ một phần mã, tên hoặc model thiết bị (với NCC khác tìm cả tên hàng không bán).'],
])
s.h3('7.2. Đọc bảng thiết bị')
s.b('Các cột: STT · Tên thiết bị · Thương hiệu · Model · Mã thiết bị · Số lượng · Serial thiết bị làm dịch vụ · '
    'NCC thiết bị · Hành động. Bảng tải hết một lần, không phân trang.')
s.b('Dòng “Tổng cộng: X thiết bị” ngay dưới tiêu đề: X là SỐ DÒNG thiết bị, không phải tổng số lượng.')
s.b('Thiết bị được chia nhóm theo Nhóm hàng hóa: dòng nhóm in đậm đánh số 1, 2…; thiết bị trong nhóm đánh số 1.1, '
    '1.2…; thiết bị chưa gắn nhóm nằm ở nhóm “Chưa phân nhóm”.')
s.b('Cột Serial là liên kết {icon:link_xemserial} — số trong ngoặc là số serial; bấm để mở Danh sách serial '
    '(mục 7.8).')
s.table([
    ['Loại dòng', 'Nhận biết và thao tác được'],
    ['Hàng Tân Phát đã bán qua hệ thống', 'Cột NCC thiết bị ghi “Tân Phát”, cột Hành động TRỐNG. Số lượng cộng từ '
     'các phiếu xuất kho / mượn / bán. Không sửa, xóa, tăng số lượng được; chỉ thêm serial cho phiếu xuất.'],
    ['Thiết bị cũ (khai thêm)', 'Cột NCC thiết bị cũng ghi “Tân Phát” nhưng cột Hành động có 3 nút '
     'Tăng số lượng, Sửa, Xóa.'],
    ['Thiết bị NCC khác (khai thêm)', 'Cột NCC thiết bị ghi “NCC khác”, có 3 nút như trên. Hàng công ty không bán '
     'thì cột Tên thiết bị hiện tên hàng không bán.'],
])
s.p('Các nút trên dòng thiết bị khai thêm: {icon:btn_tangsoluong} Tăng số lượng · {icon:btn_sua_tb} Sửa · '
    '{icon:btn_xoa_tb} Xóa.')

s.h3('7.3. Thêm thiết bị cũ')
s.step('Bước 1: Bấm {icon:btn_themtbcu}. Hệ thống mở cửa sổ “Thêm mới thiết bị cũ”.')
s.img('40_tb_themcu.png', 'Cửa sổ Thêm mới thiết bị cũ')
s.step('Bước 2: Nhập các trường:')
s.table([
    ['Trường', 'Kiểu nhập', 'Bắt buộc', 'Giá trị ban đầu', 'Ghi chú'],
    ['Trang thiết bị', 'Ô chọn, gõ mã/tên để tìm', 'Có', 'Trống',
     'Kết quả hiện dạng “Mã - Tên”. Bỏ trống báo “Vui lòng chọn trang thiết bị”.'],
    ['Số lượng', 'Ô số', 'Có', 'Trống', 'Chỉ nhận số nguyên dương. Bỏ trống báo “Vui lòng nhập số lượng”; ≤ 0 báo '
     '“Số lượng phải lớn hơn 0”.'],
    ['Ngày xuất hàng', 'Chọn ngày', 'Không', 'Trống', 'dd/mm/yyyy.'],
    ['Địa điểm sử dụng', 'Ô nhập nhiều dòng', 'Có', 'Trống', 'Bỏ trống báo “Vui lòng nhập địa điểm sử dụng”.'],
], five=True)
s.step('Bước 3: Bấm {icon:btn_luulai}. Thành công: báo “Thao tác thành công”, đóng cửa sổ và tải lại bảng. Bấm '
       '{icon:btn_dong_popup} để hủy.')
s.b('Khách hàng đã có thiết bị cũ CÙNG sản phẩm thì hệ thống báo “Đã tồn tại thiết bị của khách hàng!” — không '
    'tự cộng dồn; muốn thêm số lượng hãy dùng Tăng số lượng (mục 7.6).')
s.b('Serial không nhập ở cửa sổ này; thêm sau ở Danh sách serial (mục 7.9).')
s.b('Cần quyền “Xem khách hàng”.')

s.h3('7.4. Thêm thiết bị NCC khác')
s.step('Bước 1: Bấm {icon:btn_themtbncc}. Hệ thống mở cửa sổ “Thêm mới thiết bị NCC khác”.')
s.img('41a_tb_themncc.png', 'Cửa sổ Thêm mới thiết bị NCC khác')
s.step('Bước 2: Nếu thiết bị là hàng công ty KHÔNG bán, tích ô “Hàng công ty không bán” — form đổi như hình dưới.')
s.img('41_tb_themncc.png', 'Khi tích Hàng công ty không bán: nhập tên hàng và chọn hàng công ty tương đương')
s.step('Bước 3: Nhập các trường:')
s.table([
    ['Trường', 'Kiểu nhập', 'Bắt buộc', 'Giá trị ban đầu', 'Ghi chú'],
    ['Hàng công ty không bán', 'Ô tích', 'Không', 'Không tích', 'Quyết định cách nhập ô Trang thiết bị.'],
    ['Trang thiết bị', 'Không tích: ô chọn tìm theo mã/tên. Có tích: ô nhập tên hàng không bán', 'Có', 'Trống',
     'Bỏ trống báo “Vui lòng chọn trang thiết bị” / “Vui lòng nhập tên hàng không bán”.'],
    ['Hàng công ty tương đương', 'Bấm vào ô để mở cửa sổ “Chọn hàng hóa áp dụng”', 'Có khi đã tích', 'Trống',
     'Chỉ hiện khi đã tích. Trong cửa sổ, tìm theo tên / mã / model rồi bấm vào một dòng để chọn. Bỏ trống báo '
     '“Vui lòng chọn hàng công ty tương đương”.'],
    ['Số lượng', 'Ô số', 'Có', 'Trống', 'Như thiết bị cũ.'],
    ['Nhà cung cấp', 'Ô chọn, gõ tên/mã để tìm', 'Có dấu (*)', 'Trống', 'Chọn trong các khách hàng được tích “Là '
     'nhà cung cấp”.'],
    ['Tình trạng thiết bị', 'Ô chọn', 'Không', 'Trống', 'Còn tốt / Cần sửa chữa / Cần đại tu.'],
    ['Thời gian sử dụng', 'Ô số', 'Không', 'Trống', ''],
    ['Địa điểm sử dụng', 'Ô nhập nhiều dòng', 'Không', 'Trống', ''],
], five=True)
s.step('Bước 4: Bấm {icon:btn_luulai}. Thành công báo “Thao tác thành công”; trùng thiết bị báo “Đã tồn tại '
       'thiết bị của khách hàng!”. Bấm {icon:btn_dong_popup} để hủy. Cần quyền “Xem khách hàng”.')

s.h3('7.5. Sửa thiết bị')
s.step('Bước 1: Ở dòng thiết bị cũ / NCC khác, bấm {icon:btn_sua_tb}.')
s.step('Bước 2: Hệ thống mở cửa sổ “Cập nhật thiết bị cũ” (hoặc “Cập nhật thiết bị NCC khác”), các ô điền sẵn '
       'dữ liệu của dòng. Các trường và quy tắc giống lúc thêm (mục 7.3 / 7.4).')
s.img('43_tb_sua.png', 'Cửa sổ Cập nhật thiết bị cũ')
s.step('Bước 3: Sửa rồi bấm {icon:btn_luulai}. Thành công báo “Thao tác thành công”. Cần quyền “Sửa khách hàng”.')

s.h3('7.6. Tăng số lượng')
s.step('Bước 1: Ở dòng thiết bị cũ / NCC khác, bấm {icon:btn_tangsoluong}.')
s.step('Bước 2: Cửa sổ “Tăng số lượng” hiện tên thiết bị và Số lượng hiện tại. Nhập “Số lượng thêm (*)” — mặc '
       'định 1.')
s.img('42_tb_tangsl.png', 'Cửa sổ Tăng số lượng')
s.step('Bước 3: Bấm {icon:btn_luulai}. Số nhập vào được CỘNG DỒN vào số lượng hiện có (đang 2, thêm 3 thì thành '
       '5); hệ thống báo “Tăng số lượng thành công”. Số ≤ 0 hoặc để trống báo “Số lượng thêm phải lớn hơn 0”.')

s.h3('7.7. Xóa thiết bị')
s.step('Bước 1: Ở dòng thiết bị cũ / NCC khác, bấm {icon:btn_xoa_tb}.')
s.step('Bước 2: Trình duyệt hỏi “Bạn chắc chắn muốn xóa thiết bị này?” — bấm OK để xóa, Hủy (Cancel) để thôi.')
s.b('Xóa thành công báo “Xóa thành công”; TOÀN BỘ serial của thiết bị đó cũng bị xóa theo.')
s.b('Thiết bị có serial đã được dùng trong nghiệp vụ (bảo hành, sửa chữa, báo giá / hợp đồng dịch vụ, giao '
    'việc…) thì không xóa được: “Hàng hóa này đã tồn tại trong luồng, không thể tiếp tục thao tác”.')
s.b('Cần quyền “Sửa khách hàng”.')

s.h3('7.8. Danh sách serial')
s.step('Bước 1: Ở cột Serial của thiết bị, bấm {icon:link_xemserial}.')
s.step('Bước 2: Cửa sổ “Danh sách serial” hiện các cột: STT · Số serial · Trạng thái (Đang sử dụng / Ngừng sử '
       'dụng) · Hình ảnh · Ngày nghiệm thu · Thời gian bảo hành · Thời gian hết bảo hành · Phiếu YCXH · Hành động. '
       'Chưa có serial thì hiện “Chưa có serial”.')
s.img('45_tb_serial.png', 'Cửa sổ Danh sách serial')
s.b('Thiết bị khai thêm: có nút {icon:btn_themserial} và trên mỗi serial có {icon:btn_sua_serial} Sửa · '
    '{icon:btn_thaydoi_serial} Thay đổi · {icon:btn_xoa_serial} Xóa.')
s.b('Hàng Tân Phát đã bán: chỉ có nút “Thêm serial cho phiếu xuất” (serial gắn vào phiếu xuất gần nhất), không '
    'sửa / xóa serial được.')
s.b('Ba cột Ngày nghiệm thu, Thời gian bảo hành, Thời gian hết bảo hành chỉ có dữ liệu khi serial gắn với phiếu '
    'xuất có hợp đồng; serial thiết bị khai thêm thường để trống.')

s.h3('7.9. Thêm serial')
s.step('Bước 1: Trong cửa sổ Danh sách serial, bấm {icon:btn_themserial} (hoặc “Thêm serial cho phiếu xuất”).')
s.step('Bước 2: Cửa sổ “Thêm serial” có sẵn 1 ô “Nhập serial”. Bấm {icon:btn_themdong_serial} để thêm ô cho '
       'serial tiếp theo; bấm × đỏ cạnh ô để bỏ ô đó.')
s.img('46_serial_them.png', 'Cửa sổ Thêm serial')
s.step('Bước 3: Bấm {icon:btn_luulai}. Thành công báo “Thêm serial thành công”, danh sách serial tự cập nhật. '
       'Lỗi hiện ngay dưới ô serial bị sai:')
s.table([
    ['Thông báo', 'Nguyên nhân'],
    ['Vui lòng nhập serial', 'Còn ô serial để trống.'],
    ['Serial bị trùng trong danh sách', 'Hai ô vừa nhập cùng một số serial.'],
    ['Serial đã tồn tại / Hệ thống đã tồn tại serial này', 'Serial đã có trên hệ thống (kể cả của khách hàng khác, '
     'kể cả serial đã ngừng sử dụng).'],
    ['Số lượng serial không được vượt quá số lượng của thiết bị', 'Thiết bị khai thêm có số serial bằng số lượng '
     'rồi — tăng số lượng trước (mục 7.6).'],
    ['Thiết bị này không có phiếu xuất để gắn serial', 'Dòng hàng Tân Phát không tìm thấy phiếu xuất.'],
])

s.h3('7.10. Sửa serial và Thay đổi serial')
s.step('Sửa (gõ nhầm số serial): bấm {icon:btn_sua_serial} ở serial cần sửa → cửa sổ “Chỉnh sửa serial”, ô '
       '“Nhập serial (*)” điền sẵn số hiện tại → sửa rồi bấm {icon:btn_luulai}. Serial được sửa trực tiếp.')
s.img('47_serial_sua.png', 'Cửa sổ Chỉnh sửa serial')
s.step('Thay đổi (thay thiết bị khác cho khách): bấm {icon:btn_thaydoi_serial} → cửa sổ “Thay đổi serial” hiện '
       '“Số serial hiện tại” (khóa) và ô “Nhập serial mới (*)” → nhập rồi bấm {icon:btn_luulai}. Hệ thống tạo '
       'serial mới và chuyển serial cũ sang “Ngừng sử dụng” — danh sách còn cả hai.')
s.img('48_serial_thaydoi.png', 'Cửa sổ Thay đổi serial')
s.b('Để trống báo “Vui lòng nhập serial”; trùng serial khác báo “Serial đã tồn tại” ngay dưới ô. Thành công báo '
    '“Thao tác thành công”.')

s.h3('7.11. Xóa serial')
s.step('Bấm {icon:btn_xoa_serial} ở serial cần xóa → trình duyệt hỏi “Bạn chắc chắn muốn xóa serial này?” → '
       'bấm OK. Thành công báo “Xóa thành công”.')
s.b('Serial đã dùng trong nghiệp vụ (bảo hành, sửa chữa…) không xóa được: “Serial này đã tồn tại trong luồng, '
    'không thể xóa”.')

s.h3('7.12. In và Xuất Excel danh sách thiết bị')
s.step('In: bấm {icon:btn_in_tab}. Hệ thống mở tab mới “Danh sách trang thiết bị” (đầu trang là tiêu đề công '
       'ty của người đang in) với các cột STT · Nhóm · Tên thiết bị · Thương hiệu · Model · Mã thiết bị · Số '
       'lượng · Số serial · NCC thiết bị, rồi tự bật hộp thoại in. In đúng những dòng đang hiện trên bảng.')
s.step('Xuất Excel: bấm {icon:btn_xuatexcel_tab}. Tải về file “danh-sach-trang-thiet-bi.xls” cùng các cột như bản '
       'in; riêng cột NCC thiết bị của dòng NCC khác ghi TÊN nhà cung cấp.')

s.h3('7.13. Quyền cần có')
s.table([
    ['Thao tác', 'Quyền'],
    ['Xem, In, Xuất Excel', 'Xem khách hàng'],
    ['Thêm thiết bị cũ, Thêm thiết bị NCC khác, Tăng số lượng', 'Xem khách hàng'],
    ['Sửa, Xóa thiết bị; Thêm / Sửa / Thay đổi / Xóa serial', 'Sửa khách hàng'],
])
s.p('Các nút luôn hiển thị; thiếu quyền thì khi bấm hệ thống báo không có quyền. Khách hàng đang Khóa thì các '
    'thao tác thêm / sửa / xóa báo “Khách hàng đang bị khoá, vui lòng mở khoá trước khi cập nhật.”')

# ── Thẻ Thông tin khác ──
s.h2('8. Thẻ Thông tin khác')
s.step('Bước 1: Bấm thẻ “Thông tin khác”. Thẻ có 3 khối: Hình ảnh, Tài liệu liên quan, Video.')
s.img('35_mgr_5.png', 'Thẻ Thông tin khác')
s.h3('8.1. Hình ảnh')
s.b('Bấm {icon:btn_taianh}, chọn một hoặc nhiều file ảnh trên máy. Ảnh được tải lên ngay (nút đổi thành “Đang '
    'tải...”), hiện dạng ô vuông nhỏ. Lỗi báo “Upload ảnh thất bại”.')
s.b('Có từ 2 ảnh trở lên: kéo thả để sắp xếp thứ tự ảnh.')
s.b('Bỏ ảnh: bấm × ở góc ảnh (bỏ ngay, không hỏi xác nhận).')
s.b('Ảnh mới, thứ tự ảnh, ảnh bị bỏ CHỈ được ghi khi bấm {icon:btn_luu_tab} của thẻ.')
s.h3('8.2. Tài liệu liên quan')
s.b('Bấm {icon:btn_themtailieu_kh}, chọn một hoặc nhiều file PDF. File mới hiện kèm nhãn “Chưa lưu” và nút × để '
    'bỏ khỏi danh sách; chỉ được tải lên khi bấm {icon:btn_luu_tab}.')
s.b('Tài liệu đã lưu: bấm vào tên để mở xem ở tab mới. Nút xóa (thùng rác) cạnh tài liệu đã lưu XÓA NGAY trên '
    'hệ thống, không hỏi xác nhận — báo “Xóa tài liệu thành công”.')
s.h3('8.3. Video')
s.b('Bấm {icon:btn_themvideo} để thêm một dòng, nhập Tên video và URL (đường liên kết video). Bỏ dòng bằng nút '
    'thùng rác cuối dòng.')
s.h3('8.4. Lưu thẻ Thông tin khác')
s.step('Bấm {icon:btn_luu_tab} ở cuối thẻ (nút đổi thành “Đang lưu...”). Thành công báo “Cập nhật Thông tin khác '
       'thành công” và tải lại thẻ; lỗi báo “Cập nhật Thông tin khác thất bại”. Cần quyền “Sửa khách hàng”; khách '
       'hàng đang Khóa thì không lưu được. Lần lưu này được ghi vào Lịch sử.')
s.b('Nhắc lại: nút Lưu ở chân trang KHÔNG lưu thẻ này (xem mục 2).')

# ═════════════ PHẦN 9 / 10: ảnh nút ═════════════
p9 = find_p('Bước 2: Bấm một trong ba nút')
fill(p9, 'Bước 2: Bấm một trong ba nút {icon:btn_xuatcsv} / {icon:btn_xuatexcel_kh} / {icon:btn_xuatpdf}. Hệ '
         'thống mở cửa sổ “Chọn trường xuất Excel” (hoặc “Chọn trường xuất CSV” / “Chọn trường xuất PDF”).')
p10 = find_p('Bước 5: Sửa trực tiếp các dòng lỗi')
fill(p10, 'Bước 5: Sửa trực tiếp các dòng lỗi trên bảng rồi bấm {icon:btn_validate} lại. Nếu không muốn nhập các '
          'dòng lỗi thì bấm {icon:btn_bodongloi} để loại chúng khỏi bảng. Muốn sửa lại cả những dòng đã khoá thì '
          'bấm {icon:btn_xoavalidate}.')

# ═════════════ bảng cập nhật tài liệu ═════════════
ver = find_table('Phiên bản')
vt = Table(ver, d)
if not any(r.cells[0].text.strip() == '1.3' for r in vt.rows):
    ver.append(copy.deepcopy(vt.rows[-1]._tr))
    vt = Table(ver, d)
    for ci, txt in enumerate(['1.3', '28/09/2026', 'Đội phát triển phần mềm',
                              'Viết chi tiết theo từng bước kèm ảnh nút: cột Hành động, Chỉnh sửa (điểm khác Thêm '
                              'mới, Nhân viên phụ trách đại lý / Cấp đại lý, khối Địa chỉ giao hàng, quy tắc người '
                              'liên hệ), Khóa / Mở khóa, Lịch sử, Chi tiết, và toàn bộ màn Quản lý khách hàng (6 '
                              'thẻ, thêm / sửa / xóa thiết bị, serial, hình ảnh, tài liệu, video).']):
        cell_text(vt.rows[-1].cells[ci], txt)

# ═════════════ đánh lại số hình ═════════════
n = 0
for el in body.iterchildren(qn('w:p')):
    if re.match(r'Hình \d+:', P(el).text):
        n += 1
        done = False
        for t in el.iter(qn('w:t')):
            if t.text and re.search(r'Hình \d+', t.text):
                t.text = re.sub(r'Hình \d+', 'Hình %d' % n, t.text, count=1)
                done = True
                break
        assert done, P(el).text
d.save(OUT)
print('OK — số hình:', n)
