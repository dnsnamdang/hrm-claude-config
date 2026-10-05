"""python3 dump_tab.py <slug> [<slug>...] — tải bản MỚI NHẤT workbook testcase từ Drive, dump tab của danh mục ra <slug>/ref/testcase_tab.txt (R<hàng>: A=… ;; B=…). Chạy lại trước mỗi lần soạn tc để số hàng khớp sheet."""
import json, os, subprocess, sys
import openpyxl
H = os.path.dirname(os.path.abspath(__file__))
X = '/private/tmp/claude-501/-Users-manhcuong-Desktop-dns-HRM/97975d78-0044-4d8d-903b-30099dcfa96c/scratchpad/tc_now.xlsx'
subprocess.run(['rclone', 'backend', 'copyid', 'gdrive:', '1dbKcipbtpm-66H3nYyP-R21g3Uxdvgjj', X], capture_output=True)
wb = openpyxl.load_workbook(X)
idx = {x['slug']: x for x in json.load(open(os.path.join(H, 'index.json')))}
for slug in sys.argv[1:]:
    tab = idx[slug]['tab'].strip()
    ws = next(w for w in wb.worksheets if w.title.strip() == tab)
    out = []
    for r in range(1, ws.max_row + 1):
        cells = []
        for c in range(1, 18):
            v = ws.cell(r, c).value
            if v not in (None, ''):
                cells.append('%s=%s' % (openpyxl.utils.get_column_letter(c), str(v).replace('\n', ' / ')))
        if cells:
            out.append('R%d: %s' % (r, ' ;; '.join(cells)))
    p = os.path.join(H, slug, 'ref', 'testcase_tab.txt')
    open(p, 'w').write('\n'.join(out) + '\n')
    print(slug, tab, len(out), 'dòng →', p)
