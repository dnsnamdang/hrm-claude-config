# -*- coding: utf-8 -*-
"""Sinh "SRS - Quản lý giải pháp.docx" theo FORM CHUẨN (2026-09-24).

Chạy:  SRS_NO_WORD=1 /opt/homebrew/bin/python3 gen_srs.py
Nguồn đối chiếu code (nhánh gop_db):
  FE  pages/assign/solutions/{index,add}.vue · _id/{index,edit,manager}.vue · constants.js
      components/{SolutionForm,InfoTab,ModulesTab,OrgChartTab}.vue
      components/manager/{OverviewTab,CategoriesTable,PendingApprovalCard,ProjectInfoTab,ReviewProfilesTab,
        SolutionApprovalModal,ModuleApprovalModal,HumanResourceTab,AssignManagerModal,AssignManagerHistoryModal,
        TasksTab,IssueTab,MeetingsTab,ProgressTab,FilesTab,*UpcomingModal,*LateTasksModal}.vue
      components/modal/solution-version-modal.vue · prospective-projects/components/SolutionAdjustmentTab.vue
      components/menu-sidebar.js (menuItemsAssign › Làm giải pháp) · components/subsystem-menu/presale.js
  BE  Modules/Assign/Routes/api.php (prefix assign/solutions, middleware solutionMemberActive)
      Http/Controllers/Api/V1/SolutionController.php · Services/{SolutionService,SolutionMemberService}.php
      Http/Requests/Solution/*.php · Entities/Solution.php (STATUSES, canEdit, canDelete, isCanManage, getNextCode)
      Services/SolutionAdjustmentRequestService.php · ProspectiveProjectService::finalizeSolution
  Quyền: PermissionsTableSeeder id 1012, 1016-1019, 1044
Ảnh: shots/ (Playwright headless 1440x900, client :3002) — ảnh h_* / s_* dùng lại từ HDSD Giải pháp
(.plans/gop-db/hdsd-feedback-0210). Dữ liệu tạo thêm xem data_created.md.
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = HERE
while not os.path.isdir(os.path.join(ROOT, '.claude', 'skills', 'srs-documenter')):
    ROOT = os.path.dirname(ROOT)
sys.path.insert(0, os.path.join(ROOT, '.claude', 'skills', 'srs-documenter', 'assets'))
from srs_docx_lib import SrsDoc  # noqa: E402
from docx.shared import Inches  # noqa: E402

SHOTS = os.path.join(HERE, 'shots')


def shot(name):
    return os.path.join(SHOTS, name)


TEN_MAN = 'Quản lý giải pháp'
MENU = 'Phân hệ Công việc => Làm giải pháp => Quản lý giải pháp'
MENU_YC = 'Phân hệ CSKH trước bán => Yêu cầu giải pháp => Làm giải pháp'
MENU_DA = 'Phân hệ CSKH trước bán => Dự án TKT => Tạo giải pháp'
MQL = MENU + ' => Quản lý giải pháp'          # chặng thứ 4 = nút ở cột Hành động

A_ALL = 'Người dùng đã đăng nhập'
A_TP = 'Trưởng phòng giải pháp'
A_PM = 'PM làm giải pháp'
A_LD = 'Leader hạng mục'
A_KD = 'Nhân viên KD (dự án Tự triển khai)'

ICONS = {
    'Phân hệ Công việc': 'icon_phanhe_cv.png',
    'Phân hệ CSKH trước bán': 'icon_phanhe_presale.png',
    'Làm giải pháp': 'icon_nhom_lgp.png',
    'Quản lý giải pháp': 'icon_man_qlgp.png',
    'Yêu cầu giải pháp': 'icon_man_ycgp.png',
    'Dự án TKT': 'icon_man_duan.png',
    'Tạo giải pháp': 'icon_s_btn_taogp_row.png',
    'Tìm kiếm nâng cao': 'icon_timkiem.png',
    'Cài đặt bộ lọc': 'icon_caidat.png',
    'Cấu hình cột hiển thị': 'icon_cot.png',
    'Xuất Excel': 'icon_xuat.png',
    'Tạo mới': 'icon_taomoi.png',
    'Mã giải pháp': 'icon_ma.png',
    'Sửa': 'icon_sua.png',
    'Xóa': 'icon_xoa.png',
    'Giao cho Leader': 'icon_h_btn_giao_leader_row.png',
    'Lưu và duyệt': 'icon_h_btn_luu_duyet_row.png',
    'Tổng quan': 'icon_tab_tongquan.png',
    'Thông tin': 'icon_tab_thongtin.png',
    'Hồ sơ': 'icon_tab_hoso.png',
    'Nhân sự': 'icon_tab_nhansu.png',
    'Nhiệm vụ': 'icon_tab_nhiemvu.png',
    'Vấn đề giải pháp': 'icon_tab_vande.png',
    'Meeting': 'icon_tab_meeting.png',
    'Tiến độ': 'icon_tab_tiendo.png',
    'Files': 'icon_tab_files.png',
    'YC Điều chỉnh': 'icon_tab_ycdc.png',
    'Phân công': 'icon_phancong.png',
    'Xem lịch sử phân công': 'icon_lsphancong.png',
    'Thêm nhân sự': 'icon_themnhansu.png',
    'Sửa phân công': 'icon_suapc.png',
    'Khóa thành viên': 'icon_khoatv.png',
    'Xóa khỏi giải pháp': 'icon_xoatv.png',
    'Tạo hồ sơ trình duyệt giải pháp': 'icon_h_btn_tao_ho_so.png',
    'Duyệt': 'icon_h_btn_duyet_row.png',
    'Xem': 'icon_s_btn_xemhoso_row.png',
    'Tạo version': 'icon_h_btn_tao_version.png',
}

out = os.path.join(HERE, 'SRS - %s.docx' % TEN_MAN)
d = SrsDoc(out=out, menu=MENU, route='', full_url='', img_prefix='sol_')
d.set_menu_icons({k: shot(v) for k, v in ICONS.items()})

# "Quản lý giải pháp" vừa là MỤC MENU (chặng 3) vừa là NÚT ở cột Hành động (chặng 4) — 2 icon khác nhau
# cho cùng 1 chữ nên chọn icon theo VỊ TRÍ chặng.
_POS_ICON = {('Quản lý giải pháp', 3): shot('icon_quanly.png')}


def _menu_para(menu):
    par = d.p('Menu: ')
    for i, seg in enumerate(menu.split(' => ')):
        if i:
            par.add_run(' => ')
        for j, part in enumerate(seg.split(' / ')):
            if j:
                par.add_run(' / ')
            ic = _POS_ICON.get((part, i)) or d.menu_icons.get(part)
            par.add_run(part + (' ' if ic else ''))
            if ic:
                par.add_run().add_picture(ic, height=Inches(d.menu_icon_h))
    return par


d._menu_para = _menu_para

NO_PERM = '– Không thỏa điều kiện → nút không hiển thị; gọi thẳng chức năng thì hệ thống từ chối và dừng xử lý.'
SCOPE = '– Chỉ lấy giải pháp trong phạm vi được xem của người dùng (quy tắc BR-01).'

# ------------------------------------------------------------------ khung 1 chức năng
_N = [0]


def fn(title, intro, layout, ui, events, uc=None, rule=('- Màn Danh sách, Phân trang và UI/UX.', 'list'),
       ui_kind='input'):
    """Sinh trọn 1 mục 2.x. uc=(code, tên, nhóm, actor) → có Biểu đồ Usecase; None → bỏ mục đó.
    layout: list các phần tử ('menu', chuỗi) | ('p', chuỗi) | ('img', file, caption) | ('modal', tên)."""
    if not ui or not events:
        raise RuntimeError('Chức năng "%s" thiếu bảng giao diện / event' % title)
    _N[0] += 1
    n = _N[0]
    d.h3('2.%d %s' % (n, title))
    k = 1
    if uc:
        d.p('2.%d.%d Biểu đồ Usecase' % (n, k))
        d.uc_figure(uc[0], uc[1], uc[2], actor=uc[3], caption='Biểu đồ Use Case — %s %s' % (uc[0], uc[1]))
        k += 1
    d.p('2.%d.%d Giới thiệu' % (n, k)); k += 1
    d.rule_ref(rule[0] + ' Chỉ bổ sung các quy tắc riêng của màn Quản lý giải pháp tại phần mô tả chi tiết.',
               anchor=rule[1])
    d.intro_table(**intro)
    d.p('2.%d.%d Layout màn hình' % (n, k)); k += 1
    first = True
    for item in layout:
        if item[0] == 'menu':
            if first:
                d.p('Đường dẫn màn hình:')
                first = False
            _menu_para(item[1])
        elif item[0] == 'p':
            d.p(item[1])
        elif item[0] == 'modal':
            d.p('Cửa sổ %s được mở ngay trên màn hình theo đường dẫn ở trên.' % item[1])
        elif item[0] == 'img':
            d.figure(shot(item[1]), item[2], width_in=6.2)
    d.p('2.%d.%d Mô tả chi tiết giao diện' % (n, k)); k += 1
    if ui_kind == 'input':
        d.ui_table(ui)
    elif ui_kind == 'read':
        d.ui_table(ui, required=False)
    else:
        d.ui_table(ui, required=False, scope=False)
    d.p('2.%d.%d Danh sách event và xử lý event' % (n, k))
    d.event_table(events)


d.title_block(TEN_MAN)
d.h2('Mục lục')
d.toc()

# ========================================================= PHẦN 1
d.h1('Phần 1. Giới thiệu')
d.h2('1 Mục đích')
d.p('Tài liệu này đặc tả yêu cầu phần mềm cho màn Quản lý giải pháp (phân hệ Công việc, nhóm Làm giải pháp) '
    'và màn Quản lý giải pháp nhiều tab mở từ từng giải pháp, nhằm:')
d.bullets([
    'Là căn cứ nghiệm thu các chức năng: xem danh sách, tìm kiếm/lọc, cài đặt bộ lọc, tùy chỉnh cột, xuất Excel, '
    'tạo mới (luồng có Yêu cầu làm giải pháp và luồng dự án Tự triển khai), xem chi tiết, chỉnh sửa, xóa, PM duyệt '
    'tiếp nhận, Leader duyệt hạng mục, quản lý nhân sự, phân công PM/Leader, hồ sơ trình duyệt, tạo version mới, '
    'phân bổ tiến độ và các tab Nhiệm vụ / Vấn đề / Meeting / Files / Yêu cầu điều chỉnh.',
    'Làm rõ vòng đời trạng thái của giải pháp qua các vai trò Trưởng phòng giải pháp – PM – Leader hạng mục, và '
    'điểm khác của dự án Tự triển khai (nhân viên KD tự làm giải pháp).',
    'Làm rõ phạm vi giải pháp mỗi người dùng được xem theo quyền cấp tổ chức và vai trò trong giải pháp.',
])
d.h2('2 Thuật ngữ và viết tắt')
d.table(['Thuật ngữ', 'Mô tả'], [
    ('GP', 'Giải pháp kỹ thuật lập cho một dự án tiền khả thi (dự án TKT). Mỗi dự án chỉ có 1 giải pháp.'),
    ('YCGP', 'Yêu cầu làm giải pháp do KD gửi phòng giải pháp. Người tiếp nhận yêu cầu là người được tạo giải pháp '
     'từ yêu cầu đó.'),
    ('Trưởng phòng giải pháp (TP)', 'Người tạo giải pháp (thường là người tiếp nhận YCGP). Là người duyệt hồ sơ '
     'trình duyệt và được phân công PM.'),
    ('PM làm giải pháp (PM)', 'Người phụ trách giải pháp: duyệt tiếp nhận, giao Leader, lập hồ sơ trình duyệt, tạo '
     'version, phân bổ tiến độ.'),
    ('Leader hạng mục', 'Người phụ trách 1 hạng mục của giải pháp, duyệt hạng mục ở bước Chờ Leader duyệt.'),
    ('Hạng mục', 'Phần việc con của giải pháp (khi giải pháp “Dự án có hạng mục”), có Leader, ngày cần hoàn thành, '
     'bảng phân công nhân sự.'),
    ('Dự án Tự triển khai', 'Dự án TKT có “Cách triển khai” = Tự triển khai: KD tự làm giải pháp, không qua YCGP, '
     'không có hạng mục, không có bước duyệt PM/Leader/TP.'),
    ('Hồ sơ trình duyệt', 'Bộ hồ sơ (nội dung + tài liệu + BOM tổng hợp) PM lập để trình TP duyệt giải pháp.'),
    ('BOM tổng hợp', 'Danh mục hàng hóa tổng hợp của 1 version giải pháp; hồ sơ trình duyệt tự gắn BOM tổng hợp ở '
     'trạng thái Hoàn thành của version hiện tại.'),
    ('Version', 'Phiên bản giải pháp V1, V2…; tạo version mới khi cần làm lại giải pháp sau khi đã duyệt.'),
    ('YC điều chỉnh', 'Yêu cầu điều chỉnh giải pháp do người tạo dự án gửi sau khi giải pháp đã duyệt.'),
], widths=[1.8, 4.2])

# ========================================================= PHẦN 2
d.h1('Phần 2. Phân quyền')
d.h2('1 Danh sách quyền')
d.p('Nhóm quyền thao tác:')
d.table(['Ký hiệu', 'Tên quyền', 'Tác dụng trên màn hình'], [
    ('Q1', 'Tiếp nhận yêu cầu làm giải pháp',
     'Được mở màn Quản lý giải pháp của các giải pháp thuộc phòng tiếp nhận mình quản lý; được sửa giải pháp ở '
     'trạng thái Chờ duyệt giải pháp / Đã duyệt giải pháp (luồng có YCGP).'),
    ('Q2', 'Quản lý giải pháp',
     'Chỉ gắn cho thao tác Đóng giải pháp ở máy chủ; màn hình hiện CHƯA có nút Đóng nên quyền này chưa có tác dụng '
     'trên giao diện.'),
], widths=[0.8, 2.0, 3.2])
d.p('Phần lớn thao tác KHÔNG gắn quyền mà do VAI TRÒ của người dùng trên từng giải pháp quyết định:')
d.table(['Ký hiệu', 'Vai trò', 'Tác dụng trên màn hình'], [
    ('R1', 'Trưởng phòng giải pháp (người tạo GP)', 'Sửa/xóa giải pháp Nháp; phân công PM và Leader; quản lý '
     'thành viên; duyệt / từ chối hồ sơ trình duyệt; xử lý YC điều chỉnh.'),
    ('R2', 'PM làm giải pháp', 'Duyệt tiếp nhận (Giao cho Leader / Lưu và duyệt); sửa giải pháp khi Chờ PM duyệt, '
     'Chờ Leader duyệt, Đang triển khai; phân công Leader; quản lý thành viên; lập và trình hồ sơ; duyệt hồ sơ hạng '
     'mục; phân bổ tiến độ; xử lý YC điều chỉnh.'),
    ('R3', 'Leader hạng mục', 'Sửa và duyệt hạng mục mình phụ trách khi giải pháp Chờ Leader duyệt; quản lý thành '
     'viên của hạng mục mình.'),
    ('R4', 'Thành viên giải pháp', 'Xem giải pháp mình tham gia, mở màn Quản lý (nếu đủ điều kiện BR-08).'),
], widths=[0.8, 2.0, 3.2])
d.p('Nhóm quyền quyết định phạm vi dữ liệu:')
d.table(['Ký hiệu', 'Tên quyền', 'Phạm vi dữ liệu'], [
    ('V1', 'Xem danh sách làm giải pháp theo tổng công ty', 'Toàn bộ giải pháp.'),
    ('V2', 'Xem danh sách làm giải pháp theo công ty', 'Giải pháp thuộc công ty đang làm việc.'),
    ('V3', 'Xem danh sách làm giải pháp theo phòng ban',
     'Giải pháp thuộc phòng ban / bộ phận mình quản lý và giải pháp do mình tạo.'),
    ('V4', 'Xem danh sách làm giải pháp theo bộ phận', 'Giải pháp thuộc bộ phận mình quản lý và do mình tạo.'),
    ('–', 'Không có quyền nào', 'Giải pháp do mình tạo; cộng thêm (với mọi người) giải pháp mình là PM, Leader '
     'hạng mục hoặc thành viên.'),
], widths=[0.8, 2.2, 3.0])

d.h2('2 Ma trận phân quyền')
Y, N = '✅', '❌'
d.table(['Chức năng', 'Q1', 'R1', 'R2', 'R3', 'R4', 'Không có quyền nào'], [
    ('FR-01 Xem danh sách giải pháp', Y, Y, Y, Y, Y, '✅ (theo phạm vi)'),
    ('FR-02 Tìm kiếm và lọc', Y, Y, Y, Y, Y, Y),
    ('FR-03 Cài đặt bộ lọc', Y, Y, Y, Y, Y, Y),
    ('FR-04 Tùy chỉnh cột', Y, Y, Y, Y, Y, Y),
    ('FR-05 Xuất Excel', Y, Y, Y, Y, Y, Y),
    ('FR-06 Tạo mới – luồng có YCGP', '✅ (người tiếp nhận YCGP)', Y, N, N, N, N),
    ('FR-07 Tạo mới – dự án Tự triển khai', N, N, N, N, N, '✅ (KD có dự án đủ điều kiện)'),
    ('FR-08 Xem chi tiết giải pháp', Y, Y, Y, Y, Y, '✅ (theo phạm vi)'),
    ('FR-09 Chỉnh sửa giải pháp', '✅ (Chờ duyệt / Đã duyệt GP)', '✅ (Nháp)', Y, '✅ (HM chưa duyệt)', N, N),
    ('FR-10 Xóa giải pháp', N, '✅ (Nháp)', N, N, N, N),
    ('FR-11 PM duyệt tiếp nhận', N, N, Y, N, N, N),
    ('FR-12 Leader duyệt hạng mục', N, N, N, Y, N, N),
    ('FR-13 Mở màn Quản lý giải pháp', '✅ (phòng tiếp nhận)', Y, Y, N, N, N),
    ('FR-14 Xem tab Thông tin', Y, Y, Y, N, N, N),
    ('FR-15 Phân công PM / Leader', N, Y, '✅ (chỉ Leader)', N, N, N),
    ('FR-16 Xem lịch sử phân công', Y, Y, Y, N, N, N),
    ('FR-17 Thêm nhân sự', Y, Y, Y, N, N, N),
    ('FR-18 Sửa / Khóa / Xóa thành viên', N, Y, Y, '✅ (HM của mình)', N, N),
    ('FR-19 Danh sách hồ sơ trình duyệt', Y, Y, Y, N, N, N),
    ('FR-20 Tạo / trình hồ sơ trình duyệt', N, N, Y, N, N, N),
    ('FR-21 Duyệt / từ chối hồ sơ trình duyệt', N, Y, N, N, N, N),
    ('FR-22 PM duyệt hồ sơ hạng mục', N, N, Y, N, N, N),
    ('FR-23 Tạo version mới', Y, Y, Y, N, N, N),
    ('FR-24 Phân bổ % tiến độ', N, N, Y, N, N, N),
    ('FR-25 Tab Nhiệm vụ', Y, Y, Y, N, N, N),
    ('FR-26 Tab Vấn đề giải pháp', Y, Y, Y, N, N, N),
    ('FR-27 Tab Meeting', Y, Y, Y, N, N, N),
    ('FR-28 Tab Files', Y, Y, Y, N, N, N),
    ('FR-29 Xử lý YC điều chỉnh', N, Y, Y, N, N, N),
], widths=[2.3, 0.75, 0.6, 0.6, 0.6, 0.45, 0.7])
d.p('R1–R4 là vai trò trên từng giải pháp, không phải quyền gán cho vai trò hệ thống. Chức năng FR-13 → FR-29 nằm '
    'trên màn Quản lý giải pháp; chỉ người mở được màn đó (BR-08) mới dùng được. Với dự án Tự triển khai, nhân viên '
    'KD tạo giải pháp đồng thời là R1 và R2. Nút thao tác không dùng được thì ẩn hẳn.')

# ========================================================= PHẦN 3
d.h1('Phần 3. Đặc tả chi tiết theo từng chức năng')
d.h2('1 Sơ đồ UML tổng quan')
d.overview_figure2(
    [(A_ALL, [0, 1, 8]), (A_TP, [2, 4, 5, 9, 13, 17]), (A_PM, [4, 6, 9, 10, 11, 12, 14, 15, 16, 17]),
     (A_LD, [4, 7, 11]), (A_KD, [3, 12])],
    [('FR-01', 'Xem danh sách giải pháp', 'view'),
     ('FR-05', 'Xuất Excel', 'io'),
     ('FR-06', 'Tạo mới – luồng có YCGP', 'crud'),
     ('FR-07', 'Tạo mới – Tự triển khai', 'crud'),
     ('FR-09', 'Chỉnh sửa giải pháp', 'crud'),
     ('FR-10', 'Xóa giải pháp', 'action'),
     ('FR-11', 'PM duyệt tiếp nhận', 'action'),
     ('FR-12', 'Leader duyệt hạng mục', 'action'),
     ('FR-13', 'Quản lý giải pháp', 'view'),
     ('FR-15', 'Phân công PM / Leader', 'action'),
     ('FR-17', 'Thêm nhân sự', 'crud'),
     ('FR-18', 'Sửa / Khóa / Xóa thành viên', 'action'),
     ('FR-20', 'Tạo / trình hồ sơ', 'crud'),
     ('FR-21', 'Duyệt / từ chối hồ sơ', 'action'),
     ('FR-22', 'Duyệt hồ sơ hạng mục', 'action'),
     ('FR-23', 'Tạo version mới', 'action'),
     ('FR-24', 'Phân bổ % tiến độ', 'crud'),
     ('FR-29', 'Xử lý YC điều chỉnh', 'action')],
    [('FR-02', 'Tìm kiếm và lọc', 'view', 'extend', [0], None),
     ('FR-03', 'Cài đặt bộ lọc', 'view', 'extend', [0], None),
     ('FR-04', 'Tùy chỉnh cột', 'view', 'extend', [0], None),
     ('FR-08', 'Xem chi tiết giải pháp', 'view', 'extend', [0], None),
     ('FR-14', 'Tab Thông tin', 'view', 'extend', [8], None),
     ('FR-16', 'Lịch sử phân công', 'view', 'extend', [8], None),
     ('FR-19', 'Danh sách hồ sơ', 'view', 'extend', [8], None),
     ('FR-25', 'Tab Nhiệm vụ', 'view', 'extend', [8], None),
     ('FR-26', 'Tab Vấn đề giải pháp', 'view', 'extend', [8], None),
     ('FR-27', 'Tab Meeting', 'view', 'extend', [8], None),
     ('FR-28', 'Tab Files', 'view', 'extend', [8], None)],
    'Sơ đồ Use Case tổng quan màn %s' % TEN_MAN)

d.h2('2 Đặc tả chi tiết từng chức năng')

# ---------------------------------------------------------------- 2.1 FR-01
fn('Xem danh sách giải pháp',
   dict(ten='Xem danh sách giải pháp',
        mota='Hiển thị bảng các giải pháp người dùng được xem, mới tạo trước; mỗi dòng có các nút thao tác theo vai '
             'trò của người dùng trên giải pháp đó.',
        tacnhan=A_ALL,
        dieukien='Người dùng đã đăng nhập hệ thống.',
        chinh='1. Người dùng mở màn Quản lý giải pháp theo đường dẫn menu.\n'
              '2. Hệ thống nạp danh sách giải pháp trong phạm vi được xem, sắp xếp Ngày tạo giảm dần, 10 dòng/trang.\n'
              '3. Hệ thống hiển thị bảng “Danh sách làm giải pháp” với đủ 26 cột (mặc định hiện hết cột).\n'
              '4. Mỗi dòng hiện các nút người dùng được phép: Sửa, Xóa, Quản lý giải pháp, Giao cho Leader / Lưu và duyệt.',
        phu='• Không có giải pháp nào khớp → bảng hiển thị “Không có dữ liệu phù hợp bộ lọc.”.\n'
            '• Lỗi khi nạp → thông báo “Lỗi khi tải dữ liệu”.\n'
            '• Quay lại màn trong vòng 10 phút → khôi phục bộ lọc và trạng thái đóng/mở khối lọc đã dùng.\n'
            '• Giải pháp Nháp của người khác không hiển thị.'),
   [('menu', MENU), ('img', '01-danh-sach.png', 'Danh sách làm giải pháp lúc mới mở'),
    ('img', '01b-danh-sach-phai.png', 'Các cột bên phải và cột Hành động (cuộn ngang bảng)')],
   [('Tiêu đề trang / tiêu đề bảng', 'Label', 'Hiển thị', '–', 'Danh sách làm giải pháp', '–'),
    ('Nút Tạo mới', 'Button', 'Enable', '–', 'Hiển thị', 'Mở màn Tạo giải pháp (FR-06).'),
    ('Nút Xuất Excel', 'Button', 'Enable / Disable', '–', 'Hiển thị', 'Khóa trong lúc đang xuất (FR-05).'),
    ('Nút Cấu hình cột hiển thị', 'Icon Button', 'Enable', '–', 'Hiển thị', 'Mở cửa sổ Tùy chỉnh cột (FR-04).'),
    ('Cột STT', 'Table/Grid', 'Read-only', '–', 'Theo trang', 'Ghim trái, không ẩn được.'),
    ('Cột Mã giải pháp', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'Ghim trái, không ẩn được; bấm để xem chi tiết (FR-08). Có sắp xếp.'),
    ('Cột Tên giải pháp', 'Table/Grid', 'Read-only', '0–255 ký tự', 'Theo dữ liệu', 'Tối đa 2 dòng. Có sắp xếp.'),
    ('Cột Khách hàng', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', '“Mã - Tên khách hàng” của dự án.'),
    ('Cột Yêu cầu làm GP', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     '“Mã - Tên YCGP”, dòng phụ “Ngày gửi yêu cầu: …”. Để trống với dự án Tự triển khai.'),
    ('Cột Dự án', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', 'Tên dự án TKT.'),
    ('Cột Khách hàng cuối', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     '“Mã - Tên”, dòng phụ “Loại hình: …”.'),
    ('Cột Giai đoạn dự án', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', '–'),
    ('Cột Mức độ ưu tiên', 'Badge', 'Read-only', '–', 'Theo dữ liệu', 'Chữ và màu theo giai đoạn dự án.'),
    ('Cột Phòng làm GP / PM phụ trách GP', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', '–'),
    ('Cột Ngày hoàn thành GP', 'Table/Grid', 'Read-only', 'dd/mm/yyyy', 'Theo dữ liệu',
     '“Dự kiến (Vn): <ngày cần xong GP>”.'),
    ('Cột Version hiện tại (Tiến độ %)', 'Table/Grid', 'Read-only', '0 – 100', 'Theo dữ liệu',
     'Mã version, dòng phụ “Tiến độ: n%”.'),
    ('Cột Phòng KD phụ trách dự án / KD phụ trách chính', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'KD phụ trách có dòng phụ “SĐT: …”.'),
    ('Cột Nhóm ngành / Nhóm giải pháp / Ứng dụng', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', '–'),
    ('Cột Loại hình hoạt động KH / Lĩnh vực kinh doanh KH', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', '–'),
    ('Cột Người tạo / Người cập nhật', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', '“Tên - Mã phòng”.'),
    ('Cột Ngày tạo / Ngày cập nhật', 'Table/Grid', 'Read-only', 'dd/mm/yyyy hh:mm', 'Theo dữ liệu', 'Có sắp xếp.'),
    ('Cột Tiến trình GP', 'Badge', 'Read-only', 'Danh sách 10 giá trị', 'Theo dữ liệu', 'Chữ và màu theo BR-03.'),
    ('Cột Hành động', 'Icon Button', 'Enable / Ẩn', '–', 'Theo vai trò',
     'Không ẩn được. Sửa · Xóa · Quản lý giải pháp · Giao cho Leader / Lưu và duyệt; nút không đủ điều kiện thì ẩn.'),
    ('Phân trang', 'Pagination', 'Enable', '10 / 20 / 50 / 100', '10 dòng/trang', '–'),
    ('Trạng thái rỗng', 'Label', 'Hiển thị', '–', 'Ẩn', '“Không có dữ liệu phù hợp bộ lọc.”')],
   [('Mở màn Quản lý giải pháp', 'System',
     'During:\n' + SCOPE + '\n– Không lấy giải pháp Nháp của người khác.\n'
     'After:\n– Hiển thị bảng; lỗi → “Lỗi khi tải dữ liệu”.'),
    ('Đổi trang / số dòng mỗi trang', 'Click', 'After:\n– Nạp lại đúng trang; đổi số dòng thì về trang 1.'),
    ('Bấm tiêu đề cột có sắp xếp', 'Click',
     'After:\n– Sắp xếp tăng/giảm theo Mã giải pháp, Tên giải pháp, Ngày tạo hoặc Ngày cập nhật.'),
    ('Bấm Mã giải pháp', 'Click', 'After:\n– Mở màn Chi tiết giải pháp (FR-08).'),
    ('Bấm nút ở cột Hành động', 'Click',
     'After:\n– Sửa → màn Sửa (FR-09); Xóa → hộp xác nhận (FR-10); Quản lý giải pháp → màn Quản lý (FR-13); '
     'Giao cho Leader / Lưu và duyệt → màn Duyệt giải pháp (FR-11, FR-12).')],
   ui_kind='read')

# ---------------------------------------------------------------- 2.2 FR-02
fn('Tìm kiếm và lọc',
   dict(ten='Tìm kiếm và lọc giải pháp',
        mota='Tìm nhanh theo từ khóa và lọc nâng cao theo 15 nhóm tiêu chí.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn danh sách giải pháp.',
        chinh='1. Người dùng gõ từ khóa vào ô tìm nhanh rồi bấm “Tìm kiếm” (hoặc Enter).\n'
              '2. Hoặc bấm “Tìm kiếm nâng cao” để mở khối lọc và chọn giá trị ở các ô lọc.\n'
              '3. Chọn giá trị ở ô chọn thì hệ thống lọc ngay; ô gõ tay (Khách hàng, Khách hàng cuối) chọn xong '
              'mới lọc.\n'
              '4. Hệ thống nạp lại danh sách từ trang 1.',
        phu='• Đổi Công ty → xóa Phòng làm GP, Bộ phận. Đổi Khách hàng → xóa Dự án và chỉ còn dự án của khách hàng đó.\n'
            '• Đổi Nhóm ngành → xóa Nhóm giải pháp, Ứng dụng; đổi Nhóm giải pháp → xóa Ứng dụng.\n'
            '• Đổi Loại hình hoạt động KH → xóa Lĩnh vực kinh doanh KH. Đổi Phòng KD phụ trách → xóa KD phụ trách.\n'
            '• Chọn Giai đoạn dự án → tự điền Mức độ ưu tiên của giai đoạn đó.\n'
            '• Bấm “Làm mới” → xóa mọi điều kiện lọc, nạp lại danh sách 1 lần.'),
   [('menu', MENU + ' => Tìm kiếm nâng cao'), ('img', '02-loc.png', 'Khối Tìm kiếm nâng cao đang mở')],
   [('Ô tìm nhanh', 'Textbox', 'Enable', '–', 'Không', 'Trống',
     'Gợi ý “Tìm theo mã giải pháp, tên giải pháp, mã khách hàng, tên khách hàng”.'),
    ('Nút Tìm kiếm / Làm mới', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Nút Tìm kiếm nâng cao / Ẩn tìm kiếm nâng cao', 'Button', 'Enable', '–', '–', 'Thu gọn',
     'Đóng/mở khối lọc; danh mục của ô lọc chỉ nạp lần đầu mở khối.'),
    ('Công ty – Phòng làm GP – Bộ phận – PM phụ trách', 'Dropdown', 'Enable / Ẩn', 'Danh sách', 'Không', 'Trống',
     '4 ô; ô Công ty/Phòng/Bộ phận hiện theo quyền V1–V4. Phòng làm GP lọc theo phòng tiếp nhận của YCGP.'),
    ('Khách hàng / Khách hàng cuối', 'Textbox', 'Enable', '≥ 2 ký tự', 'Không', 'Trống',
     'Gõ tối thiểu 2 ký tự để hiện danh sách gợi ý (tên, mã • SĐT); “Không tìm thấy khách hàng” khi không có.'),
    ('Dự án', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống', 'Lọc theo Khách hàng đã chọn.'),
    ('Tiến trình giải pháp', 'Dropdown', 'Enable', 'Danh sách 10 giá trị', 'Không', 'Trống', 'Xem BR-03.'),
    ('Giai đoạn dự án', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống', 'Kèm icon ⓘ mô tả giai đoạn.'),
    ('Mức độ ưu tiên', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống', '–'),
    ('Phòng KD phụ trách chính / KD phụ trách chính', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống',
     'KD phụ trách lọc theo phòng đã chọn.'),
    ('Nhóm ngành / Nhóm giải pháp / Ứng dụng', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống',
     'Lọc theo cấp cha.'),
    ('Loại hình hoạt động KH / Lĩnh vực kinh doanh KH', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống',
     'Lĩnh vực lọc theo loại hình.'),
    ('Ngày cập nhật', 'Datepicker', 'Enable', 'dd/mm/yyyy', 'Không', 'Trống', 'Khoảng từ – đến trong 1 ô.')],
   [('Bấm Tìm kiếm / Enter ở ô tìm nhanh', 'Click',
     'After:\n– Tìm giải pháp có tên hoặc mã chứa từ khóa, hoặc dự án có mã / tên khách hàng chứa từ khóa.\n'
     '– Nạp lại từ trang 1.'),
    ('Chọn giá trị ô lọc', 'Change', 'After:\n– Lọc ngay, về trang 1.\n' + SCOPE),
    ('Gõ ô Khách hàng / Khách hàng cuối', 'Keypress',
     'After:\n– Từ 2 ký tự: tìm khách hàng, hiện “Đang tìm...” rồi danh sách gợi ý.\n– Chọn 1 khách hàng → lọc.\n'
     '– Xóa trắng ô → bỏ lọc.'),
    ('Bấm Làm mới', 'Click', 'After:\n– Xóa mọi điều kiện lọc, nạp lại trang 1.')],
   rule=('- Kịch bản tìm kiếm, Bộ lọc và Dropdown.', 'search'))

# ---------------------------------------------------------------- 2.3 FR-03
fn('Cài đặt bộ lọc',
   dict(ten='Cài đặt bộ lọc',
        mota='Chọn các ô lọc hiện trong khối Tìm kiếm nâng cao và sắp xếp thứ tự; cấu hình lưu riêng cho màn này của '
             'từng người dùng.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn danh sách giải pháp.',
        chinh='1. Người dùng bấm “Cài đặt bộ lọc”.\n2. Hệ thống mở cửa sổ với 15 trường lọc theo cấu hình hiện tại.\n'
              '3. Người dùng tích / bỏ tích và kéo biểu tượng ⠿ để sắp xếp.\n4. Bấm “Lưu”.\n'
              '5. Hệ thống lưu cấu hình, đóng cửa sổ, báo “Cập nhật thành công” và vẽ lại khối lọc.',
        phu='• Bấm “Khôi phục mặc định” → hiện đủ 15 trường theo thứ tự gốc (chưa lưu tới khi bấm Lưu).\n'
            '• Bấm “Đóng” / dấu × → không lưu.\n• Lưu lỗi → “Thao tác thất bại”.'),
   [('menu', MENU + ' => Cài đặt bộ lọc'), ('modal', 'Cài đặt bộ lọc'),
    ('img', '03-cai-dat-loc.png', 'Cửa sổ Cài đặt bộ lọc')],
   [('Tiêu đề cửa sổ', 'Label', 'Hiển thị', '–', '–', 'Cài đặt bộ lọc', '–'),
    ('Dòng hướng dẫn', 'Label', 'Hiển thị', '–', '–', 'Hiển thị',
     '“Tích chọn trường lọc muốn hiển thị; kéo ⠿ để sắp xếp thứ tự. Cài đặt được lưu theo từng màn hình.”'),
    ('Danh sách trường lọc', 'Checkbox', 'Enable', 'Danh sách 15 giá trị', 'Không', 'Theo cấu hình đã lưu',
     'Công ty – Phòng làm GP – Bộ phận – PM phụ trách · Khách hàng · Khách hàng cuối · Dự án · Tiến trình giải pháp · '
     'Giai đoạn dự án · Mức độ ưu tiên · Phòng KD phụ trách chính · KD phụ trách chính · Nhóm ngành · Nhóm giải pháp · '
     'Ứng dụng · Loại hình hoạt động khách hàng · Lĩnh vực kinh doanh khách hàng · Ngày cập nhật.'),
    ('Nút Lưu / Khôi phục mặc định / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Cài đặt bộ lọc', 'Click', 'After:\n– Mở cửa sổ với cấu hình đang áp dụng.'),
    ('Tích / bỏ tích, kéo sắp xếp', 'Click', 'After:\n– Cập nhật danh sách trong cửa sổ, chưa lưu.'),
    ('Bấm Lưu', 'Click', 'After:\n– Lưu cấu hình theo người dùng + màn; “Cập nhật thành công”; lỗi → “Thao tác thất bại”.'),
    ('Bấm Khôi phục mặc định / Đóng', 'Click', 'After:\n– Khôi phục danh sách gốc / đóng cửa sổ không lưu.')],
   rule=('- Bộ lọc và UI/UX.', 'search'))

# ---------------------------------------------------------------- 2.4 FR-04
fn('Tùy chỉnh cột',
   dict(ten='Tùy chỉnh cột hiển thị',
        mota='Ẩn/hiện và sắp xếp thứ tự các cột của bảng danh sách; cấu hình lưu theo từng người dùng.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn danh sách giải pháp.',
        chinh='1. Người dùng bấm nút Cấu hình cột hiển thị.\n2. Hệ thống mở cửa sổ “Tùy chỉnh cột” liệt kê đủ 26 cột.\n'
              '3. Người dùng tích / bỏ tích, kéo biểu tượng ☰ để đổi thứ tự.\n4. Bấm “Lưu”.\n'
              '5. Hệ thống lưu cấu hình và vẽ lại bảng.',
        phu='• Cột STT, Mã giải pháp, Hành động bị khóa: hiện xám, không bỏ tích và không kéo được.\n'
            '• Bấm “Đóng” → không lưu.'),
   [('menu', MENU + ' => Cấu hình cột hiển thị'), ('modal', 'Tùy chỉnh cột'),
    ('img', '04-tuy-chinh-cot.png', 'Cửa sổ Tùy chỉnh cột')],
   [('Tiêu đề cửa sổ', 'Label', 'Hiển thị', '–', '–', 'Tùy chỉnh cột', '–'),
    ('Danh sách cột', 'Checkbox', 'Enable / Disable', 'Danh sách 26 giá trị', 'Không', 'Hiện hết cột',
     'Mặc định hiện hết cột (ngoại lệ có chủ ý). STT, Mã giải pháp, Hành động: khóa (biểu tượng ổ khóa).'),
    ('Tay kéo ☰', 'Icon Button', 'Enable / Ẩn', '–', '–', 'Hiển thị', 'Kéo thả đổi thứ tự; cột khóa không có.'),
    ('Nút Lưu / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Cấu hình cột hiển thị', 'Click', 'After:\n– Mở cửa sổ với cấu hình đã lưu.'),
    ('Tích / bỏ tích, kéo thả', 'Click', 'After:\n– Cập nhật trong cửa sổ; cột khóa không thay đổi được.'),
    ('Bấm Lưu', 'Click', 'After:\n– Lưu cấu hình theo người dùng cho màn này; bảng vẽ lại theo cấu hình mới.'),
    ('Bấm Đóng', 'Click', 'After:\n– Đóng cửa sổ, không lưu.')],
   rule=('- Quy tắc Excel và Cấu hình cột.', 'excel'))

# ---------------------------------------------------------------- 2.5 FR-05
fn('Xuất Excel',
   dict(ten='Xuất Excel danh sách giải pháp',
        mota='Xuất ra file Excel toàn bộ giải pháp khớp bộ lọc đang áp dụng, theo các cột người dùng chọn.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn danh sách giải pháp.',
        chinh='1. Người dùng bấm “Xuất Excel”.\n2. Hệ thống mở cửa sổ “Chọn trường xuất file”, tích sẵn đúng các cột '
              'đang hiện trên bảng.\n3. Người dùng tích / bỏ tích, kéo để đổi thứ tự cột.\n4. Bấm “Xuất file”.\n'
              '5. Hệ thống tải về file danh_sach_lam_giai_phap.xlsx và báo “Xuất Excel thành công”.',
        phu='• “Chọn tất cả” / “Bỏ chọn hết” để tích nhanh.\n• Lỗi → “Lỗi khi xuất Excel”.\n'
            '• Trong lúc đang xuất, nút Xuất Excel bị khóa.'),
   [('menu', MENU + ' => Xuất Excel'), ('modal', 'Chọn trường xuất file'),
    ('img', '05-xuat.png', 'Cửa sổ Chọn trường xuất file')],
   [('Tiêu đề cửa sổ', 'Label', 'Hiển thị', '–', '–', 'Chọn trường xuất file', '–'),
    ('Danh sách trường', 'Checkbox', 'Enable', 'Danh sách 28 giá trị', 'Có (≥ 1)', 'Tích sẵn cột đang hiện',
     'Mã/Tên giải pháp, Mã/Tên khách hàng, Mã/Tên YCGP, Dự án, Khách hàng cuối, Giai đoạn, Mức độ ưu tiên, Phòng làm '
     'GP, PM phụ trách, Tiến trình GP, Ngày hoàn thành GP (dự kiến), Version hiện tại, Tiến độ (%), Phòng KD, KD phụ '
     'trách, SĐT KD, Nhóm ngành, Nhóm giải pháp, Ứng dụng, Loại hình, Lĩnh vực, Người tạo, Ngày tạo, Người cập nhật, '
     'Ngày cập nhật.'),
    ('Dòng đếm', 'Label', 'Hiển thị', '–', '–', 'Theo số tích', '“Đang chọn a/28 trường”.'),
    ('Nút Chọn tất cả / Bỏ chọn hết', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Nút Xuất file / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Xuất Excel', 'Click', 'After:\n– Mở cửa sổ chọn trường.'),
    ('Bấm Xuất file', 'Click',
     'Before:\n– Không kiểm tra quyền riêng; phạm vi dữ liệu theo BR-01.\n'
     'After:\n– Xuất tất cả dòng khớp bộ lọc (không phân trang), cột theo đúng thứ tự đã chọn.\n'
     '– Thành công: “Xuất Excel thành công”; lỗi: “Lỗi khi xuất Excel”.')],
   uc=('FR-05', 'Xuất Excel danh sách giải pháp', 'io', A_ALL),
   rule=('- Quy tắc Excel và Cấu hình cột.', 'excel'))

# ---------------------------------------------------------------- 2.6 FR-06
UI_FORM = [
    ('Tab Thông tin / Quản lý hạng mục / Sơ đồ nhân sự', 'Tab', 'Enable', '–', '–', 'Thông tin',
     'Tab Quản lý hạng mục chỉ có khi tích “Dự án có hạng mục”. Lưu bị lỗi thì tự chuyển tới tab có lỗi.'),
    ('Khách hàng / Khách hàng cuối', 'Text', 'Read-only', '–', '–', 'Theo dự án',
     '“Mã - Tên”, bấm mở chi tiết khách hàng ở tab mới. Khách hàng cuối chỉ hiện khi KH trực tiếp là KH trung gian.'),
    ('Dự án', 'Text', 'Read-only', '–', '–', 'Theo YCGP',
     'Tên dự án + KD phụ trách, Giai đoạn (ⓘ mô tả), Ngày hoàn thành dự án; bấm mở màn quản lý dự án.'),
    ('Yêu cầu làm GP', 'Dropdown', 'Enable', 'Danh sách', 'Có', 'Trống',
     'Chỉ các YCGP ở trạng thái Đã tiếp nhận mà người dùng là người tiếp nhận. Chọn xong hiển thị “Mã - Tên YCGP”, '
     'Ngày KH cần GP, Ưu tiên. Ẩn khi mở từ YCGP (đã chọn sẵn) và khi dự án Tự triển khai.'),
    ('Phòng làm GP', 'Textbox', 'Read-only', '–', '–', 'Phòng tiếp nhận YCGP', '–'),
    ('Mã GP', 'Textbox', 'Read-only', '–', '–', 'Sinh tự động', 'Theo BR-02.'),
    ('PM làm giải pháp', 'Dropdown', 'Enable / Disable', 'Danh sách nhân viên', 'Có', 'PM ghi trên YCGP',
     '“Tên - Mã phòng - Mã NV”. Khóa với dự án Tự triển khai (PM = người tạo).'),
    ('Tên GP', 'Textbox', 'Enable', '0–255 ký tự', 'Có', 'Trống', 'Không trùng tên giải pháp khác.'),
    ('Ngày cần xong GP', 'Datepicker', 'Enable', 'dd/mm/yyyy', 'Có', 'Theo YCGP', '≥ hôm nay khi giải pháp Nháp.'),
    ('Ứng dụng', 'Textbox', 'Read-only', '–', 'Có', 'Theo dự án', 'Kế thừa dự án, không sửa được.'),
    ('Nhóm ngành / Nhóm giải pháp', 'Dropdown', 'Enable', 'Danh sách theo Ứng dụng', 'Có', 'Theo YCGP',
     'ⓘ “Nếu không tìm thấy … hãy liên hệ với bộ phận Master Data để thêm mới”.'),
    ('Mô tả', 'Textarea', 'Enable', '–', 'Không', 'Trống', '–'),
    ('Dự án có hạng mục', 'Checkbox', 'Enable', '–', 'Không', 'Tích', 'Ẩn với dự án Tự triển khai.'),
    ('Bảng phân công nhân sự (không hạng mục)', 'Table/Grid', 'Enable', '–', '–', 'Trống',
     'Hiện khi bỏ tích “Dự án có hạng mục”: Phòng ban, Nhân sự*, Vai trò dự án*, Mô tả công việc, Ngày bắt đầu* '
     '(mặc định hôm nay), nút Xoá; nút “Thêm nhân sự”.'),
    ('Nút Thêm hạng mục', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Thêm 1 khối hạng mục (tab Quản lý hạng mục).'),
    ('Mã hạng mục', 'Textbox', 'Read-only', '–', 'Có', 'Sinh tự động', '“<Mã GP>_HMnn”.'),
    ('Tên hạng mục', 'Dropdown', 'Enable', 'Danh mục Hạng mục dự án', 'Có', 'Trống',
     'Nút “Tạo nhanh” mở cửa sổ Thêm hạng mục dự án (Hạng mục dự án*, Mô tả; Lưu → “Lưu thành công”).'),
    ('Leader hạng mục', 'Dropdown', 'Enable', 'Danh sách nhân viên', 'Có', 'Trống', '–'),
    ('Ngày cần hoàn thành (hạng mục)', 'Datepicker', 'Enable', 'dd/mm/yyyy', 'Có', 'Trống',
     '≥ hôm nay khi giải pháp Nháp.'),
    ('Ghi chú (hạng mục)', 'Textarea', 'Enable', '–', 'Không', 'Trống', '–'),
    ('Bảng phân công nhân sự (hạng mục)', 'Table/Grid', 'Enable', '–', '–', 'Trống',
     'Như bảng không hạng mục; danh sách Nhân sự loại Leader và người đã chọn.'),
    ('Nút Thu gọn / Mở chi tiết, Xoá hạng mục', 'Icon Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Tab Sơ đồ nhân sự', 'Table/Grid', 'Read-only', '–', '–', 'Theo dữ liệu',
     'Sơ đồ cây PM → Hạng mục → Leader → Nhân sự và bảng “Sơ đồ nhân sự dạng bảng”.'),
    ('Nút Lưu nháp / Lưu và gửi / Quay lại', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Ở chân trang.'),
]
fn('Tạo mới giải pháp – luồng có Yêu cầu làm giải pháp',
   dict(ten='Tạo mới giải pháp từ Yêu cầu làm giải pháp',
        mota='Người tiếp nhận YCGP (Trưởng phòng giải pháp) lập giải pháp cho dự án của yêu cầu, khai hạng mục và '
             'phân công nhân sự, rồi lưu nháp hoặc gửi cho PM.',
        tacnhan=A_TP + '; ' + A_ALL,
        dieukien='Người dùng là người tiếp nhận ít nhất 1 YCGP ở trạng thái Đã tiếp nhận; dự án chưa có giải pháp.',
        chinh='1. Người dùng bấm “Tạo mới” ở danh sách (hoặc biểu tượng “Làm giải pháp” ở dòng YCGP).\n'
              '2. Chọn Yêu cầu làm GP (mở từ YCGP thì đã chọn sẵn); hệ thống điền thông tin dự án, khách hàng, PM, '
              'Ngày cần xong GP, Nhóm ngành, Nhóm giải pháp từ YCGP và dự án.\n'
              '3. Người dùng nhập Tên GP, chỉnh các trường còn lại.\n'
              '4. Có hạng mục: sang tab Quản lý hạng mục, bấm “Thêm hạng mục”, chọn tên, Leader, ngày cần hoàn thành, '
              'phân công nhân sự. Không hạng mục: bỏ tích và phân công nhân sự ngay trên tab Thông tin.\n'
              '5. Bấm “Lưu nháp” (giải pháp Nháp) hoặc “Lưu và gửi” (giải pháp Chờ PM duyệt).\n'
              '6. Hệ thống lưu, báo “Đã lưu thành công!” và quay về danh sách.',
        phu='• Bỏ chọn YCGP → xóa toàn bộ thông tin đã điền từ yêu cầu.\n'
            '• Lỗi nhập liệu → viền đỏ + câu lỗi tại ô, tự chuyển tới tab có lỗi, toast “Vui lòng kiểm tra lại thông '
            'tin” (hoặc câu lỗi của trường không hiện trên form).\n'
            '• Lỗi khác → “Đã xảy ra lỗi. Vui lòng thử lại.”.',
        dacbiet='Khi lưu: YCGP chuyển sang Đang thực hiện; tạo version V1; trạng thái dự án đồng bộ theo giải pháp; '
                '“Lưu và gửi” gửi thông báo cho PM.'),
   [('menu', MENU + ' => Tạo mới'), ('menu', MENU_YC),
    ('p', 'Lối vào thứ hai mở form với YCGP đã chọn sẵn (biểu tượng chỉ hiện ở YCGP Đã tiếp nhận mà người dùng là '
          'người tiếp nhận).'),
    ('img', '06-tao-moi.png', 'Form Tạo giải pháp – tab Thông tin, đã chọn YCGP'),
    ('img', '06b-tao-moi-hang-muc.png', 'Tab Quản lý hạng mục'),
    ('img', '06c-tao-moi-so-do.png', 'Tab Sơ đồ nhân sự'),
    ('img', '06e-tao-moi-loi.png', 'Lỗi bắt buộc tại ô Ngày cần hoàn thành của hạng mục khi lưu')],
   UI_FORM,
   [('Chọn Yêu cầu làm GP', 'Change',
     'After:\n– Điền thông tin dự án, khách hàng, YCGP, Phòng làm GP, PM, Ngày cần xong GP, Nhóm ngành, Nhóm giải '
     'pháp, Ứng dụng.\n– Lỗi nạp → “Không thể tải thông tin yêu cầu giải pháp” / “Không thể tải thông tin dự án”.'),
    ('Bỏ / tích Dự án có hạng mục', 'Change',
     'After:\n– Bỏ tích: xóa toàn bộ hạng mục, ẩn tab Quản lý hạng mục, hiện bảng phân công trên tab Thông tin.\n'
     '– Tích: xóa bảng phân công trực tiếp.'),
    ('Chọn Phòng ban / Nhân sự ở bảng phân công', 'Change',
     'After:\n– Chọn phòng → danh sách nhân sự theo phòng; đổi phòng khác phòng của người đã chọn → xóa nhân sự.\n'
     '– Chọn nhân sự → tự điền phòng của người đó.'),
    ('Bấm Lưu nháp / Lưu và gửi', 'Click',
     'Before:\n– Không kiểm tra quyền hệ thống (điều kiện ở BR-06).\n'
     'During:\n– Yêu cầu làm GP, PM làm giải pháp, Tên GP, Ngày cần xong GP, Nhóm ngành, Nhóm giải pháp trống → '
     '“Bắt buộc phải nhập”.\n– Tên GP trùng → “Đã tồn tại trên hệ thống”; quá 255 ký tự → “Vui lòng nhập tối đa 255 ký '
     'tự.”.\n– Ngày cần xong GP / ngày hạng mục trước hôm nay → “Phải lớn hơn hoặc bằng ngày hiện tại”.\n'
     '– Hạng mục thiếu Tên / Leader / Ngày; nhân sự thiếu Nhân sự / Vai trò / Ngày bắt đầu → “Bắt buộc phải nhập”.\n'
     '– Dự án đã có giải pháp → “Dự án này đã có giải pháp. Không thể tạo giải pháp khác.”\n'
     '– Dự án cha → “Dự án cha không làm giải pháp kỹ thuật. Vui lòng tạo giải pháp trên dự án con.”\n'
     '– Có lỗi → không thực hiện bước After.\n'
     'After:\n– Lưu giải pháp (Nháp hoặc Chờ PM duyệt), hạng mục, nhân sự; sinh mã; tạo version V1.\n'
     '– YCGP → Đang thực hiện; đồng bộ trạng thái dự án.\n– Lưu và gửi: thông báo cho PM (BR-15).\n'
     '– “Đã lưu thành công!”, về danh sách.'),
    ('Bấm Quay lại', 'Click', 'After:\n– Quay lại màn trước.')],
   uc=('FR-06', 'Tạo mới giải pháp – luồng có YCGP', 'crud', A_TP),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.7 FR-07
fn('Tạo mới giải pháp – dự án Tự triển khai',
   dict(ten='Tạo mới giải pháp cho dự án Tự triển khai (sale tự làm)',
        mota='Nhân viên KD tự lập giải pháp cho dự án Tự triển khai của mình, không qua YCGP, không có hạng mục.',
        tacnhan=A_KD + '; ' + A_ALL,
        dieukien='Dự án Tự triển khai đang ở trạng thái Thu thập thông tin dự án, có tích “Có làm GP”, chưa có giải '
                 'pháp và phiếu thu thập thông tin đã nhập đủ các trường bắt buộc.',
        chinh='1. Ở danh sách Dự án TKT, người dùng bấm biểu tượng “Tạo giải pháp” ở dòng dự án.\n'
              '2. Hệ thống mở form Tạo giải pháp, điền thông tin dự án / khách hàng; PM = chính người dùng (khóa); '
              'không có ô Yêu cầu làm GP, không có ô “Dự án có hạng mục”.\n'
              '3. Người dùng nhập Tên GP, Ngày cần xong GP, Nhóm ngành, Nhóm giải pháp, Mô tả và bảng phân công nhân '
              'sự.\n4. Bấm “Lưu nháp” hoặc “Lưu và gửi”.\n'
              '5. Hệ thống lưu, “Đã lưu thành công!”, về danh sách giải pháp.',
        phu='• Lưu và gửi → giải pháp vào thẳng Đang triển khai (bỏ qua duyệt PM/Leader).\n'
            '• Dự án không đủ điều kiện → “Dự án không đủ điều kiện để tạo giải pháp.”.\n'
            '• Phiếu thu thập chưa đủ → “Phiếu thu thập thông tin chưa nhập đủ các trường bắt buộc. Vui lòng hoàn '
            'thiện phiếu trước khi tạo giải pháp.”'),
   [('menu', MENU_DA), ('img', 's_04_sol_add.png', 'Form Tạo giải pháp của dự án Tự triển khai')],
   [r for r in UI_FORM if not r[0].startswith(('Yêu cầu làm GP', 'Dự án có hạng mục', 'Nút Thêm hạng mục',
                                                  'Mã hạng mục', 'Tên hạng mục', 'Leader hạng mục',
                                                  'Ngày cần hoàn thành', 'Ghi chú (hạng mục)',
                                                  'Bảng phân công nhân sự (hạng mục)', 'Nút Thu gọn'))],
   [('Mở form từ Dự án TKT', 'System',
     'After:\n– Nạp dự án; đặt PM = người dùng, bỏ YCGP, “không có hạng mục”.'),
    ('Bấm Lưu nháp / Lưu và gửi', 'Click',
     'Before:\n– Kiểm tra dự án đủ điều kiện (BR-06); không đủ → báo lỗi tại ô Dự án và dừng.\n'
     'During:\n– Tên GP, Ngày cần xong GP, Nhóm ngành, Nhóm giải pháp trống → “Bắt buộc phải nhập”.\n'
     '– Tên GP trùng → “Đã tồn tại trên hệ thống”.\n– Nhân sự thiếu trường bắt buộc → “Bắt buộc phải nhập”.\n'
     'After:\n– Lưu Nháp hoặc Đang triển khai; tạo version V1; đồng bộ trạng thái dự án.\n'
     '– “Đã lưu thành công!”, về danh sách.')],
   uc=('FR-07', 'Tạo mới giải pháp – Tự triển khai', 'crud', A_KD),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.8 FR-08
fn('Xem chi tiết giải pháp',
   dict(ten='Xem chi tiết giải pháp',
        mota='Hiển thị toàn bộ thông tin giải pháp ở chế độ chỉ đọc, gồm 3 tab Thông tin / Quản lý hạng mục / Sơ đồ '
             'nhân sự, kèm các nút thao tác được phép.',
        tacnhan=A_ALL, dieukien='Giải pháp nằm trong phạm vi được xem; người dùng không bị khóa khỏi giải pháp.',
        chinh='1. Người dùng bấm Mã giải pháp ở danh sách.\n'
              '2. Hệ thống mở màn “Chi tiết giải pháp: <mã> (<trạng thái>)” với mọi ô ở chế độ chỉ đọc.\n'
              '3. Người dùng chuyển giữa các tab để xem.',
        phu='• Người dùng đã bị khóa khỏi giải pháp → “Bạn đã bị khóa khỏi giải pháp này nên không thao tác được. '
            'Vui lòng liên hệ PM hoặc Trưởng phòng giải pháp.” và quay về danh sách.\n'
            '• Lỗi nạp → “Không thể tải thông tin giải pháp”.\n'
            '• Leader hạng mục mở giải pháp Chờ Leader duyệt → tự chuyển sang tab Quản lý hạng mục, mở và cuộn tới '
            'hạng mục của mình.'),
   [('menu', MENU + ' => Mã giải pháp'), ('img', '09-chi-tiet.png', 'Màn Chi tiết giải pháp (giải pháp Nháp)')],
   [('Tiêu đề trang', 'Label', 'Hiển thị', '–', 'Chi tiết giải pháp: <mã> (<trạng thái>)', '–'),
    ('Các tab và mọi ô của form', 'Text', 'Read-only', '–', 'Theo dữ liệu', 'Như FR-06, tất cả ở chế độ chỉ đọc.'),
    ('Nút Sửa', 'Button', 'Enable / Ẩn', '–', 'Ẩn khi thiếu quyền', 'Hiện khi được sửa (BR-07).'),
    ('Nút Xóa', 'Button', 'Enable / Ẩn', '–', 'Ẩn khi thiếu quyền', 'Hiện khi được xóa (BR-07).'),
    ('Nút Giao cho Leader / Lưu và duyệt', 'Button', 'Enable / Ẩn', '–', 'Ẩn',
     'Cùng điều kiện với cột Hành động (FR-11, FR-12).'),
    ('Nút Quay lại', 'Button', 'Enable', '–', 'Hiển thị', 'Về màn trước.')],
   [('Mở màn chi tiết', 'System',
     'Before:\n– Thành viên đã bị khóa → báo lý do, về danh sách.\nAfter:\n– Hiển thị dữ liệu, chọn tab theo vai trò.'),
    ('Bấm Sửa', 'Click', 'After:\n– Mở màn Sửa giải pháp (FR-09).'),
    ('Bấm Xóa', 'Click', 'After:\n– Mở hộp xác nhận xóa (FR-10); xóa xong về danh sách.'),
    ('Bấm Giao cho Leader / Lưu và duyệt', 'Click', 'After:\n– Mở màn Duyệt giải pháp.')],
   ui_kind='read', rule=('- Màn Xem chi tiết và Phân quyền.', 'detail'))

# ---------------------------------------------------------------- 2.9 FR-09
fn('Chỉnh sửa giải pháp',
   dict(ten='Chỉnh sửa giải pháp',
        mota='Cập nhật giải pháp; nút ở chân trang thay đổi theo trạng thái giải pháp và vai trò người dùng.',
        tacnhan=A_TP + '; ' + A_PM + '; ' + A_LD,
        dieukien='Người dùng được sửa giải pháp theo BR-07.',
        chinh='1. Người dùng bấm “Sửa” ở danh sách hoặc ở màn chi tiết.\n'
              '2. Hệ thống mở màn “Sửa giải pháp” với dữ liệu hiện tại.\n3. Người dùng chỉnh thông tin.\n'
              '4. Bấm nút lưu tương ứng trạng thái:\n'
              '   – Nháp: “Lưu nháp” (giữ Nháp) hoặc “Lưu và gửi” (→ Chờ PM duyệt).\n'
              '   – Chờ PM duyệt (PM): “Lưu” (giữ trạng thái).\n'
              '   – Chờ Leader duyệt (PM): “Lưu hạng mục” — thêm hạng mục mới, Leader vẫn phải duyệt.\n'
              '   – Chờ Leader duyệt (Leader có hạng mục chưa duyệt): “Lưu”.\n'
              '   – Đang triển khai (PM, có hạng mục): “Lưu và duyệt hạng mục” — hạng mục mới thêm được duyệt luôn; '
              'nút “Tạo/Sửa hồ sơ trình duyệt giải pháp” mở cửa sổ hồ sơ (FR-20).\n'
              '5. Hệ thống lưu, báo thành công và quay về danh sách.',
        phu='• Không còn quyền sửa khi mở màn → “Bạn không có quyền!” và về danh sách.\n'
            '• Từ trạng thái Chờ PM duyệt trở đi (trừ Đóng): khóa các ô Thông tin GP (PM, Tên, Ngày, Nhóm…) và bảng '
            'phân công trực tiếp; hạng mục đã lưu bị khóa khi giải pháp Đang triển khai (chỉ sửa hạng mục mới thêm).\n'
            '• Leader ở bước Chờ Leader duyệt: tab Thông tin chỉ đọc; chỉ sửa hạng mục mình làm Leader (không đổi tên '
            'và Leader của hạng mục đó).\n'
            '• Dữ liệu đã bị người/tab khác đổi trạng thái → “Dữ liệu đã được cập nhật bởi người khác hoặc tab khác. '
            'Vui lòng tải lại trang!” rồi tự tải lại sau 2 giây.',
        dacbiet='Mã GP chốt lúc tạo, không đổi khi sửa. Đổi Ngày cần xong GP thì cập nhật ngày kết thúc version '
                'hiện tại.'),
   [('menu', MENU + ' => Sửa'), ('img', '08-sua.png', 'Màn Sửa giải pháp (giải pháp Nháp)')],
   UI_FORM + [
       ('Nút Lưu / Lưu hạng mục / Lưu và duyệt hạng mục', 'Button', 'Enable / Ẩn', '–', '–', 'Theo trạng thái',
        'Xem dòng sự kiện chính.'),
       ('Nút Tạo / Sửa hồ sơ trình duyệt giải pháp', 'Button', 'Enable / Ẩn', '–', '–', 'Ẩn',
        'Chỉ khi Đang triển khai; chữ “Sửa…” khi hồ sơ mới nhất đang Nháp.')],
   [('Mở màn Sửa', 'System', 'Before:\n– Không được sửa → “Bạn không có quyền!”, về danh sách.'),
    ('Bấm nút lưu', 'Click',
     'Before:\n– Kiểm tra quyền sửa (BR-07); không có → “Bạn không có quyền chỉnh sửa giải pháp này!” rồi về danh '
     'sách sau 1,5 giây.\n– Trạng thái đã đổi → thông báo dữ liệu đã cập nhật, tải lại trang.\n'
     'During:\n– Như FR-06 (bắt buộc, trùng tên, ngày ≥ hôm nay chỉ áp khi giải pháp đang Nháp).\n'
     'After:\n– Lưu giải pháp, hạng mục, nhân sự; cập nhật trạng thái version.\n'
     '– Lưu và gửi từ Nháp → Chờ PM duyệt, thông báo PM.\n'
     '– Thông báo “Đã cập nhật thành công!” (hoặc “Đã lưu hạng mục thành công!”, “Đã lưu và duyệt thành công!”), '
     'về danh sách.\n– Lỗi khác → “Đã xảy ra lỗi. Vui lòng thử lại.”')],
   uc=('FR-09', 'Chỉnh sửa giải pháp', 'crud', A_PM),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.10 FR-10
fn('Xóa giải pháp',
   dict(ten='Xóa giải pháp',
        mota='Xóa hẳn giải pháp Nháp do chính mình tạo.',
        tacnhan=A_TP, dieukien='Giải pháp ở trạng thái Nháp và người dùng là người tạo.',
        chinh='1. Người dùng bấm biểu tượng Xóa ở dòng giải pháp (hoặc nút Xóa ở màn chi tiết).\n'
              '2. Hệ thống hiện hộp “Xác nhận xóa”: “Bạn có chắc muốn xóa giải pháp ‘<tên>’?”.\n'
              '3. Người dùng bấm “Xóa”.\n4. Hệ thống xóa, báo “Xóa thành công” và nạp lại danh sách.',
        phu='• Bấm “Hủy” → đóng hộp, không xóa.\n'
            '• Không còn quyền xóa → “Bạn không có quyền xóa giải pháp này!”.'),
   [('menu', MENU + ' => Xóa'), ('modal', 'Xác nhận xóa'), ('img', '10-xoa.png', 'Hộp Xác nhận xóa giải pháp')],
   [('Tiêu đề', 'Label', 'Hiển thị', 'Xác nhận xóa', '–'),
    ('Nội dung', 'Label', 'Hiển thị', 'Theo dữ liệu', '“Bạn có chắc muốn xóa giải pháp ‘<tên giải pháp>’?”'),
    ('Nút Xóa', 'Button', 'Enable', 'Hiển thị', 'Màu đỏ.'),
    ('Nút Hủy', 'Button', 'Enable', 'Hiển thị', '–')],
   [('Bấm Xóa (xác nhận)', 'Click',
     'Before:\n– Kiểm tra giải pháp Nháp và người dùng là người tạo; không thỏa → “Bạn không có quyền xóa giải pháp '
     'này!” và dừng.\nAfter:\n– Xóa giải pháp và nhân sự trực tiếp; YCGP gắn kèm quay về Đã tiếp nhận.\n'
     '– “Xóa thành công”, nạp lại danh sách (từ màn chi tiết thì về danh sách).\n– Lỗi → “Lỗi khi xóa”.'),
    ('Bấm Hủy', 'Click', 'After:\n– Đóng hộp xác nhận.')],
   uc=('FR-10', 'Xóa giải pháp', 'action', A_TP), ui_kind='confirm',
   rule=('- Thông báo, Quy tắc Xóa.', 'delete'))

# ---------------------------------------------------------------- 2.11 FR-11
fn('PM duyệt tiếp nhận giải pháp',
   dict(ten='PM duyệt tiếp nhận (Giao cho Leader / Lưu và duyệt)',
        mota='PM rà giải pháp được giao, bổ sung hạng mục / nhân sự rồi giao cho các Leader hạng mục duyệt, hoặc duyệt '
             'thẳng khi giải pháp không có hạng mục.',
        tacnhan=A_PM, dieukien='Giải pháp ở trạng thái Chờ PM duyệt và người dùng là PM.',
        chinh='1. PM bấm biểu tượng “Giao cho Leader” (có hạng mục) hoặc “Lưu và duyệt” (không hạng mục) ở dòng giải '
              'pháp hoặc ở màn chi tiết.\n2. Hệ thống mở màn “Duyệt giải pháp” (form như FR-09).\n'
              '3. PM rà soát, thêm/sửa hạng mục, Leader, nhân sự.\n4. Bấm “Giao cho Leader” / “Lưu và duyệt”.\n'
              '5. Hệ thống lưu, “Đã cập nhật thành công!”, về danh sách.',
        phu='• Giao cho Leader khi chưa có hạng mục nào → “Vui lòng thêm ít nhất 1 hạng mục trước khi giao cho Leader”.\n'
            '• Giao cho Leader → giải pháp Chờ Leader duyệt, thông báo các Leader.\n'
            '• Lưu và duyệt (không hạng mục) → giải pháp Đang triển khai.'),
   [('menu', MENU + ' => Giao cho Leader / Lưu và duyệt'),
    ('img', 'h_05_pm_duyet.png', 'Màn Duyệt giải pháp của PM (giải pháp có hạng mục)')],
   [('Form giải pháp', 'Text', 'Enable / Disable', '–', '–', 'Theo dữ liệu',
     'Như FR-09; Thông tin GP bị khóa, hạng mục và nhân sự sửa được.'),
    ('Nút Giao cho Leader / Lưu và duyệt', 'Button', 'Enable', '–', '–', 'Hiển thị',
     '“Giao cho Leader” khi giải pháp có hạng mục, ngược lại “Lưu và duyệt”.'),
    ('Nút Quay lại', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Giao cho Leader', 'Click',
     'Before:\n– Người dùng là PM, giải pháp Chờ PM duyệt.' + '\n' + NO_PERM + '\n'
     'During:\n– Chưa có hạng mục → “Vui lòng thêm ít nhất 1 hạng mục trước khi giao cho Leader”.\n'
     '– Lỗi nhập liệu như FR-06.\n'
     'After:\n– Giải pháp → Chờ Leader duyệt; thông báo cho từng Leader (BR-15).\n– “Đã cập nhật thành công!”.'),
    ('Bấm Lưu và duyệt (không hạng mục)', 'Click',
     'After:\n– Giải pháp → Đang triển khai; ghi ngày bắt đầu version; đồng bộ trạng thái dự án.\n'
     '– “Đã cập nhật thành công!”.')],
   uc=('FR-11', 'PM duyệt tiếp nhận giải pháp', 'action', A_PM),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.12 FR-12
fn('Leader duyệt hạng mục',
   dict(ten='Leader duyệt hạng mục',
        mota='Leader rà hạng mục mình phụ trách (ngày, ghi chú, nhân sự) rồi duyệt; khi mọi hạng mục đã duyệt, giải '
             'pháp tự sang Đang triển khai.',
        tacnhan=A_LD, dieukien='Giải pháp Chờ Leader duyệt và người dùng là Leader của ít nhất 1 hạng mục chưa duyệt.',
        chinh='1. Leader bấm biểu tượng “Lưu và duyệt” ở dòng giải pháp (hoặc ở màn chi tiết).\n'
              '2. Hệ thống mở màn Duyệt giải pháp ở tab Quản lý hạng mục, mở sẵn hạng mục của Leader.\n'
              '3. Leader chỉnh ngày cần hoàn thành, ghi chú, phân công nhân sự của hạng mục mình.\n'
              '4. Bấm “Lưu và duyệt”.\n5. Hệ thống đánh dấu các hạng mục của Leader là Đã duyệt, “Đã lưu và duyệt '
              'thành công!”, về danh sách.',
        phu='• Hạng mục cuối cùng được duyệt → giải pháp tự chuyển Đang triển khai.\n'
            '• Leader chỉ muốn lưu (chưa duyệt) → vào Sửa, bấm “Lưu”.\n'
            '• Hạng mục của Leader khác và tab Thông tin ở chế độ chỉ đọc.'),
   [('menu', MENU + ' => Lưu và duyệt'), ('img', 'h_x_leader_duyet.png', 'Màn Duyệt giải pháp của Leader hạng mục')],
   [('Khối hạng mục của Leader', 'Table/Grid', 'Enable', '–', '–', 'Mở sẵn, viền nổi bật',
     'Sửa được Ngày cần hoàn thành, Ghi chú, bảng nhân sự; Tên hạng mục và Leader bị khóa.'),
    ('Khối hạng mục khác', 'Table/Grid', 'Read-only', '–', '–', 'Thu gọn', '–'),
    ('Nút Lưu và duyệt', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Nút Quay lại', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Lưu và duyệt', 'Click',
     'Before:\n– Người dùng là Leader của hạng mục chưa duyệt.\n' + NO_PERM + '\n'
     'During:\n– Lỗi nhập liệu như FR-06 (ngày, nhân sự bắt buộc).\n'
     'After:\n– Đánh dấu các hạng mục của Leader Đã duyệt (ghi ngày duyệt).\n'
     '– Mọi hạng mục đã duyệt → giải pháp Đang triển khai, ghi ngày bắt đầu version, đồng bộ trạng thái dự án.\n'
     '– “Đã lưu và duyệt thành công!”, về danh sách.')],
   uc=('FR-12', 'Leader duyệt hạng mục', 'action', A_LD),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.13 FR-13
fn('Mở màn Quản lý giải pháp',
   dict(ten='Mở màn Quản lý giải pháp và xem Tổng quan',
        mota='Màn toàn trang theo dõi 1 giải pháp: thanh đầu trang (tên, mã, version, tiến độ, nút thao tác) và 10 tab '
             'Tổng quan · Thông tin · Hồ sơ · Nhân sự · Nhiệm vụ · Vấn đề giải pháp · Meeting · Tiến độ · Files · YC Điều '
             'chỉnh. Tab Tổng quan hiển thị cảnh báo, hạng mục, nhiệm vụ cần duyệt và biểu đồ.',
        tacnhan=A_TP + '; ' + A_PM + '; ' + A_ALL,
        dieukien='Giải pháp không ở Nháp / Đóng và người dùng đủ điều kiện quản lý (BR-08).',
        chinh='1. Người dùng bấm biểu tượng “Quản lý giải pháp” ở dòng giải pháp.\n'
              '2. Hệ thống mở màn Quản lý ở tab Tổng quan.\n'
              '3. Người dùng xem 5 thẻ cảnh báo, bảng “Hạng mục trong dự án”, khối “Nhiệm vụ cần duyệt”, biểu đồ '
              '“Phân bổ trạng thái nhiệm vụ” và “Mức độ ưu tiên của Issue”.\n'
              '4. Bấm 1 thẻ cảnh báo để mở danh sách chi tiết; bấm tên hạng mục để mở màn quản lý hạng mục; bấm 1 '
              'nhiệm vụ cần duyệt để mở cửa sổ nhiệm vụ.',
        phu='• Người dùng bị khóa khỏi giải pháp → báo lý do, về danh sách giải pháp.\n'
            '• Biểu đồ không có số liệu → “Không có dữ liệu”; không có nhiệm vụ cần duyệt → “Không có nhiệm vụ cần '
            'duyệt”.\n• Bấm “Quay lại” → về màn trước.'),
   [('menu', MQL), ('img', '11-quan-ly-tong-quan.png', 'Màn Quản lý giải pháp – tab tổng quan (tab đầu tiên)'),
    ('img', '12-canh-bao.png', 'Cửa sổ chi tiết của thẻ cảnh báo “Nhiệm vụ sắp tới hạn”'),
    ('img', 's_06_sol_mgr.png', 'Tab tổng quan của giải pháp dự án Tự triển khai (không có hạng mục)')],
   [('Tên giải pháp, Mã, Version, Tiến độ', 'Label', 'Hiển thị', '0 – 100', 'Theo dữ liệu',
     'Màu nhãn tiến độ: ≥80% xanh lá, ≥50% xanh dương, ≥20% cam, còn lại đỏ.'),
    ('Nút Tạo / Sửa hồ sơ trình duyệt giải pháp', 'Button', 'Enable / Ẩn', '–', 'Ẩn', 'FR-20.'),
    ('Nút Tạo version', 'Button', 'Enable / Ẩn', '–', 'Ẩn', 'FR-23.'),
    ('Nút Quay lại', 'Button', 'Enable', '–', 'Hiển thị', '–'),
    ('Thanh 10 tab', 'Tab', 'Enable', '–', 'Tổng quan', '–'),
    ('5 thẻ cảnh báo', 'Badge', 'Enable', '≥ 0', 'Theo dữ liệu',
     'Nhiệm vụ sắp tới hạn (trong n ngày) · Hạng mục nhiều nhiệm vụ trễ (từ n nhiệm vụ trễ) · Nhân sự nhiều nhiệm vụ '
     'trễ · Vấn đề sắp tới hạn · Meeting sắp tới lịch; n lấy từ cấu hình hạn.'),
    ('Cửa sổ chi tiết cảnh báo', 'Modal', 'Read-only', '–', 'Ẩn',
     'Bộ lọc + bảng (vd Mã-Tên nhiệm vụ, Hạng mục, Trạng thái, Ưu tiên, Người làm, Người duyệt, Hạn hoàn thành); '
     'nút “Xem chi tiết” ở dòng; nút Đóng.'),
    ('Bảng Hạng mục trong dự án', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'STT, Hạng mục (link), Trạng thái, Version, Lead, Thời gian, Nhân sự, Tiến độ.'),
    ('Khối Nhiệm vụ cần duyệt', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'Nhiệm vụ người dùng đang được duyệt: tên, người làm, Deadline.'),
    ('Biểu đồ Phân bổ trạng thái nhiệm vụ / Mức độ ưu tiên của Issue', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'Biểu đồ tròn / cột.')],
   [('Mở màn Quản lý', 'System',
     'Before:\n– Thành viên bị khóa → “Bạn đã bị khóa khỏi giải pháp này nên không thao tác được. Vui lòng liên hệ PM '
     'hoặc Trưởng phòng giải pháp.”, về danh sách.\nAfter:\n– Nạp giải pháp, hồ sơ, hạng mục, số đếm cảnh báo, thống '
     'kê nhiệm vụ, vấn đề.'),
    ('Bấm thẻ cảnh báo', 'Click', 'After:\n– Mở cửa sổ danh sách tương ứng; “Xem chi tiết” mở nhiệm vụ / vấn đề / '
     'meeting / hạng mục.'),
    ('Bấm nhiệm vụ cần duyệt', 'Click', 'After:\n– Mở cửa sổ nhiệm vụ (chi tiết theo SRS Nhiệm vụ).'),
    ('Đổi tab', 'Click', 'After:\n– Hiển thị nội dung tab (FR-14 → FR-29).')],
   ui_kind='read')

# ---------------------------------------------------------------- 2.14 FR-14
fn('Xem tab Thông tin',
   dict(ten='Xem thông tin dự án và yêu cầu làm giải pháp',
        mota='Hiển thị chỉ đọc 4 khối: Thông tin khách hàng, Yêu cầu làm giải pháp, Dự án, Nguồn vốn & kỳ vọng tài '
             'chính.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn Quản lý giải pháp.',
        chinh='1. Người dùng bấm tab “Thông tin”.\n2. Hệ thống hiển thị 4 khối thông tin.',
        phu='• Trường không có dữ liệu hiển thị “—”.\n• Khối Khách hàng cuối chỉ hiện khi KH trực tiếp là KH trung gian.'),
   [('menu', MQL + ' => Thông tin'), ('img', '13-tab-thong-tin.png', 'Tab Thông tin')],
   [('Thông tin khách hàng', 'Text', 'Read-only', '–', 'Theo dự án',
     'Khách hàng (link), Người liên hệ, Chức vụ, Email, SĐT; khối KH cuối tương tự.'),
    ('Yêu cầu làm giải pháp', 'Text', 'Read-only', '–', 'Theo YCGP',
     'Mã YCLGP (link), Nội dung, Người tạo YC, Ngày tạo, Trạng thái (badge).'),
    ('Dự án', 'Text', 'Read-only', '–', 'Theo dự án', 'Tên (link), PM chính, Deadline, Tiến trình nội bộ (badge).'),
    ('Nguồn vốn & kỳ vọng tài chính', 'Text', 'Read-only', '≥ 0', 'Theo dự án',
     'Nguồn vốn, Ngân sách dự kiến, Giá trị hợp đồng dự kiến, Lợi nhuận kỳ vọng.')],
   [('Bấm tab Thông tin', 'Click', 'After:\n– Hiển thị dữ liệu đã nạp cùng giải pháp.'),
    ('Bấm link khách hàng / YCGP / dự án', 'Click', 'After:\n– Mở màn tương ứng.')],
   ui_kind='read', rule=('- Màn Xem chi tiết và Phân quyền.', 'detail'))

# ---------------------------------------------------------------- 2.15 FR-15
fn('Phân công PM / Leader',
   dict(ten='Phân công quản lý (đổi PM / đổi Leader hạng mục)',
        mota='Thay PM của giải pháp hoặc Leader của 1 hạng mục, ghi lý do và giao vai trò mới cho người cũ.',
        tacnhan=A_TP + '; ' + A_PM,
        dieukien='Đổi PM: người dùng là người tạo giải pháp. Đổi Leader: người tạo hoặc PM. Giải pháp không ở Đóng.',
        chinh='1. Ở tab Nhân sự, bấm “Phân công”.\n2. Hệ thống mở cửa sổ “Phân công quản lý”.\n'
              '3. Chọn Loại phân công (Phân công PM / Phân công Leader); với Leader chọn Hạng mục.\n'
              '4. Hệ thống hiện PM/Leader hiện tại; chọn Vai trò mới cho người cũ (bỏ trống = rời dự án), Hạng mục cho '
              'PM cũ (nếu cần), PM/Leader mới và nhập Ghi chú ≥ 50 ký tự.\n'
              '5. Bấm “Xác nhận phân công”, xác nhận ở hộp “Xác nhận phân công”.\n'
              '6. Hệ thống đổi người, ghi lịch sử, thông báo người cũ và người mới, nạp lại tab Nhân sự.',
        phu='• Chỉ được 1 loại phân công → loại đó được chọn sẵn, loại kia bị khóa.\n'
            '• Lỗi máy chủ → câu lỗi hiện trong khung đỏ trong cửa sổ.'),
   [('menu', MQL + ' => Nhân sự => Phân công'), ('modal', 'Phân công quản lý'),
    ('img', 'h_04_doi_pm.png', 'Cửa sổ Phân công quản lý – đổi PM')],
   [('Loại phân công', 'Radio', 'Enable / Disable', 'Phân công PM / Phân công Leader', 'Có', 'Theo quyền',
     '“Vui lòng chọn loại phân công.”'),
    ('Hạng mục', 'Dropdown', 'Enable / Ẩn', 'Danh sách hạng mục', 'Có khi Phân công Leader', 'Trống',
     '“Vui lòng chọn hạng mục.”'),
    ('PM / Leader hiện tại', 'Textbox', 'Read-only', '–', '–', 'Theo dữ liệu', '–'),
    ('Vai trò mới cho PM / Leader cũ', 'Dropdown', 'Enable', 'Danh mục vai trò dự án', 'Không', 'Trống',
     'Gợi ý “Không chọn = rời dự án”.'),
    ('Hạng mục cho PM cũ', 'Dropdown', 'Enable / Ẩn', 'Danh sách hạng mục', 'Có khi đổi PM, GP có hạng mục và có '
     'chọn vai trò mới', 'Trống', '“Vui lòng chọn hạng mục cho PM cũ.”'),
    ('PM / Leader mới', 'Dropdown', 'Enable', 'Nhân viên cùng Phòng làm GP, khác người hiện tại', 'Có', 'Trống',
     '“Vui lòng chọn người mới.”'),
    ('Ghi chú', 'Textarea', 'Enable', '≥ 50 ký tự', 'Có', 'Trống',
     'Bộ đếm “n/50 ký tự tối thiểu”; “Vui lòng nhập ghi chú.” / “Ghi chú phải có ít nhất 50 ký tự (hiện n).”'),
    ('Nút Xác nhận phân công / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Phân công', 'Click', 'After:\n– Mở cửa sổ, chọn sẵn loại phân công nếu chỉ có 1 loại.'),
    ('Bấm Xác nhận phân công', 'Click',
     'Before:\n– Đổi PM mà không phải người tạo → “Bạn không có quyền phân công PM.”; đổi Leader mà không phải người '
     'tạo/PM → “Bạn không có quyền phân công Leader.”; giải pháp Đóng → “Giải pháp đã đóng, không thể phân công.”\n'
     'During:\n– Kiểm tra các ô bắt buộc và ghi chú ≥ 50 ký tự (câu lỗi như cột Mô tả).\n'
     '– Hỏi xác nhận “Xác nhận phân công PM/Leader mới: <tên>? Hệ thống sẽ gửi thông báo cho … cũ và … mới.”\n'
     'After:\n– Đổi PM của giải pháp / Leader của hạng mục; người cũ có vai trò mới thì thêm vào nhân sự.\n'
     '– Ghi lịch sử phân công; gửi thông báo cho người cũ và người mới (BR-15); đóng cửa sổ, nạp lại tab.')],
   uc=('FR-15', 'Phân công PM / Leader', 'action', A_TP),
   rule=('- Thông báo và UI/UX.', 'notice'))

# ---------------------------------------------------------------- 2.16 FR-16
fn('Xem lịch sử phân công',
   dict(ten='Xem lịch sử phân công quản lý',
        mota='Liệt kê các lần đổi PM / Leader của giải pháp, mới nhất trước.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở tab Nhân sự.',
        chinh='1. Bấm “Xem lịch sử phân công”.\n2. Hệ thống mở cửa sổ “Lịch sử phân công quản lý” với bảng lịch sử.',
        phu='• Chưa có lần phân công nào → “Chưa có lịch sử phân công.”.'),
   [('menu', MQL + ' => Nhân sự => Xem lịch sử phân công'), ('modal', 'Lịch sử phân công quản lý'),
    ('img', '18-lich-su-phan-cong.png', 'Cửa sổ Lịch sử phân công quản lý (giải pháp chưa đổi PM/Leader)')],
   [('Bảng lịch sử', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'Ngày, Loại (badge PM/Leader), Hạng mục, Người cũ, Người mới, Vai trò mới (người cũ), Ghi chú, Người thực hiện.'),
    ('Trạng thái rỗng', 'Label', 'Hiển thị', '–', 'Ẩn', '“Chưa có lịch sử phân công.”'),
    ('Nút Đóng', 'Button', 'Enable', '–', 'Hiển thị', '–')],
   [('Bấm Xem lịch sử phân công', 'Click', 'After:\n– Nạp và hiển thị lịch sử, mới nhất trước.'),
    ('Bấm Đóng', 'Click', 'After:\n– Đóng cửa sổ.')],
   ui_kind='read', rule=('- Quy tắc ghi lịch sử.', 'history'))

# ---------------------------------------------------------------- 2.17 FR-17
fn('Thêm nhân sự',
   dict(ten='Thêm nhân sự vào giải pháp / hạng mục',
        mota='Bổ sung thành viên vào 1 hạng mục (giải pháp có hạng mục) hoặc vào thẳng giải pháp.',
        tacnhan=A_PM + '; ' + A_TP, dieukien='Người dùng đang ở tab Nhân sự của màn Quản lý.',
        chinh='1. Bấm “Thêm nhân sự”.\n2. Hệ thống mở cửa sổ “Thêm nhân sự vào hạng mục” (hoặc “… vào dự án”).\n'
              '3. Chọn Hạng mục (nếu có), Thành viên, Vai trò dự án, nhập Mô tả công việc, Ngày bắt đầu.\n'
              '4. Bấm “Thêm nhân sự”.\n5. Hệ thống thêm thành viên, nạp lại bảng và sơ đồ nhân sự.',
        phu='• Thành viên đã có trong hạng mục → “Nhân sự đã có trong hạng mục này.”; đã có trong dự án → “Nhân sự đã '
            'có trong dự án.”.\n• Bấm “Huỷ” → đóng cửa sổ.'),
   [('menu', MQL + ' => Nhân sự => Thêm nhân sự'), ('modal', 'Thêm nhân sự'),
    ('img', 'h_03_them_nhan_su.png', 'Cửa sổ Thêm nhân sự vào hạng mục')],
   [('Hạng mục', 'Dropdown', 'Enable / Ẩn', 'Danh sách hạng mục', 'Có khi giải pháp có hạng mục', 'Trống',
     '“Vui lòng chọn hạng mục”.'),
    ('Thành viên', 'Dropdown', 'Enable / Disable', 'Nhân viên chưa có trong hạng mục / giải pháp', 'Có', 'Trống',
     'Khóa với gợi ý “Vui lòng chọn hạng mục trước” khi chưa chọn hạng mục.'),
    ('Vai trò dự án', 'Dropdown', 'Enable', 'Danh mục vai trò dự án', 'Có', 'Trống', '–'),
    ('Mô tả công việc', 'Textarea', 'Enable', '–', 'Không', 'Trống', '–'),
    ('Ngày bắt đầu', 'Datepicker', 'Enable', 'dd/mm/yyyy, ≥ hôm nay', 'Có', 'Trống', '–'),
    ('Nút Thêm nhân sự / Huỷ', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Thêm nhân sự (trong cửa sổ)', 'Click',
     'Before:\n– Không kiểm tra quyền riêng (xem mục điểm lưu ý cuối tài liệu).\n'
     'During:\n– Thành viên / Vai trò / Ngày bắt đầu trống → “Bắt buộc phải nhập”; thiếu hạng mục → “Vui lòng chọn '
     'hạng mục”; Ngày bắt đầu trước hôm nay → “Phải lớn hơn hoặc bằng ngày hiện tại”.\n'
     '– Hạng mục không thuộc giải pháp → “Hạng mục không tồn tại hoặc không thuộc giải pháp này.”; trùng → câu ở '
     'dòng sự kiện phụ.\n'
     'After:\n– Thêm thành viên, đóng cửa sổ, nạp lại bảng nhân sự.')],
   uc=('FR-17', 'Thêm nhân sự', 'crud', A_PM),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.18 FR-18
fn('Sửa phân công / Khóa / Mở khóa / Xóa thành viên',
   dict(ten='Quản lý thành viên giải pháp',
        mota='Sửa phân công, khóa / mở khóa hoặc xóa 1 thành viên khỏi giải pháp; dòng PM và Leader không có nút '
             '(đổi bằng FR-15).',
        tacnhan=A_TP + '; ' + A_PM + '; ' + A_LD,
        dieukien='Giải pháp không Đóng; người dùng là người tạo / PM (mọi thành viên) hoặc Leader (thành viên hạng mục '
                 'của mình).',
        chinh='1. Ở bảng tab Nhân sự, bấm “Sửa phân công”, “Khóa thành viên” / “Mở khóa thành viên” hoặc “Xóa khỏi giải '
              'pháp” ở cột Thao tác.\n'
              '2. Sửa: cửa sổ “Sửa phân công thành viên” (Hạng mục, Vai trò dự án, Mô tả công việc, Ngày bắt đầu, Ngày '
              'kết thúc) → “Lưu” → “Đã cập nhật phân công”.\n'
              '3. Khóa: hộp “Xác nhận khóa thành viên” → “Khóa” → “Đã khóa thành viên”.\n'
              '4. Mở khóa: hộp “Xác nhận mở khóa” → “Mở khóa” → “Đã mở khóa thành viên”.\n'
              '5. Xóa: hộp “Xác nhận xóa thành viên” → “Xóa” → “Đã xóa thành viên khỏi giải pháp”.',
        phu='• Khóa khi thành viên còn công việc tồn đọng → mở cửa sổ “Bàn giao công việc tồn đọng trước khi khóa thành '
            'viên” (danh sách Task, Issue chưa xong, tình trạng bàn giao) với nút “Nhắc bàn giao” (→ “Đã gửi nhắc bàn '
            'giao cho thành viên”) và “Mở màn Bàn giao”.\n'
            '• Xóa thành viên đã phát sinh BOM/Task/Issue → “Thành viên đã phát sinh dữ liệu (BOM/Task/Issue). Vui lòng '
            'sử dụng chức năng Khóa để bảo toàn dữ liệu hệ thống.”\n'
            '• Không có quyền trên thành viên → “Bạn không có quyền thao tác trên thành viên này.”'),
   [('menu', MQL + ' => Nhân sự => Sửa phân công / Khóa thành viên / Xóa khỏi giải pháp'),
    ('img', '15-tab-nhan-su.png', 'Tab Nhân sự với cột Thao tác'),
    ('img', '19-sua-phan-cong.png', 'Cửa sổ Sửa phân công thành viên'),
    ('img', '20-khoa-thanh-vien.png', 'Hộp Xác nhận khóa thành viên')],
   [('Bảng nhân sự', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'STT, Thành viên (ảnh, tên, email), Hạng mục, Vai trò, Ngày bắt đầu, Ngày kết thúc, Nhiệm vụ đang phụ trách, '
     'Trạng thái (badge), Thao tác.'),
    ('Nút Sửa phân công / Khóa – Mở khóa / Xóa', 'Icon Button', 'Enable / Ẩn', '–', 'Ẩn khi thiếu quyền',
     'Ẩn ở dòng PM / Leader và khi giải pháp Đóng.'),
    ('Cửa sổ Sửa phân công', 'Modal', 'Enable', '–', 'Theo dữ liệu',
     'Hạng mục* (khi có hạng mục), Vai trò dự án, Mô tả công việc, Ngày bắt đầu, Ngày kết thúc; nút Lưu / Hủy.'),
    ('Hộp xác nhận Khóa / Mở khóa / Xóa', 'Modal', 'Enable', '–', 'Ẩn',
     'Khóa: “Khóa ‘<tên>’? Thành viên sẽ không tạo thêm BOM, Task hay Issue trong giải pháp này nữa. Dữ liệu cũ vẫn '
     'giữ nguyên.” · Mở khóa: “Mở khóa ‘<tên>’? Thành viên tham gia lại giải pháp như bình thường.” · Xóa: “Xóa '
     '‘<tên>’ khỏi danh sách nhân sự giải pháp?”'),
    ('Sơ đồ cấu trúc nhân sự', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', 'Cây PM → hạng mục → Leader → nhân sự.')],
   [('Bấm Lưu (Sửa phân công)', 'Click',
     'Before:\n– Kiểm tra quyền trên thành viên.\nDuring:\n– Ngày kết thúc trước ngày bắt đầu → “Ngày kết thúc phải từ '
     'ngày bắt đầu trở đi.”; trùng hạng mục → “Nhân sự đã có trong hạng mục này.”\nAfter:\n– Lưu, “Đã cập nhật phân công”.'),
    ('Xác nhận Khóa', 'Click',
     'During:\n– Còn công việc tồn đọng → “Thành viên còn n công việc chưa hoàn tất. Vui lòng bàn giao trước khi khóa.” '
     'và mở cửa sổ bàn giao.\nAfter:\n– Khóa thành viên; thành viên không truy cập được giải pháp nữa.'),
    ('Xác nhận Mở khóa', 'Click', 'After:\n– Mở khóa, “Đã mở khóa thành viên”.'),
    ('Xác nhận Xóa', 'Click', 'During:\n– Đã phát sinh dữ liệu → chặn với câu ở dòng sự kiện phụ.\nAfter:\n– Xóa thành viên.'),
    ('Bấm Nhắc bàn giao', 'Click', 'After:\n– Gửi thông báo nhắc; không còn việc tồn đọng → “Thành viên không còn công '
     'việc tồn đọng.”')],
   uc=('FR-18', 'Sửa / Khóa / Xóa thành viên', 'action', A_PM),
   rule=('- Quy tắc ghi lịch sử, Khóa / Mở khóa.', 'history'))

# ---------------------------------------------------------------- 2.19 FR-19
fn('Danh sách hồ sơ trình duyệt',
   dict(ten='Xem danh sách hồ sơ trình duyệt (tab Hồ sơ)',
        mota='Liệt kê hồ sơ trình duyệt giải pháp và hồ sơ trình duyệt hạng mục của giải pháp, kèm nút Xem / Sửa / Duyệt.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn Quản lý giải pháp.',
        chinh='1. Bấm tab “Hồ sơ”.\n2. Hệ thống hiển thị bảng “Danh sách hồ sơ trình duyệt”, mới tạo trước.\n'
              '3. Lọc theo Trạng thái, Loại, Hạng mục, Version HM, Version GP, Ngày gửi, Hạn duyệt hoặc tìm theo mã hồ sơ.\n'
              '4. Bấm “Xem” để mở hồ sơ chỉ đọc; “Sửa” / “Duyệt” khi đủ điều kiện.',
        phu='• Loại = Giải pháp → ẩn ô Hạng mục / Version HM; Loại = Hạng mục → ẩn Version GP.\n'
            '• Lỗi nạp → “Lỗi khi tải dữ liệu hồ sơ trình duyệt”.'),
   [('menu', MQL + ' => Hồ sơ'), ('img', 'h_x_tab_ho_so.png', 'Tab Hồ sơ – danh sách hồ sơ trình duyệt')],
   [('Ô tìm nhanh + Bộ lọc hồ sơ trình duyệt', 'Textbox', 'Enable', '–', 'Trống',
     '“Tìm theo mã hồ sơ”; lọc Trạng thái (Nháp, Chờ duyệt, Đã duyệt, Không duyệt, Hết hiệu lực), Loại (Tất cả / '
     'Giải pháp / Hạng mục), Hạng mục, Version HM, Version GP, Ngày gửi, Hạn duyệt.'),
    ('Bảng hồ sơ', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'STT, Mã / Nội dung hồ sơ (kèm nút), Tên hồ sơ, Loại (badge Giải pháp/Hạng mục), Hạng mục, Version, Trạng thái, '
     'Ngày gửi, Hạn duyệt.'),
    ('Nút Xem', 'Icon Button', 'Enable', '–', 'Hiển thị', 'Mở hồ sơ chỉ đọc.'),
    ('Nút Sửa / Duyệt', 'Icon Button', 'Enable / Ẩn', '–', 'Ẩn',
     'Sửa: người tạo hồ sơ, hồ sơ Nháp. Duyệt: hồ sơ giải pháp Chờ duyệt và người dùng là người tạo giải pháp; hồ sơ '
     'hạng mục Chờ duyệt và người dùng là PM.'),
    ('Phân trang', 'Pagination', 'Enable', '–', '10 dòng/trang', '–')],
   [('Bấm tab Hồ sơ / đổi bộ lọc', 'Click', 'After:\n– Nạp lại danh sách từ trang 1.'),
    ('Bấm Xem / Sửa / Duyệt', 'Click', 'After:\n– Mở cửa sổ hồ sơ tương ứng (FR-20, FR-21, FR-22).')],
   ui_kind='read', rule=('- Màn Danh sách, Phân trang và UI/UX.', 'list'))

# ---------------------------------------------------------------- 2.20 FR-20
fn('Tạo / trình hồ sơ trình duyệt giải pháp',
   dict(ten='Tạo, sửa và trình hồ sơ trình duyệt giải pháp',
        mota='PM lập hồ sơ (tên, nội dung, tài liệu), hệ thống tự gắn BOM tổng hợp Hoàn thành của version hiện tại; '
             'lưu nháp hoặc trình duyệt cho Trưởng phòng. Dự án Tự triển khai: “Lưu & Duyệt” tự duyệt luôn.',
        tacnhan=A_PM + '; ' + A_KD,
        dieukien='Giải pháp Đang triển khai và người dùng là PM.',
        chinh='1. Bấm “Tạo hồ sơ trình duyệt giải pháp” (hoặc “Sửa…” khi hồ sơ gần nhất đang Nháp) ở đầu màn Quản lý.\n'
              '2. Hệ thống mở cửa sổ “Tạo hồ sơ trình duyệt”, hiển thị BOM tổng hợp sẽ gắn và thông tin tham chiếu.\n'
              '3. Nhập Tên hồ sơ, Nội dung trình duyệt, thêm file (tên tài liệu, loại, người thực hiện, file), chọn Hạn duyệt.\n'
              '4. Bấm “Lưu” (Nháp) hoặc “Lưu & Trình duyệt” (dự án Tự triển khai: “Lưu & Duyệt”).\n'
              '5. Hệ thống lưu, báo “Đã lưu hồ sơ trình duyệt giải pháp.” / “Đã lưu và gửi trình duyệt hồ sơ giải pháp.”, '
              'đóng cửa sổ và nạp lại.',
        phu='• Luồng thường chưa có BOM tổng hợp Hoàn thành → cảnh báo đỏ “Chưa có BOM tổng hợp ở trạng thái Hoàn thành '
            'cho version giải pháp này. Vui lòng lập BOM tổng hợp trước khi gửi hồ sơ.”\n'
            '• Tự triển khai chưa có BOM → cảnh báo vàng, vẫn lưu được nhưng bắt buộc có tài liệu đính kèm.\n'
            '• Hồ sơ đã có thì hiện khối bình luận.',
        dacbiet='Trình duyệt: hồ sơ Chờ duyệt, giải pháp → Chờ duyệt giải pháp, BOM → chờ duyệt, thông báo Trưởng phòng. '
                'Tự triển khai: hồ sơ Đã duyệt, giải pháp → Đã duyệt giải pháp, BOM → đã duyệt.'),
   [('menu', MQL + ' => Tạo hồ sơ trình duyệt giải pháp'), ('modal', 'Tạo hồ sơ trình duyệt'),
    ('img', 'h_08_tao_ho_so.png', 'Cửa sổ Tạo hồ sơ trình duyệt (luồng có YCGP)'),
    ('img', 's_07_hoso_nobom.png', 'Hồ sơ của dự án Tự triển khai khi chưa có BOM tổng hợp – nút Lưu & Duyệt')],
   [('Tên hồ sơ', 'Textbox', 'Enable', '0–255 ký tự', 'Có', 'Trống', '“Vui lòng nhập tên hồ sơ.”'),
    ('Nội dung trình duyệt', 'Textarea', 'Enable', '–', 'Có', 'Trống', 'Trình soạn thảo định dạng.'),
    ('Danh sách các file', 'Table/Grid', 'Enable', '–', 'Có (luồng thường); Có khi không có BOM (Tự triển khai)',
     'Trống', 'Tên tài liệu*, Loại tài liệu*, Người thực hiện*, File đính kèm*, Dung lượng; nút Thêm file.'),
    ('BOM tổng hợp gắn vào hồ sơ', 'Label', 'Read-only', '–', '–', 'Tự tìm',
     '“Mã — Tên BOM” (mở tab mới) hoặc cảnh báo khi chưa có.'),
    ('Thông tin tham chiếu', 'Text', 'Read-only', '–', '–', 'Theo giải pháp',
     'Dự án, Giải pháp, Hạng mục, PM phụ trách, Ngày trình duyệt, Thời gian thực hiện.'),
    ('Hạn duyệt', 'Datepicker', 'Enable / Ẩn', 'dd/mm/yyyy, ≥ hôm nay', 'Có (luồng thường)', 'Trống',
     'Ẩn với dự án Tự triển khai.'),
    ('Trưởng phòng duyệt', 'Text', 'Read-only / Ẩn', '–', '–', 'Theo dữ liệu', 'Ẩn với dự án Tự triển khai.'),
    ('Nút Lưu & Trình duyệt (Lưu & Duyệt) / Lưu / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Lưu / Lưu & Trình duyệt', 'Click',
     'Before:\n– Giải pháp phải Đang triển khai: “Chỉ được thao tác hồ sơ trình duyệt khi giải pháp đang ở trạng thái '
     'Đang triển khai.”\nDuring:\n– Tên, Nội dung, Hạn duyệt (luồng thường) trống → câu bắt buộc; Hạn duyệt trước hôm '
     'nay → “Hạn duyệt phải lớn hơn hoặc bằng ngày hôm nay.”; file thiếu trường → “Bắt buộc phải nhập”.\n'
     '– Trình duyệt mà chưa có BOM tổng hợp Hoàn thành → “Chưa có BOM tổng hợp Hoàn thành cho version giải pháp này. '
     'Vui lòng lập BOM tổng hợp trước khi gửi hồ sơ.”\n'
     '– Tự triển khai, không BOM, không tài liệu → “Vui lòng đính kèm ít nhất 1 tài liệu khi hồ sơ không đính kèm BOM list.”\n'
     '– Lỗi nhập liệu → “Vui lòng kiểm tra lại giữ liệu nhập”.\n'
     'After:\n– Lưu hồ sơ (mã HS.TD.<mã GP>.<n>), file; trình duyệt thì gắn BOM, đổi trạng thái hồ sơ / giải pháp / '
     'BOM / dự án như dòng Yêu cầu đặc biệt; thông báo Trưởng phòng.\n– Lỗi khác → “Không thể lưu hồ sơ trình duyệt '
     'giải pháp.”')],
   uc=('FR-20', 'Tạo / trình hồ sơ trình duyệt', 'crud', A_PM),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.21 FR-21
fn('Duyệt / từ chối hồ sơ trình duyệt giải pháp',
   dict(ten='Trưởng phòng duyệt / từ chối hồ sơ trình duyệt',
        mota='Trưởng phòng giải pháp (người tạo giải pháp) duyệt hoặc từ chối hồ sơ đang Chờ duyệt.',
        tacnhan=A_TP, dieukien='Hồ sơ giải pháp ở trạng thái Chờ duyệt; người dùng là người tạo giải pháp; dự án không '
                               'phải Tự triển khai.',
        chinh='1. Ở tab Hồ sơ, bấm “Duyệt” ở dòng hồ sơ Chờ duyệt.\n'
              '2. Hệ thống mở cửa sổ “Duyệt hồ sơ trình duyệt <mã>” chỉ đọc.\n'
              '3. Bấm “Duyệt” → “Đã duyệt hồ sơ trình duyệt.”\n'
              '4. Hoặc bấm “Từ chối”, nhập Lý do từ chối trong hộp “Xác nhận từ chối hồ sơ <mã>?”, bấm “Đồng ý” → “Đã '
              'không duyệt hồ sơ trình duyệt.”',
        phu='• Bấm “Không” ở hộp từ chối → quay lại cửa sổ hồ sơ.\n'
            '• Hồ sơ bị từ chối hiển thị khối “Lý do từ chối” khi mở lại.'),
   [('menu', MQL + ' => Hồ sơ => Duyệt'), ('modal', 'Duyệt hồ sơ trình duyệt'),
    ('img', 'h_09_tp_duyet.png', 'Cửa sổ Duyệt hồ sơ trình duyệt'),
    ('img', 'h_x_tu_choi.png', 'Hộp xác nhận từ chối hồ sơ')],
   [('Nội dung hồ sơ', 'Text', 'Read-only', '–', '–', 'Theo dữ liệu', 'Như FR-20, chỉ đọc; có khối bình luận.'),
    ('Nút Duyệt', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Nút Từ chối', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Màu đỏ, mở hộp nhập lý do.'),
    ('Lý do từ chối', 'Textarea', 'Enable', '0–1000 ký tự', 'Có', 'Trống', '“Vui lòng nhập lý do không duyệt.”'),
    ('Nút Đồng ý / Không / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Duyệt', 'Click',
     'Before:\n– Không phải người tạo giải pháp → “Bạn không có quyền duyệt hồ sơ trình duyệt này.”; hồ sơ không Chờ '
     'duyệt → “Chỉ được duyệt hồ sơ đang ở trạng thái chờ duyệt.”; Tự triển khai → “Dự án tự triển khai không có bước '
     'duyệt hồ sơ riêng.”\nAfter:\n– Hồ sơ Đã duyệt; giải pháp → Đã duyệt giải pháp; BOM đồng bộ; ghi ngày duyệt version; '
     'đồng bộ dự án; thông báo KD chính của dự án.'),
    ('Bấm Đồng ý (từ chối)', 'Click',
     'During:\n– Lý do trống → “Vui lòng nhập lý do không duyệt.”; quá 1000 ký tự → “Lý do không duyệt không được vượt '
     'quá 1000 ký tự.”\nAfter:\n– Hồ sơ Không duyệt (lưu lý do); giải pháp quay về Đang triển khai; BOM đồng bộ.\n'
     '– Lỗi → “Không thể xử lý duyệt hồ sơ trình duyệt.”')],
   uc=('FR-21', 'Duyệt / từ chối hồ sơ trình duyệt', 'action', A_TP),
   rule=('- Thông báo và UI/UX.', 'notice'))

# ---------------------------------------------------------------- 2.22 FR-22
fn('PM duyệt hồ sơ hạng mục',
   dict(ten='PM duyệt / từ chối hồ sơ trình duyệt hạng mục',
        mota='Leader trình hồ sơ hạng mục ở màn Hạng mục giải pháp; PM duyệt hoặc từ chối ngay tại tab Hồ sơ của màn '
             'Quản lý giải pháp.',
        tacnhan=A_PM, dieukien='Hồ sơ hạng mục ở trạng thái Chờ duyệt và người dùng là PM của giải pháp.',
        chinh='1. Ở tab Hồ sơ, bấm “Duyệt” ở dòng hồ sơ Loại = Hạng mục.\n'
              '2. Hệ thống mở cửa sổ “Duyệt hồ sơ trình duyệt: <mã>”.\n'
              '3. Bấm “Duyệt” → “Đã duyệt hồ sơ trình duyệt.”; hoặc “Từ chối”, nhập lý do → “Đã từ chối hồ sơ trình duyệt.”',
        phu='• Lỗi → “Không thể xử lý. Vui lòng thử lại.”\n• Chi tiết lập hồ sơ hạng mục xem SRS Hạng mục giải pháp.'),
   [('menu', MQL + ' => Hồ sơ => Duyệt'), ('modal', 'Duyệt hồ sơ trình duyệt hạng mục'),
    ('img', 'h_x_pm_duyet_hm.png', 'PM duyệt hạng mục / hồ sơ hạng mục')],
   [('Nội dung hồ sơ hạng mục', 'Text', 'Read-only', '–', '–', 'Theo dữ liệu', 'Tên, nội dung, file, thông tin hạng mục.'),
    ('Nút Duyệt / Từ chối / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Duyệt / Từ chối chỉ khi hồ sơ Chờ duyệt.'),
    ('Lý do từ chối', 'Textarea', 'Enable', '–', 'Có', 'Trống', 'Trong hộp xác nhận từ chối.')],
   [('Bấm Duyệt', 'Click', 'Before:\n– Người dùng là PM, hồ sơ Chờ duyệt.\n' + NO_PERM + '\nAfter:\n– Hồ sơ hạng mục '
     'Đã duyệt; nạp lại tab Hồ sơ và giải pháp.'),
    ('Bấm Đồng ý (từ chối)', 'Click', 'During:\n– Lý do trống → báo lỗi tại ô.\nAfter:\n– Hồ sơ Không duyệt, lưu lý do.')],
   uc=('FR-22', 'PM duyệt hồ sơ hạng mục', 'action', A_PM),
   rule=('- Thông báo và UI/UX.', 'notice'))

# ---------------------------------------------------------------- 2.23 FR-23
fn('Tạo version mới',
   dict(ten='Tạo version mới của giải pháp',
        mota='Mở phiên bản mới (V2, V3…) để làm lại giải pháp sau khi đã duyệt (thường sau YC điều chỉnh).',
        tacnhan=A_PM + '; ' + A_TP,
        dieukien='Giải pháp ở trạng thái Đã duyệt giải pháp, Chờ làm giá hoặc Đã duyệt giá.',
        chinh='1. Bấm “Tạo version” ở đầu màn Quản lý.\n2. Hệ thống mở cửa sổ “Tạo phiên bản mới” với Mã phiên bản = '
              'V<hiện tại + 1>.\n3. Nhập Mô tả (tùy chọn), chọn Ngày kết thúc.\n4. Bấm “Tạo mới”.\n'
              '5. Hệ thống tạo version, báo “Tạo phiên bản mới thành công”, nạp lại màn.',
        phu='• Bấm “Đóng” → không tạo.\n• Lỗi nhập liệu → “Vui lòng kiểm tra lại thông tin”.',
        dacbiet='Sau khi tạo: giải pháp về Đang triển khai, tiến độ về 0%, hồ sơ đã duyệt chuyển Hết hiệu lực, dự án về '
                'Đang làm giải pháp.'),
   [('menu', MQL + ' => Tạo version'), ('modal', 'Tạo phiên bản mới'),
    ('img', 'h_10_tao_version.png', 'Cửa sổ Tạo phiên bản mới')],
   [('Mã phiên bản', 'Textbox', 'Read-only', '–', '–', 'V<n+1>', '–'),
    ('Mô tả', 'Textarea', 'Enable', '–', 'Không', 'Trống', 'Gợi ý “Mô tả phiên bản (tuỳ chọn)”.'),
    ('Ngày kết thúc', 'Datepicker', 'Enable', 'dd/mm/yyyy, ≥ hôm nay', 'Có', 'Trống', '–'),
    ('Nút Tạo mới / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Tạo mới', 'Click',
     'Before:\n– Giải pháp không ở trạng thái cho phép → “Bạn không có quyền tạo phiên bản mới cho giải pháp này!”\n'
     'During:\n– Ngày kết thúc trống → “Bắt buộc phải nhập”; trước hôm nay → “Ngày kết thúc phải lớn hơn hoặc bằng ngày '
     'hiện tại”.\nAfter:\n– Lưu ảnh chụp nhân sự và tiến độ version cũ; tạo version mới; cập nhật như dòng Yêu cầu đặc biệt.\n'
     '– “Tạo phiên bản mới thành công”; lỗi khác → “Tạo mới thất bại”.')],
   uc=('FR-23', 'Tạo version mới', 'action', A_PM),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.24 FR-24
fn('Phân bổ % tiến độ',
   dict(ten='Thiết lập phân bổ % tiến độ (tab Tiến độ)',
        mota='PM nhập trọng số (%) cho từng hạng mục (giải pháp có hạng mục) hoặc từng nhiệm vụ (không hạng mục); tiến '
             'độ giải pháp tính theo trọng số × tiến độ thực tế.',
        tacnhan=A_PM, dieukien='Người dùng đang ở tab Tiến độ; chỉ PM được nhập.',
        chinh='1. Bấm tab “Tiến độ”.\n2. Hệ thống hiển thị bảng “Phân bổ tiến độ theo hạng mục” (hoặc “… theo nhiệm vụ”) '
              'và dòng “Đã phân bổ: n% - Chưa phân bổ: m%”.\n3. PM nhập ô Phân bổ (%).\n'
              '4. Sau 1 giây ngừng gõ, hệ thống tự lưu và hiện “Đã lưu”.',
        phu='• Tổng > 100% → “Tổng trọng số không được vượt quá 100%”, không lưu.\n• Lưu lỗi → “Lỗi khi lưu”.\n'
            '• Không phải PM → ô Phân bổ chỉ hiển thị số %.\n'
            '• Chế độ theo nhiệm vụ có bộ lọc: tìm mã/tên nhiệm vụ, Người thực hiện, Ngày giao, Trạng thái.'),
   [('menu', MQL + ' => Tiến độ'), ('img', 'h_11_tien_do.png', 'Tab Tiến độ – phân bổ theo hạng mục')],
   [('Dòng Đã phân bổ / Chưa phân bổ', 'Label', 'Hiển thị', '0 – 100', '–', 'Theo dữ liệu',
     'Đã phân bổ xanh lá khi đủ 100%; Chưa phân bổ đỏ khi > 0.'),
    ('Bảng hạng mục', 'Table/Grid', 'Read-only', '–', '–', 'Theo dữ liệu',
     'STT, Mã hạng mục, Tên hạng mục, Leader, Deadline, Version, Phân bổ (%), Tiến độ TT (%), Trạng thái.'),
    ('Bảng nhiệm vụ (không hạng mục)', 'Table/Grid', 'Read-only', '–', '–', 'Theo dữ liệu',
     'STT, Mã nhiệm vụ, Tên công việc, Người thực hiện, Version, Phân bổ (%), Tiến độ TT (%), Trạng thái.'),
    ('Ô Phân bổ (%)', 'Number', 'Enable / Read-only', '0 – 100', 'Không', 'Theo dữ liệu',
     'Viền đỏ khi < 0 hoặc > 100. Chỉ PM sửa được.'),
    ('Ghi chú', 'Label', 'Hiển thị', '–', '–', 'Hiển thị', '“Thay đổi được lưu tự động”.')],
   [('Nhập ô Phân bổ (%)', 'Change',
     'During:\n– Tính lại tổng; tổng > 100 → “Tổng trọng số không được vượt quá 100%”, không lưu.\n'
     'After:\n– Sau 1 giây tự lưu; máy chủ chỉ nhận khi người dùng là PM ghi trên YCGP (“Bạn không có quyền cập nhật '
     'trọng số”); tính lại tiến độ giải pháp; hiện “Đã lưu” / “Lỗi khi lưu”.')],
   uc=('FR-24', 'Phân bổ % tiến độ', 'crud', A_PM),
   rule=('- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX.', 'create'))

# ---------------------------------------------------------------- 2.25-2.27 tab Nhiệm vụ / Vấn đề / Meeting
fn('Tab Nhiệm vụ',
   dict(ten='Xem và thao tác nhiệm vụ của giải pháp',
        mota='Danh sách nhiệm vụ gắn với giải pháp; các thao tác nhiệm vụ mở bằng chính cửa sổ của màn Nhiệm vụ (chi '
             'tiết tham chiếu SRS Nhiệm vụ).',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn Quản lý giải pháp.',
        chinh='1. Bấm tab “Nhiệm vụ”.\n2. Hệ thống hiển thị bảng “Danh sách nhiệm vụ”, ô Quá hạn, 4 nút lọc nhanh.\n'
              '3. Người dùng tìm/lọc, bấm “Tạo mới”, “Xuất Excel”, hoặc các nút ở dòng (Xem, Sửa, Nhập kết quả, Duyệt, '
              'Lịch sử chỉnh sửa, Xoá).',
        phu='• Nút Tạo mới chỉ hiện khi giải pháp ở Chờ Leader duyệt, Đang triển khai, Chờ duyệt GP, Đã duyệt GP, Chờ '
            'làm giá, Đã duyệt giá.\n• Xuất Excel tải file danh_sach_nhiem_vu.xls; lỗi → “Lỗi khi xuất Excel”.'),
   [('menu', MQL + ' => Nhiệm vụ'), ('img', 'h_06_nhiem_vu.png', 'Tab Nhiệm vụ')],
   [('Bộ lọc danh sách nhiệm vụ', 'Textbox', 'Enable', '–', 'Trống',
     'Tìm mã/tên; lọc Hạng mục, Người thực hiện, Người giao, Trạng thái, Ưu tiên, Người duyệt kết quả, Ngày cập nhật, '
     'Hạn hoàn thành, Version giải pháp.'),
    ('Nút lọc nhanh Nhiệm vụ tôi làm / tôi giao / tôi theo dõi / tôi duyệt kết quả', 'Icon Button', 'Enable', '–',
     'Không bật', '–'),
    ('Nút Tạo mới / Xuất Excel / Tùy chỉnh cột', 'Button', 'Enable / Ẩn', '–', 'Theo trạng thái GP', '–'),
    ('Bảng nhiệm vụ', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'Mã-Tên nhiệm vụ, Hạng mục/Module, Version GP, Trạng thái, Ưu tiên, Người làm, … ; nút ở dòng theo vai trò.')],
   [('Bấm tab / đổi bộ lọc', 'Click', 'After:\n– Nạp nhiệm vụ của giải pháp.'),
    ('Bấm Tạo mới / nút ở dòng', 'Click', 'After:\n– Mở cửa sổ nhiệm vụ tương ứng (SRS Nhiệm vụ), lưu xong nạp lại.'),
    ('Bấm Xuất Excel', 'Click', 'After:\n– Tải file; “Xuất Excel thành công” / “Lỗi khi xuất Excel”.')],
   ui_kind='read')

fn('Tab Vấn đề giải pháp',
   dict(ten='Xem và thao tác vấn đề của giải pháp',
        mota='Danh sách vấn đề (issue) gắn với giải pháp; thao tác mở bằng cửa sổ của màn Vấn đề (chi tiết tham chiếu '
             'SRS Vấn đề).',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn Quản lý giải pháp.',
        chinh='1. Bấm tab “Vấn đề giải pháp”.\n2. Hệ thống hiển thị bảng “Danh sách Vấn đề”, ô Quá hạn, 3 nút lọc nhanh.\n'
              '3. Người dùng tìm/lọc, “Tạo mới”, “Xuất Excel”, hoặc Xem / Sửa / Lịch sử / Xoá ở dòng.',
        phu='• Nút Tạo mới theo cùng danh sách trạng thái với tab Nhiệm vụ.\n• Xuất Excel tải file danh_sach_issue.xls.'),
   [('menu', MQL + ' => Vấn đề giải pháp'), ('img', 'h_07_van_de.png', 'Tab Vấn đề giải pháp')],
   [('Bộ lọc danh sách Vấn đề', 'Textbox', 'Enable', '–', 'Trống',
     'Tìm mã/tên; lọc Hạng mục, Người xử lý, Người tạo, Người duyệt, Trạng thái, Ưu tiên, Ngày tạo, Version giải pháp.'),
    ('Nút lọc nhanh Vấn đề tôi làm / tôi báo cáo / tôi theo dõi', 'Icon Button', 'Enable', '–', 'Không bật', '–'),
    ('Nút Tạo mới / Xuất Excel / Tùy chỉnh cột', 'Button', 'Enable / Ẩn', '–', 'Theo trạng thái GP', '–'),
    ('Bảng vấn đề', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'Mã-Tên Vấn đề, Dự án / Module, Version GP, Người xử lý, Ưu tiên, Trạng thái, …')],
   [('Bấm tab / đổi bộ lọc', 'Click', 'After:\n– Nạp vấn đề của giải pháp.'),
    ('Bấm Tạo mới / nút ở dòng', 'Click', 'After:\n– Mở cửa sổ vấn đề (SRS Vấn đề), lưu xong nạp lại.'),
    ('Bấm Xuất Excel', 'Click', 'After:\n– Tải file danh sách vấn đề.')],
   ui_kind='read')

fn('Tab Meeting',
   dict(ten='Xem meeting của giải pháp',
        mota='Danh sách cuộc họp gắn với dự án của giải pháp; tạo mới / xem / sửa / in mở sang màn Meeting (chi tiết '
             'tham chiếu SRS Tổng hợp meeting).',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn Quản lý giải pháp.',
        chinh='1. Bấm tab “Meeting”.\n2. Hệ thống hiển thị bảng “Danh sách Meetings”.\n'
              '3. Bấm “Tạo mới” để sang màn tạo meeting gắn sẵn dự án; hoặc Xem / Chỉnh sửa / In / Tạo phiếu công tác '
              'khác ở dòng.',
        phu='• Không có meeting → bảng trống “Không có dữ liệu”.\n• Nút Tạo mới chỉ hiện khi giải pháp gắn dự án.'),
   [('menu', MQL + ' => Meeting'), ('img', '24-tab-meeting.png', 'Tab Meeting')],
   [('Bộ lọc danh sách Meetings', 'Textbox', 'Enable', '–', 'Trống',
     'Tìm mã/tên; lọc Loại Meeting, Hình thức, Trạng thái, Có biên bản?, Người cập nhật gần nhất, Version giải pháp.'),
    ('Nút Tạo mới / Xuất Excel / Tùy chỉnh cột', 'Button', 'Enable / Ẩn', '–', 'Hiển thị', '–'),
    ('Bảng meeting', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'Mã / Tên Meeting, Version GP, Thời gian, Thời lượng họp, Loại & Hình thức, Khách hàng, Thành phần tham dự, '
     'Trạng thái, Biên bản.')],
   [('Bấm tab / đổi bộ lọc', 'Click', 'After:\n– Nạp meeting của giải pháp.'),
    ('Bấm Tạo mới / Xem / Chỉnh sửa', 'Click', 'After:\n– Chuyển sang màn Meeting tương ứng; In mở bản in ở tab mới.'),
    ('Bấm Xuất Excel', 'Click', 'After:\n– Tải file danh_sach_cuoc_hop.xls.')],
   ui_kind='read')

# ---------------------------------------------------------------- 2.28 Files
fn('Tab Files',
   dict(ten='Xem tài liệu của giải pháp',
        mota='Tổng hợp mọi file đính kèm phát sinh trong giải pháp (hồ sơ, nhiệm vụ, vấn đề…), cho xem trước và tải về.',
        tacnhan=A_ALL, dieukien='Người dùng đang ở màn Quản lý giải pháp.',
        chinh='1. Bấm tab “Files”.\n2. Hệ thống hiển thị bảng “Danh sách Files”.\n'
              '3. Người dùng tìm/lọc, bấm Xem trước (file xem được) hoặc Tải xuống; bấm liên kết để mở hồ sơ / nhiệm vụ / '
              'vấn đề gốc.',
        phu='• Không có file → “Không có file phù hợp bộ lọc.”'),
   [('menu', MQL + ' => Files'), ('img', '26-tab-files.png', 'Tab Files')],
   [('Bộ lọc danh sách Files', 'Textbox', 'Enable', '–', 'Trống',
     '“Tìm theo tên file, tên tài liệu, mã liên kết”; lọc Nguồn, Nhóm tài liệu, Hạng mục, Version, Người tạo.'),
    ('Bảng file', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu',
     'STT, Tên tài liệu, File, Nguồn, Nhóm tài liệu, Hạng mục, Version, Liên kết, Người tạo, Ngày tạo, Dung lượng, Thao tác.'),
    ('Nút Xem trước / Tải xuống', 'Icon Button', 'Enable / Ẩn', '–', 'Hiển thị', 'Xem trước chỉ với file xem được.')],
   [('Bấm tab / đổi bộ lọc', 'Click', 'After:\n– Nạp danh sách file.'),
    ('Bấm Xem trước / Tải xuống', 'Click', 'After:\n– Mở cửa sổ xem trước / mở file ở tab mới.'),
    ('Bấm liên kết', 'Click', 'After:\n– Mở hồ sơ trình duyệt hoặc cửa sổ nhiệm vụ / vấn đề gốc.')],
   ui_kind='read')

# ---------------------------------------------------------------- 2.29 YC điều chỉnh
fn('Xử lý yêu cầu điều chỉnh giải pháp',
   dict(ten='Tiếp nhận / từ chối yêu cầu điều chỉnh giải pháp',
        mota='Xem các yêu cầu điều chỉnh giải pháp do người tạo dự án gửi; Trưởng phòng hoặc PM tiếp nhận hoặc từ chối.',
        tacnhan=A_TP + '; ' + A_PM,
        dieukien='Yêu cầu ở trạng thái Đã gửi; người dùng là Trưởng phòng (người tạo GP) hoặc PM.',
        chinh='1. Bấm tab “YC Điều chỉnh”.\n2. Hệ thống hiển thị bảng “Yêu cầu điều chỉnh giải pháp”.\n'
              '3. Bấm mã yêu cầu hoặc “Xem chi tiết” để xem nội dung, file.\n'
              '4. Bấm “Tiếp nhận” → xác nhận “Bạn có chắc muốn tiếp nhận yêu cầu điều chỉnh này?” → “Đã tiếp nhận yêu '
              'cầu điều chỉnh”.\n5. Hoặc “Từ chối” → nhập Lý do từ chối → “Gửi” → “Đã từ chối yêu cầu điều chỉnh”.',
        phu='• Không có yêu cầu → “Chưa có yêu cầu điều chỉnh giải pháp nào.”\n'
            '• Lý do trống → “Lý do từ chối không được để trống”.\n• Lỗi → “Có lỗi xảy ra, vui lòng thử lại”.',
        dacbiet='Tiếp nhận: các yêu cầu tính giá bán của dự án đang Chờ/Đang xác định giá chuyển Dừng; báo giá đang lập / '
                'chờ duyệt chuyển Dừng và báo người lập báo giá. Sau đó PM tạo version mới (FR-23) để làm lại.'),
   [('menu', MQL + ' => YC Điều chỉnh'), ('img', '27-tab-ycdc.png', 'Tab YC Điều chỉnh')],
   [('Bảng yêu cầu điều chỉnh', 'Table/Grid', 'Read-only', '–', '–', 'Theo dữ liệu',
     'STT, Mã yêu cầu (link), Version, Người yêu cầu, Ngày gửi, Trạng thái (Đã gửi / Tiếp nhận / Từ chối), Hành động.'),
    ('Nút Xem chi tiết / Tiếp nhận / Từ chối', 'Icon Button', 'Enable / Ẩn', '–', '–', 'Theo điều kiện',
     'Tiếp nhận / Từ chối chỉ khi Đã gửi và người dùng là TP hoặc PM.'),
    ('Cửa sổ Chi tiết yêu cầu điều chỉnh', 'Modal', 'Read-only', '–', '–', 'Ẩn', 'Nội dung, file, người xử lý, lý do.'),
    ('Lý do từ chối', 'Textarea', 'Enable', '–', 'Có', 'Trống', 'Trong cửa sổ “Từ chối yêu cầu điều chỉnh”.'),
    ('Nút Gửi / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–')],
   [('Bấm Tiếp nhận (xác nhận)', 'Click',
     'Before:\n– Nút chỉ hiện khi Đã gửi và người dùng là TP/PM.\nAfter:\n– Yêu cầu → Tiếp nhận; dừng yêu cầu tính giá / '
     'báo giá liên quan; “Đã tiếp nhận yêu cầu điều chỉnh”.'),
    ('Bấm Gửi (từ chối)', 'Click',
     'During:\n– Lý do trống → “Lý do từ chối không được để trống”.\nAfter:\n– Yêu cầu → Từ chối, lưu lý do; “Đã từ chối '
     'yêu cầu điều chỉnh”.')],
   uc=('FR-29', 'Xử lý YC điều chỉnh', 'action', A_TP),
   rule=('- Thông báo và UI/UX.', 'notice'))

# ========================================================= PHẦN 4
d.h1('Phần 4. Quy tắc nghiệp vụ')
d.p('Quy tắc áp dụng: chỉ ghi các quy tắc đặc thù của màn Quản lý giải pháp; quy tắc chung xem SRS Các quy tắc chung.')
d.rule_table([
    ('BR-01', 'Phạm vi giải pháp được xem',
     ['– Theo quyền cấp cao nhất đang có: V1 toàn bộ; V2 công ty đang làm việc; V3 phòng ban/bộ phận quản lý + do mình '
      'tạo; V4 bộ phận quản lý + do mình tạo; không có quyền: do mình tạo.',
      '– Cộng thêm với mọi người: giải pháp mình là PM, Leader hạng mục hoặc thành viên.',
      '– Giải pháp Nháp chỉ người tạo nhìn thấy.',
      '– Thành viên đã bị khóa không xem / thao tác được giải pháp đó.'],
     ['Xem danh sách', 'Xuất Excel', 'Xem chi tiết']),
    ('BR-02', 'Mã tự sinh',
     ['– Mã GP: <Mã dự án>_GP<STT 2 số>, STT = tổng số giải pháp + 1, tăng tiếp nếu trùng; không sửa sau khi tạo.',
      '– Mã hạng mục: <Mã GP>_HMnn theo thứ tự hạng mục.',
      '– Mã hồ sơ: HS.TD.<Mã GP>.<số thứ tự hồ sơ>. Version: V1 khi tạo, +1 mỗi lần tạo version.'],
     ['Tạo mới', 'Tạo hồ sơ', 'Tạo version']),
    ('BR-03', 'Trạng thái giải pháp',
     ['– Nháp (xám) · Chờ PM duyệt (cam) · Chờ Leader duyệt (cam) · Đang triển khai (xanh dương) · Chờ duyệt giải pháp '
      '(tím) · Đã duyệt giải pháp (xanh lá) · Đã duyệt giá (xanh lá) · Chờ làm giá (cam) · Chốt giải pháp (tím) · Đóng (xám).'],
     ['Xem danh sách', 'Tìm kiếm']),
    ('BR-04', 'Luồng trạng thái – dự án có YCGP',
     ['– Nháp → (Lưu và gửi) Chờ PM duyệt.',
      '– Chờ PM duyệt → PM “Giao cho Leader” → Chờ Leader duyệt (bắt buộc ≥ 1 hạng mục); hoặc “Lưu và duyệt” khi không '
      'có hạng mục → Đang triển khai.',
      '– Chờ Leader duyệt → khi mọi hạng mục đã được Leader duyệt → Đang triển khai.',
      '– Đang triển khai → PM trình hồ sơ → Chờ duyệt giải pháp → TP duyệt → Đã duyệt giải pháp; TP từ chối → Đang triển khai.',
      '– Đã duyệt giải pháp → Chờ làm giá (khi lập yêu cầu tính giá / báo giá) → Đã duyệt giá (khi báo giá được duyệt).',
      '– Chốt giải pháp: thực hiện ở màn Dự án TKT (nút Chốt giải pháp, chọn hồ sơ đã duyệt).',
      '– Tạo version mới từ Đã duyệt giải pháp / Chờ làm giá / Đã duyệt giá → Đang triển khai.',
      '– Mỗi lần đổi trạng thái đồng bộ trạng thái version hiện tại và trạng thái dự án TKT.'],
     ['Tạo mới', 'Chỉnh sửa', 'PM duyệt', 'Leader duyệt', 'Hồ sơ trình duyệt', 'Tạo version']),
    ('BR-05', 'Luồng dự án Tự triển khai (sale tự làm)',
     ['– Không có YCGP, không có hạng mục; PM luôn là người tạo (nhân viên KD).',
      '– Lưu và gửi → Đang triển khai (bỏ qua duyệt PM/Leader).',
      '– Hồ sơ trình duyệt: nút “Lưu & Duyệt”, không có Hạn duyệt / Trưởng phòng; hồ sơ tự Đã duyệt, giải pháp → Đã duyệt '
      'giải pháp ngay.',
      '– BOM tổng hợp không bắt buộc, nhưng không có BOM thì bắt buộc ≥ 1 tài liệu đính kèm.',
      '– Không có bước TP duyệt / từ chối hồ sơ.'],
     ['Tạo mới', 'Hồ sơ trình duyệt']),
    ('BR-06', 'Điều kiện tạo giải pháp',
     ['– Mỗi dự án chỉ 1 giải pháp; dự án cha không làm giải pháp.',
      '– Luồng thường: chọn được YCGP khi YCGP Đã tiếp nhận và người dùng là người tiếp nhận; tạo xong YCGP → Đang thực '
      'hiện; xóa giải pháp Nháp thì YCGP quay về Đã tiếp nhận.',
      '– Tự triển khai: dự án ở Thu thập thông tin dự án, “Có làm GP”, phiếu thu thập đủ trường bắt buộc.'],
     ['Tạo mới', 'Xóa']),
    ('BR-07', 'Điều kiện Sửa / Xóa',
     ['– Sửa: Nháp – người tạo; Chờ PM duyệt – PM; Chờ Leader duyệt – PM hoặc Leader có hạng mục chưa duyệt; Đang triển '
      'khai – PM; Chờ duyệt / Đã duyệt giải pháp – người có Q1 thuộc phòng tiếp nhận YCGP.',
      '– Xóa: chỉ giải pháp Nháp, chỉ người tạo; xóa hẳn.',
      '– Nút không dùng được thì ẩn ở cả danh sách và màn chi tiết.'],
     ['Chỉnh sửa', 'Xóa']),
    ('BR-08', 'Điều kiện mở màn Quản lý giải pháp',
     ['– Không mở được khi giải pháp Nháp hoặc Đóng.',
      '– Mở được: PM, người tạo giải pháp; luồng “Triển khai theo phòng”: người có Q1 thuộc phòng tiếp nhận YCGP; luồng '
      'liên phòng ban: người có Q1 thuộc phòng quản lý giải pháp. Dự án Tự triển khai: chỉ PM / người tạo.'],
     ['Quản lý giải pháp']),
    ('BR-09', 'Ràng buộc ngày',
     ['– Ngày cần xong GP và ngày cần hoàn thành hạng mục ≥ hôm nay chỉ khi giải pháp đang Nháp / tạo mới.',
      '– Ngày bắt đầu khi Thêm nhân sự, Hạn duyệt hồ sơ, Ngày kết thúc version: ≥ hôm nay.',
      '– Đổi Ngày cần xong GP cập nhật ngày kết thúc của version hiện tại.'],
     ['Tạo mới', 'Chỉnh sửa', 'Thêm nhân sự', 'Hồ sơ', 'Tạo version']),
    ('BR-10', 'Hồ sơ trình duyệt và BOM',
     ['– Chỉ lập / trình khi giải pháp Đang triển khai; chỉ PM thấy nút.',
      '– Tự gắn BOM tổng hợp ở trạng thái Hoàn thành mới nhất của version hiện tại; luồng thường bắt buộc có khi trình.',
      '– Trạng thái BOM đi theo hồ sơ: trình → chờ duyệt; duyệt → đã duyệt; từ chối → theo quyết định.',
      '– Người duyệt hồ sơ giải pháp là người tạo giải pháp (Trưởng phòng); hồ sơ hạng mục do PM duyệt.',
      '– Màn Quản lý giải pháp không có tab BOM riêng; BOM lập ở màn BOM Giải pháp.'],
     ['Hồ sơ trình duyệt', 'Duyệt hồ sơ']),
    ('BR-11', 'Version mới',
     ['– Cho phép khi Đã duyệt giải pháp / Chờ làm giá / Đã duyệt giá.',
      '– Lưu ảnh chụp nhân sự + tiến độ version cũ; hồ sơ đã duyệt → Hết hiệu lực; tiến độ về 0%; giải pháp → Đang triển khai.'],
     ['Tạo version']),
    ('BR-12', 'Phân công PM / Leader',
     ['– Đổi PM: chỉ người tạo giải pháp; đổi Leader: người tạo hoặc PM; giải pháp Đóng không phân công được.',
      '– PM/Leader mới phải cùng Phòng làm GP và khác người hiện tại; ghi chú ≥ 50 ký tự.',
      '– Người cũ có vai trò mới → được thêm vào nhân sự (PM cũ vào hạng mục đã chọn nếu GP có hạng mục); không chọn → rời dự án.',
      '– Mỗi lần phân công ghi 1 dòng lịch sử.'],
     ['Phân công', 'Lịch sử phân công']),
    ('BR-13', 'Thành viên giải pháp',
     ['– Người tạo / PM quản lý mọi thành viên; Leader chỉ thành viên hạng mục của mình; dòng PM, Leader không có thao tác.',
      '– Khóa bị chặn khi thành viên còn Task / Issue chưa hoàn tất — phải bàn giao trước.',
      '– Xóa bị chặn khi thành viên đã phát sinh BOM / Task / Issue — dùng Khóa.',
      '– Thành viên bị khóa không tạo được BOM, Task, Issue và không truy cập được giải pháp.'],
     ['Thêm nhân sự', 'Quản lý thành viên']),
    ('BR-14', 'Phân bổ tiến độ',
     ['– Có hạng mục: trọng số theo hạng mục; không hạng mục: theo nhiệm vụ.',
      '– Tổng trọng số ≤ 100%; tự lưu sau 1 giây; lưu xong tính lại tiến độ giải pháp.'],
     ['Tiến độ']),
    ('BR-15', 'Thông báo',
     ['– Giải pháp sang Chờ PM duyệt → PM: “Bạn có giải pháp cần duyệt: <mã> - <tên>”.',
      '– Sang Chờ Leader duyệt → các Leader: “Bạn có giải pháp cần duyệt hạng mục: <mã> - <tên>”.',
      '– Trình hồ sơ → Trưởng phòng: “Hồ sơ trình duyệt <mã HS> của giải pháp <mã> - <tên> đã được gửi trình duyệt”.',
      '– TP duyệt hồ sơ → KD chính dự án: “… đã được duyệt”.',
      '– Phân công: người cũ “Bạn đã được thay thế vị trí PM/Leader giải pháp …”, người mới “Bạn được phân công làm '
      'PM/Leader giải pháp …”.',
      '– Không gửi cho chính người thao tác; bấm thông báo mở giải pháp.'],
     ['Tạo mới', 'Chỉnh sửa', 'Hồ sơ', 'Phân công']),
    ('BR-16', 'Tạo nhiệm vụ / vấn đề trong giải pháp',
     ['– Chỉ khi giải pháp ở Chờ Leader duyệt, Đang triển khai, Chờ duyệt GP, Đã duyệt GP, Chờ làm giá, Đã duyệt giá.'],
     ['Tab Nhiệm vụ', 'Tab Vấn đề']),
    ('BR-17', 'Đóng giải pháp',
     ['– Máy chủ có thao tác Đóng (quyền Q2, không áp cho Nháp / Chốt giải pháp) nhưng màn hình chưa có nút Đóng.'],
     ['–']),
])

d.save(update_fields=os.environ.get('SRS_NO_WORD') != '1')
