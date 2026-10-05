# -*- coding: utf-8 -*-
"""Sinh "SRS - Báo giá.docx" theo FORM CHUẨN (2026-09-24).

Chạy:  SRS_NO_WORD=1 /opt/homebrew/bin/python3 gen_srs.py
Nguồn đối chiếu code (nhánh gop_db):
  FE  pages/assign/quotations/{index.vue, create.vue, _id/index.vue, _id/edit.vue}
      pages/assign/quotations/components/{QuotationImportModal, QuotationProductSearchModal,
      QuotationProductEditModal, QuotationCopyPreviewModal}.vue
      components/assign/quotation/{QuotationSubmitModal, QuotationRejectModal, QuotationHistoryModal,
      QuotationPrintConfigModal, QuotationPrintPreview, QuotationLowPriceWarningModal,
      VatBulkApplyToolbar, VatFirstEntryPromptModal}.vue · components/assign/SystemInfoSection.vue
      utils/mixins/Assign/QuotationCopyMixin.js
      components/subsystem-menu/{presale.js, sale-hub.js, sale.js} · components/subsystems.js
  BE  Modules/Assign/Routes/api.php (prefix assign/quotations)
      Http/Controllers/Api/V1/QuotationController.php · Services/QuotationService.php
      Http/Requests/Quotation/{QuotationStoreRequest, QuotationRejectRequest, QuotationFinalizeRequest}
      Entities/Quotation.php (STATUS_*, TYPE_*, LEVEL_*) · app/Helper/PermissionHelper.php
  Quyền: PermissionsTableSeeder id 1080–1086, 1091, 1092, 1173 (nhóm "Báo giá")
Ảnh: shots/ (Playwright headless 1440x900, client :3002) — dữ liệu mẫu xem data_created.md.
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = HERE
while not os.path.isdir(os.path.join(ROOT, '.claude', 'skills', 'srs-documenter')):
    ROOT = os.path.dirname(ROOT)
sys.path.insert(0, os.path.join(ROOT, '.claude', 'skills', 'srs-documenter', 'assets'))
from srs_docx_lib import SrsDoc  # noqa: E402
from docx.shared import Inches  # noqa: E402

SHOTS = os.path.join(HERE, 'shots')


def shot(name):
    return os.path.join(SHOTS, name)


TEN_MAN = 'Báo giá'
MENU_A = 'Phân hệ CSKH trước bán => Báo giá'
MENU_B = 'Phân hệ Bán hàng => Bán hàng => Báo giá => Danh sách báo giá'

A_SALE = 'Người lập báo giá'
A_DUYET = 'Người duyệt giá (TP / BGĐ)'
A_ALL = 'Người dùng đã đăng nhập'

ICONS = {
    'Phân hệ CSKH trước bán': 'icon_phanhe_presale.png',
    'Phân hệ Bán hàng': 'icon_phanhe_bh.png',
    'Danh sách báo giá': 'icon_bh_dsbaogia.png',
    'Tạo báo giá': 'icon_taomoi.png',
    'Tìm kiếm nâng cao': 'icon_timkiem.png',
    'Cài đặt bộ lọc': 'icon_caidat.png',
    'Cấu hình cột hiển thị': 'icon_cot.png',
    'Mã báo giá': 'icon_ma.png',
    'Sao chép báo giá': 'icon_saochep.png',
    'In báo giá': 'icon_in.png',
    'Lịch sử phê duyệt': 'icon_lichsu.png',
    'Xóa': 'icon_xoa.png',
    'Thêm mới': 'icon_themmoi_hang.png',
    'Thêm nhóm': 'icon_themnhom.png',
    'Nhân bản': 'icon_nhanban.png',
    'Thêm con': 'icon_themcon.png',
    'Import Excel': 'icon_import.png',
    'Lưu nháp': 'icon_luunhap.png',
    'Gửi duyệt': 'icon_guiduyet.png',
    'Duyệt': 'icon_duyet.png',
    'Duyệt & chuyển BGĐ': 'icon_duyet_bgd.png',
    'BGĐ duyệt': 'icon_bgd_duyet.png',
    'Từ chối': 'icon_tuchoi.png',
    'Lưu ghi chú': 'icon_luughichu.png',
    'Lập hợp đồng': 'icon_laphd.png',
    'Xem lịch sử': 'icon_xemlichsu.png',
    'Sao chép': 'icon_saochep_ct.png',
    'Xoá': 'icon_xoa_ct.png',
}

out = os.path.join(HERE, 'SRS - %s.docx' % TEN_MAN)
d = SrsDoc(out=out, menu=MENU_A, route='', full_url='', img_prefix='quot_')
d.set_menu_icons({k: shot(v) for k, v in ICONS.items()})


def _icon_for(segs, i, seg):
    """Cùng 1 chữ nhưng khác phần tử trên giao diện -> chọn icon theo VỊ TRÍ / ngữ cảnh."""
    if seg == 'Báo giá':
        if segs[0] == 'Phân hệ CSKH trước bán':
            return shot('icon_presale_baogia.png')
        return shot('icon_bh_baogia.png')
    if seg == 'Bán hàng' and i == 1:
        return shot('icon_bh_nhom.png')
    in_detail = 'Mã báo giá' in segs[:i]
    if seg == 'Xuất Excel':
        return shot('icon_xuat_ct2.png' if in_detail else 'icon_xuat.png')
    if seg == 'In':
        return shot('icon_in_ct.png')
    if seg == 'Sửa / Làm giá':
        return shot('icon_sua_ct.png' if in_detail else 'icon_sua.png')
    return d.menu_icons.get(seg)


def _menu_para(menu):
    """Dòng "Menu:" — KHÔNG tách theo " / " (nút "Sửa / Làm giá" là 1 nút);
    nhiều nút thay thế nhau ghi bằng " | " và in ra là " / "."""
    par = d.p('Menu: ')
    segs = menu.split(' => ')
    for i, seg in enumerate(segs):
        if i:
            par.add_run(' => ')
        for j, part in enumerate(seg.split(' | ')):
            if j:
                par.add_run(' / ')
            ic = _icon_for(segs, i, part)
            par.add_run(part + (' ' if ic else ''))
            if ic:
                par.add_run().add_picture(ic, height=Inches(d.menu_icon_h))
    return par


d._menu_para = _menu_para


def lay(suffix, shots, note=None):
    """Mục Layout: liệt kê đủ 2 lối vào menu rồi tới ảnh."""
    d.layout(menu=MENU_A + suffix)
    _menu_para(MENU_B + suffix)
    if note:
        d.p(note)
    for png, cap in shots:
        d.figure(shot(png), cap, width_in=6.2)


NO_PERM = ('– Không thỏa điều kiện → nút không hiển thị; gọi thẳng chức năng thì hệ thống từ chối kèm '
           'thông báo và dừng xử lý.')
OWNER = ('– Chỉ người lập báo giá, khi báo giá ở trạng thái Đang tạo. Sai trạng thái → "Báo giá không ở '
         'trạng thái Đang tạo, không thể sửa."; không phải người lập → "Chỉ người tạo báo giá mới có thể sửa."')

COUNTER = {'n': 0}


def fr(code, ten, group, actor, anchor, tail, intro, layout, ui, events, mode='full', usecase=True):
    """1 chức năng 2.x — 5 mục con (hoặc 4 nếu không có Biểu đồ Usecase)."""
    COUNTER['n'] += 1
    n = COUNTER['n']
    d.h3('2.%d %s' % (n, ten))
    k = 1
    if usecase:
        d.p('2.%d.%d Biểu đồ Usecase' % (n, k))
        d.uc_figure(code, ten, group, actor=actor)
        k += 1
    d.p('2.%d.%d Giới thiệu' % (n, k))
    k += 1
    d.rule_ref(tail, anchor=anchor)
    d.intro_table(ten=ten, **intro)
    d.p('2.%d.%d Layout màn hình' % (n, k))
    k += 1
    lay(*layout)
    d.p('2.%d.%d Mô tả chi tiết giao diện' % (n, k))
    k += 1
    if not ui:
        raise ValueError('Thiếu bảng giao diện: ' + ten)
    if mode == 'full':
        d.ui_table(ui)
    elif mode == 'ro':
        d.ui_table(ui, required=False)
    else:
        d.ui_table(ui, required=False, scope=False)
    d.p('2.%d.%d Danh sách event và xử lý event' % (n, k))
    if not events:
        raise ValueError('Thiếu bảng event: ' + ten)
    d.event_table(events)


d.title_block(TEN_MAN)
d.h2('Mục lục')
d.toc()

# ======================================================================= PHẦN 1
d.h1('Phần 1. Giới thiệu')
d.h2('1 Mục đích')
d.p('Tài liệu này đặc tả yêu cầu phần mềm cho màn Báo giá (danh sách báo giá, màn Tạo báo giá / Làm giá và màn '
    'Chi tiết báo giá) — có lối vào ở phân hệ CSKH trước bán và phân hệ Bán hàng, nhằm:')
d.bullets([
    'Là căn cứ nghiệm thu các chức năng: xem danh sách, tìm kiếm/lọc, cài đặt bộ lọc, tùy chỉnh cột, xuất Excel, '
    'tạo mới (nạp hàng từ BOM tổng hợp hoặc báo giá tự lập), làm giá hàng hoá/dịch vụ, thêm hàng tạm, giảm giá, '
    'tổng hợp giá trị và cấp duyệt dự kiến, điều khoản báo giá, import Excel, lưu nháp, gửi duyệt, sửa, xem chi '
    'tiết, duyệt, từ chối, xóa, sao chép, in, xuất Excel báo giá, ghi chú kinh doanh, lập hợp đồng, lịch sử.',
    'Làm rõ cách hệ thống xác định cấp duyệt C1/C2/C3 theo Cấu hình duyệt giá (tối đa 3 điều kiện) và luồng '
    'trạng thái của báo giá: Đang tạo → Chờ TP duyệt → Chờ BGĐ duyệt → Đã duyệt → Trúng thầu.',
    'Làm rõ ai được thao tác gì: người lập báo giá, Sale phụ trách dự án, Trưởng phòng / Ban giám đốc duyệt giá, '
    'và phạm vi báo giá mỗi người được xem.',
])
d.p('Màn Báo giá chờ duyệt (hàng chờ của người duyệt) có SRS riêng "SRS - Báo giá chờ duyệt"; các nút Duyệt / '
    'Từ chối được đặc tả trong tài liệu này vì nằm trên màn Chi tiết báo giá. Ngưỡng cấp duyệt được đặc tả ở '
    '"SRS - Cấu hình duyệt giá".')

d.h2('2 Thuật ngữ và viết tắt')
d.table(['Thuật ngữ', 'Mô tả'], [
    ('BG', 'Báo giá. Mã tự sinh dạng BG-YYYY-NNNNN, chỉ cấp khi báo giá được lưu lần đầu.'),
    ('Báo giá từ BOM', 'Loại báo giá "Từ BOM": hàng hoá nạp từ BOM tổng hợp đã duyệt của dự án; cấu trúc hàng '
     'hoá bị khoá, chỉ sửa giá.'),
    ('Báo giá tự lập', 'Loại báo giá "Tự nhập": dự án không có BOM tổng hợp đã duyệt; người lập tự thêm nhóm, '
     'hàng hoá ERP, hàng tạm và dịch vụ.'),
    ('Báo giá tổng', 'Báo giá của dự án cha, gộp từ báo giá con. Hiện chung danh sách (badge "Báo giá tổng") '
     'nhưng mở màn riêng, không thuộc tài liệu này.'),
    ('Hàng ERP', 'Hàng hoá có trong danh mục hàng hoá ERP; giá bán lấy theo Bảng giá, VAT lấy theo danh mục, '
     'không sửa được.'),
    ('Hàng tạm', 'Hàng hoá chưa có trong danh mục ERP, thêm tạm vào báo giá; mã tự sinh dạng HHBG + số khi lưu.'),
    ('YCBG', 'Yêu cầu xây dựng giá (màn Yêu cầu tính giá bán). Báo giá lập từ YCBG hiện mã YCBG.'),
    ('TSLN', 'Tỷ suất lợi nhuận (%) = (Thành tiền bán − Thành tiền nhập) / Thành tiền nhập × 100.'),
    ('Mức sàn', 'Tỷ suất lợi nhuận mức sàn của công ty — dưới mức này số TSLN hiển thị màu đỏ.'),
    ('C1 / C2 / C3', 'Cấp duyệt: C1 người lập tự duyệt; C2 Trưởng phòng duyệt; C3 Trưởng phòng duyệt & chuyển '
     'Ban giám đốc duyệt.'),
    ('GG', 'Giảm giá: GG mặt hàng (theo từng dòng) hoặc GG tổng (theo tổng đơn, phân bổ xuống từng dòng).'),
    ('Sale phụ trách', 'Nhân viên kinh doanh phụ trách chính của dự án TKT.'),
    ('Báo giá Mỏ neo', 'Báo giá đã duyệt của dự án cha; báo giá dự án con phải khớp về tiền tệ, giảm giá, đơn giá.'),
], widths=[1.6, 4.4])

# ======================================================================= PHẦN 2
d.h1('Phần 2. Phân quyền')
d.h2('1 Danh sách quyền')
d.p('Nhóm quyền thao tác:')
d.table(['Ký hiệu', 'Tên quyền', 'Tác dụng trên màn hình'], [
    ('Q1', 'Xây dựng giá bán theo công ty', 'Lập / sao chép báo giá từ YCBG của dự án triển khai chéo phòng hoặc '
     'tự triển khai.'),
    ('Q2', 'Xây dựng giá bán theo phòng', 'Lập / sao chép báo giá từ YCBG của dự án triển khai theo phòng, YCBG '
     'phải cùng phòng ban với người dùng.'),
    ('Q3', 'Trưởng phòng duyệt giá Bom giải pháp', 'Hiện nút Duyệt / Duyệt & chuyển BGĐ / Từ chối với báo giá '
     'Chờ TP duyệt thuộc phòng ban mình quản lý.'),
    ('Q4', 'Ban giám đốc duyệt giá Bom giải pháp', 'Hiện nút BGĐ duyệt / Từ chối với báo giá Chờ BGĐ duyệt cùng '
     'công ty.'),
    ('Q5', 'Xem giá vốn hàng hoá', 'Xem Giá nhập, Thành tiền nhập, TSLN của hàng ERP, cấp duyệt dự kiến; được chọn '
     'hàng ERP làm hàng con.'),
    ('Q6', 'Cho phép thêm giảm giá trong báo giá', 'Được chọn phương thức giảm giá và nhập / xoá giảm giá.'),
], widths=[0.8, 2.0, 3.2])
d.p('Ngoài quyền, nhiều thao tác gắn với VAI TRÒ trên dữ liệu: người lập báo giá (sửa, xoá, gửi duyệt, import, '
    'lập hợp đồng), Sale phụ trách dự án (tạo báo giá tự lập, sao chép báo giá tự lập, ghi chú kinh doanh). Màn '
    'danh sách, xem chi tiết, in, xuất Excel và lịch sử không gắn quyền thao tác riêng — chỉ giới hạn theo phạm vi '
    'dữ liệu dưới đây.')
d.p('Nhóm quyền quyết định phạm vi dữ liệu:')
d.table(['Ký hiệu', 'Tên quyền', 'Phạm vi dữ liệu'], [
    ('V1', 'Xem danh sách Báo giá theo tổng công ty', 'Toàn bộ báo giá.'),
    ('V2', 'Xem danh sách Báo giá theo công ty', 'Báo giá thuộc công ty đang làm việc + báo giá do mình lập.'),
    ('V3', 'Xem danh sách Báo giá theo phòng ban', 'Báo giá thuộc phòng ban / bộ phận mình quản lý + báo giá do '
     'mình lập.'),
    ('V4', 'Xem danh sách Báo giá theo bộ phận', 'Báo giá thuộc bộ phận mình quản lý + báo giá do mình lập.'),
    ('—', 'Không có V1–V4', 'Chỉ báo giá do mình lập.'),
], widths=[0.8, 2.4, 2.8])

d.h2('2 Ma trận phân quyền')
Y, N = '✅', '❌'
d.table(['Chức năng', 'Q1', 'Q2', 'Q3', 'Q4', 'Q5', 'Q6', 'Không có quyền nào'], [
    ('FR-01 Xem danh sách báo giá', Y, Y, Y, Y, Y, Y, Y + ' (theo V1–V4)'),
    ('FR-02 Tìm kiếm và lọc', Y, Y, Y, Y, Y, Y, Y),
    ('FR-03 Cài đặt bộ lọc', Y, Y, Y, Y, Y, Y, Y),
    ('FR-04 Tùy chỉnh cột', Y, Y, Y, Y, Y, Y, Y),
    ('FR-05 Xuất Excel danh sách', Y, Y, Y, Y, Y, Y, Y),
    ('FR-06 Tạo mới báo giá', Y + ' (từ YCBG)', Y + ' (từ YCBG)', N, N, N, N, Y + ' (Sale phụ trách, tự lập)'),
    ('FR-07 Làm giá hàng hoá, dịch vụ', Y, Y, N, N, Y, N, Y + ' (người lập)'),
    ('FR-08 Thêm hàng hoá, hàng tạm, dịch vụ', Y, Y, N, N, Y, N, Y + ' (người lập)'),
    ('FR-09 Áp dụng giảm giá', N, N, N, N, N, Y + ' (người lập)', N),
    ('FR-10 Tổng hợp giá trị và cấp duyệt dự kiến', Y, Y, Y, Y, Y + ' (thấy cấp duyệt)', Y, Y),
    ('FR-11 Điều khoản báo giá, ghi chú nội bộ', Y, Y, N, N, N, N, Y + ' (người lập)'),
    ('FR-12 Import Excel', Y, Y, N, N, N, N, Y + ' (người lập)'),
    ('FR-13 Lưu nháp', Y, Y, N, N, N, N, Y + ' (người lập)'),
    ('FR-14 Gửi duyệt', Y, Y, N, N, N, N, Y + ' (người lập)'),
    ('FR-15 Chỉnh sửa báo giá', Y, Y, N, N, N, N, Y + ' (người lập, Đang tạo)'),
    ('FR-16 Xem chi tiết', Y, Y, Y, Y, Y, Y, Y + ' (theo V1–V4)'),
    ('FR-17 Duyệt báo giá', N, N, Y, Y, N, N, N),
    ('FR-18 Từ chối báo giá', N, N, Y, Y, N, N, N),
    ('FR-19 Xóa báo giá', N, N, N, N, N, N, Y + ' (người lập, Đang tạo)'),
    ('FR-20 Sao chép báo giá', Y + ' (BG từ YCBG)', Y + ' (BG từ YCBG)', N, N, N, N, Y + ' (Sale phụ trách)'),
    ('FR-21 In báo giá', Y, Y, Y, Y, Y, Y, Y),
    ('FR-22 Xuất Excel báo giá', Y, Y, Y, Y, Y, Y, Y),
    ('FR-23 Ghi chú kinh doanh', N, N, N, N, N, N, Y + ' (Sale phụ trách, Đã duyệt)'),
    ('FR-24 Lập hợp đồng', N, N, N, N, N, N, Y + ' (người lập, Trúng thầu)'),
    ('FR-25 Xem lịch sử', Y, Y, Y, Y, Y, Y, Y),
], widths=[2.0, 0.55, 0.55, 0.5, 0.5, 0.55, 0.55, 0.8])
d.p('Ghi chú: các ô "Không có quyền nào" mang điều kiện vai trò (người lập, Sale phụ trách) áp cho mọi người dùng '
    'thỏa vai trò đó, kể cả có thêm Q1–Q6.')

# ======================================================================= PHẦN 3
d.h1('Phần 3. Đặc tả chi tiết theo từng chức năng')
d.h2('1 Sơ đồ UML tổng quan')
MAINS = [
    ('FR-01', 'Xem danh sách báo giá', 'view'),        # 0
    ('FR-05', 'Xuất Excel danh sách', 'io'),           # 1
    ('FR-06', 'Tạo mới báo giá', 'crud'),              # 2
    ('FR-07', 'Làm giá hàng hoá, dịch vụ', 'crud'),    # 3
    ('FR-08', 'Thêm hàng hoá, hàng tạm, dịch vụ', 'crud'),  # 4
    ('FR-09', 'Áp dụng giảm giá', 'crud'),             # 5
    ('FR-10', 'Tổng hợp giá trị, cấp duyệt', 'crud'),  # 6
    ('FR-11', 'Điều khoản báo giá', 'crud'),           # 7
    ('FR-12', 'Import Excel', 'io'),                   # 8
    ('FR-13', 'Lưu nháp', 'crud'),                     # 9
    ('FR-14', 'Gửi duyệt', 'action'),                  # 10
    ('FR-15', 'Chỉnh sửa báo giá', 'crud'),            # 11
    ('FR-17', 'Duyệt báo giá', 'action'),              # 12
    ('FR-18', 'Từ chối báo giá', 'action'),            # 13
    ('FR-19', 'Xóa báo giá', 'action'),                # 14
    ('FR-20', 'Sao chép báo giá', 'crud'),             # 15
    ('FR-21', 'In báo giá', 'io'),                     # 16
    ('FR-22', 'Xuất Excel báo giá', 'io'),             # 17
    ('FR-23', 'Ghi chú kinh doanh', 'crud'),           # 18
    ('FR-24', 'Lập hợp đồng', 'action'),               # 19
]
d.overview_figure2(
    [(A_SALE, [0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11, 14, 15, 16, 17, 18, 19]),
     (A_DUYET, [0, 12, 13, 16, 17])],
    MAINS,
    [('FR-02', 'Tìm kiếm và lọc', 'view', 'extend', [0], None),
     ('FR-03', 'Cài đặt bộ lọc', 'view', 'extend', [0], None),
     ('FR-04', 'Tùy chỉnh cột', 'view', 'extend', [0], None),
     ('FR-16', 'Xem chi tiết', 'view', 'extend', [0], None),
     ('FR-25', 'Xem lịch sử', 'view', 'extend', [0], None)],
    'Sơ đồ Use Case tổng quan màn %s' % TEN_MAN)

d.h2('2 Đặc tả chi tiết từng chức năng')

# --------------------------------------------------------------------- FR-01
fr('FR-01', 'Xem danh sách báo giá', 'view', A_ALL, 'list',
   '- Màn Danh sách, Phân trang, Sắp xếp và UI/UX. Chỉ bổ sung các quy tắc riêng của màn Báo giá tại phần mô '
   'tả chi tiết.',
   dict(mota='Hiển thị danh sách báo giá (cả báo giá thường và báo giá tổng) trong phạm vi được xem, mới tạo ở '
        'trên, 20 dòng / trang.',
        tacnhan='Người lập báo giá, Người duyệt giá; Người dùng đã đăng nhập',
        dieukien='Người dùng đã đăng nhập.',
        chinh='1. Người dùng vào menu Báo giá (CSKH trước bán) hoặc Danh sách báo giá (Bán hàng).\n'
              '2. Hệ thống khôi phục bộ lọc đã dùng gần nhất (lưu 10 phút) rồi nạp danh sách theo phạm vi xem.\n'
              '3. Hệ thống hiển thị bảng báo giá kèm cột Hành động theo từng dòng.',
        phu='• Không có dòng nào khớp → "Không có dữ liệu phù hợp bộ lọc."\n'
            '• Lỗi tải → thông báo "Lỗi tải dữ liệu".\n'
            '• Bấm mã báo giá thường → mở Chi tiết báo giá; mã báo giá tổng → mở màn báo giá tổng.'),
   ('', [('01-danh-sach.png', 'Màn danh sách báo giá (phần đầu bảng)'),
         ('01d-menu-them.png', 'Danh sách — phần cuối bảng, cột Trạng thái, Hành động và menu ⋮')]),
   [('Tiêu đề bảng', 'Label', 'Hiển thị', '–', 'Danh sách báo giá', '–'),
    ('Nút Tạo báo giá', 'Button', 'Enable', '–', 'Hiển thị', 'Mở màn Tạo báo giá (FR-06).'),
    ('Nút Xuất Excel', 'Button', 'Enable / Disable', '–', 'Hiển thị', 'Khoá trong lúc đang xuất (FR-05).'),
    ('Nút Cấu hình cột hiển thị', 'Icon Button', 'Enable', '–', 'Hiển thị', 'FR-04.'),
    ('STT', 'Text', 'Read-only', '≥ 1', 'Theo trang', 'Cột ghim trái, khoá.'),
    ('Mã báo giá', 'Link', 'Read-only', '–', 'Theo dữ liệu', 'Cột ghim trái, khoá, sắp xếp được; bấm mở chi tiết.'),
    ('Loại báo giá', 'Badge', 'Read-only', '3 giá trị', 'Theo dữ liệu', '"Từ BOM" / "Tự nhập" / "Báo giá tổng".'),
    ('BOM list', 'Text', 'Read-only', '–', 'Theo dữ liệu', '"Mã BOM - Tên BOM", tối đa 2 dòng.'),
    ('Mã YCBG', 'Text', 'Read-only', '–', 'Theo dữ liệu', 'Mã yêu cầu xây dựng giá (nếu lập từ YCBG).'),
    ('Dự án TKT', 'Text', 'Read-only', '–', 'Theo dữ liệu', '"Mã dự án - Tên dự án".'),
    ('Khách hàng', 'Text', 'Read-only', '–', 'Theo dữ liệu', 'Tên khách hàng lưu trên báo giá.'),
    ('Giai đoạn dự án', 'Text', 'Read-only', '–', 'Theo dữ liệu', '–'),
    ('Tiền tệ', 'Text', 'Read-only', '–', 'Theo dữ liệu', 'Mã loại tiền.'),
    ('Tổng giá trị sau VAT', 'Number', 'Read-only', '≥ 0', 'Theo dữ liệu', 'Căn phải, định dạng 1,234,567; sắp '
     'xếp được.'),
    ('Cấp duyệt', 'Badge', 'Read-only', '3 giá trị', 'Theo dữ liệu', 'Nhãn và màu do hệ thống trả về.'),
    ('Người duyệt / Ngày duyệt', 'Text', 'Read-only', 'dd/mm/yyyy hh:mm', 'Theo dữ liệu', 'Ngày duyệt sắp xếp '
     'được.'),
    ('Đồng bộ ERP', 'Badge', 'Read-only', '2 giá trị', 'Trống', '"Đã đồng bộ" (xanh) / "Thất bại" (đỏ, rê '
     'chuột xem lỗi).'),
    ('Người tạo / Ngày tạo / Người cập nhật / Ngày cập nhật', 'Text', 'Read-only', 'dd/mm/yyyy hh:mm',
     'Theo dữ liệu', 'Ngày tạo, ngày cập nhật sắp xếp được.'),
    ('Trạng thái', 'Badge', 'Read-only', '7 giá trị', 'Theo dữ liệu', 'Đang tạo · Chờ TP duyệt · Chờ BGĐ duyệt · '
     'Đã duyệt · Đóng · Dừng · Trúng thầu. Chữ và màu do hệ thống trả về.'),
    ('Hành động', 'Icon Button', 'Enable / Ẩn', '–', 'Theo dòng', 'Cột cuối, khoá. 2 nút đầu hiện trực tiếp, còn '
     'lại vào menu ⋮: Sửa / Làm giá, Xóa (chỉ người lập + Đang tạo), Sao chép báo giá (theo FR-20), In báo giá, '
     'Lịch sử phê duyệt.'),
    ('Phân trang', 'Pagination', 'Enable', '–', '20 dòng / trang', 'Đổi trang, đổi số dòng / trang.'),
    ('Trạng thái rỗng', 'Label', 'Hiển thị', '–', 'Ẩn', '"Không có dữ liệu phù hợp bộ lọc."'),
   ],
   [('Mở màn hình', 'System',
     'Before:\n– Đọc 4 quyền V1–V4 để xác định phạm vi (xem quy tắc BR-01).\n'
     'After:\n– Khôi phục bộ lọc đã lưu; nạp danh sách trang 1, sắp xếp mặc định theo Ngày tạo giảm dần; nạp cấu '
     'hình cột và danh mục bộ lọc ở nền.\n– Lỗi → "Lỗi tải dữ liệu".'),
    ('Bấm tiêu đề cột có sắp xếp', 'Click', 'After:\n– Đổi chiều sắp xếp, quay về trang 1 và nạp lại.'),
    ('Đổi trang / số dòng', 'Click', 'After:\n– Nạp lại đúng trang / số dòng đã chọn.'),
    ('Bấm Mã báo giá', 'Click', 'After:\n– Mở màn Chi tiết báo giá (FR-16); báo giá tổng mở màn báo giá tổng.'),
    ('Bấm ⋮ trên dòng', 'Click', 'After:\n– Mở menu các hành động còn lại của dòng.'),
   ], mode='ro', usecase=False)

# --------------------------------------------------------------------- FR-02
fr('FR-02', 'Tìm kiếm và lọc', 'view', A_ALL, 'search',
   '- Kịch bản tìm kiếm, Bộ lọc, Dropdown. Chỉ bổ sung các quy tắc riêng của màn Báo giá tại phần mô tả chi tiết.',
   dict(mota='Tìm nhanh theo mã báo giá / tên khách hàng và lọc nâng cao theo 14 điều kiện.',
        tacnhan='Người lập báo giá, Người duyệt giá; Người dùng đã đăng nhập',
        dieukien='Người dùng đang ở màn danh sách báo giá.',
        chinh='1. Người dùng nhập từ khoá vào ô tìm nhanh rồi bấm Tìm kiếm.\n'
              '2. Hoặc bấm "Tìm kiếm nâng cao" để mở khối bộ lọc, chọn điều kiện.\n'
              '3. Chọn giá trị ở ô chọn → danh sách tự nạp lại; ô gõ tay chỉ tìm khi bấm Tìm kiếm.\n'
              '4. Bấm "Làm mới" để xoá toàn bộ điều kiện.',
        phu='• Chưa chọn Dự án TKT → ô Giải pháp bị khoá; chưa chọn Giải pháp → Version giải pháp rỗng.\n'
            '• Đổi Dự án → xoá Giải pháp và Version đã chọn.\n'
            '• Ô Công ty / Phòng ban / Bộ phận khoá theo phạm vi quyền V1–V4.'),
   (' => Tìm kiếm nâng cao', [('02-loc.png', 'Khối Tìm kiếm nâng cao đang mở')]),
   [('Ô tìm nhanh', 'Textbox', 'Enable', '0–255 ký tự', 'Không', 'Trống',
     'Gợi ý "Tìm theo mã báo giá, tên khách hàng"; tìm chứa chuỗi.'),
    ('Nút Tìm kiếm', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Nạp lại từ trang 1.'),
    ('Nút Làm mới', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Xoá mọi điều kiện, nạp lại 1 lần.'),
    ('Nút Tìm kiếm nâng cao', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Mở / đóng khối lọc.'),
    ('Công ty / Phòng ban / Bộ phận / Người lập', 'Dropdown', 'Enable / Disable', 'Danh sách', 'Không', 'Trống',
     'Khoá theo phạm vi quyền (biểu tượng ổ khoá). Người lập = người tạo báo giá.'),
    ('Dự án TKT', 'Dropdown', 'Enable', 'Tìm từ xa, 20 kết quả', 'Không', 'Trống', '"Mã - Tên dự án".'),
    ('Giải pháp', 'Dropdown', 'Enable / Disable', 'Tìm từ xa', 'Không', 'Trống', 'Khoá khi chưa chọn Dự án TKT.'),
    ('Version giải pháp', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống', 'Theo giải pháp đã chọn.'),
    ('Khách hàng', 'Dropdown', 'Enable', 'Tìm từ xa', 'Không', 'Trống', '–'),
    ('Trạng thái', 'Dropdown', 'Enable', 'Danh sách 7 giá trị', 'Không', 'Trống', 'Như cột Trạng thái.'),
    ('Loại báo giá', 'Dropdown', 'Enable', 'Danh sách 2 giá trị', 'Không', 'Trống',
     '"Báo giá thường" / "Báo giá tổng"; trống = cả hai.'),
    ('Cấp duyệt', 'Dropdown', 'Enable', 'Danh sách 3 giá trị', 'Không', 'Trống',
     '"Cấp 1 — Tự duyệt" / "Cấp 2 — TP duyệt" / "Cấp 3 — BGĐ duyệt".'),
    ('Người duyệt', 'Dropdown', 'Enable', 'Danh sách nhân viên', 'Không', 'Trống', 'Người duyệt cuối cùng.'),
    ('Giai đoạn dự án', 'Dropdown', 'Enable', 'Danh sách', 'Không', 'Trống',
     'Kèm biểu tượng ⓘ mô tả giai đoạn; giai đoạn đã khoá vẫn hiện nếu đang lọc.'),
    ('Ngày tạo', 'Datepicker', 'Enable', 'dd/mm/yyyy, từ – đến', 'Không', 'Trống', 'Một ô khoảng ngày.'),
   ],
   [('Bấm Tìm kiếm / Enter ở ô tìm nhanh', 'Click / Keypress',
     'After:\n– Về trang 1, nạp danh sách theo từ khoá + điều kiện đang chọn, giữ phạm vi BR-01.'),
    ('Chọn giá trị ở ô chọn', 'Change', 'After:\n– Tự nạp lại danh sách từ trang 1.'),
    ('Chọn Dự án TKT', 'Change', 'After:\n– Xoá Giải pháp, Version giải pháp; mở khoá ô Giải pháp.'),
    ('Chọn Giải pháp', 'Change', 'After:\n– Nạp danh sách Version của giải pháp; xoá Version đang chọn.'),
    ('Bấm Làm mới', 'Click', 'After:\n– Xoá toàn bộ điều kiện lọc, nạp lại đúng 1 lần.'),
   ], usecase=False)

# --------------------------------------------------------------------- FR-03
fr('FR-03', 'Cài đặt bộ lọc', 'view', A_ALL, 'search',
   '- Bộ lọc, Cấu hình hiển thị. Chỉ bổ sung các quy tắc riêng của màn Báo giá tại phần mô tả chi tiết.',
   dict(mota='Cho người dùng chọn ô lọc nào được hiển thị và thứ tự các ô trong khối Tìm kiếm nâng cao.',
        tacnhan='Người dùng đã đăng nhập',
        dieukien='Người dùng đang ở màn danh sách báo giá.',
        chinh='1. Người dùng bấm "Cài đặt bộ lọc".\n'
              '2. Tích / bỏ tích từng trường lọc, kéo biểu tượng ⋮⋮ để đổi thứ tự.\n'
              '3. Bấm Lưu → hệ thống lưu cấu hình cho người dùng và báo "Cập nhật thành công".',
        phu='• Bấm "Khôi phục mặc định" → trở về đủ 11 trường theo thứ tự mặc định (chưa ghi, phải bấm Lưu).\n'
            '• Lưu lỗi → "Thao tác thất bại".'),
   (' => Tìm kiếm nâng cao => Cài đặt bộ lọc', [('03-cai-dat-loc.png', 'Cửa sổ Cài đặt bộ lọc')]),
   [('Tiêu đề cửa sổ', 'Label', 'Hiển thị', '–', '–', 'Cài đặt bộ lọc', '–'),
    ('Danh sách trường lọc', 'Checkbox', 'Enable', '11 trường', 'Không', 'Theo cấu hình đã lưu',
     'Công ty – Phòng ban – Bộ phận – Người lập · Dự án TKT · Giải pháp · Version giải pháp · Khách hàng · Trạng '
     'thái · Loại báo giá · Cấp duyệt · Người duyệt · Giai đoạn dự án · Ngày tạo.'),
    ('Biểu tượng kéo thả', 'Icon', 'Enable', '–', '–', 'Hiển thị', 'Kéo để đổi thứ tự hiển thị.'),
    ('Nút Lưu', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Lưu cấu hình theo người dùng.'),
    ('Nút Khôi phục mặc định', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Nút Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Đóng không lưu.'),
   ],
   [('Bấm Lưu', 'Click', 'After:\n– Lưu cấu hình bộ lọc cho người dùng ở màn Báo giá; đóng cửa sổ; hiển thị '
     '"Cập nhật thành công".\n– Lỗi → "Thao tác thất bại".'),
    ('Bấm Khôi phục mặc định', 'Click', 'After:\n– Đưa danh sách về mặc định trên cửa sổ, chưa ghi.'),
    ('Bấm Đóng / ×', 'Click', 'After:\n– Đóng cửa sổ, không lưu.'),
   ], usecase=False)

# --------------------------------------------------------------------- FR-04
fr('FR-04', 'Tùy chỉnh cột', 'view', A_ALL, 'excel',
   '- Quy tắc Excel và Cấu hình cột. Chỉ bổ sung các quy tắc riêng của màn Báo giá tại phần mô tả chi tiết.',
   dict(mota='Cho người dùng ẩn / hiện và sắp xếp lại thứ tự các cột của bảng danh sách báo giá.',
        tacnhan='Người dùng đã đăng nhập',
        dieukien='Người dùng đang ở màn danh sách báo giá.',
        chinh='1. Người dùng bấm biểu tượng Cấu hình cột hiển thị.\n'
              '2. Tích / bỏ tích cột, kéo để đổi thứ tự.\n'
              '3. Bấm Lưu → bảng áp dụng ngay, hệ thống báo "Cập nhật thành công".',
        phu='• Cột STT, Mã báo giá, Hành động luôn hiển thị, không bỏ tích / kéo được (biểu tượng ổ khoá).\n'
            '• Lưu lỗi → "Thao tác thất bại".'),
   (' => Cấu hình cột hiển thị', [('04-tuy-chinh-cot.png', 'Cửa sổ Tuỳ chỉnh cột')]),
   [('Tiêu đề cửa sổ', 'Label', 'Hiển thị', '–', '–', 'Tuỳ chỉnh cột', '–'),
    ('Danh sách cột', 'Checkbox', 'Enable', '20 cột', 'Không', 'Mặc định hiện hết',
     'STT, Mã báo giá, Loại báo giá, BOM list, Mã YCBG, Dự án TKT, Khách hàng, Giai đoạn dự án, Tiền tệ, Tổng giá '
     'trị sau VAT, Cấp duyệt, Người duyệt, Ngày duyệt, Đồng bộ ERP, Người tạo, Ngày tạo, Người cập nhật, Ngày cập '
     'nhật, Trạng thái, Hành động.'),
    ('Cột khoá', 'Checkbox', 'Disable', '3 cột', '–', 'Đã tích', 'STT, Mã báo giá, Hành động.'),
    ('Biểu tượng kéo thả', 'Icon', 'Enable', '–', '–', 'Hiển thị', 'Không áp dụng cho cột khoá.'),
    ('Nút Lưu', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Lưu cấu hình cột theo người dùng.'),
    ('Nút Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
   ],
   [('Bấm Lưu', 'Click', 'After:\n– Áp cột mới cho bảng, lưu cấu hình theo người dùng; hiển thị "Cập nhật thành '
     'công".\n– Lỗi → "Thao tác thất bại".'),
    ('Kéo thả cột', 'Drag', 'After:\n– Đổi thứ tự trên cửa sổ, chưa ghi.'),
   ], usecase=False)

# --------------------------------------------------------------------- FR-05
fr('FR-05', 'Xuất Excel danh sách', 'io', A_ALL, 'excel',
   '- Quy tắc Excel và Cấu hình cột. Chỉ bổ sung các quy tắc riêng của màn Báo giá tại phần mô tả chi tiết.',
   dict(mota='Xuất toàn bộ báo giá đang thỏa bộ lọc (không giới hạn theo trang) ra file Excel theo các trường được '
        'chọn.',
        tacnhan='Người lập báo giá, Người duyệt giá; Người dùng đã đăng nhập',
        dieukien='Người dùng đang ở màn danh sách báo giá.',
        chinh='1. Người dùng bấm "Xuất Excel".\n'
              '2. Cửa sổ "Chọn trường xuất file" mở, mặc định tích các cột đang hiện trên bảng.\n'
              '3. Người dùng tích / bỏ tích, kéo để đổi thứ tự cột trong file, bấm "Xuất file".\n'
              '4. Hệ thống tải file danh_sach_bao_gia.xlsx và báo "Xuất Excel thành công".',
        phu='• Không chọn trường nào → nút Xuất file không thao tác được.\n'
            '• Lỗi xuất → "Lỗi khi xuất Excel".'),
   (' => Xuất Excel', [('05-xuat.png', 'Cửa sổ Chọn trường xuất file')]),
   [('Tiêu đề cửa sổ', 'Label', 'Hiển thị', '–', '–', 'Chọn trường xuất file', '–'),
    ('Danh sách trường', 'Checkbox', 'Enable', '20 trường', 'Có (≥ 1)', 'Theo cột đang hiện',
     'Mã báo giá, Tên báo giá, Loại báo giá, Mã YCBG, Dự án TKT, Mã BOM, Khách hàng, Version giải pháp, Giai đoạn '
     'dự án, Tiền tệ, Tổng giá trị sau VAT, Cấp duyệt, Người duyệt, Ngày duyệt, Trạng thái, Trạng thái đồng bộ ERP, '
     'Người tạo, Ngày tạo, Người cập nhật, Ngày cập nhật.'),
    ('Nút Chọn tất cả / Bỏ chọn hết', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Dòng đếm', 'Label', 'Hiển thị', '–', '–', 'Đang chọn n/20 trường', '–'),
    ('Nút Xuất file', 'Button', 'Enable / Disable', '–', '–', 'Hiển thị', 'Khoá khi đang xuất.'),
    ('Nút Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
   ],
   [('Bấm Xuất Excel', 'Click', 'After:\n– Mở cửa sổ Chọn trường xuất file.'),
    ('Bấm Xuất file', 'Click',
     'Before:\n– Áp đúng phạm vi xem BR-01 và bộ lọc đang chọn, lấy tất cả dòng.\n'
     'After:\n– Tải file danh_sach_bao_gia.xlsx, tiêu đề "Danh sách báo giá", cột theo đúng thứ tự đã chọn; hiển '
     'thị "Xuất Excel thành công".\n– Lỗi → "Lỗi khi xuất Excel".'),
   ])

# --------------------------------------------------------------------- FR-06
fr('FR-06', 'Tạo mới báo giá', 'crud', A_SALE, 'create',
   '- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Lập báo giá cho 1 dự án TKT: chọn dự án → hệ thống điền thông tin khách hàng, tiền tệ, giai đoạn; '
        'nếu dự án có BOM tổng hợp đã duyệt thì nạp hàng hoá từ BOM, ngược lại là báo giá tự lập.',
        tacnhan='Sale phụ trách dự án / Người có quyền xây dựng giá (Q1, Q2)',
        dieukien='Người dùng là Sale phụ trách của ít nhất 1 dự án TKT (danh sách Dự án chỉ gồm dự án mình phụ '
                 'trách, không gồm dự án cha).',
        chinh='1. Người dùng bấm "Tạo báo giá"; màn Tạo báo giá mở với Mã báo giá "(Chưa tạo)".\n'
              '2. Chọn Dự án → hệ thống nạp Khách hàng, MST, Địa chỉ, Người liên hệ, SĐT, Email, Loại tiền tệ, '
              'Giai đoạn dự án hiện tại và Hiệu lực báo giá.\n'
              '3. Hệ thống tìm BOM tổng hợp đã duyệt của dự án: có đúng 1 BOM → tự chọn và nạp toàn bộ hàng hoá, '
              'dịch vụ; nhiều BOM → người dùng chọn ở ô "BOM tổng hợp"; không có BOM → báo giá tự lập.\n'
              '4. Người dùng nhập Giao hàng (ngày), Bảo hành, chọn Bảng giá, làm giá (FR-07) và điều khoản (FR-11).\n'
              '5. Bấm "Lưu nháp" (FR-13) hoặc "Gửi duyệt" (FR-14).',
        phu='• Dự án có BOM mà chưa chọn tiền tệ → "Dự án này có BOM. Vui lòng chọn loại tiền tệ trước để tính tỷ '
            'giá quy đổi khi tải sản phẩm từ BOM."\n'
            '• Chọn lại Dự án khi đã có hàng → hỏi "Việc chọn lại Dự án sẽ xóa toàn bộ thông tin hàng hoá/dịch vụ '
            'trên báo giá!".\n'
            '• Chọn lại BOM khi đã có hàng → hỏi "Việc chọn lại BOM sẽ xóa toàn bộ thông tin hàng hoá/dịch vụ trên '
            'báo giá!".\n'
            '• Dự án con → Loại tiền tệ, Bảng giá, phương thức giảm giá kế thừa dự án cha và bị khoá; nếu dự án cha '
            'có Báo giá Mỏ neo thì hiện dải nhắc đồng bộ.',
        dacbiet='Báo giá chỉ được ghi xuống hệ thống khi bấm Lưu nháp hoặc Xác nhận ở popup Gửi duyệt; rời màn '
                'trước đó thì không có bản ghi nào được tạo.'),
   (' => Tạo báo giá', [('06-tao-moi-trong.png', 'Màn Tạo báo giá lúc mới mở'),
                        ('06b-tao-moi-bom.png', 'Chọn dự án có 1 BOM tổng hợp đã duyệt — hệ thống tự nạp hàng từ BOM')]),
   [('Mã báo giá', 'Text', 'Read-only', '–', '–', '(Chưa tạo)', 'Cấp khi lưu: BG-YYYY-NNNNN.'),
    ('Dự án', 'Dropdown', 'Enable', 'Dự án mình phụ trách', 'Có', 'Trống',
     'Gợi ý "Tìm dự án..."; hiển thị "Mã — Tên dự án". Thiếu khi lưu → "Vui lòng chọn dự án trước khi tạo báo '
     'giá".'),
    ('BOM tổng hợp', 'Dropdown', 'Enable / Ẩn', 'BOM tổng hợp đã duyệt của dự án', 'Không', 'Tự chọn khi có 1 BOM',
     'Chưa chọn dự án → "Chọn dự án trước"; không có BOM → "Không có BOM tổng hợp đã duyệt".'),
    ('Khách hàng, MST, Địa chỉ, Người liên hệ, SĐT liên hệ', 'Text', 'Read-only', '–', '–', 'Theo dự án',
     'Lấy từ dự án TKT, lưu kèm báo giá.'),
    ('Email khách hàng', 'Textbox', 'Enable', 'Email, ≤ 255 ký tự', 'Không', 'Theo dự án',
     'Gợi ý "Nhập email khách hàng"; sửa tay được.'),
    ('Giao hàng (ngày)', 'Number', 'Enable', '> 0', 'Có khi gửi duyệt', 'Trống',
     'Thiếu khi gửi duyệt → "Vui lòng nhập thời gian giao hàng".'),
    ('Hiệu lực báo giá', 'Text', 'Read-only', 'dd/mm/yyyy', '–', 'Hôm nay + số ngày hiệu lực',
     'Rút về ngày đổi giá sớm nhất của hàng ERP (BR-15). Khác ngày đã lưu → hiện "(đã lưu: dd/mm/yyyy)".'),
    ('Loại tiền tệ', 'Dropdown', 'Enable / Disable', 'Danh sách tiền tệ', 'Có khi gửi duyệt', 'Theo dự án',
     'Khoá ở màn Sửa và với dự án con. Thiếu → "Vui lòng chọn loại tiền tệ".'),
    ('Bảo hành (tháng)', 'Number', 'Enable', '≥ 0', 'Không', 'Trống', '–'),
    ('Bảng giá', 'Dropdown', 'Enable / Disable', 'Danh sách 6 bảng giá', 'Có', 'Bán lẻ',
     'Khoá ở màn Sửa và với dự án con. Thiếu → "Vui lòng chọn bảng giá".'),
    ('Giai đoạn dự án', 'Dropdown', 'Enable', 'Danh sách', 'Có khi gửi duyệt', 'Giai đoạn hiện tại của dự án',
     'Kèm ⓘ mô tả giai đoạn. Thiếu → "Vui lòng chọn giai đoạn dự án".'),
    ('Thanh công cụ VAT / GG / Làm tròn', 'Toolbar', 'Enable', '–', '–', 'Hiển thị', 'Xem FR-07, FR-09.'),
    ('Bảng Chi tiết sản phẩm', 'Table/Grid', 'Enable', '–', '–', 'Trống / Theo BOM', 'Xem FR-07, FR-08.'),
    ('Khối Tổng hợp giá trị báo giá', 'Table/Grid', 'Read-only', '–', '–', '0', 'Xem FR-10.'),
    ('Khối Thanh toán & Ghi chú nội bộ', 'Section', 'Enable', '–', '–', 'Trống', 'Xem FR-11.'),
    ('Thanh chân trang', 'Toolbar', 'Hiển thị', '–', '–', 'Hiển thị',
     'Nhập · Bán · (GG · Sau GG) · ~VND · LN · Cấp duyệt; nút Quay lại, Lưu nháp, Gửi duyệt.'),
   ],
   [('Mở màn Tạo báo giá', 'System', 'After:\n– Khởi tạo báo giá trống trạng thái Đang tạo; nạp nền danh sách '
     'dự án mình phụ trách, tiền tệ, bảng giá, ĐVT, cấu hình duyệt giá, mẫu điều khoản.'),
    ('Chọn Dự án', 'Change',
     'Before:\n– Đã có hàng hoá / dịch vụ → hỏi xác nhận "Xác nhận thay đổi Dự án"; Huỷ → giữ dự án cũ.\n'
     'After:\n– Nạp thông tin khách hàng, tiền tệ, giai đoạn; dự án con → khoá tiền tệ / bảng giá / phương thức GG, '
     'kiểm tra Báo giá Mỏ neo.\n– Tìm BOM tổng hợp đã duyệt: 1 BOM → nạp hàng; nhiều → chờ chọn; 0 → báo giá tự lập.'
     '\n– Lỗi → "Không tải được thông tin dự án".'),
    ('Chọn BOM tổng hợp', 'Change',
     'Before:\n– Chưa có tiền tệ → "Vui lòng chọn loại tiền tệ trước để tính tỷ giá quy đổi khi tải sản phẩm từ '
     'BOM.", dừng.\n– Đã có hàng → hỏi "Xác nhận thay đổi BOM".\n'
     'After:\n– Xoá toàn bộ hàng cũ, nạp nhóm, hàng hoá (giá nhập / giá bán ERP quy theo tỷ giá, VAT ERP; hàng '
     'tạm giá = 0) và dịch vụ (giá = 0) từ BOM.\n– Lỗi → "Không tải được sản phẩm BOM".'),
    ('Đổi Loại tiền tệ', 'Change', 'Before:\n– Đã có giá > 0 → hỏi "Các giá trị tiền tệ đã nhập sẽ không thay đổi '
     'khi bạn đổi loại tiền tệ ! Bạn có chắc chắc muốn đổi?".\nAfter:\n– VNĐ / tiền không phần lẻ → làm tròn Số '
     'nguyên (0); tiền có phần lẻ → Mặc định 2 số lẻ.'),
    ('Đổi Bảng giá', 'Change', 'Before:\n– Có hàng ERP → hỏi "Đổi bảng giá sẽ cập nhật lại đơn giá bán của các '
     'hàng hoá ERP theo bảng giá mới. Chiết khấu và hàng tạm giữ nguyên. Bạn có chắc muốn đổi?".\n'
     'After:\n– Tính lại đơn giá bán + VAT hàng ERP theo bảng giá mới.'),
    ('Bấm Quay lại', 'Click', 'After:\n– Về danh sách báo giá.'),
   ])

# --------------------------------------------------------------------- FR-07
fr('FR-07', 'Làm giá hàng hoá, dịch vụ', 'crud', A_SALE, 'create',
   '- Validate dữ liệu, Thông báo và UI/UX. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Nhập / điều chỉnh giá trên bảng Chi tiết sản phẩm gồm phần A — Hàng hoá (theo nhóm, cây cha – con) '
        'và phần B — Dịch vụ & Chi phí khác; hệ thống tính thành tiền, TSLN, VAT theo từng dòng.',
        tacnhan='Người lập báo giá',
        dieukien='Báo giá đang ở màn Tạo hoặc màn Sửa (người lập, trạng thái Đang tạo).',
        chinh='1. Người dùng nhập Giá nhập (hàng tạm), Giá bán, SL (báo giá tự lập), chọn ĐVT (hàng ERP có nhiều '
              'ĐVT), nhập VAT (hàng tạm, dịch vụ) và Ghi chú.\n'
              '2. Hệ thống tính lại Thành tiền nhập, Thành tiền bán, TSLN, Tiền VAT, Thành tiền sau VAT, tổng phần A/B '
              'và dòng TỔNG.\n'
              '3. Có thể áp VAT đồng loạt cho hàng tạm (ô VAT + "Tất cả" / "VAT=0").',
        phu='• Lần đầu nhập VAT > 0 cho 1 dòng trong khi còn dòng 0% → popup "Áp dụng VAT đồng loạt?" với 2 lựa chọn '
            '"Chỉ dòng này" / "Áp dụng tất cả dòng còn 0%".\n'
            '• Hàng cha có hàng con: Giá nhập cha tự tính từ hàng con, khoá.\n'
            '• Nút "Ẩn cột chi tiết" / "Hiện cột chi tiết" ẩn hiện Model, Thương hiệu, Xuất xứ, Thông số kỹ thuật, '
            'Ghi chú; "Thu gọn tất cả" / "Mở rộng tất cả" đóng mở các nhóm.\n'
            '• Làm tròn (tiền khác VNĐ): chọn mức rồi bấm "Áp dụng" → hỏi "Bạn có chắc muốn làm tròn đến … cho toàn '
            'bộ đơn giá? Thao tác này không thể hoàn tác." VNĐ khoá ở Số nguyên (0).',
        dacbiet='Không tự sửa số người dùng nhập: giá trị ngoài khoảng báo đỏ ngay dưới ô (BR-08).'),
   (' => Sửa / Làm giá', [('07b-bang-gia.png', 'Bảng Chi tiết sản phẩm — cột giá (đã ẩn cột chi tiết)'),
                          ('07c-bang-gia-phai.png', 'Bảng Chi tiết sản phẩm — cột TSLN, VAT, Thành tiền sau VAT')],
    'Hoặc ở màn Tạo báo giá (FR-06), phần Chi tiết sản phẩm.'),
   [('STT', 'Text', 'Read-only', '–', '–', 'Tự đánh', 'Cha "1", con "1.1"; nhóm đánh số La Mã.'),
    ('Mã', 'Text', 'Read-only', '–', '–', 'Theo dữ liệu', 'Hàng tạm chưa lưu để trống; ⚠ "Giá bán thay đổi ngày '
     'dd/mm/yyyy" khi hàng ERP sắp đổi giá.'),
    ('Tên hàng', 'Text', 'Read-only', '–', '–', 'Theo dữ liệu', 'Badge "DV" nếu là dịch vụ; dòng "Ghi chú nội bộ" '
     'của hàng ERP.'),
    ('Model / Thương hiệu / Xuất xứ / Thông số kỹ thuật', 'Text', 'Read-only', '–', '–', 'Theo dữ liệu',
     'Ẩn khi bấm "Ẩn cột chi tiết".'),
    ('Ghi chú', 'Textbox', 'Enable', '–', 'Không', 'Trống', 'Gợi ý "Ghi chú".'),
    ('SL', 'Number', 'Enable / Read-only', '> 0', 'Có', 'Theo dữ liệu', 'Chỉ nhập ở báo giá tự lập; hàng con theo '
     'công thức ghép bộ tự nhân theo SL cha.'),
    ('ĐVT', 'Dropdown', 'Enable / Read-only', 'ĐVT của hàng ERP', 'Có', 'Theo dữ liệu',
     'Đổi ĐVT → giá nhập, giá bán tính lại theo ĐVT mới.'),
    ('Giá nhập (tiền tệ)', 'Number', 'Enable / Disable', '> 0 (hàng tạm)', 'Có với hàng tạm', 'Theo dữ liệu',
     'Khoá với hàng ERP và cha có con. Không có Q5 → hàng ERP hiện "—". Lỗi: "Giá nhập phải > 0", "Giá vốn cha '
     'phải ≥ tổng giá vốn con".'),
    ('Thành tiền nhập', 'Number', 'Read-only', '–', '–', 'Tự tính', 'Giá nhập × SL.'),
    ('Giá bán (tiền tệ)', 'Number', 'Enable / Disable', '> 0', 'Có', 'Theo dữ liệu',
     'Khoá với hàng ERP. Lỗi: "Giá / thành tiền bán phải > 0", "Giá bán hàng cha phải ≥ tổng giá bán hàng con".'),
    ('Thành tiền bán', 'Number', 'Read-only', '–', '–', 'Tự tính', 'Giá bán × SL (trừ giảm giá — FR-09).'),
    ('Tỷ suất LN', 'Number', 'Read-only', '%', '–', 'Tự tính', 'Đỏ khi < mức sàn, xanh khi ≥ mức sàn.'),
    ('VAT(%)', 'Number', 'Enable / Disable', '0 – 100', 'Không', 'ERP: theo danh mục; tạm: 0',
     'ⓘ "VAT chỉ áp dụng trên dòng CHA hoặc SP độc lập. SP con được cộng gộp vào CHA." Khoá với hàng ERP. Lỗi: '
     '"VAT phải trong 0–100%".'),
    ('Tiền VAT / Thành tiền sau VAT', 'Number', 'Read-only', '–', '–', 'Tự tính', 'Tính trên thành tiền sau GG.'),
    ('Dòng dịch vụ / chi phí (phần B)', 'Table/Grid', 'Enable', '–', '–', 'Theo dữ liệu',
     'Biểu tượng phân loại Dịch vụ / Chi phí khác, badge "Tỷ lệ giá vốn: x%"; SL = 1; giá nhập = giá bán × tỷ lệ '
     'giá vốn (khoá) nếu danh mục có tỷ lệ.'),
    ('Ô VAT + nút Tất cả / VAT=0', 'Toolbar', 'Enable / Disable', '0 – 100', 'Không', 'Trống',
     'Áp VAT cho mọi hàng tạm cấp cha / chỉ dòng đang 0%; bỏ qua hàng ERP và dịch vụ.'),
    ('Làm tròn + nút Áp dụng', 'Dropdown', 'Enable / Disable', '7 mức (-3 … 2)', 'Không',
     'VNĐ: Số nguyên (0)', 'ⓘ "Khi bấm Áp dụng, hệ thống tự động tính toán giá trị làm tròn dựa trên cấu hình, giá '
     'trị làm tròn áp dụng với các cột GIÁ NHẬP, GIÁ BÁN".'),
    ('Dòng TỔNG', 'Number', 'Read-only', '–', '–', 'Tự tính', 'Tổng nhập, tổng bán, TSLN hàng hoá, VAT, sau VAT.'),
   ],
   [('Nhập Giá nhập / Giá bán / SL', 'Change',
     'After:\n– Tính lại thành tiền, TSLN, VAT, tổng phần A/B, khối Tổng hợp và cấp duyệt dự kiến (FR-10).\n'
     '– Dịch vụ có tỷ lệ giá vốn: giá nhập = giá bán × tỷ lệ.'),
    ('Đổi ĐVT', 'Change', 'After:\n– Lấy lại giá nhập, giá bán theo ĐVT mới.'),
    ('Nhập VAT 0 → > 0 lần đầu', 'Change', 'After:\n– Còn dòng 0% → mở popup "Áp dụng VAT đồng loạt?"; chọn "Áp dụng '
     'tất cả dòng còn 0%" → "Đã áp x% cho n dòng".'),
    ('Bấm Tất cả / VAT=0', 'Click', 'After:\n– Áp VAT cho hàng tạm cấp cha; hiển thị "Đã áp x% cho n dòng".'),
    ('Bấm Áp dụng (Làm tròn)', 'Click', 'Before:\n– Hỏi xác nhận "Xác nhận làm tròn giá".\nAfter:\n– Làm tròn giá '
     'nhập (trừ cha có con) và giá bán (trừ hàng con), dịch vụ; hiển thị "Đã làm tròn đơn giá".'),
   ])

# --------------------------------------------------------------------- FR-08
fr('FR-08', 'Thêm hàng hoá, hàng tạm, dịch vụ', 'crud', A_SALE, 'create',
   '- Màn Thêm mới, Validate dữ liệu, Thông báo và UI/UX. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Với báo giá tự lập: thêm nhóm / nhóm con, thêm hàng hoá ERP, hàng tạm (tái sử dụng hoặc nhập thủ '
        'công), hàng con, nhân bản, sửa, xoá dòng và kéo thả sắp xếp. Mọi loại báo giá: thêm / chọn lại / xoá dịch vụ, '
        'chi phí ở phần B.',
        tacnhan='Người lập báo giá',
        dieukien='Màn Tạo / Sửa báo giá, đã chọn Loại tiền tệ. Thêm hàng hoá chỉ có ở báo giá tự lập (báo giá từ BOM '
                 'khoá cấu trúc).',
        chinh='1. Người dùng bấm "Thêm mới" (phần A, trong nhóm) → popup "Thêm hàng hoá" liệt kê hàng ERP có lọc, '
              'phân trang.\n'
              '2. Chọn nhóm hàng, tích hàng cần thêm, bấm "Thêm n hàng hoá" → hàng vào bảng, giá theo Bảng giá và tỷ giá.\n'
              '3. Hàng chưa có trên ERP: bấm "Thêm hàng tạm" → tab "Chọn từ kho hàng tạm (Tái sử dụng)" hoặc "Thêm mới '
              'thủ công" → Lưu / Lưu và tiếp tục.\n'
              '4. Phần B bấm "Thêm mới" → popup "Thêm dịch vụ / chi phí" chọn từ danh mục.',
        phu='• Chưa chọn tiền tệ → "Vui lòng chọn loại tiền tệ trước khi thêm." (hoặc "… trước khi thêm sản phẩm.").\n'
            '• Hàng đã có trong danh sách → hỏi cộng dồn số lượng hay tạo dòng mới.\n'
            '• Hàng con trùng mã hàng cha → "Hàng con không được trùng mã với hàng cha (…): …".\n'
            '• Thêm nhóm: popup "Thêm nhóm sản phẩm", "Tên nhóm" bắt buộc; có nhóm thì mọi hàng phải thuộc nhóm.\n'
            '• Xoá hàng cha có con → "Xoá "…" sẽ xoá cả n hàng con kèm theo. Bạn có chắc?".\n'
            '• Nhân bản hàng tạm → "Đã nhân bản hàng tạm (dùng chung mã). Chỉnh số lượng nếu cần rồi Lưu."'),
   (' => Sửa / Làm giá => Thêm mới',
    [('08-them-hang.png', 'Popup Thêm hàng hoá'),
     ('08e-hang-tam-thucong.png', 'Thêm hàng tạm — tab Thêm mới thủ công'),
     ('08h-them-dich-vu.png', 'Popup Thêm dịch vụ / chi phí'),
     ('08f-sua-hang-tam.png', 'Popup Sửa hàng hoá (hàng tạm)')],
    'Các nút Thêm nhóm, Thêm nhóm con, Nhân bản, Thêm con, Sửa, biểu tượng thùng rác nằm ngay trên bảng Chi tiết '
    'sản phẩm của báo giá tự lập.'),
   [('Nút Thêm nhóm / Thêm nhóm con', 'Button', 'Enable / Ẩn', '–', '–', 'Ẩn với báo giá từ BOM',
     'Mở popup Thêm nhóm sản phẩm (Tên nhóm bắt buộc, tối đa 255 ký tự); bút chì đổi tên, thùng rác xoá nhóm.'),
    ('Nút Thêm mới (phần A / trong nhóm)', 'Button', 'Enable / Ẩn', '–', '–', 'Ẩn với báo giá từ BOM',
     'Mở popup Thêm hàng hoá.'),
    ('Popup Thêm hàng hoá — Nhóm hàng', 'Dropdown', 'Enable', 'Nhóm của báo giá', 'Có khi có nhóm',
     'Nhóm đang đứng', '"Hàng hoá sẽ được thêm vào nhóm này". Thiếu → "Vui lòng chọn nhóm hàng".'),
    ('Popup Thêm hàng hoá — Bộ lọc', 'Textbox', 'Enable', '–', 'Không', 'Trống',
     'Tìm theo tên, mã, model hàng hoá; Tìm kiếm nâng cao; Cài đặt bộ lọc.'),
    ('Popup Thêm hàng hoá — Bảng hàng ERP', 'Table/Grid', 'Enable', '20 dòng / trang', '–', 'Theo dữ liệu',
     'Ảnh, Loại hàng hoá, Tên, Lĩnh vực, Chương, Model, Mã hàng, Nguồn hàng, Giá niêm yết, Bảo hành, VAT, Định mức '
     'đàm phán giá, SL tồn có thể bán, SL KM có thể xuất, SL có thể LR, Ghi chú, Tính chất, Nguồn.'),
    ('Nút Thêm n hàng hoá', 'Button', 'Enable / Disable', '–', '–', 'Khoá khi chưa tích',
     'Thêm, không đóng popup; "Đã thêm n hàng hoá vào phiếu."'),
    ('Nút Thêm hàng tạm', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Mở khung Thêm hàng tạm 2 tab.'),
    ('Tab Chọn từ kho hàng tạm (Tái sử dụng)', 'Table/Grid', 'Enable', 'Hàng tạm của dự án', '–', 'Theo dữ liệu',
     'Tìm nhanh theo mã / tên; nút "Thêm n hàng hóa".'),
    ('Tab Thêm mới thủ công — Tên hàng hoá', 'Textbox', 'Enable', '0–255 ký tự', 'Có', 'Trống',
     'Thiếu → "Tên là bắt buộc".'),
    ('Đơn vị tính', 'Dropdown', 'Enable', 'Danh sách', 'Có', 'Trống', 'Thiếu → "Đơn vị tính là bắt buộc".'),
    ('Model', 'Dropdown', 'Enable', 'Tìm từ xa', 'Không', 'Trống', '⊕ thêm nhanh model.'),
    ('Thương hiệu', 'Dropdown', 'Enable', 'Tìm từ xa', 'Có', 'Trống', '⊕ thêm nhanh. Thiếu → "Thương hiệu là bắt '
     'buộc".'),
    ('Xuất xứ', 'Dropdown', 'Enable', 'Tìm từ xa', 'Có', 'Trống', '⊕ thêm nhanh. Thiếu → "Xuất xứ là bắt buộc".'),
    ('Số lượng cần dùng / Ghi chú / Đặc điểm – Thông số kỹ thuật', 'Number / Textarea', 'Enable', '≥ 0', 'Không',
     'SL 1', '–'),
    ('Nút Lưu / Lưu và tiếp tục / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Lưu và tiếp tục giữ popup, xoá '
     'trắng form.'),
    ('Nút Nhân bản / Thêm con / Sửa (dưới tên hàng tạm)', 'Button', 'Enable / Ẩn', '–', '–', 'Hiện khi rê chuột',
     'Nhân bản chỉ hàng tạm không có con; Thêm con không áp hàng ERP; Sửa mở popup "Sửa hàng hoá".'),
    ('Popup Sửa hàng hoá', 'Modal', 'Enable', '–', 'Tên, ĐVT có', 'Theo dòng',
     'Tên hàng hoá, Mã (khoá), Đơn vị tính, Model, Thương hiệu, Xuất xứ, Đặc điểm / Thông số kỹ thuật, Ghi chú.'),
    ('Ẩn con / Hiện con (cha ERP có con)', 'Button', 'Enable', '–', '–', 'Hiện con',
     'Ẩn / hiện hàng con khi in, xuất.'),
    ('Biểu tượng kéo / thùng rác (cột Thao tác)', 'Icon Button', 'Enable / Ẩn', '–', '–', 'Hiển thị',
     'Kéo đổi thứ tự nhóm, hàng, dịch vụ; xoá dòng (hỏi "Xác nhận xóa").'),
    ('Phần B — Thêm mới / bút chì / thùng rác', 'Button', 'Enable / Ẩn', '–', '–', 'Hiển thị khi sửa được',
     'Thêm dịch vụ, "Chọn lại dịch vụ / chi phí", "Xoá dịch vụ".'),
   ],
   [('Bấm Thêm mới / Thêm con / Thêm mới (phần B)', 'Click',
     'Before:\n– Chưa chọn tiền tệ → "Vui lòng chọn loại tiền tệ trước khi thêm.", dừng.\nAfter:\n– Mở popup ở tab '
     'tương ứng (Hàng hoá / Dịch vụ & Chi phí).'),
    ('Bấm Thêm n hàng hoá', 'Click',
     'During:\n– Hàng con trùng mã cha → "Hàng con không được trùng mã với hàng cha (…): …", dừng.\n– Trùng hàng đã '
     'có → hỏi cộng dồn số lượng hay tạo dòng mới.\nAfter:\n– Thêm vào nhóm đã chọn; giá bán, giá nhập, VAT lấy theo '
     'ERP + bảng giá + tỷ giá; "Đã thêm n hàng hoá vào phiếu."'),
    ('Bấm Lưu (Thêm hàng tạm thủ công)', 'Click',
     'During:\n– Thiếu Tên / ĐVT / Thương hiệu / Xuất xứ / Nhóm → báo đỏ dưới ô, không thêm.\nAfter:\n– Thêm dòng hàng '
     'tạm (mã HHBG sinh khi lưu báo giá); Lưu và tiếp tục giữ popup.'),
    ('Chọn dịch vụ / chi phí', 'Click', 'After:\n– Thêm dòng phần B, VAT theo danh mục; "Đã thêm: <tên>".'),
    ('Bấm Nhân bản', 'Click', 'After:\n– Thêm 1 dòng cùng hàng tạm ngay dưới, dùng chung mã; thông báo "Đã nhân bản '
     'hàng tạm (dùng chung mã). Chỉnh số lượng nếu cần rồi Lưu."'),
    ('Bấm thùng rác trên dòng', 'Click', 'Before:\n– Hỏi "Bạn có chắc muốn xoá "<tên>" khỏi báo giá?" (cha có con: '
     '"Xoá "<tên>" sẽ xoá cả n hàng con kèm theo. Bạn có chắc?").\nAfter:\n– Xoá dòng (và con) trên màn; chỉ ghi khi '
     'Lưu.'),
    ('Xoá nhóm', 'Click', 'After:\n– Xoá nhóm và nhóm con; hàng thuộc nhóm chuyển sang nhóm đầu còn lại.'),
   ])

# --------------------------------------------------------------------- FR-09
fr('FR-09', 'Áp dụng giảm giá', 'crud', A_SALE, 'create',
   '- Validate dữ liệu, Thông báo và UI/UX. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Chọn phương thức giảm giá của báo giá: Không GG, GG mặt hàng (nhập % hoặc ₫ theo từng dòng) hoặc GG '
        'tổng (khai các khoản giảm giá theo loại, phân bổ xuống từng dòng).',
        tacnhan='Người lập báo giá có quyền Q6',
        dieukien='Màn Tạo / Sửa, người dùng có quyền Q6. Dự án con kế thừa phương thức của dự án cha (khoá).',
        chinh='1. Chọn ô "GG:" = GG mặt hàng → bảng hiện cột GG(%), GG(₫), Đơn giá sau GG; nhập % hoặc ₫ (tự quy đổi '
              'lẫn nhau).\n'
              '2. Hoặc chọn GG tổng → khối "Giảm giá tổng đơn hàng": bấm "Thêm khoản GG", chọn Loại GG, Kiểu (% / ₫), '
              'Giá trị.\n'
              '3. Bấm "Phân bổ tự động" → cột "Phân bổ GG" được ghi theo tỷ lệ giá trị từng dòng (gồm cả chi phí vận '
              'chuyển); có thể sửa tay từng dòng.',
        phu='• Không có Q6 → ô GG khoá, biểu tượng ổ khoá "Bạn không có quyền áp dụng giảm giá trong báo giá"; khối GG '
            'tổng chỉ hiện để xem khi báo giá đã có khoản GG.\n'
            '• Đổi phương thức khi đã có dữ liệu → hỏi "Thay đổi phương thức sẽ xóa dữ liệu giảm giá hiện tại. Tiếp '
            'tục?".\n'
            '• Đã phân bổ nhưng giá đổi / tổng GG đổi → hiện nút "Phân bổ lại".\n'
            '• Nút "Xóa giảm giá" (GG mặt hàng) / "Xóa phân bổ" (GG tổng) xoá hàng loạt sau khi xác nhận.'),
   (' => Sửa / Làm giá', [('09-gg-mat-hang.png', 'GG mặt hàng — cột GG(%), GG(₫), Đơn giá sau GG'),
                          ('09b-gg-tong.png', 'GG tổng — khối Giảm giá tổng đơn hàng'),
                          ('09d-gg-tong-bang.png', 'GG tổng — cột GG phân bổ tự động và Phân bổ GG')],
    'Ô "GG:" nằm trên thanh công cụ ngay trên bảng Chi tiết sản phẩm.'),
   [('Ô GG', 'Dropdown', 'Enable / Disable', 'Không GG / GG mặt hàng / GG tổng', 'Không', 'Không GG',
     'Khoá khi không có Q6 hoặc dự án con (ổ khoá "Kế thừa từ dự án cha, không sửa được").'),
    ('GG(%)', 'Number', 'Enable', '0 – 100', 'Không', '0', 'Lỗi "GG (%) phải trong 0–100%".'),
    ('GG(₫)', 'Number', 'Enable', '≤ đơn giá bán', 'Không', '0', 'Lỗi "GG không được lớn hơn đơn giá bán".'),
    ('Đơn giá sau GG', 'Number', 'Read-only', '–', '–', 'Tự tính', 'Giá bán − GG(₫).'),
    ('Khối Giảm giá tổng đơn hàng — Loại GG', 'Dropdown', 'Enable', 'Danh mục loại giảm giá', 'Có', 'Trống',
     '"Vui lòng chọn loại GG".'),
    ('Kiểu', 'Dropdown', 'Enable', '% / ₫', 'Có', '%', '–'),
    ('Giá trị % / ₫', 'Number', 'Enable', '> 0; % ≤ 100', 'Có', '0', '"Giá trị phải > 0", "Giá trị tối đa 100%".'),
    ('Thành tiền GG / Tổng GG', 'Number', 'Read-only', '–', '–', 'Tự tính', '% tính trên tổng trước VAT gồm vận '
     'chuyển.'),
    ('Nút Thêm khoản GG / thùng rác', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Nút Phân bổ tự động / Phân bổ lại', 'Button', 'Enable / Ẩn', '–', '–', 'Ẩn khi Tổng GG = 0', '–'),
    ('Dòng Đã phân bổ', 'Label', 'Hiển thị', '–', '–', 'Ẩn', '"Đã phân bổ: x / y" ✓ (xanh) hoặc "— Chênh lệch: z" '
     '(đỏ).'),
    ('Cột GG phân bổ tự động', 'Number', 'Read-only', '–', '–', 'Tự tính', 'ⓘ "Tham khảo — tự tính theo tỷ lệ giá '
     'trị".'),
    ('Cột Phân bổ GG', 'Number', 'Enable / Disable', '≥ 0', 'Không', '0', 'Khoá khi tổng GG ≤ 0; tiêu đề hiện "Còn: '
     '…".'),
    ('Nút Xóa giảm giá / Xóa phân bổ', 'Button', 'Enable / Disable', '–', '–', 'Khoá khi không có dữ liệu', '–'),
   ],
   [('Đổi ô GG', 'Change', 'Before:\n– Đã có dữ liệu GG → hỏi "Xác nhận đổi phương thức giảm giá"; Huỷ → giữ cũ.\n'
     'After:\n– Xoá toàn bộ dữ liệu GG cũ (cả phân bổ vận chuyển), đổi cột bảng.'),
    ('Nhập GG(%) / GG(₫)', 'Change', 'After:\n– Quy đổi % ↔ ₫ theo đơn giá bán; tính lại thành tiền, VAT, tổng.'),
    ('Bấm Phân bổ tự động / Phân bổ lại', 'Click', 'Before:\n– Hỏi "Thao tác này sẽ ghi đè toàn bộ giá trị cột "Phân '
     'bổ GG" bằng giá trị phân bổ tự động (theo tỷ lệ giá trị). Bạn có chắc chắn?".\nAfter:\n– Ghi cột Phân bổ GG; '
     '"Đã cập nhật cột Phân bổ GG theo phân bổ tự động".'),
    ('Bấm Xóa giảm giá / Xóa phân bổ', 'Click', 'Before:\n– Hỏi "Bạn có chắc chắn muốn xóa toàn bộ dữ liệu giảm '
     'giá?" / "… xóa toàn bộ giá trị phân bổ giảm giá? Tất cả giá trị tại cột Phân bổ GG sẽ được đặt về 0".\nAfter:\n'
     '– Đưa về 0; "Đã xóa toàn bộ dữ liệu giảm giá" / "Đã xóa toàn bộ phân bổ giảm giá" (khoản GG tổng giữ nguyên).'),
    ('Lưu khi không có Q6', 'System', 'Before:\n– Thay đổi phương thức hoặc số tiền GG → "Bạn không có quyền áp '
     'dụng giảm giá trong báo giá", không lưu. Giữ nguyên GG cũ thì vẫn lưu được.'),
   ])

# --------------------------------------------------------------------- FR-10
fr('FR-10', 'Tổng hợp giá trị, chi phí vận chuyển và cấp duyệt dự kiến', 'crud', A_SALE, 'create',
   '- Validate dữ liệu và UI/UX. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Khối "Tổng hợp giá trị báo giá" tổng hợp theo 5 nhóm (I Hàng hoá, II Dịch vụ, III Chi phí, IV Chi phí '
        'vận chuyển, V Tổng), cho nhập chi phí vận chuyển; thanh chân trang hiển thị TSLN và cấp duyệt dự kiến cập '
        'nhật ngay khi sửa giá.',
        tacnhan='Người lập báo giá',
        dieukien='Màn Tạo / Sửa báo giá.',
        chinh='1. Hệ thống tính từng nhóm: Thành tiền nhập, Thành tiền trước VAT, (Giảm giá, Thành tiền sau GG), Thuế '
              'VAT, Thành tiền sau VAT.\n'
              '2. Người dùng nhập dòng IV: giá nhập vận chuyển, chi phí vận chuyển (trước VAT), GG vận chuyển (% / ₫ '
              'khi GG mặt hàng, hoặc Phân bổ GG khi GG tổng); VAT vận chuyển cố định 8%.\n'
              '3. Hệ thống tính TSLN trước GG, sau GG và cấp duyệt dự kiến C1 / C2 / C3 theo Cấu hình duyệt giá.',
        phu='• Không có Q5 → cột Thành tiền nhập hiện "—", ẩn LN và Cấp duyệt ở chân trang.\n'
            '• Tổng sau GG = 0 → cấp duyệt hiện "—".\n'
            '• Giá nhập / chi phí vận chuyển < 0 → "Giá nhập không hợp lệ" / "Chi phí vận chuyển không hợp lệ"; GG '
            'vận chuyển > chi phí → "GG vận chuyển không được lớn hơn chi phí".'),
   (' => Sửa / Làm giá', [('10-tong-hop.png', 'Khối Tổng hợp giá trị báo giá'),
                          ('10b-cap-duyet-c3.png', 'Tăng giá nhập vận chuyển — chân trang chuyển sang "C3 — TP & BGĐ"')]),
   [('STT / Nhóm chi phí', 'Text', 'Read-only', '5 dòng', '–', 'Cố định', 'I Hàng hoá · II Dịch vụ · III Chi phí · IV '
     'Chi phí vận chuyển · V Tổng giá trị báo giá.'),
    ('Thành tiền nhập', 'Number', 'Read-only', '–', '–', 'Tự tính', '"—" khi không có Q5.'),
    ('Giá nhập vận chuyển (dòng IV)', 'Number', 'Enable', '≥ 0', 'Không', '0', 'Người lập nhập được kể cả không có Q5.'),
    ('Chi phí vận chuyển (Thành tiền trước VAT, dòng IV)', 'Number', 'Enable', '≥ 0', 'Không', '0', '–'),
    ('Giảm giá / Thành tiền sau GG (trước VAT)', 'Number', 'Enable / Read-only', '≤ chi phí vận chuyển', 'Không', '0',
     'Chỉ hiện khi có phương thức GG.'),
    ('Thuế VAT vận chuyển', 'Number', 'Disable', '8%', '–', '8%', 'Cố định, không sửa.'),
    ('Thành tiền sau VAT / dòng V', 'Number', 'Read-only', '–', '–', 'Tự tính', 'Dòng V kèm mã tiền tệ.'),
    ('TSLN trước GG / TSLN sau GG', 'Label', 'Read-only', '%', '–', 'Tự tính', 'Xanh ≥ 0, đỏ < 0; chỉ hiện khi có Q5.'),
    ('Chân trang: Nhập / Bán / GG / Sau GG / ~VND / LN', 'Label', 'Read-only', '–', '–', 'Tự tính',
     '~VND chỉ hiện khi tiền tệ khác VNĐ; LN màu theo mức sàn.'),
    ('Chân trang: Cấp duyệt', 'Badge', 'Read-only', 'C1 / C2 / C3', '–', 'Tự tính',
     '"C1 — Tự duyệt", "C2 — TP", "C3 — TP & BGĐ"; chỉ hiện khi có Q5.'),
    ('Dòng tỷ giá', 'Label', 'Read-only', '–', '–', 'Ẩn với VNĐ', '"Tỷ giá: 1 <mã> = x VND (dd/mm/yyyy) - Bảng giá '
     'bán lẻ".'),
   ],
   [('Sửa bất kỳ giá / SL / GG / vận chuyển', 'Change',
     'After:\n– Tính lại 5 dòng tổng hợp, TSLN và cấp duyệt dự kiến = cấp cao hơn giữa cấp theo giá trị đơn hàng sau '
     'GG quy VNĐ và cấp theo TSLN (BR-07).'),
    ('Nhập GG vận chuyển % / ₫', 'Change', 'After:\n– Quy đổi lẫn nhau theo chi phí vận chuyển.'),
   ])

# --------------------------------------------------------------------- FR-11
fr('FR-11', 'Điều khoản báo giá và ghi chú nội bộ', 'crud', A_SALE, 'create',
   '- Validate dữ liệu, Thông báo và UI/UX. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Khối "Thanh toán & Ghi chú nội bộ": soạn Điều khoản báo giá (chọn nhanh từ mẫu điều khoản) và ghi chú '
        'chỉ dùng nội bộ.',
        tacnhan='Người lập báo giá',
        dieukien='Màn Tạo / Sửa báo giá.',
        chinh='1. Người dùng chọn "Mẫu điều khoản báo giá" → nội dung mẫu điền vào ô Điều khoản báo giá.\n'
              '2. Chỉnh sửa nội dung bằng trình soạn thảo; có thể dùng biến {{VAT_NOTE}}, {{VAN_CHUYEN_NOTE}}… được '
              'thay bằng số liệu thật khi xem / in.\n'
              '3. Nhập Ghi chú nội bộ (không in ra cho khách).',
        phu='• Ô điều khoản đang có nội dung mà chọn mẫu khác → hỏi "Nếu bạn thay đổi mẫu điều khoản, thông tin điều '
            'khoản báo giá sẽ bị thay đổi! Bạn có chắc chắn muốn thực hiện?"; Huỷ → giữ nguyên.\n'
            '• Gửi duyệt khi điều khoản trống → "Vui lòng nhập điều khoản thanh toán".'),
   (' => Sửa / Làm giá', [('11-dieu-khoan.png', 'Khối Thanh toán & Ghi chú nội bộ')]),
   [('Tiêu đề khối', 'Label', 'Hiển thị', '–', '–', 'Thanh toán & Ghi chú nội bộ', 'Bấm để thu gọn / mở.'),
    ('Mẫu điều khoản báo giá', 'Dropdown', 'Enable / Ẩn', 'Danh mục mẫu điều khoản', 'Không', 'Trống',
     'Ẩn khi không sửa được.'),
    ('Điều khoản báo giá', 'Textarea (soạn thảo)', 'Enable / Disable', 'HTML', 'Có khi gửi duyệt', 'Trống',
     'Bỏ thẻ định dạng rồi mới kiểm trống.'),
    ('Ghi chú nội bộ (chỉ nội bộ)', 'Textarea (soạn thảo)', 'Enable / Disable', 'HTML', 'Không', 'Trống', '–'),
   ],
   [('Chọn Mẫu điều khoản', 'Change', 'Before:\n– Ô đang có nội dung → hỏi "Xác nhận thay đổi mẫu điều khoản".\n'
     'After:\n– Điền nội dung mẫu vào ô Điều khoản báo giá.'),
    ('Gửi duyệt khi điều khoản trống', 'System', 'During:\n– Báo "Vui lòng nhập điều khoản thanh toán" dưới ô, cuộn '
     'tới ô lỗi.'),
   ])

# --------------------------------------------------------------------- FR-12
fr('FR-12', 'Import Excel', 'io', A_SALE, 'excel',
   '- Quy tắc Excel (Import). Chỉ bổ sung các quy tắc riêng của màn Báo giá tại phần mô tả chi tiết.',
   dict(mota='Nạp hàng hoá, dịch vụ, giảm giá và chi phí vận chuyển từ file Excel lên lưới báo giá; dữ liệu chỉ được '
        'ghi khi bấm Lưu báo giá.',
        tacnhan='Người lập báo giá',
        dieukien='Màn Sửa báo giá (đã lưu ít nhất 1 lần), người lập, trạng thái Đang tạo.',
        chinh='1. Bấm "Import Excel" → popup "Import báo giá từ Excel".\n'
              '2. Bấm "Tải file mẫu" (mẫu theo phương thức GG của báo giá), điền file, bấm "Chọn file Excel".\n'
              '3. Bấm "Load lên bảng" → lưới hiển thị từng dòng; bấm "Validate" → hệ thống kiểm tra.\n'
              '4. Hợp lệ hết → "Dữ liệu hợp lệ. Bấm "Import" để ghi đè báo giá."; bấm "Import" → chọn phương thức.\n'
              '5. Hệ thống đổ dữ liệu lên lưới: "Đã nạp dữ liệu từ file Excel. Kiểm tra lại rồi bấm "Lưu báo giá" để '
              'chốt."',
        phu='• Chưa lưu báo giá → "Vui lòng lưu báo giá trước khi import."\n'
            '• File không phải .xlsx / .xls → "Định dạng file không hỗ trợ. Vui lòng chọn file Excel (.xlsx, .xls)."\n'
            '• Còn dòng lỗi → không Import được (tất cả hoặc không); xem "Xem chi tiết n lỗi" → popup "Import thất bại: '
            'Phát hiện n lỗi dữ liệu" với nút Sao chép lỗi / Tải File lỗi.\n'
            '• Báo giá từ BOM chỉ có "Import từng phần" ("Báo giá kế thừa BOM có cấu trúc cố định — chỉ hỗ trợ cập '
            'nhật từng phần.").\n'
            '• File thuộc báo giá khác → popup "Phát hiện dữ liệu thuộc báo giá khác", nút Sao chép.\n'
            '• Model / Thương hiệu / Xuất xứ chưa có → cảnh báo "Tạo mới danh mục" (không chặn).'),
   (' => Sửa / Làm giá => Import Excel', [('12c-import-validate.png', 'Popup Import — sau khi Validate hợp lệ'),
                                          ('12d-import-phuongthuc.png', 'Popup Chọn phương thức import')]),
   [('Nút Chọn file Excel', 'Button', 'Enable', '.xlsx, .xls', 'Có', '–', 'Hiện tên file đã chọn.'),
    ('Nút Tải file mẫu', 'Button', 'Enable', '–', '–', 'Hiển thị', '3 mẫu: Không GG / GG mặt hàng / GG tổng.'),
    ('Nút Load lên bảng', 'Button', 'Enable / Disable', '–', '–', 'Khoá khi chưa chọn file', '–'),
    ('Nút Validate', 'Button', 'Enable / Disable', '–', '–', 'Khoá khi lưới trống', '–'),
    ('Nút Chỉ dòng lỗi / Hiện tất cả', 'Button', 'Enable / Disable', '–', '–', 'Khoá khi lưới trống', '–'),
    ('Thống kê Tổng / Hợp lệ / Lỗi / Tạo mới danh mục', 'Badge', 'Hiển thị', '–', '–', 'Ẩn trước Validate', '–'),
    ('Lưới dữ liệu', 'Table/Grid', 'Enable / Disable', '–', '–', '"Chưa có dữ liệu. Chọn file Excel rồi bấm "Load '
     'lên bảng"."', 'Cột: #, Trạng thái, Loại*, Có tính doanh thu, Nhóm hàng cha / con, STT*, Mã hàng cha, Mã hàng*, '
     'Tên hàng*, Model, Thương hiệu*, Xuất xứ*, Thông số kỹ thuật, Ghi chú, ĐVT*, Số lượng*, Đơn giá nhập, Thành '
     'tiền nhập, Đơn giá bán, (cột GG), Thành tiền bán, TSLN, Thuế VAT, Giá trị VAT, Thành tiền sau VAT. Ô sửa '
     'được; dòng hợp lệ khoá.'),
    ('Nút Import', 'Button', 'Enable / Disable', '–', '–', 'Khoá khi chưa Validate / còn lỗi',
     'Rê chuột: "Bấm Validate để kiểm tra dữ liệu trước khi Import." / "Còn dòng lỗi — sửa file rồi Validate lại."'),
    ('Popup Chọn phương thức import', 'Modal', 'Enable', '2 lựa chọn', 'Có', 'Import từng phần',
     '"Import từng phần" — giữ dòng không có trong file, cập nhật dòng khớp, thêm dòng mới; "Thay thế hoàn toàn" — '
     'xoá toàn bộ rồi đổ lại, "Thao tác không thể hoàn tác."'),
    ('Nút Làm mới / Đóng', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
   ],
   [('Bấm Import Excel', 'Click', 'Before:\n– Báo giá chưa có mã → "Vui lòng lưu báo giá trước khi import.", '
     'dừng.\nAfter:\n– Mở popup trống.'),
    ('Bấm Load lên bảng', 'Click', 'During:\n– Sai định dạng file / không đọc được → thông báo đỏ trong popup.\n'
     'After:\n– Đổ dữ liệu file lên lưới.'),
    ('Bấm Validate', 'Click', 'Before:\n– Chỉ người lập + Đang tạo; báo giá ngoài phạm vi xem → không tìm thấy.\n'
     'During:\n– Kiểm tra bắt buộc, kiểu số, mã hàng, nhóm, cha – con; lỗi gắn theo dòng + cột.\nAfter:\n– Hợp lệ → '
     'khoá các dòng, "Dữ liệu hợp lệ. Bấm "Import" để ghi đè báo giá."; không ghi dữ liệu.'),
    ('Bấm Import → chọn phương thức', 'Click', 'After:\n– Tự tạo Model / Thương hiệu / Xuất xứ còn thiếu; đổ dữ liệu '
     'lên lưới theo phương thức; "Đã nạp dữ liệu từ file Excel. Kiểm tra lại rồi bấm "Lưu báo giá" để chốt."'),
    ('Bấm Sao chép lỗi / Tải File lỗi', 'Click', 'After:\n– "Đã sao chép n lỗi vào clipboard." / "Đã tải file lỗi: '
     '<tên>".'),
   ])

# --------------------------------------------------------------------- FR-13
fr('FR-13', 'Lưu nháp', 'crud', A_SALE, 'create',
   '- Màn Thêm mới, Validate dữ liệu, Thông báo. Logic ghi lịch sử theo SRS Các quy tắc chung - Quy tắc ghi lịch sử.',
   dict(mota='Lưu báo giá ở trạng thái Đang tạo mà chưa gửi duyệt; chỉ kiểm tra khoảng giá trị, không bắt buộc đủ '
        'thông tin.',
        tacnhan='Người lập báo giá',
        dieukien='Màn Tạo hoặc Sửa báo giá (người lập, Đang tạo).',
        chinh='1. Người dùng bấm "Lưu nháp".\n'
              '2. Hệ thống kiểm tra VAT, % giảm giá trong 0–100 và khoản GG tổng.\n'
              '3. Lưu báo giá: lần đầu sinh mã BG-YYYY-NNNNN và mã HHBG cho hàng tạm; ghi lịch sử.\n'
              '4. Hiển thị "Đã tạo báo giá" (lần đầu) / "Đã lưu báo giá" và quay về danh sách.',
        phu='• Có lỗi → báo đỏ tại ô, cuộn tới ô lỗi đầu, không lưu.\n'
            '• Báo giá lệch Báo giá Mỏ neo → vẫn lưu, thông báo "Lưu ý: … Hệ thống vẫn lưu nháp, nhưng cần đồng bộ '
            'trước khi Trình duyệt."\n'
            '• Dự án cha → "Dự án cha không lập báo giá riêng. Vui lòng tạo báo giá trên dự án con, hoặc dùng Báo giá '
            'tổng ở màn dự án cha."\n'
            '• Lỗi khác → nội dung lỗi hoặc "Lỗi lưu".'),
   (' => Sửa / Làm giá => Lưu nháp', [('07-lam-gia.png', 'Màn Làm giá — nút Lưu nháp ở chân trang')]),
   [('Nút Lưu nháp', 'Button', 'Enable / Ẩn', '–', '–', 'Hiện khi sửa được', 'Hiện trạng thái đang xử lý khi lưu.'),
    ('Thông báo lỗi tại ô', 'Toast / Alert', 'Hiển thị', '–', '–', 'Ẩn', 'Viền đỏ + chữ đỏ dưới ô sai.'),
    ('Thông báo thành công', 'Toast / Alert', 'Hiển thị', '–', '–', 'Ẩn', '"Đã tạo báo giá" / "Đã lưu báo giá".'),
   ],
   [('Bấm Lưu nháp', 'Click',
     'Before:\n' + OWNER + '\n– Màn Tạo: phải là Sale phụ trách dự án ("Bạn không phải Sale phụ trách dự án này") '
     'hoặc có quyền xây dựng giá khi lập từ YCBG.\n'
     'During:\n– VAT ngoài 0–100 → "VAT phải trong 0–100%".\n– GG % ngoài 0–100 → "GG (%) phải trong 0–100%".\n'
     '– Khoản GG tổng thiếu loại / giá trị ≤ 0 → "Vui lòng chọn loại GG" / "Giá trị phải > 0".\n– Màn Tạo chưa chọn '
     'dự án → "Vui lòng chọn dự án trước khi tạo báo giá".\n– Có lỗi → không thực hiện After.\n'
     'After:\n– Lưu thông tin chung, hàng hoá, dịch vụ, giảm giá, vận chuyển; hàng ERP lấy lại giá + VAT theo ERP '
     '(trừ khi đã chọn Từ chối cập nhật giá — FR-15).\n– Ghi lịch sử "Tạo báo giá" / "Lưu nháp" kèm trường thay đổi.\n'
     '– Hiển thị "Đã tạo báo giá" / "Đã lưu báo giá", về danh sách.'),
   ])

# --------------------------------------------------------------------- FR-14
fr('FR-14', 'Gửi duyệt', 'action', A_SALE, 'create',
   '- Validate dữ liệu, Thông báo. Ngưỡng cấp duyệt theo SRS Cấu hình duyệt giá.',
   dict(mota='Lưu đầy đủ, kiểm tra điều kiện, tính cấp duyệt C1/C2/C3 rồi gửi báo giá đi duyệt (hoặc tự duyệt với C1).',
        tacnhan='Người lập báo giá',
        dieukien='Người lập, báo giá Đang tạo (hoặc đang tạo mới).',
        chinh='1. Bấm "Gửi duyệt" → hệ thống lưu ngầm với kiểm tra đầy đủ (màn Tạo: chỉ xem trước, chưa ghi).\n'
              '2. Có mặt hàng đơn giá bán ≤ 1.000 → popup cảnh báo, chọn "Tiếp tục gửi duyệt" hoặc "Quay lại".\n'
              '3. Popup "Gửi duyệt báo giá" hiển thị Tổng giá nhập (Q5), Tổng giá bán, Tổng giảm giá, Tổng bán sau GG, '
              'Tỷ suất LN (Q5), Cấp duyệt dự kiến và thông điệp theo cấp.\n'
              '4. C1: bấm "Xác nhận duyệt" → báo giá Đã duyệt ("Đã tự duyệt báo giá"). C2/C3: bấm "Xác nhận gửi" → '
              'Chờ TP duyệt ("Đã gửi duyệt báo giá").\n'
              '5. Quay về danh sách.',
        phu='• Thiếu thông tin / giá → báo đỏ tại ô, cuộn tới lỗi, không mở popup.\n'
            '• Lệch Báo giá Mỏ neo → popup "Không thể trình duyệt — vi phạm quy tắc đồng bộ" liệt kê từng lỗi.\n'
            '• Màn Tạo bị chặn → popup "Không thể gửi duyệt báo giá" với lý do.\n'
            '• Báo giá chỉ gồm hàng / dịch vụ ERP, đơn giá > 1.000, không giảm giá → "Tự động duyệt" (cấp 1).\n'
            '• Không tính được cấp → "Không tính được cấp duyệt. Vui lòng kiểm tra giá sản phẩm."',
        dacbiet='Màn Tạo: tạo + gửi duyệt trong cùng 1 lượt, lỗi thì không để lại báo giá nháp.'),
   (' => Sửa / Làm giá => Gửi duyệt',
    [('13b-gui-duyet-loi.png', 'Gửi duyệt khi thiếu thông tin — báo lỗi tại ô'),
     ('14b-canh-bao-gia-thap.png', 'Cảnh báo mặt hàng có đơn giá ≤ 1.000'),
     ('14-gui-duyet-c1.png', 'Popup Gửi duyệt báo giá — Cấp 1 (tự duyệt)'),
     ('14c-gui-duyet-c3.png', 'Popup Gửi duyệt báo giá — Cấp 3 (2 bước duyệt)')]),
   [('Nút Gửi duyệt', 'Button', 'Enable / Ẩn', '–', '–', 'Hiện khi sửa được', '–'),
    ('Popup cảnh báo giá thấp', 'Modal', 'Hiển thị', '–', '–', 'Ẩn', 'Tiêu đề "Cảnh báo: Tồn tại mặt hàng có đơn giá '
     '<= 1.000 vnđ"; bảng STT, Mã hàng hóa, Tên hàng hóa, Đơn giá bán; nút "Tiếp tục gửi duyệt", "Quay lại".'),
    ('Tổng giá nhập / Tỷ suất LN', 'Number', 'Read-only', '–', '–', 'Ẩn khi không có Q5', '–'),
    ('Tổng giá bán / Tổng giảm giá / Tổng bán sau GG', 'Number', 'Read-only', '–', '–', 'Theo báo giá',
     'Giảm giá chỉ hiện khi > 0.'),
    ('Cấp duyệt dự kiến', 'Text', 'Read-only', 'Cấp 1 – 3 / Tự động duyệt', '–', 'Tự tính', '–'),
    ('Thông điệp theo cấp', 'Toast / Alert', 'Hiển thị', '–', '–', 'Theo cấp',
     'C1 "Theo quy chế, bạn có thể tự duyệt báo giá này."; C2 "Theo quy chế, báo giá cần gửi tới Trưởng phòng '
     'duyệt."; C3 "Theo quy chế, báo giá cần gửi qua 2 cấp duyệt:" + 1 Trưởng phòng duyệt & chuyển BGĐ → 2 Ban '
     'giám đốc duyệt.'),
    ('Nút Xác nhận duyệt / Xác nhận gửi', 'Button', 'Enable', '–', '–', 'Theo cấp', 'C1 "Xác nhận duyệt", C2/C3 '
     '"Xác nhận gửi".'),
    ('Nút Huỷ', 'Button', 'Enable', '–', '–', 'Hiển thị', 'Đóng, báo giá giữ Đang tạo.'),
   ],
   [('Bấm Gửi duyệt', 'Click',
     'Before:\n' + OWNER + '\n'
     'During:\n– Giao hàng trống → "Vui lòng nhập thời gian giao hàng"; tiền tệ / bảng giá / giai đoạn / điều khoản '
     'trống → "Vui lòng chọn loại tiền tệ" / "Vui lòng chọn bảng giá" / "Vui lòng chọn giai đoạn dự án" / "Vui lòng '
     'nhập điều khoản thanh toán".\n– Giá bán / thành tiền bán ≤ 0, giá nhập hàng tạm ≤ 0, giá cha < tổng con, GG > '
     'đơn giá, vận chuyển âm → báo đỏ tại ô.\n– Hệ thống kiểm tra lại: "Vui lòng nhập điều khoản thanh toán trước khi '
     'gửi duyệt.", "Vui lòng chọn giai đoạn dự án trước khi gửi duyệt.", "Có n sản phẩm/dịch vụ chưa hợp lệ (giá bán '
     '> 0, số lượng > 0, giảm giá ≤ đơn giá bán).", "KHÔNG THỂ TRÌNH DUYỆT — báo giá không khớp với Báo giá Mỏ neo …".'
     '\n– Có lỗi → không thực hiện After.\n'
     'After:\n– Lưu báo giá; tính cấp duyệt (BR-07); mở popup Gửi duyệt báo giá.'),
    ('Bấm Xác nhận gửi (C2/C3)', 'Click',
     'After:\n– Đồng bộ giai đoạn về dự án; tính hiệu lực báo giá; trạng thái → Chờ TP duyệt, lưu cấp duyệt, xoá lý '
     'do từ chối cũ.\n– Ghi lịch sử "Gửi duyệt" kèm cấp.\n– Gửi thông báo cho Trưởng phòng duyệt giá quản lý phòng ban '
     'của báo giá.\n– "Đã gửi duyệt báo giá", về danh sách.'),
    ('Bấm Xác nhận duyệt (C1)', 'Click',
     'After:\n– Trạng thái → Đã duyệt, người duyệt = người lập; YCBG → Đã có báo giá; dự án → Thương thảo giá; giải '
     'pháp → Đã duyệt giá; đồng bộ ERP.\n– Ghi lịch sử "Tự duyệt", gửi thông báo duyệt (BR-11).\n– "Đã tự duyệt báo '
     'giá", về danh sách.\n– Cấp thay đổi giữa lúc xem và xác nhận → "Cấp duyệt của báo giá vừa thay đổi thành cấp 1 '
     '(tự duyệt). Vui lòng bấm Gửi duyệt lại để xác nhận."'),
   ])

# --------------------------------------------------------------------- FR-15
fr('FR-15', 'Chỉnh sửa báo giá', 'crud', A_SALE, 'create',
   '- Màn Chỉnh sửa, Validate dữ liệu, Thông báo. Logic ghi lịch sử theo SRS Các quy tắc chung - Quy tắc ghi lịch sử.',
   dict(mota='Mở màn "Làm giá" của báo giá đang tạo để sửa thông tin chung, giá, giảm giá, điều khoản rồi Lưu nháp / '
        'Gửi duyệt.',
        tacnhan='Người lập báo giá',
        dieukien='Người dùng là người lập, báo giá Đang tạo (kể cả báo giá bị từ chối).',
        chinh='1. Bấm "Sửa / Làm giá" trên dòng danh sách hoặc ở chân màn chi tiết.\n'
              '2. Màn "Làm giá: <mã> (Đang tạo)" mở với dữ liệu đã lưu; Dự án, Loại tiền tệ, Bảng giá chỉ đọc.\n'
              '3. Hệ thống so đơn giá hàng ERP với thời giá; có lệch → hỏi cập nhật.\n'
              '4. Người dùng sửa theo FR-07 → FR-12, bấm Lưu nháp hoặc Gửi duyệt.',
        phu='• Không phải người lập / không Đang tạo → nút không hiển thị; mở thẳng màn thì hiện dải vàng "Báo giá ở '
            'trạng thái (…) — không cho phép sửa." và khoá mọi ô.\n'
            '• Popup "Đơn giá hàng hoá đã thay đổi": "Đồng ý" → cập nhật đơn giá theo thời giá; "Từ chối" → giữ đơn giá '
            'đã lưu khi Lưu.\n'
            '• Báo giá sao chép (Đang tạo) được đổi Dự án — chỉ sang dự án chưa có BOM tổng hợp đã duyệt; hỏi "Xác nhận '
            'đổi Dự án", giữ nguyên hàng hoá / dịch vụ / giảm giá, chỉ ghi khi Lưu.'),
   (' => Sửa / Làm giá', [('07-lam-gia.png', 'Màn Làm giá báo giá tự lập')]),
   [('Tiêu đề', 'Label', 'Hiển thị', '–', '–', 'Làm giá: <mã> (trạng thái)', '–'),
    ('Dải cảnh báo không cho sửa', 'Toast / Alert', 'Hiển thị', '–', '–', 'Ẩn', 'Khi không thoả điều kiện sửa.'),
    ('Mã báo giá / YCBG', 'Text', 'Read-only', '–', '–', 'Theo dữ liệu', 'Bấm mã YCBG mở popup chi tiết YCBG.'),
    ('Dự án', 'Link / Dropdown', 'Read-only / Enable', '–', 'Có với bản sao', 'Theo dữ liệu',
     'Chỉ báo giá sao chép Đang tạo mới đổi được.'),
    ('Giải pháp / Hạng mục / BOM', 'Link', 'Read-only', '–', '–', 'Theo dữ liệu', 'Chỉ báo giá từ BOM.'),
    ('Loại tiền tệ / Bảng giá', 'Dropdown', 'Disable', '–', '–', 'Theo dữ liệu', 'Khoá ở màn Sửa.'),
    ('Các trường còn lại', 'Theo FR-06 → FR-11', 'Enable', '–', '–', 'Theo dữ liệu', '–'),
    ('Popup Đơn giá hàng hoá đã thay đổi', 'Modal', 'Hiển thị', '≤ 10 dòng liệt kê', '–', 'Ẩn',
     '"Có sự thay đổi về đơn giá của một số hàng hóa. Bạn có muốn cập nhật đơn giá theo thời giá hiện tại không?" + '
     'danh sách "<tên>: giá cũ → giá mới VNĐ"; nút Đồng ý / Từ chối.'),
    ('Nút Quay lại / Lưu nháp / Gửi duyệt', 'Button', 'Enable / Ẩn', '–', '–', 'Theo quyền', 'Lưu nháp, Gửi duyệt '
     'ẩn khi không sửa được.'),
   ],
   [('Mở màn Sửa', 'System', 'Before:\n– Báo giá ngoài phạm vi xem → "Không tìm thấy báo giá."\nAfter:\n– Nạp dữ '
     'liệu; kiểm tra đơn giá hàng ERP, lệch → popup hỏi cập nhật.\n– Lỗi → "Không tải được báo giá".'),
    ('Bấm Đồng ý / Từ chối (cập nhật giá)', 'Click', 'After:\n– Đồng ý: thay đơn giá bán + VAT hàng ERP theo thời '
     'giá; GG tổng đã phân bổ → nhắc phân bổ lại.\n– Từ chối: giữ đơn giá, khi lưu hệ thống không áp giá ERP mới.'),
    ('Đổi Dự án (bản sao)', 'Change', 'Before:\n– Hỏi xác nhận.\nDuring:\n– Dự án có BOM tổng hợp đã duyệt → "Dự án '
     '"<tên>" đã có BOM tổng hợp đã duyệt nên không thể chọn cho báo giá sao chép. Vui lòng chọn dự án khác hoặc tạo '
     'báo giá mới từ BOM của dự án đó."\nAfter:\n– Lấy lại khách hàng, tiền tệ, bảng giá; ngắt BOM khi Lưu.'),
    ('Bấm Lưu nháp / Gửi duyệt', 'Click', 'Như FR-13, FR-14.'),
   ])

# --------------------------------------------------------------------- FR-16
fr('FR-16', 'Xem chi tiết báo giá', 'view', A_ALL, 'detail',
   '- Màn Xem chi tiết và Phân quyền. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Màn "Chi tiết báo giá: <mã> (trạng thái) Cấp n" ở chế độ chỉ đọc: thông tin chung, chi tiết báo giá, '
        'giảm giá, tổng hợp, điều khoản, ghi chú, lịch sử và các nút thao tác theo trạng thái.',
        tacnhan='Người lập báo giá, Người duyệt giá; Người dùng đã đăng nhập',
        dieukien='Báo giá thuộc phạm vi xem (BR-01) hoặc đang nằm trong hàng chờ duyệt của người dùng.',
        chinh='1. Bấm mã báo giá ở danh sách (hoặc mở từ Báo giá chờ duyệt / thông báo).\n'
              '2. Hệ thống hiển thị các khối thông tin và các nút ở chân trang theo điều kiện.',
        phu='• Báo giá bị từ chối → dải đỏ "Đã bị từ chối: <lý do>".\n'
            '• Có Báo giá Mỏ neo → dải xác nhận đồng bộ 100% hoặc cảnh báo lệch n điểm.\n'
            '• Không thuộc phạm vi → "Không tìm thấy báo giá."'),
   (' => Mã báo giá', [('15-xem-chi-tiet.png', 'Màn Chi tiết báo giá (báo giá bị từ chối, đang tạo)')]),
   [('Tiêu đề', 'Label', 'Hiển thị', '–', 'Chi tiết báo giá: <mã> (trạng thái) + nhãn cấp duyệt', '–'),
    ('Dải Đã bị từ chối', 'Toast / Alert', 'Hiển thị', '–', 'Ẩn', 'Khi có lý do từ chối.'),
    ('Thông tin chung', 'Section', 'Read-only', '–', 'Mở', 'Người lập, Ngày tạo; Mã báo giá, YCBG, BOM list, Dự án, '
     'Giải pháp (+version), Hạng mục, Khách hàng, MST, Người liên hệ, SĐT, Địa chỉ, Email, Hiệu lực, Giao hàng, Bảo '
     'hành, Loại tiền tệ + tỷ giá + Bảng giá, Giai đoạn dự án (ⓘ), Ngày lập, Người lập, Ghi chú yêu cầu.'),
    ('Chi tiết báo giá', 'Table/Grid', 'Read-only', '–', 'Theo dữ liệu', 'Nhãn phương thức GG; cột như FR-07; nút '
     'Thu gọn / Mở rộng tất cả, Ẩn / Hiện cột chi tiết. Không có Q5 → giá vốn "—".'),
    ('Giảm giá tổng đơn hàng', 'Table/Grid', 'Read-only', '–', 'Ẩn nếu không GG tổng', 'Loại GG, Giá trị, Tổng GG.'),
    ('Tổng hợp giá trị báo giá', 'Table/Grid', 'Read-only', '5 dòng', 'Theo dữ liệu', 'Như FR-10; TSLN trước / sau '
     'GG chỉ hiện khi có Q5.'),
    ('Điều khoản báo giá / Ghi chú nội bộ', 'Text', 'Read-only', '–', 'Theo dữ liệu', 'Biến điều khoản đã thay số '
     'liệu thật.'),
    ('Ghi chú Kinh doanh', 'Text / Textarea', 'Read-only / Enable', '–', 'Chỉ khi Đã duyệt', 'Xem FR-23.'),
    ('Khối Lịch sử', 'Section', 'Read-only', '–', 'Thu gọn', 'Xem FR-25.'),
    ('Nút chân trang', 'Button', 'Enable / Ẩn', '–', 'Theo điều kiện', 'Quay lại · Xuất Excel · In · Sao chép · Sửa / '
     'Làm giá + Xoá (người lập, Đang tạo) · Duyệt / Duyệt & chuyển BGĐ / BGĐ duyệt / Từ chối (FR-17, 18) · Lập hợp '
     'đồng / Xem hợp đồng (FR-24).'),
   ],
   [('Mở màn chi tiết', 'System', 'Before:\n– Kiểm tra phạm vi xem; ngoài phạm vi → "Không tìm thấy báo giá."\n'
     'After:\n– Nạp báo giá, ngày đổi giá hàng ERP (⚠ cạnh mã), kiểm tra Báo giá Mỏ neo.\n– Lỗi → "Không tải được '
     'báo giá".'),
    ('Bấm Quay lại', 'Click', 'After:\n– Về trang trước; mở trực tiếp thì về danh sách (vào từ Báo giá chờ duyệt thì '
     'về đó).'),
    ('Bấm mã YCBG / BOM / Dự án / Giải pháp', 'Click', 'After:\n– YCBG mở popup chi tiết YCBG; các liên kết khác mở '
     'tab mới.'),
   ], mode='ro', usecase=False)

# --------------------------------------------------------------------- FR-17
fr('FR-17', 'Duyệt báo giá', 'action', A_DUYET, 'detail',
   '- Màn Xem chi tiết và Phân quyền, Thông báo. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Trưởng phòng duyệt báo giá Chờ TP duyệt (C2 → Đã duyệt; C3 → chuyển BGĐ), Ban giám đốc duyệt báo giá '
        'Chờ BGĐ duyệt.',
        tacnhan='Trưởng phòng duyệt giá (Q3), Ban giám đốc duyệt giá (Q4)',
        dieukien='Q3 + báo giá Chờ TP duyệt thuộc phòng ban mình quản lý; hoặc Q4 + báo giá Chờ BGĐ duyệt cùng công '
                 'ty.',
        chinh='1. Người duyệt mở chi tiết báo giá (thường từ màn Báo giá chờ duyệt).\n'
              '2. Bấm "Duyệt" (C2) / "Duyệt & chuyển BGĐ" (C3) / "BGĐ duyệt".\n'
              '3. Xác nhận ở hộp "Xác nhận duyệt".\n'
              '4. Hệ thống cập nhật trạng thái, gửi thông báo, quay về danh sách nguồn.',
        phu='• Huỷ ở hộp xác nhận → không thay đổi.\n'
            '• Không quản lý phòng ban của báo giá → "Bạn không quản lý phòng ban của báo giá này."\n'
            '• Khác công ty → "Bạn không thuộc công ty của báo giá này."\n'
            '• Sai trạng thái → "Báo giá không ở trạng thái Chờ TP duyệt." / "… Chờ BGĐ duyệt."'),
   (' => Mã báo giá => Duyệt | Duyệt & chuyển BGĐ | BGĐ duyệt',
    [('21-duyet-tp.png', 'Chi tiết báo giá Chờ TP duyệt (C2) — nút Duyệt, Từ chối'),
     ('21b-xac-nhan-duyet.png', 'Hộp Xác nhận duyệt'),
     ('21d-bgd-duyet.png', 'Chi tiết báo giá Chờ BGĐ duyệt — nút BGĐ duyệt')]),
   [('Nút Duyệt / Duyệt & chuyển BGĐ', 'Button', 'Enable / Ẩn', 'Ẩn', 'Hiện khi có Q3 và báo giá Chờ TP duyệt; nhãn '
     '"Duyệt & chuyển BGĐ" với C3.'),
    ('Nút BGĐ duyệt', 'Button', 'Enable / Ẩn', 'Ẩn', 'Hiện khi có Q4 và báo giá Chờ BGĐ duyệt.'),
    ('Hộp Xác nhận duyệt', 'Modal', 'Hiển thị', 'Ẩn', 'Nội dung "Duyệt báo giá?" / "Duyệt & chuyển BGĐ?" / "BGĐ duyệt '
     'báo giá?"; nút Duyệt, Huỷ.'),
   ],
   [('Bấm Duyệt (C2)', 'Click',
     'Before:\n– Kiểm tra Q3, quản lý phòng ban của báo giá, trạng thái Chờ TP duyệt.\n' + NO_PERM + '\n'
     'After:\n– Trạng thái → Đã duyệt, ghi người / ngày duyệt; YCBG → Đã có báo giá; dự án → Thương thảo giá; giải '
     'pháp → Đã duyệt giá; đồng bộ ERP.\n– Ghi lịch sử "TP duyệt"; thông báo cho người lập, Sale phụ trách, TP duyệt '
     'giá của phòng.\n– "Đã duyệt báo giá", quay về danh sách nguồn.'),
    ('Bấm Duyệt & chuyển BGĐ (C3)', 'Click',
     'After:\n– Trạng thái → Chờ BGĐ duyệt, ghi TP duyệt; ghi lịch sử "TP duyệt & chuyển BGĐ"; thông báo cho BGĐ duyệt '
     'giá cùng công ty "… TP đã duyệt & chuyển BGĐ <mã>".\n– "Đã duyệt báo giá".'),
    ('Bấm BGĐ duyệt', 'Click',
     'Before:\n– Kiểm tra Q4, cùng công ty, trạng thái Chờ BGĐ duyệt.\n'
     'After:\n– Như duyệt C2 (Đã duyệt + cập nhật dự án, giải pháp, YCBG, đồng bộ ERP); ghi lịch sử "BGĐ duyệt".\n'
     '– "BGĐ đã duyệt báo giá".'),
    ('Lỗi khi duyệt', 'System', 'After:\n– Hiển thị nội dung lỗi hoặc "Lỗi duyệt".'),
   ], mode='confirm')

# --------------------------------------------------------------------- FR-18
fr('FR-18', 'Từ chối báo giá', 'action', A_DUYET, 'detail',
   '- Thông báo, Validate dữ liệu. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Người duyệt trả báo giá về Đang tạo kèm lý do; người lập sửa rồi gửi duyệt lại.',
        tacnhan='Trưởng phòng duyệt giá (Q3), Ban giám đốc duyệt giá (Q4)',
        dieukien='(Q3 + Chờ TP duyệt + quản lý phòng ban) hoặc (Q4 + Chờ BGĐ duyệt + cùng công ty).',
        chinh='1. Bấm "Từ chối" ở chân màn chi tiết.\n'
              '2. Popup "Từ chối báo giá": nhập "Lý do từ chối".\n'
              '3. Bấm "Xác nhận từ chối" → báo giá về Đang tạo; "Đã từ chối báo giá"; quay về danh sách nguồn.',
        phu='• Lý do trống → "Vui lòng nhập lý do từ chối".\n'
            '• Sai trạng thái → "Báo giá không ở trạng thái chờ duyệt."\n'
            '• Lỗi → nội dung lỗi hoặc "Không thể từ chối".'),
   (' => Mã báo giá => Từ chối', [('22-tu-choi.png', 'Popup Từ chối báo giá'),
                                  ('22b-tu-choi-loi.png', 'Bấm Xác nhận từ chối khi chưa nhập lý do')]),
   [('Lý do từ chối', 'Textarea', 'Enable', '≤ 1000 ký tự', 'Có', 'Trống', 'Gợi ý "Nhập lý do từ chối...".'),
    ('Thông báo lỗi', 'Toast / Alert', 'Hiển thị', '–', '–', 'Ẩn', '"Vui lòng nhập lý do từ chối".'),
    ('Nút Xác nhận từ chối', 'Button', 'Enable / Disable', '–', '–', 'Hiển thị', '"Đang gửi..." khi đang xử lý.'),
    ('Nút Huỷ', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
   ],
   [('Bấm Xác nhận từ chối', 'Click',
     'Before:\n– Kiểm tra quyền + phòng ban / công ty như FR-17.\n' + NO_PERM + '\n'
     'During:\n– Lý do trống → "Vui lòng nhập lý do từ chối", dừng.\n'
     'After:\n– Trạng thái → Đang tạo; lưu lý do; xoá ngày gửi, TP duyệt, người duyệt, cấp duyệt.\n– Ghi lịch sử "Từ '
     'chối" kèm lý do.\n– Thông báo cho người lập (và TP đã duyệt nếu từ chối ở bước BGĐ): "<tên> từ chối báo giá '
     '<mã>. Lý do: <lý do>".\n– "Đã từ chối báo giá", quay về danh sách nguồn.'),
   ])

# --------------------------------------------------------------------- FR-19
fr('FR-19', 'Xóa báo giá', 'action', A_SALE, 'delete',
   '- Quy tắc Xóa, Thông báo. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Xoá hẳn báo giá đang tạo cùng toàn bộ dòng giá và lịch sử.',
        tacnhan='Người lập báo giá',
        dieukien='Người lập, báo giá Đang tạo.',
        chinh='1. Bấm biểu tượng Xóa trên dòng (hoặc "Xoá" ở chân màn chi tiết).\n'
              '2. Hộp "Xác nhận xoá": "Bạn có chắc muốn xoá báo giá \'<mã>\'?".\n'
              '3. Bấm Xoá → "Đã xoá báo giá"; danh sách nạp lại (từ màn chi tiết thì về danh sách).',
        phu='• Bấm Huỷ → không xoá.\n'
            '• Sai trạng thái / không phải người lập → "Báo giá không ở trạng thái Đang tạo, không thể sửa." / "Chỉ '
            'người tạo báo giá mới có thể sửa."\n'
            '• Lỗi khác → nội dung lỗi hoặc "Xoá thất bại".'),
   (' => Xóa', [('16-xoa.png', 'Hộp Xác nhận xoá')], 'Hoặc ở màn chi tiết: Mã báo giá => Xoá.'),
   [('Tiêu đề', 'Label', 'Hiển thị', 'Xác nhận xoá', '–'),
    ('Nội dung', 'Label', 'Hiển thị', 'Theo báo giá', '"Bạn có chắc muốn xoá báo giá \'<mã>\'?"'),
    ('Nút Xoá', 'Button', 'Enable', 'Hiển thị', 'Màu đỏ.'),
    ('Nút Huỷ', 'Button', 'Enable', 'Hiển thị', '–'),
   ],
   [('Bấm Xoá', 'Click',
     'Before:\n' + OWNER + '\n'
     'After:\n– Xoá báo giá, các dòng giá và lịch sử báo giá.\n– Báo giá lập từ YCBG và YCBG không còn báo giá nào '
     '→ YCBG về "Chờ xây dựng giá".\n– "Đã xoá báo giá"; lỗi → "Xoá thất bại".'),
    ('Bấm Huỷ', 'Click', 'After:\n– Đóng hộp, không xoá.'),
   ], mode='confirm')

# --------------------------------------------------------------------- FR-20
fr('FR-20', 'Sao chép báo giá', 'crud', A_SALE, 'create',
   '- Màn Thêm mới, Thông báo. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Tạo phiên bản mới từ 1 báo giá ở bất kỳ trạng thái: mở màn Tạo báo giá điền sẵn nội dung báo giá nguồn '
        'với giá, VAT ERP và tỷ giá hiện hành; chỉ tạo bản ghi khi bấm Lưu.',
        tacnhan='Sale phụ trách dự án / Người có quyền xây dựng giá (Q1, Q2)',
        dieukien='Báo giá thuộc phạm vi xem. Báo giá tự lập: người dùng là Sale phụ trách dự án; báo giá từ YCBG: có '
                 'Q1 / Q2 theo kiểu triển khai dự án.',
        chinh='1. Bấm "Sao chép báo giá" (menu ⋮ ở danh sách) hoặc "Sao chép" (chân màn chi tiết).\n'
              '2. Hệ thống so dữ liệu ERP: không thay đổi → mở thẳng màn Tạo; có thay đổi → popup "Phát hiện thay đổi '
              'dữ liệu từ ERP".\n'
              '3. Màn Tạo hiện dòng "Sao chép từ <mã> — báo giá mới chỉ được tạo sau khi bấm Lưu".\n'
              '4. Người dùng có thể đổi Dự án, sửa giá rồi Lưu nháp / Gửi duyệt.',
        phu='• Popup thay đổi ERP: bảng Loại thay đổi (Thay đổi giá / Thay đổi VAT / Thay đổi cấu trúc), Mã / Tên vật '
            'tư, Thông tin cũ (V1), Thông tin mới (Cập nhật), Hành động hệ thống; nút "Xác nhận Sao chép báo giá", "Hủy '
            'bỏ". Không có Q5 → không hiện dòng thay đổi giá vốn.\n'
            '• Không đủ quyền → "Bạn không phải Sale phụ trách dự án này" / "Bạn không có quyền xây dựng giá bán theo '
            '…".\n'
            '• Rời màn khi chưa lưu bản sao → cảnh báo chưa lưu.'),
   (' => Sao chép báo giá', [('18b-ban-sao.png', 'Màn Tạo báo giá điền sẵn từ báo giá nguồn')],
    'Hoặc ở màn chi tiết: Mã báo giá => Sao chép.'),
   [('Nút Sao chép báo giá / Sao chép', 'Icon Button', 'Enable / Ẩn', '–', '–', 'Theo quyền', 'Mọi trạng thái.'),
    ('Popup Phát hiện thay đổi dữ liệu từ ERP', 'Modal', 'Hiển thị', '–', '–', 'Ẩn khi không có thay đổi',
     'Dòng phụ "Báo giá này kế thừa cấu trúc từ <mã>. Hệ thống đã tự động cập nhật giá mới nhất và phát hiện các thay '
     'đổi sau".'),
    ('Dòng Sao chép từ', 'Text', 'Read-only', '–', '–', 'Mã báo giá nguồn', '–'),
    ('Mã báo giá', 'Text', 'Read-only', '–', '–', '(Chưa tạo)', 'Mã mới cấp khi lưu.'),
    ('Dự án', 'Dropdown', 'Enable', 'Dự án mình phụ trách chưa có BOM tổng hợp đã duyệt (+ dự án nguồn)', 'Có',
     'Dự án nguồn', '–'),
    ('Các trường còn lại', 'Theo FR-06', 'Enable', '–', '–', 'Theo báo giá nguồn', '–'),
   ],
   [('Bấm Sao chép', 'Click',
     'Before:\n– Báo giá ngoài phạm vi → "Không tìm thấy báo giá."\nAfter:\n– Lấy danh sách thay đổi ERP; lỗi → '
     '"Không kiểm tra được thay đổi từ ERP".\n– Không thay đổi → mở màn Tạo; có → mở popup.'),
    ('Bấm Xác nhận Sao chép báo giá', 'Click', 'After:\n– Mở màn Tạo điền sẵn: bản sao luôn là báo giá tự lập, giữ '
     'nhóm, hàng hoá (mã hàng tạm giữ nguyên), dịch vụ, giảm giá; lấy lại giá + VAT ERP và tỷ giá hiện tại; không kế '
     'thừa YCBG, BOM, giải pháp; xoá vết duyệt.\n– Lỗi → nội dung lỗi hoặc "Không tải được dữ liệu sao chép", về danh '
     'sách.'),
    ('Bấm Lưu nháp / Gửi duyệt trên bản sao', 'Click', 'After:\n– Tạo báo giá mới Đang tạo, ghi nguồn sao chép; tiếp '
     'theo như FR-13 / FR-14.'),
   ])

# --------------------------------------------------------------------- FR-21
fr('FR-21', 'In báo giá', 'io', A_ALL, 'list',
   '- Màn In / Xem trước bản in. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Chọn cột cần in rồi xem trước và in báo giá theo mẫu "BÁO GIÁ HÀNG HÓA" có letterhead công ty.',
        tacnhan='Người lập báo giá, Người duyệt giá; Người dùng đã đăng nhập',
        dieukien='Báo giá thuộc phạm vi xem.',
        chinh='1. Bấm "In báo giá" (menu ⋮) hoặc "In" (chân màn chi tiết).\n'
              '2. Popup "Cấu hình in báo giá": tích "Hiện hàng hoá cấp con", chọn cột (mặc định tất cả trừ Mã hàng hoá).\n'
              '3. Bấm "Xem trước" → cửa sổ "Xem trước báo giá"; bấm "In".',
        phu='• Không chọn cột nào → "Vui lòng chọn ít nhất 1 cột để in".\n'
            '• Không tải được báo giá (từ danh sách) → "Không tải được chi tiết báo giá".'),
   (' => In báo giá', [('17-in-cauhinh.png', 'Popup Cấu hình in báo giá'),
                       ('17b-in-xemtruoc.png', 'Xem trước báo giá')],
    'Hoặc ở màn chi tiết: Mã báo giá => In.'),
   [('Hiện hàng hoá cấp con', 'Checkbox', 'Enable', '–', 'Không', 'Đã tích', '–'),
    ('Chọn tất cả', 'Checkbox', 'Enable', '–', 'Không', 'Theo danh sách', '–'),
    ('Danh sách cột', 'Checkbox', 'Enable', '17–20 cột theo phương thức GG', 'Có (≥ 1)', 'Tất cả trừ Mã hàng hoá',
     'STT, Tên hàng hoá, Mã hàng hoá, Model, Thương hiệu, Xuất xứ, Đơn vị tính, Thông số kỹ thuật, Ghi chú, Số lượng, '
     'Đơn giá bán, Thành tiền bán, [GG (%), GG (₫), Đơn giá sau GG] hoặc [GG phân bổ], VAT (%), Tiền VAT, Thành tiền '
     'sau VAT, Hình ảnh, Thời gian bảo hành.'),
    ('Nút Huỷ / Xem trước', 'Button', 'Enable', '–', '–', 'Hiển thị', '–'),
    ('Bản xem trước', 'Modal', 'Read-only', '–', '–', '–', 'Letterhead công ty; Kính gửi, Tên đơn vị, Điện thoại / '
     'Email, Hiệu lực, Bảo hành, Báo giá số, Dự án, Ngày, Đại diện kinh doanh, Giao hàng, Loại tiền tệ; bảng hàng theo '
     'nhóm; nút In.'),
   ],
   [('Bấm Xem trước', 'Click', 'During:\n– Chưa chọn cột → "Vui lòng chọn ít nhất 1 cột để in".\nAfter:\n– Mở cửa '
     'sổ Xem trước báo giá theo cấu hình.'),
    ('Bấm In', 'Click', 'After:\n– Mở hộp thoại in của trình duyệt.'),
   ])

# --------------------------------------------------------------------- FR-22
fr('FR-22', 'Xuất Excel báo giá', 'io', A_ALL, 'excel',
   '- Quy tắc Excel. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Xuất dữ liệu 1 báo giá ra file Excel theo đúng mẫu Import (dùng lại để nhập ngược).',
        tacnhan='Người lập báo giá, Người duyệt giá; Người dùng đã đăng nhập',
        dieukien='Báo giá thuộc phạm vi xem; ở màn Sửa chỉ khi báo giá đã lưu.',
        chinh='1. Bấm "Xuất Excel" ở chân màn chi tiết (hoặc trên bảng màn Làm giá).\n'
              '2. Hệ thống tải file <mã báo giá>_<dd-mm-yyyy>.xlsx.',
        phu='• Lỗi → "Không xuất được file Excel. Vui lòng thử lại."\n'
            '• Không có Q5 → các cột giá vốn trong file để trống.'),
   (' => Mã báo giá => Xuất Excel', [('15-xem-chi-tiet.png', 'Nút Xuất Excel ở chân màn chi tiết')],
    'Hoặc ở màn Làm giá: nút Xuất Excel trên bảng Chi tiết sản phẩm.'),
   [('Nút Xuất Excel', 'Button', 'Enable / Disable', '–', '–', 'Hiển thị', '"Đang xuất..." khi đang tải.'),
    ('File Excel', 'File', '–', '22 / 25 / 24 cột', '–', '–', 'Bộ cột theo phương thức GG: Không GG / GG mặt hàng / '
     'GG tổng.'),
   ],
   [('Bấm Xuất Excel', 'Click', 'Before:\n– Kiểm tra phạm vi xem.\nAfter:\n– Tải file <mã>_<dd-mm-yyyy>.xlsx.\n– Lỗi '
     '→ "Không xuất được file Excel. Vui lòng thử lại."'),
   ])

# --------------------------------------------------------------------- FR-23
fr('FR-23', 'Ghi chú kinh doanh', 'crud', A_SALE, 'history',
   '- Thông báo, Quy tắc ghi lịch sử. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Sale phụ trách dự án ghi chú kinh doanh trên báo giá đã duyệt (theo dõi thương thảo với khách).',
        tacnhan='Sale phụ trách dự án',
        dieukien='Báo giá Đã duyệt; người dùng là Sale phụ trách dự án.',
        chinh='1. Ở màn chi tiết, mục "Ghi chú Kinh doanh" hiện ô nhập.\n'
              '2. Nhập nội dung, bấm "Lưu ghi chú".\n'
              '3. Hệ thống lưu, ghi lịch sử, báo "Đã lưu ghi chú".',
        phu='• Người khác → chỉ xem nội dung (hoặc "—").\n'
            '• Sai trạng thái → "Chỉ có thể sửa ghi chú khi báo giá đã duyệt."\n'
            '• Không phải Sale phụ trách → "Chỉ NV Kinh doanh phụ trách dự án mới có thể cập nhật ghi chú."'),
   (' => Mã báo giá => Lưu ghi chú', [('23-ghi-chu-kd.png', 'Mục Ghi chú Kinh doanh trên báo giá Đã duyệt')]),
   [('Ghi chú Kinh doanh', 'Textarea', 'Enable / Read-only', '–', 'Không', 'Theo dữ liệu',
     'Gợi ý "Nhập ghi chú kinh doanh..."; chỉ hiện khi Đã duyệt.'),
    ('Nút Lưu ghi chú', 'Button', 'Enable / Ẩn', '–', '–', 'Hiện với Sale phụ trách', '"Đang lưu..." khi xử lý.'),
   ],
   [('Bấm Lưu ghi chú', 'Click', 'Before:\n– Kiểm tra Đã duyệt + Sale phụ trách.\n' + NO_PERM + '\nAfter:\n– Lưu '
     'ghi chú; ghi lịch sử "Cập nhật ghi chú KD" (cũ → mới); "Đã lưu ghi chú".\n– Lỗi → nội dung lỗi hoặc "Lỗi lưu".'),
   ])

# --------------------------------------------------------------------- FR-24
fr('FR-24', 'Lập hợp đồng', 'action', A_SALE, 'detail',
   '- Màn Xem chi tiết và Phân quyền. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Từ báo giá Trúng thầu, người lập mở màn lập Hợp đồng (dự án TKT) điền sẵn theo báo giá; khi đã có hợp '
        'đồng thì nút đổi thành "Xem hợp đồng".',
        tacnhan='Người lập báo giá',
        dieukien='Báo giá ở trạng thái Trúng thầu (chốt ở tab Báo giá của dự án TKT, xem BR-14), chưa có hợp đồng, '
                 'người dùng là người lập.',
        chinh='1. Mở chi tiết báo giá Trúng thầu.\n'
              '2. Bấm "Lập hợp đồng" → chuyển sang màn tạo Hợp đồng gắn báo giá này.',
        phu='• Báo giá đã có hợp đồng → nút "Xem hợp đồng" mở hợp đồng.\n'
            '• Không phải người lập / chưa Trúng thầu → không hiện nút.'),
   (' => Mã báo giá => Lập hợp đồng', [('24-lap-hop-dong.png', 'Báo giá Trúng thầu — nút Lập hợp đồng')]),
   [('Nút Lập hợp đồng', 'Button', 'Enable / Ẩn', '–', '–', 'Ẩn', 'Hiện khi Trúng thầu, chưa có hợp đồng, người lập.'),
    ('Nút Xem hợp đồng', 'Button', 'Enable / Ẩn', '–', '–', 'Ẩn', 'Hiện khi đã có hợp đồng.'),
   ],
   [('Bấm Lập hợp đồng', 'Click', 'After:\n– Mở màn tạo Hợp đồng với báo giá đã chọn (đặc tả ở SRS Hợp đồng).'),
    ('Bấm Xem hợp đồng', 'Click', 'After:\n– Mở màn chi tiết hợp đồng.'),
   ])

# --------------------------------------------------------------------- FR-25
fr('FR-25', 'Xem lịch sử', 'view', A_ALL, 'history',
   '- Quy tắc ghi lịch sử. Chỉ bổ sung các quy tắc riêng tại phần mô tả chi tiết.',
   dict(mota='Xem dòng thời gian thao tác của báo giá: popup "Lịch sử báo giá" ở danh sách và khối "Lịch sử" cuối màn '
        'chi tiết.',
        tacnhan='Người dùng đã đăng nhập',
        dieukien='Báo giá thuộc phạm vi xem.',
        chinh='1. Ở danh sách bấm "Lịch sử phê duyệt" (menu ⋮) → popup "Lịch sử báo giá".\n'
              '2. Ở màn chi tiết bấm "Xem lịch sử" ở khối Lịch sử → nạp lần đầu, có Bộ lọc và Làm mới.\n'
              '3. Mỗi mục: hành động (màu), cấp, người thực hiện, trạng thái cũ → mới, lý do, thay đổi trường / sản phẩm / '
              'dịch vụ / giảm giá, thời điểm.',
        phu='• Chưa có lịch sử → "Chưa có lịch sử" / "Chưa có lịch sử thao tác nào."\n'
            '• Lỗi tải → thông báo đỏ + nút "Thử lại".'),
   (' => Lịch sử phê duyệt', [('20-lichsu-popup.png', 'Popup Lịch sử báo giá'),
                              ('20b-lich-su-chi-tiet.png', 'Khối Lịch sử ở màn chi tiết')],
    'Hoặc ở màn chi tiết: Mã báo giá => Xem lịch sử.'),
   [('Hành động', 'Badge', 'Read-only', '15 loại', 'Theo dữ liệu', 'Tạo báo giá, Lưu nháp, Gửi duyệt, Tự duyệt, TP '
     'duyệt, TP duyệt & chuyển BGĐ, BGĐ duyệt, Từ chối, Cập nhật ghi chú KD, Áp dụng VAT đồng loạt, Import giá, Đóng '
     'theo dự án, Cập nhật giảm giá, Chốt báo giá, Hủy chốt.'),
    ('Người thực hiện', 'Text', 'Read-only', '–', 'Theo dữ liệu', '"Hệ thống" nếu tự động.'),
    ('Trạng thái cũ → mới', 'Text', 'Read-only', '–', 'Theo dữ liệu', 'Chỉ khi trạng thái đổi.'),
    ('Lý do / Thay đổi', 'Text', 'Read-only', '–', 'Theo dữ liệu', 'Giá trị cũ → mới; "(trống)" nếu rỗng; sản phẩm '
     'sửa, dịch vụ thêm / xoá, giảm giá.'),
    ('Thời điểm', 'Text', 'Read-only', 'dd/mm/yyyy hh:mm', 'Theo dữ liệu', '–'),
    ('Nút Đóng / Xem lịch sử / Thu gọn / Làm mới / Bộ lọc', 'Button', 'Enable', '–', 'Hiển thị', '–'),
   ],
   [('Bấm Lịch sử phê duyệt / Xem lịch sử', 'Click', 'After:\n– Nạp lịch sử báo giá, mới nhất ở trên.'),
    ('Bấm Làm mới / Bộ lọc', 'Click', 'After:\n– Nạp lại / lọc theo loại hành động, người thực hiện, thời gian.'),
   ], mode='ro', usecase=False)

# ======================================================================= PHẦN 4
d.h1('Phần 4. Quy tắc nghiệp vụ')
d.rule_ref('. Phần này chỉ ghi các quy tắc đặc thù của màn Báo giá; không lặp lại các quy tắc đã có trong SRS quy '
           'tắc chung.', anchor='list', head='Quy tắc áp dụng',
           lead='Các quy tắc nghiệp vụ dùng chung được định nghĩa tại SRS Các quy tắc chung ')
d.rule_table([
    ('BR-01', 'Phạm vi xem báo giá', [
        '– V1 xem tất cả; V2 công ty đang làm việc; V3 phòng ban / bộ phận mình quản lý; V4 bộ phận mình quản lý; '
        'V2–V4 luôn thấy thêm báo giá do mình lập; không có V1–V4 chỉ thấy báo giá do mình lập.',
        '– Người duyệt mở được báo giá đang chờ chính mình duyệt dù thiếu V1–V4.',
        '– Danh sách, chi tiết, sao chép, xuất Excel cùng phạm vi; ngoài phạm vi báo "Không tìm thấy báo giá."',
    ], ['Xem danh sách', 'Xem chi tiết', 'Xuất Excel', 'Sao chép']),
    ('BR-02', 'Ai được lập báo giá', [
        '– Báo giá tự lập: chỉ Sale phụ trách dự án.',
        '– Báo giá từ YCBG: dự án triển khai theo phòng cần Q2 và YCBG cùng phòng ban; dự án khác cần Q1.',
        '– Dự án cha không lập báo giá thường (dùng Báo giá tổng).',
    ], ['Tạo mới', 'Sao chép']),
    ('BR-03', 'Chỉ người lập sửa báo giá Đang tạo', [
        '– Sửa, Lưu nháp, Gửi duyệt, Import, Xoá: chỉ người lập khi báo giá Đang tạo (kể cả bị từ chối).',
        '– Trạng thái khác: không hiện nút; màn Sửa khoá toàn bộ.',
    ], ['Chỉnh sửa', 'Xóa', 'Gửi duyệt', 'Import Excel']),
    ('BR-04', 'Báo giá từ BOM và báo giá tự lập', [
        '– Chọn dự án: có 1 BOM tổng hợp đã duyệt → tự nạp; nhiều → chọn; không có → tự lập.',
        '– Báo giá từ BOM khoá cấu trúc: chỉ sửa giá, VAT hàng tạm, ghi chú, giảm giá, vận chuyển và dịch vụ.',
        '– Chọn lại dự án / BOM xoá toàn bộ hàng đang có (có hỏi xác nhận).',
    ], ['Tạo mới', 'Làm giá', 'Thêm hàng hoá']),
    ('BR-05', 'Giá và VAT hàng ERP', [
        '– Giá bán theo Bảng giá đã chọn, giá nhập theo giá vốn ERP, quy đổi theo tỷ giá chốt lúc tạo; VAT theo danh '
        'mục; không sửa tay.',
        '– Khi lưu, hệ thống áp lại giá + VAT ERP; mở màn Sửa có lệch thì hỏi cập nhật, Từ chối thì giữ đơn giá đã lưu.',
        '– Không có Q5: giá vốn hàng ERP hiển thị "—".',
    ], ['Làm giá', 'Chỉnh sửa']),
    ('BR-06', 'Công thức tổng và tỷ suất lợi nhuận', [
        '– Thành tiền bán = Giá bán × SL − giảm giá dòng (GG mặt hàng: GG₫ × SL; GG tổng: Phân bổ GG).',
        '– VAT tính trên thành tiền sau giảm giá; chỉ dòng cha / độc lập; vận chuyển VAT 8%.',
        '– Tổng nhập = hàng hoá + dịch vụ + giá nhập vận chuyển; TSLN = (Tổng sau GG − Tổng nhập) / Tổng nhập × 100.',
        '– TSLN dưới mức sàn hiển thị đỏ.',
    ], ['Làm giá', 'Tổng hợp giá trị']),
    ('BR-07', 'Xác định cấp duyệt', [
        '– Cấp = cấp CAO NHẤT của 3 điều kiện trong Cấu hình duyệt giá: giá trị đơn hàng sau GG quy VNĐ (trước thuế), '
        'tỷ suất lợi nhuận tổng, tỷ suất lợi nhuận dòng hàng tạm thấp nhất; không khớp → Cấp 3.',
        '– Báo giá chỉ gồm hàng / dịch vụ ERP, đơn giá bán > 1.000, không giảm giá → tự động duyệt (Cấp 1).',
        '– Chân trang màn Làm giá chỉ ước tính theo 2 điều kiện đầu; cấp chính thức tính lại khi gửi duyệt.',
    ], ['Tổng hợp giá trị', 'Gửi duyệt']),
    ('BR-08', 'Điều kiện gửi duyệt', [
        '– Bắt buộc: giao hàng > 0, tiền tệ, bảng giá, giai đoạn dự án, điều khoản báo giá.',
        '– Mọi dòng: giá bán và thành tiền bán > 0; hàng tạm giá nhập > 0; giá cha ≥ tổng con; GG ≤ đơn giá; dịch vụ có '
        'tên, SL > 0, giá bán > 0.',
        '– Phải khớp Báo giá Mỏ neo (dự án con); đơn giá ≤ 1.000 chỉ cảnh báo.',
        '– Lưu nháp chỉ kiểm VAT, % GG trong 0–100 và khoản GG tổng.',
    ], ['Lưu nháp', 'Gửi duyệt']),
    ('BR-09', 'Luồng trạng thái', [
        '– Đang tạo → (C1) Đã duyệt; → (C2/C3) Chờ TP duyệt → (C2) Đã duyệt / (C3) Chờ BGĐ duyệt → Đã duyệt.',
        '– TP duyệt khi quản lý phòng ban của báo giá; BGĐ duyệt khi cùng công ty.',
        '– Từ chối ở bước chờ duyệt → Đang tạo, xoá vết duyệt, giữ lý do.',
        '– Đóng, Dừng do hệ thống đặt theo dự án / giải pháp, không thao tác trên màn này.',
    ], ['Gửi duyệt', 'Duyệt', 'Từ chối']),
    ('BR-10', 'Tác động khi báo giá được duyệt', [
        '– Dự án TKT → Thương thảo giá; giải pháp → Đã duyệt giá; YCBG → Đã có báo giá; đồng bộ sang ERP.',
        '– Xoá báo giá lập từ YCBG mà YCBG không còn báo giá → YCBG về Chờ xây dựng giá.',
    ], ['Duyệt', 'Gửi duyệt', 'Xóa']),
    ('BR-11', 'Thông báo', [
        '– Gửi duyệt: TP duyệt giá quản lý phòng ban của báo giá.',
        '– TP duyệt & chuyển BGĐ: BGĐ duyệt giá cùng công ty.',
        '– Duyệt xong: người lập, Sale phụ trách, TP duyệt giá của phòng.',
        '– Từ chối: người lập (+ TP đã duyệt nếu từ chối ở bước BGĐ), kèm lý do.',
    ], ['Gửi duyệt', 'Duyệt', 'Từ chối']),
    ('BR-12', 'Giảm giá', [
        '– Chỉ người có Q6 được đổi phương thức / số tiền giảm giá; không có Q6 vẫn lưu được báo giá giữ nguyên GG cũ.',
        '– Dự án con kế thừa phương thức GG của dự án cha.',
        '– GG tổng: % tính trên tổng trước VAT gồm vận chuyển; phân bổ theo tỷ lệ giá trị từng dòng.',
    ], 'Áp dụng giảm giá'),
    ('BR-13', 'Sao chép báo giá', [
        '– Bản sao luôn là báo giá tự lập, mã mới cấp khi Lưu; lấy giá + VAT ERP và tỷ giá hiện tại; không kế thừa YCBG, '
        'BOM, giải pháp, vết duyệt.',
        '– Chỉ đổi sang dự án chưa có BOM tổng hợp đã duyệt.',
    ], 'Sao chép'),
    ('BR-14', 'Trúng thầu và hợp đồng', [
        '– Chốt / Hủy chốt Trúng thầu thực hiện ở tab Báo giá của dự án TKT: chỉ báo giá Đã duyệt; mỗi dự án 1 báo giá '
        'trúng thầu ("Dự án đã có báo giá trúng thầu, vui lòng hủy chốt trước."); bắt buộc file xác nhận khách hàng; '
        'hủy chốt bắt buộc lý do, báo giá về Đã duyệt.',
        '– Lập hợp đồng chỉ từ báo giá Trúng thầu, do người lập báo giá, khi chưa có hợp đồng.',
    ], ['Lập hợp đồng', 'Lịch sử']),
    ('BR-15', 'Hiệu lực báo giá', [
        '– = Hôm nay + số ngày hiệu lực cấu hình (mặc định 30 ngày), rút về ngày đổi giá sớm nhất của hàng ERP cấp cha; '
        'tính lại khi gửi duyệt.',
    ], ['Tạo mới', 'Gửi duyệt']),
])

d.save(update_fields=os.environ.get('SRS_NO_WORD') != '1')
