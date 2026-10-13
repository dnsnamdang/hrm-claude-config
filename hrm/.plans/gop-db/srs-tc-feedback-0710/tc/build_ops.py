# plans_<book>/NN.json + workbook đã tải -> ops/<book>_NN/{index.json, e00.json.., 00.json..} + runner JS
# python3 build_ops.py cd 16      (book: cd | qlda)
import json, sys, os, re
import openpyxl
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from gen_js import tsv

S = os.path.dirname(os.path.abspath(__file__))
SRC = {'cd': '../src/tc_chuyendoi.xlsx', 'qlda': '../src/tc_qlda.xlsx'}
JDIR = '/Users/manhcuong/Desktop/dns/HRM/.playwright-mcp/tcw2'
_wb = {}


def wb(book):
    if book not in _wb:
        _wb[book] = openpyxl.load_workbook(os.path.join(S, SRC[book]))
    return _wb[book]


def cval(ws, r, c):
    v = ws.cell(r, c).value
    return '' if v is None else str(v)


def merge_of(ws, col, r):
    for m in ws.merged_cells.ranges:
        if m.min_col <= col <= m.max_col and m.min_row <= r <= m.max_row:
            return m
    return None


def is_case(v):
    return bool(re.match(r'^(TC|=)', v or ''))


IDC_G = [3]


def group_start(ws, r):
    x = r
    while x > 7 and is_case(cval(ws, x - 1, IDC_G[0])):
        x -= 1
    return x


from openpyxl.utils import get_column_letter as CL


def colmap(ws):
    for r in range(1, 30):
        hdr = {c: cval(ws, r, c).strip().lower() for c in range(1, 20)}
        if 'tc id' in hdr.values():
            m = {'hdr_row': r}
            for c, h in hdr.items():
                if h == 'tc id': m['id'] = c
                elif h.startswith('chức năng') and 'func' not in m: m['func'] = c
                elif h.startswith('priority'): m['priority'] = c
                elif h.startswith('tiền điều kiện'): m['pre'] = c
                elif h.startswith('bước'): m['steps'] = c
                elif h.startswith('test data') and 'data' not in m: m['data'] = c
                elif h.startswith('expected'): m['expected'] = c
                elif 'check' in h: m.setdefault('checks', []).append(c)
            return m
    raise SystemExit('không thấy hàng TC ID')


def build(book, nn):
    p = json.load(open(os.path.join(S, f'plans_{book}', f'{nn}.json')))
    w = wb(book)
    ws = w.worksheets[int(nn)]
    assert ws.title.strip() == p['tab'].strip(), (ws.title, p['tab'])
    cm = colmap(ws)
    IDC = cm['id']; IDC_G[0] = IDC; IDL = CL(IDC)
    width = cm['expected'] - IDC + 1
    chk = f"{CL(min(cm['checks']))}{{r}}:{CL(max(cm['checks']))}{{r}}" if cm.get('checks') else None
    def rowvals(c):
        a = [''] * width
        for k in ('id', 'func', 'priority', 'pre', 'steps', 'data', 'expected'):
            if k in cm: a[cm[k] - IDC] = c[k]
        return a
    tabkey = ws.title.strip()[:20]
    edits = []
    for e in p.get('edits', []):
        from openpyxl.utils import column_index_from_string as CI
        col = CI(e['col'])
        cur = cval(ws, e['row'], col)
        assert cur.replace('\r', '').startswith(e['old_prefix'][:30].replace('\r', '')), (nn, e['row'], cur[:60], e['old_prefix'][:60])
        edits.append({'row': e['row'], 'cell': f"{e['col']}{e['row']}", 'expect_c': cval(ws, e['row'], IDC), 'idcol': IDL,
                      'expect_prefix': cur[:25], 'text': tsv([[e['new']]]), 'clear': chk.format(r=e['row']) if chk else None})
    ops = []
    nextval = {}
    order = sorted(enumerate(p['inserts']), key=lambda t: (-t[1]['after_row'], 0 if t[1].get('new_group') else 1, -t[0]))
    for _, ins in order:
        r = ins['after_row']
        cases = ins['cases']
        n = len(cases)
        if not n:
            continue
        expect = cval(ws, r, IDC)
        assert is_case(expect), (nn, r, expect)
        rows = [rowvals(c) for c in cases]
        op = {'tab': tabkey, 'idcol': IDL, 'after_row': r, 'expect_c': expect, 'expect_next': cval(ws, r + 1, IDC),
              'tsv': tsv(rows), 'last_id': cases[-1]['id'], 'first_id': cases[0]['id'], 'merges': [],
              'header_from': None, 'ab_tsv': None}
        if not ins.get('new_group'):
            op['n_rows'] = n
            op['paste_at'] = f'{IDL}{r + 1}'
            op['last_case_row'] = r + n
            for col, L in ((1, 'A'), (2, 'B')):
                m = merge_of(ws, col, r)
                if m:  # chỉ gộp lại khi ô gốc đang gộp; tab không gộp thì để trống A/B
                    op['merges'].append(f'{L}{m.min_row}:{L}{r + n}')
        else:
            h = group_start(ws, r) - 1
            op['n_rows'] = n + 1
            op['header_from'] = h
            op['header_title'] = ins['group_title']
            op['paste_at'] = f'{IDL}{r + 2}'
            op['first_case_row'] = r + 2
            op['last_case_row'] = r + 1 + n
            op['ab_tsv'] = tsv([[p.get('module', ''), ins['group_name']]])
            gm = merge_of(ws, 1, group_start(ws, r))
            op['merges'] = [f'A{r + 2}:A{r + 1 + n}', f'B{r + 2}:B{r + 1 + n}'] if (n > 1 and gm) else []
        if r in nextval:
            op['expect_next'] = nextval[r]
        nextval[r] = op['header_title'] if op['header_from'] else cases[0]['id']
        ops.append(op)
    out = os.path.join(S, 'ops', f'{book}_{nn}')
    os.makedirs(out, exist_ok=True)
    for f in os.listdir(out):
        os.remove(os.path.join(out, f))
    json.dump({'tab': tabkey, 'count': len(ops), 'ecount': len(edits)}, open(os.path.join(out, 'index.json'), 'w'), ensure_ascii=False)
    for i, e in enumerate(edits):
        json.dump(e, open(os.path.join(out, f'e{i:02d}.json'), 'w'), ensure_ascii=False)
    for i, o in enumerate(ops):
        json.dump(o, open(os.path.join(out, f'{i:02d}.json'), 'w'), ensure_ascii=False)
    os.makedirs(JDIR, exist_ok=True)
    tpl = open(os.path.join(S, 'runner2_tpl.js')).read()
    open(os.path.join(JDIR, f'run_{book}_{nn}.js'), 'w').write(tpl.replace('__KEY__', f'{book}_{nn}'))
    return edits, ops


if __name__ == '__main__':
    book = sys.argv[1]
    for nn in sys.argv[2:]:
        edits, ops = build(book, nn)
        print(book, nn, 'edits', len(edits), '| inserts', [(o['after_row'], o['n_rows'], o['first_id'], o['merges']) for o in ops])
