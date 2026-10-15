import openpyxl,sys
wb=openpyxl.load_workbook(sys.argv[1])
ws=wb[sys.argv[2]]
print('merged:',[str(m) for m in ws.merged_cells.ranges][:60])
last=0
for r in range(1,ws.max_row+1):
  vals=[ws.cell(r,c).value for c in range(1,ws.max_column+1)]
  if any(v not in (None,'') for v in vals): last=r
for r in range(1,last+1):
  vals=[ws.cell(r,c).value for c in range(1,min(ws.max_column,14)+1)]
  if any(v not in (None,'') for v in vals):
    print(r,'|',' | '.join((str(v).replace('\n','⏎')[:90] if v is not None else '') for v in vals))
