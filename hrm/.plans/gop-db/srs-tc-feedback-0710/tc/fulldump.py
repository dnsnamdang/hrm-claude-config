import openpyxl,sys,os,re
wb=openpyxl.load_workbook(sys.argv[1])
out=sys.argv[2]
for i,ws in enumerate(wb.worksheets[1:],1):
  lines=[]
  lines.append('TAB: '+ws.title)
  lines.append('MERGED A/B: '+', '.join(sorted([str(m) for m in ws.merged_cells.ranges if m.min_col<=2 and m.min_row>=17],key=lambda s:int(re.sub(r'\D','',s.split(':')[0])))))
  last=0
  for r in range(1,ws.max_row+1):
    if any(ws.cell(r,c).value not in (None,'') for c in range(1,15)): last=r
  for r in range(1,last+1):
    vals=[ws.cell(r,c).value for c in range(1,15)]
    if any(v not in (None,'') for v in vals):
      lines.append(f'R{r}: '+' || '.join(f'{chr(64+j)}={str(v)}' for j,v in enumerate(vals,1) if v not in (None,'') and not hasattr(v,'text')))
  open(os.path.join(out,f'{i:02d}.txt'),'w').write('\n'.join(lines))
  print(i,ws.title,'last',last)
