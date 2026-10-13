import openpyxl,json,glob,sys
old=openpyxl.load_workbook('tc/tc.xlsx'); new=openpyxl.load_workbook('tc3/tc.xlsx')
bad=0
def rows(ws):
    out=[]
    for r in range(17,ws.max_row+1):
        v=[ws.cell(r,c).value for c in range(3,15)]
        if v[0] is not None: out.append(tuple(None if x is None else str(x) for x in v))
    return out
for f in sorted(glob.glob('plans/*.json')):
    nn=int(f.split('/')[-1][:2]); p=json.load(open(f))
    wo,wn=old.worksheets[nn],new.worksheets[nn]
    ro,rn=rows(wo),rows(wn)
    planned={c['id'] for i in p['inserts'] for c in i['cases']}
    titles={i['group_title'] for i in p['inserts'] if i.get('new_group')}
    # case mới có trong bản mới
    ids_n=[r[0] for r in rn]
    miss=[x for x in planned if x not in ids_n]
    miss_t=[t for t in titles if t not in ids_n]
    # dòng cũ giữ nguyên (so toàn bộ C..N), theo thứ tự
    rest=[r for r in rn if not (r[0] in planned and r[0] not in [x[0] for x in ro]) and r[0] not in titles]
    changed= ro!=rest
    newK=[r for r in rn if r[0] in planned and r[0] not in [x[0] for x in ro] and any(r[i] for i in (7,8,9,10))]
    ab=[]
    status='OK' if not(miss or miss_t or changed or newK) else 'XEM'
    if status!='OK': bad+=1
    print(nn, wn.title[:28], status, 'thiếu',len(miss),'thiếu nhóm',len(miss_t),'cũ đổi',changed,'mới có J-M',len(newK))
    if changed:
        import difflib
        for d in list(difflib.unified_diff([str(x)[:120] for x in ro],[str(x)[:120] for x in rest],lineterm='',n=0))[:8]: print('   ',d)
print('BAD',bad)
