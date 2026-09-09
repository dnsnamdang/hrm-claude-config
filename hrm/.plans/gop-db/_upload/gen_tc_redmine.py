# -*- coding: utf-8 -*-
"""Sinh testcase cho 4 task Redmine, dong goi theo dung bo cot cua 2 tab dich.

File dich la GOOGLE SHEETS THAT (khong phai .xlsx) nen khong ghi cell truc tiep duoc — file
nay xuat ra mot .xlsx 2 sheet, moi sheet la KHOI DONG MOI dung thu tu cot cua tab tuong ung
de nguoi dung boi + dan thang vao dong trong dau tien.

    Tab "Dự án tiền khả thi" (gid 739303646) — cot A..I:
        Module | Nhóm chức năng | TC ID | Chức năng | Priority | Tiền điều kiện |
        Bước thực hiện | Test Data | Expected Result (chi tiết)
        -> task 10789, 10797, 10898; TC ID chay tiep TC-ROLE-340 (dong cuoi dang co la 339)

    Tab "11.Meeting" (gid 1614285068) — cot A..J (co 1 cot TRONG o vi tri C):
        Module | Nhóm chức năng | (trống) | TC ID | Chức năng | Priority | Tiền điều kiện |
        Bước thực hiện | Test Data | Expected Result (chi tiết)
        -> task 11014; TC ID chay tiep TC_17.001 (dong cuoi dang co la TC_16.012)

Nguon yeu cau: quanly.dnsmedia.vn/issues/{10789,10797,10898,11014} (doc 03/09/2026).

Chay:  python3 .plans/gop-db/_upload/gen_tc_redmine.py
"""
import io
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Testcase 4 task Redmine (dán vào Testcase _Quản lý dự án).xlsx")

FONT = 'Times New Roman'
THIN = Side(style='thin', color='BFBFBF')
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
HEAD = PatternFill('solid', fgColor='4472C4')
BANNER_TITLE = 'UPDATE NỘI DUNG 03/09/2026'

# ===================================================================== TASK 10789
T10789 = ('Giảm giá báo giá', 'http://quanly.dnsmedia.vn/issues/10789')
TC_10789 = [
    ('Quyền mới xuất hiện trong màn phân quyền', 'P0',
     'Đăng nhập bằng tài khoản Quản trị.',
     '1. Vào Cấu hình hệ thống → Phân quyền người dùng → mở một vai trò\n'
     '2. Tìm nhóm quyền của phân hệ Quản lý báo giá',
     'Vai trò: Nhân viên kinh doanh',
     '- Có quyền mới tên đúng “Cho phép thêm giảm giá trong báo giá”\n'
     '- Tích chọn và lưu vai trò thành công'),
    ('Gán quyền cho một người và bỏ quyền của người khác', 'P0',
     'Có 2 tài khoản kinh doanh: User A và User B, cùng lập được báo giá.',
     '1. Gán quyền “Cho phép thêm giảm giá trong báo giá” cho User A\n'
     '2. Bỏ quyền đó của User B\n3. Lưu lại',
     'User A: có quyền · User B: không có quyền',
     '- Lưu thành công, danh sách quyền của từng người hiển thị đúng như vừa gán'),
    ('Người CÓ quyền: ô GG mở và chọn được loại giảm giá', 'P0',
     'Đăng nhập User A, mở màn Tạo mới báo giá.',
     '1. Quan sát trường GG\n2. Mở danh sách chọn loại giảm giá',
     'User A',
     '- Trường GG bật, chọn được\n- Danh sách có đủ các loại giảm giá theo cấu hình'),
    ('Người CÓ quyền: nhập giảm giá theo tổng hóa đơn', 'P0',
     'User A đang ở màn Tạo mới báo giá, đã có ít nhất 2 dòng hàng.',
     '1. Chọn loại giảm giá theo tổng hóa đơn\n2. Nhập mức giảm\n3. Bấm Lưu',
     'Giảm giá tổng: 5%',
     '- Hiện đúng ô nhập tương ứng loại vừa chọn\n- Lưu thành công\n'
     '- Tổng tiền sau giảm giá tính đúng theo mức vừa nhập'),
    ('Người CÓ quyền: nhập giảm giá trên từng dòng hàng', 'P0',
     'Như trên.',
     '1. Chọn loại giảm giá theo mặt hàng\n2. Nhập mức giảm cho dòng thứ nhất\n3. Bấm Lưu',
     'Dòng 1: giảm 10%',
     '- Cột giảm giá của dòng cho nhập\n- Lưu thành công\n'
     '- Thành tiền của đúng dòng đó giảm theo, các dòng khác giữ nguyên'),
    ('Người CÓ quyền: sửa lại mức giảm giá đã lưu', 'P1',
     'Có báo giá do User A lập, đã có giảm giá tổng 5%.',
     '1. Mở màn Cập nhật báo giá\n2. Đổi mức giảm thành 8%\n3. Lưu',
     'Giảm giá tổng: 5% → 8%',
     '- Lưu thành công, tổng tiền tính lại theo mức mới'),
    ('Người KHÔNG có quyền: ô GG bị khóa và mặc định Không giảm giá', 'P0',
     'Đăng nhập User B, mở màn Tạo mới báo giá.',
     '1. Quan sát trường GG\n2. Thử bấm vào ô',
     'User B',
     '- Trường GG bị khóa, không chọn được\n- Giá trị hiển thị là “Không giảm giá”\n'
     '- Không có ô nhập giảm giá tổng và không nhập được giảm giá trên dòng hàng'),
    ('Người KHÔNG có quyền: mở báo giá cũ đang có giảm giá', 'P0',
     'Có báo giá do User A lập, đang có giảm giá tổng 5%. Đăng nhập User B.',
     '1. Mở màn Cập nhật báo giá đó\n2. Quan sát phần giảm giá',
     'Báo giá có giảm giá tổng 5%',
     '- Vẫn nhìn thấy mức giảm giá đang có để biết số tiền\n'
     '- Nhưng ô giảm giá bị khóa, không sửa được'),
    ('Người KHÔNG có quyền: gọi thẳng chức năng lưu kèm giảm giá', 'P0',
     'Đăng nhập User B; dùng công cụ kiểm thử để gọi thẳng chức năng lưu báo giá, bỏ qua giao diện.',
     '1. Gửi yêu cầu lưu báo giá có mức giảm giá lớn hơn 0',
     'Giảm giá tổng: 5%',
     '- Hệ thống từ chối lưu\n- Báo đúng thông báo “Bạn không có quyền áp dụng giảm giá trong '
     'báo giá”\n- ⚠️ Báo giá KHÔNG được ghi vào hệ thống'),
    ('Người KHÔNG có quyền: lưu báo giá không giảm giá', 'P0',
     'Đăng nhập User B.',
     '1. Lập báo giá bình thường, để GG là Không giảm giá\n2. Bấm Lưu',
     'Không giảm giá',
     '- Lưu thành công như bình thường — ⚠️ quyền này chỉ chặn phần giảm giá, không chặn việc '
     'lập báo giá'),
    ('Thu hồi quyền khi người dùng đang mở form', 'P1',
     'User A đang mở màn Cập nhật báo giá và đã nhập giảm giá 5%.',
     '1. Quản trị bỏ quyền của User A\n2. User A bấm Lưu',
     'Giảm giá tổng: 5%',
     '- Hệ thống từ chối lưu và báo không có quyền áp dụng giảm giá\n'
     '- Dữ liệu đã nhập vẫn còn trên màn hình, không bị mất'),
    ('Quyền không ảnh hưởng các thao tác khác của báo giá', 'P1',
     'Đăng nhập User B.',
     '1. Mở báo giá của mình, sửa số lượng và đơn giá một dòng hàng\n2. Lưu',
     'Số lượng: 2 → 3',
     '- Lưu thành công, chỉ riêng phần giảm giá là bị khóa'),
]

# ===================================================================== TASK 10797
T10797 = ('Cảnh báo giá bán thấp & duyệt', 'http://quanly.dnsmedia.vn/issues/10797')
TC_10797 = [
    ('Cảnh báo khi có hàng bán đơn giá dưới mức 1.000đ', 'P0',
     'Báo giá có 3 dòng hàng bán, trong đó 1 dòng đơn giá 500đ.',
     '1. Mở báo giá\n2. Bấm Trình duyệt',
     'Dòng 2: đơn giá 500đ',
     '- Dòng vi phạm được tô màu cam trên lưới\n'
     '- Hiện popup tiêu đề “Cảnh báo: Tồn tại mặt hàng có đơn giá <= 1.000 vnđ đồng”'),
    ('Nội dung popup ghi đúng số lượng mặt hàng vi phạm', 'P0',
     'Báo giá có 3 dòng đơn giá lần lượt 500đ, 800đ và 50.000đ.',
     '1. Bấm Trình duyệt\n2. Đọc nội dung popup',
     '2 dòng vi phạm',
     '- Câu thông báo ghi đúng số 2\n'
     '- Bảng trong popup liệt kê đủ 2 dòng, mỗi dòng có mã hàng hóa và tên hàng hóa'),
    ('Nút Quay lại chỉnh sửa', 'P0',
     'Đang hiện popup cảnh báo.',
     '1. Bấm “Quay lại chỉnh sửa”',
     '—',
     '- Popup đóng\n- Màn hình giữ nguyên như trước khi bấm Trình duyệt, dữ liệu đã nhập còn '
     'nguyên\n- Báo giá KHÔNG được gửi duyệt'),
    ('Nút Tiếp tục trình duyệt', 'P0',
     'Đang hiện popup cảnh báo.',
     '1. Bấm “Tiếp tục trình duyệt”',
     '—',
     '- Popup đóng và báo giá đi tiếp sang bước xét duyệt\n'
     '- Trạng thái báo giá đổi đúng theo luồng ở các trường hợp bên dưới'),
    ('Giá đúng 1.000đ vẫn bị cảnh báo', 'P0',
     'Báo giá có 1 dòng hàng bán đơn giá đúng 1.000đ.',
     '1. Bấm Trình duyệt',
     'Đơn giá: 1.000đ',
     '- ⚠️ Vẫn hiện cảnh báo — điều kiện là nhỏ hơn HOẶC BẰNG 1.000đ'),
    ('Giá 1.001đ không bị cảnh báo', 'P0',
     'Báo giá có 1 dòng hàng bán đơn giá 1.001đ, các dòng còn lại đều trên 1.000đ.',
     '1. Bấm Trình duyệt',
     'Đơn giá: 1.001đ',
     '- Không hiện popup cảnh báo, không dòng nào bị tô cam'),
    ('Hàng khuyến mại không bị tính là vi phạm', 'P0',
     'Báo giá có 1 dòng hàng khuyến mại đơn giá 0đ và các dòng hàng bán đều trên 1.000đ.',
     '1. Bấm Trình duyệt',
     'Dòng khuyến mại: 0đ',
     '- ⚠️ Không cảnh báo — điều kiện chỉ xét hàng hóa BÁN, bỏ qua hàng khuyến mại'),
    ('Nhiều dòng vi phạm cùng lúc', 'P1',
     'Báo giá có 5 dòng hàng bán, trong đó 4 dòng đơn giá dưới 1.000đ.',
     '1. Bấm Trình duyệt\n2. Đối chiếu popup với lưới',
     '4 dòng vi phạm',
     '- Cả 4 dòng đều tô cam\n- Popup ghi số 4 và liệt kê đủ 4 dòng'),
    ('Tự động duyệt khi báo giá sạch', 'P0',
     'Báo giá chỉ gồm hàng hóa và dịch vụ đã khai trên ERP, đơn giá đều trên 1.000đ, không có '
     'hàng tạm, không có dịch vụ tạm, không giảm giá dòng, không giảm giá tổng.',
     '1. Bấm Trình duyệt',
     '—',
     '- Không hiện cảnh báo\n- Báo giá được duyệt ngay, trạng thái ghi nhận là nhân viên kinh '
     'doanh tự phê duyệt\n- Không xuất hiện trong danh sách chờ duyệt của cấp trên'),
    ('Có hàng tạm thì đi luồng duyệt phân cấp', 'P0',
     'Báo giá có 1 dòng hàng tạm, các điều kiện khác đều sạch.',
     '1. Bấm Trình duyệt',
     '1 dòng hàng tạm',
     '- ⚠️ KHÔNG tự động duyệt\n- Báo giá chuyển sang chờ duyệt theo đúng luồng phân cấp hiện '
     'hành'),
    ('Có dịch vụ tạm thì đi luồng duyệt phân cấp', 'P0',
     'Báo giá có 1 dòng dịch vụ tạm, các điều kiện khác đều sạch.',
     '1. Bấm Trình duyệt',
     '1 dòng dịch vụ tạm',
     '- Không tự động duyệt, chuyển sang chờ duyệt theo luồng phân cấp'),
    ('Có giảm giá theo mặt hàng thì đi luồng duyệt phân cấp', 'P0',
     'Báo giá có 1 dòng được giảm giá 10%, các điều kiện khác đều sạch.',
     '1. Bấm Trình duyệt',
     'Dòng 1 giảm 10%',
     '- Không tự động duyệt, chuyển sang chờ duyệt theo luồng phân cấp'),
    ('Có giảm giá theo tổng đơn thì đi luồng duyệt phân cấp', 'P0',
     'Báo giá có giảm giá tổng 5%, các điều kiện khác đều sạch.',
     '1. Bấm Trình duyệt',
     'Giảm giá tổng 5%',
     '- Không tự động duyệt, chuyển sang chờ duyệt theo luồng phân cấp'),
    ('Vừa có cảnh báo giá thấp vừa có hàng tạm', 'P1',
     'Báo giá có 1 dòng đơn giá 500đ và 1 dòng hàng tạm.',
     '1. Bấm Trình duyệt\n2. Bấm “Tiếp tục trình duyệt”',
     '—',
     '- Hiện cảnh báo trước\n- Sau khi tiếp tục, báo giá đi luồng duyệt phân cấp chứ không tự '
     'động duyệt'),
]

# ===================================================================== TASK 10898
T10898 = ('Tỷ suất LN hàng/dịch vụ tạm', 'http://quanly.dnsmedia.vn/issues/10898')
TC_10898 = [
    ('Cấu hình duyệt giá có thêm điều kiện mới', 'P0',
     'Đăng nhập tài khoản quản trị cấu hình duyệt giá.',
     '1. Mở màn cấu hình hạn mức phê duyệt báo giá\n2. Quan sát danh sách điều kiện',
     '—',
     '- Có điều kiện mới “Theo tỷ suất LN dòng hàng tạm (%)” bên cạnh 2 điều kiện cũ\n'
     '- Khai được mức và lưu thành công'),
    ('Gửi duyệt: hệ thống quét đủ cả 3 điều kiện', 'P0',
     'Đã cấu hình đủ 3 điều kiện: giá trị đơn hàng, tỷ suất LN tổng, tỷ suất LN dòng hàng tạm.',
     '1. Lập báo giá có hàng tạm thỏa cả 3 điều kiện ở các cấp khác nhau\n2. Bấm Trình duyệt',
     'Giá trị: cấp 1 · LN tổng: cấp 2 · LN hàng tạm: cấp 3',
     '- ⚠️ Báo giá đi theo cấp duyệt CAO NHẤT trong 3 điều kiện (cấp 3), không phải cấp đầu tiên '
     'khớp'),
    ('Chỉ điều kiện tỷ suất LN dòng hàng tạm bị vi phạm', 'P0',
     'Báo giá có giá trị nhỏ và tỷ suất LN tổng đạt, riêng một dòng hàng tạm có tỷ suất LN thấp '
     'hơn mức cấu hình.',
     '1. Bấm Trình duyệt',
     'LN dòng hàng tạm: 3% (mức yêu cầu 10%)',
     '- Báo giá vẫn phải đi duyệt theo cấp mà điều kiện tỷ suất LN dòng hàng tạm quy định'),
    ('Báo giá không có hàng tạm', 'P0',
     'Báo giá chỉ gồm hàng hóa ERP, không có dòng hàng tạm nào.',
     '1. Bấm Trình duyệt',
     '—',
     '- Điều kiện tỷ suất LN dòng hàng tạm không áp dụng\n- Cấp duyệt xác định theo 2 điều kiện '
     'còn lại'),
    ('Nhiều dòng hàng tạm, lấy dòng xấu nhất', 'P1',
     'Báo giá có 3 dòng hàng tạm, tỷ suất LN lần lượt 25%, 12% và 4%.',
     '1. Bấm Trình duyệt',
     'LN thấp nhất: 4%',
     '- Cấp duyệt xác định theo dòng có tỷ suất thấp nhất (4%)'),
    ('Người lập báo giá xem được giá vốn hàng tạm', 'P0',
     'User A là người lập báo giá, KHÔNG được phân quyền xem giá vốn hàng hóa.',
     '1. Đăng nhập User A\n2. Mở màn Xem chi tiết báo giá có hàng tạm\n3. Quan sát dòng hàng tạm',
     'User A: không có quyền xem giá vốn',
     '- ⚠️ Vẫn thấy đủ Giá nhập hàng tạm, Giá bán hàng tạm và Tỷ suất lợi nhuận dòng hàng tạm'),
    ('Người phê duyệt xem được giá vốn hàng tạm', 'P0',
     'User C là người phê duyệt báo giá đó, KHÔNG có quyền xem giá vốn.',
     '1. Đăng nhập User C\n2. Mở màn Phê duyệt báo giá\n3. Quan sát dòng hàng tạm',
     'User C: không có quyền xem giá vốn',
     '- Vẫn thấy đủ 3 thông tin như trên'),
    ('Người ngoài không xem được giá vốn hàng tạm', 'P0',
     'User D không phải người lập, không phải người duyệt, không có quyền xem giá vốn.',
     '1. Đăng nhập User D\n2. Mở báo giá có hàng tạm',
     'User D',
     '- ⚠️ KHÔNG thấy Giá nhập hàng tạm và Tỷ suất lợi nhuận dòng hàng tạm'),
    ('Hiển thị ở màn Tạo mới', 'P1',
     'User A đang lập báo giá và vừa thêm một dòng hàng tạm.',
     '1. Nhập giá nhập và giá bán cho dòng hàng tạm\n2. Quan sát',
     'Giá nhập 100.000 · Giá bán 130.000',
     '- Hiện tỷ suất lợi nhuận của dòng, tính đúng theo giá vừa nhập'),
    ('Hiển thị ở màn Cập nhật', 'P1',
     'Báo giá có hàng tạm, User A mở màn Cập nhật.',
     '1. Sửa giá bán dòng hàng tạm\n2. Quan sát tỷ suất',
     'Giá bán: 130.000 → 110.000',
     '- Tỷ suất lợi nhuận tính lại ngay theo giá mới'),
    ('Người có quyền xem giá vốn thì không đổi gì', 'P2',
     'User E có quyền xem giá vốn hàng hóa.',
     '1. Mở báo giá có hàng tạm',
     'User E',
     '- Vẫn xem được như trước, tính năng mới không làm mất quyền cũ'),
    ('Đối chiếu tỷ suất với tài liệu mô phỏng', 'P1',
     'Có báo giá hàng tạm với số liệu đúng như file mô phỏng đính kèm task.',
     '1. So từng dòng giữa màn hình và file mô phỏng',
     'Theo file mô phỏng của task',
     '- Tỷ suất lợi nhuận từng dòng và cấp duyệt xác định được khớp với file mô phỏng'),
]

# ===================================================================== TASK 11014
T11014 = ('Biên bản & tự hủy cuộc họp', 'http://quanly.dnsmedia.vn/issues/11014')
TC_11014 = [
    ('Màn cấu hình có 2 tham số mới', 'P0',
     'Đăng nhập tài khoản quản trị cấu hình hệ thống.',
     '1. Mở màn cấu hình hệ thống, phần Meeting\n2. Quan sát',
     '—',
     '- Có ô “Số giờ gửi cảnh báo trước thời điểm hủy tự động”, mặc định 3\n'
     '- Có ô “Số ngày chặn cập nhật biên bản sau khi cuộc họp kết thúc”, mặc định 1\n'
     '- Sửa và lưu được'),
    ('Nhập biên bản trước hạn thì lưu bình thường', 'P0',
     'Cuộc họp kết thúc lúc 10:00 hôm nay, hạn nhập biên bản là 10:00 ngày mai.',
     '1. Người chủ trì mở cuộc họp\n2. Nhập biên bản và bấm hoàn thành',
     'Thời điểm thao tác: trước 10:00 ngày mai',
     '- Lưu biên bản thành công\n- Cuộc họp chuyển sang trạng thái hoàn thành, không bị hủy'),
    ('Cảnh báo trước thời điểm hủy', 'P0',
     'Cuộc họp kết thúc 10:00 hôm nay, hạn 10:00 ngày mai, cấu hình cảnh báo trước 3 giờ.',
     '1. Chờ tới 07:00 ngày mai\n2. Kiểm tra chuông thông báo của người tạo cuộc họp',
     'Cấu hình cảnh báo: 3 giờ',
     '- Người tạo cuộc họp nhận được thông báo\n'
     '- Nội dung đúng khuôn: ⏰ [MET] Nhắc nhở: {Tên cuộc họp}. Sẽ bị HỦY nếu không HOÀN THÀNH '
     'cuộc họp. Vui lòng cập nhật biên bản và HOÀN THÀNH trước lúc {Thời điểm}'),
    ('Đổi số giờ cảnh báo thì mốc nhắc đổi theo', 'P1',
     'Đổi cấu hình cảnh báo thành 5 giờ, cuộc họp có hạn 10:00 ngày mai.',
     '1. Chờ tới 05:00 ngày mai\n2. Kiểm tra thông báo',
     'Cấu hình cảnh báo: 5 giờ',
     '- Thông báo nhắc nhở gửi lúc 05:00, không phải 07:00'),
    ('Quá hạn chưa hoàn thành thì tự hủy', 'P0',
     'Cuộc họp kết thúc 10:00 hôm qua, người chủ trì chưa hoàn thành biên bản.',
     '1. Chờ qua 10:00 hôm nay\n2. Mở danh sách meeting',
     '—',
     '- Trạng thái cuộc họp tự chuyển thành “Hủy”\n'
     '- Người tạo và các thành viên nội bộ tham gia đều nhận được thông báo'),
    ('Sau khi bị hủy thì không cập nhật biên bản được nữa', 'P0',
     'Cuộc họp vừa bị hệ thống tự hủy.',
     '1. Người chủ trì mở cuộc họp\n2. Thử nhập và lưu biên bản',
     '—',
     '- ⚠️ Hệ thống chặn, không cho lưu biên bản\n- Nêu rõ lý do cuộc họp đã bị hủy do quá hạn'),
    ('Cuộc họp bị hủy không tính vào KPI', 'P0',
     'Nhân viên X có 1 cuộc họp bị hệ thống tự hủy trong kỳ.',
     '1. Mở báo cáo KPI / báo cáo meeting theo nhân viên của kỳ đó',
     'Nhân viên X',
     '- ⚠️ Cuộc họp bị hủy KHÔNG được cộng vào số liệu KPI của nhân viên'),
    ('Hoàn thành đúng lúc sát hạn', 'P1',
     'Cuộc họp có hạn 10:00, người chủ trì hoàn thành biên bản lúc 09:59.',
     '1. Lưu và hoàn thành biên bản lúc 09:59\n2. Chờ qua 10:00',
     'Thời điểm hoàn thành: 09:59',
     '- Cuộc họp giữ trạng thái hoàn thành, không bị hủy'),
    ('Chỉ lưu nháp biên bản mà chưa hoàn thành', 'P0',
     'Người chủ trì có nhập nội dung biên bản nhưng chưa bấm hoàn thành, đã quá hạn.',
     '1. Chờ qua hạn\n2. Kiểm tra trạng thái cuộc họp',
     '—',
     '- ⚠️ Vẫn bị tự hủy — điều kiện là HOÀN THÀNH biên bản, không phải chỉ nhập nội dung'),
    ('Đổi số ngày chốt biên bản', 'P1',
     'Đổi cấu hình số ngày chặn cập nhật thành 2 ngày.',
     '1. Tạo cuộc họp kết thúc 10:00 hôm nay\n2. Kiểm tra hạn nhập biên bản',
     'Cấu hình: 2 ngày',
     '- Hạn nhập biên bản là 10:00 sau 2 ngày, thông báo nhắc và mốc tự hủy dịch theo'),
    ('Cuộc họp bị hủy tay trước hạn', 'P2',
     'Người tạo chủ động hủy cuộc họp trước khi tới hạn.',
     '1. Hủy cuộc họp\n2. Chờ qua hạn nhập biên bản',
     '—',
     '- Không phát sinh thêm thông báo nhắc nhở và không bị hủy lần hai'),
    ('Nhiều cuộc họp cùng tới hạn', 'P1',
     'Có 3 cuộc họp cùng kết thúc 10:00 hôm qua, đều chưa hoàn thành biên bản.',
     '1. Chờ qua 10:00 hôm nay\n2. Kiểm tra cả 3',
     '3 cuộc họp',
     '- Cả 3 đều chuyển sang “Hủy” và đều gửi thông báo, không sót cuộc nào'),
]


def sheet(wb, title, tasks, tc_start, prefix, module, blank_col_c):
    """Dung mot sheet la KHOI DONG MOI de dan vao tab dich.

    blank_col_c=True: tab dich co 1 cot TRONG o vi tri C (tab Meeting).
    """
    ws = wb.create_sheet(title)
    cols = (['Module', 'Nhóm chức năng', '(cột trống)', 'TC ID', 'Chức năng', 'Priority',
             'Tiền điều kiện', 'Bước thực hiện', 'Test Data', 'Expected Result (chi tiết)']
            if blank_col_c else
            ['Module', 'Nhóm chức năng', 'TC ID', 'Chức năng', 'Priority', 'Tiền điều kiện',
             'Bước thực hiện', 'Test Data', 'Expected Result (chi tiết)'])
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(1, i, c)
        cell.font = Font(name=FONT, size=12, bold=True, color='FFFFFF')
        cell.fill = HEAD
        cell.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
        cell.border = BOX
    widths = ([14, 26, 8, 14, 34, 9, 34, 38, 24, 52] if blank_col_c
              else [14, 26, 14, 34, 9, 34, 38, 24, 52])
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    ws.row_dimensions[1].height = 30

    r = 2
    n = tc_start
    for (group, link), items in tasks:
        # Dong BANNER phan cach giua cac task — dung quy uoc co san cua file dich:
        #   tab TKT     : cot A = "UPDATE ...", cot D = "Task <link>"
        #   tab Meeting : cot A = "UPDATE ...", cot G = "Task <link>"
        banner = [''] * len(cols)
        banner[0] = BANNER_TITLE
        banner[6 if blank_col_c else 3] = 'Task ' + link
        for i, v in enumerate(banner, start=1):
            cell = ws.cell(r, i, v)
            cell.font = Font(name=FONT, size=12, bold=True)
            cell.alignment = WRAP
            cell.border = BOX
        ws.row_dimensions[r].height = 22
        r += 1
        for idx, (fn, pri, pre, steps, data, exp) in enumerate(items):
            tc_id = prefix % n
            n += 1
            vals = ([module if idx == 0 else '', group if idx == 0 else '', '',
                     tc_id, fn, pri, pre, steps, data, exp] if blank_col_c else
                    [module if idx == 0 else '', group if idx == 0 else '',
                     tc_id, fn, pri, pre, steps, data, exp])
            for i, v in enumerate(vals, start=1):
                cell = ws.cell(r, i, v)
                cell.font = Font(name=FONT, size=12)
                cell.alignment = WRAP
                cell.border = BOX
            ws.row_dimensions[r].height = 58
            r += 1
    return ws, n - 1


wb = Workbook()
wb.remove(wb.active)

ws1, last1 = sheet(wb, 'Dán vào tab Dự án tiền khả thi',
                   [(T10789, TC_10789), (T10797, TC_10797), (T10898, TC_10898)],
                   340, 'TC-ROLE-%d', 'Dự án tiền KT', blank_col_c=False)
ws2, last2 = sheet(wb, 'Dán vào tab 11.Meeting',
                   [(T11014, TC_11014)], 1, 'TC_17.%03d', 'Meeting', blank_col_c=True)

wb.save(OUT)

# --- xuat CSV (dung de Tep > Nhap > Them vao trang tinh hien tai) ---
import csv
for ws, name in ((ws1, 'tc_duan_tkt.csv'), (ws2, 'tc_meeting.csv')):
    path = os.path.join(HERE, name)
    with io.open(path, 'w', encoding='utf-8', newline='') as f:
        w = csv.writer(f, quoting=csv.QUOTE_ALL)
        for row in ws.iter_rows(min_row=2, values_only=True):   # bo dong tieu de
            w.writerow(['' if v is None else v for v in row])
    print('CSV:', path)
print('Da tao:', OUT)
print('Tab Du an tien kha thi : %d TC (TC-ROLE-340 -> TC-ROLE-%d)'
      % (len(TC_10789) + len(TC_10797) + len(TC_10898), last1))
print('Tab Meeting           : %d TC (TC_17.001 -> TC_17.%03d)' % (len(TC_11014), last2))
