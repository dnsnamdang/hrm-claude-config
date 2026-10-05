"""
Helper sửa TẠI CHỖ file .docx do Google Docs xuất ra (HDSD / SRS của nhóm danh mục địa lý).

Vì sao không dựng lại file từ khung mẫu: 2 tài liệu này đã được QA chỉnh trực tiếp trên Drive
(HDSD v1.1 21/09/2026, SRS 28/08/2026). Dựng lại = mất hết phần QA sửa. Ở đây chỉ chèn/sửa đúng
chỗ cần, mọi định dạng lấy bằng cách CLONE một đoạn/bảng có sẵn trong chính tài liệu đó.
"""
import copy

from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.enum.text import WD_ALIGN_PARAGRAPH


def para_text(p_el):
    return Paragraph(p_el, None).text.strip()


def _strip_bookmarks(p_el):
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd'):
        for el in p_el.findall(qn(tag)):
            p_el.remove(el)


def set_para_text(p_el, text, doc=None):
    """Giữ nguyên định dạng của run đầu tiên, xoá các run còn lại rồi gán text mới."""
    par = Paragraph(p_el, doc._body if doc else None)
    runs = par.runs
    for r in runs[1:]:
        r._r.getparent().remove(r._r)
    if runs:
        runs[0].text = text
    else:
        par.add_run(text)
    return par


def clone_para(tmpl_p, text, doc, bookmark=None, bm_id=None):
    """Nhân bản 1 đoạn mẫu -> đoạn mới cùng style, text mới (kèm bookmark nếu là heading)."""
    new = copy.deepcopy(tmpl_p)
    _strip_bookmarks(new)
    set_para_text(new, text, doc)
    if bookmark:
        start = new.makeelement(qn('w:bookmarkStart'), {})
        start.set(qn('w:id'), str(bm_id))
        start.set(qn('w:name'), bookmark)
        start.set(qn('w:colFirst'), '0')
        start.set(qn('w:colLast'), '0')
        end = new.makeelement(qn('w:bookmarkEnd'), {})
        end.set(qn('w:id'), str(bm_id))
        pPr = new.find(qn('w:pPr'))
        if pPr is not None:
            pPr.addnext(end)
            pPr.addnext(start)
        else:
            new.insert(0, end)
            new.insert(0, start)
    return new


def clone_image_para(tmpl_p, image_path, doc, width_inches=6.0):
    """Đoạn chỉ chứa 1 ảnh, canh giữa — mọi ảnh nội dung rộng 6.0 inch như các ảnh sẵn có."""
    new = copy.deepcopy(tmpl_p)
    _strip_bookmarks(new)
    par = Paragraph(new, doc._body)
    for r in par.runs:
        r._r.getparent().remove(r._r)
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.add_run().add_picture(image_path, width=Inches(width_inches))
    return new


def clone_table(tmpl_tbl, rows, doc):
    """Nhân bản bảng mẫu rồi đổ dữ liệu `rows` (list các list chuỗi, dòng 0 là tiêu đề)."""
    new_el = copy.deepcopy(tmpl_tbl._tbl)
    tbl = Table(new_el, doc._body)
    # số dòng: thêm/bớt cho khớp
    while len(tbl.rows) > len(rows):
        new_el.remove(tbl.rows[-1]._tr)
    while len(tbl.rows) < len(rows):
        new_el.append(copy.deepcopy(tbl.rows[-1]._tr))
    for r_idx, row_data in enumerate(rows):
        for c_idx, value in enumerate(row_data):
            cell = tbl.rows[r_idx].cells[c_idx]
            paras = cell.paragraphs
            for p in paras[1:]:
                p._p.getparent().remove(p._p)
            set_para_text(paras[0]._p, value, doc)
    return new_el


def append_before_sectpr(doc, elements):
    body = doc.element.body
    sectPr = body.find(qn('w:sectPr'))
    for el in elements:
        if sectPr is not None:
            sectPr.addprevious(el)
        else:
            body.append(el)


def toc_paragraphs(doc):
    sdt = doc.element.body.findall(qn('w:sdt'))[0]
    return sdt.find(qn('w:sdtContent')), sdt.find(qn('w:sdtContent')).findall(qn('w:p'))


def clone_toc_entry(tmpl_p, text, anchor, page):
    """Nhân bản 1 dòng mục lục: đổi liên kết, chữ và số trang."""
    new = copy.deepcopy(tmpl_p)
    hl = new.find(qn('w:hyperlink'))
    hl.set(qn('w:anchor'), anchor)
    runs = hl.findall(qn('w:r'))
    for r in runs[1:]:
        hl.remove(r)
    r = runs[0]
    for t in r.findall(qn('w:t')) + r.findall(qn('w:tab')):
        r.remove(t)
    t1 = r.makeelement(qn('w:t'), {})
    t1.set(qn('xml:space'), 'preserve')
    t1.text = text
    tab = r.makeelement(qn('w:tab'), {})
    t2 = r.makeelement(qn('w:t'), {})
    t2.set(qn('xml:space'), 'preserve')
    t2.text = str(page)
    r.append(t1)
    r.append(tab)
    r.append(t2)
    return new


def replace_image_blob(doc, para_index, new_path, nth=0):
    """Thay ẢNH SẴN CÓ trong một ĐOẠN bằng file mới (giữ nguyên vị trí + kích thước đã khai).

    ⚠️ Chỉ đường theo ĐOẠN chứ không theo `doc.inline_shapes[i]`: danh sách inline_shapes bỏ qua
    ảnh neo nổi (logo trang bìa) nên số thứ tự lệch, rất dễ ghi đè nhầm vào icon menu nhỏ nằm
    giữa dòng văn bản.
    """
    p_el = doc.paragraphs[para_index]._p
    blips = p_el.findall('.//' + qn('a:blip'))
    assert len(blips) > nth, 'Đoạn %d không có đủ ảnh' % para_index
    rId = blips[nth].get(qn('r:embed'))
    part = doc.part.related_parts[rId]
    with open(new_path, 'rb') as f:
        part._blob = f.read()
    return part.partname
