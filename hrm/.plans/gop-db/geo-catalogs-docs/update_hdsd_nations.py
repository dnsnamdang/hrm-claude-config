"""
Cập nhật HDSD_Danh mục quốc gia.docx cho khớp code hiện tại (nhánh gop_db, 22/09/2026).

Bổ sung 4 chức năng còn thiếu (Xem chi tiết, Xuất Excel, Import Excel, Tuỳ chỉnh cột) và sửa các
chỗ tài liệu mô tả sai (ô tìm nhanh, danh sách cột, điều kiện ẩn nút Xóa, thanh công cụ).

Chạy:  python3 update_hdsd_nations.py <file_vao.docx> <file_ra.docx>
"""
import os
import sys

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_edit_helper import (  # noqa: E402
    append_before_sectpr, clone_image_para, clone_para, clone_table,
    clone_toc_entry, para_text, replace_image_blob, set_para_text,
    toc_paragraphs,
)

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'hdsd_nations_shots')

# Số trang của các mục MỚI, điền vào mục lục. Đo bằng cách xuất bản PDF của chính file kết quả
# (LibreOffice) rồi trừ đi độ lệch so với số trang mục lục cũ. Mục lục vẫn là field thật: mở bằng
# Word / Google Docs bấm cập nhật là ra số trang chuẩn của máy đang xem.
TRANG = {
    '_heading=h.hdsdp7xemchitiet': 17,
    '_heading=h.hdsdp7m1': 17,
    '_heading=h.hdsdp7m2': 17,
    '_heading=h.hdsdp8xuatexcel': 19,
    '_heading=h.hdsdp8m1': 19,
    '_heading=h.hdsdp8m2': 19,
    '_heading=h.hdsdp9import': 21,
    '_heading=h.hdsdp9m1': 21,
    '_heading=h.hdsdp9m2': 23,
    '_heading=h.hdsdp10cot': 25,
}


def shot(name):
    p = os.path.join(SHOTS, name)
    assert os.path.exists(p), 'Thiếu ảnh: ' + p
    return p


def main(src, out):
    doc = Document(src)
    paras = doc.paragraphs

    T_H1 = paras[32]._p          # "PHẦN 1: TRUY CẬP VÀ BỐ CỤC MÀN HÌNH"
    T_H2 = paras[33]._p          # "1. Truy cập màn hình"
    T_BODY = paras[39]._p        # đoạn văn thường
    T_CAP = paras[37]._p         # "Hình 1: ..."
    T_IMG = paras[36]._p         # đoạn chỉ có ảnh
    T_TBL2 = doc.tables[3]       # bảng 2 cột (Ô | Cách dùng)
    # Đoạn Heading 1 RỖNG chỉ chứa ngắt trang — tài liệu dùng nó để đẩy mỗi PHẦN sang trang mới
    # (style Heading 1 KHÔNG tự ngắt trang, thiếu đoạn này là phần mới dính vào cuối phần trước).
    T_PAGEBREAK = paras[132]._p

    # ------------------------------------------------------------------ #
    # 1. Ảnh màn danh sách (Hình 1) — thanh công cụ nay có thêm 3 nút mới #
    # ------------------------------------------------------------------ #
    # Đoạn 36 = đoạn chỉ chứa ảnh của Hình 1
    replace_image_blob(doc, 36, shot('01-danh-sach.png'))

    # ---------------------------------------- #
    # 2. Bảng "Cập nhật tài liệu" — thêm v1.2  #
    # ---------------------------------------- #
    tbl = doc.tables[1]
    import copy
    tbl._tbl.append(copy.deepcopy(tbl.rows[-1]._tr))
    row = tbl.rows[-1]
    for cell, value in zip(row.cells, [
        '1.2', '22/09/2026', 'Đội phát triển phần mềm',
        'Bổ sung Xem chi tiết, Xuất Excel, Import Excel, Tuỳ chỉnh cột; '
        'cập nhật danh sách cột và phạm vi ô tìm kiếm nhanh.',
    ]):
        for p in cell.paragraphs[1:]:
            p._p.getparent().remove(p._p)
        set_para_text(cell.paragraphs[0]._p, value, doc)

    # ------------------------------------------------- #
    # 3. Sửa các đoạn mô tả sai ở PHẦN 1 và PHẦN 2      #
    # ------------------------------------------------- #
    sua_doan = {
        'Thanh công cụ — nút Tạo mới.':
            'Thanh công cụ — nút Tạo mới, nút Xuất Excel, nút Import Excel và biểu tượng '
            'Tuỳ chỉnh cột.',
        'Biểu tượng thùng rác — xóa bản ghi.':
            'Biểu tượng thùng rác — xóa bản ghi. Nút này chỉ hiện khi quốc gia CHƯA có Khu vực '
            'nào trực thuộc; đã có Khu vực bên dưới thì nút biến mất.',
    }
    for p in doc.paragraphs:
        t = p.text.strip()
        if t in sua_doan:
            set_para_text(p._p, sua_doan.pop(t), doc)
    assert not sua_doan, 'Không tìm thấy đoạn cần sửa: ' + str(list(sua_doan))

    # Bảng "Các cột của bảng danh sách" — dựng lại cho đủ 10 cột
    cot_rows = [
        ['Cột', 'Nội dung'],
        ['STT', 'Số thứ tự, chạy liên tục qua các trang. Luôn hiển thị.'],
        ['Tên quốc gia', 'Tên đầy đủ. Luôn hiển thị, sắp xếp được. Bấm vào tên để mở cửa sổ '
                         'Xem quốc gia (xem PHẦN 7).'],
        ['Mã quốc gia', 'Mã do người dùng đặt, sắp xếp được.'],
        ['Mã bưu chính', 'Mặc định ẩn, bật ở cửa sổ Tuỳ chỉnh cột (xem PHẦN 10).'],
        ['Người cập nhật', 'Người sửa bản ghi gần nhất. Mặc định ẩn.'],
        ['Ngày cập nhật', 'Thời điểm sửa gần nhất. Mặc định ẩn, sắp xếp được.'],
        ['Người tạo', 'Người đã thêm bản ghi.'],
        ['Ngày tạo', 'Thời điểm thêm, sắp xếp được.'],
        ['Trạng thái', 'Hoạt động hoặc Khóa.'],
        ['Hành động', 'Sửa, Xóa hiện thẳng; Khóa / Mở khóa và Lịch sử nằm trong nút ba chấm.'],
    ]
    cu = doc.tables[2]._tbl
    moi = clone_table(doc.tables[2], cot_rows, doc)
    cu.addnext(moi)
    cu.getparent().remove(cu)

    # Bảng ô tìm kiếm / lọc — ô tìm nhanh chỉ tìm theo TÊN
    tim_rows = [
        ['Ô', 'Cách dùng'],
        ['Ô tìm kiếm nhanh', 'Tìm theo tên quốc gia. Gõ một phần tên cũng ra kết quả. '
                             'Ô này KHÔNG tìm theo mã quốc gia.'],
        ['Trạng thái', 'Bỏ trống thì hiện cả hai trạng thái.'],
    ]
    cu = doc.tables[3]._tbl
    moi = clone_table(T_TBL2, tim_rows, doc)
    cu.addnext(moi)
    cu.getparent().remove(cu)
    T_TBL2 = doc.tables[3]

    # ------------------------------------------------------- #
    # 4. Thêm 4 phần mới vào cuối tài liệu (PHẦN 7 -> PHẦN 10) #
    # ------------------------------------------------------- #
    els = []
    toc_moi = []          # (cấp, text, anchor, trang tạm)
    bm_id = [9000]

    def h1(text, anchor):
        bm_id[0] += 1
        import copy as _copy
        ngat_trang = _copy.deepcopy(T_PAGEBREAK)
        for tag in ('w:bookmarkStart', 'w:bookmarkEnd'):
            for el in ngat_trang.findall(qn(tag)):
                ngat_trang.remove(el)
        els.append(ngat_trang)
        els.append(clone_para(T_H1, text, doc, bookmark=anchor, bm_id=bm_id[0]))
        toc_moi.append((1, text, anchor))

    def h2(text, anchor):
        bm_id[0] += 1
        els.append(clone_para(T_H2, text, doc, bookmark=anchor, bm_id=bm_id[0]))
        toc_moi.append((2, text, anchor))

    def para(text):
        els.append(clone_para(T_BODY, text, doc))

    def hinh(file_name, caption):
        els.append(clone_image_para(T_IMG, shot(file_name), doc))
        els.append(clone_para(T_CAP, caption, doc))

    def bang(rows):
        els.append(clone_table(T_TBL2, rows, doc))

    # ---------------- PHẦN 7: XEM CHI TIẾT ---------------- #
    h1('PHẦN 7: XEM CHI TIẾT QUỐC GIA', '_heading=h.hdsdp7xemchitiet')
    h2('1. Mở cửa sổ xem chi tiết', '_heading=h.hdsdp7m1')
    para('Bước 1: Truy cập vào màn hình danh mục quốc gia.')
    para('Bước 2: Bấm vào TÊN quốc gia ở cột Tên quốc gia. Hệ thống mở cửa sổ “Xem quốc gia”.')
    hinh('07-xem-chi-tiet.png', 'Hình 10: Cửa sổ Xem quốc gia ở chế độ chỉ đọc')
    para('Cửa sổ hiển thị Tên quốc gia, Mã quốc gia và Mã bưu chính ở chế độ chỉ đọc: không gõ '
         'được, không có nút Lưu. Quốc gia đang ở trạng thái Khóa vẫn xem được bình thường.')
    para('Bấm Đóng để quay về danh sách. Bộ lọc và trang đang xem giữ nguyên.')
    h2('2. Xem lịch sử ngay trong cửa sổ', '_heading=h.hdsdp7m2')
    para('Bước 1: Trong cửa sổ Xem quốc gia, bấm “Xem lịch sử” ở khối Lịch sử.')
    para('Bước 2: Danh sách các lần thay đổi hiện ra, mới nhất ở trên cùng.')
    hinh('08-xem-lich-su-trong-popup.png',
         'Hình 11: Khối Lịch sử mở ngay trong cửa sổ Xem quốc gia')
    para('Con số bên cạnh chữ Lịch sử là số lần thay đổi đã ghi nhận của bản ghi đó. Bấm “Bộ lọc” '
         'để thu hẹp danh sách, “Làm mới” để nạp lại, “Thu gọn” để đóng khối lịch sử.')
    para('Bản ghi chưa từng sửa thì khối này hiện “Chưa có lịch sử thao tác nào.”.')

    # ---------------- PHẦN 8: XUẤT EXCEL ---------------- #
    h1('PHẦN 8: XUẤT DANH SÁCH RA EXCEL', '_heading=h.hdsdp8xuatexcel')
    h2('1. Các bước xuất file', '_heading=h.hdsdp8m1')
    para('Bước 1: Lọc danh sách theo đúng dữ liệu cần lấy (xem PHẦN 2). File xuất ra chạy theo '
         'ĐÚNG bộ lọc đang áp dụng.')
    para('Bước 2: Bấm nút Xuất Excel (nút màu xanh lá trên thanh công cụ của bảng). Hệ thống mở '
         'cửa sổ “Chọn trường xuất file”.')
    hinh('05-xuat-excel.png', 'Hình 12: Cửa sổ Chọn trường xuất file')
    para('Bước 3: Tích chọn các trường cần xuất. Mặc định hệ thống tích sẵn đúng những cột đang '
         'hiển thị trên bảng, nên file xuất ra giống thứ Người dùng đang nhìn.')
    para('Bước 4: Bấm Xuất file. Hệ thống tải về file danh_muc_quoc_gia.xlsx.')
    h2('2. Tác dụng các thành phần trên cửa sổ', '_heading=h.hdsdp8m2')
    bang([
        ['Thành phần', 'Tác dụng'],
        ['Ô Trường xuất', 'Chọn nhiều trường: Tên quốc gia, Mã quốc gia, Mã bưu chính, Trạng '
                          'thái, Người tạo, Ngày tạo, Người cập nhật, Ngày cập nhật.'],
        ['Dòng “Thứ tự cột trong file”', 'Cho biết thứ tự cột sẽ ghi ra file. Thứ tự chạy theo '
                                         'trình tự Người dùng tích chọn; muốn đổi vị trí thì bỏ '
                                         'chọn rồi chọn lại theo trình tự mong muốn.'],
        ['Chọn tất cả', 'Tích hết các trường.'],
        ['Bỏ chọn hết', 'Bỏ tích toàn bộ. Khi không còn trường nào được chọn, nút Xuất file mờ đi '
                        'và không bấm được.'],
        ['Xuất file', 'Sinh file Excel và tải về máy.'],
        ['Đóng', 'Đóng cửa sổ, không xuất gì.'],
    ])
    para('Lưu ý: file xuất chứa TOÀN BỘ kết quả của bộ lọc, không chỉ các dòng của trang đang '
         'xem.')

    # ---------------- PHẦN 9: IMPORT ---------------- #
    h1('PHẦN 9: NHẬP QUỐC GIA TỪ FILE EXCEL', '_heading=h.hdsdp9import')
    h2('1. Các bước nhập dữ liệu', '_heading=h.hdsdp9m1')
    para('Bước 1: Truy cập màn hình danh mục quốc gia, bấm nút Import Excel (nút màu cam trên '
         'thanh công cụ). Hệ thống mở cửa sổ “Import quốc gia”.')
    hinh('02-import-buoc1.png', 'Hình 13: Cửa sổ Import quốc gia khi vừa mở')
    para('Bước 2: Bấm Tải file mẫu để lấy file Mau_import_quoc_gia.xlsx, gồm các cột: STT, Tên '
         'quốc gia (bắt buộc), Mã quốc gia (bắt buộc), Mã bưu chính. Điền dữ liệu vào file này.')
    para('Bước 3: Bấm Chọn file Excel, chọn file vừa điền, rồi bấm Load lên bảng. Dữ liệu hiện '
         'lên bảng xem trước và ô Tổng cho biết đọc được bao nhiêu dòng.')
    hinh('03-import-load.png', 'Hình 14: Dữ liệu đã được load lên bảng xem trước')
    para('Bước 4: Bấm Validate để hệ thống kiểm tra từng dòng. Bước này CHƯA ghi dữ liệu vào hệ '
         'thống.')
    hinh('04-import-validate.png', 'Hình 15: Kết quả kiểm tra — mỗi dòng lỗi nêu rõ lý do')
    para('Sau khi kiểm tra, hệ thống hiện ba con số: Tổng, Hợp lệ, Lỗi. Dòng lỗi được tô nền đỏ '
         'kèm lý do ngay dưới ô; dòng hợp lệ bị khoá lại, không sửa được nữa.')
    para('Bước 5: Sửa trực tiếp các dòng lỗi trên bảng rồi bấm Validate lại. Nếu không muốn nhập '
         'các dòng lỗi thì bấm Bỏ dòng lỗi để loại chúng khỏi bảng. Muốn sửa lại cả những dòng đã '
         'khoá thì bấm Xoá trạng thái validate.')
    para('Bước 6: Bấm Import. Hệ thống chỉ ghi các dòng hợp lệ, báo số dòng thành công / thất bại '
         'rồi đóng cửa sổ và nạp lại danh sách.')
    h2('2. Các lỗi thường gặp khi kiểm tra dữ liệu', '_heading=h.hdsdp9m2')
    bang([
        ['Thông báo', 'Nguyên nhân và cách xử lý'],
        ['Tên quốc gia không được để trống', 'Dòng đó bỏ trống cột Tên quốc gia.'],
        ['Tên quốc gia đã tồn tại trong hệ thống', 'Danh mục đã có quốc gia trùng tên. Kiểm tra '
                                                   'lại danh sách trước khi nhập.'],
        ['Tên quốc gia bị trùng với dòng … trong file', 'Trong chính file đang nhập có hai dòng '
                                                        'cùng tên. Xoá bớt một dòng.'],
        ['Tên quốc gia tối đa 255 ký tự', 'Tên quá dài, rút gọn lại.'],
        ['Mã quốc gia không được để trống', 'Dòng đó bỏ trống cột Mã quốc gia.'],
        ['Mã quốc gia đã tồn tại trong hệ thống', 'Danh mục đã có quốc gia dùng mã này.'],
        ['Mã quốc gia bị trùng với dòng … trong file', 'Hai dòng trong file cùng một mã.'],
        ['Mã bưu chính chỉ được nhập chữ số', 'Ô Mã bưu chính có chữ cái hoặc ký tự khác.'],
        ['Mã bưu chính tối đa 50 chữ số', 'Mã bưu chính vượt quá 50 chữ số.'],
    ])
    para('Lưu ý: quốc gia nhập từ Excel luôn được tạo ở trạng thái Hoạt động (file mẫu không có '
         'cột Trạng thái), người tạo là chính người thực hiện nhập. Mã quốc gia được tự động '
         'chuyển thành chữ IN HOA.')

    # ---------------- PHẦN 10: TUỲ CHỈNH CỘT ---------------- #
    h1('PHẦN 10: TUỲ CHỈNH CỘT HIỂN THỊ', '_heading=h.hdsdp10cot')
    para('Bước 1: Ở thanh công cụ của bảng, bấm biểu tượng cột (nút ngoài cùng bên phải, cạnh nút '
         'Import Excel). Hệ thống mở cửa sổ “Tuỳ chỉnh cột”.')
    hinh('06-cau-hinh-cot.png', 'Hình 16: Cửa sổ Tuỳ chỉnh cột')
    para('Bước 2: Tích để hiện cột, bỏ tích để ẩn cột. Kéo biểu tượng ba gạch ở cuối mỗi dòng để '
         'đổi thứ tự cột.')
    para('Bước 3: Bấm Lưu để áp dụng; bảng vẽ lại ngay theo cấu hình mới. Bấm Đóng nếu muốn bỏ '
         'các thay đổi vừa chỉnh.')
    para('Ba cột STT, Tên quốc gia và Hành động bị khoá (hiển thị xám, có biểu tượng ổ khoá): '
         'không bỏ tích và không kéo đổi vị trí được.')
    para('Cấu hình cột được ghi nhớ riêng cho từng người dùng, lần sau vào màn hình vẫn giữ '
         'nguyên.')

    append_before_sectpr(doc, els)

    # ------------------------------- #
    # 5. Bổ sung dòng vào MỤC LỤC     #
    # ------------------------------- #
    content, toc_ps = toc_paragraphs(doc)
    tmpl_h1 = toc_ps[27]      # dòng cấp 1 ("PHẦN 6: ...")
    tmpl_h2 = toc_ps[26]      # dòng cấp 2 ("3. Mở khóa")
    for cap, text, anchor in toc_moi:
        tmpl = tmpl_h1 if cap == 1 else tmpl_h2
        content.append(clone_toc_entry(tmpl, text, anchor, TRANG.get(anchor, 0)))

    doc.save(out)

    # ---------- kiểm tra ---------- #
    check = Document(out)
    texts = [p.text.strip() for p in check.paragraphs]
    for must in ['PHẦN 7: XEM CHI TIẾT QUỐC GIA', 'PHẦN 8: XUẤT DANH SÁCH RA EXCEL',
                 'PHẦN 9: NHẬP QUỐC GIA TỪ FILE EXCEL', 'PHẦN 10: TUỲ CHỈNH CỘT HIỂN THỊ',
                 'Hình 16: Cửa sổ Tuỳ chỉnh cột']:
        assert must in texts, 'Thiếu: ' + must
    print('OK — số ảnh:', len(check.inline_shapes), '| số bảng:', len(check.tables))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
