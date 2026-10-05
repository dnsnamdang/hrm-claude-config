"""
Cập nhật "SRS - Danh mục quốc gia.docx" cho khớp code hiện tại (nhánh gop_db, 22/09/2026).

Việc chính:
  · sửa đánh số mục con bị lẫn giữa các chức năng (2.8 đang mang số 2.10.x, 2.9 mang 2.11.5,
    2.10 có hai mục 2.10.1, 2.11 mang 2.4.x) và đánh số hình bị trùng;
  · chèn ảnh thật cho 2.8 Import và 2.9 Xuất Excel (đang ghi "Thêm hình sau" / "Thêm sau"),
    thay ảnh cũ ở 2.1 / 2.10 / 2.11 bằng ảnh chụp bản hiện tại;
  · viết lại các bảng mô tả giao diện + event của 2.8 → 2.11 cho đúng màn thật
    (bỏ Xuất PDF, bỏ giới hạn 1.000 dòng, bỏ nút Khôi phục mặc định, bỏ nút Quay lại…);
  · ô tìm kiếm nhanh chỉ tìm theo TÊN (BE chỉ lọc theo `name`);
  · bổ sung FR-08 → FR-11 vào ma trận phân quyền.

Chạy:  python3 update_srs_nations.py <file_vao.docx> <file_ra.docx>
"""
import os
import sys

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from docx_edit_helper import (  # noqa: E402
    clone_image_para, clone_para, clone_table, replace_image_blob, set_para_text,
)

SHOTS = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'hdsd_nations_shots')
CHUA_PHAN_QUYEN = ('Người dùng đã đăng nhập. Màn hình hiện chưa gắn quyền riêng; sẽ cập nhật sau '
                   'khi có ma trận phân quyền chi tiết.')


def shot(name):
    p = os.path.join(SHOTS, name)
    assert os.path.exists(p), 'Thiếu ảnh: ' + p
    return p


def sdt_tables(doc):
    """7 bảng của tài liệu nằm trong khối sdt (Google Docs bọc lại), python-docx không thấy."""
    from docx.table import Table
    out = []
    for sdt in (c for c in doc.element.body.iterchildren() if c.tag == qn('w:sdt')):
        content = sdt.find(qn('w:sdtContent'))
        for tbl in content.findall(qn('w:tbl')):
            out.append(Table(tbl, doc._body))
    return out


def thay_bang(doc, tbl, rows):
    cu = tbl._tbl
    moi = clone_table(tbl, rows, doc)
    cu.addnext(moi)
    cu.getparent().remove(cu)


def main(src, out):
    doc = Document(src)
    paras = doc.paragraphs
    T_CAP = paras[29]._p     # "Hình 1: Sơ đồ Use Case..."
    T_IMG = paras[28]._p     # đoạn chỉ chứa ảnh

    # ---------------------------------------------------------- #
    # 1. Đánh số lại các mục con bị lẫn (sửa theo VỊ TRÍ đoạn)     #
    # ---------------------------------------------------------- #
    danh_so = {
        158: '2.8.3 Layout màn hình',
        164: '2.8.4 Mô tả chi tiết giao diện',
        166: '2.8.5 Danh sách event và xử lý event',
        183: '2.9.5 Danh sách event và xử lý event',
        188: '2.10.2 Giới thiệu',
        191: '2.10.3 Layout màn hình',
        197: '2.10.4 Mô tả chi tiết giao diện',
        199: '2.10.5 Danh sách event và xử lý event',
        205: '2.11.2 Layout màn hình',
        211: '2.11.3 Mô tả chi tiết giao diện',
        213: '2.11.4 Danh sách event và xử lý event',
    }
    for idx, text in danh_so.items():
        cu = paras[idx].text.strip()
        assert cu and cu[0].isdigit(), 'Đoạn %d không phải tiêu đề mục: %r' % (idx, cu)
        set_para_text(paras[idx]._p, text, doc)

    # ------------------------------------------- #
    # 2. Đánh số lại hình + đổi ảnh cho đúng bản   #
    # ------------------------------------------- #
    caption = {
        154: 'Hình 15: Biểu đồ Use Case — Nhập quốc gia từ file Excel',
        163: 'Hình 16: Cửa sổ Import quốc gia khi vừa mở',
        171: 'Hình 19: Biểu đồ Use Case — Xuất danh sách ra Excel',
        180: 'Hình 20: Cửa sổ Chọn trường xuất file',
        187: 'Hình 21: Biểu đồ Use Case — Xem chi tiết quốc gia',
        196: 'Hình 22: Cửa sổ Xem quốc gia ở chế độ chỉ đọc',
        210: 'Hình 23: Cửa sổ Tuỳ chỉnh cột — STT, Tên quốc gia và Hành động bị khoá',
    }
    for idx, text in caption.items():
        assert paras[idx].text.strip().startswith('Hình'), 'Đoạn %d không phải caption' % idx
        set_para_text(paras[idx]._p, text, doc)

    # Ảnh màn danh sách (Hình 2) + 2 ảnh popup ở 2.10 / 2.11: chụp lại theo bản hiện tại
    for para_idx, file_name in ((39, '01-danh-sach.png'), (195, '08-xem-lich-su-trong-popup.png'),
                                (209, '06-cau-hinh-cot.png')):
        replace_image_blob(doc, para_idx, shot(file_name))

    # Chèn ảnh thật thay 2 dòng "Thêm hình sau" / "Thêm sau"
    # 2.8 — cửa sổ Import: 3 ảnh (mở, load lên bảng, validate)
    p162 = paras[162]._p
    assert paras[162].text.strip() == 'Thêm hình sau'
    p162.addprevious(clone_image_para(T_IMG, shot('02-import-buoc1.png'), doc))
    p162.getparent().remove(p162)
    p163 = paras[163]._p          # caption Hình 16 (đã đổi ở trên)
    for img, cap in (('03-import-load.png', 'Hình 17: Dữ liệu đã được load lên bảng xem trước'),
                     ('04-import-validate.png',
                      'Hình 18: Kết quả Validate — mỗi dòng lỗi nêu rõ lý do')):
        p163.addnext(clone_para(T_CAP, cap, doc))
        p163.addnext(clone_image_para(T_IMG, shot(img), doc))

    # 2.9 — cửa sổ Chọn trường xuất file
    p179 = paras[179]._p
    assert paras[179].text.strip() == 'Thêm sau'
    p179.addprevious(clone_image_para(T_IMG, shot('05-xuat-excel.png'), doc))
    p179.getparent().remove(p179)

    # ----------------------------------------- #
    # 3. Sửa nội dung các đoạn mô tả sai         #
    # ----------------------------------------- #
    sua = {
        'Modal Chọn trường xuất Excel được mở ngay trên màn hình danh sách theo đường dẫn ở trên.':
            'Cửa sổ Chọn trường xuất file được mở ngay trên màn hình danh sách theo đường dẫn ở '
            'trên.',
    }
    for p in doc.paragraphs:
        t = p.text.strip()
        if t in sua:
            set_para_text(p._p, sua.pop(t), doc)

    # ------------------------------------------------- #
    # 4. Ma trận phân quyền — bổ sung FR-08 đến FR-11    #
    # ------------------------------------------------- #
    mt = doc.tables[2]
    rows = [[c.text.strip() for c in r.cells] for r in mt.rows]
    cho = 'Chờ cập nhật theo ma trận phân quyền chi tiết'
    rows += [
        ['FR-08 Nhập quốc gia từ file Excel', cho],
        ['FR-09 Xuất danh sách ra Excel', cho],
        ['FR-10 Xem chi tiết quốc gia', cho],
        ['FR-11 Tuỳ chỉnh cột hiển thị', cho],
    ]
    thay_bang(doc, mt, rows)

    # ------------------------------------------------------------------ #
    # 5. 2.1.3 — bảng mô tả giao diện màn danh sách (đủ nút thanh công cụ) #
    # ------------------------------------------------------------------ #
    thay_bang(doc, doc.tables[4], [
        ['STT', 'Tên đối tượng', 'Loại', 'Trạng thái', 'Mô tả'],
        ['1', 'STT', 'Table/Grid', 'Read-only', 'Số thứ tự, chạy liên tục qua các trang. Luôn '
                                                'hiển thị, không ẩn được.'],
        ['2', 'Tên quốc gia', 'Table/Grid', 'Read-only', 'Tên đầy đủ. Luôn hiển thị, sắp xếp '
                                                         'được. Bấm vào tên mở cửa sổ Xem quốc '
                                                         'gia (xem 2.10).'],
        ['3', 'Mã quốc gia', 'Table/Grid', 'Read-only', 'Mã do người dùng đặt, sắp xếp được.'],
        ['4', 'Mã bưu chính', 'Table/Grid', 'Read-only', 'Mặc định ẩn, bật ở cửa sổ Tuỳ chỉnh '
                                                         'cột.'],
        ['5', 'Người cập nhật', 'Table/Grid', 'Read-only', 'Người sửa bản ghi gần nhất. Mặc định '
                                                           'ẩn.'],
        ['6', 'Ngày cập nhật', 'Table/Grid', 'Read-only', 'Thời điểm sửa gần nhất, dạng '
                                                          'dd/mm/yyyy hh:mm. Mặc định ẩn, sắp '
                                                          'xếp được.'],
        ['7', 'Người tạo', 'Table/Grid', 'Read-only', 'Người đã thêm bản ghi.'],
        ['8', 'Ngày tạo', 'Table/Grid', 'Read-only', 'Thời điểm thêm, dạng dd/mm/yyyy hh:mm, '
                                                     'sắp xếp được.'],
        ['9', 'Trạng thái', 'Table/Grid', 'Read-only', 'Hoạt động hoặc Khóa.'],
        ['10', 'Hành động', 'Table/Grid', 'Read-only', 'Sửa, Xóa hiện thẳng; Khóa / Mở khóa và '
                                                       'Lịch sử nằm trong nút ba chấm. Bản ghi '
                                                       'đã Khóa không còn Sửa và Xóa; quốc gia '
                                                       'đã có Khu vực trực thuộc không còn Xóa.'],
        ['11', 'Nút Tạo mới', 'Button', 'Enable', 'Mở cửa sổ thêm mới quốc gia.'],
        ['12', 'Nút Xuất Excel', 'Button', 'Enable', 'Mở cửa sổ Chọn trường xuất file (xem 2.9).'],
        ['13', 'Nút Import Excel', 'Button', 'Enable', 'Mở cửa sổ Import quốc gia (xem 2.8).'],
        ['14', 'Biểu tượng Tuỳ chỉnh cột', 'Icon Button', 'Enable', 'Mở cửa sổ bật/tắt và sắp '
                                                                    'xếp cột (xem 2.11).'],
        ['15', 'Phân trang', 'Pagination', 'Enable', 'Nút về đầu / lùi / số trang / tiến / về '
                                                     'cuối và ô chọn số dòng mỗi trang.'],
        ['16', 'Trạng thái rỗng', 'Label', 'Hiển thị', 'Hiện “Không có dữ liệu phù hợp bộ lọc.” '
                                                       'khi danh sách trống.'],
    ])

    # ---------------------------------------------- #
    # 6. 2.2.3 — ô tìm kiếm nhanh chỉ tìm theo TÊN    #
    # ---------------------------------------------- #
    thay_bang(doc, doc.tables[7], [
        ['STT', 'Tên đối tượng', 'Loại', 'Trạng thái', 'Phạm vi', 'Mô tả'],
        ['1', 'Ô tìm kiếm nhanh', 'Textbox', 'Enable', '0–255 ký tự',
         'Tìm theo TÊN quốc gia (khớp một phần). Ô này KHÔNG tìm theo mã quốc gia.'],
        ['2', 'Trạng thái', 'Dropdown', 'Enable', 'Hoạt động / Khóa',
         'Bỏ trống thì hiện cả hai trạng thái. Xoá lựa chọn bằng dấu x trong ô.'],
        ['3', 'Nút Tìm kiếm', 'Button', 'Enable', '–', 'Áp dụng các tiêu chí.'],
        ['4', 'Nút Làm mới', 'Button', 'Enable', '–', 'Xóa hết tiêu chí VÀ nạp lại danh sách '
                                                      'ngay.'],
    ])

    # ------------------------------------- #
    # 7. 2.8 Import — 3 bảng                #
    # ------------------------------------- #
    thay_bang(doc, doc.tables[24], [
        ['Mục', 'Nội dung'],
        ['Tên chức năng', 'Nhập quốc gia từ file Excel'],
        ['Mô tả', 'Thêm nhiều quốc gia cùng lúc từ file Excel theo mẫu. Có bước kiểm tra dữ liệu '
                  '(Validate) trước khi ghi để người dùng sửa lỗi mà không tạo ra bản ghi rác.'],
        ['Tác nhân', CHUA_PHAN_QUYEN],
        ['Điều kiện ban đầu', 'Đang ở màn Danh mục quốc gia và đã chuẩn bị file theo mẫu.'],
        ['Dòng sự kiện chính',
         '1. Người dùng bấm Import Excel.\n2. Bấm Tải file mẫu và điền dữ liệu vào file.\n'
         '3. Bấm Chọn file Excel rồi bấm Load lên bảng.\n4. Bấm Validate; hệ thống kiểm tra từng '
         'dòng, đánh dấu dòng lỗi kèm lý do và khoá các dòng hợp lệ.\n5. Sửa dòng lỗi rồi Validate '
         'lại, hoặc bấm Bỏ dòng lỗi.\n6. Bấm Import; hệ thống ghi các dòng hợp lệ và báo tổng số '
         'dòng, số dòng thành công, số dòng thất bại.'],
        ['Dòng sự kiện phụ',
         '• File sai định dạng, hỏng hoặc đang mở bằng Excel → báo lỗi không đọc được file.\n'
         '• Sheet không có dòng dữ liệu → báo “Sheet dữ liệu trống hoặc thiếu dòng.”.\n'
         '• Toàn bộ dòng lỗi → nút Import không dùng được, không ghi bản ghi nào.\n'
         '• Một phần dòng lỗi → các dòng hợp lệ vẫn được ghi, dòng lỗi thì không.'],
        ['Yêu cầu đặc biệt',
         'Bước Validate KHÔNG ghi dữ liệu. Bản ghi tạo từ import luôn ở trạng thái Hoạt động (file '
         'mẫu không có cột Trạng thái); mã quốc gia được tự động chuyển thành chữ IN HOA. Khi bấm '
         'Import, máy chủ kiểm tra lại toàn bộ dòng gửi lên chứ không tin kết quả Validate.'],
    ])

    thay_bang(doc, doc.tables[25], [
        ['STT', 'Tên đối tượng', 'Loại', 'Trạng thái', 'Bắt buộc', 'Giá trị ban đầu', 'Mô tả'],
        ['1', 'Nút Import Excel', 'Button', 'Enable', '–', 'Hiển thị',
         'Nằm trên thanh công cụ của bảng danh sách. Mở cửa sổ Import quốc gia.'],
        ['2', 'Nút Tải file mẫu', 'Button', 'Enable', '–', 'Hiển thị',
         'Tải file Mau_import_quoc_gia.xlsx gồm STT, Tên quốc gia, Mã quốc gia, Mã bưu chính; hai '
         'cột bắt buộc có dấu sao đỏ.'],
        ['3', 'Nút Chọn file Excel', 'Button', 'Enable', 'Có', 'Chưa chọn file',
         'Chọn file .xlsx từ máy; tên file hiện ngay cạnh nút.'],
        ['4', 'Nút Load lên bảng', 'Button', 'Enable / Disable', '–', 'Mờ khi chưa chọn file',
         'Đọc file và đổ dữ liệu lên bảng xem trước.'],
        ['5', 'Nút Validate', 'Button', 'Enable / Disable', '–', 'Mờ khi bảng chưa có dòng nào',
         'Kiểm tra từng dòng; dòng hợp lệ bị khoá lại không sửa được nữa.'],
        ['6', 'Nút Import', 'Button', 'Enable / Disable', '–', 'Mờ khi chưa có dòng hợp lệ',
         'Ghi các dòng hợp lệ vào hệ thống.'],
        ['7', 'Bảng xem trước', 'Table/Grid', 'Enable', '–', 'Trống',
         'Ba cột Tên quốc gia, Mã quốc gia, Mã bưu chính. Sửa trực tiếp được ở các dòng lỗi.'],
        ['8', 'Ô Tổng / Hợp lệ / Lỗi', 'Label', 'Read-only', '–', 'Tổng: 0',
         'Tổng số dòng đọc được; sau khi Validate hiện thêm số dòng hợp lệ và số dòng lỗi.'],
        ['9', 'Thông báo lỗi theo dòng', 'Label', 'Hiển thị', '–', 'Ẩn',
         'Dòng lỗi được tô nền đỏ, lý do ghi ngay dưới ô sai kèm số dòng của file.'],
        ['10', 'Nút Bỏ dòng lỗi', 'Button', 'Enable', '–', 'Hiển thị',
         'Loại toàn bộ dòng đang lỗi khỏi bảng xem trước.'],
        ['11', 'Nút Xoá trạng thái validate', 'Button', 'Enable', '–', 'Hiển thị',
         'Bỏ khoá các dòng hợp lệ để sửa lại toàn bộ bảng.'],
        ['12', 'Nút Làm mới / Đóng', 'Button', 'Enable', '–', 'Hiển thị',
         'Xoá trạng thái hiện tại / đóng cửa sổ.'],
    ])

    thay_bang(doc, doc.tables[26], [
        ['STT', 'Event', 'Loại event', 'Xử lý event'],
        ['1', 'Bấm Import Excel', 'Click', 'After:\n– Mở cửa sổ Import quốc gia ở trạng thái '
                                           'rỗng.'],
        ['2', 'Bấm Tải file mẫu', 'Click', 'After:\n– Tải về file mẫu có đủ các trường như màn '
                                           'Thêm mới.'],
        ['3', 'Bấm Load lên bảng', 'Click',
         'During:\n– File sai định dạng / không đọc được → hiển thị thông báo file không hợp lệ.\n'
         '– Sheet không có dòng dữ liệu → hiển thị “Sheet dữ liệu trống hoặc thiếu dòng.”.\n'
         'After:\n– Đổ dữ liệu lên bảng xem trước và cập nhật ô Tổng.'],
        ['4', 'Bấm Validate', 'Click',
         'During:\n– Kiểm tra từng dòng theo đúng quy tắc của chức năng Thêm mới: bắt buộc Tên và '
         'Mã quốc gia; chống trùng tên, trùng mã với dữ liệu đã có VÀ giữa các dòng trong file; '
         'Mã bưu chính chỉ nhận chữ số, tối đa 50 chữ số; Tên quốc gia tối đa 255 ký tự.\n'
         'After:\n– Đánh dấu dòng lỗi kèm lý do theo số dòng, khoá các dòng hợp lệ, hiện số dòng '
         'Hợp lệ / Lỗi.\n– KHÔNG ghi bất kỳ bản ghi nào ở bước này.'],
        ['5', 'Bấm Import', 'Click',
         'During:\n– Máy chủ kiểm tra lại toàn bộ dòng gửi lên và bỏ qua các dòng lỗi.\n'
         '– Không còn dòng hợp lệ nào → dừng và báo không có dòng nào hợp lệ để import.\n'
         'After:\n– Ghi các dòng hợp lệ thành quốc gia mới ở trạng thái Hoạt động, người tạo là '
         'người đang thực hiện.\n– Hiển thị kết quả tổng / thành công / thất bại, đóng cửa sổ và '
         'nạp lại danh sách.'],
    ])

    # ------------------------------------- #
    # 8. 2.9 Xuất Excel — 3 bảng            #
    # ------------------------------------- #
    thay_bang(doc, doc.tables[27], [
        ['Mục', 'Nội dung'],
        ['Tên chức năng', 'Xuất danh sách quốc gia ra Excel'],
        ['Mô tả', 'Xuất kết quả đang lọc ra file Excel, cho phép chọn trường và thứ tự cột.'],
        ['Tác nhân', CHUA_PHAN_QUYEN],
        ['Điều kiện ban đầu', 'Đang ở màn Danh mục quốc gia.'],
        ['Dòng sự kiện chính',
         '1. Người dùng bấm Xuất Excel.\n2. Hệ thống mở cửa sổ Chọn trường xuất file, tích sẵn '
         'đúng các cột đang hiển thị trên bảng.\n3. Người dùng chọn các trường theo trình tự mong '
         'muốn.\n4. Bấm Xuất file; hệ thống sinh file danh_muc_quoc_gia.xlsx theo đúng bộ lọc đang '
         'áp dụng và đúng thứ tự trường đã chọn.'],
        ['Dòng sự kiện phụ',
         '• Bỏ chọn hết trường → nút Xuất file mờ đi, không bấm được.\n'
         '• Bấm Đóng → đóng cửa sổ, không xuất gì.'],
        ['Yêu cầu đặc biệt', 'File xuất chứa toàn bộ kết quả lọc, không giới hạn ở trang đang '
                             'xem. Màn hình chỉ xuất Excel, không có xuất PDF.'],
    ])

    thay_bang(doc, doc.tables[28], [
        ['STT', 'Tên đối tượng', 'Loại', 'Trạng thái', 'Giá trị ban đầu', 'Mô tả'],
        ['1', 'Nút Xuất Excel', 'Button', 'Enable', 'Hiển thị',
         'Nằm trên thanh công cụ của bảng danh sách. Mở cửa sổ Chọn trường xuất file.'],
        ['2', 'Ô Trường xuất', 'Dropdown chọn nhiều', 'Enable', 'Tích sẵn các cột đang hiển thị '
                                                                'trên bảng',
         'Tám trường: Tên quốc gia, Mã quốc gia, Mã bưu chính, Trạng thái, Người tạo, Ngày tạo, '
         'Người cập nhật, Ngày cập nhật.'],
        ['3', 'Dòng “Thứ tự cột trong file”', 'Label', 'Hiển thị', 'Theo trường đang chọn',
         'Cho biết thứ tự cột sẽ ghi ra file. Muốn đổi vị trí thì bỏ chọn rồi chọn lại theo trình '
         'tự mong muốn.'],
        ['4', 'Nút Chọn tất cả / Bỏ chọn hết', 'Button', 'Enable', 'Hiển thị',
         'Tích hết / bỏ tích toàn bộ trường.'],
        ['5', 'Nút Xuất file', 'Button', 'Enable / Disable', 'Mờ khi chưa chọn trường nào',
         'Sinh file Excel và tải về máy.'],
        ['6', 'Nút Đóng', 'Button', 'Enable', 'Hiển thị', 'Đóng cửa sổ, không xuất.'],
    ])

    thay_bang(doc, doc.tables[29], [
        ['STT', 'Event', 'Loại event', 'Xử lý event'],
        ['1', 'Bấm Xuất Excel', 'Click', 'After:\n– Mở cửa sổ Chọn trường xuất file, tích sẵn '
                                         'đúng các cột đang hiển thị trên bảng.'],
        ['2', 'Chọn / bỏ chọn trường', 'Change',
         'During:\n– Ghi nhận thứ tự chọn để quyết định thứ tự cột trong file.\n'
         '– Bỏ chọn hết thì nút Xuất file mờ đi.'],
        ['3', 'Bấm Xuất file', 'Click',
         'During:\n– Sinh file theo ĐÚNG bộ lọc đang áp dụng và đúng thứ tự trường đã chọn.\n'
         'After:\n– Tải file danh_muc_quoc_gia.xlsx về máy và hiển thị thông báo xuất thành '
         'công.'],
        ['4', 'Bấm Đóng', 'Click', 'After:\n– Đóng cửa sổ, không xuất gì.'],
    ])

    # ------------------------------------------------------------------- #
    # 9. 2.10 Xem chi tiết + 2.11 Tuỳ chỉnh cột (bảng nằm trong khối sdt)  #
    # ------------------------------------------------------------------- #
    sdt_tbls = sdt_tables(doc)
    assert len(sdt_tbls) == 7, 'Số bảng trong khối sdt đổi: %d' % len(sdt_tbls)
    t_210_mo, t_210_ui, t_210_ev, t_211_mo, t_211_ui, t_211_ev = sdt_tbls[:6]

    thay_bang(doc, t_210_mo, [
        ['Mục', 'Nội dung'],
        ['Tên chức năng', 'Xem chi tiết quốc gia'],
        ['Mô tả', 'Hiển thị toàn bộ thông tin của một quốc gia ở chế độ chỉ đọc, kèm lịch sử thay '
                  'đổi của chính bản ghi đó.'],
        ['Tác nhân', CHUA_PHAN_QUYEN],
        ['Điều kiện ban đầu', 'Đang ở màn Danh mục quốc gia.'],
        ['Dòng sự kiện chính',
         '1. Người dùng bấm vào TÊN quốc gia trên bảng danh sách.\n2. Hệ thống mở cửa sổ “Xem quốc '
         'gia” ngay trên màn danh sách.\n3. Mọi thông tin hiển thị ở chế độ chỉ đọc; khối Lịch sử '
         'mở / thu gọn được.'],
        ['Dòng sự kiện phụ',
         '• Quốc gia đã Khóa vẫn xem được bình thường.\n'
         '• Bấm Đóng → quay lại danh sách, giữ nguyên bộ lọc và trang đang xem.'],
        ['Yêu cầu đặc biệt', 'Đây là cửa sổ mở chồng lên màn danh sách, KHÔNG phải một trang '
                             'riêng: không có nút Lưu và không có nút Quay lại.'],
    ])

    thay_bang(doc, t_210_ui, [
        ['STT', 'Tên đối tượng', 'Loại', 'Trạng thái', 'Giá trị ban đầu', 'Mô tả'],
        ['1', 'Tên quốc gia, Mã quốc gia, Mã bưu chính', 'Textbox', 'Disable (chỉ đọc)',
         'Dữ liệu hiện tại', 'Không gõ được; cửa sổ không có nút Lưu.'],
        ['2', 'Khối Lịch sử', 'Vùng mở rộng', 'Enable', 'Thu gọn',
         'Bấm “Xem lịch sử” để mở; bên trong có nút Bộ lọc, Làm mới, Thu gọn. Con số cạnh chữ Lịch '
         'sử là số lần thay đổi đã ghi nhận.'],
        ['3', 'Trạng thái rỗng của lịch sử', 'Label', 'Hiển thị', 'Ẩn',
         'Hiện “Chưa có lịch sử thao tác nào.” khi bản ghi chưa từng thay đổi.'],
        ['4', 'Nút Đóng', 'Button', 'Enable', 'Hiển thị', 'Đóng cửa sổ, quay về danh sách.'],
    ])

    thay_bang(doc, t_210_ev, [
        ['STT', 'Event', 'Loại event', 'Xử lý event'],
        ['1', 'Bấm tên quốc gia', 'Click', 'After:\n– Mở cửa sổ Xem quốc gia ở chế độ chỉ đọc.'],
        ['2', 'Bấm Xem lịch sử', 'Click',
         'After:\n– Nạp danh sách thay đổi của bản ghi, mới nhất trên cùng.\n– Chưa có thay đổi '
         'nào thì hiện “Chưa có lịch sử thao tác nào.”.'],
        ['3', 'Bấm Đóng', 'Click', 'After:\n– Đóng cửa sổ; danh sách phía sau giữ nguyên bộ lọc '
                                   'và trang.'],
    ])

    thay_bang(doc, t_211_mo, [
        ['Mục', 'Nội dung'],
        ['Tên chức năng', 'Tuỳ chỉnh cột hiển thị'],
        ['Mô tả', 'Cho phép bật / tắt và sắp xếp các cột của bảng danh sách. Cấu hình ghi nhớ theo '
                  'từng người dùng.'],
        ['Tác nhân', CHUA_PHAN_QUYEN],
        ['Điều kiện ban đầu', 'Đang ở màn Danh mục quốc gia.'],
        ['Dòng sự kiện chính',
         '1. Người dùng bấm biểu tượng Tuỳ chỉnh cột.\n2. Hệ thống hiển thị danh sách cột kèm ô '
         'tích và tay nắm kéo.\n3. Người dùng tích / bỏ tích, kéo đổi thứ tự.\n4. Bấm Lưu; bảng vẽ '
         'lại theo cấu hình mới.'],
        ['Dòng sự kiện phụ',
         '• Cột STT, Tên quốc gia và Hành động bị khoá: không bỏ tích và không kéo đổi vị trí '
         'được.\n• Bấm Đóng → bỏ mọi thay đổi chưa lưu.'],
        ['Yêu cầu đặc biệt', 'Cửa sổ chỉ có hai nút Lưu và Đóng, không có nút khôi phục cấu hình '
                             'mặc định.'],
    ])

    thay_bang(doc, t_211_ui, [
        ['STT', 'Tên đối tượng', 'Loại', 'Trạng thái', 'Giá trị ban đầu', 'Mô tả'],
        ['1', 'Biểu tượng Tuỳ chỉnh cột', 'Icon Button', 'Enable', 'Hiển thị',
         'Nút ngoài cùng bên phải thanh công cụ của bảng.'],
        ['2', 'Danh sách cột', 'Table/Grid', 'Enable', 'Theo cấu hình đã lưu',
         'Mỗi dòng gồm ô tích, tên cột và tay nắm kéo.'],
        ['3', 'Ô tích cột STT, Tên quốc gia, Hành động', 'Checkbox', 'Disable', 'Đã tích',
         'Cột bắt buộc: hiển thị xám kèm biểu tượng ổ khoá, không bỏ tích được.'],
        ['4', 'Ô tích các cột còn lại', 'Checkbox', 'Enable', 'Theo cấu hình đã lưu',
         'Tích để hiện, bỏ tích để ẩn cột.'],
        ['5', 'Tay nắm kéo', 'Icon', 'Enable', 'Hiển thị',
         'Kéo thả để đổi thứ tự cột; dòng bị khoá không kéo được.'],
        ['6', 'Nút Lưu', 'Button', 'Enable', 'Hiển thị', 'Ghi cấu hình, vẽ lại bảng và đóng cửa '
                                                         'sổ.'],
        ['7', 'Nút Đóng', 'Button', 'Enable', 'Hiển thị', 'Đóng cửa sổ, bỏ các thay đổi chưa '
                                                          'lưu.'],
    ])

    thay_bang(doc, t_211_ev, [
        ['STT', 'Event', 'Loại event', 'Xử lý event'],
        ['1', 'Bấm biểu tượng Tuỳ chỉnh cột', 'Click', 'After:\n– Mở cửa sổ, nạp cấu hình cột đã '
                                                       'lưu của chính người dùng.'],
        ['2', 'Tích / bỏ tích cột', 'Change',
         'During:\n– Cột STT, Tên quốc gia và Hành động không đổi được.\n– Các cột khác cập nhật '
         'trạng thái tạm.'],
        ['3', 'Kéo đổi thứ tự cột', 'Drag', 'After:\n– Cập nhật thứ tự tạm; các dòng bị khoá giữ '
                                            'nguyên vị trí.'],
        ['4', 'Bấm Lưu', 'Click', 'After:\n– Ghi cấu hình theo từng người dùng, vẽ lại bảng theo '
                                  'cột và thứ tự mới.'],
    ])

    doc.save(out)

    # ---------------- kiểm tra ---------------- #
    check = Document(out)
    texts = [p.text.strip() for p in check.paragraphs]
    for must in ['2.8.3 Layout màn hình', '2.8.4 Mô tả chi tiết giao diện',
                 '2.8.5 Danh sách event và xử lý event', '2.9.5 Danh sách event và xử lý event',
                 '2.11.2 Layout màn hình', 'Hình 22: Cửa sổ Xem quốc gia ở chế độ chỉ đọc']:
        assert must in texts, 'Thiếu: ' + must
    for khong_duoc_con in ['Thêm hình sau', 'Thêm sau', '2.10.6 Danh sách event và xử lý event',
                           '2.11.5 Danh sách event và xử lý event']:
        assert khong_duoc_con not in texts, 'Vẫn còn: ' + khong_duoc_con
    # "2.4.3 Layout màn hình" chỉ được còn ĐÚNG 1 lần (của mục 2.4 Chỉnh sửa)
    assert texts.count('2.4.3 Layout màn hình') == 1, 'Mục 2.11 vẫn mang số 2.4.x'
    print('OK — số ảnh:', len(check.inline_shapes), '| số bảng (ngoài sdt):', len(check.tables))


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2])
