# Sinh 5 file import cây 4 cấp theo mẫu import của 4 màn HRM (header = alias của importColumns)
# python3 -I sinh_file_import_cay.py <file17> <thư mục ra>
import openpyxl,sys,re,os,collections
from openpyxl.styles import Font,PatternFill
src,outdir=sys.argv[1:3]
# Header PHẢI có đủ mọi cột importColumns của màn (kể cả cột không bắt buộc) — thiếu 1 cột là trình đọc báo 'File không đúng mẫu'.
# Sheet tên 'Data', cột đầu STT (khuôn file mẫu buildImportTemplate).
clean=lambda s: re.sub(r'\s+',' ',str(s).strip()) if s is not None else ''
rs=list(openpyxl.load_workbook(src,read_only=True,data_only=True)['NewERP Green Power'].iter_rows(values_only=True))[1:]
paths=sorted({tuple(clean(x) for x in r[1:5]) for r in rs if all(clean(x) for x in r[1:5])},key=lambda p:[x.casefold() for x in p])
BARCODE={'Máy móc - thiết bị':1,'Dây chuyền - hệ thống':1,'Phần mềm':1,'Dụng cụ - đồ nghề':2,
         'Vật tư - linh kiện - phụ kiện':4,'Phụ tùng - linh kiện ô tô - xe máy':4}
def codes(items,prefix):
    seen={}
    for it in items:
        if it not in seen: seen[it]=f'{prefix}.{len(seen)+1:04d}'
    return seen
tc=codes([p[0] for p in paths],'TCHH'); ncn=codes([p[1] for p in paths],'NCN')
nsp=codes([p[2] for p in paths],'NSP'); lsp=codes([p[3] for p in paths],'LSP')
par=lambda k: {p[k]:p[k-1] for p in paths}
H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='1ABC9C')
def write(name,head,rows,widths):
    head=['STT']+head+['Diễn giải']; rows=[(i+1,*r,'') for i,r in enumerate(rows)]; widths=[6]+widths+[20]
    wb=openpyxl.Workbook(); ws=wb.active; ws.title='Data'; ws.append(head)
    for c in ws[1]: c.font=H; c.fill=F
    for r in rows: ws.append(r)
    for i,w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
    ws.freeze_panes='A2'; wb.save(os.path.join(outdir,name)); print(name,len(rows))
missing=[n for n in tc if n not in BARCODE]; assert not missing,missing
write('1_Tinh_chat_hang_hoa.xlsx',['Mã tính chất hàng hóa *','Tính chất hàng hóa *','Mẫu in barcode (1-6) *','Trạng thái *'],
      [(c,n,BARCODE[n],'Hoạt động') for n,c in tc.items()],[22,40,22,14])
p1=par(1); write('2_Nhom_chuc_nang.xlsx',['Mã nhóm chức năng *','Nhóm chức năng *','Mã tính chất hàng hóa *','Trạng thái *'],
      [(c,n,tc[p1[n]],'Hoạt động') for n,c in ncn.items()],[20,40,22,14])
p2=par(2); write('3_Nhom_san_pham.xlsx',['Mã nhóm sản phẩm *','Nhóm sản phẩm *','Mã nhóm chức năng *','Trạng thái *'],
      [(c,n,ncn[p2[n]],'Hoạt động') for n,c in nsp.items()],[20,45,20,14])
p3=par(3); L=[(c,n,nsp[p3[n]],'Không bắt buộc','','Hoạt động') for n,c in lsp.items()]
hl=['Mã loại sản phẩm *','Loại sản phẩm *','Mã nhóm sản phẩm *','Serial *','Số ngày tính tồn lâu','Trạng thái *']
write('4a_Loai_san_pham_1-500.xlsx',hl,L[:500],[20,50,20,16,18,14])
write('4b_Loai_san_pham_501-%d.xlsx'%len(L),hl,L[500:],[20,50,20,16,18,14])
# Bảng tra đủ đường để soát
wb=openpyxl.Workbook(); ws=wb.active; ws.title='Cay 4 cap'
ws.append(['Mã TC','Tính chất','Mã NCN','Nhóm chức năng','Mã NSP','Nhóm sản phẩm','Mã LSP','Loại sản phẩm'])
for c in ws[1]: c.font=H; c.fill=F
for p in paths: ws.append([tc[p[0]],p[0],ncn[p[1]],p[1],nsp[p[2]],p[2],lsp[p[3]],p[3]])
for i,w in enumerate([11,32,10,32,10,40,10,50]): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:H{ws.max_row}'; wb.save(os.path.join(outdir,'0_Bang_tra_cay_4_cap.xlsx'))
print('TC',len(tc),'NCN',len(ncn),'NSP',len(nsp),'LSP',len(lsp),'paths',len(paths))
