# Hàng hoá đang dùng (products.status<>0) KHÔNG có mã trong sheet NewERP Green Power (file 17/10)
# python3 -I xuat_hang_ngoai_file.py <file17> <inuse.tsv> <groups_full.tsv> <out.xlsx>
import openpyxl,sys,re,collections
from openpyxl.styles import Font,PatternFill,Alignment
f17,itsv,gt,out=sys.argv[1:5]
n=lambda s: re.sub(r'\s+',' ',str(s).strip()).casefold()
rs=list(openpyxl.load_workbook(f17,read_only=True,data_only=True)['NewERP Green Power'].iter_rows(values_only=True))[1:]
codes={n(r[6]) for r in rs if r[6]}
g2l=collections.defaultdict(set)
for r in rs:
    if r[8] not in (None,'') and r[4]: g2l[int(float(r[8]))].add(re.sub(r'\s+',' ',str(r[4]).strip()))
G={}
for l in open(gt,encoding='utf-8'):
    p=l.rstrip('\n').split('\t'); G[int(p[0])]=(p[1],int(p[2]))
rows=[]
for l in open(itsv,encoding='utf-8'):
    pid,code,name,gid,brand,manu=l.rstrip('\n').split('\t')
    if n(code) in codes: continue
    g=int(gid) if gid else None
    L=g2l.get(g,set())
    kind='Nhóm cũ ứng 1 Loại' if len(L)==1 else ('Nhóm cũ ứng nhiều Loại' if L else 'Nhóm cũ chưa có Loại')
    rows.append((int(pid),code,name,brand,manu,g or '',G.get(g,('',''))[0] if g else '',kind,len(L),next(iter(L)) if len(L)==1 else ''))
rows.sort(key=lambda r:(r[7],r[6],r[2]))
o=openpyxl.Workbook(); ws=o.active; ws.title='Hàng ngoài file'
H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='1ABC9C')
c=collections.Counter(r[7] for r in rows)
ws.append([f'{len(rows):,} hàng đang dùng (DB hrm_erp_giacty 19/10, products.status<>0) có mã KHÔNG nằm trong sheet NewERP Green Power — file ERP_final 17/10. '.replace(',','.')+' · '.join(f'{k}: {v:,}'.replace(',','.') for k,v in sorted(c.items()))])
ws['A1'].font=Font(italic=True)
ws.append(['ID hàng','Mã hàng','Tên hàng','Thương hiệu','Hãng','ID nhóm cũ','Nhóm hàng hoá cũ','Tình trạng nhóm cũ','Số Loại của nhóm (theo file)','Loại hàng hoá (nhóm ứng 1 Loại)'])
for x in ws[2]: x.font=H; x.fill=F; x.alignment=Alignment(wrap_text=True,vertical='center')
for r in rows: ws.append(list(r))
for i,w in enumerate([9,22,55,14,22,9,38,22,12,36]): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
ws.freeze_panes='D3'; ws.auto_filter.ref=f'A2:J{ws.max_row}'
o.save(out); print(len(rows),c)
