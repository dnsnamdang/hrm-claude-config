# Xuất hàng đang dùng nhóm c (nhóm cũ ứng nhiều Loại) + d (nhóm cũ chưa có Loại) cho nghiệp vụ gán Loại SP
# python3 -I xuat_hang_c_d.py <file17> <groups.tsv> <inuse.tsv> <out.xlsx>
import openpyxl,sys,collections,re
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.datavalidation import DataValidation
f17,gtsv,itsv,out=sys.argv[1:5]
norm=lambda s: re.sub(r'\s+',' ',str(s).strip()).casefold() if s is not None and str(s).strip()!='' else None
clean=lambda s: re.sub(r'\s+',' ',str(s).strip()) if s is not None else ''
rs=list(openpyxl.load_workbook(f17,read_only=True,data_only=True)['NewERP Green Power'].iter_rows(values_only=True))[1:]
G={}
for l in open(gtsv,encoding='utf-8'):
    i,st,nm,*_=l.rstrip('\n').split('\t'); G[int(i)]=nm
path={}; lcount=collections.Counter(); codes=set(); g2l=collections.defaultdict(collections.Counter)
for r in rs:
    if not r[4]: continue
    l=norm(r[4]); path.setdefault(l,tuple(clean(x) for x in r[1:5])); lcount[l]+=1
    if r[6]: codes.add(clean(r[6]).casefold())
    g=r[8]
    if g not in (None,''): g2l[int(float(g))][l]+=1
C=[];D=[];CAND={}
for line in open(itsv,encoding='utf-8'):
    pid,code,name,gid,brand,manu=line.rstrip('\n').split('\t')
    if code.strip().casefold() in codes: continue
    gid=int(gid) if gid else None
    if gid in g2l and len(g2l[gid])==1: continue
    gname=G.get(gid,'') if gid else ''
    if gid in g2l:
        mc=g2l[gid].most_common(); tot=sum(n for _,n in mc)
        for l,_ in mc: C.append((int(pid),code,name,brand,manu,gid,gname,path[l][3]))
    else:
        D.append((int(pid),code,name,brand,manu,gid or '',gname,''))
C.sort(key=lambda x:(x[6],x[2],x[0])); D.sort(key=lambda x:(x[6],x[2]))
o=openpyxl.Workbook(); o.remove(o.active)
H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='1ABC9C'); FILL=PatternFill('solid',fgColor='FFF2CC')
NT=len(path)+2
def sheet(name,note,head,data,widths,fillcol,per_group=False):
    ws=o.create_sheet(name); ws.append([note]); ws['A1'].font=Font(italic=True); ws.append(list(head))
    for c in ws[2]: c.font=H; c.fill=F; c.alignment=Alignment(wrap_text=True,vertical='center')
    for d in data: ws.append(list(d))
    for i,w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
    for row in ws.iter_rows(min_row=3):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    L=openpyxl.utils.get_column_letter(fillcol)
    for i in range(3,ws.max_row+1): ws[f'{L}{i}'].fill=FILL
    if per_group:
        GR=PatternFill('solid',fgColor='F2F2F2'); prev=None; on=False
        for i in range(3,ws.max_row+1):
            ws[f'{L}{i}'].fill=PatternFill()
            pid=ws.cell(i,1).value
            if pid!=prev: on=not on; prev=pid
            if on:
                for c in ws[i]: c.fill=GR
    else:
        dv=DataValidation(type='list',formula1=f"='Cây 4 cấp'!$A$3:$A${NT}",allow_blank=True,showErrorMessage=True,
                          errorTitle='Loại SP không có trong cây',error='Chọn Loại SP từ danh sách (sheet "Cây 4 cấp").')
        ws.add_data_validation(dv); dv.add(f'{L}3:{L}{ws.max_row}')
    ws.freeze_panes='D3'; ws.auto_filter.ref=f'A2:{openpyxl.utils.get_column_letter(len(head))}{ws.max_row}'
gc=len({x[5] for x in C}); gd=len({x[5] for x in D})
sheet('C. Nhóm cũ nhiều Loại',f'{len({x[0] for x in C}):,} hàng đang dùng (thuộc {gc} nhóm cũ), {len(C):,} dòng — nhóm cũ ứng với >1 Loại hàng hoá trong file 17/10 ⇒ mỗi hàng lặp 1 dòng cho mỗi Loại của nhóm. Xoá các dòng sai, giữ lại 1 dòng/hàng. Dòng cùng 1 hàng cùng màu nền.'.replace(',','.'),
      ['ID hàng','Mã hàng','Tên hàng','Thương hiệu','Hãng','ID nhóm cũ','Nhóm hàng hoá cũ','Loại hàng hoá'],C,[9,20,55,14,18,9,40,40],8,per_group=True)
sheet('D. Nhóm cũ chưa có Loại',f'{len(D):,} hàng đang dùng (thuộc {gd} nhóm cũ) — nhóm cũ không xuất hiện trong file 17/10. Điền cột vàng "Loại hàng hoá" từ danh sách (toàn bộ cây).'.replace(',','.'),
      ['ID hàng','Mã hàng','Tên hàng','Thương hiệu','Hãng','ID nhóm cũ','Nhóm hàng hoá cũ','Loại hàng hoá'],D,[9,20,55,14,18,9,40,40],8)
ws=o.create_sheet('Cây 4 cấp'); ws.append([f'{len(path)} Loại SP theo sheet NewERP Green Power (17/10). Loại SP chỉ thuộc 1 Nhóm SP ⇒ chọn Loại là đủ, cấp cha tự suy ra.']); ws['A1'].font=Font(italic=True)
ws.append(['Loại SP','Nhóm SP','Nhóm chức năng','Tính chất','Số dòng hàng trong file'])
for c in ws[2]: c.font=H; c.fill=F
for l,p in sorted(path.items(),key=lambda x:(x[1][0],x[1][1],x[1][2],x[1][3])): ws.append([p[3],p[2],p[1],p[0],lcount[l]])
for i,w in enumerate([40,35,30,30,12]): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:E{ws.max_row}'
o.active=0
o.save(out); print('C',len(C),gc,'D',len(D),'gd',gd,'Loai',len(path))
