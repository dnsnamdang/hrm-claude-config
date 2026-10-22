# Rà soát sheet "NewERP Green Power" (file 17/10) — chỉ đọc file + 2 dump DB (groups.tsv, products.tsv)
# python3 -I ra_soat_1710.py <file17> <file16> <groups.tsv> <products.tsv> <out.xlsx>
import openpyxl,sys,collections,re,json
f17,f16,gtsv,ptsv,out=sys.argv[1:6]
norm=lambda s: re.sub(r'\s+',' ',str(s).strip()).casefold() if s is not None and str(s).strip()!='' else None
clean=lambda s: re.sub(r'\s+',' ',str(s).strip()) if s is not None else ''
LV=['Tính chất','Nhóm chức năng','Nhóm sản phẩm','Loại sản phẩm']
def load(f,ci,gi):
    rs=list(openpyxl.load_workbook(f,read_only=True,data_only=True)['NewERP Green Power'].iter_rows(values_only=True))
    out=[]
    for ex,r in enumerate(rs[1:],start=2):
        if not any(r[1:6]): continue
        g=r[gi]
        try: g=int(float(g)) if g not in (None,'') else None
        except: g=('?',g)
        out.append(dict(stt=r[0],ex=ex,lv=[r[1],r[2],r[3],r[4]],name=r[5],code=clean(r[ci]),gid=g))
    return out
A=load(f17,6,8); B=load(f16,7,9)
G={}
for l in open(gtsv,encoding='utf-8'):
    i,st,nm,used,tot=l.rstrip('\n').split('\t'); G[int(i)]=(int(st),nm,int(used) if used not in ('','NULL') else 0,int(tot))
P={}
for l in open(ptsv,encoding='utf-8'):
    c,i,g,st,pt=l.rstrip('\n').split('\t'); P.setdefault(c.strip().casefold(),[]).append((int(i),int(g) if g not in ('NULL','') else None,int(st)))
def multi_parent(rows):
    res={}
    for k in (1,2,3):
        par=collections.defaultdict(collections.Counter); disp={}
        for r in rows:
            c,p=norm(r['lv'][k]),norm(r['lv'][k-1])
            if c and p: par[c][p]+=1; disp[c]=clean(r['lv'][k]); disp[p]=clean(r['lv'][k-1])
        for c,v in par.items():
            if len(v)>1: res[(LV[k],c)]=(disp[c],[(disp[p],n) for p,n in v.most_common()])
    return res
MA,MB=multi_parent(A),multi_parent(B)
# Cây
tree=[set() for _ in range(4)]; paths=set()
for r in A:
    l=[norm(x) for x in r['lv']]
    for k in range(4):
        if l[k]: tree[k].add(l[k])
    if all(l): paths.add(tuple(l))
missing=[r for r in A if not all(norm(x) for x in r['lv'])]
# Tên dài/khoảng trắng thừa
ws_issue=[(r['stt'],r['ex'],LV[k],repr(r['lv'][k])) for r in A for k in range(4) if r['lv'][k] is not None and str(r['lv'][k])!=clean(r['lv'][k])]
# Mã hàng
codes=collections.defaultdict(list)
for r in A:
    if r['code']: codes[r['code'].casefold()].append(r)
dupcode=[(c,len(v),'; '.join(sorted({clean(x['lv'][3]) for x in v}))) for c,v in codes.items() if len(v)>1]
dupconf=[d for d in dupcode if ';' in d[2]]
code_db=[c for c in codes if c in P]
code_used=[c for c in code_db if any(st!=0 for _,_,st in P[c])]
code_notdb=[(v[0]['stt'],v[0]['ex'],v[0]['code'],clean(v[0]['name']),clean(v[0]['lv'][3]),v[0]['gid']) for c,v in codes.items() if c not in P]
nocode=[r for r in A if not r['code']]
# Nhóm cũ -> Loại
g2l=collections.defaultdict(collections.Counter); l2g=collections.defaultdict(collections.Counter); ld={}
badgid=[]; blankgid=0
for r in A:
    l=norm(r['lv'][3]); g=r['gid']
    if g is None: blankgid+=1; continue
    if isinstance(g,tuple): badgid.append((r['stt'],r['ex'],str(g[1]))); continue
    if l: ld[l]=clean(r['lv'][3]); g2l[g][l]+=1; l2g[l][g]+=1
# Lệch nhóm: dòng có mã trên DB nhưng ID nhóm cũ ghi ≠ group_id DB
gmis=[]
for r in A:
    c=r['code'].casefold()
    if c in P and isinstance(r['gid'],int):
        dbg={g for _,g,_ in P[c]}
        if r['gid'] not in dbg: gmis.append((r['stt'],r['ex'],r['code'],r['gid'],', '.join(map(str,sorted(x for x in dbg if x is not None)))))
# Độ phủ hàng đang dùng
inuse=[(c,i,g) for c,v in P.items() for i,g,st in v if st!=0]
cov=collections.Counter(); left_groups=collections.Counter()
for c,i,g in inuse:
    if c in codes: cov['a. theo mã trong file']+=1
    elif g in g2l and len(g2l[g])==1: cov['b. nhóm cũ 1–1 Loại']+=1
    elif g in g2l: cov['c. nhóm cũ nhiều Loại (cần chọn)']+=1; left_groups[g]+=1
    else: cov['d. nhóm cũ chưa có Loại']+=1; left_groups[g]+=1
s2=[(g,G[g][1] if g in G else '(không có DB)',G[g][2] if g in G else '',len(v),'\n'.join(f'{ld[l]} ({n})' for l,n in v.most_common())) for g,v in sorted(g2l.items()) if len(v)>1]
s4=[(g,G[g][1],G[g][2],G[g][3]) for g in sorted(G) if g not in g2l and G[g][2]>0]
s5=[(g,'; '.join(ld[l] for l in v)) for g,v in g2l.items() if g not in G]
# So 16/10
keysA,keysB=set(MA),set(MB)
fixed=sorted(keysB-keysA); still=sorted(keysA&keysB); new=sorted(keysA-keysB)
treeB=[set(norm(r['lv'][k]) for r in B if norm(r['lv'][k])) for k in range(4)]
# ---------- Ghi
from openpyxl.styles import Font,PatternFill,Alignment
o=openpyxl.Workbook(); o.remove(o.active)
H=Font(bold=True,color='FFFFFF'); F=PatternFill('solid',fgColor='1ABC9C')
def sheet(name,note,head,data,widths):
    ws=o.create_sheet(name[:31]); ws.append([note]); ws['A1'].font=Font(italic=True); ws.append(list(head))
    for c in ws[2]: c.font=H; c.fill=F; c.alignment=Alignment(wrap_text=True,vertical='center')
    for d in data: ws.append(list(d))
    for i,w in enumerate(widths): ws.column_dimensions[openpyxl.utils.get_column_letter(i+1)].width=w
    for row in ws.iter_rows(min_row=3):
        for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    ws.freeze_panes='A3'; ws.auto_filter.ref=f'A2:{openpyxl.utils.get_column_letter(len(head))}{max(ws.max_row,2)}'
summ=[('Dòng hàng',len(A),len(B))]+[(f'Số {LV[k]} (theo tên)',len(tree[k]),len(treeB[k])) for k in range(4)]+[
 ('Tên cấp con nằm dưới >1 cha',len(keysA),len(keysB)),('Dòng thiếu cấp',len(missing),''),
 ('Mã hàng khác nhau trong file',len(codes),''),('Mã trùng (≥2 dòng)',len(dupcode),''),('Mã trùng mà khác Loại',len(dupconf),''),
 ('Mã khớp products.code',len(code_db),''),('Mã khớp hàng đang dùng',len(code_used),''),('Dòng không có mã',len(nocode),''),
 ('Dòng trống ID nhóm cũ',blankgid,''),('Nhóm cũ ứng >1 Loại',len(s2),''),('ID nhóm cũ ghi trong file ≠ group_id trên DB',len(gmis),''),
 ('Nhóm DB còn hàng đang dùng nhưng chưa có Loại',len(s4),''),('ID nhóm trong file không có DB',len(s5),''),
 ('— Hàng đang dùng trên DB',len(inuse),'')]+[('   '+k,v,'') for k,v in sorted(cov.items())]
sheet('0. Tổng hợp','File 17/10 so với 16/10. Cây trên DB hrm_erp_giacty hiện RỖNG (product_types 0) ⇒ chưa có gì để so với DB ngoài hàng + nhóm.',['Chỉ số','17/10','16/10'],summ,[50,12,12])
sheet('1a. Cấp con nhiều cha',f'{len(keysA)} tên vi phạm "1 cấp con chỉ thuộc 1 cha". Cột "So 16/10": còn từ 16/10 hay mới phát sinh.',['Cấp','Tên','Số cha','Các cha (số dòng)','So 16/10'],
      [(k[0],MA[k][0],len(MA[k][1]),'\n'.join(f'{p} ({n})' for p,n in MA[k][1]),'còn' if k in keysB else 'MỚI') for k in sorted(keysA)],[16,45,8,70,10])
sheet('1b. Đã sửa so 16-10',f'{len(fixed)} tên từng vi phạm ở file 16/10, nay chỉ còn 1 cha.',['Cấp','Tên','Các cha ở 16/10'],[(k[0],MB[k][0],'\n'.join(f'{p} ({n})' for p,n in MB[k][1])) for k in fixed],[16,45,70])
sheet('1c. Dòng thiếu cấp',f'{len(missing)} dòng trống ≥1 cấp.',['STT','Dòng Excel']+LV+['Tên hàng','Mã'],[(r['stt'],r['ex'],*[clean(x) for x in r['lv']],clean(r['name']),r['code']) for r in missing],[8,8,22,25,30,35,50,22])
sheet('1d. Khoảng trắng thừa',f'{len(ws_issue)} ô tên cấp có khoảng trắng đầu/cuối/kép (nhập sẽ tự cắt).',['STT','Dòng Excel','Cấp','Giá trị gốc'],ws_issue,[8,8,16,60])
sheet('2a. Mã trùng',f'{len(dupcode)} mã xuất hiện ≥2 dòng; {len(dupconf)} mã trong đó gán Loại khác nhau (phải chọn 1).',['Mã','Số dòng','Các Loại'],sorted(dupcode,key=lambda x:(';' not in x[2],x[0])),[25,8,70])
sheet('2b. Mã không có DB',f'{len(code_notdb)} mã trong file không có ở products.code (hàng mới / sai mã).',['STT','Dòng Excel','Mã','Tên hàng','Loại SP','ID nhóm cũ'],code_notdb,[8,8,25,50,35,10])
sheet('2c. Nhóm file ≠ nhóm DB',f'{len(gmis)} dòng: mã có trên DB nhưng "ID nhóm hàng cũ" ghi trong file khác group_id của hàng trên DB.',['STT','Dòng Excel','Mã','ID nhóm (file)','group_id (DB)'],gmis,[8,8,25,12,14])
sheet('3a. Nhóm cũ → nhiều Loại',f'{len(s2)} nhóm cũ ứng >1 Loại SP.',['ID nhóm','Tên nhóm DB','Hàng đang dùng','Số Loại','Các Loại (số dòng)'],s2,[10,40,12,8,70])
sheet('3b. Nhóm DB chưa có Loại',f'{len(s4)} nhóm trên DB còn hàng đang dùng nhưng không xuất hiện trong file.',['ID nhóm','Tên nhóm','Hàng đang dùng','Tổng hàng'],sorted(s4,key=lambda x:-x[2]),[10,50,12,12])
sheet('3c. ID nhóm không có DB',f'{len(s5)} ID nhóm trong file không có ở bảng groups.',['ID nhóm','Loại SP'],s5,[12,70])
o.save(out)
for s in summ: print(*s,sep=' | ')
print('fixed',len(fixed),'still',len(still),'new',len(new))
print('new:',[(k[0],MA[k][0],MA[k][1]) for k in new][:20])
print('ws',len(ws_issue),'badgid',badgid[:5])
