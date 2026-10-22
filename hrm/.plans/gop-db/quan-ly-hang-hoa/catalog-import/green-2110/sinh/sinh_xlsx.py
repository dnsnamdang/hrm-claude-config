import sys,json
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill,Alignment,Border,Side
from openpyxl.utils import get_column_letter
d=json.load(open(sys.argv[1])); out=sys.argv[2]
F=Font(name='Arial',size=10); FB=Font(name='Arial',size=10,bold=True); FH=Font(name='Arial',size=10,bold=True,color='1F3B33')
HF=PatternFill('solid',fgColor='D8E7E1'); NEW=PatternFill('solid',fgColor='FFF2CC'); WARN=PatternFill('solid',fgColor='FCE4D6')
th=Side(style='thin',color='BFBFBF')
wb=Workbook()
def sheet(title,head,rows,widths,mark=None):
    ws=wb.create_sheet(title); ws.append(head)
    for c in ws[1]: c.font=FH; c.fill=HF; c.alignment=Alignment(wrap_text=True,vertical='center'); c.border=Border(bottom=th)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2):
        for c in row: c.font=F; c.alignment=Alignment(vertical='top')
        if mark:
            for col,fill in mark(row): row[col].fill=fill
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w
    ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
    return ws
up=d['upd']
# 2. Danh muc moi
dm=[[n['cap_ten'],n['code'],n['name'],n['pcode'],n['pname'],n['psrc'],n['so_dong'],', '.join(map(str,n['dong'])),', '.join(n['bien_the'])] for n in d['new']]
sheet('Danh mục mới',['Cấp','Mã mới','Tên','Mã cha','Tên cha','Cha là','Số dòng file dùng','Dòng trong file','Cách viết khác trong file'],dm,[15,11,40,11,40,9,10,30,25],
      lambda row:[(5,NEW)] if row[5].value=='MỚI' else [])
# 3. Hang gan loai
def hrow(u): p=u['path']; return [u['code'],u['name'],u['cty'],u['status'],u['model'],u['r'],p[0][1],p[1][1],p[2][1],p[3][0],p[3][1],u['how'],', '.join(u['lech']),u['file'][0],u['file'][1],u['file'][2],u['file'][3]]
H=['Mã hàng','Tên hàng','Công ty','Trạng thái','Model','Dòng file','Tính chất','Nhóm chức năng','Nhóm sản phẩm','Mã Loại SP','Loại sản phẩm','Nguồn Loại','Lệch so với file ở cấp','File: Tính chất','File: Nhóm CN','File: Nhóm SP','File: Loại SP']
W=[22,40,12,15,16,8,26,26,30,11,30,18,24,26,26,30,30]
def mk(row):
    m=[]
    if row[11].value=='Loại mới tạo': m.append((10,NEW))
    if row[12].value: m.append((12,WARN))
    if row[3].value!='Đang dùng': m.append((3,WARN))
    return m
sheet('Hàng gán Loại',H,[hrow(u) for u in up],W,mk)
sheet('Lệch cây so với file',H,[hrow(u) for u in up if u['lech']],W,mk)
sheet('Dòng file bỏ qua',['Dòng file','Tên hàng','Model','Lý do','File: Tính chất','File: Nhóm CN','File: Nhóm SP','File: Loại SP'],
      [[s['r'],s['ten'],s['model'],s['ly_do'],*s['file']] for s in d['skip']],[8,45,16,55,26,26,30,30])
sheet('Mã đã có Loại (giữ)',['Mã hàng','Tên hàng','Công ty','Trạng thái','Model','Dòng file','Loại SP hiện tại'],
      [[k['code'],k['name'],k['cty'],k['status'],k['model'],k['r'],k['loai_hien_tai']] for k in d['keep']],[22,45,12,15,16,8,35])
# 1. Tong quan
ws=wb.worksheets[0]; ws.title='Tổng quan'
L=[('Cập nhật danh mục 4 cấp cho hàng ETEK GREEN — 21/10/2026',FB),
   ('Nguồn: "Copy of UPDATE hàng green.xlsx" (sheet Thành phẩm, 855 dòng) · đối chiếu trên bản sao prod hrm_erp_21_10 (dump 21/10 20:25) · SQL: danh_muc_green_2110.sql',F),
   ('',F),('LUẬT ĐÃ CHỐT',FB),
   ('1. So tên coi "," và "-" là một (không phân biệt hoa thường, dấu thanh); danh mục mới đặt tên kiểu gạch ngang theo quy ước DB.',F),
   ('2. Loại SP đã có trên DB ⇒ lấy luôn 4 cấp đang có trên DB, bỏ qua cha ghi trong file. Nhóm SP / Nhóm CN đã có ⇒ dùng cha trên DB.',F),
   ('3. Bỏ qua: model chưa có trên DB · Loại SP = "0" · danh mục MỚI mà file đặt dưới nhiều cha khác nhau.',F),
   ('4. Mã hàng đã có Loại SP ⇒ giữ nguyên. Gán cả mã hàng đã xoá (status 0).',F),
   ('5. Loại SP mới: tồn lâu 180 ngày, VAT trống, mẫu tem trống, không bắt buộc serial, không thuộc tính, Hoạt động; người tạo DNS Admin.',F),
   ('6. Danh mục mới ghi lịch sử "Tạo mới"; gán Loại cho hàng KHÔNG ghi lịch sử hàng hoá (như go-live 20/10).',F),
   ('',F),('SỐ LIỆU',FB)]
for t,f in L:
    ws.append([t]); ws.cell(ws.max_row,1).font=f
stats=[('Dòng trong file',855,None),
 ('Dòng bị bỏ qua',"=COUNTA('Dòng file bỏ qua'!A:A)-1",None),
 ('   · Model chưa có trên DB',"=COUNTIF('Dòng file bỏ qua'!D:D,\"Model chưa có*\")",None),
 ('   · Loại SP = \"0\"',"=COUNTIF('Dòng file bỏ qua'!D:D,\"Loại SP trong file*\")",None),
 ('   · Danh mục mới nằm dưới nhiều cha',"=COUNTIF('Dòng file bỏ qua'!D:D,\"*nhiều cha*\")",None),
 ('Nhóm sản phẩm tạo mới',"=COUNTIF('Danh mục mới'!A:A,\"Nhóm sản phẩm\")",None),
 ('Loại sản phẩm tạo mới',"=COUNTIF('Danh mục mới'!A:A,\"Loại sản phẩm\")",None),
 ('Mã hàng được gán Loại',"=COUNTA('Hàng gán Loại'!A:A)-1",None),
 ('   · dùng Loại có sẵn trên DB',"=COUNTIF('Hàng gán Loại'!L:L,\"Loại có sẵn trên DB\")",None),
 ('   · dùng Loại mới tạo',"=COUNTIF('Hàng gán Loại'!L:L,\"Loại mới tạo\")",None),
 ('   · trong đó mã đã xoá (status 0)',"=COUNTIF('Hàng gán Loại'!D:D,\"Đã xoá*\")",None),
 ('   · cây gán lệch so với file (nghiệp vụ rà)',"=COUNTA('Lệch cây so với file'!A:A)-1",None),
 ('Mã hàng đã có Loại — giữ nguyên',"=COUNTA('Mã đã có Loại (giữ)'!A:A)-1",None)]
for t,v,_ in stats:
    ws.append([t,v]); ws.cell(ws.max_row,1).font=F; ws.cell(ws.max_row,2).font=FB
ws.append([]); ws.append(['Màu: vàng = danh mục MỚI tạo · cam = lệch so với file / mã đã xoá — cần nghiệp vụ rà.']); ws.cell(ws.max_row,1).font=F
ws.column_dimensions['A'].width=60; ws.column_dimensions['B'].width=12
wb.calculation.fullCalcOnLoad=True
wb.save(out); print('ok')
