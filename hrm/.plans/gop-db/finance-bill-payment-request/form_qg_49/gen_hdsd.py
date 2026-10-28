# -*- coding: utf-8 -*-
"""HDSD bản GỌN — Phiếu đề nghị thanh toán (form chuẩn Danh mục quốc gia, 28/10/2026).

Chạy: python gen_hdsd.py  → out/HDSD_Phiếu đề nghị thanh toán.docx (+ .pdf, out/render/ để soát)
Khuôn: màn 27 / màn 48 (finance-bill-income-request/form_qg_48/gen_hdsd.py); phần chung ở
../../phieu-theo-form-quoc-gia/hdsd_phieu_lib.py. Thuật ngữ + dòng phiên bản cũ đọc từ bản Drive.
Icon riêng chụp ở hdsd_shots/icons (capture.js); icon trùng nhãn dùng lại của màn 27.
"""
import os
import sys

import docx

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, '..', '..', 'phieu-theo-form-quoc-gia'))
from hdsd_phieu_lib import PhieuHdsd, finalize_word, toc_lines, render_check  # noqa: E402

TEN = 'Phiếu đề nghị thanh toán'
SHOTS = os.path.join(HERE, 'hdsd_shots')
OUT = os.path.join(HERE, 'out', 'HDSD_%s.docx' % TEN)
DRIVE = os.path.join(HERE, 'drive_0710', 'HDSD_Drive.docx')
ICON_27 = os.path.join(HERE, '..', '..', 'finance-product-import-direct-transfer', 'form_qg', 'hdsd_shots', 'icons')


def drive_tables():
    d = docx.Document(DRIVE)
    t0, t1 = d.tables[0], d.tables[1]
    thuat_ngu = [[r.cells[0].text.strip(), r.cells[1].text.strip()] for r in t0.rows[1:]]
    cu = [[r.cells[0].text.strip(), r.cells[1].text.strip(), r.cells[2].text.strip()] for r in t1.rows[1:]]
    return thuat_ngu, cu


QUYEN = [
    ['Xem danh sách, xem chi tiết, in, xuất Excel, xem lịch sử',
     'Một trong các quyền: Xem tất cả phiếu đề nghị thanh toán của tổng công ty / của công ty / của phòng ban / của '
     'bộ phận. Phiếu do chính mình tạo luôn xem được.'],
    ['Lập phiếu', 'Không cần quyền riêng.'],
    ['Sửa, xóa phiếu', 'Người tạo phiếu, phiếu đang ở trạng thái Đang tạo hoặc Không duyệt.'],
    ['Duyệt, từ chối phiếu',
     'Quyền duyệt của đúng cấp đang giữ phiếu: Trưởng phòng / Kế toán công nợ / Kế toán trưởng / Ban giám đốc duyệt '
     'đề nghị thanh toán; cùng công ty với phiếu.'],
    ['Tạo phiếu chi, tạo ủy nhiệm chi', 'Quyền Kế toán thanh toán; phiếu đang ở trạng thái Chờ tạo phiếu chi.'],
]


def keep_table_together(w, first_header):
    """keepNext cho mọi đoạn của bảng (trừ dòng cuối) → dòng tiêu đề không bị bỏ lại cuối trang trước."""
    from docx.oxml.ns import qn
    for tbl in w.doc.element.body.iter(qn('w:tbl')):
        trs = tbl.findall(qn('w:tr'))
        if not trs or ''.join(t.text or '' for t in trs[0].iter(qn('w:t'))).strip()[:len(first_header)] != first_header:
            continue
        for tr in trs[:-1]:
            for p in tr.iter(qn('w:p')):
                ppr = p.find(qn('w:pPr'))
                if ppr is None:
                    ppr = p.makeelement(qn('w:pPr'), {})
                    p.insert(0, ppr)
                if ppr.find(qn('w:keepNext')) is None:
                    ppr.append(ppr.makeelement(qn('w:keepNext'), {}))


def build():
    thuat_ngu, cu = drive_tables()
    phien_ban = cu + [['1.1', '28/10/2026', 'Chuyển sang form chuẩn']]
    d = PhieuHdsd(TEN, SHOTS, icon_dirs=[os.path.join(SHOTS, 'icons'), ICON_27])
    d.tong_quan(thuat_ngu, phien_ban, QUYEN, os.path.join(HERE, 'uml', 'overview_srs.png'))
    keep_table_together(d.w, 'Chức năng')

    # PHẦN 1
    d.truy_cap(
        menu=[('Tài chính', 'menu_phanhe'), ('Khởi tạo', 'menu_khoitao'),
              ('Khởi tạo phiếu yêu cầu - Công nợ - Thu - Chi', 'menu_nhom'), ('Đề nghị thanh toán', 'menu_muc')],
        o_nhanh='mã phiếu',
        bo_loc=[('Công ty', 'o_congty'), ('Phòng ban', 'o_phongban'), ('Loại chi', 'o_loaichi'),
                ('Hình thức thanh toán', 'o_hinhthuc'), ('Trạng thái', 'o_trangthai'), ('Lý do chi', 'o_lydochi'),
                ('Khách hàng', 'o_khachhang'), ('Nhà cung cấp', 'o_ncc'), ('Người lập', 'o_nguoilap'),
                ('Số tiền đề nghị', 'o_sotien'), ('Ngày lập', 'o_ngaylap')])

    # PHẦN 2 — Lập phiếu
    d.phan('LẬP PHIẾU ĐỀ NGHỊ THANH TOÁN')
    d.step(d.vao)
    d.step(['Bước 2: Bấm ', '@btn_taomoi'])
    d.img('10_create', 'Màn hình Thêm phiếu đề nghị thanh toán')
    d.step('Bước 3: Nhập thông tin chung:')
    d.sub('Loại chi (bắt buộc): Chi trả nhà cung cấp (mặc định), Chi trả lại khách hàng, Chi thưởng thực hiện hợp '
          'đồng hoặc Thanh toán chi phí vận chuyển NCC.')
    d.sub('Hình thức thanh toán (bắt buộc khi gửi duyệt): TM (mặc định) hoặc CK.')
    d.sub('Loại tiền (bắt buộc khi gửi duyệt): mặc định VNĐ; chọn ngoại tệ thì Tỷ giá tự điền và sửa được.')
    d.sub('Tỷ giá (VND) (bắt buộc khi gửi duyệt): số lớn hơn 0; loại tiền VNĐ thì cố định bằng 1.')
    d.sub('Đến ngày (bắt buộc, chỉ với loại Thanh toán chi phí vận chuyển NCC): chọn ngày dd/mm/yyyy.')
    d.sub('Nhà cung cấp / Khách hàng (bắt buộc khi hình thức CK hoặc loại vận chuyển): bấm vào ô rồi chọn trong cửa sổ.')
    d.sub('Lý do chi (bắt buộc khi gửi duyệt, tùy chọn khi lưu nháp): nhập chữ.')
    d.step('Bước 4: Với hình thức CK, khối Thông tin ngân hàng tự điền theo hồ sơ của nhà cung cấp / khách hàng đã '
           'chọn (chỉ xem); có nhiều tài khoản thì chọn ở ô Tài khoản ngân hàng.')
    d.img('13_create_ck', 'Hình thức chuyển khoản: ô Nhà cung cấp và khối Thông tin ngân hàng')
    d.step(['Bước 5: Ở khối Chi tiết, bấm ', '@btn_themdong', ' để thêm một dòng, rồi nhập từng dòng:'])
    d.sub('Nhà cung cấp / Khách hàng của dòng (bắt buộc khi hình thức TM): bấm vào ô, tìm rồi bấm vào dòng cần chọn.')
    d.sub('Hợp đồng (bắt buộc): bấm vào ô, bấm vào dòng hợp đồng cần chi; Số tiền còn nợ tự điền. Hợp đồng đã có ở '
          'dòng khác không chọn lại được.')
    d.sub('Số tiền đề nghị chi (bắt buộc khi gửi duyệt): số từ 1 trở lên; dấu phẩy ngăn hàng nghìn, ví dụ 1,200,000.')
    d.sub('Ghi chú của dòng (tùy chọn).')
    d.sub(['Hoặc bấm ', '@btn_xoadong', ' để bỏ dòng khỏi phiếu.'])
    d.img('11_supplier', 'Cửa sổ Chọn nhà cung cấp')
    d.step(['Bước 6: Riêng loại Thanh toán chi phí vận chuyển NCC: chọn Nhà cung cấp và Đến ngày, bấm ',
            '@btn_laydulieu', ', tích các chuyến xe cần thanh toán và nhập số tiền cho các dòng đã tích.'])
    d.step(['Bước 7: Ở khối File đính kèm, bấm ', '@btn_themtailieu', ', bấm ', '@btn_chontep',
            ' và chọn tệp (pdf, png, jpg, jpeg, doc, docx, xls, xlsx, zip; tối đa 20MB). Bắt buộc ít nhất 1 tệp với '
            'loại Chi trả nhà cung cấp khi gửi duyệt.'])
    d.step(['Bước 8: Bấm ', '@btn_luuguiduyet', ', chọn ', '@btn_confirm_xacnhan', ' để gửi Trưởng phòng duyệt,'])
    d.sub(['Hoặc ', '@btn_luunhap', ' để lưu nháp (chỉ cần chọn Loại chi).'])
    d.sub(['Hoặc ', '@btn_luutieptuc', ' để lưu nháp rồi lập tiếp phiếu khác.'])
    d.sub(['Hoặc ', '@btn_quaylai', ' để không lưu thông tin.'])

    # PHẦN 3 — Chỉnh sửa
    d.phan('CHỈNH SỬA')
    d.step(d.vao)
    d.step(['Bước 2: Bấm nút ', '@btn_sua_row', ' ở cột Hành động của phiếu Đang tạo hoặc Không duyệt do mình tạo.'])
    d.sub(['Hoặc bấm vào Mã phiếu để mở màn chi tiết, chọn ', '@btn_sua'])
    d.step('Bước 3: Hệ thống hiển thị màn hình “Sửa phiếu đề nghị thanh toán”.')
    d.img('16_edit', 'Màn hình Sửa phiếu đề nghị thanh toán')
    d.step('Bước 4: Nhập thông tin cần sửa (quy tắc nhập như Lập phiếu ở PHẦN 2).')
    d.step(['Bước 5: Bấm ', '@btn_luuguiduyet', ' để lưu và gửi duyệt lại từ cấp Trưởng phòng,'])
    d.sub(['Hoặc ', '@btn_luunhap', ' để lưu nháp.'])
    d.sub(['Hoặc ', '@btn_quaylai', ' nếu muốn hủy.'])

    # PHẦN 4 — Xóa (Xóa nằm trong nút ba chấm khi dòng có trên 3 nút)
    d.phan('XÓA')
    d.step(['Bước 1: Ở cột Hành động của phiếu Đang tạo hoặc Không duyệt do mình tạo, bấm ', '@btn_bacham_row',
            ' rồi chọn ', '@item_xoa'])
    d.sub(['Hoặc bấm vào Mã phiếu để mở màn chi tiết, chọn ', '@btn_xoa'])
    d.step('Bước 2: Hệ thống hiện hộp xác nhận “Xác nhận xóa”')
    d.img('15_delete', 'Hộp xác nhận xóa phiếu')
    d.step(['Bước 3: Bấm ', '@btn_confirm_xoa', ' để xác nhận xóa'])
    d.sub(['Hoặc bấm ', '@btn_huy_confirm', ' nếu bấm nhầm'])

    # PHẦN 5 — Xem chi tiết
    d.phan('XEM CHI TIẾT PHIẾU')
    d.step(d.vao)
    d.step('Bước 2: Bấm vào Mã phiếu. Hệ thống mở màn “Chi tiết phiếu đề nghị thanh toán”.')
    d.img('17b_detail_view', 'Màn hình Chi tiết phiếu đề nghị thanh toán')
    d.step(['Bước 3: Bấm ', '@btn_quaylai', ' để về màn danh sách.'])

    # PHẦN 6 — Duyệt theo cấp
    d.phan('DUYỆT PHIẾU')
    d.step(['Bước 1: Từ phân hệ Tài chính ', '@menu_phanhe', ', tại menu Chờ duyệt ', '@menu_choduyet', ', nhóm '
            'Phê duyệt - Công nợ - Thu - Chi ', '@menu_nhom_cd', ', chọn Phiếu đề nghị thanh toán chờ duyệt ',
            '@menu_muc_cd'])
    d.img('03_pending', 'Màn hình Phiếu đề nghị thanh toán chờ duyệt')
    d.step(['Bước 2: Bấm ', '@btn_duyet_row', ' ở cột Hành động (hoặc bấm vào Mã phiếu) để mở màn chi tiết.'])
    d.step('Bước 3: Kiểm tra thông tin, file đính kèm; nhập số tiền chấp thuận của từng dòng ở cột tiền của cấp mình.')
    d.step(['Bước 4: Bấm ', '@btn_duyet', '. Hệ thống hiện hộp xác nhận “Xác nhận duyệt”.'])
    d.sub(['Hoặc (cấp Kế toán trưởng) bấm ', '@btn_chuyenduyet', ' để chuyển phiếu lên Ban giám đốc duyệt.'])
    d.img('18_approve', 'Hộp xác nhận duyệt phiếu')
    d.step(['Bước 5: Bấm ', '@btn_confirm_duyet', ' (hoặc ', '@btn_confirm_chuyenduyet', ') để xác nhận'])
    d.sub(['Hoặc bấm ', '@btn_huy_confirm', ' nếu bấm nhầm'])

    # PHẦN 7 — Từ chối
    d.phan('TỪ CHỐI PHIẾU')
    d.step('Bước 1: Mở màn chi tiết của phiếu đang chờ cấp mình duyệt (như Bước 1, 2 PHẦN 6).')
    d.step(['Bước 2: Bấm ', '@btn_tuchoi', '. Hệ thống mở cửa sổ “Từ chối phiếu”.'])
    d.img('19_reject', 'Cửa sổ Từ chối phiếu')
    d.step('Bước 3: Nhập Ghi chú của cấp mình (bắt buộc) và Lý do không duyệt (tùy chọn, hiện trên thông báo gửi '
           'người tạo).')
    d.step(['Bước 4: Bấm ', '@btn_confirm_tuchoi', ' để xác nhận. Cấp Trưởng phòng từ chối thì phiếu về Đang tạo; '
            'các cấp sau từ chối thì phiếu sang Không duyệt.'])
    d.sub(['Hoặc bấm ', '@btn_dong_reject', ' nếu muốn hủy.'])

    # PHẦN 8 — Tạo phiếu chi / ủy nhiệm chi
    d.phan('TẠO PHIẾU CHI / ỦY NHIỆM CHI')
    d.step('Bước 1: Mở màn chi tiết của phiếu ở trạng thái Chờ tạo phiếu chi (từ màn chờ duyệt ở PHẦN 6).')
    d.img('23_waitpay', 'Màn hình chi tiết phiếu Chờ tạo phiếu chi')
    d.step(['Bước 2: Phiếu hình thức TM: bấm ', '@btn_taophieuchi', '. Hệ thống mở màn lập Phiếu chi, gắn sẵn phiếu '
            'đề nghị này.'])
    d.sub(['Hoặc phiếu hình thức CK: bấm ', '@btn_taounc', ' để mở màn lập Phiếu ủy nhiệm chi.'])
    d.sub(['Hoặc bấm ', '@btn_taounc_row', ' ở cột Hành động của phiếu trên màn chờ duyệt.'])
    d.step('Bước 3: Lập và gửi duyệt chứng từ chi ở màn mới mở.')

    # PHẦN 9 — In
    d.phan('IN PHIẾU')
    d.step(['Bước 1: Ở cột Hành động, bấm ', '@btn_in_row'])
    d.sub(['Hoặc tại màn chi tiết phiếu, bấm ', '@btn_in'])
    d.step('Bước 2: Hệ thống mở cửa sổ xem trước bản in.')
    d.img('21_print', 'Cửa sổ Xem trước phiếu đề nghị thanh toán')
    d.step(['Bước 3: Bấm ', '@btn_in_preview', ' để in.'])

    # PHẦN 10 — Xuất Excel (1 phiếu, tải thẳng tệp — không có cửa sổ chọn trường)
    d.phan('XUẤT EXCEL PHIẾU')
    d.step(['Bước 1: Ở cột Hành động, bấm ', '@btn_xuatexcel_row'])
    d.sub(['Hoặc bấm ', '@btn_bacham_row', ' rồi chọn ', '@item_xuatexcel'])
    d.sub(['Hoặc tại màn chi tiết phiếu, bấm ', '@btn_xuatexcel_ct'])
    d.img('22_export', 'Mục Xuất Excel trong nút ba chấm của dòng phiếu')
    d.step('Bước 2: Hệ thống tải về tệp De-nghi-thanh-toan-<mã phiếu>.xlsx và báo “Xuất Excel thành công”.')
    d.step('Bước 3: Mở file Excel đã tải.')

    # PHẦN 11..13
    d.lich_su('Mã phiếu', img='20_history_row', cap='Cửa sổ Lịch sử thay đổi của phiếu', bo_loc=False)
    d.cai_dat_bo_loc()
    d.tuy_chinh_cot()
    return d.save(OUT)


if __name__ == '__main__':
    out = build()
    pdf = os.path.splitext(out)[0] + '.pdf'
    size = finalize_word(out, pdf)
    print('docx', out, size)
    print('TOC:')
    for line in toc_lines(out):
        print('  ', line)
    n, warn, sheets = render_check(pdf, os.path.join(HERE, 'out', 'render'))
    print('pages', n)
    for w_ in warn:
        print('WARN', w_)
    for s in sheets:
        print(s)
