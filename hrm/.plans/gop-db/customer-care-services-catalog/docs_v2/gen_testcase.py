# -*- coding: utf-8 -*-
"""Tab testcase "12.Danh mục gói bảo dưỡng" — dựng trên KHUÔN tab "13. DM quốc gia" của workbook
"Testcase chuyển đổi.xlsx" (khối mô tả dòng 1-10, TEST SUMMARY dòng 11-16, tiêu đề dòng 17,
nhóm chức năng tô nền, dropdown DNS/TP check, cột R kết quả cuối).

Chạy:  python3 gen_testcase.py "<Testcase chuyển đổi.xlsx>"
Kết quả: out/testcase - Danh mục gói bảo dưỡng.xlsx (1 tab) — dùng để duyệt trước khi chèn vào
workbook chung.
"""
import os
import sys
from copy import copy

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
TAB = '12.Danh mục gói bảo dưỡng'
MAU = '13. DM quốc gia'
MODULE = 'DM gói bảo dưỡng'
NGAY = '25/09/2026'

MOTA = [
    ('1. Mục đích tính năng',
     'Quản lý danh mục gói bảo dưỡng của phân hệ CSKH sau bán: mỗi gói gồm thông tin chung, bảng nội dung kiểm tra bảo dưỡng theo từng cấp, giá bán theo từng công ty, hàng hóa áp dụng và file PDF. Gói là nguồn dữ liệu cho báo giá dịch vụ và phiếu yêu cầu dịch vụ.\nNgười dùng làm được: xem danh sách, tìm kiếm, lọc, tạo mới, sửa, xem chi tiết, nhân bản, xóa, khóa/mở khóa, import Excel, xuất Excel, in phiếu, tùy chỉnh cột và xem lịch sử.'),
    ('2. Đối tượng được tính / hiển thị',
     '- Toàn bộ gói bảo dưỡng của hệ thống, gồm CẢ gói Hoạt động lẫn gói đã Khóa, không phân theo công ty/phòng ban.\n- Bộ lọc Trạng thái để trống thì hiện cả hai nhóm; đây là mặc định khi vào màn.\n- Gói đã Khóa vẫn xem chi tiết, in, nhân bản và xem lịch sử được.'),
    ('3. Đối tượng bị ẩn / không tính',
     '- Nút Tạo mới, Import Excel, Nhân bản: ẩn khi không có quyền Thêm.\n- Nút Sửa: ẩn khi không có quyền Sửa hoặc gói đang Khóa.\n- Nút Xóa: ẩn khi không có quyền Xóa, gói đang Khóa, hoặc gói ĐÃ ĐƯỢC SỬ DỤNG (đã gắn hàng hóa / đã có trong báo giá dịch vụ).\n- Nút Mở khóa: chỉ hiện với gói Khóa và tài khoản có quyền Sửa.\n- Gói đã Khóa không còn chọn được ở màn nghiệp vụ khác.'),
    ('4. Bộ lọc thời gian áp dụng cho',
     'Màn hình KHÔNG có bộ lọc khoảng thời gian. Yếu tố thời gian chỉ có ở cột Ngày tạo, Ngày cập nhật (sắp xếp được) và bộ lọc Từ ngày – Đến ngày trong cửa sổ Lịch sử thay đổi.'),
    ('5. Cấu trúc dữ liệu / cây phân cấp',
     'Một gói gồm 5 khối:\n1. Thông tin chung: tên, mã, định mức đàm phán giá, VAT, công ty quản lý, trạng thái, ghi chú, hệ số giá bán.\n2. Danh mục kiểm tra bảo dưỡng định kỳ dạng MA TRẬN: mỗi DÒNG là một nội dung kiểm tra (ĐVT, SL), mỗi CỘT là một cấp bảo dưỡng; ô giao chọn ghi chú kiểm tra. Dưới ma trận: Định mức công, Hệ số công nghệ, Giá vốn, Giá công thức, Giá bán cơ sở, Gợi ý hàng hóa, Giá bán theo công ty.\n3. Giá vốn theo công ty: hệ số giá bán riêng từng công ty.\n4. Áp dụng cho hàng hóa (gom theo nhóm hàng).\n5. File đính kèm PDF (bắt buộc ≥ 1).\nDanh mục phụ thuộc: Cấp dịch vụ bảo dưỡng, Ghi chú kiểm tra bảo dưỡng, Đơn vị tính, Công ty, Hàng hóa, Nhóm hàng.'),
    ('6. Quy tắc cộng dồn / deduplicate',
     '- Tên gói và Mã gói DUY NHẤT toàn hệ thống; mã lưu CHỮ IN HOA, chỉ gồm chữ không dấu, số, dấu - và _.\n- Một cấp bảo dưỡng chỉ chọn một lần trong một gói.\n- Khi sửa, gói đang sửa được loại khỏi phép kiểm tra trùng.\n- Ô tìm nhanh quét đồng thời Tên và Mã; bản ghi khớp cả hai vẫn chỉ hiện một dòng.'),
    ('7. Phân quyền cấp',
     'Màn hình dùng 3 quyền thuộc nhóm “Danh mục dịch vụ bảo dưỡng”:\n- “Thêm danh mục gói bảo dưỡng” → Tạo mới, Import Excel, Nhân bản.\n- “Sửa danh mục gói bảo dưỡng” → Sửa, Mở khóa (Khóa qua ô Trạng thái của form Sửa).\n- “Xóa danh mục gói bảo dưỡng” → Xóa.\nXem danh sách, xem chi tiết, In, Xuất Excel, Lịch sử KHÔNG đòi quyền. Không phân quyền theo cấp công ty/phòng ban.'),
    ('8. Cách tính các ô thống kê',
     '- “Hiển thị a–b / N”: N là tổng số gói khớp bộ lọc.\n- Số dòng/trang 5 / 10 / 20 / 50 / 100, mặc định 10; đổi số dòng thì về trang 1.\n- Giá vốn = Đơn giá công của công ty quản lý × Định mức công × Hệ số công nghệ; Giá công thức = Giá vốn × Hệ số giá bán gói; Giá bán cơ sở mặc định = Giá công thức; Giá bán theo công ty = Giá bán cơ sở × Hệ số của công ty.\n- Số hiển thị chuẩn quốc tế: 1,234,567.89.'),
    ('9. Ghi chú đọc bảng',
     'Các bẫy dễ sai nhất, QA đọc trước khi test:\n- Màn Thêm mới / Sửa là TRANG RIÊNG, không phải popup.\n- Không có nút Khóa riêng: khóa bằng cách đổi Trạng thái = Khóa trong trang Sửa rồi Lưu. Mở khóa có nút riêng.\n- Gói đã gắn hàng hóa hoặc đã có trong báo giá là “đã được sử dụng” → KHÔNG có nút Xóa (ẩn hẳn).\n- Xuất Excel lấy TOÀN BỘ danh mục, không áp bộ lọc trên màn hình — đúng thiết kế.\n- Gói tạo bằng Import chưa có file PDF → muốn Sửa phải đính kèm PDF trước.\n- Mã gói tự chuyển chữ in hoa ngay khi gõ.'),
]

# (nhóm, [ (Chức năng, Priority, Tiền điều kiện, Bước, Test data, Expected), ... ])
G = []


def grp(name, rows):
    G.append((name, rows))


grp('PHÂN QUYỀN & TRUY CẬP', [
    ('Tài khoản KHÔNG có quyền nào vẫn xem được danh sách', 'P0', 'Tài khoản đã đăng nhập, không có 3 quyền của màn.', '1. Đăng nhập\n2. Vào menu CSKH sau bán → Danh mục → Gói bảo dưỡng', '', '- Vào được màn hình, danh sách hiện đủ gói\n- KHÔNG có nút Tạo mới, Import Excel\n- Có nút Xuất Excel và biểu tượng Cấu hình cột\n- Trên dòng chỉ có In và Lịch sử'),
    ('Tài khoản không quyền vẫn xem chi tiết, in, xuất Excel, xem lịch sử', 'P1', 'Tài khoản như trên.', '1. Bấm mã một gói\n2. Bấm In, Xem lịch sử ở chân trang\n3. Quay lại, bấm Xuất Excel → Xuất file', '', '- Mở được chi tiết, xem trước phiếu in, lịch sử\n- Tải được file Excel\n- Chân trang chi tiết KHÔNG có Sửa, Xóa, Nhân bản, Mở khóa'),
    ('Chỉ có quyền Thêm', 'P0', 'Tài khoản CHỈ có quyền “Thêm danh mục gói bảo dưỡng”.', '1. Vào màn danh sách\n2. Quan sát thanh công cụ và cột Hành động', '', '- Có Tạo mới, Import Excel\n- Dòng có Nhân bản, In, Lịch sử\n- KHÔNG có Sửa, Xóa, Mở khóa'),
    ('Chỉ có quyền Sửa', 'P0', 'Tài khoản CHỈ có quyền “Sửa danh mục gói bảo dưỡng”; có gói Hoạt động và gói Khóa.', '1. Vào màn danh sách\n2. Quan sát cột Hành động của gói Hoạt động và gói Khóa', '', '- Gói Hoạt động: có Sửa\n- Gói Khóa: có Mở khóa\n- KHÔNG có Tạo mới, Import Excel, Nhân bản, Xóa'),
    ('Chỉ có quyền Xóa', 'P0', 'Tài khoản CHỈ có quyền “Xóa danh mục gói bảo dưỡng”; có gói chưa được sử dụng.', '1. Vào màn danh sách\n2. Quan sát cột Hành động', '', '- Có Xóa ở gói Hoạt động chưa được sử dụng\n- KHÔNG có Tạo mới, Sửa, Mở khóa, Nhân bản'),
    ('Vào thẳng đường dẫn Thêm mới khi không có quyền Thêm', 'P1', 'Tài khoản không có quyền Thêm.', '1. Gõ đường dẫn /customer-care/services/create\n2. Nhập đủ dữ liệu, bấm Lưu', '', '- Hệ thống từ chối, báo không có quyền\n- Không phát sinh gói mới'),
    ('Chưa đăng nhập thì không vào được màn hình', 'P0', 'Trình duyệt chưa đăng nhập.', '1. Gõ đường dẫn /customer-care/services', '', '- Chuyển sang màn Đăng nhập, không hiển thị dữ liệu'),
])

grp('HIỂN THỊ TRANG & TRUY CẬP', [
    ('Mở màn hình từ menu', 'P0', 'Đã đăng nhập, danh mục có dữ liệu.', '1. Chọn phân hệ CSKH sau bán\n2. Menu Danh mục → Gói bảo dưỡng', '', '- Tiêu đề “Danh mục gói bảo dưỡng”\n- Khối Bộ lọc danh sách ở trên, bảng ở dưới\n- Dòng “Hiển thị 1–10 / N”'),
    ('Kiểm tra đủ các cột mặc định', 'P0', 'Tài khoản chưa chỉnh cấu hình cột.', '1. Đọc tiêu đề các cột', '', '- Hiện: STT, Mã, Tên gói bảo dưỡng, Người tạo, Ngày tạo, Trạng thái, Hành động\n- KHÔNG hiện: Công ty quản lý gói bảo dưỡng, Người cập nhật, Ngày cập nhật'),
    ('Thanh công cụ', 'P0', 'Tài khoản đủ 3 quyền.', '1. Quan sát góc phải trên bảng', '', '- Có Tạo mới, Xuất Excel, Import Excel và biểu tượng Cấu hình cột hiển thị'),
    ('Bấm mã gói mở trang chi tiết', 'P0', 'Có gói DOC-GBD-03.', '1. Bấm mã DOC-GBD-03', 'Gói: DOC-GBD-03', '- Mở trang “Chi tiết gói bảo dưỡng: DOC-GBD-03”\n- Mở được bằng tab mới (chuột phải)'),
    ('Biểu tượng ⓘ hiển thị giá theo cấp', 'P1', 'Gói đã khai cấp và định mức công.', '1. Rê chuột vào biểu tượng ⓘ cạnh tên gói', '', '- Hiện danh sách “Cấp: giá bán”, số có dấu phẩy ngăn nghìn (VD 1,680,000)'),
    ('Gói chưa khai cấp không có biểu tượng ⓘ', 'P2', 'Có gói chưa có cấp.', '1. Quan sát cột Tên gói', '', '- Chỉ hiện tên gói, không có ⓘ'),
    ('Nhãn trạng thái', 'P0', 'Có gói Hoạt động và gói Khóa.', '1. Quan sát cột Trạng thái', '', '- Hoạt động: nhãn xanh\n- Khóa: nhãn đỏ'),
    ('Nút hành động của gói Hoạt động CHƯA được sử dụng', 'P0', 'Đủ quyền; gói chưa gắn hàng hóa, chưa có trong báo giá.', '1. Quan sát cột Hành động\n2. Mở nút ba chấm', '', '- Hiện thẳng: Sửa, Xóa\n- Ba chấm: Nhân bản, In, Lịch sử'),
    ('Nút hành động của gói Hoạt động ĐÃ được sử dụng', 'P0', 'Đủ quyền; gói đã gắn hàng hóa.', '1. Quan sát cột Hành động\n2. Mở nút ba chấm', '', '- Hiện thẳng: Sửa, Nhân bản\n- Ba chấm: In, Lịch sử\n- KHÔNG có Xóa'),
    ('Nút hành động của gói đang Khóa', 'P0', 'Đủ quyền; gói đang Khóa.', '1. Quan sát cột Hành động\n2. Mở nút ba chấm', '', '- Hiện thẳng: Mở khóa, Nhân bản\n- Ba chấm: In, Lịch sử\n- KHÔNG có Sửa, Xóa'),
    ('Hiển thị khi không có dữ liệu khớp', 'P1', 'Danh mục có dữ liệu.', '1. Gõ vào ô tìm nhanh chuỗi không tồn tại', 'Từ khóa: zzzkhongtontai999', '- Bảng hiện thông báo không có dữ liệu\n- Tổng = 0'),
])

grp('DANH SÁCH, SẮP XẾP & PHÂN TRANG', [
    ('Mặc định xếp gói mới nhất lên đầu', 'P1', 'Có gói tạo ở nhiều thời điểm.', '1. Vào màn hình, đọc cột Ngày tạo', '', '- Ngày tạo giảm dần'),
    ('Số thứ tự liên tục qua các trang', 'P0', 'Có > 1 trang.', '1. Ghi STT dòng cuối trang 1\n2. Sang trang 2', '', '- STT nối tiếp, không quay về 1'),
    ('Đổi số dòng mỗi trang', 'P1', 'Đang ở trang 2.', '1. Đổi Số dòng/trang sang 20', '', '- Có các mức 5/10/20/50/100\n- Bảng hiện 20 dòng, quay về trang 1'),
    ('Sắp xếp theo Mã', 'P1', 'Có ≥ 5 gói.', '1. Bấm tiêu đề cột Mã\n2. Bấm lần 2', '', '- Lần 1 tăng dần, lần 2 giảm dần'),
    ('Sắp xếp theo Tên gói', 'P1', 'Có ≥ 5 gói.', '1. Bấm tiêu đề cột Tên gói bảo dưỡng hai lần', '', '- Đảo chiều đúng thứ tự chữ cái tiếng Việt'),
    ('Sắp xếp theo Ngày tạo / Ngày cập nhật', 'P2', 'Đã bật cột Ngày cập nhật.', '1. Bấm tiêu đề cột Ngày tạo, rồi Ngày cập nhật', '', '- Sắp xếp đúng theo thời gian'),
    ('Cột không sắp xếp được', 'P2', '', '1. Bấm tiêu đề cột Người tạo, Trạng thái', '', '- Không đổi thứ tự, không có biểu tượng sắp xếp'),
    ('Sắp xếp và bộ lọc giữ nguyên khi chuyển trang', 'P1', 'Đang lọc Trạng thái = Hoạt động, sắp xếp theo Mã.', '1. Sang trang 2', '', '- Bộ lọc và thứ tự giữ nguyên, không trùng/sót bản ghi'),
])

grp('TÌM KIẾM & LỌC DANH SÁCH', [
    ('Tìm nhanh theo tên', 'P0', 'Có gói tên chứa “điều hòa”.', '1. Gõ “điều hòa” vào ô tìm nhanh\n2. Chờ bảng nạp lại', 'Từ khóa: điều hòa', '- Chỉ còn gói có tên chứa từ khóa\n- Tổng số bản ghi đổi theo'),
    ('Tìm nhanh theo mã', 'P0', 'Có gói DOC-GBD-01..03.', '1. Gõ “DOC-GBD”', 'Từ khóa: DOC-GBD', '- Hiện đúng các gói có mã chứa DOC-GBD'),
    ('Tự lọc khi ngừng gõ, không cần bấm Tìm kiếm', 'P1', '', '1. Gõ từ khóa và dừng khoảng 1 giây', '', '- Danh sách tự nạp lại'),
    ('Không phân biệt chữ hoa chữ thường', 'P1', '', '1. Gõ “doc-gbd”\n2. Gõ “DOC-GBD”', '', '- Hai lần cho cùng kết quả'),
    ('Kết quả xếp theo độ khớp', 'P2', 'Có gói mã “DOC-GBD-01” và gói tên chứa “DOC-GBD-01” ở giữa.', '1. Gõ “DOC-GBD-01”', '', '- Gói trùng khít mã đứng đầu'),
    ('Lọc theo Trạng thái = Khóa', 'P0', 'Có gói Khóa.', '1. Chọn Trạng thái = Khóa', '', '- Chỉ hiện gói Khóa, lọc ngay không cần bấm Tìm kiếm'),
    ('Lọc theo Người tạo', 'P1', 'Có gói do DNS Admin tạo.', '1. Chọn Người tạo = DNS Admin', '', '- Chỉ hiện gói do người đó tạo\n- Ô chọn hiển thị dạng “Tên - Mã phòng - Mã nhân viên”'),
    ('Kết hợp nhiều tiêu chí', 'P1', '', '1. Gõ “DOC-GBD”\n2. Chọn Trạng thái = Hoạt động', '', '- Chỉ gói thỏa ĐỒNG THỜI hai tiêu chí'),
    ('Lọc khi đang ở trang giữa', 'P1', 'Đang ở trang 3.', '1. Chọn một tiêu chí lọc', '', '- Bảng quay về trang 1'),
    ('Nút Làm mới', 'P0', 'Đang áp dụng đủ tiêu chí.', '1. Bấm Làm mới', '', '- Mọi ô lọc trở về rỗng\n- Danh sách nạp lại đầy đủ ngay'),
    ('Ô lọc dùng nhãn nổi, không có placeholder trùng nhãn', 'P2', '', '1. Quan sát ô Trạng thái, Người tạo trước và sau khi chọn', '', '- Nhãn nằm trong ô, bay lên khi đã chọn\n- Ô tìm nhanh có placeholder “Tìm theo tên hoặc mã gói bảo dưỡng...”'),
])

grp('THÊM MỚI', [
    ('Mở trang Thêm mới', 'P0', 'Có quyền Thêm.', '1. Bấm Tạo mới', '', '- Mở TRANG “Thêm gói bảo dưỡng” (không phải popup)\n- Có 5 khối: Thông tin chung, Danh mục kiểm tra bảo dưỡng định kỳ, Giá vốn theo công ty, Áp dụng cho hàng hóa, File đính kèm (PDF)\n- Chân trang: Lưu, Lưu và tiếp tục, Quay lại'),
    ('Giá trị mặc định', 'P0', '', '1. Quan sát các ô khi vừa mở', '', '- Công ty quản lý = công ty của người đăng nhập\n- Trạng thái = Hoạt động\n- Bảng kiểm tra trống “Không có danh mục kiểm tra bảo dưỡng”\n- Giá vốn theo công ty: mọi công ty, hệ số 1'),
    ('Thêm mới đầy đủ và lưu thành công', 'P0', 'Có file PDF hợp lệ.', '1. Nhập Tên, Mã\n2. Thêm 1 dòng nội dung (ĐVT, SL), 1 cột cấp, chọn ghi chú, nhập Định mức công\n3. Chọn 2 hàng hóa, đính kèm PDF\n4. Bấm Lưu', 'Tên: Gói bảo dưỡng điều hòa (test N)\nMã: TEST-GBD-N\nCấp 1 (6T), Định mức công 2, Hệ số công nghệ 1.5, Hệ số giá bán 1.2', '- Báo “Tạo gói bảo dưỡng thành công”\n- Quay về danh sách, gói mới ở đầu, trạng thái Hoạt động'),
    ('Lịch sử Tạo mới sau khi thêm', 'P1', 'Vừa tạo gói ở TC trên.', '1. Mở Lịch sử của gói', '', '- Có mốc “Tạo mới” kèm người thực hiện, thời gian'),
    ('Lưu và tiếp tục', 'P0', 'Đã nhập đủ dữ liệu hợp lệ.', '1. Bấm Lưu và tiếp tục', '', '- Báo tạo thành công\n- Ở lại trang Thêm mới, form được làm trắng'),
    ('Thiếu Tên gói', 'P0', '', '1. Để trống Tên\n2. Bấm Lưu', '', '- Lỗi đỏ dưới ô Tên “Bắt buộc phải nhập”\n- Toast “Vui lòng kiểm tra lại dữ liệu nhập”\n- Trang không chuyển, dữ liệu còn nguyên'),
    ('Thiếu Mã gói', 'P0', 'Đã nhập Tên.', '1. Để trống Mã\n2. Bấm Lưu', '', '- Lỗi “Bắt buộc phải nhập” dưới ô Mã'),
    ('Mã tự chuyển chữ in hoa', 'P1', '', '1. Gõ mã chữ thường', 'Mã: doc-gbd-09', '- Ô hiển thị DOC-GBD-09 ngay khi gõ'),
    ('Mã sai định dạng', 'P0', '', '1. Gõ mã có dấu tiếng Việt / khoảng trắng / ký tự đặc biệt', 'Mã: mã gói bảo dưỡng\nMã: GBD@01', '- Báo ngay dưới ô “Chỉ gồm chữ không dấu, số, dấu - và _”'),
    ('Lỗi định dạng mã mất khi nhập lại đúng', 'P1', 'Đang có lỗi định dạng mã.', '1. Sửa lại mã đúng định dạng', 'Mã: GBD_01-A', '- Lỗi đỏ biến mất ngay'),
    ('Trùng Tên gói', 'P0', 'Đã có gói tên X.', '1. Nhập Tên = X\n2. Bấm Lưu', '', '- Lỗi “Đã tồn tại” dưới ô Tên, không tạo gói'),
    ('Trùng Mã gói (kể cả khác hoa thường)', 'P0', 'Đã có gói mã BD-HBTT8TAY.', '1. Nhập Mã = bd-hbtt8tay\n2. Bấm Lưu', '', '- Lỗi “Đã tồn tại” dưới ô Mã'),
    ('Tên / Ghi chú vượt 255 ký tự', 'P2', '', '1. Nhập 256 ký tự vào Tên (hoặc Ghi chú)\n2. Bấm Lưu', '', '- Báo “Tối đa 255 ký tự”'),
    ('Định mức đàm phán giá / VAT / Hệ số giá bán ngoài giới hạn', 'P1', '', '1. Nhập Định mức đàm phán giá = 100\n2. VAT = 101\n3. Hệ số giá bán = 0.5\n4. Bấm Lưu', '', '- Báo lỗi dưới từng ô (tối đa 99 / tối đa 100 / không được nhỏ hơn 1)\n- Giữ nguyên số đã gõ, không tự sửa'),
    ('VAT, Hệ số giá bán để trống', 'P2', '', '1. Để trống VAT và Hệ số giá bán\n2. Lưu thành công, mở chi tiết', '', '- VAT = 0, Hệ số giá bán = 1'),
    ('Thiếu nội dung / ĐVT / SL của dòng', 'P1', 'Đã thêm 1 dòng.', '1. Để trống ĐVT\n2. Bấm Lưu', '', '- Báo lỗi tại ô ĐVT, không lưu'),
    ('Thiếu ghi chú kiểm tra ở ô giao', 'P1', 'Có 1 dòng, 1 cột cấp.', '1. Không chọn ghi chú ở ô giao\n2. Bấm Lưu', '', '- Báo lỗi tại ô giao, không lưu'),
    ('Thiếu Định mức công của cấp', 'P1', 'Có 1 cột cấp.', '1. Để trống Định mức công\n2. Bấm Lưu', '', '- Báo “Bắt buộc phải nhập” tại Định mức công'),
    ('Chọn trùng cấp ở hai cột', 'P0', 'Cột 1 đã chọn Cấp 1 (6T).', '1. Thêm cột 2, chọn Cấp 1 (6T)', '', '- Báo “Cấp bảo dưỡng này đã được chọn ở cột khác”, không nhận'),
    ('Xóa cột cấp', 'P1', 'Có 2 cột cấp.', '1. Bấm × ở tiêu đề cột 2\n2. Bấm Xác nhận', '', '- Hỏi “Xóa cột … sẽ xóa toàn bộ ghi chú và giá của cột này. Bạn có chắc chắn?”\n- Xác nhận: cột biến mất; Hủy: giữ nguyên'),
    ('Xóa dòng nội dung', 'P2', 'Có 2 dòng.', '1. Bấm thùng rác ở dòng 2', '', '- Dòng 2 biến mất, STT đánh lại'),
    ('Giá tự tính', 'P0', 'Công ty quản lý có đơn giá công 700,000.', '1. Nhập Định mức công 2, Hệ số công nghệ 1.5, Hệ số giá bán gói 1.2', '', '- Giá vốn 2,100,000\n- Giá công thức 2,520,000\n- Giá bán cơ sở 2,520,000 (sửa tay được)\n- Giá bán theo công ty = Giá bán cơ sở × hệ số công ty'),
    ('Sửa tay Giá bán cơ sở', 'P1', 'Đã có giá tự tính.', '1. Sửa Giá bán cơ sở = 3,000,000\n2. Lưu, mở chi tiết', '', '- Giá bán cơ sở giữ 3,000,000'),
    ('Hệ số giá bán theo công ty', 'P1', '', '1. Nhập hệ số 1.15 cho một công ty khác công ty quản lý', '', '- Giá bán theo công ty đó = Giá bán cơ sở × 1.15\n- Dòng công ty quản lý bị khóa, luôn = 1'),
    ('Chọn hàng hóa áp dụng', 'P1', '', '1. Bấm Chọn hàng hóa\n2. Tích 2 hàng hóa, bấm “Thêm 2 hàng hoá”\n3. Đóng popup', '', '- Hàng hóa hiện trong khối, gom theo Nhóm hàng'),
    ('Chọn nhóm hàng', 'P2', '', '1. Bấm Chọn nhóm hàng, chọn 1 nhóm', '', '- Toàn bộ hàng hóa của nhóm được thêm'),
    ('Chưa đính kèm file PDF', 'P0', 'Nhập đủ dữ liệu khác.', '1. Không thêm tài liệu\n2. Bấm Lưu', '', '- Báo “Bắt buộc phải đính kèm ít nhất 1 file PDF”, không lưu'),
    ('Đính kèm file không phải PDF', 'P1', '', '1. Thêm tài liệu, chọn file .docx/.png', '', '- Hệ thống từ chối “Chỉ nhận file PDF”'),
    ('File PDF được lưu vào gói', 'P0', 'Đã đính kèm 1 PDF và lưu thành công.', '1. Mở chi tiết gói vừa tạo', '', '- Khối File đính kèm hiển thị đúng file đã chọn, bấm mở được'),
    ('Bấm Lưu nhiều lần liên tiếp', 'P1', 'Đã nhập đủ dữ liệu.', '1. Bấm Lưu 2 lần thật nhanh', '', '- Chỉ tạo ĐÚNG MỘT gói'),
    ('Quay lại khi chưa nhập gì', 'P2', '', '1. Bấm Quay lại', '', '- Về danh sách, không hỏi'),
    ('Quay lại khi đã nhập dở', 'P1', 'Đã nhập Tên.', '1. Bấm Quay lại', '', '- Hỏi “Thông tin chưa lưu” – “Bạn có thông tin chưa lưu. Có chắc chắn muốn thoát?”\n- Thoát: về danh sách, không lưu\n- Ở lại: giữ nguyên dữ liệu'),
])

grp('CHỈNH SỬA', [
    ('Mở trang Sửa', 'P0', 'Gói Hoạt động, có quyền Sửa.', '1. Bấm Sửa ở dòng gói', '', '- Mở trang “Sửa gói bảo dưỡng” với dữ liệu hiện tại\n- Chân trang: Lưu, Nhân bản, Quay lại (không có Lưu và tiếp tục)'),
    ('Sửa và lưu thành công', 'P0', '', '1. Đổi Tên\n2. Bấm Lưu', 'Tên mới: …(sửa)', '- Báo “Cập nhật gói bảo dưỡng thành công”, về danh sách\n- Tên trên bảng đổi, Người/Ngày cập nhật cập nhật'),
    ('Giữ nguyên Tên, Mã của chính gói', 'P0', '', '1. Không đổi Tên/Mã, đổi Ghi chú\n2. Lưu', '', '- Lưu được, không báo trùng'),
    ('Sửa trùng Tên với gói khác', 'P0', '', '1. Đổi Tên = tên gói khác\n2. Lưu', '', '- Báo “Đã tồn tại”'),
    ('Mã sai định dạng khi sửa', 'P1', '', '1. Đổi Mã = “Mã sửa”', '', '- Báo “Chỉ gồm chữ không dấu, số, dấu - và _”'),
    ('Sửa bảng nội dung / hệ số / hàng hóa ghi lịch sử', 'P1', '', '1. Thêm 1 dòng nội dung, đổi hệ số 1 công ty\n2. Lưu\n3. Mở Lịch sử', '', '- Có mốc “Thay đổi thông tin” liệt kê dòng thêm mới / thay đổi của bảng con'),
    ('Lưu mà không thay đổi gì', 'P2', '', '1. Mở Sửa, bấm Lưu ngay', '', '- Báo thành công, KHÔNG phát sinh mốc lịch sử mới'),
    ('Bỏ cột cấp ĐÃ phát sinh báo giá', 'P1', 'Gói có cấp đã dùng trong báo giá dịch vụ.', '1. Xóa cột cấp đó\n2. Lưu', '', '- Báo “Không thể xóa cấp dịch vụ đã được sử dụng!”, dữ liệu giữ nguyên'),
    ('Gói Khóa không có nút Sửa', 'P0', 'Gói đang Khóa.', '1. Quan sát cột Hành động và chân trang chi tiết', '', '- Không có nút Sửa ở cả 2 nơi'),
    ('Gõ thẳng đường dẫn Sửa của gói Khóa', 'P1', 'Gói đang Khóa (id X).', '1. Gõ /customer-care/services/X/edit', '', '- Báo “Gói bảo dưỡng đang bị khoá, vui lòng mở khoá trước khi sửa.”\n- Chuyển về trang Chi tiết'),
    ('Rời trang Sửa khi có thay đổi chưa lưu', 'P1', '', '1. Đổi Tên\n2. Bấm Quay lại', '', '- Hỏi “Thông tin chưa lưu” (Thoát / Ở lại)'),
    ('Gói import (chưa có PDF) phải đính kèm PDF mới sửa được', 'P1', 'Gói tạo bằng Import.', '1. Mở Sửa, đổi Ghi chú\n2. Lưu', '', '- Báo “Bắt buộc phải đính kèm ít nhất 1 file PDF”\n- Đính kèm PDF rồi Lưu → thành công'),
])

grp('XÓA', [
    ('Hộp xác nhận xóa', 'P0', 'Gói Hoạt động chưa được sử dụng.', '1. Bấm Xóa', '', '- Hộp “Xác nhận xóa”: “Bạn có chắc chắn muốn xóa gói bảo dưỡng ‘<tên>’? Hành động này không thể hoàn tác.”\n- Nút Xóa, Hủy'),
    ('Hủy xóa', 'P1', '', '1. Bấm Xóa\n2. Bấm Hủy', '', '- Không có gì thay đổi'),
    ('Xóa thành công', 'P0', '', '1. Bấm Xóa\n2. Bấm Xóa trong hộp', '', '- Báo “Xóa gói bảo dưỡng thành công”, gói biến mất, tổng giảm 1'),
    ('Xóa từ trang Chi tiết', 'P1', '', '1. Mở chi tiết gói chưa được sử dụng\n2. Bấm Xóa ở chân trang, xác nhận', '', '- Xóa thành công, quay về danh sách'),
    ('Gói đã gắn hàng hóa không có nút Xóa', 'P0', 'Gói đã gắn hàng hóa.', '1. Quan sát cột Hành động và chân trang chi tiết', '', '- Không có nút Xóa (ẩn hẳn, không mờ)'),
    ('Gói đã dùng ở báo giá dịch vụ không có nút Xóa', 'P0', 'Gói đã có trong báo giá dịch vụ.', '1. Quan sát cột Hành động', '', '- Không có nút Xóa'),
    ('Gói Khóa không có nút Xóa', 'P0', 'Gói đang Khóa.', '1. Quan sát cột Hành động', '', '- Không có nút Xóa'),
    ('Xóa dòng cuối của trang cuối', 'P2', 'Trang cuối chỉ còn 1 gói.', '1. Xóa gói đó', '', '- Danh sách tự lùi về trang trước, không hiện trang trắng'),
])

grp('KHÓA / MỞ KHÓA', [
    ('Khóa gói qua ô Trạng thái trang Sửa', 'P0', 'Gói Hoạt động.', '1. Mở Sửa\n2. Chọn Trạng thái = Khóa\n3. Lưu', '', '- Cột Trạng thái hiện nhãn đỏ “Khóa”\n- Gói vẫn nằm trong danh sách'),
    ('Sau khi Khóa mất nút Sửa và Xóa', 'P0', 'Gói vừa Khóa.', '1. Quan sát cột Hành động', '', '- Còn: Mở khóa, Nhân bản, In, Lịch sử'),
    ('Hộp xác nhận Mở khóa', 'P0', 'Gói Khóa.', '1. Bấm Mở khóa', '', '- “Xác nhận mở khóa”: “Bạn có chắc chắn muốn mở khóa gói bảo dưỡng ‘<tên>’?”; nút Mở khóa, Hủy'),
    ('Mở khóa thành công', 'P0', '', '1. Bấm Mở khóa\n2. Xác nhận', '', '- Báo “Mở khóa gói bảo dưỡng thành công”\n- Trạng thái Hoạt động, hiện lại Sửa'),
    ('Hủy Mở khóa', 'P1', '', '1. Bấm Mở khóa\n2. Bấm Hủy', '', '- Trạng thái không đổi'),
    ('Mở khóa từ trang Chi tiết', 'P1', 'Gói Khóa.', '1. Mở chi tiết\n2. Bấm Mở khóa ở chân trang, xác nhận', '', '- Mở khóa thành công'),
    ('Gói Khóa không chọn được ở màn khác', 'P1', 'Gói X đang Khóa.', '1. Mở màn báo giá dịch vụ, ô chọn gói bảo dưỡng', '', '- Không thấy gói X trong danh sách chọn'),
    ('Khóa / Mở khóa ghi lịch sử', 'P1', '', '1. Khóa rồi Mở khóa một gói\n2. Mở Lịch sử', '', '- Có mốc “Khóa” (Trạng thái: Hoạt động → Khóa) và “Mở khóa”'),
    ('Khóa gói trong lúc người khác đang sửa', 'P2', 'User A mở Sửa gói X; user B khóa gói X.', '1. User A bấm Lưu', '', '- Báo gói đang bị khoá, không lưu'),
])

grp('LỊCH SỬ THAY ĐỔI', [
    ('Mở lịch sử từ danh sách', 'P0', '', '1. Mở ba chấm → Lịch sử', '', '- Cửa sổ “Lịch sử thay đổi: <mã> - <tên>”, mốc mới nhất trên cùng'),
    ('Mở lịch sử từ trang chi tiết', 'P1', '', '1. Mở chi tiết\n2. Bấm Xem lịch sử', '', '- Khối Lịch sử cuối trang mở ra'),
    ('Mỗi mốc nêu đủ thông tin', 'P1', '', '1. Đọc một mốc “Thay đổi thông tin”', '', '- Thời điểm dd/mm/yyyy hh:mm, loại, người thực hiện + phòng ban, giá trị cũ → mới'),
    ('Lọc theo loại hành động', 'P2', 'Gói có nhiều loại mốc.', '1. Bấm Bộ lọc, chọn Thay đổi trạng thái', '', '- Chỉ còn mốc Khóa / Mở khóa'),
    ('Lọc theo người thực hiện, thời gian', 'P2', '', '1. Chọn Người thực hiện, Từ ngày – Đến ngày', '', '- Chỉ hiện mốc thỏa điều kiện; Làm mới xóa bộ lọc'),
    ('Lịch sử của gói tạo bằng Import', 'P2', 'Gói import.', '1. Mở Lịch sử', '', '- Có mốc Tạo mới, người thực hiện là người import'),
    ('Đóng cửa sổ Lịch sử', 'P2', '', '1. Bấm Đóng', '', '- Danh sách phía sau giữ nguyên bộ lọc và trang'),
])

grp('XEM CHI TIẾT', [
    ('Chi tiết hiển thị đủ 5 khối, chỉ đọc', 'P0', '', '1. Bấm mã gói', '', '- Đủ 5 khối như Thêm mới, mọi ô không gõ được\n- Có khối Lịch sử cuối trang'),
    ('Nút chân trang khớp với danh sách', 'P0', 'Gói Hoạt động chưa được sử dụng; đủ quyền.', '1. So nút chân trang chi tiết với nút trên dòng danh sách', '', '- Chi tiết có: Sửa, In, Nhân bản, Xóa, Quay lại (+ Xem lịch sử)\n- Không có nút nào ẩn ở danh sách mà hiện ở chi tiết'),
    ('Chi tiết gói Khóa', 'P1', 'Gói Khóa.', '1. Mở chi tiết', '', '- Chân trang: In, Nhân bản, Mở khóa, Quay lại; không có Sửa, Xóa'),
    ('Dữ liệu chi tiết khớp dữ liệu đã tạo', 'P0', 'Gói vừa tạo ở nhóm Thêm mới.', '1. Đối chiếu từng khối', '', '- Khớp Tên, Mã, số liệu giá, hàng hóa, file'),
    ('Quay lại về đúng nơi đi vào', 'P2', '', '1. Từ danh sách đang lọc, mở chi tiết\n2. Bấm Quay lại', '', '- Về danh sách, giữ bộ lọc'),
])

grp('NHÂN BẢN', [
    ('Nhân bản từ danh sách', 'P0', 'Có quyền Thêm.', '1. Bấm Nhân bản ở dòng gói', '', '- Mở trang “Sao chép gói bảo dưỡng” điền sẵn dữ liệu (kể cả hàng hóa, file)\n- Dòng nhắc “Đang sao chép từ gói … — hãy đổi tên/mã trước khi lưu.”'),
    ('Lưu bản sao chưa đổi Tên, Mã', 'P0', '', '1. Bấm Lưu ngay', '', '- Báo “Đã tồn tại” ở Tên và Mã'),
    ('Lưu bản sao sau khi đổi Tên, Mã', 'P0', '', '1. Đổi Tên, Mã\n2. Lưu', '', '- Tạo gói mới trạng thái Hoạt động; gói nguồn không đổi'),
    ('Nhân bản từ gói Khóa', 'P1', 'Gói nguồn đang Khóa.', '1. Nhân bản, đổi Tên/Mã, Lưu', '', '- Gói mới ở trạng thái Hoạt động'),
    ('Nhân bản từ chân trang chi tiết mở tab mới', 'P2', '', '1. Mở chi tiết, bấm Nhân bản', '', '- Trang Sao chép mở ở tab mới'),
])

grp('TÙY CHỈNH CỘT', [
    ('Mở cửa sổ Tuỳ chỉnh cột', 'P1', '', '1. Bấm biểu tượng cột', '', '- Liệt kê 10 cột kèm trạng thái bật/tắt'),
    ('Cột bị khóa', 'P1', '', '1. Thử bỏ tích / kéo STT, Mã, Hành động', '', '- Không bỏ tích, không kéo được'),
    ('Bật cột ẩn mặc định và lưu', 'P1', '', '1. Tích Công ty quản lý gói bảo dưỡng, Người cập nhật\n2. Lưu', '', '- Bảng hiện thêm 2 cột'),
    ('Cấu hình được ghi nhớ', 'P2', 'Đã lưu cấu hình.', '1. Tải lại trang / đăng nhập lại', '', '- Cấu hình giữ nguyên, không ảnh hưởng tài khoản khác'),
    ('Đóng không lưu', 'P2', '', '1. Đổi cấu hình\n2. Bấm Đóng', '', '- Bảng giữ nguyên'),
])

grp('XUẤT DANH SÁCH EXCEL', [
    ('Mở cửa sổ Chọn trường xuất file', 'P0', '', '1. Bấm Xuất Excel', '', '- Cửa sổ “Chọn trường xuất file” có 6 trường: Mã, Tên gói bảo dưỡng, Người tạo, Ngày tạo, Trạng thái, Công ty quản lý\n- Tích sẵn đúng các cột đang hiển thị'),
    ('Xuất file với trường mặc định', 'P0', '', '1. Bấm Xuất file', '', '- Tải Danh_sach_goi_bao_duong.xlsx, báo “Xuất Excel thành công”\n- File có đủ các cột đã tích (có cột Mã và Trạng thái) + cột Giá gói bảo dưỡng'),
    ('Bỏ tích một trường', 'P1', '', '1. Bỏ tích Người tạo\n2. Xuất file', '', '- File không có cột Người tạo'),
    ('Thứ tự cột theo thứ tự kéo', 'P2', '', '1. Kéo Trạng thái lên đầu\n2. Xuất file', '', '- Cột Trạng thái đứng đầu trong file'),
    ('Chọn tất cả / Bỏ chọn hết', 'P2', '', '1. Bấm Bỏ chọn hết\n2. Bấm Chọn tất cả', '', '- Bỏ chọn hết: nút Xuất file không bấm được\n- Chọn tất cả: tích đủ 6 trường'),
    ('Xuất lấy toàn bộ danh mục, không theo bộ lọc', 'P1', 'Đang lọc còn 3 gói.', '1. Xuất file', '', '- File chứa TOÀN BỘ gói của danh mục'),
    ('Dữ liệu file khớp hệ thống', 'P0', '', '1. Đối chiếu 5 dòng đầu file với màn hình', '', '- Mã, tên, trạng thái, người tạo, ngày tạo khớp\n- Giá theo cấp đúng, định dạng 1,234,000\n- Mã toàn số giữ số 0 ở đầu'),
])

grp('IMPORT TỪ FILE EXCEL', [
    ('Mở cửa sổ Import', 'P0', 'Có quyền Thêm.', '1. Bấm Import Excel', '', '- Cửa sổ “Import gói bảo dưỡng”: Chọn file Excel, Tải file mẫu, Load lên bảng, Validate, Import, Làm mới, Đóng\n- Load/Validate/Import mờ khi chưa có file'),
    ('Tải file mẫu', 'P0', '', '1. Bấm Tải file mẫu', '', '- Tải Mau_import_goi_bao_duong.xlsx gồm 5 sheet: 1. Gói bảo dưỡng, 2. Cấp bảo dưỡng, 3. Nội dung kiểm tra, 4. Hệ số theo công ty, 5. Hàng hoá; cột bắt buộc có dấu *'),
    ('Chọn file không phải Excel', 'P1', '', '1. Chọn file .pdf', '', '- Báo “Vui lòng chọn file .xlsx hoặc .xls”'),
    ('File thiếu sheet / thiếu cột', 'P1', 'File mẫu đã xóa sheet 2.', '1. Chọn file, Load lên bảng', '', '- Báo “File không đúng mẫu: thiếu sheet …”'),
    ('Load file đúng mẫu', 'P0', 'File có 4 gói.', '1. Chọn file\n2. Load lên bảng', '', '- Dữ liệu hiện theo 5 tab, Tổng = 4'),
    ('Validate — mã sai định dạng', 'P0', 'Sheet 1 có Mã “Mã lỗi Việt”.', '1. Validate', '', '- Dòng lỗi tô đỏ “Mã gói: Chỉ gồm chữ không dấu, số, dấu - và _”'),
    ('Validate — mã đã tồn tại', 'P0', 'Sheet 1 có Mã BD-HBTT8TAY.', '1. Validate', '', '- “Mã gói đã tồn tại trong hệ thống”'),
    ('Validate — trùng mã trong file', 'P1', 'Hai dòng cùng mã.', '1. Validate', '', '- “Mã gói bị trùng với dòng N trong file”'),
    ('Validate — công ty không có trong danh mục', 'P1', '', '1. Validate', 'Công ty: Công ty không tồn tại', '- “Công ty quản lý “Công ty không tồn tại” không có trong danh mục”'),
    ('Validate — VAT ngoài giới hạn', 'P1', '', '1. Validate', 'VAT: 150', '- “VAT tối đa 100”'),
    ('Validate — cấp không có trong danh mục', 'P1', 'Sheet 2 cấp “Cấp 9 không có”.', '1. Validate', '', '- Lỗi tại sheet 2; gói ở sheet 1 báo “Sheet “2. Cấp bảo dưỡng” dòng N có lỗi”'),
    ('Validate — ghi chú kiểm tra không có trong danh mục', 'P1', 'Sheet 3 ghi chú “Ghi chú không có”.', '1. Validate, mở tab 3', '', '- “Ghi chú kiểm tra “Ghi chú không có” không có trong danh mục”'),
    ('Validate — mã gói ở sheet con không có ở sheet 1', 'P2', '', '1. Validate', '', '- “Mã gói “X” không có ở sheet “1. Gói bảo dưỡng””'),
    ('Tab sheet có lỗi hiện số lỗi', 'P2', 'File có lỗi ở sheet 1–3.', '1. Validate', '', '- Tab lỗi hiện “n lỗi” màu đỏ; ô Hợp lệ / Lỗi đúng số'),
    ('Còn dòng lỗi thì không Import được', 'P0', 'Còn dòng lỗi.', '1. Quan sát nút Import', '', '- Nút Import không bấm được'),
    ('Sửa dòng lỗi rồi Validate lại', 'P1', '', '1. Sửa mã lỗi thành mã hợp lệ\n2. Validate', '', '- Dòng hết lỗi, dòng hợp lệ bị khóa'),
    ('Bỏ dòng lỗi', 'P0', 'Có 2 gói lỗi.', '1. Bấm Bỏ dòng lỗi\n2. Validate', '', '- “Đã bỏ N dòng lỗi (kèm dòng con của gói bị bỏ). Hãy bấm Validate lại.”\n- Sau Validate: “Tất cả N gói bảo dưỡng đã hợp lệ. Có thể import ngay.”'),
    ('Xoá trạng thái validate', 'P2', 'Đã Validate.', '1. Bấm Xoá trạng thái validate', '', '- Dòng hợp lệ được mở khóa, thông báo lỗi biến mất'),
    ('Import thành công', 'P0', 'Tất cả hợp lệ.', '1. Bấm Import', '', '- “Import thành công N gói bảo dưỡng.”\n- Cửa sổ đóng, danh sách có gói mới, trạng thái Hoạt động, mã in hoa'),
    ('Dữ liệu import đầy đủ', 'P0', 'Vừa import gói DOC-GBD-01.', '1. Mở chi tiết gói', '', '- Có đủ cấp, nội dung kiểm tra, ghi chú, hệ số công ty, hàng hóa như file\n- VAT trống → 0, Hệ số giá bán trống → 1\n- Khối File đính kèm trống'),
    ('Đóng cửa sổ khi đã tải dữ liệu', 'P2', 'Đã Load lên bảng.', '1. Bấm Đóng', '', '- Hỏi “Thông tin chưa lưu”; Thoát thì đóng, Ở lại thì giữ'),
])

grp('IN PHIẾU', [
    ('Mở xem trước phiếu từ danh sách', 'P1', '', '1. Ba chấm → In', '', '- Cửa sổ “Xem trước gói bảo dưỡng <mã>” hiện phiếu trên khổ giấy'),
    ('Nội dung phiếu', 'P1', 'Gói có 1 cấp, 1 nội dung, ghi chú.', '1. Đọc phiếu', '', '- “DANH MỤC KIỂM TRA BẢO DƯỠNG ĐỊNH KỲ”, “TÊN DỊCH VỤ: <tên in hoa>”\n- Bảng STT, Nội dung, SL, cột cấp, Kiểm tra (Có/Không), Ghi chú\n- Ghi chú gói, bảng ký hiệu, chỗ ký Kỹ thuật viên / Khách hàng'),
    ('In từ trang chi tiết', 'P2', '', '1. Mở chi tiết, bấm In', '', '- Mở cùng cửa sổ xem trước'),
    ('Bấm In trên cửa sổ xem trước', 'P2', '', '1. Bấm In', '', '- Mở hộp thoại in của trình duyệt'),
])

grp('RÀNG BUỘC VỚI CÁC CHỨC NĂNG KHÁC', [
    ('Gói mới dùng được ở báo giá dịch vụ', 'P1', 'Vừa tạo gói X.', '1. Mở màn báo giá dịch vụ, chọn gói bảo dưỡng', '', '- Chọn được gói X, giá theo cấp khớp'),
    ('Chọn gói vào báo giá thì gói mất nút Xóa', 'P1', 'Gói X chưa được sử dụng.', '1. Dùng gói X ở một báo giá dịch vụ\n2. Quay lại danh mục', '', '- Gói X không còn nút Xóa'),
    ('Khóa cấp / ghi chú / ĐVT đang dùng trong gói', 'P2', 'ĐVT “Bộ” đang dùng trong gói X.', '1. Khóa ĐVT “Bộ” ở danh mục ĐVT\n2. Mở Sửa gói X', '', '- Ô ĐVT vẫn hiển thị “Bộ” kèm 🔒, không mất dữ liệu khi lưu'),
    ('Đổi tên danh mục phụ thuộc', 'P2', 'Cấp “Cấp 1 (6T)” dùng trong gói X.', '1. Đổi tên cấp ở danh mục Cấp dịch vụ\n2. Mở chi tiết gói X', '', '- Hiện tên cấp mới'),
])


def cp(sc, dc):
    """Chép định dạng ô giữa 2 workbook (không dùng _style vì chỉ số style khác nhau)."""
    if sc.has_style:
        dc.font = copy(sc.font)
        dc.fill = copy(sc.fill)
        dc.border = copy(sc.border)
        dc.alignment = copy(sc.alignment)
        dc.number_format = sc.number_format
        dc.protection = copy(sc.protection)


def copy_row_style(ws, src_row, dst_row, ncol=18):
    for c in range(1, ncol + 1):
        s = ws.cell(src_row, c)
        d = ws.cell(dst_row, c)
        if s.has_style:
            d._style = copy(s._style)


def build(book, out):
    src_wb = openpyxl.load_workbook(book)
    src = src_wb[MAU]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TAB
    # chép dòng 1-19 + dòng nhóm 18 làm mẫu định dạng
    for row in src.iter_rows(min_row=1, max_row=19):
        for cell in row:
            d = ws.cell(cell.row, cell.column, cell.value if cell.row <= 17 else None)
            cp(cell, d)
    for letter, dim in src.column_dimensions.items():
        if dim.width:
            ws.column_dimensions[letter].width = dim.width
    ws.column_dimensions['R'].width = 14
    for r in range(1, 18):
        if src.row_dimensions[r].height:
            ws.row_dimensions[r].height = src.row_dimensions[r].height
    for rng in src.merged_cells.ranges:
        if rng.max_row <= 17:
            ws.merge_cells(str(rng))
    # khối mô tả
    ws['A1'] = 'MÔ TẢ TÍNH NĂNG (đọc trước khi xem testcase)'
    for i, (k, v) in enumerate(MOTA):
        ws.cell(2 + i, 1, k)
        ws.cell(2 + i, 2, v)
    ws['A11'] = 'Testcase _ Danh mục gói bảo dưỡng - Cập nhật ngày %s' % NGAY

    # dữ liệu
    r = 18
    n_sec = 0
    dv_dns = DataValidation(type='list', formula1='"Passed,Failed,Pending,Not Executed"', allow_blank=True)
    dv_tp = DataValidation(type='list', formula1='"P,F,PE"', allow_blank=True)
    ws.add_data_validation(dv_dns)
    ws.add_data_validation(dv_tp)
    roman = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII', 'VIII', 'IX', 'X', 'XI', 'XII', 'XIII', 'XIV', 'XV', 'XVI']
    first_data = None
    for name, rows in G:
        n_sec += 1
        for c in range(1, 19):
            cp(src.cell(18, c), ws.cell(r, c))
        ws.cell(r, 3, '%s. %s' % (roman[n_sec - 1], name))
        ws.merge_cells(start_row=r, start_column=3, end_row=r, end_column=17)
        r += 1
        start = r
        for j, (fn, pri, pre, steps, data, exp) in enumerate(rows):
            for c in range(1, 19):
                cp(src.cell(19, c), ws.cell(r, c))
            ws.cell(r, 1, MODULE)
            ws.cell(r, 2, name)
            ws.cell(r, 3, 'TC_%02d.%03d' % (n_sec, j + 1))
            ws.cell(r, 4, fn)
            ws.cell(r, 5, pri)
            ws.cell(r, 6, pre)
            ws.cell(r, 7, steps)
            ws.cell(r, 8, data)
            ws.cell(r, 9, exp)
            ws.cell(r, 18, '=IF(COUNTIF(O%d:Q%d,"P")>0,"P","F")' % (r, r))
            first_data = first_data or r
            r += 1
        end = r - 1
        ws.merge_cells(start_row=start, start_column=1, end_row=end, end_column=1)
        ws.merge_cells(start_row=start, start_column=2, end_row=end, end_column=2)
        dv_dns.add('K%d:M%d' % (start, end))
        dv_tp.add('O%d:Q%d' % (start, end))
    last = r - 1
    # TEST SUMMARY: công thức theo đúng vùng dữ liệu của tab này
    ws['K11'] = '=COUNTIF(K18:M%d,"Passed")' % last
    ws['K12'] = '=COUNTIF(K18:M%d,"Failed")' % last
    ws['K13'] = '=COUNTIF(K18:M%d,"Pending")' % last
    ws['K14'] = '=K16-K15'
    ws['K15'] = '=COUNTIF(K18:K%d,"<>")' % last
    ws['K16'] = '=COUNTIF(I18:I%d,"*")' % last
    ws['P11'] = '=COUNTIF(R18:R%d,"P")' % last
    ws['P12'] = '=COUNTIF(R18:R%d,"F")' % last
    ws['P13'] = '=COUNTIF(R18:R%d,"PE")' % last
    ws['P14'] = '=K16-P15'
    ws['P15'] = '=SUM(P11:Q13)'
    total = sum(len(x[1]) for x in G)
    p0 = sum(1 for x in G for t in x[1] if t[1] == 'P0')
    wb.save(out)
    return total, p0, out


if __name__ == '__main__':
    out = os.path.join(HERE, 'out', 'testcase - Danh mục gói bảo dưỡng.xlsx')
    print(build(sys.argv[1], out))
