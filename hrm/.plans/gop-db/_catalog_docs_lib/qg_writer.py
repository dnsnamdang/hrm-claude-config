# -*- coding: utf-8 -*-
"""Dựng HDSD / SRS mới trên KHUÔN của tài liệu Danh mục quốc gia (bản TPE đã duyệt 23/09/2026).

Vì sao không dùng hdsd_engine / srs_docx_lib: bộ tài liệu Quốc gia đã được QA chỉnh trực tiếp
trên Google Docs (font, lề, màu chú thích, bullet, bảng, cách chèn icon nút ngay trong câu…).
Team chốt đó là bản "đúng" để các danh mục khác bám theo. Cách chắc nhất để giống 100% là
lấy CHÍNH file đó làm vỏ: giữ trang bìa + mục lục + style, xoá thân bài, rồi dựng thân bài mới
bằng cách NHÂN BẢN các đoạn/bảng mẫu có sẵn trong file.

Mỗi đoạn mẫu được chọn theo "vai trò" (heading, đoạn thường, bullet, bước, ảnh, chú thích,
bảng N cột) — xem `roles` do từng generator truyền vào.
"""
import copy
import os
import re

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.table import Table
from docx.text.paragraph import Paragraph


def _strip_bookmarks(el):
    for tag in ('w:bookmarkStart', 'w:bookmarkEnd'):
        for b in el.findall('.//' + qn(tag)):
            b.getparent().remove(b)


def _clear_runs(p_el):
    for child in list(p_el):
        if child.tag in (qn('w:r'), qn('w:hyperlink'), qn('w:ins'), qn('w:del'),
                         qn('w:smartTag'), qn('w:fldSimple')):
            p_el.remove(child)


def _first_text_run(p_el):
    for r in p_el.iter(qn('w:r')):
        if r.find(qn('w:t')) is not None and (r.find(qn('w:t')).text or '').strip():
            return r
    for r in p_el.iter(qn('w:r')):
        if r.find(qn('w:drawing')) is None:
            return r
    return None


class QgWriter(object):
    """roles: dict vai_trò -> chỉ số phần tử body (w:p / w:tbl) trong file mẫu."""

    def __init__(self, template, roles, body_from, cover_replace=None):
        self.doc = Document(template)
        body = self.doc.element.body
        els = list(body.iterchildren())
        self.tmpl = {}
        for k, idx in roles.items():
            self.tmpl[k] = copy.deepcopy(els[idx])
        # run mẫu để lấy định dạng chữ cho từng vai trò có chữ
        self.run_tmpl = {}
        for k, el in self.tmpl.items():
            if el.tag == qn('w:p'):
                r = _first_text_run(el)
                if r is not None:
                    r = copy.deepcopy(r)
                    for c in list(r):
                        if c.tag != qn('w:rPr'):
                            r.remove(c)
                    self.run_tmpl[k] = r
        # xoá thân bài cũ (giữ phần trước body_from: bìa + "MỤC LỤC" + khối mục lục)
        self.sectPr = body.find(qn('w:sectPr'))
        for el in els[body_from:]:
            if el is not self.sectPr:
                body.remove(el)
        if cover_replace:
            for p in self.doc.paragraphs[:body_from]:
                for old, new in cover_replace.items():
                    if old in p.text:
                        for t in p._p.iter(qn('w:t')):
                            if t.text and old in t.text:
                                t.text = t.text.replace(old, new)
        self.bm = 9000
        self.fig = 0
        self.headings = []   # (level, text, bookmark)

    # ------------------------------------------------------------ tiện ích
    def _append(self, el):
        self.sectPr.addprevious(el)
        return el

    def _run(self, role, text, bold=False, italic=False, color=None):
        src = self.run_tmpl.get(role)
        r = copy.deepcopy(src if src is not None else self.run_tmpl['p'])
        rpr = r.find(qn('w:rPr'))
        if rpr is None:
            rpr = r.makeelement(qn('w:rPr'), {})
            r.insert(0, rpr)
        for tag, on in ((qn('w:b'), bold), (qn('w:i'), italic)):
            for e in rpr.findall(tag):
                rpr.remove(e)
            if on:
                e = rpr.makeelement(tag, {})
                e.set(qn('w:val'), '1')
                rpr.insert(0, e)
        if color:
            for e in rpr.findall(qn('w:color')):
                rpr.remove(e)
            e = rpr.makeelement(qn('w:color'), {})
            e.set(qn('w:val'), color)
            rpr.append(e)
        t = r.makeelement(qn('w:t'), {})
        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
        t.text = text
        r.append(t)
        return r

    def _fill(self, p_el, role, segs):
        """segs: chuỗi, hoặc list gồm chuỗi / ('b', text) / ('i', text) / ('img', path, cao_inch)."""
        if isinstance(segs, str):
            segs = [segs]
        par = Paragraph(p_el, self.doc._body)
        for s in segs:
            if isinstance(s, str):
                # **đậm** trong chuỗi
                for j, part in enumerate(re.split(r'\*\*', s)):
                    if part:
                        p_el.append(self._run(role, part, bold=(j % 2 == 1)))
            elif s[0] == 'b':
                p_el.append(self._run(role, s[1], bold=True))
            elif s[0] == 'i':
                p_el.append(self._run(role, s[1], italic=True))
            elif s[0] == 'img':
                h = s[2] if len(s) > 2 else 0.22
                par.add_run().add_picture(s[1], height=Inches(h))
        return p_el

    def _para(self, role, segs):
        el = copy.deepcopy(self.tmpl[role])
        _strip_bookmarks(el)
        _clear_runs(el)
        self._fill(el, role, segs)
        return self._append(el)

    def _bookmark(self, el):
        self.bm += 1
        name = '_heading=h.gbd%d' % self.bm
        s = el.makeelement(qn('w:bookmarkStart'), {})
        s.set(qn('w:id'), str(self.bm))
        s.set(qn('w:name'), name)
        e = el.makeelement(qn('w:bookmarkEnd'), {})
        e.set(qn('w:id'), str(self.bm))
        ppr = el.find(qn('w:pPr'))
        (ppr.addnext if ppr is not None else (lambda x: el.insert(0, x)))(e)
        (ppr.addnext if ppr is not None else (lambda x: el.insert(0, x)))(s)
        return name

    # ------------------------------------------------------------ API
    def page_break(self):
        el = copy.deepcopy(self.tmpl['pagebreak'])
        _strip_bookmarks(el)
        self._append(el)

    def h1(self, text, new_page=True):
        if new_page:
            self.page_break()
        el = self._para('h1', text)
        self.headings.append((1, text, self._bookmark(el)))

    def h2(self, text):
        el = self._para('h2', text)
        self.headings.append((2, text, self._bookmark(el)))

    def h3(self, text):
        role = 'h3' if 'h3' in self.tmpl else 'h2'
        el = self._para(role, text)
        self.headings.append((3, text, self._bookmark(el)))

    def p(self, segs):
        return self._para('p', segs)

    def bullet(self, segs):
        return self._para('bullet', segs)

    def step(self, segs):
        return self._para('step', segs)

    def raw(self, role, segs):
        """Đoạn theo vai trò bất kỳ đã khai trong roles (vd 'sub' = tiêu đề con in đậm)."""
        return self._para(role, segs)

    def clone_texts(self, role, texts):
        """Nhân bản NGUYÊN đoạn mẫu (giữ liên kết, định dạng từng run), chỉ thay chữ của các
        w:t theo thứ tự. Dùng cho đoạn 'Quy tắc chung: … SRS_Các quy tắc chung_VN_1.0 …'."""
        el = copy.deepcopy(self.tmpl[role])
        _strip_bookmarks(el)
        ts = list(el.iter(qn('w:t')))
        for t, v in zip(ts, texts):
            if v is not None:
                t.text = v
        return self._append(el)

    def blank(self):
        el = copy.deepcopy(self.tmpl['blank'])
        _strip_bookmarks(el)
        _clear_runs(el)
        self._append(el)

    def image(self, path, width=6.0, caption=None):
        assert os.path.exists(path), 'Thiếu ảnh: %s' % path
        el = copy.deepcopy(self.tmpl['img'])
        _strip_bookmarks(el)
        _clear_runs(el)
        self._append(el)
        from PIL import Image
        w, h = Image.open(path).size
        # ảnh cắt nhỏ (1 cụm nút, 1 ô lọc) không phóng to quá kích thước thật (~110 px/inch)
        width = min(width, max(w / 110.0, 2.0))
        # ảnh dọc quá dài thì khống chế chiều cao 7.5 inch
        if h / float(w) * width > 7.5:
            Paragraph(el, self.doc._body).add_run().add_picture(path, height=Inches(7.5))
        else:
            Paragraph(el, self.doc._body).add_run().add_picture(path, width=Inches(width))
        if caption:
            self.fig += 1
            self._para('cap', 'Hình %d: %s' % (self.fig, caption))
        return self.fig

    def table(self, rows, widths=None):
        """rows[0] là dòng tiêu đề. Chọn bảng mẫu có số cột >= số cột cần, bớt cột thừa."""
        n = len(rows[0])
        cands = sorted((k for k in self.tmpl if k.startswith('tbl')), key=lambda k: int(k[3:]))
        key = next((k for k in cands if int(k[3:]) >= n), cands[-1])
        tbl_el = copy.deepcopy(self.tmpl[key])
        _strip_bookmarks(tbl_el)
        tbl = Table(tbl_el, self.doc._body)
        m = int(key[3:])
        # bớt cột thừa (xoá cột cuối cho tới khi đủ n)
        grid = tbl_el.find(qn('w:tblGrid'))
        for _ in range(m - n):
            cols = grid.findall(qn('w:gridCol'))
            grid.remove(cols[-1])
            for tr in tbl_el.findall(qn('w:tr')):
                tcs = tr.findall(qn('w:tc'))
                tr.remove(tcs[-1])
        # đủ số dòng: dòng 0 là tiêu đề, nhân bản dòng 1 làm dòng dữ liệu
        trs = tbl_el.findall(qn('w:tr'))
        body_tmpl = copy.deepcopy(trs[1] if len(trs) > 1 else trs[0])
        for tr in trs[1:]:
            tbl_el.remove(tr)
        for _ in rows[1:]:
            tbl_el.append(copy.deepcopy(body_tmpl))
        # độ rộng cột: tổng giữ nguyên theo bảng mẫu
        total = sum(int(float(c.get(qn('w:w')))) for c in grid.findall(qn('w:gridCol')))
        tblw = tbl_el.find(qn('w:tblPr')).find(qn('w:tblW'))
        if tblw is not None and tblw.get(qn('w:type')) == 'dxa' and int(float(tblw.get(qn('w:w')))) > total:
            total = int(float(tblw.get(qn('w:w'))))
        # bảng luôn rộng bằng vùng chữ của trang (như bảng trong tài liệu mẫu)
        sec = self.doc.sections[-1]
        text_w = int((sec.page_width - sec.left_margin - sec.right_margin) / 635)  # EMU -> dxa
        total = max(total, text_w)
        if tblw is not None:
            tblw.set(qn('w:w'), str(total))
            tblw.set(qn('w:type'), 'dxa')
        widths = widths or [1.0] * n
        s = float(sum(widths))
        ws = [int(total * w / s) for w in widths]
        for c, w in zip(grid.findall(qn('w:gridCol')), ws):
            c.set(qn('w:w'), str(w))
        for tr in tbl_el.findall(qn('w:tr')):
            for tc, w in zip(tr.findall(qn('w:tc')), ws):
                tcw = tc.find(qn('w:tcPr')).find(qn('w:tcW'))
                if tcw is not None:
                    tcw.set(qn('w:w'), str(w))
                    tcw.set(qn('w:type'), 'dxa')
        tbl = Table(tbl_el, self.doc._body)
        for r_i, row in enumerate(rows):
            for c_i, val in enumerate(row):
                cell = tbl.rows[r_i].cells[c_i]
                ps = cell.paragraphs
                for extra in ps[1:]:
                    extra._p.getparent().remove(extra._p)
                p0 = ps[0]._p
                role_run = _first_text_run(p0)
                rtmpl = copy.deepcopy(role_run) if role_run is not None else None
                _clear_runs(p0)
                lines = val if isinstance(val, list) else str(val).split('\n')
                for li, line in enumerate(lines):
                    target = p0 if li == 0 else copy.deepcopy(p0)
                    if li > 0:
                        _clear_runs(target)
                        cell._tc.append(target)
                    for j, part in enumerate(re.split(r'\*\*', line)):
                        if not part:
                            continue
                        if rtmpl is not None:
                            r = copy.deepcopy(rtmpl)
                            for c in list(r):
                                if c.tag != qn('w:rPr'):
                                    r.remove(c)
                        else:
                            r = p0.makeelement(qn('w:r'), {})
                        if j % 2 == 1:
                            rpr = r.find(qn('w:rPr'))
                            if rpr is None:
                                rpr = r.makeelement(qn('w:rPr'), {})
                                r.insert(0, rpr)
                            b = rpr.makeelement(qn('w:b'), {})
                            b.set(qn('w:val'), '1')
                            rpr.insert(0, b)
                        t = r.makeelement(qn('w:t'), {})
                        t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
                        t.text = part
                        r.append(t)
                        target.append(r)
        self._append(tbl_el)
        self.blank()

    # ------------------------------------------------------------ mục lục
    def rebuild_toc(self, levels=(1, 2)):
        """Dựng lại các dòng mục lục (khối w:sdt) theo heading mới; số trang để trống,
        Word/Google Docs cập nhật lại khi mở. Nhân bản dòng mẫu cấp 1 / cấp 2 có sẵn."""
        sdt = self.doc.element.body.find(qn('w:sdt'))
        if sdt is None:
            return
        content = sdt.find(qn('w:sdtContent'))
        ps = content.findall(qn('w:p'))
        entries = [p for p in ps if p.find(qn('w:hyperlink')) is not None]
        if not entries:
            return
        lvl_tmpl = {}
        for p in entries:
            ind = p.find('.//' + qn('w:ind'))
            left = int(ind.get(qn('w:left'), '0')) if ind is not None else 0
            lvl = 1 if left == 0 else 2
            lvl_tmpl.setdefault(lvl, p)
        first = entries[0]
        anchor_pos = first
        # Giữ nguyên mã trường TOC: run "begin/instr/separate" nằm trước hyperlink ở dòng đầu,
        # run "end" nằm sau hyperlink ở dòng cuối — mất 1 trong 2 là Word không nhận ra mục lục.
        begin_runs = []
        for c in list(first):
            if c.tag == qn('w:hyperlink'):
                break
            if c.tag == qn('w:r'):
                begin_runs.append(c)
        end_runs = []
        for p in entries:
            hl = p.find(qn('w:hyperlink'))
            after = False
            for c in list(p):
                if c is hl:
                    after = True
                    continue
                if after and c.tag == qn('w:r') and c.find(qn('w:fldChar')) is not None:
                    end_runs.append(c)
        new = []
        for lvl, text, bm in self.headings:
            if lvl not in levels:
                continue
            t = copy.deepcopy(lvl_tmpl.get(min(lvl, 2), lvl_tmpl[1]))
            if lvl >= 3:
                ind = t.find('.//' + qn('w:ind'))
                if ind is not None:
                    ind.set(qn('w:left'), '720')
            hl = t.find(qn('w:hyperlink'))
            hl.set(qn('w:anchor'), bm)
            runs = hl.findall(qn('w:r'))
            for r in runs[1:]:
                hl.remove(r)
            r = runs[0]
            for c in list(r):
                if c.tag != qn('w:rPr'):
                    r.remove(c)
            t1 = r.makeelement(qn('w:t'), {})
            t1.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
            t1.text = text
            r.append(t1)
            r.append(r.makeelement(qn('w:tab'), {}))
            # bỏ run fldChar thừa bị nhân bản theo dòng mẫu
            for c in list(t):
                if c.tag == qn('w:r') and (c.find(qn('w:fldChar')) is not None
                                           or c.find(qn('w:instrText')) is not None):
                    t.remove(c)
            new.append(t)
        if new:
            hl0 = new[0].find(qn('w:hyperlink'))
            for r in begin_runs:
                hl0.addprevious(r)
            for r in end_runs:
                new[-1].append(r)
        for p in entries[1:]:
            content.remove(p)
        for t in new:
            anchor_pos.addprevious(t)
        content.remove(anchor_pos)

    def _purge_orphan_images(self):
        """Bỏ quan hệ ảnh không còn được thân bài tham chiếu (ảnh của tài liệu mẫu đã xoá),
        nếu không file .docx vẫn mang theo toàn bộ ảnh cũ."""
        from lxml import etree
        part = self.doc.part
        xml = etree.tostring(part._element, encoding='unicode')
        used = set(re.findall(r'r:(?:embed|link|id)="([^"]+)"', xml))
        for rId, rel in list(part.rels.items()):
            if rel.reltype.endswith('/image') and rId not in used:
                part.drop_rel(rId) if hasattr(part, 'drop_rel') else part.rels.pop(rId)

    def save(self, path):
        self._purge_orphan_images()
        # File mẫu (xuất từ Google Docs) bật "nhúng font"; Word lưu lại sẽ nhét ~15MB font vào file.
        st = self.doc.settings.element
        for e in st.findall(qn('w:embedTrueTypeFonts')) + st.findall(qn('w:saveSubsetFonts')):
            st.remove(e)
        self.doc.core_properties.title = os.path.splitext(os.path.basename(path))[0]
        self.doc.save(path)
        return path


def strip_embedded_fonts(path):
    """Gỡ font nhúng khỏi file .docx đã lưu (Word lưu lại file mẫu Google Docs sẽ nhúng ~15MB
    font). Chạy SAU bước Word cập nhật mục lục."""
    import shutil
    import tempfile
    import zipfile
    tmp = tempfile.mktemp(suffix='.docx')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            name = item.filename
            if name.startswith('word/fonts/'):
                continue
            data = zin.read(name)
            if name == 'word/fontTable.xml':
                data = re.sub(rb'<w:embed(Regular|Bold|Italic|BoldItalic)\b[^>]*/>', b'', data)
            elif name == 'word/_rels/fontTable.xml.rels':
                data = re.sub(rb'<Relationship [^>]*fonts/[^>]*/>', b'', data)
            elif name == 'word/settings.xml':
                data = re.sub(rb'<w:(embedTrueTypeFonts|saveSubsetFonts)\b[^>]*/>', b'', data)
            elif name == '[Content_Types].xml':
                data = re.sub(rb'<Default Extension="odttf"[^>]*/>', b'', data)
            zout.writestr(item, data)
    shutil.move(tmp, path)
    return os.path.getsize(path)
