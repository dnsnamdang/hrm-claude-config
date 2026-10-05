# -*- coding: utf-8 -*-
"""Cập nhật HDSD_PheDuyet_YeuCauGiaiPhap.docx (bản Drive) — GIỮ nội dung cũ, CHÈN phần thiếu,
sửa tại chỗ những chỗ lệch code (02/10/2026).

Chạy: /opt/homebrew/bin/python3 gen_hdsd.py
Nguồn: ../drive_orig/HDSD_PheDuyet_YeuCauGiaiPhap.docx (không sửa) → ./HDSD_PheDuyet_YeuCauGiaiPhap.docx
"""
import copy, io, os, re
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from docx.text.paragraph import Paragraph
from lxml import etree
from PIL import Image

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, '..', 'drive_orig', 'HDSD_PheDuyet_YeuCauGiaiPhap.docx')
OUT = os.path.join(HERE, 'HDSD_PheDuyet_YeuCauGiaiPhap.docx')
S = os.path.join(HERE, 'shots') + '/'
I = os.path.join(HERE, 'icons') + '/'
W14 = '{http://schemas.microsoft.com/office/word/2010/wordml}paraId'

d = Document(SRC)
body = d.element.body
B = list(body)                      # chỉ số theo bản gốc
TPL = {k: copy.deepcopy(B[k]) for k in (110, 111, 112, 114, 68, 115, 116)}


# ---------------------------------------------------------------- helpers
def _strip_ids(el):
    for e in el.iter():
        if W14 in e.attrib:
            del e.attrib[W14]
    return el


def set_text(p, text):
    """Đặt lại chữ của đoạn: giữ rPr của run chữ đầu tiên, bỏ các run chữ khác."""
    runs = [r for r in p.iter(qn('w:r')) if r.find(qn('w:t')) is not None]
    if not runs:
        raise ValueError('no text run')
    first = runs[0]
    for r in runs[1:]:
        r.getparent().remove(r)
    for t in first.findall(qn('w:t'))[1:]:
        first.remove(t)
    t = first.find(qn('w:t'))
    t.text = text
    t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')


def ptext(p):
    return ''.join(x.text or '' for x in p.iter(qn('w:t')))


def replace_in(p, old, new):
    t = ptext(p)
    assert old in t, (old, t[:80])
    set_text(p, t.replace(old, new))


def cell_set(tc, text):
    p = tc.find(qn('w:p'))
    if p.find('.//' + qn('w:t')) is None:      # ô trống → thêm run
        r = etree.SubElement(p, qn('w:r'))
        etree.SubElement(r, qn('w:t'))
    set_text(p, text)


def table_rows(tbl):
    return tbl.findall(qn('w:tr'))


def row_set(tr, vals):
    for tc, v in zip(tr.findall(qn('w:tc')), vals):
        cell_set(tc, v)


def rebuild_rows(tbl, rows, tpl_row_idx=1):
    """Thay toàn bộ dòng dữ liệu (giữ dòng tiêu đề) bằng `rows`, nhân bản dòng mẫu."""
    trs = table_rows(tbl)
    tpl = copy.deepcopy(trs[tpl_row_idx])
    for tr in trs[1:]:
        tbl.remove(tr)
    for vals in rows:
        tr = _strip_ids(copy.deepcopy(tpl))
        row_set(tr, vals)
        tbl.append(tr)


def new_par(key, segs=None, text=None):
    p = _strip_ids(copy.deepcopy(TPL[key]))
    if key == 111:                                   # H1 mới: sang trang bằng chính tiêu đề
        ppr = p.find(qn('w:pPr'))
        pb = etree.SubElement(ppr, qn('w:pageBreakBefore'))
        ppr.remove(pb); ppr.insert(1, pb)
    if text is not None:
        set_text(p, text)
        return p
    if segs is not None:
        runs = [r for r in p.findall(qn('w:r'))]
        rpr_src = None
        for r in runs:
            if r.find(qn('w:t')) is not None and rpr_src is None:
                rpr_src = r.find(qn('w:rPr'))
            p.remove(r)
        para = Paragraph(p, d._body)
        for s in segs:
            if isinstance(s, tuple) and s[0] == 'img':
                run = para.add_run()
                run.add_picture(s[1], height=Inches(s[2] if len(s) > 2 else 0.22))
            else:
                bold = isinstance(s, tuple) and s[0] == 'b'
                txt = s[1] if isinstance(s, tuple) else s
                r = etree.SubElement(p, qn('w:r'))
                if rpr_src is not None:
                    r.append(copy.deepcopy(rpr_src))
                if bold:
                    rpr = r.find(qn('w:rPr'))
                    if rpr is None:
                        rpr = etree.SubElement(r, qn('w:rPr'))
                    etree.SubElement(rpr, qn('w:b'))
                t = etree.SubElement(r, qn('w:t'))
                t.text = txt
                t.set('{http://www.w3.org/XML/1998/namespace}space', 'preserve')
    return p


def img_par(path, max_h_in=8.2):
    p = _strip_ids(copy.deepcopy(TPL[115]))
    for r in p.findall(qn('w:r')):
        p.remove(r)
    ppr = p.find(qn('w:pPr'))
    kn = etree.SubElement(ppr, qn('w:keepNext'))      # ảnh dính chú thích
    ppr.remove(kn); ppr.insert(0, kn)
    w, h = Image.open(path).size
    width = 6.0
    if width * h / w > max_h_in:
        width = max_h_in * w / h
    Paragraph(p, d._body).add_run().add_picture(path, width=Inches(width))
    return p


ICON = lambda n: ('img', I + n + '.png')


def caption(text):
    return new_par(116, text=text)


def H1(text):
    return new_par(111, text=text)


def H2(text):
    return new_par(114, text=text)


def STEP(*segs):
    return new_par(68, segs=list(segs))


def OR(*segs):
    """Dòng 'Hoặc …' thụt lề dưới bước."""
    p = new_par(112, segs=['Hoặc '] + list(segs))
    ppr = p.find(qn('w:pPr'))
    ind = etree.SubElement(ppr, qn('w:ind'))
    ind.set(qn('w:left'), '720')
    return p


def insert_before(ref, els):
    for e in els:
        ref.addprevious(e)


def replace_image(par, path, max_h_in=8.2, crop_h=None):
    blip = par.find('.//' + qn('a:blip'))
    rid = blip.get(qn('r:embed'))
    im = Image.open(path)
    if crop_h and im.height > crop_h:
        im = im.crop((0, 0, im.width, crop_h))
    buf = io.BytesIO(); im.save(buf, 'PNG')
    d.part.related_parts[rid]._blob = buf.getvalue()
    w, h = im.size
    cx = int(Inches(6.0))
    cy = int(cx * h / w)
    if cy > Inches(max_h_in):
        cy = int(Inches(max_h_in)); cx = int(cy * w / h)
    for ext in par.iter(qn('wp:extent')):
        ext.set('cx', str(cx)); ext.set('cy', str(cy))
    for ext in par.iter(qn('a:ext')):
        if ext.get('cx'):
            ext.set('cx', str(cx)); ext.set('cy', str(cy))
    ppr = par.find(qn('w:pPr'))
    if ppr.find(qn('w:keepNext')) is None:
        kn = etree.SubElement(ppr, qn('w:keepNext')); ppr.remove(kn); ppr.insert(0, kn)


# ================================================================ 1. SỬA CHỖ LỆCH (tại chỗ)
log = []

# --- Lịch sử cập nhật tài liệu: thêm dòng 1.1
tbl = B[28]
tr = _strip_ids(copy.deepcopy(table_rows(tbl)[1]))
row_set(tr, ['1.1', '02/10/2026',
             'Bổ sung các chức năng Tạo mới, Sửa, Xóa yêu cầu làm giải pháp và Cập nhật thông tin bổ sung, gửi lại yêu cầu; '
             'cập nhật đường vào menu, bố cục danh sách, bộ lọc, nút Từ chối, xuất Excel và toàn bộ ảnh màn hình theo phiên bản hiện tại',
             'Ban Dự án'])
tbl.append(tr)

# --- Giới thiệu chung
body.remove(B[32]); log.append('Bỏ dòng "Đường dẫn truy cập: https://…" (tài liệu không ghi URL).')
set_text(B[33], 'Vị trí trên menu: Phân hệ Công việc → Phê duyệt → Giải pháp - Dự án → Yêu cầu làm giải pháp')
log.append('Vị trí menu: "Dự án & Giao việc → Phê duyệt → Yêu cầu giải pháp" → "Công việc → Phê duyệt → Giải pháp - Dự án → Yêu cầu làm giải pháp".')

# --- Quyền
t = B[38]
r1 = table_rows(t)[1]
cell_set(r1.findall(qn('w:tc'))[1],
         'Điều kiện bắt buộc để: nhìn thấy mục menu "Phê duyệt → Giải pháp - Dự án → Yêu cầu làm giải pháp", mở được màn hình, '
         'nhận được thông báo khi có YC mới, bấm nút "Tiếp nhận", "Từ chối" và gửi "Yêu cầu bổ sung".')
tr = _strip_ids(copy.deepcopy(r1))
row_set(tr, ['(Không cần quyền riêng)',
             'Tạo mới yêu cầu làm giải pháp từ dự án TKT do chính mình tạo. Nút Sửa, Xóa, Hủy yêu cầu chỉ hiện với NGƯỜI TẠO yêu cầu, theo trạng thái (xem Phần 6, Phần 7).'])
t.append(tr)
log.append('Bảng quyền: thêm "Từ chối" vào quyền Tiếp nhận; thêm dòng quyền cho Tạo mới/Sửa/Xóa.')
set_text(B[46], 'Hình 1: Vị trí màn hình trên menu — Công việc → Phê duyệt → Giải pháp - Dự án → Yêu cầu làm giải pháp')

# --- Vòng đời
t = B[53]
for tr in table_rows(t):
    tcs = tr.findall(qn('w:tc'))
    if ptext(tcs[0]).startswith('Từ chối'):
        cell_set(tcs[1], 'Người tiếp nhận bấm "Từ chối" ở màn chi tiết (bắt buộc nhập lý do). Dự án TKT quay về bước "Thu thập thông tin dự án", KD nhận thông báo.')
        cell_set(tcs[2], 'Người tiếp nhận')
        tr2 = _strip_ids(copy.deepcopy(tr))
        row_set(tr2, ['Đã hủy', 'Người tạo hủy yêu cầu khi chưa được tiếp nhận (bắt buộc nhập lý do hủy). Dự án TKT quay về bước "Thu thập thông tin dự án".', 'KD (người tạo)'])
        tr.addnext(tr2)
log.append('Bảng vòng đời: "Từ chối" đã dùng thật (nút ở màn chi tiết); thêm trạng thái "Đã hủy".')
set_text(B[55], 'Lưu ý quan trọng: tại danh sách chờ duyệt, mỗi dòng chỉ có hai thao tác — "Tiếp nhận" (đồng ý làm) và "Yêu cầu bổ sung thông tin" '
                '(trả về KD để làm rõ). Nút "Từ chối" (bắt buộc nhập lý do) chỉ đặt ở màn chi tiết, để người tiếp nhận đọc kỹ hồ sơ trước khi bác bỏ. '
                'Hồ sơ chỉ thiếu thông tin thì dùng "Yêu cầu bổ sung" thay vì từ chối.')
log.append('Đoạn "màn hình này KHÔNG có nút Từ chối" → đã có nút Từ chối ở màn chi tiết.')
replace_in(B[63], ' (đúng như ví dụ trong ảnh chụp ở các phần sau)', '')

# --- PHẦN 2: truy cập + bố cục
set_text(B[69], 'Bước 2: Vào phân hệ "Công việc".')
set_text(B[70], 'Bước 3: Trên menu bên trái, bấm "Phê duyệt". Hệ thống mở bảng chức năng Phê duyệt.')
set_text(B[71], 'Bước 4: Chọn nhóm "Giải pháp - Dự án", bấm "Yêu cầu làm giải pháp". Hệ thống mở màn hình "Yêu cầu làm giải pháp chờ duyệt".')
log.append('PHẦN 2.1: các bước truy cập theo menu mới.')
set_text(B[76], 'Khối 1 — Bộ lọc danh sách: ô tìm nhanh, nút "Tìm kiếm", nút "Làm mới", nút "Cài đặt bộ lọc" (chọn ô lọc hiển thị) và nút "Tìm kiếm nâng cao" để mở/đóng phần lọc chi tiết.')
set_text(B[77], 'Khối 2 — Danh sách: thanh công cụ (nút "Xuất Excel" và nút "Cấu hình cột hiển thị" để ẩn/hiện, sắp xếp cột), bảng dữ liệu 19 cột có thanh cuộn ngang, và thanh phân trang ở dưới cùng.')
set_text(B[78], 'Hai cột đầu (STT và "Mã yêu cầu") được ghim cố định bên trái, nên khi cuộn ngang để xem các cột phía sau vẫn biết đang đọc dòng nào.')
log.append('PHẦN 2.2: thêm nút Cài đặt bộ lọc, Cấu hình cột; bảng 19 cột; cột ghim "Mã yêu cầu".')

# --- PHẦN 3
set_text(B[81], 'Ô tìm nhanh nằm ngay dưới tiêu đề bộ lọc, gợi ý "Tìm theo mã yêu cầu, tên yêu cầu".')
set_text(B[85], 'Bấm nút "Tìm kiếm nâng cao" ở góc phải khối bộ lọc để mở thêm các ô lọc. Bấm lần nữa (nút đổi chữ thành "Ẩn tìm kiếm nâng cao") để thu gọn. '
                'Nút "Cài đặt bộ lọc" cho phép bật/tắt và sắp xếp lại các ô lọc này.')
set_text(B[87], 'Hình 3: Bộ lọc nâng cao')
rebuild_rows(B[88], [
    ['Công ty / Phòng ban / Bộ phận', 'Danh sách chọn', 'Lọc theo đơn vị của nhân viên lập yêu cầu. Ô bị khoá theo phạm vi quyền xem của tài khoản.'],
    ['Nhân viên gửi yêu cầu', 'Danh sách chọn, có ô tìm kiếm, cho phép xóa lựa chọn', 'Chọn đúng nhân viên KD đã lập yêu cầu.'],
    ['Giai đoạn dự án', 'Danh sách chọn, có biểu tượng ⓘ xem mô tả giai đoạn', 'Lọc các yêu cầu có giai đoạn dự án đã chọn.'],
    ['Ngày tạo', 'Chọn khoảng ngày (từ – đến, dd/mm/yyyy)', 'Lọc các yêu cầu có ngày tạo nằm trong khoảng đã chọn (tính cả 2 đầu).'],
])
set_text(B[90], 'Các ô lọc trên tự động tìm kiếm ngay khi chọn — không cần bấm nút "Tìm kiếm". Có thể kết hợp nhiều ô lọc với ô tìm nhanh; các điều kiện được cộng dồn (VÀ).')
set_text(B[92], 'Nút "Làm mới": xóa toàn bộ điều kiện đang đặt (ô tìm nhanh và các ô lọc nâng cao), đưa sắp xếp về mặc định và tải lại danh sách từ trang 1.')
set_text(B[95], 'Các cột có biểu tượng sắp xếp (Mã yêu cầu, Tên yêu cầu, Ngày gửi YC, Ngày cần tiếp nhận YC, Ngày chốt GP, Ngày KH cần GP, Ngày tạo, Ngày cập nhật): '
                'bấm vào tiêu đề cột để đảo chiều tăng/giảm. Khi đổi sắp xếp, hệ thống tự quay về trang 1.')
set_text(B[96], 'Thanh phân trang dưới bảng: chọn số dòng/trang (5, 10, 20, 50, 100 — mặc định 10) và chuyển trang bằng các nút «, ‹, số trang, ›, ».')
log.append('PHẦN 3: bộ lọc nâng cao 4 ô (Công ty/Phòng ban/Bộ phận, Nhân viên gửi yêu cầu, Giai đoạn dự án, Ngày tạo); cột sắp xếp; số dòng/trang có 100.')

# --- PHẦN 4: cột
rebuild_rows(B[100], [
    ['1', 'STT', 'Số thứ tự dòng, tính liên tục theo trang đang xem.'],
    ['2', 'Mã yêu cầu', 'Mã yêu cầu (ví dụ TPE.YCP.TC.26.0912), chữ xanh — bấm vào để mở màn chi tiết.'],
    ['3', 'Tên yêu cầu', 'Tên yêu cầu; dòng phụ bên dưới là nhãn hạn xử lý có màu (xem 4.1).'],
    ['4', 'Dự án TKT', 'Mã - Tên dự án tiền khả thi mà yêu cầu này phục vụ.'],
    ['5', 'Khách hàng', 'Mã - Tên khách hàng của dự án.'],
    ['6', 'Giai đoạn dự án', 'Giai đoạn của dự án TKT, ví dụ "1.Giai đoạn nghiên cứu dự án".'],
    ['7', 'Mức độ ưu tiên', 'Nhãn màu theo mức ưu tiên của giai đoạn dự án (Thấp / Bình Thường / Cao…). Đây chính là yếu tố quyết định hạn tiếp nhận.'],
    ['8', 'Ngày gửi YC', 'Ngày và giờ KD bấm gửi yêu cầu (mốc bắt đầu tính hạn).'],
    ['9', 'Ngày cần tiếp nhận YC', 'Hạn phải xử lý yêu cầu, có cả ngày và giờ. Để trống nếu dự án chưa có mức ưu tiên.'],
    ['10', 'Ngày chốt GP', 'Ngày chốt giải pháp — chỉ có giá trị sau khi yêu cầu đã được tiếp nhận.'],
    ['11', 'Ngày KH cần GP', 'Ngày khách hàng cần có giải pháp (do KD nhập).'],
    ['12', 'Phòng KD', 'Phòng ban của nhân viên đã lập yêu cầu.'],
    ['13', 'Phòng tiếp nhận YC', 'Phòng ban được chỉ định xử lý yêu cầu.'],
    ['14', 'Người tạo', 'Nhân viên KD lập yêu cầu.'],
    ['15', 'Ngày tạo', 'Thời điểm tạo yêu cầu.'],
    ['16', 'Người cập nhật', 'Người sửa yêu cầu gần nhất.'],
    ['17', 'Ngày cập nhật', 'Thời điểm sửa gần nhất.'],
    ['18', 'Tiến trình YC', 'Trạng thái hiện tại của yêu cầu. Trên màn này luôn là "Chờ tiếp nhận".'],
    ['19', 'Hành động', 'Các nút thao tác trên dòng (xem 4.2).'],
])
set_text(B[102], '4.1. Nhãn hạn xử lý (dòng phụ ở cột Tên yêu cầu)')
set_text(B[107], '4.2. Các thao tác trên từng dòng')
rebuild_rows(B[108], [
    ['Mã yêu cầu (chữ xanh)', '—', 'Bấm để mở màn chi tiết yêu cầu (xem Phần 8). Có thể giữ Ctrl (hoặc Cmd) khi bấm để mở ở tab mới.'],
    ['Hình tờ giấy có dấu +', 'Yêu cầu bổ sung thông tin', 'Mở màn chi tiết và nhảy thẳng vào tab "Phiếu thu thập thông tin" để soạn câu hỏi bổ sung (xem Phần 10).'],
    ['Hình khay nhận', 'Tiếp nhận', 'Mở cửa sổ Tiếp nhận yêu cầu ngay tại danh sách (xem Phần 9).'],
])
log.append('PHẦN 4: bảng cột theo 19 cột hiện tại (tách Mã/Tên, thêm Người tạo/Ngày tạo/Người cập nhật/Ngày cập nhật/Hành động, bỏ "Trạng thái YC", "Người tiếp nhận YC", "Mã GP", "PM làm GP"); '
           'thao tác dòng còn 2 nút, bỏ nút "con mắt" — mở chi tiết bằng cách bấm Mã yêu cầu.')

# --- PHẦN 5 cũ → 8: Xem chi tiết
set_text(B[111], 'PHẦN 8: XEM CHI TIẾT YÊU CẦU')
set_text(B[112], 'Bấm vào Mã yêu cầu trên một dòng để mở màn "Chi tiết yêu cầu giải pháp". Toàn bộ thông tin ở màn này là CHỈ ĐỌC — người tiếp nhận không sửa được nội dung yêu cầu do KD lập.')
for k, old, new in [(114, '5.1.', '8.1.'), (120, '5.2.', '8.2.'), (124, '5.3.', '8.3.'), (128, '5.4.', '8.4.')]:
    replace_in(B[k], old, new)
set_text(B[119], 'Cuối màn có thanh nút cố định: "Tiếp nhận" và "Từ chối" (chỉ hiện khi bạn đủ điều kiện tiếp nhận); "Sửa", "Xóa", "Hủy yêu cầu làm giải pháp" '
                 '(chỉ hiện với người tạo yêu cầu, theo trạng thái — xem Phần 6, Phần 7); và nút "Quay lại" để trở về danh sách. Cuối trang có mục "Lịch sử" ghi lại các lần thay đổi của yêu cầu.')
replace_in(B[135], 'Phần 7', 'Phần 10')
log.append('PHẦN Xem chi tiết: mở bằng Mã yêu cầu; thanh nút cuối màn có thêm Từ chối/Sửa/Xóa/Hủy yêu cầu + mục Lịch sử.')

# --- PHẦN 6 cũ → 9: Tiếp nhận
set_text(B[139], 'PHẦN 9: TIẾP NHẬN YÊU CẦU')
for k, old, new in [(144, '6.1.', '9.1.'), (150, '6.2.', '9.2.'), (156, '6.3.', '9.3.'), (163, '6.4.', '9.4.')]:
    replace_in(B[k], old, new)
for tr in table_rows(B[148]):
    tcs = tr.findall(qn('w:tc'))
    if ptext(tcs[0]) == 'Mã - Tên dự án':
        cell_set(tcs[0], 'Tên dự án')
log.append('PHẦN Tiếp nhận: ô "Mã - Tên dự án" trên cửa sổ nay là "Tên dự án".')

# --- PHẦN 7 cũ → 10: Yêu cầu bổ sung
set_text(B[167], 'PHẦN 10: YÊU CẦU BỔ SUNG THÔNG TIN')
for k, old, new in [(168, '7.1.', '10.1.'), (170, '7.2.', '10.2.'), (180, '7.3.', '10.3.')]:
    replace_in(B[k], old, new)
replace_in(B[186], 'Sau khi KD trả lời và gửi lại', 'Sau khi KD trả lời và gửi lại (xem Phần 11)')

# --- PHẦN 8 cũ → 12: Xuất Excel
set_text(B[189], 'PHẦN 12: XUẤT EXCEL')
set_text(B[191], 'Bước 1: Đặt bộ lọc mong muốn (tìm nhanh, các ô lọc nâng cao).')
set_text(B[192], 'Bước 2: Bấm "Xuất Excel". Hệ thống mở cửa sổ "Chọn trường xuất file", tích sẵn các cột đang hiện trên bảng.')
set_text(B[193], 'Bước 3: Tích chọn/bỏ chọn trường, kéo biểu tượng ☰ để đổi vị trí cột, rồi bấm "Xuất file". Trình duyệt tải về tệp yeu_cau_lam_giai_phap_cho_duyet.xlsx và hệ thống hiện thông báo "Xuất Excel thành công".')
set_text(B[196], 'Gồm đúng các trường đã chọn (tối đa 18 trường), kèm dòng tiêu đề "Danh sách yêu cầu làm giải pháp chờ duyệt".')
set_text(B[197], 'Mã dự án TKT và Tên dự án TKT, Mã khách hàng và Tên khách hàng được tách thành các cột riêng.')
body.remove(B[198])
set_text(B[199], 'Dữ liệu xuất theo đúng phạm vi bạn được xem trên màn hình.')
log.append('PHẦN Xuất Excel: nay mở cửa sổ "Chọn trường xuất file", tệp .xlsx, tối đa 18 trường; bỏ mô tả cột gộp/cột "Cập nhật" cũ.')

# --- PHẦN 9 → 13, 10 → 14, 11 → 15
set_text(B[201], 'PHẦN 13: THÔNG BÁO TỰ ĐỘNG LIÊN QUAN')
set_text(B[209], 'PHẦN 14: SAU KHI TIẾP NHẬN — CÁC BƯỚC TIẾP THEO')
set_text(B[219], 'PHẦN 15: CÂU HỎI THƯỜNG GẶP')
for tr in table_rows(B[220]):
    tcs = tr.findall(qn('w:tc'))
    if ptext(tcs[0]).startswith('Không thấy mục'):
        cell_set(tcs[0], 'Không thấy mục "Phê duyệt → Giải pháp - Dự án → Yêu cầu làm giải pháp" trên menu')

# --- Hình cũ 4..12 → 7..15 + thay ảnh theo giao diện hiện tại
for k, old_n in [(116, 4), (122, 5), (126, 6), (130, 7), (137, 8), (146, 9), (155, 10), (174, 11), (178, 12)]:
    replace_in(B[k], 'Hình %d:' % old_n, 'Hình %d:' % (old_n + 3))
for k, f, crop in [(45, 'h01_menu', None), (73, 'h02_pending', None), (86, 'h03_filter', None),
                   (115, 'h04_detail_req', None), (121, 'h05_tab_tkt', 1700), (125, 'h06_tab_meet', None),
                   (129, 'h07_tab_phieu', None), (136, 'h08_history', None), (145, 'h09_receive', None),
                   (154, 'h10_pm', None), (173, 'h11_add_q', None), (177, 'h12_q_added', None)]:
    replace_image(B[k], S + f + '.png', crop_h=crop)
log.append('Thay toàn bộ 12 ảnh màn hình cũ bằng ảnh chụp giao diện hiện tại (cùng nội dung, cùng vị trí).')

# ================================================================ 2. CHÈN PHẦN MỚI
ACCESS = ['Bước 1: Từ phân hệ ', ICON('menu_cskh'), ', tại menu bên trái chọn ', ICON('menu_ycgp'),
          '. Hệ thống mở màn "Danh sách yêu cầu làm giải pháp".']

new5 = [
    H1('PHẦN 5: TẠO MỚI YÊU CẦU LÀM GIẢI PHÁP'),
    STEP(*ACCESS),
    STEP('Bước 2: Bấm ', ICON('btn_taomoi'), '. Hệ thống mở màn "Tạo yêu cầu giải pháp".'),
    OR('tại danh sách Dự án TKT, bấm ', ICON('btn_taoycgp_row'), ' (Tạo yêu cầu làm giải pháp) ở cột Hành động của dự án.'),
    img_par(S + 'n10_create.png'),
    caption('Hình 4: Màn Tạo yêu cầu giải pháp'),
    STEP('Bước 3: Chọn Dự án tiền khả thi. Hệ thống tự điền các thông tin kế thừa từ dự án.'),
    STEP('Bước 4: Nhập/chọn đầy đủ thông tin bắt buộc: Tên yêu cầu, Phòng tiếp nhận yêu cầu, Giai đoạn dự án, Ngày KH cần giải pháp, Ngày KH cần báo giá.'),
    STEP('Bước 5: Bấm ', ICON('btn_luunhap'), ' để lưu ở trạng thái Nháp.'),
    OR('bấm ', ICON('btn_luuvagui'), ', chọn ', ICON('btn_xacnhan'), ' tại hộp "Xác nhận lưu và gửi" để gửi sang phòng tiếp nhận (trạng thái Chờ tiếp nhận).'),
    OR('bấm ', ICON('btn_quaylai'), ' để thoát không lưu.'),
]
new6 = [
    H1('PHẦN 6: SỬA YÊU CẦU LÀM GIẢI PHÁP'),
    STEP(*ACCESS),
    STEP('Bước 2: Bấm ', ICON('btn_sua_row'), ' ở cột Hành động của yêu cầu đang ở trạng thái Nháp hoặc Yêu cầu bổ sung (chỉ người tạo yêu cầu thấy nút này).'),
    OR('bấm vào Mã yêu cầu ', ICON('link_ma_yc'), ' để mở màn chi tiết, rồi bấm ', ICON('btn_sua_footer'), '.'),
    img_par(S + 'n11_edit.png'),
    caption('Hình 5: Màn Sửa yêu cầu giải pháp'),
    STEP('Bước 3: Sửa thông tin cần thay đổi (quy tắc nhập như Tạo mới).'),
    STEP('Bước 4: Bấm ', ICON('btn_luunhap'), ' để lưu lại bản nháp (nút chỉ có khi yêu cầu đang Nháp).'),
    OR('bấm ', ICON('btn_luuvagui'), ', chọn ', ICON('btn_xacnhan'), ' để lưu và gửi đi.'),
    OR('bấm ', ICON('btn_quaylai'), ' để thoát không lưu.'),
]
new7 = [
    H1('PHẦN 7: XÓA YÊU CẦU LÀM GIẢI PHÁP'),
    STEP(*ACCESS),
    STEP('Bước 2: Bấm ', ICON('btn_xoa_row'), ' ở cột Hành động (chỉ hiện với yêu cầu ở trạng thái Nháp, do chính bạn tạo, chưa có giải pháp).'),
    OR('bấm vào Mã yêu cầu ', ICON('link_ma_yc'), ' để mở màn chi tiết, rồi bấm ', ICON('btn_xoa_footer'), '.'),
    STEP('Bước 3: Hệ thống hiện hộp xác nhận "Xác nhận xóa".'),
    img_par(S + 'n12_delete.png'),
    caption('Hình 6: Hộp xác nhận xóa yêu cầu làm giải pháp'),
    STEP('Bước 4: Bấm ', ICON('btn_confirm_xoa'), ' để xóa yêu cầu.'),
    OR('bấm ', ICON('btn_confirm_huy'), ' nếu bấm nhầm.'),
]
new11 = [
    H1('PHẦN 11: CẬP NHẬT THÔNG TIN BỔ SUNG VÀ GỬI LẠI YÊU CẦU'),
    STEP('Bước 1: Khi nhận thông báo "Thông báo yêu cầu bổ sung câu hỏi mới", từ phân hệ ', ICON('menu_cskh'),
         ' chọn ', ICON('menu_duantkt'), ', bấm vào Mã dự án ', ICON('link_ma_duan'), ' của dự án có yêu cầu bị trả về.'),
    STEP('Bước 2: Chọn thẻ ', ICON('tab_thuthap'), ', kéo xuống mục "Thông tin bổ sung".'),
    img_par(S + 'n16_update.png'),
    caption('Hình 16: Trả lời câu hỏi bổ sung trong phiếu thu thập thông tin của dự án TKT'),
    STEP('Bước 3: Nhập câu trả lời vào ô dưới từng câu hỏi bổ sung.'),
    STEP('Bước 4: Bấm ', ICON('btn_luuphieu'), '.'),
    STEP('Bước 5: Từ phân hệ ', ICON('menu_cskh'), ' chọn ', ICON('menu_ycgp'), ', bấm ', ICON('btn_sua_row'),
         ' ở dòng yêu cầu đang ở trạng thái "Yêu cầu bổ sung".'),
    STEP('Bước 6: Bổ sung thông tin yêu cầu nếu cần, bấm ', ICON('btn_luuvagui'), ', chọn ', ICON('btn_xacnhan'),
         '. Yêu cầu quay về trạng thái "Chờ tiếp nhận".'),
]

# chèn 5-6-7 trước đoạn ngắt trang đứng trước PHẦN 8 (Xem chi tiết)
insert_before(B[110], new5 + new6 + new7)
# chèn 11 trước đoạn ngắt trang đứng trước PHẦN 12 (Xuất Excel)
insert_before(B[188], new11)

d.save(OUT)

# ---------------------------------------------------------------- kiểm số Hình / PHẦN
doc = Document(OUT)
hs = [ptext(p) for p in doc.element.body.findall(qn('w:p')) if re.match(r'^Hình \d+:', ptext(p))]
ps = [ptext(p) for p in doc.element.body.findall(qn('w:p')) if re.match(r'^PHẦN \d+:', ptext(p))]
print('\n'.join(ps)); print(); print('\n'.join(hs))
nums = [int(re.match(r'^Hình (\d+)', h).group(1)) for h in hs]
assert nums == list(range(1, len(nums) + 1)), nums
with open(os.path.join(HERE, 'changes_log.txt'), 'w') as f:
    f.write('\n'.join('- ' + x for x in log))
print('OK', OUT)
