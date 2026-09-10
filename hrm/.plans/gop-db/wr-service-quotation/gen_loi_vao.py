# -*- coding: utf-8 -*-
"""
Sinh bảng tra CÁC LỐI VÀO của 5 màn luồng Dịch vụ (ERP -> HRM).

Nguồn số liệu — đọc trực tiếp, không suy đoán:
  · Mục menu ERP: TanPhatDev/resources/views/layouts/topmenubar.blade.php
  · Điều kiện lọc: searchByFilter() của model ERP + applyScope() của service HRM
  · Số bản ghi   : gọi API HRM thật bằng tài khoản quản trị (28/08/2026, DB gộp)
"""
import sys
try: sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception: pass

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

WB = Workbook()

TIEU_DE_FILL = PatternFill('solid', fgColor='4472C4')
NHOM_FILL = PatternFill('solid', fgColor='D6E4F0')
CANH_BAO_FILL = PatternFill('solid', fgColor='FFF2CC')
VIEN = Border(*[Side(style='thin', color='BFBFBF')] * 4)
F = 'Times New Roman'

def style_header(ws, row, cols):
    for c in range(1, cols + 1):
        o = ws.cell(row=row, column=c)
        o.font = Font(name=F, size=12, bold=True, color='FFFFFF')
        o.fill = TIEU_DE_FILL
        o.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        o.border = VIEN
    ws.row_dimensions[row].height = 34

def style_body(ws, r1, r2, cols):
    for r in range(r1, r2 + 1):
        for c in range(1, cols + 1):
            o = ws.cell(row=r, column=c)
            o.font = Font(name=F, size=12)
            o.alignment = Alignment(vertical='center', wrap_text=True)
            o.border = VIEN

# ═══════════════════════════════ SHEET 1 — Bảng tra lối vào ═══════════════════════════════
ws = WB.active
ws.title = 'Lối vào từng màn'

ws['A1'] = 'CÁC LỐI VÀO CỦA 5 MÀN LUỒNG DỊCH VỤ — vào từ menu nào, hiện ra gì'
ws.merge_cells('A1:H1')
ws['A1'].font = Font(name=F, size=15, bold=True, color='FFFFFF')
ws['A1'].fill = TIEU_DE_FILL
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 30

ws['A2'] = ('Cùng MỘT màn nhưng vào từ mục menu khác nhau thì danh sách hiện ra KHÁC HẲN — khác nhau ở đoạn tham số phía sau đường dẫn. '
            'Mở nhầm lối vào rồi kết luận "màn thiếu dữ liệu" là chuyện đã xảy ra. Số liệu đo ngày 28/08/2026 bằng tài khoản quản trị (DNS Admin) trên dữ liệu gộp.')
ws.merge_cells('A2:H2')
ws['A2'].font = Font(name=F, size=11, italic=True)
ws['A2'].fill = CANH_BAO_FILL
ws['A2'].alignment = Alignment(vertical='center', wrap_text=True)
ws.row_dimensions[2].height = 42

HEADER = ['Màn hình', 'Vào từ đâu (ERP)', 'Vào từ đâu (HRM)', 'Đường dẫn HRM',
          'Danh sách hiện ra những phiếu nào', 'Điều kiện lọc cụ thể',
          'Quản trị thấy bao nhiêu', 'Ghi chú']
for i, h in enumerate(HEADER, 1):
    ws.cell(row=4, column=i, value=h)
style_header(ws, 4, len(HEADER))

DU_LIEU = [
    # (man, erp, hrm, url, hien ra, dieu kien, so luong, ghi chu)
    ('1. Yêu cầu kiểm tra sửa chữa – bảo hành',
     'Khởi tạo → Lắp đặt - BH - SC → "Yêu cầu kiểm tra sửa chữa - bảo hành"',
     'Bán hàng → Bán dịch vụ → "Yêu cầu sửa chữa - bảo hành"',
     '/customer-care/warranty-repair-requests',
     'CHỈ phiếu do chính người đang đăng nhập lập (kể cả phiếu còn ở trạng thái Đang tạo)',
     'Người tạo = tôi',
     '6',
     'Đây là phạm vi MẶC ĐỊNH khi vào đường dẫn không kèm tham số. Không phụ thuộc quyền xem theo cấp.'),
    ('1. Yêu cầu kiểm tra sửa chữa – bảo hành',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Phiếu yêu cầu kiểm tra sửa chữa - bảo hành"',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Yêu cầu kiểm tra sửa chữa - bảo hành"',
     '/customer-care/warranty-repair-requests?type=all',
     'Toàn bộ phiếu trong phạm vi quyền của tôi, cộng thêm phiếu gửi về phòng tôi tiếp nhận',
     'Theo 3 cấp quyền (tổng công ty / công ty / phòng ban), HOẶC người tạo = tôi, HOẶC phòng tiếp nhận = phòng tôi. Phiếu Đang tạo của người khác luôn bị ẩn.',
     '5.371',
     'Lối vào chính của người phụ trách. Người KHÔNG có quyền xem theo cấp nào thì vào đây vẫn chỉ thấy phiếu của mình.'),
    ('1. Yêu cầu kiểm tra sửa chữa – bảo hành',
     'Khởi tạo → Lắp đặt - BH -SC → "Phiếu yêu cầu sửa chữa - bảo hành" (trang tổng quan)',
     '(HRM chưa có mục này)',
     '/customer-care/warranty-repair-requests?type=all',
     'Giống dòng ngay trên',
     'Như trên',
     '5.371',
     'ERP có 2 mục cùng trỏ ?type=all ở hai chỗ khác nhau. HRM gộp làm một.'),
    ('1. Yêu cầu kiểm tra sửa chữa – bảo hành',
     '(không có mục menu — link từ màn phiếu)',
     '(không có mục menu)',
     '/customer-care/warranty-repair-requests?type=waiting_handle',
     'Phiếu đang Chờ xử lý gửi về đúng phòng tôi tiếp nhận',
     'Trạng thái = Chờ xử lý VÀ phòng tiếp nhận = phòng tôi. Không có quyền "Xử lý yêu cầu sửa chữa" thì chỉ còn phiếu của mình.',
     '0',
     'Danh sách việc phòng tôi phải làm. Rỗng nghĩa là phòng không có phiếu nào đang chờ.'),

    ('2. Phiếu xử lý yêu cầu',
     '(không có mục menu vào route trần)',
     '(không có mục menu)',
     '/customer-care/warranty-repair-handle-requests',
     'CHỈ phiếu do chính tôi lập',
     'Người tạo = tôi',
     '5',
     'Phạm vi mặc định.'),
    ('2. Phiếu xử lý yêu cầu',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Phiếu xử lý yêu cầu"',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Phiếu xử lý yêu cầu"',
     '/customer-care/warranty-repair-handle-requests?type=all',
     'Toàn bộ phiếu trong phạm vi quyền của tôi',
     'Theo 3 cấp quyền, HOẶC người tạo = tôi. Phiếu Đang tạo của người khác bị ẩn.',
     '5.258',
     'Lối vào chính.'),
    ('2. Phiếu xử lý yêu cầu',
     '(không có mục menu)',
     '(không có mục menu)',
     '/customer-care/warranty-repair-handle-requests?type=waiting_information',
     'Phiếu đang Chờ cung cấp thông tin — việc cần tôi làm tiếp',
     'Trạng thái = Chờ cung cấp thông tin, và phải có quyền "Tạo phiếu cung cấp thông tin"; không có quyền thì chỉ còn phiếu của mình.',
     '22',
     'Hộp việc.'),

    ('3. Phiếu cung cấp thông tin làm báo giá',
     '(không có mục menu vào route trần)',
     '(không có mục menu)',
     '/customer-care/wr-information-requests',
     'CHỈ phiếu do chính tôi lập',
     'Người tạo = tôi',
     '1',
     'Phạm vi mặc định.'),
    ('3. Phiếu cung cấp thông tin làm báo giá',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Phiếu cung cấp thông tin làm báo giá"',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Phiếu cung cấp thông tin làm báo giá"',
     '/customer-care/wr-information-requests?permission=all',
     'Toàn bộ phiếu trong phạm vi quyền của tôi',
     'Theo 3 cấp quyền, HOẶC người tạo = tôi.',
     '4.980',
     '⚠️ Màn này ERP đặt tên tham số là permission (không phải type). HRM nhận CẢ HAI tên.'),
    ('3. Phiếu cung cấp thông tin làm báo giá',
     'Kinh doanh → Báo giá dịch vụ SC-BD-BT → "Phiếu cung cấp thông tin làm báo giá"',
     'Bán hàng → Bán dịch vụ → "Phiếu cung cấp thông tin làm báo giá"',
     '/customer-care/wr-information-requests?type=waiting_create_quotation',
     'Phiếu đang chờ CHÍNH TÔI làm báo giá',
     'Đủ 3 điều kiện cùng lúc: (1) trạng thái = Chờ làm báo giá; (2) chính tôi là người lập PHIẾU YÊU CẦU gốc; (3) phiếu có phần cần báo giá — ít nhất 1 thiết bị sửa chữa hoặc 1 thiết bị bảo dưỡng.',
     '0',
     'Hộp việc "chờ tôi báo giá". Toàn hệ thống chỉ 164/3.759 phiếu Chờ làm báo giá qua được điều kiện (3); phần còn lại chỉ có hàng bảo hành nên không phải báo giá.'),

    ('4. Báo giá dịch vụ',
     '(không có mục menu vào route trần)',
     '(không có mục menu)',
     '/customer-care/wr-quotations',
     'CHỈ báo giá do chính tôi lập',
     'Người tạo = tôi',
     '0',
     'Phạm vi mặc định.'),
    ('4. Báo giá dịch vụ',
     'Kinh doanh → Báo giá dịch vụ SC-BD-BT → "Danh sách báo giá"',
     'Bán hàng → Bán dịch vụ → "Danh sách báo giá"',
     '/customer-care/wr-quotations?type=all',
     'Toàn bộ báo giá trong phạm vi quyền của tôi',
     'Theo 3 cấp quyền, HOẶC người tạo = tôi. Báo giá Đang tạo của người khác bị ẩn.',
     '3.585',
     '⚠️ ERP dùng tên tham số permission ở màn này. HRM nhận cả type lẫn permission.'),

    ('5. Phiếu bảo hành',
     '(không có mục menu vào route trần)',
     '(không có mục menu)',
     '/customer-care/wr-warranties',
     'CHỈ phiếu do chính tôi lập',
     'Người tạo = tôi',
     '0',
     'Phiếu bảo hành sinh TỰ ĐỘNG từ phiếu cung cấp thông tin, nên "người tạo" là người lập phiếu đó.'),
    ('5. Phiếu bảo hành',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Phiếu bảo hành"',
     'CSKH → Kiểm tra bảo hành sửa chữa → "Phiếu bảo hành"',
     '/customer-care/wr-warranties?type=all',
     'Toàn bộ phiếu bảo hành trong phạm vi quyền của tôi',
     'Theo 3 cấp quyền, HOẶC người tạo = tôi.',
     '3.631',
     'Lối vào chính.'),
]

r = 5
man_truoc = None
for row in DU_LIEU:
    if row[0] != man_truoc:      # dải phân cách giữa các màn
        ws.cell(row=r, column=1, value=row[0])
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=len(HEADER))
        o = ws.cell(row=r, column=1)
        o.font = Font(name=F, size=12, bold=True, color='1F4E79')
        o.fill = NHOM_FILL
        o.alignment = Alignment(vertical='center')
        o.border = VIEN
        ws.row_dimensions[r].height = 24
        man_truoc = row[0]
        r += 1
    for i, v in enumerate(row[1:], 2):
        ws.cell(row=r, column=i, value=v)
    ws.cell(row=r, column=1, value='')
    r += 1

style_body(ws, 5, r - 1, len(HEADER))
for c, w in zip('ABCDEFGH', [4, 46, 40, 46, 40, 52, 15, 46]):
    ws.column_dimensions[c].width = w
ws.freeze_panes = None   # KHÔNG đóng băng — user cuộn tự do

# ═══════════════════════════════ SHEET 2 — 5 điều cần nhớ ═══════════════════════════════
ws2 = WB.create_sheet('Điều cần nhớ')
ws2['A1'] = 'NĂM ĐIỀU DỄ HIỂU NHẦM'
ws2.merge_cells('A1:C1')
ws2['A1'].font = Font(name=F, size=15, bold=True, color='FFFFFF')
ws2['A1'].fill = TIEU_DE_FILL
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 30

for i, h in enumerate(['#', 'Điều dễ hiểu nhầm', 'Thực tế'], 1):
    ws2.cell(row=3, column=i, value=h)
style_header(ws2, 3, 3)

GHI_NHO = [
    ('Vào màn thấy ít phiếu là hệ thống lỗi / mất dữ liệu',
     'Nhiều khả năng đang vào từ lối "chỉ phiếu của mình". Kiểm tra lại bạn bấm từ menu nào — mỗi lối vào hiển thị một phạm vi khác nhau.'),
    ('Đường dẫn kèm ?type=all thì ai mở cũng thấy hết',
     'Không. Đường dẫn chỉ quyết định CÁCH XEM, không cấp quyền. Người không có quyền xem theo cấp nào mà mở link này thì vẫn chỉ thấy phiếu của chính mình.'),
    ('Bấm "Làm mới" sẽ đưa về danh sách đầy đủ',
     '"Làm mới" chỉ xoá các điều kiện lọc đang áp dụng, KHÔNG đổi phạm vi đang xem. Muốn đổi phạm vi thì bấm lại mục menu tương ứng.'),
    ('Sửa tay tham số trên thanh địa chỉ để xem thêm dữ liệu',
     'Giá trị lạ sẽ bị bỏ qua, hệ thống quay về phạm vi mặc định. Không lỗi, cũng không lộ thêm dữ liệu.'),
    ('Hộp việc rỗng nghĩa là chức năng hỏng',
     'Các lối vào dạng hộp việc (chờ xử lý / chờ cung cấp thông tin / chờ làm báo giá) lọc rất chặt theo người và theo trạng thái. Rỗng thường chỉ có nghĩa là hiện không có việc nào đến lượt bạn.'),
]
r = 4
for i, (a, b) in enumerate(GHI_NHO, 1):
    ws2.cell(row=r, column=1, value=i)
    ws2.cell(row=r, column=2, value=a)
    ws2.cell(row=r, column=3, value=b)
    r += 1
style_body(ws2, 4, r - 1, 3)
for c, w in zip('ABC', [5, 52, 92]):
    ws2.column_dimensions[c].width = w
for rr in range(4, r):
    ws2.row_dimensions[rr].height = 46

OUT = 'Loi vao man hinh - luong Dich vu.xlsx'
WB.save(OUT)
print('Da tao:', OUT)
print('  Sheet 1: %d dong du lieu' % len(DU_LIEU))
print('  Sheet 2: %d dieu can nho' % len(GHI_NHO))
