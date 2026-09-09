# -*- coding: utf-8 -*-
"""Sinh file Excel giai thich cach tinh VAT o bang "III - Tong hop bao gia".

Man: Phieu cung cap thong tin lam bao gia (/customer-care/wr-information-requests/{id})
Nguon: hrm-client/utils/wrServiceQuotationMoney.js — grandTotal() / warrantyTotals() /
       repairMaintainTotals() / costTotals() / lineTotals() / workTotals()

Chay:  python3 .plans/gop-db/wr-service-quotation/gen_giai_thich_vat.py
"""
import os
import sys

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "Cách tính VAT - Tổng hợp báo giá.xlsx")

FONT = "Times New Roman"
HEAD_FILL = PatternFill("solid", fgColor="4472C4")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")
IN_FILL = PatternFill("solid", fgColor="E2EFDA")      # o cho tester tu nhap
SUM_FILL = PatternFill("solid", fgColor="DDEBF7")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_C = Alignment(wrap_text=True, vertical="center", horizontal="center")
RIGHT = Alignment(horizontal="right", vertical="center")


def put(ws, r, c, v, bold=False, size=12, fill=None, align=WRAP, border=True,
        color=None, fmt=None):
    cell = ws.cell(r, c, v)
    cell.font = Font(name=FONT, size=size, bold=bold,
                     color=color or ("FFFFFF" if fill is HEAD_FILL else "000000"))
    cell.alignment = align
    if fill:
        cell.fill = fill
    if border:
        cell.border = BOX
    if fmt:
        cell.number_format = fmt
    return cell


def title(ws, r, text, cols):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=cols)
    put(ws, r, 1, text, bold=True, size=14, align=Alignment(vertical="center"), border=False)
    ws.row_dimensions[r].height = 24


def header(ws, r, labels, widths):
    for i, (lab, w) in enumerate(zip(labels, widths), start=1):
        put(ws, r, i, lab, bold=True, fill=HEAD_FILL, align=WRAP_C)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[r].height = 32


wb = Workbook()

# ============================================================ SHEET 1
ws = wb.active
ws.title = "1. Cách tính"
ws.sheet_view.showGridLines = False

title(ws, 1, "Bảng “III - Tổng hợp báo giá” — dòng VAT được tính thế nào", 4)
ws.merge_cells("A2:D2")
put(ws, 2, 1, "Màn hình: Phiếu cung cấp thông tin làm báo giá  ·  Khối D, bảng III  ·  "
              "Cập nhật 03/09/2026", size=11, border=False,
    align=Alignment(vertical="center"))

ws.merge_cells("A4:D4")
put(ws, 4, 1,
    "Ô VAT KHÔNG phải là một phép nhân duy nhất lên tổng tiền. Nó là TỔNG VAT CỦA 4 NHÓM, "
    "và mỗi nhóm lấy thuế suất ở một chỗ khác nhau:",
    bold=True, fill=NOTE_FILL)
ws.row_dimensions[4].height = 34

header(ws, 6, ["Nhóm tiền", "Lấy % VAT ở đâu", "Nhân với số nào", "Điểm dễ hiểu nhầm"],
       [26, 34, 40, 42])

rows = [
    ("1. Bảo hành\n(bảng I của khối D)",
     "% VAT khai ở ĐẦU PHIẾU — một mức duy nhất cho cả nhóm này.",
     "(Tổng thành tiền của các hạng mục bảo hành − Chi phí được bảo hành) "
     "+ phần bảo hành của các dòng chi phí.",
     "Đổi ô VAT ở đầu phiếu chỉ làm đổi nhóm này, không ảnh hưởng nhóm Sửa chữa – Bảo dưỡng."),
    ("2. Sửa chữa – Bảo dưỡng\n(bảng II của khối D)",
     "% VAT của TỪNG DÒNG: dòng công, dòng dịch vụ, dòng vật tư, dòng gói bảo dưỡng — "
     "mỗi dòng một mức riêng lấy từ danh mục lúc lập phiếu.",
     "Thành tiền của dòng SAU KHI trừ chiết khấu %.",
     "Hai dòng cùng một thiết bị vẫn có thể chịu hai mức VAT khác nhau."),
    ("3. Các khoản chi phí liên quan",
     "% VAT của từng dòng chi phí. ĐỂ TRỐNG thì hệ thống hiểu là 8%, KHÔNG phải 0%.",
     "Số tiền sửa chữa của dòng chi phí đó.",
     "Chi phí không có chiết khấu, nên số trước và sau chiết khấu bằng nhau."),
    ("4. Chi phí vận chuyển thiết bị vật tư",
     "Như nhóm 3 — để trống cũng là 8%.",
     "Số tiền sửa chữa của dòng vận chuyển.",
     "Tách riêng khỏi nhóm 3 trên giao diện nhưng cộng chung vào cùng một ô VAT."),
]
r = 7
for row in rows:
    for i, v in enumerate(row, start=1):
        put(ws, r, i, v)
    ws.row_dimensions[r].height = 62
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
put(ws, r, 1, "Ba dòng cuối của bảng III khớp nhau như sau:", bold=True, border=False)
r += 1
for lab, val in [
    ("Thành tiền (cột Phải thanh toán)", "Tổng tiền phải trả TRƯỚC thuế của cả 4 nhóm."),
    ("VAT", "Tổng VAT của 4 nhóm ở bảng trên."),
    ("Tổng thanh toán", "Thành tiền + VAT."),
]:
    put(ws, r, 1, lab, bold=True)
    ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    put(ws, r, 2, val)
    ws.row_dimensions[r].height = 22
    r += 1

r += 1
ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
put(ws, r, 1,
    "Lưu ý khi kiểm thử: số ở bảng III được TÍNH LẠI ngay trên trình duyệt từ các dòng chi tiết "
    "mỗi lần mở phiếu, không đọc từ số tiền đã lưu. Sửa đơn giá, số lượng, % chiết khấu hay % VAT "
    "của bất kỳ dòng nào thì bảng tổng đổi theo ngay, chưa cần bấm Lưu.",
    fill=NOTE_FILL)
ws.row_dimensions[r].height = 46

# ============================================================ SHEET 2
ws2 = wb.create_sheet("2. Ví dụ có số")
ws2.sheet_view.showGridLines = False

title(ws2, 1, "Ví dụ minh hoạ — ô nền xanh là số tự nhập, các ô còn lại tự tính", 6)
ws2.merge_cells("A2:F2")
put(ws2, 2, 1, "Tester đổi số ở cột nền xanh để đối chiếu với phiếu thật.", size=11,
    border=False, align=Alignment(vertical="center"))

put(ws2, 4, 1, "% VAT khai ở đầu phiếu", bold=True)
ws2.merge_cells("B4:C4")
put(ws2, 4, 2, 10, fill=IN_FILL, align=RIGHT, fmt="0")
ws2.column_dimensions["A"].width = 34

header(ws2, 6, ["Nhóm / dòng", "Thành tiền", "Chiết khấu (%)", "Sau chiết khấu",
                "% VAT áp dụng", "Tiền VAT"], [34, 16, 16, 18, 16, 18])

MONEY = '#,##0'
r = 7
# --- nhom 1: bao hanh
put(ws2, r, 1, "1. Bảo hành — hạng mục A", bold=True)
put(ws2, r, 2, 5000000, fill=IN_FILL, align=RIGHT, fmt=MONEY)
put(ws2, r, 3, "Chi phí được bảo hành:", align=RIGHT)
put(ws2, r, 4, 5000000, fill=IN_FILL, align=RIGHT, fmt=MONEY)
put(ws2, r, 5, "=$B$4", align=RIGHT, fmt="0")
put(ws2, r, 6, "=(B7-D7)*E7/100", align=RIGHT, fmt=MONEY)
ws2.row_dimensions[r].height = 20

r = 8
put(ws2, r, 1, "2. Sửa chữa – dòng công", bold=True)
put(ws2, r, 2, 2000000, fill=IN_FILL, align=RIGHT, fmt=MONEY)
put(ws2, r, 3, 10, fill=IN_FILL, align=RIGHT, fmt="0")
put(ws2, r, 4, "=B8-B8*C8/100", align=RIGHT, fmt=MONEY)
put(ws2, r, 5, 8, fill=IN_FILL, align=RIGHT, fmt="0")
put(ws2, r, 6, "=D8*E8/100", align=RIGHT, fmt=MONEY)

r = 9
put(ws2, r, 1, "2. Sửa chữa – dòng vật tư", bold=True)
put(ws2, r, 2, 3000000, fill=IN_FILL, align=RIGHT, fmt=MONEY)
put(ws2, r, 3, 0, fill=IN_FILL, align=RIGHT, fmt="0")
put(ws2, r, 4, "=B9-B9*C9/100", align=RIGHT, fmt=MONEY)
put(ws2, r, 5, 10, fill=IN_FILL, align=RIGHT, fmt="0")
put(ws2, r, 6, "=D9*E9/100", align=RIGHT, fmt=MONEY)

r = 10
put(ws2, r, 1, "3. Chi phí liên quan — để TRỐNG ô VAT", bold=True)
put(ws2, r, 2, 1000000, fill=IN_FILL, align=RIGHT, fmt=MONEY)
put(ws2, r, 3, "không có chiết khấu", align=RIGHT)
put(ws2, r, 4, "=B10", align=RIGHT, fmt=MONEY)
put(ws2, r, 5, 8, align=RIGHT, fmt="0")
put(ws2, r, 6, "=D10*E10/100", align=RIGHT, fmt=MONEY)

r = 11
put(ws2, r, 1, "4. Chi phí vận chuyển", bold=True)
put(ws2, r, 2, 500000, fill=IN_FILL, align=RIGHT, fmt=MONEY)
put(ws2, r, 3, "không có chiết khấu", align=RIGHT)
put(ws2, r, 4, "=B11", align=RIGHT, fmt=MONEY)
put(ws2, r, 5, 8, fill=IN_FILL, align=RIGHT, fmt="0")
put(ws2, r, 6, "=D11*E11/100", align=RIGHT, fmt=MONEY)

for i in (7, 8, 9, 10, 11):
    ws2.row_dimensions[i].height = 20

r = 13
put(ws2, r, 1, "Thành tiền (trước thuế)", bold=True, fill=SUM_FILL)
put(ws2, r, 2, "", fill=SUM_FILL)
put(ws2, r, 3, "", fill=SUM_FILL)
put(ws2, r, 4, "=SUM(D8:D11)+(B7-D7)", bold=True, fill=SUM_FILL, align=RIGHT, fmt=MONEY)
put(ws2, r, 5, "", fill=SUM_FILL)
put(ws2, r, 6, "", fill=SUM_FILL)

r = 14
put(ws2, r, 1, "VAT (dòng VAT của bảng III)", bold=True, fill=SUM_FILL)
for c in (2, 3, 4, 5):
    put(ws2, r, c, "", fill=SUM_FILL)
put(ws2, r, 6, "=SUM(F7:F11)", bold=True, fill=SUM_FILL, align=RIGHT, fmt=MONEY)

r = 15
put(ws2, r, 1, "Tổng thanh toán", bold=True, fill=SUM_FILL)
for c in (2, 3, 4, 5):
    put(ws2, r, c, "", fill=SUM_FILL)
put(ws2, r, 6, "=D13+F14", bold=True, fill=SUM_FILL, align=RIGHT, fmt=MONEY)

r = 17
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
put(ws2, r, 1,
    "Dòng 1 (Bảo hành): chiết khấu KHÔNG nhập theo %, mà là số tiền gõ ở cột “Chi phí được bảo "
    "hành”. Ví dụ trên để bằng đúng thành tiền nên phần bảo hành còn 0 đồng và VAT của nhóm này "
    "cũng bằng 0 — đúng thực tế các phiếu đang chạy.",
    fill=NOTE_FILL)
ws2.row_dimensions[r].height = 46

r = 19
ws2.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
put(ws2, r, 1,
    "Dòng 3 để trống ô VAT nhưng vẫn tính 8% — đây là điểm khác biệt hay bị báo là lỗi. Hành vi "
    "này bê nguyên từ phần mềm ERP.",
    fill=NOTE_FILL)
ws2.row_dimensions[r].height = 32

# ============================================================ SHEET 3
ws3 = wb.create_sheet("3. Cách đối chiếu")
ws3.sheet_view.showGridLines = False

title(ws3, 1, "Cách tự đối chiếu một phiếu thật", 3)
header(ws3, 3, ["Bước", "Làm gì", "Kết quả mong đợi"], [10, 56, 56])

steps = [
    ("1", "Mở phiếu, ghi lại % VAT khai ở đầu phiếu.",
     "Đây là mức áp cho riêng nhóm Bảo hành."),
    ("2", "Ở bảng II (Sửa chữa – Bảo dưỡng), cộng cột “Phải thanh toán” của từng dòng rồi nhân "
          "với % VAT của chính dòng đó.",
     "Ra phần VAT của nhóm 2. Lưu ý mỗi dòng một mức VAT riêng."),
    ("3", "Ở hai bảng chi phí, lấy số tiền sửa chữa nhân 8% (hoặc % ghi trên dòng nếu có).",
     "Ra phần VAT của nhóm 3 và 4."),
    ("4", "Ở bảng I (Bảo hành), lấy (Thành tiền − Chi phí được bảo hành) nhân % VAT của phiếu.",
     "Ra phần VAT của nhóm 1; thường bằng 0 vì hai số này bằng nhau."),
    ("5", "Cộng bốn phần vừa tính.",
     "Phải bằng đúng ô VAT ở bảng III. Lệch thì chụp màn hình cả 3 bảng gửi lại để soi dòng nào."),
]
r = 4
for s in steps:
    for i, v in enumerate(s, start=1):
        put(ws3, r, i, v, align=WRAP_C if i == 1 else WRAP)
    ws3.row_dimensions[r].height = 46
    r += 1

r += 1
ws3.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
put(ws3, r, 1,
    "Ba chỗ hay làm lệch số khi đối chiếu tay: (1) lấy nhầm % VAT của phiếu áp cho nhóm Sửa chữa; "
    "(2) coi ô VAT để trống của dòng chi phí là 0% thay vì 8%; (3) nhân VAT lên số TRƯỚC chiết "
    "khấu thay vì sau chiết khấu.",
    fill=NOTE_FILL)
ws3.row_dimensions[r].height = 46

wb.save(OUT)
print("Da tao:", OUT)
