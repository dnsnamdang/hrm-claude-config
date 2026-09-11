# -*- coding: utf-8 -*-
"""Sinh file testcase cho Danh mục Lĩnh vực Công ty kinh doanh.

Chạy: python3 docs/srs/generate_testcase_lvctkd.py  (từ thư mục HRM/)
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Linh vuc Cong ty kinh doanh"

# === STYLES ===
thin_border = Border(
    left=Side(style='thin', color='FF000000'),
    right=Side(style='thin', color='FF000000'),
    top=Side(style='thin', color='FF000000'),
    bottom=Side(style='thin', color='FF000000'),
)
title_font = Font(bold=True, size=14, color='FF1F4E79')
header_font = Font(bold=True, size=11, color='FFFFFFFF')
header_fill = PatternFill(start_color='FF4472C4', end_color='FF4472C4', fill_type='solid')
header_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
section_font = Font(bold=True, size=11, color='FF1F4E79')
section_fill = PatternFill(start_color='FFD6E4F0', end_color='FFD6E4F0', fill_type='solid')
data_font = Font(size=11, color='FF000000')
data_alignment = Alignment(vertical='top', wrap_text=True)
summary_font = Font(size=11, color='FF000000')

col_widths = {'A': 14, 'B': 26, 'C': 18, 'D': 44, 'E': 10,
              'F': 34, 'G': 58, 'H': 26, 'I': 12, 'J': 62,
              'K': 15, 'L': 14, 'M': 18}
for c, w in col_widths.items():
    ws.column_dimensions[c].width = w

ws.merge_cells('A1:E1')
ws['A1'] = 'Testcase _ Danh mục Lĩnh vực Công ty kinh doanh (Assign / Danh mục)'
ws['A1'].font = title_font
ws.merge_cells('F1:I1')
ws['F1'] = 'TEST SUMMARY'
ws['F1'].font = summary_font

for i, (label, formula) in enumerate([
    ('Số trường hợp kiểm thử đạt (P):', '=COUNTIF(L8:L500,"Passed")'),
    ('Số trường hợp kiểm thử không đạt (F):', '=COUNTIF(L8:L500,"Failed")'),
    ('Số trường hợp kiểm thử đang xem xét (PE):', '=COUNTIF(L8:L500,"Pending")'),
    ('Số trường hợp kiểm thử chưa thực hiện:', '=COUNTIF(L8:L500,"Not Executed")'),
    ('Tổng số trường hợp kiểm thử:', '=COUNTA(L8:L500)'),
]):
    ws.cell(row=i + 1, column=10, value=label).font = summary_font
    ws.cell(row=i + 1, column=11, value=formula).font = summary_font

headers = ['Module', 'Nhóm chức năng', 'TC ID', 'Chức năng', 'Priority',
           'Tiền điều kiện', 'Bước thực hiện', 'Test Data', 'Test Data',
           'Expected Result (chi tiết)', 'KQ thực tế', 'Status', 'Ghi chú']
ws.row_dimensions[6].height = 30
for idx, h in enumerate(headers, 1):
    cell = ws.cell(row=6, column=idx, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = header_alignment
    cell.border = thin_border

current_row = 7
MODULE = 'Dự án & Giao việc'
GROUP = 'DM Lĩnh vực Công ty KD'
PREFIX = 'LVCTKD'
_counter = {}


def section(title):
    global current_row
    ws.merge_cells(f'C{current_row}:M{current_row}')
    cell = ws.cell(row=current_row, column=3, value=title)
    cell.font = section_font
    cell.fill = section_fill
    cell.border = thin_border
    current_row += 1


def tc(sec, func, priority, precondition, steps, test_data, expected, note=''):
    """Tự đánh số thứ tự trong section."""
    global current_row
    _counter[sec] = _counter.get(sec, 0) + 1
    tc_id = f'{PREFIX}_{sec:03d}.{_counter[sec]:03d}'
    values = [MODULE, GROUP, tc_id, func, priority,
              precondition, steps, test_data, '', expected, '', 'Not Executed', note]
    for idx, val in enumerate(values, 1):
        cell = ws.cell(row=current_row, column=idx, value=val)
        cell.font = data_font
        cell.alignment = data_alignment
        cell.border = thin_border
    current_row += 1


# ======================================================================
# I. MENU & PHÂN QUYỀN
# ======================================================================
section('I. MENU & PHÂN QUYỀN')

tc(1, 'Hiển thị mục menu với quyền Quản lý', 'P0',
   'User có quyền "Quản lý danh mục lĩnh vực Công ty kinh doanh"',
   '1. Đăng nhập\n2. Mở menu trái\n3. Vào nhóm "Danh mục"',
   '',
   'Thấy mục "Lĩnh vực Công ty kinh doanh"\nMục nằm NGAY TRƯỚC mục "Nhóm ngành"\nClick vào điều hướng tới /assign/internal-business-scopes')

tc(1, 'Hiển thị mục menu với quyền Xem', 'P0',
   'User CHỈ có quyền "Xem danh mục lĩnh vực Công ty kinh doanh"',
   '1. Đăng nhập\n2. Mở nhóm "Danh mục"',
   '',
   'Vẫn thấy mục "Lĩnh vực Công ty kinh doanh"\nVào được màn danh sách')

tc(1, 'Ẩn mục menu khi KHÔNG có quyền nào', 'P0',
   'User không có cả 2 quyền (dùng tài khoản fixture e2e_nocatalog@test.local)',
   '1. Đăng nhập\n2. Mở nhóm "Danh mục"',
   '',
   'KHÔNG thấy mục "Lĩnh vực Công ty kinh doanh" trong menu')

tc(1, 'Chặn truy cập trực tiếp URL khi không có quyền', 'P0',
   'User không có cả 2 quyền',
   '1. Gõ thẳng URL /assign/internal-business-scopes lên thanh địa chỉ\n2. Enter',
   '',
   'API danh sách trả 403\nBảng rỗng\nKHÔNG hiện toast lỗi (tránh nhiễu)')

tc(1, 'Ẩn nút Tạo mới / Import với quyền Xem', 'P0',
   'User CHỈ có quyền Xem',
   '1. Vào màn danh sách\n2. Quan sát khu vực nút phía trên bảng',
   '',
   'KHÔNG có nút "Tạo mới"\nKHÔNG có nút "Import Excel"\nVẪN có nút "Xuất Excel"\nNút bị ẩn HẲN, không phải disable')

tc(1, 'Ẩn toàn bộ nút Hành động với quyền Xem', 'P0',
   'User CHỈ có quyền Xem, danh sách có ≥ 1 bản ghi Hoạt động',
   '1. Vào màn danh sách\n2. Quan sát cột "Hành động"',
   '',
   'Cột Hành động trống hoàn toàn: không có Sửa / Khoá / Mở khoá / Xoá')

tc(1, 'Quyền Xem vẫn mở được modal Xem chi tiết', 'P1',
   'User CHỈ có quyền Xem',
   '1. Vào màn danh sách\n2. Click vào giá trị cột "Mã"',
   '',
   'Mở modal "Xem chi tiết lĩnh vực Công ty kinh doanh"\nMọi ô chỉ đọc\nFooter chỉ có nút Đóng')

tc(1, 'Cờ quyền fail-closed khi request quyền lỗi', 'P1',
   'Giả lập API quyền trả lỗi / chậm',
   '1. Vào màn danh sách khi store quyền chưa có dữ liệu\n2. Quan sát nút',
   '',
   'canManage = false → các nút quản lý bị ẩn\nKhông có trường hợp nút hiện nhầm rồi mất')

tc(1, 'Thu hồi quyền giữa chừng', 'P2',
   'User đang mở màn, admin gỡ quyền ở tab khác',
   '1. Ở màn danh sách\n2. Admin gỡ quyền\n3. Bấm Tìm kiếm để gọi lại API',
   '',
   'API trả 403\nBảng rỗng\nKHÔNG hiện toast lỗi')

# ======================================================================
# II. MÀN DANH SÁCH — HIỂN THỊ
# ======================================================================
section('II. MÀN DANH SÁCH — HIỂN THỊ')

tc(2, 'Hiển thị đủ 9 cột đúng thứ tự', 'P0',
   'User có quyền Quản lý, có ≥ 1 bản ghi',
   '1. Vào /assign/internal-business-scopes\n2. Quan sát header bảng',
   '',
   'Đúng 9 cột theo thứ tự: STT · Mã · Tên lĩnh vực Công ty kinh doanh · Người tạo · Ngày tạo · Người cập nhật · Ngày cập nhật · Trạng thái · Hành động\nCột Hành động ở CUỐI cùng')

tc(2, 'Tiêu đề trang & tiêu đề bảng', 'P1',
   'Đang ở màn danh sách',
   '1. Quan sát tiêu đề trang (tab trình duyệt) và tiêu đề bảng',
   '',
   'Tab trình duyệt: "Danh sách lĩnh vực Công ty kinh doanh"\nTiêu đề bảng: "Danh sách lĩnh vực Công ty kinh doanh"')

tc(2, 'Cột Mã là link mở modal Xem', 'P0',
   'Có ≥ 1 bản ghi',
   '1. Click vào giá trị cột "Mã" của dòng bất kỳ',
   'LVKDNB.OTO',
   'Mở modal "Xem chi tiết lĩnh vực Công ty kinh doanh" đúng bản ghi vừa click\nKHÔNG có nút "Xem" riêng ở cột Hành động')

tc(2, 'Badge trạng thái Hoạt động', 'P0',
   'Có bản ghi status = 1',
   '1. Quan sát cột Trạng thái',
   '',
   'Hiện badge "Hoạt động" (variant brand — màu chủ đạo)')

tc(2, 'Badge trạng thái Khoá', 'P0',
   'Có bản ghi status = 2',
   '1. Quan sát cột Trạng thái',
   '',
   'Hiện badge "Khoá" (variant required — màu đỏ)')

tc(2, 'Định dạng ngày KHÔNG có giây', 'P0',
   'Có ≥ 1 bản ghi',
   '1. Quan sát cột Ngày tạo, Ngày cập nhật',
   '',
   'Hiển thị dạng dd/mm/yyyy HH:mm (VD 22/08/2026 14:30)\nKHÔNG hiển thị giây')

tc(2, 'Cột Người tạo/Người cập nhật chỉ hiện TÊN', 'P0',
   'Có ≥ 1 bản ghi',
   '1. Quan sát cột Người tạo và Người cập nhật',
   '',
   'Chỉ hiện họ tên nhân viên\nKHÔNG kèm mã nhân viên, không kèm email')

tc(2, 'Ô trống hiển thị dấu gạch ngang', 'P1',
   'Bản ghi có updated_by = NULL (VD tạo bằng seeder)',
   '1. Quan sát cột Người cập nhật của bản ghi đó',
   '',
   'Hiện ký tự "—" thay vì để trống')

tc(2, 'Chữ trong ô để thường, không in đậm', 'P2',
   'Có ≥ 1 bản ghi',
   '1. Quan sát cột Tên',
   '',
   'Chữ hiển thị thường (font-weight normal), không bold')

tc(2, 'Bảng loading ngay khi vào màn', 'P1',
   'Mạng chậm (throttle Slow 3G)',
   '1. Vào màn danh sách\n2. Quan sát ngay lập tức',
   '',
   'Spinner hiện NGAY, không chờ request quyền xong\nRequest danh sách là request bắn ĐẦU TIÊN')

tc(2, 'Thông báo khi không có dữ liệu', 'P1',
   'Bộ lọc không khớp bản ghi nào',
   '1. Nhập từ khoá không tồn tại "zzzzzz"\n2. Bấm Tìm kiếm',
   'zzzzzz',
   'Bảng hiện "Không có dữ liệu phù hợp bộ lọc."\nKhông có dòng dữ liệu nào')

tc(2, 'Toast khi API danh sách lỗi 500', 'P2',
   'Giả lập API trả 500',
   '1. Vào màn danh sách',
   '',
   'Toast đỏ "Lỗi khi tải dữ liệu"\nBảng rỗng, spinner tắt')

tc(2, 'Sắp xếp mặc định id DESC', 'P0',
   'Có ≥ 3 bản ghi tạo ở các thời điểm khác nhau',
   '1. Vào màn danh sách (chưa lọc, chưa sort)',
   '',
   'Bản ghi tạo MỚI NHẤT nằm ở dòng đầu tiên')

# ======================================================================
# III. TÌM KIẾM & LỌC
# ======================================================================
section('III. TÌM KIẾM & LỌC')

tc(3, 'Tìm nhanh theo Mã', 'P0',
   'Có bản ghi LVKDNB.OTO — Ô tô',
   '1. Nhập "OTO" vào ô tìm nhanh\n2. Bấm Tìm kiếm',
   'OTO',
   'Danh sách chỉ còn bản ghi có mã chứa "OTO"\nVề trang 1')

tc(3, 'Tìm nhanh theo Tên', 'P0',
   'Có bản ghi tên "Ô tô"',
   '1. Nhập "Ô tô"\n2. Bấm Tìm kiếm',
   'Ô tô',
   'Trả về bản ghi có tên chứa "Ô tô"')

tc(3, 'Tìm nhanh theo TÊN NGƯỜI TẠO', 'P0',
   'Bản ghi do "Nguyễn Văn A" tạo',
   '1. Nhập "Nguyễn Văn A"\n2. Bấm Tìm kiếm',
   'Nguyễn Văn A',
   'Trả về các bản ghi do người đó tạo\n(BE tìm qua employee_infos.fullname bằng EXISTS)')

tc(3, 'Placeholder ô tìm nhanh', 'P2',
   'Đang ở màn danh sách',
   '1. Quan sát ô tìm nhanh',
   '',
   'Placeholder: "Tìm theo mã, tên lĩnh vực Công ty kinh doanh, người tạo"')

tc(3, 'Ô tìm nhanh KHÔNG auto-search', 'P1',
   'Đang ở màn danh sách',
   '1. Gõ ký tự vào ô tìm nhanh\n2. Đợi 3 giây, KHÔNG bấm Tìm kiếm\n3. Quan sát Network',
   '',
   'KHÔNG có request danh sách nào được gọi\nDanh sách giữ nguyên')

tc(3, 'Tìm kiếm bằng phím Enter', 'P1',
   'Đang ở màn danh sách',
   '1. Gõ từ khoá vào ô tìm nhanh\n2. Nhấn Enter',
   'OTO',
   'Gọi API tìm kiếm, kết quả lọc đúng')

tc(3, 'Tìm không dấu vẫn khớp', 'P1',
   'Có bản ghi tên "Ô tô"',
   '1. Nhập "o to"\n2. Bấm Tìm kiếm',
   'o to',
   'Vẫn khớp bản ghi "Ô tô" (collation _ci bỏ dấu)')

tc(3, 'Sắp theo độ khớp — khớp đúng dấu lên trước', 'P1',
   'Có 2 bản ghi: "Ô tô" và "Oto điện"',
   '1. Nhập "Ô tô"\n2. Bấm Tìm kiếm\n3. Quan sát thứ tự',
   'Ô tô',
   'Bản ghi khớp ĐÚNG DẤU xếp TRƯỚC bản ghi chỉ khớp nhờ bỏ dấu')

tc(3, 'Sắp theo độ khớp — trùng khít lên đầu', 'P1',
   'Có bản ghi "Ô tô" và "Ô tô điện" và "Phụ tùng Ô tô"',
   '1. Nhập "Ô tô"\n2. Bấm Tìm kiếm',
   'Ô tô',
   'Thứ tự: trùng khít → bắt đầu bằng → khớp đầu từ → chỉ chứa\nCùng mức thì tên NGẮN hơn lên trước')

tc(3, 'Không sắp độ khớp khi từ khoá < 2 ký tự', 'P2',
   'Có nhiều bản ghi',
   '1. Nhập "O" (1 ký tự)\n2. Bấm Tìm kiếm',
   'O',
   'Kết quả sắp theo id DESC (không áp thuật toán độ khớp)')

tc(3, 'Đã bấm sort cột thì KHÔNG sắp theo độ khớp', 'P2',
   'Có nhiều bản ghi',
   '1. Bấm sort cột Mã tăng dần\n2. Nhập từ khoá\n3. Bấm Tìm kiếm',
   'OTO',
   'Kết quả sắp theo cột Mã tăng dần, không theo độ khớp')

tc(3, 'Mở panel lọc nâng cao', 'P0',
   'Đang ở màn danh sách (panel mặc định thu gọn)',
   '1. Bấm nút mở rộng bộ lọc',
   '',
   'Hiện đủ 7 ô: Mã · Tên lĩnh vực Công ty kinh doanh · Trạng thái · Người tạo · Người cập nhật · Cập nhật từ · Cập nhật đến')

tc(3, 'Lọc theo Mã (auto-search)', 'P0',
   'Panel nâng cao đang mở',
   '1. Nhập "LVKDNB.O" vào ô Mã\n2. KHÔNG bấm nút nào',
   'LVKDNB.O',
   'Danh sách TỰ nạp lại (auto-search)\nVề trang 1\nChỉ còn bản ghi có mã chứa chuỗi đó')

tc(3, 'Lọc theo Tên (auto-search)', 'P0',
   'Panel nâng cao đang mở',
   '1. Nhập "Ô tô" vào ô Tên',
   'Ô tô',
   'Danh sách tự lọc theo tên')

tc(3, 'Lọc theo Trạng thái = Hoạt động', 'P0',
   'Có cả bản ghi Hoạt động và Khoá',
   '1. Chọn Trạng thái = "Hoạt động"',
   'Hoạt động',
   'Chỉ hiện bản ghi status = 1')

tc(3, 'Lọc theo Trạng thái = Khoá', 'P0',
   'Có cả bản ghi Hoạt động và Khoá',
   '1. Chọn Trạng thái = "Khoá"',
   'Khoá',
   'Chỉ hiện bản ghi status = 2')

tc(3, 'Xoá điều kiện Trạng thái', 'P1',
   'Đang lọc Trạng thái = Khoá',
   '1. Bấm dấu X trên ô Trạng thái',
   '',
   'Điều kiện được gỡ, danh sách nạp lại đầy đủ')

tc(3, 'Lọc theo Người tạo', 'P0',
   'Panel nâng cao đang mở',
   '1. Chọn 1 nhân viên ở ô Người tạo',
   'Nguyễn Văn A',
   'Chỉ hiện bản ghi do nhân viên đó tạo')

tc(3, 'Lọc theo Người cập nhật', 'P1',
   'Panel nâng cao đang mở',
   '1. Chọn 1 nhân viên ở ô Người cập nhật',
   'Nguyễn Văn A',
   'Chỉ hiện bản ghi do nhân viên đó cập nhật gần nhất')

tc(3, 'Lọc khoảng ngày cập nhật', 'P0',
   'Có bản ghi cập nhật ở nhiều ngày khác nhau',
   '1. Chọn "Cập nhật từ" = 01/08/2026\n2. Chọn "Cập nhật đến" = 22/08/2026',
   'từ 01/08/2026 đến 22/08/2026',
   'Chỉ hiện bản ghi có updated_at nằm trong khoảng (so theo NGÀY, bao gồm 2 đầu mút)')

tc(3, 'Lọc chỉ có "Cập nhật từ"', 'P1',
   'Panel nâng cao đang mở',
   '1. Chọn "Cập nhật từ" = 22/08/2026, để trống ô đến',
   '',
   'Hiện bản ghi cập nhật từ ngày đó trở về sau')

tc(3, 'Kết hợp nhiều điều kiện lọc', 'P0',
   'Panel nâng cao đang mở',
   '1. Nhập Tên = "Ô"\n2. Chọn Trạng thái = Hoạt động\n3. Chọn Người tạo',
   '',
   'Các điều kiện áp ĐỒNG THỜI (AND), kết quả thoả cả 3')

tc(3, 'Nút Làm mới xoá hết điều kiện VÀ nạp lại', 'P0',
   'Đang có từ khoá tìm nhanh + vài ô lọc nâng cao',
   '1. Bấm nút "Làm mới"\n2. Quan sát Network + bảng',
   '',
   'Mọi ô lọc + ô tìm nhanh về rỗng\nVề trang 1\nGọi API danh sách ĐÚNG 1 LẦN (không 0 lần, không 2 lần)\nBảng hiện lại đầy đủ dữ liệu')

tc(3, 'Lọc xong về trang 1', 'P1',
   'Đang ở trang 3 của danh sách',
   '1. Nhập điều kiện lọc bất kỳ',
   '',
   'Danh sách nhảy về trang 1')

# ======================================================================
# IV. SẮP XẾP & PHÂN TRANG
# ======================================================================
section('IV. SẮP XẾP & PHÂN TRANG')

tc(4, 'Sort theo Mã tăng/giảm', 'P0',
   'Có ≥ 3 bản ghi',
   '1. Click header cột "Mã" (lần 1)\n2. Click lần 2',
   '',
   'Lần 1: sắp tăng dần theo mã\nLần 2: sắp giảm dần\nIcon sort đổi chiều tương ứng')

tc(4, 'Sort theo Tên', 'P0',
   'Có ≥ 3 bản ghi',
   '1. Click header cột "Tên lĩnh vực Công ty kinh doanh"',
   '',
   'Danh sách sắp theo tên A→Z')

tc(4, 'Sort theo Ngày tạo', 'P1',
   'Có ≥ 3 bản ghi',
   '1. Click header cột "Ngày tạo"',
   '',
   'Danh sách sắp theo created_at')

tc(4, 'Sort theo Ngày cập nhật', 'P1',
   'Có ≥ 3 bản ghi',
   '1. Click header cột "Ngày cập nhật"',
   '',
   'Danh sách sắp theo updated_at')

tc(4, 'Các cột KHÔNG cho sort', 'P2',
   'Có ≥ 1 bản ghi',
   '1. Thử click header STT, Người tạo, Người cập nhật, Trạng thái, Hành động',
   '',
   'Không có icon sort, click không đổi thứ tự')

tc(4, 'Sort bằng tham số ngoài whitelist', 'P1',
   'Có ≥ 2 bản ghi',
   '1. Gọi API với sort_by=abc (dùng Postman hoặc sửa URL)',
   'sort_by=abc',
   'Bỏ qua tham số lạ, kết quả rơi về id DESC, KHÔNG lỗi 500')

tc(4, 'Đổi số dòng/trang', 'P0',
   'Có > 10 bản ghi',
   '1. Đổi page size từ 10 sang 25',
   '25',
   'Bảng hiện tối đa 25 dòng\nVề trang 1')

tc(4, 'Chuyển trang', 'P0',
   'Có > 10 bản ghi',
   '1. Bấm sang trang 2',
   '',
   'Hiện tập bản ghi kế tiếp\nCột STT đánh số tiếp nối (11, 12, ...)')

tc(4, 'STT đánh số đúng theo trang', 'P1',
   'Có > 10 bản ghi, page size = 10',
   '1. Sang trang 2\n2. Quan sát cột STT dòng đầu',
   '',
   'STT dòng đầu trang 2 = 11 (không reset về 1)')

tc(4, 'Race condition khi bấm nhanh 2 lượt lọc', 'P2',
   'Mạng chậm',
   '1. Nhập điều kiện A\n2. Ngay lập tức nhập điều kiện B\n3. Đợi cả 2 response về',
   '',
   'Bảng hiển thị kết quả của lượt gọi CUỐI CÙNG (B)\nKhông bị response cũ ghi đè (loadSeq)')

# ======================================================================
# V. TẠO MỚI — VALIDATE Ô MÃ
# ======================================================================
section('V. TẠO MỚI — VALIDATE Ô MÃ')

tc(5, 'Mở modal Tạo mới', 'P0',
   'User có quyền Quản lý',
   '1. Bấm nút "Tạo mới"',
   '',
   'Mở modal tiêu đề "Tạo mới lĩnh vực Công ty kinh doanh"\nÔ Mã đã có sẵn tiền tố "LVKDNB."\nÔ Tên rỗng\nTrạng thái mặc định "Hoạt động"\nFooter: Lưu · Lưu & Tiếp tục · Đóng')

tc(5, 'Bố cục 1 hàng đủ 12 cột', 'P2',
   'Modal Tạo mới đang mở',
   '1. Quan sát bố cục form',
   '',
   '1 hàng gồm 3 ô: Mã (3 cột) · Tên (6 cột) · Trạng thái (3 cột)\nKhông có ô nào nằm lẻ 1 dòng')

tc(5, 'Tiền tố LVKDNB. không xoá được', 'P0',
   'Modal Tạo mới đang mở',
   '1. Click vào ô Mã\n2. Nhấn Backspace nhiều lần',
   '',
   'Tiền tố "LVKDNB." luôn còn nguyên, không xoá được')

tc(5, 'Giới hạn hậu tố 4 ký tự trên UI', 'P0',
   'Modal Tạo mới đang mở',
   '1. Gõ "ABCDEFGH" vào phần hậu tố',
   'ABCDEFGH',
   'Ô chỉ nhận 4 ký tự đầu "ABCD" (maxlength = 4)')

tc(5, 'Để trống hậu tố → báo Bắt buộc phải nhập', 'P0',
   'Modal Tạo mới đang mở',
   '1. Để ô Mã chỉ có "LVKDNB."\n2. Nhập Tên hợp lệ\n3. Bấm Lưu',
   'LVKDNB.',
   'Lỗi dưới ô Mã: "Bắt buộc phải nhập."\nÔ Mã viền đỏ\nKHÔNG gọi API\nKHÔNG hiện câu lỗi định dạng')

tc(5, 'Hậu tố có ký tự đặc biệt', 'P0',
   'Modal Tạo mới đang mở',
   '1. Nhập hậu tố "A@#"\n2. Nhập Tên hợp lệ\n3. Bấm Lưu',
   'LVKDNB.A@#',
   'Lỗi: "Hậu tố tối đa 4 ký tự, chỉ gồm chữ không dấu (A-Z), số (0-9) và dấu gạch dưới (_)."')

tc(5, 'Hậu tố có dấu tiếng Việt', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập hậu tố "Ôtô"\n2. Bấm Lưu',
   'LVKDNB.Ôtô',
   'Báo lỗi định dạng, không cho lưu')

tc(5, 'Hậu tố chấp nhận dấu gạch dưới', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập hậu tố "A_1"\n2. Nhập Tên hợp lệ\n3. Bấm Lưu',
   'LVKDNB.A_1',
   'Lưu thành công, mã trong danh sách là LVKDNB.A_1')

tc(5, 'Hậu tố toàn số', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập hậu tố "1234"\n2. Nhập Tên hợp lệ\n3. Bấm Lưu',
   'LVKDNB.1234',
   'Lưu thành công')

tc(5, 'Hậu tố chữ thường tự viết HOA', 'P0',
   'Modal Tạo mới đang mở',
   '1. Nhập hậu tố "oto"\n2. Nhập Tên "Ô tô test"\n3. Bấm Lưu\n4. Xem lại danh sách',
   'oto',
   'Bản ghi lưu với mã "LVKDNB.OTO" (viết HOA)')

tc(5, 'Hậu tố 1 ký tự (biên dưới)', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập hậu tố "A"\n2. Nhập Tên hợp lệ\n3. Bấm Lưu',
   'LVKDNB.A',
   'Lưu thành công (1 ký tự là hợp lệ)')

tc(5, 'Hậu tố 4 ký tự (biên trên)', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập hậu tố "ABCD"\n2. Nhập Tên hợp lệ\n3. Bấm Lưu',
   'LVKDNB.ABCD',
   'Lưu thành công')

tc(5, 'Mã trùng bản ghi đã có', 'P0',
   'Đã tồn tại LVKDNB.OTO',
   '1. Tạo mới với mã LVKDNB.OTO\n2. Nhập Tên khác\n3. Bấm Lưu',
   'LVKDNB.OTO',
   'API trả 422\nLỗi dưới ô Mã: "Mã lĩnh vực Công ty kinh doanh đã tồn tại"\nModal KHÔNG đóng\nToast đỏ "Bạn chưa nhập đầy đủ thông tin"')

tc(5, 'Mã trùng khác hoa thường', 'P1',
   'Đã tồn tại LVKDNB.OTO',
   '1. Tạo mới với hậu tố "oto"\n2. Bấm Lưu',
   'LVKDNB.oto',
   'Vẫn báo trùng (do BE tự upper trước khi kiểm)')

tc(5, 'Khoảng trắng ngay sau dấu chấm bị loại', 'P2',
   'Modal Tạo mới đang mở (hoặc gọi API trực tiếp)',
   '1. Gửi code = "LVKDNB. OTO"',
   'LVKDNB. OTO',
   'BE chuẩn hoá thành LVKDNB.OTO trước khi validate')

# ======================================================================
# VI. TẠO MỚI — VALIDATE Ô TÊN
# ======================================================================
section('VI. TẠO MỚI — VALIDATE Ô TÊN')

tc(6, 'Để trống Tên', 'P0',
   'Modal Tạo mới đang mở',
   '1. Nhập Mã hợp lệ\n2. Để trống Tên\n3. Bấm Lưu',
   '',
   'Lỗi dưới ô Tên: "Bắt buộc phải nhập"\nÔ Tên viền đỏ\nKHÔNG gọi API')

tc(6, 'Tên chỉ gồm khoảng trắng', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập Mã hợp lệ\n2. Nhập Tên = "   " (3 dấu cách)\n3. Bấm Lưu',
   '"   "',
   'BE trim → rỗng → báo "Bắt buộc phải nhập"')

tc(6, 'Tên 255 ký tự (biên trên)', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập Tên đúng 255 ký tự\n2. Bấm Lưu',
   'Chuỗi 255 ký tự',
   'Lưu thành công')

tc(6, 'Tên 256 ký tự (vượt biên)', 'P0',
   'Modal Tạo mới đang mở',
   '1. Nhập Tên 256 ký tự\n2. Bấm Lưu',
   'Chuỗi 256 ký tự',
   'Lỗi: "Tên lĩnh vực Công ty kinh doanh tối đa 255 ký tự"')

tc(6, 'Tên chứa dấu phẩy', 'P0',
   'Modal Tạo mới đang mở',
   '1. Nhập Tên = "Ô tô, xe máy"\n2. Bấm Lưu',
   'Ô tô, xe máy',
   'Lỗi: "Tên không được chứa ký tự dấu phẩy (,) và dấu hai chấm (:)"')

tc(6, 'Tên chứa dấu hai chấm', 'P0',
   'Modal Tạo mới đang mở',
   '1. Nhập Tên = "Ô tô: điện"\n2. Bấm Lưu',
   'Ô tô: điện',
   'Lỗi: "Tên không được chứa ký tự dấu phẩy (,) và dấu hai chấm (:)"')

tc(6, 'Tên trùng bản ghi đã có', 'P0',
   'Đã tồn tại bản ghi tên "Ô tô"',
   '1. Tạo mới với Mã khác nhưng Tên = "Ô tô"\n2. Bấm Lưu',
   'Ô tô',
   'API trả 422\nLỗi dưới ô Tên: "Tên lĩnh vực Công ty kinh doanh đã tồn tại"\nModal KHÔNG đóng')

tc(6, 'Tên trùng khác hoa thường', 'P1',
   'Đã tồn tại bản ghi tên "Ô tô"',
   '1. Tạo mới với Tên = "ô tô"\n2. Bấm Lưu',
   'ô tô',
   'Vẫn báo trùng (MySQL collation _ci)')

tc(6, 'Tên có khoảng trắng đầu/cuối bị trim', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập Tên = "  Ô tô mới  "\n2. Bấm Lưu\n3. Xem lại danh sách',
   '"  Ô tô mới  "',
   'Bản ghi lưu với tên "Ô tô mới" (đã trim 2 đầu)')

tc(6, 'Tên có ký tự Unicode / emoji', 'P2',
   'Modal Tạo mới đang mở',
   '1. Nhập Tên = "Ô tô 🚗"\n2. Bấm Lưu',
   'Ô tô 🚗',
   'Lưu được (không có rule cấm), hiển thị đúng ở danh sách')

# ======================================================================
# VII. TẠO MỚI — VALIDATE ĐỒNG THỜI & LƯU
# ======================================================================
section('VII. TẠO MỚI — VALIDATE ĐỒNG THỜI & LƯU')

tc(7, 'Bấm Lưu khi form TRỐNG → 2 lỗi đồng thời', 'P0',
   'Modal Tạo mới vừa mở, chưa nhập gì',
   '1. Bấm Lưu ngay',
   '',
   'HIỆN CÙNG LÚC 2 lỗi:\n- Dưới ô Mã: "Bắt buộc phải nhập."\n- Dưới ô Tên: "Bắt buộc phải nhập"\nCẢ 2 ô đều viền đỏ\nCon trỏ focus vào ô Mã (ô lỗi đầu tiên)\nKHÔNG gọi API')

tc(7, 'Sai cả Mã lẫn Tên → báo cùng lúc', 'P0',
   'Modal Tạo mới đang mở',
   '1. Nhập Mã = "LVKDNB.@@"\n2. Nhập Tên = "Ô tô, xe"\n3. Bấm Lưu',
   'LVKDNB.@@ / Ô tô, xe',
   'Hiện đồng thời lỗi định dạng Mã + lỗi ký tự cấm ở Tên\nKhông phải sửa xong ô này mới lòi lỗi ô kia')

tc(7, 'Focus ô lỗi đầu tiên khi chỉ Tên sai', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập Mã hợp lệ\n2. Để trống Tên\n3. Bấm Lưu',
   '',
   'Focus nhảy vào ô Tên')

tc(7, 'Lỗi BE 422 tự mất khi sửa lại ô đó', 'P0',
   'Vừa bị lỗi 422 "Tên đã tồn tại"',
   '1. Sửa lại ô Tên thành giá trị khác\n2. Quan sát',
   '',
   'Lỗi 422 cũ dưới ô Tên BIẾN MẤT ngay khi gõ\nKhông che mất validate realtime')

tc(7, 'Lỗi BE hiện đúng dưới từng ô', 'P0',
   'Đã có LVKDNB.OTO tên "Ô tô"',
   '1. Tạo mới với Mã = LVKDNB.OTO và Tên = "Ô tô"\n2. Bấm Lưu',
   'LVKDNB.OTO / Ô tô',
   'Hiện ĐỒNG THỜI 2 lỗi 422: "Mã ... đã tồn tại" dưới ô Mã, "Tên ... đã tồn tại" dưới ô Tên')

tc(7, 'Lưu thành công', 'P0',
   'Modal Tạo mới đang mở, dữ liệu hợp lệ chưa trùng',
   '1. Nhập Mã = LVKDNB.TEST\n2. Nhập Tên = "Lĩnh vực test"\n3. Bấm Lưu',
   'LVKDNB.TEST / Lĩnh vực test',
   'Toast xanh "Thêm mới thành công"\nModal ĐÓNG\nDanh sách nạp lại, bản ghi mới ở ĐẦU danh sách\nTrạng thái = Hoạt động')

tc(7, 'Metadata tự ghi khi tạo', 'P0',
   'Vừa tạo bản ghi mới',
   '1. Quan sát cột Người tạo / Ngày tạo của bản ghi vừa tạo\n2. Kiểm SQL: select code, created_by, updated_by from internal_business_scopes',
   '',
   'Người tạo = tên tài khoản đang đăng nhập\nNgày tạo = thời điểm hiện tại, format dd/mm/yyyy HH:mm\ncreated_by và updated_by KHÁC NULL')

tc(7, 'Chọn Trạng thái = Khoá khi tạo mới', 'P1',
   'Modal Tạo mới đang mở',
   '1. Nhập Mã, Tên hợp lệ\n2. Đổi Trạng thái sang "Khoá"\n3. Bấm Lưu',
   'status = Khoá',
   'Bản ghi tạo ra với trạng thái Khoá\nDòng đó chỉ còn nút "Mở khoá"')

tc(7, 'Lưu & Tiếp tục giữ modal mở', 'P0',
   'Modal Tạo mới đang mở, dữ liệu hợp lệ',
   '1. Nhập Mã + Tên\n2. Bấm "Lưu & Tiếp tục"',
   '',
   'Toast "Thêm mới thành công"\nModal VẪN MỞ\nForm reset về trống (Mã chỉ còn tiền tố, Tên rỗng, Trạng thái Hoạt động)\nDanh sách nền đã nạp lại')

tc(7, 'Nút Lưu & Tiếp tục chỉ có ở chế độ Tạo', 'P1',
   'Có bản ghi để sửa',
   '1. Mở modal Sửa\n2. Quan sát footer',
   '',
   'Footer chỉ có "Lưu" và "Đóng"\nKHÔNG có "Lưu & Tiếp tục"')

tc(7, 'Đóng modal không lưu', 'P1',
   'Modal Tạo mới, đã nhập dữ liệu',
   '1. Bấm "Đóng"\n2. Mở lại modal Tạo mới',
   '',
   'Modal đóng, KHÔNG tạo bản ghi\nMở lại thì form đã reset sạch, không còn dữ liệu cũ và lỗi cũ')

tc(7, 'Đóng modal bằng dấu X', 'P2',
   'Modal Tạo mới đang mở',
   '1. Bấm dấu X góc phải header',
   '',
   'Modal đóng, không tạo bản ghi, form được reset')

tc(7, 'Chặn double-submit', 'P1',
   'Modal Tạo mới, dữ liệu hợp lệ, mạng chậm',
   '1. Bấm Lưu\n2. Bấm Lưu lần 2 ngay lập tức',
   '',
   'Nút Lưu bị disable trong lúc đang gửi\nChỉ tạo ra ĐÚNG 1 bản ghi')

# ======================================================================
# VIII. SỬA
# ======================================================================
section('VIII. SỬA')

tc(8, 'Mở modal Sửa', 'P0',
   'Có bản ghi Hoạt động, user có quyền Quản lý',
   '1. Bấm biểu tượng Sửa trên dòng',
   '',
   'Gọi GET /{id} nạp dữ liệu mới nhất\nModal tiêu đề "Sửa lĩnh vực Công ty kinh doanh"\nMã + Tên + Trạng thái điền đúng giá trị hiện tại')

tc(8, 'Header hiện chip người/ngày cập nhật', 'P1',
   'Bản ghi đã từng được cập nhật',
   '1. Mở modal Sửa\n2. Quan sát header',
   '',
   'Header hiện chip metadata: người cập nhật + ngày cập nhật gần nhất')

tc(8, 'Cuối body hiện Người tạo / Ngày tạo', 'P1',
   'Bản ghi bất kỳ',
   '1. Mở modal Sửa\n2. Cuộn xuống cuối body',
   '',
   'Hiện block metadata: Người tạo + Ngày tạo\nKHÔNG có ô nhập cho 2 trường này')

tc(8, 'Sửa Tên thành công', 'P0',
   'Bản ghi Hoạt động',
   '1. Mở modal Sửa\n2. Đổi Tên thành "Ô tô sửa"\n3. Bấm Lưu',
   'Ô tô sửa',
   'Toast "Cập nhật thành công"\nModal đóng\nDanh sách hiện tên mới')

tc(8, 'Sửa Mã thành công', 'P0',
   'Bản ghi Hoạt động',
   '1. Mở modal Sửa\n2. Đổi hậu tố mã\n3. Bấm Lưu',
   'LVKDNB.NEW',
   'Lưu thành công, danh sách hiện mã mới')

tc(8, 'Metadata cập nhật sau khi Sửa', 'P0',
   'Bản ghi do user A tạo, đăng nhập bằng user B có quyền Quản lý',
   '1. Sửa bản ghi\n2. Về danh sách quan sát cột Người/Ngày cập nhật\n3. Kiểm SQL updated_by',
   '',
   'Người cập nhật = user B\nNgày cập nhật = thời điểm vừa sửa\nNgười tạo VẪN là user A (không đổi)')

tc(8, 'Không cho sửa trùng Mã bản ghi khác', 'P0',
   'Có 2 bản ghi A và B',
   '1. Mở Sửa bản ghi A\n2. Đổi mã A thành mã của B\n3. Bấm Lưu',
   '',
   '422 "Mã lĩnh vực Công ty kinh doanh đã tồn tại", modal không đóng')

tc(8, 'Không cho sửa trùng Tên bản ghi khác', 'P0',
   'Có 2 bản ghi A và B',
   '1. Mở Sửa A\n2. Đổi tên A thành tên của B\n3. Bấm Lưu',
   '',
   '422 "Tên lĩnh vực Công ty kinh doanh đã tồn tại"')

tc(8, 'Giữ nguyên Mã/Tên của CHÍNH bản ghi đó', 'P0',
   'Bản ghi A',
   '1. Mở Sửa A\n2. KHÔNG đổi gì, bấm Lưu',
   '',
   'Lưu thành công, KHÔNG báo trùng với chính nó (unique có loại trừ id)')

tc(8, 'Nút Sửa bị ẩn khi bản ghi đã Khoá', 'P0',
   'Có bản ghi status = 2',
   '1. Quan sát cột Hành động của dòng đã Khoá',
   '',
   'KHÔNG có nút Sửa (is_can_edit = false)\nChỉ còn nút "Mở khoá"')

tc(8, 'Ô Trạng thái bị disable khi bản ghi đang Khoá', 'P1',
   'Bản ghi đang Khoá — mở qua modal Xem',
   '1. Bấm vào Mã của dòng đã Khoá',
   '',
   'Ô Trạng thái bị disable, không đổi được')

tc(8, 'Sửa bản ghi vừa bị người khác KHOÁ (2 tab)', 'P0',
   '2 tab cùng mở danh sách',
   '1. Tab A: khoá bản ghi X\n2. Tab B: bấm Sửa X (danh sách chưa refresh) rồi Lưu',
   '',
   'API trả 423\nToast "Bản ghi đang bị khoá, vui lòng mở khoá trước khi cập nhật."')

tc(8, 'Sửa bản ghi vừa bị người khác XOÁ (2 tab)', 'P0',
   '2 tab cùng mở danh sách',
   '1. Tab A: xoá bản ghi X\n2. Tab B: bấm Sửa X',
   '',
   'API trả 404\nToast "Dữ liệu đã thay đổi, vui lòng tải lại"\nModal đóng lại')

# ======================================================================
# IX. XEM CHI TIẾT
# ======================================================================
section('IX. XEM CHI TIẾT')

tc(9, 'Mở modal Xem bằng cách bấm Mã', 'P0',
   'Có ≥ 1 bản ghi',
   '1. Click vào giá trị cột Mã',
   '',
   'Modal tiêu đề "Xem chi tiết lĩnh vực Công ty kinh doanh"')

tc(9, 'Mọi ô ở chế độ chỉ đọc', 'P0',
   'Modal Xem đang mở',
   '1. Thử gõ vào ô Mã, ô Tên\n2. Thử đổi ô Trạng thái',
   '',
   'Cả 3 ô đều disable, không sửa được giá trị')

tc(9, 'Footer chỉ có nút Đóng', 'P0',
   'Modal Xem đang mở',
   '1. Quan sát footer',
   '',
   'Chỉ có nút "Đóng"\nKHÔNG có Lưu, KHÔNG có Lưu & Tiếp tục')

tc(9, 'Không có dấu * bắt buộc ở chế độ Xem', 'P2',
   'Modal Xem đang mở',
   '1. Quan sát nhãn ô Mã và ô Tên',
   '',
   'Không hiện dấu * đỏ (Required chỉ hiện khi không phải chế độ Xem)')

tc(9, 'Xem bản ghi đã Khoá', 'P1',
   'Có bản ghi status = 2',
   '1. Bấm vào Mã của bản ghi đã Khoá',
   '',
   'Modal mở bình thường, Trạng thái hiện "Khoá"')

# ======================================================================
# X. KHOÁ / MỞ KHOÁ
# ======================================================================
section('X. KHOÁ / MỞ KHOÁ')

tc(10, 'Khoá bản ghi thành công', 'P0',
    'Bản ghi Hoạt động, KHÔNG có Nhóm ngành nào trỏ tới',
    '1. Bấm biểu tượng Khoá\n2. Xác nhận trên modal',
    '',
    'Modal xác nhận hiện đúng câu: "Bạn có chắc muốn khoá lĩnh vực Công ty kinh doanh \'<tên>\'?"\nToast "Khoá thành công"\nBadge đổi sang "Khoá" (đỏ)')

tc(10, 'Nút Sửa/Xoá biến mất sau khi Khoá', 'P0',
    'Vừa khoá 1 bản ghi',
    '1. Quan sát cột Hành động của dòng đó',
    '',
    'Chỉ còn nút "Mở khoá"\nKhông còn Sửa, không còn Xoá')

tc(10, 'Mở khoá thành công', 'P0',
    'Bản ghi đang Khoá',
    '1. Bấm "Mở khoá"\n2. Xác nhận',
    '',
    'Tiêu đề modal: "Xác nhận mở khoá"\nToast "Mở khoá thành công"\nBadge về "Hoạt động"\nNút Sửa/Xoá hiện lại')

tc(10, 'Huỷ modal xác nhận Khoá', 'P1',
    'Bản ghi Hoạt động',
    '1. Bấm Khoá\n2. Bấm "Hủy"',
    '',
    'Modal đóng, trạng thái bản ghi KHÔNG đổi')

tc(10, 'Khoá ghi lại Người/Ngày cập nhật', 'P0',
    'Bản ghi Hoạt động',
    '1. Ghi lại giá trị Người/Ngày cập nhật\n2. Khoá bản ghi\n3. Quan sát lại',
    '',
    'Người cập nhật = user đang đăng nhập\nNgày cập nhật = thời điểm khoá\n(khoá cũng tính là 1 lần cập nhật)')

tc(10, 'ẨN nút Khoá khi còn Nhóm ngành ĐANG HOẠT ĐỘNG', 'P0',
    'Có Nhóm ngành đang Hoạt động trỏ tới lĩnh vực này',
    '1. Quan sát cột Hành động của lĩnh vực đó',
    '',
    'KHÔNG có nút Khoá (is_can_lock_update = false)\nVẫn có nút Sửa')

tc(10, 'CHO PHÉP khoá khi Nhóm ngành liên quan đều đã Khoá', 'P1',
    'Lĩnh vực có 1 Nhóm ngành, Nhóm ngành đó đang ở trạng thái Khoá',
    '1. Quan sát nút Khoá\n2. Bấm Khoá + xác nhận',
    '',
    'Nút Khoá HIỆN\nKhoá thành công')

tc(10, 'Gọi thẳng API lock khi đang bị chặn', 'P0',
    'Lĩnh vực có Nhóm ngành đang Hoạt động',
    '1. Dùng Postman gọi GET /assign/internal-business-scopes/{id}/lock',
    '',
    'HTTP 400\nMessage "Dữ liệu đang được sử dụng, vui lòng tải lại"\nTrạng thái KHÔNG đổi')

tc(10, 'Mở khoá luôn được phép', 'P1',
    'Lĩnh vực đang Khoá và có nhiều Nhóm ngành trỏ tới',
    '1. Bấm Mở khoá + xác nhận',
    '',
    'Mở khoá thành công (không có điều kiện chặn nào)')

# ======================================================================
# XI. XOÁ
# ======================================================================
section('XI. XOÁ')

tc(11, 'Xoá bản ghi thành công', 'P0',
    'Bản ghi Hoạt động, KHÔNG có Nhóm ngành nào trỏ tới',
    '1. Bấm biểu tượng Xoá\n2. Xác nhận',
    '',
    'Modal xác nhận: "Bạn có chắc muốn xóa lĩnh vực Công ty kinh doanh \'<tên>\'?"\nToast "Xoá thành công"\nBản ghi biến mất khỏi danh sách')

tc(11, 'Huỷ modal xác nhận Xoá', 'P1',
    'Bản ghi Hoạt động',
    '1. Bấm Xoá\n2. Bấm "Hủy"',
    '',
    'Modal đóng, bản ghi VẪN CÒN')

tc(11, 'ẨN nút Xoá khi có Nhóm ngành trỏ tới', 'P0',
    'Lĩnh vực có ≥ 1 Nhóm ngành (bất kể trạng thái) trỏ tới',
    '1. Quan sát cột Hành động',
    '',
    'KHÔNG có nút Xoá (is_can_delete = false)')

tc(11, 'ẨN nút Xoá khi bản ghi đang Khoá', 'P0',
    'Bản ghi status = 2',
    '1. Quan sát cột Hành động',
    '',
    'KHÔNG có nút Xoá')

tc(11, 'Gọi thẳng API delete khi đang được sử dụng', 'P0',
    'Lĩnh vực có Nhóm ngành trỏ tới',
    '1. Postman: DELETE /assign/internal-business-scopes/{id}',
    '',
    'HTTP 400 "Dữ liệu đang được sử dụng, vui lòng tải lại"\nBản ghi KHÔNG bị xoá')

tc(11, 'Gọi thẳng API delete bản ghi đang Khoá', 'P0',
    'Bản ghi status = 2, không có Nhóm ngành',
    '1. Postman: DELETE /assign/internal-business-scopes/{id}',
    '',
    'HTTP 423 "Bản ghi đang bị khoá, vui lòng mở khoá trước khi xoá."')

tc(11, 'Xoá bản ghi đã bị người khác xoá', 'P1',
    '2 tab cùng mở danh sách',
    '1. Tab A: xoá bản ghi X\n2. Tab B: xoá lại X',
    '',
    '404 → toast "Dữ liệu đã thay đổi, vui lòng tải lại"')

tc(11, 'Xoá xong danh sách nạp lại đúng', 'P1',
    'Đang ở trang 2, xoá dòng cuối cùng của trang',
    '1. Xoá bản ghi\n2. Quan sát',
    '',
    'Danh sách nạp lại, phân trang cập nhật đúng tổng số')

# ======================================================================
# XII. XUẤT EXCEL
# ======================================================================
section('XII. XUẤT EXCEL')

tc(12, 'Xuất Excel không lọc', 'P0',
    'Có ≥ 3 bản ghi, user có quyền',
    '1. Bấm "Xuất Excel"\n2. Mở file tải về',
    '',
    'Tải được file tên "danh_sach_linh_vuc_cong_ty_kinh_doanh.xls"\nFile mở được bằng Excel')

tc(12, 'Đúng 8 cột theo thứ tự', 'P0',
    'Vừa tải file export',
    '1. Mở file, quan sát dòng header',
    '',
    'STT · Mã · Tên lĩnh vực Công ty kinh doanh · Trạng thái · Người tạo · Ngày tạo · Người cập nhật · Ngày cập nhật')

tc(12, 'Tiêu đề trong file', 'P1',
    'Vừa tải file export',
    '1. Quan sát dòng tiêu đề phía trên bảng',
    '',
    'Ghi "Danh sách lĩnh vực Công ty kinh doanh"')

tc(12, 'Export theo ĐÚNG bộ lọc đang áp', 'P0',
    'Có 10 bản ghi, đang lọc Trạng thái = Khoá (còn 3 bản ghi)',
    '1. Bấm Xuất Excel\n2. Đếm số dòng trong file',
    'status = Khoá',
    'File chỉ chứa 3 bản ghi đang lọc, không phải toàn bộ 10')

tc(12, 'Export KHÔNG bị giới hạn phân trang', 'P0',
    'Có 30 bản ghi, đang xem trang 1 với page size = 10',
    '1. Bấm Xuất Excel\n2. Đếm số dòng',
    '',
    'File chứa đủ 30 bản ghi (không chỉ 10 dòng của trang hiện tại)')

tc(12, 'Export khi bộ lọc không ra dòng nào', 'P2',
    'Lọc bằng từ khoá không tồn tại',
    '1. Bấm Xuất Excel',
    'zzzzzz',
    'File tải về chỉ có phần header, không có dòng dữ liệu, không lỗi')

tc(12, 'Quyền Xem vẫn xuất được Excel', 'P0',
    'User CHỈ có quyền Xem',
    '1. Bấm Xuất Excel',
    '',
    'Tải file thành công (route export nhận cả 2 quyền)')

tc(12, 'Tên file đúng trên Safari', 'P1',
    'Dùng trình duyệt Safari (hoặc webview iOS)',
    '1. Bấm Xuất Excel\n2. Quan sát tên file trong thư mục Downloads',
    '',
    'File có tên + đuôi .xls đúng\nKHÔNG ra file UUID không đuôi (do tải trực tiếp qua ?token=, không dùng blob)')

tc(12, 'Định dạng ngày trong file export', 'P2',
    'Vừa tải file',
    '1. Quan sát cột Ngày tạo / Ngày cập nhật',
    '',
    'Dạng dd/mm/yyyy HH:mm, không có giây')

# ======================================================================
# XIII. IMPORT EXCEL
# ======================================================================
section('XIII. IMPORT EXCEL')

tc(13, 'Mở modal Import', 'P0',
    'User có quyền Quản lý',
    '1. Bấm "Import Excel"',
    '',
    'Mở modal tiêu đề "Import Lĩnh vực Công ty kinh doanh"\nPhụ đề: "Chỉ nhập Mã và Tên • Validate xong dòng hợp lệ sẽ bị khoá"')

tc(13, 'Tải file mẫu', 'P0',
    'Modal Import đang mở',
    '1. Bấm "Tải file mẫu"\n2. Mở file tải về',
    '',
    'Tải được Mau_import_LinhVucKinhDoanhNoiBo.xlsx\nDòng 1 header: STT · Mã lĩnh vực Công ty kinh doanh · Tên lĩnh vực Công ty kinh doanh\nDòng 2 là dòng hướng dẫn\nDòng 3-4 là 2 dòng mẫu')

tc(13, 'Nạp file mẫu nguyên bản', 'P0',
    'Đã tải file mẫu',
    '1. Upload lại chính file mẫu\n2. Bấm "Load lên bảng"',
    'File mẫu gốc',
    'Bảng nạp ĐÚNG 2 dòng (dòng hướng dẫn bị bỏ qua nhờ skipRows = 1)')

tc(13, 'File mẫu import được thật', 'P0',
    'Đã nạp 2 dòng mẫu, DB chưa có 2 mã đó',
    '1. Bấm Validate\n2. Bấm Import',
    '',
    '2 dòng đều hợp lệ\nImport thành công 2 bản ghi')

tc(13, 'Nhận file có header TÊN CŨ', 'P1',
    'File Excel dùng header "Mã lĩnh vực kinh doanh nội bộ" / "Tên lĩnh vực kinh doanh nội bộ"',
    '1. Upload file\n2. Bấm Load lên bảng',
    'File mẫu bản cũ',
    'Vẫn khớp cột đúng nhờ aliases, nạp được dữ liệu')

tc(13, 'Báo lỗi khi file thiếu cột bắt buộc', 'P1',
    'File Excel chỉ có cột Mã, thiếu cột Tên',
    '1. Upload + Load lên bảng',
    '',
    'Modal báo không khớp header / thiếu trường bắt buộc')

tc(13, 'Validate — dòng hợp lệ', 'P0',
    'Bảng đã nạp 1 dòng Mã + Tên hợp lệ, chưa trùng',
    '1. Bấm "Validate"',
    'LVKDNB.IM1 / Lĩnh vực import 1',
    'Trả validCount = 1, invalidCount = 0\nDòng được đánh dấu hợp lệ và bị khoá lại (không sửa được nữa)\nToast "Validate thành công"')

tc(13, 'Validate — Mã sai định dạng', 'P0',
    'Bảng đã nạp dòng có mã sai',
    '1. Bấm Validate',
    'ABC.123',
    'Dòng báo lỗi: "Mã phải có dạng LVKDNB. + tối đa 4 ký tự (A-Z, 0-9, _)"')

tc(13, 'Validate — Mã để trống', 'P0',
    'Bảng có dòng thiếu mã',
    '1. Bấm Validate',
    '(trống)',
    'Lỗi "Mã bắt buộc phải nhập"')

tc(13, 'Validate — Mã đã có trong DB', 'P0',
    'DB đã có LVKDNB.OTO',
    '1. Nạp dòng có mã LVKDNB.OTO\n2. Bấm Validate',
    'LVKDNB.OTO',
    'Lỗi "Mã đã tồn tại trong hệ thống"')

tc(13, 'Validate — Mã trùng giữa 2 dòng trong file', 'P0',
    'File có dòng 1 và dòng 2 cùng mã',
    '1. Nạp + Validate',
    'Dòng 1 và 2 đều LVKDNB.AA',
    'Dòng 2 báo "Mã bị trùng với dòng 1 trong file"\nDòng 1 vẫn hợp lệ')

tc(13, 'Validate — Tên để trống', 'P0',
    'Dòng có mã hợp lệ nhưng thiếu tên',
    '1. Bấm Validate',
    '',
    'Lỗi "Tên bắt buộc phải nhập"')

tc(13, 'Validate — Tên > 255 ký tự', 'P1',
    'Dòng có tên 256 ký tự',
    '1. Bấm Validate',
    'Chuỗi 256 ký tự',
    'Lỗi "Tên tối đa 255 ký tự"')

tc(13, 'Validate — Tên chứa dấu phẩy / hai chấm', 'P1',
    'Dòng có tên "Ô tô, xe"',
    '1. Bấm Validate',
    'Ô tô, xe',
    'Lỗi "Tên không được chứa dấu phẩy (,) và dấu hai chấm (:)"')

tc(13, 'Validate — Tên đã có trong DB', 'P0',
    'DB đã có bản ghi tên "Ô tô"',
    '1. Nạp dòng tên "ô tô" (khác hoa thường)\n2. Validate',
    'ô tô',
    'Lỗi "Tên đã tồn tại trong hệ thống" (so không phân biệt hoa/thường)')

tc(13, 'Validate — Tên trùng giữa 2 dòng trong file', 'P0',
    'File có 2 dòng cùng tên',
    '1. Nạp + Validate',
    '',
    'Dòng sau báo "Tên bị trùng với dòng N trong file"')

tc(13, 'Nút Import chỉ bật khi hết dòng lỗi', 'P0',
    'File có 1 dòng hợp lệ + 1 dòng lỗi, đã Validate',
    '1. Quan sát nút "Import"',
    '',
    'Nút Import bị khoá/không bấm được khi còn dòng lỗi')

tc(13, 'Sửa dòng lỗi rồi validate lại', 'P0',
    'Đang có 1 dòng lỗi sau khi Validate',
    '1. Sửa lại giá trị dòng lỗi ngay trên bảng\n2. Bấm Validate lại',
    '',
    'Dòng chuyển sang hợp lệ\nNút Import bật lên')

tc(13, 'Import toàn bộ hợp lệ', 'P0',
    'Đã validate 3 dòng đều hợp lệ',
    '1. Bấm Import',
    '',
    'Toast xanh "Import thành công 3 lĩnh vực Công ty kinh doanh"\nModal đóng\nDanh sách có thêm 3 bản ghi, đều ở trạng thái Hoạt động')

tc(13, 'Import một phần (HTTP 207)', 'P0',
    'Gọi thẳng API import với 3 dòng: 2 hợp lệ + 1 lỗi',
    '1. Postman: POST /import với payload hỗn hợp',
    '',
    'HTTP 207\nMessage "Import thành công 2/3 lĩnh vực Công ty kinh doanh. 1 lĩnh vực thất bại"\nFE hiện toast màu cảnh báo')

tc(13, 'Import không có dòng nào hợp lệ', 'P0',
    'Payload toàn dòng lỗi',
    '1. Postman: POST /import',
    '',
    'HTTP 400 "Không có dữ liệu hợp lệ để import"\nKHÔNG ghi bản ghi nào')

tc(13, 'Import quá 1000 dòng', 'P1',
    'Payload 1001 phần tử',
    '1. Postman: POST /import',
    '1001 dòng',
    'HTTP 422 "Import danh sách lĩnh vực Công ty kinh doanh tối đa 1000 phần tử"')

tc(13, 'Import mảng rỗng', 'P1',
    'Payload internal_business_scopes = []',
    '1. Postman: POST /import',
    '[]',
    'HTTP 422 "Import danh sách lĩnh vực Công ty kinh doanh phải có ít nhất 1 phần tử"')

tc(13, 'Import thiếu key internal_business_scopes', 'P2',
    'Payload rỗng {}',
    '1. Postman: POST /import',
    '{}',
    'HTTP 422 "Danh sách lĩnh vực Công ty kinh doanh là bắt buộc"')

tc(13, 'Import với quyền Xem', 'P0',
    'User CHỈ có quyền Xem',
    '1. Postman: POST /import với token của user đó',
    '',
    'HTTP 403, không ghi dữ liệu')

tc(13, 'Import ghi trong 1 transaction', 'P2',
    'Payload nhiều dòng, giả lập lỗi giữa chừng',
    '1. Gây lỗi ở dòng cuối\n2. Kiểm DB',
    '',
    'Không có bản ghi nào bị ghi dở dang (rollback)')

# ======================================================================
# XIV. API — KIỂM TRỰC TIẾP
# ======================================================================
section('XIV. API — KIỂM TRỰC TIẾP (Postman / curl)')

tc(14, 'GET danh sách trả 200 + cấu trúc đúng', 'P0',
    'Token của user có quyền',
    '1. GET /api/v1/assign/internal-business-scopes',
    '',
    'HTTP 200\nCó data[] và meta (current_page, per_page, total, last_page, from, to)')

tc(14, 'Resource trả đủ field', 'P0',
    'Có ≥ 1 bản ghi',
    '1. GET danh sách\n2. Kiểm 1 phần tử trong data',
    '',
    'Có đủ: id, code, name, status, status_text, created_by_name, created_at, updated_by_name, updated_at, is_can_edit, is_can_delete, is_can_lock_update, scopes_count')

tc(14, 'GET danh sách KHÔNG token', 'P0',
    'Không gửi Authorization header',
    '1. GET /api/v1/assign/internal-business-scopes',
    '',
    'HTTP 401')

tc(14, 'GET danh sách với user KHÔNG quyền', 'P0',
    'Token của e2e_nocatalog@test.local',
    '1. GET danh sách',
    '',
    'HTTP 403')

tc(14, 'GET chi tiết id không tồn tại', 'P1',
    'Token hợp lệ',
    '1. GET /assign/internal-business-scopes/99999999',
    '',
    'HTTP 404')

tc(14, 'POST tạo mới với user quyền Xem', 'P0',
    'Token quyền Xem',
    '1. POST / với payload hợp lệ',
    '',
    'HTTP 403, không tạo bản ghi')

tc(14, 'PUT sửa với user quyền Xem', 'P0',
    'Token quyền Xem',
    '1. PUT /{id}',
    '',
    'HTTP 403')

tc(14, 'DELETE với user quyền Xem', 'P0',
    'Token quyền Xem',
    '1. DELETE /{id}',
    '',
    'HTTP 403')

tc(14, 'lock/unlock với user quyền Xem', 'P0',
    'Token quyền Xem',
    '1. GET /{id}/lock\n2. GET /{id}/unlock',
    '',
    'Cả 2 đều HTTP 403')

tc(14, 'POST có id trong body → cập nhật', 'P1',
    'Bản ghi id = 12 đang Hoạt động',
    '1. POST / với body {id:12, code, name}',
    '',
    'HTTP 200, bản ghi 12 được cập nhật (không tạo bản ghi mới)')

tc(14, 'POST có id KHÔNG tồn tại', 'P1',
    'Token hợp lệ',
    '1. POST / với body {id: 99999999, ...}',
    '',
    'HTTP 404 "Dữ liệu đã thay đổi, vui lòng tải lại"')

tc(14, 'Lỗi 422 giữ nguyên (không bị nuốt thành 400)', 'P0',
    'Token hợp lệ',
    '1. POST / với code sai định dạng',
    'code = "ABC"',
    'HTTP 422 kèm object errors theo từng field\nKHÔNG phải 400')

tc(14, 'status ngoài 1/2', 'P1',
    'Token hợp lệ',
    '1. POST / với status = 9',
    'status = 9',
    'HTTP 422 "Trạng thái không hợp lệ"')

tc(14, 'GET /getAll chỉ trả bản ghi Hoạt động', 'P0',
    'DB có cả bản ghi Hoạt động và Khoá',
    '1. GET /assign/internal-business-scopes/getAll',
    '',
    'Chỉ trả bản ghi status = 1\nSắp theo name tăng dần\nMỗi phần tử có is_locked')

tc(14, 'GET /getAll với include_ids', 'P1',
    'Bản ghi id = 5 đang Khoá',
    '1. GET /getAll?include_ids[]=5',
    'include_ids = [5]',
    'Trả cả bản ghi id = 5 dù đang Khoá, kèm is_locked = true\n(để màn Sửa không mất giá trị đã chọn)')

tc(14, 'Phân trang per_page', 'P1',
    'Có > 5 bản ghi',
    '1. GET /?per_page=5&page=2',
    '',
    'Trả tối đa 5 phần tử, meta.current_page = 2')

# ======================================================================
# XV. LIÊN KẾT VỚI NHÓM NGÀNH
# ======================================================================
section('XV. LIÊN KẾT VỚI NHÓM NGÀNH')

tc(15, 'Ô chọn Lĩnh vực trong modal Nhóm ngành', 'P0',
    'Có ≥ 1 lĩnh vực đang Hoạt động',
    '1. Vào Danh mục › Nhóm ngành\n2. Bấm Tạo mới',
    '',
    'Có ô "Lĩnh vực Công ty kinh doanh" kèm dấu * bắt buộc\nPlaceholder "Chọn lĩnh vực Công ty kinh doanh"\nDanh sách chọn chỉ có lĩnh vực đang Hoạt động')

tc(15, 'Bắt buộc chọn Lĩnh vực khi tạo Nhóm ngành', 'P0',
    'Modal Tạo Nhóm ngành đang mở',
    '1. Nhập Mã + Tên, KHÔNG chọn Lĩnh vực\n2. Bấm Lưu',
    '',
    'Báo lỗi "Bắt buộc phải chọn" dưới ô Lĩnh vực')

tc(15, 'Cột Lĩnh vực trong danh sách Nhóm ngành', 'P0',
    'Có Nhóm ngành đã gắn lĩnh vực',
    '1. Vào danh sách Nhóm ngành',
    '',
    'Có cột "Lĩnh vực Công ty kinh doanh" hiện đúng tên lĩnh vực\nBản ghi chưa gắn hiện "—"')

tc(15, 'Bộ lọc Lĩnh vực ở màn Nhóm ngành', 'P1',
    'Danh sách Nhóm ngành có nhiều lĩnh vực khác nhau',
    '1. Mở lọc nâng cao\n2. Chọn 1 Lĩnh vực',
    '',
    'Chỉ hiện Nhóm ngành thuộc lĩnh vực đó')

tc(15, 'Không chọn được lĩnh vực đã Khoá', 'P0',
    'Có lĩnh vực đang Khoá',
    '1. Mở modal Tạo Nhóm ngành\n2. Mở ô chọn Lĩnh vực',
    '',
    'Lĩnh vực đang Khoá KHÔNG xuất hiện trong danh sách chọn')

tc(15, 'Sửa Nhóm ngành có lĩnh vực đã bị Khoá', 'P0',
    'Nhóm ngành X gắn lĩnh vực Y, sau đó Y bị Khoá',
    '1. Mở modal Sửa Nhóm ngành X',
    '',
    'Ô Lĩnh vực VẪN hiện giá trị Y (nhờ include_ids), có dấu hiệu đã khoá\nKhông bị mất giá trị đang chọn')

tc(15, 'Gán lĩnh vực đã khoá qua API', 'P1',
    'Lĩnh vực Y đang Khoá',
    '1. Postman: tạo Nhóm ngành với internal_business_scope_id = Y',
    '',
    'HTTP 422 "Lĩnh vực Công ty kinh doanh \'<tên>\' đã bị khoá, vui lòng chọn lĩnh vực khác"')

tc(15, 'Gán lĩnh vực không tồn tại', 'P1',
    'Token hợp lệ',
    '1. Postman: tạo Nhóm ngành với internal_business_scope_id = 99999999',
    '',
    'HTTP 422 "Lĩnh vực Công ty kinh doanh không tồn tại"')

tc(15, 'Import Nhóm ngành theo MÃ lĩnh vực', 'P0',
    'File mẫu Nhóm ngành có cột "Mã lĩnh vực Công ty kinh doanh *"',
    '1. Tải file mẫu Nhóm ngành\n2. Điền mã LVKDNB.OTO\n3. Validate + Import',
    'LVKDNB.OTO',
    'Import thành công, Nhóm ngành gắn đúng lĩnh vực')

tc(15, 'Import Nhóm ngành với mã lĩnh vực không tồn tại', 'P0',
    'Đang import Nhóm ngành',
    '1. Điền mã LVKDNB.XXX (không có trong DB)\n2. Validate',
    'LVKDNB.XXX',
    'Dòng báo lỗi "Lĩnh vực Công ty kinh doanh \'LVKDNB.XXX\' không tồn tại"')

tc(15, 'Import Nhóm ngành với mã lĩnh vực đã Khoá', 'P1',
    'Lĩnh vực LVKDNB.OLD đang Khoá',
    '1. Điền mã đó vào file import Nhóm ngành\n2. Validate',
    'LVKDNB.OLD',
    'Dòng báo lỗi "Lĩnh vực Công ty kinh doanh \'LVKDNB.OLD\' đã bị khoá"')

tc(15, 'scopes_count phản ánh đúng số Nhóm ngành', 'P1',
    'Lĩnh vực có 3 Nhóm ngành',
    '1. GET chi tiết lĩnh vực đó',
    '',
    'scopes_count = 3\nis_can_delete = false')

# ======================================================================
# XVI. GIAO DIỆN & E2E TỔNG HỢP
# ======================================================================
section('XVI. GIAO DIỆN & E2E TỔNG HỢP')

tc(16, 'Không còn chữ "nội bộ" trên giao diện', 'P0',
    'User có quyền Quản lý',
    '1. Rà toàn bộ màn: menu, tiêu đề bảng, nhãn cột, bộ lọc, modal, toast, modal Import',
    '',
    'Mọi chỗ đều ghi "Lĩnh vực Công ty kinh doanh"\nKHÔNG còn chuỗi "lĩnh vực kinh doanh nội bộ" ở bất kỳ đâu người dùng nhìn thấy')

tc(16, 'Tên quyền trong màn Phân quyền', 'P0',
    'Đăng nhập tài khoản admin',
    '1. Vào màn Phân quyền\n2. Tìm nhóm "Danh mục"',
    '',
    'Hiện 2 quyền: "Quản lý danh mục lĩnh vực Công ty kinh doanh" và "Xem danh mục lĩnh vực Công ty kinh doanh"')

tc(16, 'Đỏ chỉ dùng cho lỗi validate', 'P2',
    'Màn danh sách + modal',
    '1. Rà màu sắc toàn màn',
    '',
    'Màu đỏ chỉ xuất hiện ở: dấu * bắt buộc, thông báo lỗi validate, badge Khoá, nút Xoá\nKhông dùng đỏ trang trí')

tc(16, 'Hiển thị trên màn hình nhỏ', 'P2',
    'Thu nhỏ cửa sổ còn 1366px',
    '1. Quan sát bảng',
    '',
    'Cột STT và Mã dính (sticky) khi cuộn ngang\nBảng cuộn ngang được, không vỡ layout')

tc(16, 'E2E: Tạo → Sửa → Khoá → Mở khoá → Xoá', 'P0',
    'User có quyền Quản lý',
    '1. Tạo LVKDNB.E2E "Lĩnh vực E2E"\n2. Sửa tên thành "Lĩnh vực E2E sửa"\n3. Khoá\n4. Mở khoá\n5. Xoá',
    'LVKDNB.E2E',
    'Cả 5 bước đều thành công\nSau bước 2: Người/Ngày cập nhật đổi\nSau bước 3: chỉ còn nút Mở khoá\nSau bước 5: bản ghi biến mất')

tc(16, 'E2E: Tạo lĩnh vực → gắn Nhóm ngành → thử xoá/khoá', 'P0',
    'User có quyền Quản lý cả 2 danh mục',
    '1. Tạo lĩnh vực mới\n2. Tạo Nhóm ngành gắn lĩnh vực đó\n3. Quay lại màn Lĩnh vực, quan sát cột Hành động\n4. Khoá Nhóm ngành\n5. Quan sát lại',
    '',
    'Bước 3: mất cả nút Xoá và nút Khoá\nBước 5: nút Khoá hiện lại (chỉ còn Nhóm ngành đã khoá), nút Xoá vẫn ẩn')

tc(16, 'E2E: Import → Xuất Excel đối chiếu', 'P1',
    'DB sạch',
    '1. Import 3 bản ghi từ file mẫu\n2. Xuất Excel\n3. Đối chiếu',
    '',
    'File export có đủ 3 bản ghi vừa import, đúng Mã/Tên/Trạng thái Hoạt động')

tc(16, 'E2E: 2 user song song', 'P1',
    '2 trình duyệt, 2 tài khoản đều có quyền Quản lý',
    '1. User A tạo bản ghi X\n2. User B refresh, thấy X\n3. User B sửa X\n4. User A refresh',
    '',
    'User A thấy Người cập nhật = User B\nDữ liệu nhất quán giữa 2 phiên')

tc(16, 'Dọn dữ liệu test sau khi chạy', 'P1',
    'Đã chạy xong bộ testcase',
    '1. Xoá toàn bộ bản ghi có mã bắt đầu LVKDNB.E2E / LVKDNB.TEST / LVKDNB.IM',
    '',
    'DB trở lại trạng thái trước khi test')

# === DATA VALIDATION ===
dv = DataValidation(type='list', formula1='"Passed,Failed,Pending,Not Executed"', allow_blank=True)
dv.error = 'Chỉ chọn: Passed, Failed, Pending, Not Executed'
dv.errorTitle = 'Giá trị không hợp lệ'
ws.add_data_validation(dv)
dv.add(f'L8:L{current_row}')

ws.freeze_panes = 'A7'

output_path = 'docs/srs/linh-vuc-cong-ty-kinh-doanh-testcases.xlsx'
wb.save(output_path)
print(f'Saved: {output_path}')

# === XUẤT THÊM BẢN MARKDOWN để review trong git ===
md = ['# Testcase — Danh mục Lĩnh vực Công ty kinh doanh', '',
      '| Thông tin | Chi tiết |', '|---|---|',
      '| Module | Dự án & Giao việc (Assign) |',
      '| Màn hình | Danh mục › Lĩnh vực Công ty kinh doanh (`/assign/internal-business-scopes`) |',
      '| SRS | `docs/srs/linh-vuc-cong-ty-kinh-doanh.md` |',
      '| File Excel | `docs/srs/linh-vuc-cong-ty-kinh-doanh-testcases.xlsx` |',
      '| Sinh bởi | `docs/srs/generate_testcase_lvctkd.py` |', '']

_total = 0
for r in range(7, ws.max_row + 1):
    module = ws.cell(row=r, column=1).value
    tcid = ws.cell(row=r, column=3).value
    if module is None and tcid:
        md += ['', f'## {tcid}', '',
               '| TC ID | Chức năng | Ưu tiên | Tiền điều kiện | Bước thực hiện | Test Data | Kết quả mong đợi |',
               '|---|---|:---:|---|---|---|---|']
        continue
    if not tcid:
        continue
    _total += 1
    cells = [ws.cell(row=r, column=c).value or '' for c in (3, 4, 5, 6, 7, 8, 10)]
    cells = [str(x).replace('|', '\\|').replace('\n', '<br>') for x in cells]
    md.append('| ' + ' | '.join(cells) + ' |')

md.insert(9, f'| Tổng số testcase | **{_total}** |')
md.insert(10, '')
with open('docs/srs/linh-vuc-cong-ty-kinh-doanh-testcases.md', 'w', encoding='utf-8') as f:
    f.write('\n'.join(md) + '\n')
print(f'Saved: docs/srs/linh-vuc-cong-ty-kinh-doanh-testcases.md ({_total} testcases)')
