# -*- coding: utf-8 -*-
"""Doi chieu mot file SRS voi BAN MAU `assets/SRS_MAU.docx` — chay duoc doc lap.

    python3 srs_selfcheck.py "<duong dan>/SRS - <Ten man>.docx"

Sinh ra sau su co 2026-09-03: mot phien lam viec dai da sinh SRS bang API CU (so do use case
ve phang, Phan 4 dang doan BR, muc Layout ghi URL) va tester tra ve 3 lan. Nguyen nhan: nguoi
viet doc SKILL.md o dau phien — truoc luc skill duoc cap nhat — roi khong doc lai, va chi doc
CHU trong file mau chu khong mo ANH so do ra xem.

Bo kiem nay DOC THANG ban mau moi lan chay, nen skill doi thi kiem tra doi theo, khong phu
thuoc vao viec nguoi viet co nho hay khong. `SrsDoc.save()` goi ham `check()` nay tu dong.

Tra ve 0 neu dat, 1 neu co loi. Loi in ra kem CACH SUA.
"""
import os
import re
import sys
import zipfile

from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

ASSETS = os.path.dirname(os.path.abspath(__file__))
MAU = os.path.join(ASSETS, 'SRS_MAU.docx')

W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def _txt(el):
    """Doc chu ke ca khi bi boc trong <w:sdt> (file tai tu Google Docs)."""
    return ''.join(t.text or '' for t in el.iter(W + 't')).strip()


def _blocks(doc):
    for c in doc.element.body.iterchildren():
        if c.tag == qn('w:p'):
            yield Paragraph(c, doc)
        elif c.tag == qn('w:tbl'):
            yield Table(c, doc)


def _uml_kinds(path):
    """Doc dau "srs-uml" trong metadata cua moi anh PNG nhung trong file .docx."""
    import io as _io

    from PIL import Image

    kinds = set()
    with zipfile.ZipFile(path) as z:
        for n in z.namelist():
            if not n.startswith('word/media/') or not n.lower().endswith('.png'):
                continue
            try:
                with Image.open(_io.BytesIO(z.read(n))) as im:
                    k = (im.info or {}).get('srs-uml')
            except Exception:                      # anh chup man hinh, khong phai so do
                continue
            if k:
                kinds.add(k)
    return kinds


def _profile(path):
    """Rut cac dac trung hinh thuc cua mot file SRS."""
    d = Document(path)
    paras = [p.text.strip() for p in d.paragraphs]
    heads = {}                      # tieu de bang -> so lan xuat hien
    ui_cols = {}                    # bang "Mo ta chi tiet giao dien" -> so cot
    rule_tbl = None
    for b in _blocks(d):
        if not isinstance(b, Table):
            continue
        h = [_txt(c._tc) for c in b.rows[0].cells]
        key = ' | '.join(h)
        heads[key] = heads.get(key, 0) + 1
        if h[:2] == ['STT', 'Tên đối tượng']:
            ui_cols[len(b.columns)] = ui_cols.get(len(b.columns), 0) + 1
        if h[:2] == ['STT', 'Mã quy tắc']:
            rule_tbl = h
    with zipfile.ZipFile(path) as z:
        media = [n for n in z.namelist() if n.startswith('word/media/')]
    return {
        'paras': paras,
        'heads': heads,
        'ui_cols': ui_cols,
        'rule_tbl': rule_tbl,
        'menu': sum(1 for t in paras if t.startswith('Menu: ')),
        'rule_ref': sum(1 for t in paras if t.startswith('Quy tắc chung: ')),
        'rule_lead': sum(1 for t in paras if t.startswith('Quy tắc áp dụng: ')),
        'links': len([r for r in Document(path).part.rels.values()
                      if r.reltype.endswith('/hyperlink')]),
        'media': len(media),
        'fn': sum(1 for p in Document(path).paragraphs
                  if p.style.name == 'Heading 3' and re.match(r'^2\.\d+\s', p.text.strip())),
    }


def check(path, verbose=True):
    """So file SRS voi ban mau. Tra ve danh sach loi (rong = dat)."""
    me = _profile(path)
    mau = _profile(MAU)
    errs = []

    # --- 4 diem cua form 2026-08-28 ---
    if me['menu'] < me['fn']:
        errs.append('Mục Layout: chỉ có %d dòng "Menu: " cho %d chức năng. '
                    'Dùng d.layout(menu=MENU + " => Tạo mới", shot=...)' % (me['menu'], me['fn']))
    if any('URL đầy đủ' in t for t in me['paras']):
        errs.append('Còn dòng "URL đầy đủ" của form 2026-08-17 — form hiện hành ghi đường dẫn '
                    'MENU, bỏ hẳn URL.')
    if me['rule_ref'] < me['fn']:
        errs.append('Chỉ %d/%d mục Giới thiệu có đoạn "Quy tắc chung: ". '
                    'Dùng d.rule_ref("- <nhóm quy tắc>. ...", anchor="list")'
                    % (me['rule_ref'], me['fn']))
    if me['links'] < me['rule_ref']:
        errs.append('Đoạn "Quy tắc chung" phải là HYPERLINK thật sang SRS quy tắc chung '
                    '(%d đoạn nhưng chỉ %d liên kết) — d.rule_ref() tự chèn, đừng viết d.p() tay.'
                    % (me['rule_ref'], me['links']))
    if me['rule_lead'] != 1:
        errs.append('Phần 4 phải mở đầu bằng ĐÚNG 1 câu "Quy tắc áp dụng: " '
                    '(d.rule_ref(..., head="Quy tắc áp dụng", lead=...)); đang có %d.'
                    % me['rule_lead'])
    if me['rule_tbl'] != mau['rule_tbl']:
        errs.append('Phần 4 phải là BẢNG 5 cột %s — dùng d.rule_table([...]), không viết dạng '
                    '"BR-0N — <tên>" + gạch đầu dòng.' % mau['rule_tbl'])

    # --- so do use case tong quan: phai co phan cap ---
    # Moi so do do srs_uml_render sinh ra deu mang dau "srs-uml" trong metadata PNG:
    #   overview-flat      = ban ve phang cua form cu (draw_overview)
    #   overview-hierarchy = ban co phan cap cua form 2026-08-28 (draw_overview2)
    # Doc dau nay chac chan hon doan theo ty le anh — ban phang cung co the rat rong.
    kinds = _uml_kinds(path)
    if 'overview-flat' in kinds:
        errs.append('Sơ đồ tổng quan đang vẽ PHẲNG (mọi use case nối thẳng tới actor) — form '
                    'hiện hành vẽ có phân cấp: chỉ MÀN HÌNH thật nối actor, thao tác trên màn đó '
                    'nối «include»/«extend» vào màn cha. Dùng '
                    'd.overview_figure2(actors, mains, subs, caption).')
    elif 'overview-hierarchy' not in kinds:
        errs.append('Không tìm thấy sơ đồ Use Case tổng quan sinh bằng d.overview_figure2() — '
                    'ảnh chèn tay hoặc dùng hàm cũ đều không đạt.')

    # --- bo cot bang giao dien: chi nhan bo cot co trong ban mau ---
    lech = sorted(set(me['ui_cols']) - set(mau['ui_cols']))
    if lech:
        errs.append('Bảng "Mô tả chi tiết giao diện" có bộ cột %s không có trong bản mẫu (mẫu '
                    'dùng %s cột). Chức năng chỉ đọc = 7 cột (required=False), có nhập liệu = 8 '
                    'cột, hộp xác nhận = 6 cột (required=False, scope=False).'
                    % (lech, sorted(mau['ui_cols'])))

    # --- muc da bo cua form cu ---
    for s in ('Tổng quan', 'Mini-Spec', 'Tiêu chí nghiệm thu', 'Ngoài phạm vi',
              'Chức năng liên quan', 'Route (FE)'):
        if any(s in t for t in me['paras']):
            errs.append('Còn mục đã bỏ của form cũ: "%s".' % s)

    if verbose:
        name = os.path.basename(path)
        print('--- Đối chiếu với bản mẫu %s' % os.path.basename(MAU))
        print('    %s: %d chức năng | Menu %d | Quy tắc chung %d | hyperlink %d | '
              'bảng giao diện %s | Phần 4 %s'
              % (name, me['fn'], me['menu'], me['rule_ref'], me['links'],
                 dict(sorted(me['ui_cols'].items())),
                 'bảng 5 cột' if me['rule_tbl'] == mau['rule_tbl'] else 'SAI'))
        if errs:
            print('!!! %d điểm chưa đúng bản mẫu:' % len(errs))
            for e in errs:
                print('  - %s' % e)
        else:
            print('    => khớp bản mẫu.')
    return errs


if __name__ == '__main__':
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    if len(sys.argv) < 2:
        print(__doc__)
        sys.exit(2)
    bad = 0
    for f in sys.argv[1:]:
        bad += 1 if check(f) else 0
    sys.exit(1 if bad else 0)
